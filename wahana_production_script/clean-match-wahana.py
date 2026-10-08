import os
import re

import numpy as np
import pandas as pd

# ─────────────────────────────────────────────────────────────────────────────
# CLEANING SUSPEND
# ─────────────────────────────────────────────────────────────────────────────

# ─────────────────────────────────────────────────────────────────────────────
# KONFIGURASI
# ─────────────────────────────────────────────────────────────────────────────

INPUT_FILE  = os.path.join("input", "3b. Database Suspense 150826.xlsx")
OUTPUT_FILE = os.path.join("output", "rev_wahana_output_suspend.xlsx")
SHEET_NAME  = "Sheet1"   # sheet yang berisi data mentah (row 1-2 kosong/judul, header di row 3)

# Kolom Cedant pada data Suspend
COMPNAME_FILTER_COL   = "CEDANT NAME"
COMPNAME_FILTER_VALUE = "PT.ASURANSI WAHANA TATA"


# Business rule: kalau hasil breakdown (insured/polis/slip) > 5 bagian,
# tidak usah dipecah per kolom -> digabung lagi jadi satu kolom.
MAX_BREAKDOWN = 5

# Regex untuk sub-removal dalam nama insured
_INSURED_TAIL_RE = re.compile(
    r"""
    \bAS\b\s+(?:THE\s+)?(?:PRINCIPAL|OFF-TAKER|MAINTENANCE|CONTRACTOR).*
  | \bBEING\s+(?:THE\s+)?(?:PRINCIPAL|OFF-TAKER).*
  | \bAND\s+ALL\s+SUBSIDIARI.*
  | \bINCLUDING\s+ALL\s+SUBSIDIARI.*
  | \bINCLUDING\s+ANY\s+SUBSIDIAR.*
  | \bCOMPRISING\s+OF.*
  | \bINSTALLMENT\b.*
  | \bRELATED\s+COMPANY\b.*
  | \bPURCHASED\s+OR\s+OTHERWISE\b.*
  | \bWPC\s+DATE\s*:.*
    """,
    re.IGNORECASE | re.VERBOSE,
)

_INSURED_JUNK_WORDS = frozenset({
    "SHANGHAI", "PR OF CHINA", "CHINA", "INDONESIA", "JAKARTA",
    "OFFICERS", "EMPLOYEES", "ALL OTHER CONTRACTORS",
    "SUB-CONTRACTORS", "SUB CONTRACTORS",
    "COMPANIES", "AFFILIATED", "AFFILIATES",
    "CORPORATIONS AND INCLUDING PARTNERSHIP",
    "JOINT VENTURES AND AGREEMENT OR BY LAW",
    "AS THEIR RESPECTIVE INTEREST MAY APPEAR",
    "AS THEIR RESPECTIVE INTERESTS MAY APPEAR",
    "SUBSIDIARY", "SUBSIDIARIES", "ANY SUBSIDIARY COMPANY",
    "RELATED COMPANY",
    "FOR THEIRS RESPECTIVE RIGHTS AND INTEREST",
    "MIGRASI AS400", "THE PRINCIPAL", "PRINCIPAL", "OWNER",
})

# Gelar / sapaan yang ikut dihapus dari nama insured (Bapak, Ibu, Ny, Mr, Mrs, Ms, dll)
_INSURED_TITLE_RE = re.compile(
    r"\b(?:BAPAK|IBU|BPK|NY\.?|SDR\.?|SDRI\.?|MR\.?|MRS\.?|MS\.?)\b\s*",
    re.IGNORECASE,
)


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN FUNCTIONS
# ─────────────────────────────────────────────────────────────────────────────

def _cap_breakdown(parts: list, sep: str = "; ") -> list:
    """
    Business rule: jika jumlah bagian hasil breakdown > MAX_BREAKDOWN,
    jangan dipecah per kolom -> gabungkan lagi jadi satu nilai.
    """
    parts = [p for p in parts if p]
    if len(parts) > MAX_BREAKDOWN:
        return [sep.join(parts)]
    return parts


def clean_polis(val) -> list:
    """Hapus suffix numerik di akhir nomor polis (misal: -01, -02/03)."""
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val or val == "-":
        return []

    # Hapus koma (mis. "000.6005.201.2025.002555.00," -> "000.6005.201.2025.002555.00")
    val = val.replace(",", "")

    # Kalau ada lebih dari satu nomor polis dipisah "+", proses masing-masing
    raw_parts = [p.strip() for p in val.split("+") if p.strip()]
    if not raw_parts:
        return []

    cleaned_parts = []
    for p in raw_parts:
        cur = p
        while True:
            stripped = re.sub(r"-\d+(?:/\d+)?$", "", cur)
            if stripped == cur:
                break
            cur = stripped
        if cur:
            cur = cur.replace(".", "")  # hapus titik dari nomor polis
            cleaned_parts.append(cur)

    return _cap_breakdown(cleaned_parts)

# ─────────────────────────────────────────────────────────────────────────────
# CLEAN Certificate
# ─────────────────────────────────────────────────────────────────────────────

def clean_certificate(polis_ori):
    """
    Mengambil Certificate dari POLIS ORI.

    RULE:
    - Polis harus mempunyai "-"
    - Satu suffix angka:
        ...-000156 -> 000156
        ...-00073  -> 000073
        ...-0045   -> 000045
    - Beberapa suffix:
        ...-110-114-109 -> 000110 SD 000109
        ...-00383-00045-00031 -> 000383 SD 000031
    - Suffix > 6 digit -> blank
    - Suffix bukan angka -> blank
    - Polis tanpa "-" -> blank
    """

    if pd.isna(polis_ori):
        return ""

    polis = str(polis_ori).strip()

    if not polis:
        return ""

    # ============================================================
    # HARUS ADA DASH
    # ============================================================

    if "-" not in polis:
        return ""

    # Ambil bagian setelah dash pertama
    parts = polis.split("-")

    if len(parts) < 2:
        return ""

    suffixes = [p.strip() for p in parts[1:]]

    # ============================================================
    # SEMUA SUFFIX HARUS ANGKA
    # ============================================================

    if not all(re.fullmatch(r"\d+", x) for x in suffixes):
        return ""

    # ============================================================
    # SINGLE Certificate
    # ============================================================

    if len(suffixes) == 1:

        cert = suffixes[0]

        # Maksimal 6 digit
        if len(cert) > 6:
            return ""

        return cert.zfill(6)

    # ============================================================
    # RANGE Certificate
    # ============================================================

    first = suffixes[0]
    last = suffixes[-1]

    # Maksimal 6 digit
    if len(first) > 6 or len(last) > 6:
        return ""

    return f"{first.zfill(6)} SD {last.zfill(6)}"

def clean_slip(val) -> list:
    """Kembalikan nilai slip (tanpa titik), pecah kalau ada lebih dari satu (+)."""
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val or val == "-":
        return []

    raw_parts = [p.strip().replace(".", "") for p in val.split("+") if p.strip()]
    return _cap_breakdown(raw_parts)


def _remove_polis_slip_from_text(text: str, polis_ori, slip_ori) -> str:
    """Hapus nomor polis & slip dari teks insured secara agresif."""
    if pd.notna(polis_ori):
        for token in [str(polis_ori).strip()] + clean_polis(polis_ori):
            if token and token != "-":
                text = text.replace(token, "")

    if pd.notna(slip_ori):
        for token in [str(slip_ori).strip()] + clean_slip(slip_ori):
            if token and token != "-":
                text = text.replace(token, "")

    return text


def _normalize_insured_part(p: str) -> str:
    """Terapkan sub-removal dan normalisasi pada satu bagian nama insured."""
    p = _INSURED_TAIL_RE.sub("", p)
    p = _INSURED_TITLE_RE.sub("", p)
    p = re.sub(r"\bKB\b",    "", p, flags=re.IGNORECASE)
    p = re.sub(r"\bA\.?W\.?\b", "", p, flags=re.IGNORECASE)
    p = re.sub(r"\(\s*\)",   "", p)
    # Trim leading/trailing AND/OR
    p = re.sub(r"^(?:AND|OR)\b\s*", "", p, flags=re.IGNORECASE)
    p = re.sub(r"\s*\b(?:AND|OR)$", "", p, flags=re.IGNORECASE)
    # Trim leading/trailing non-alphanumeric
    p = re.sub(r"^[^a-zA-Z0-9(]+", "", p)
    p = re.sub(r"[^a-zA-Z0-9)]+$", "", p)
    return p.strip()


def _is_valid_insured_part(p: str) -> bool:
    """Return True jika bagian insured layak dipertahankan."""
    if len(p) <= 2:
        return False
    up = p.upper()
    if re.match(r"^[\d\/\-\.]+$", up):
        return False
    if re.search(r"\b(?:NO\.\s*\d+|BUILDING|FLOOR|ROOM|ROAD|STREET|TOWER|KAV\.?|BLOK)\b", up):
        return False
    if re.search(r"\b(?:PLTGU|PLTMH|PLTU|POWER PLANT|COMBINED CYCLE|MW|HYDRO ELECTRIC)\b", up):
        return False
    return up not in _INSURED_JUNK_WORDS


def clean_insured(val, polis_ori, slip_ori) -> list:
    """Bersihkan nama insured dengan menghapus nomor polis/slip, gelar, dan junk words."""
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val:
        return []

    val = _remove_polis_slip_from_text(val, polis_ori, slip_ori)

    # Hapus tail info (mis. "AS PRINCIPAL...", "WPC Date : ...") SEBELUM split,
    # karena split_pattern (":" dll) bisa memotong pola ini duluan.
    val = _INSURED_TAIL_RE.sub("", val)

    # Normalisasi sebelum split
    # Hapus semua isi dalam kurung, mis. "(FACULTATIVE)", "(12.34/56)", dll
    val = re.sub(r"\([^)]*\)", "", val)
    val = re.sub(r"\b(?:AND|AN|OR)\s*/\s*(?:AND|OR)\b", ",", val, flags=re.IGNORECASE)
    val = re.sub(r"\bAND\s+OR\b", ",", val, flags=re.IGNORECASE)
    val = re.sub(r"\bCO\.,?\s*LTD\.?\b", ",", val, flags=re.IGNORECASE)
    for pattern in [r"\bTBK\.?\b", r"\(PERSERO\)", r"\bPERSERO\b", r"\bLTD\.?\b", r"\(FCI\.\s*I\)"]:
        val = re.sub(pattern, "", val, flags=re.IGNORECASE)

    # Delimiter pemisah insured: , / QQ and/or & – + (sesuai catatan rules)
    split_pattern = (
        r"\bQQ\b|/|,|&|\+"
        r"|-(?!\s*(?:19|20)\d{2}\b)"
        r"|\d+\.|\bPT\.?\b|\bCV\.?\b|:|;"
    )
    parts = re.split(split_pattern, val, flags=re.IGNORECASE)

    cleaned = []
    for p in parts:
        p = _normalize_insured_part(p)
        
        if _is_valid_insured_part(p):
            cleaned.append(p.upper())

    return _cap_breakdown(cleaned)


# ─────────────────────────────────────────────────────────────────────────────
# COLUMN EXPANSION HELPER
# ─────────────────────────────────────────────────────────────────────────────

def _expand_clean_columns(
    df: pd.DataFrame,
    all_lists: list,
    prefix: str,
    max_cols: int,
) -> list:
    """
    Tambahkan kolom clean_{prefix}_1 .. N langsung setelah kolom ori.
    max_cols otomatis dibatasi MAX_BREAKDOWN karena _cap_breakdown()
    sudah menggabungkan hasil > MAX_BREAKDOWN jadi 1 kolom.
    Returns daftar nama kolom baru yang ditambahkan.
    """
    added = []
    for i in range(1, max_cols + 1):
        col_name = f"clean {prefix} {i}"
        df[col_name] = [lst[i - 1] if i - 1 < len(lst) else None for lst in all_lists]
        added.append(col_name)
    return added


# ─────────────────────────────────────────────────────────────────────────────
# PROCESS DATA
# ─────────────────────────────────────────────────────────────────────────────

def process_data(input_file: str, output_file: str) -> None:
    print(f"[1/5] Membaca data dari: {input_file} (sheet: {SHEET_NAME}) ...")
    df = pd.read_excel(input_file, sheet_name=SHEET_NAME, header=2)
    print(f"      Total baris keseluruhan: {len(df):,}")

        # ========================================================
    # FILTER STATUS - HANYA SUSPENSE
    # ADJUSTED DIBUANG PERMANEN DARI OUTPUT
    # ========================================================

    if "STATUS" not in df.columns:
        print("\n[ERROR] Kolom 'STATUS' tidak ditemukan!")
        print(f"       Kolom tersedia: {list(df.columns)}")
        return

    total_sebelum_status = len(df)

    df = df[
        df["STATUS"]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
        .eq("SUSPENSE")
    ].copy()

    total_suspense = len(df)
    total_adjusted = total_sebelum_status - total_suspense

    print(f"      Data SUSPENSE : {total_suspense:,} baris")
    print(f"      Data ADJUSTED : {total_adjusted:,} baris dibuang")

    if COMPNAME_FILTER_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{COMPNAME_FILTER_COL}' tidak ditemukan!")
        print(f"        Kolom tersedia: {list(df.columns)}")
        return

    df[COMPNAME_FILTER_COL] = df[COMPNAME_FILTER_COL].astype(str).str.strip()
    is_wahana = df[COMPNAME_FILTER_COL] == COMPNAME_FILTER_VALUE
    print(f"[2/5] Filter comp_name '{COMPNAME_FILTER_VALUE}': {is_wahana.sum():,} baris WAHANA dari total {len(df):,} baris.")

    # Output cleaning hanya untuk Cedant WAHANA.
    df = df.loc[is_wahana].copy()
    is_wahana = pd.Series(True, index=df.index)
    print(f"      Output cleaning Suspend dibatasi ke WAHANA: {len(df):,} baris.")

    # Rename kolom asli -> _ori
    rename_map = {"INSURED": "insured_ori", "POLIS": "polis_ori", "SLIP NO": "slip_ori"}
    df.rename(columns={k: v for k, v in rename_map.items() if k in df.columns}, inplace=True)

    print("[3/5] Menjalankan proses cleaning hanya untuk baris WAHANA (baris lain dibiarkan apa adanya) ...")

    all_clean_polis = []
    all_clean_slip = []
    all_clean_ins = []
    all_certificates = []
    max_polis = max_slip = max_ins = 1

    for idx, (_, row) in enumerate(df.iterrows()):
        p_ori = row.get("polis_ori", "")
        s_ori = row.get("slip_ori", "")
        i_ori = row.get("insured_ori", "")

        # Hanya cleaning baris WAHANA, non-WAHANA dikosongkan (dibiarkan apa adanya)
        if is_wahana.iloc[idx]:
            c_polis = clean_polis(p_ori)
            c_slip  = clean_slip(s_ori)
            c_ins   = clean_insured(i_ori, p_ori, s_ori)
            c_certificate = clean_certificate(p_ori)
        else:
            c_polis, c_slip, c_ins = [], [], []
            c_certificate = ""

        max_polis = max(max_polis, len(c_polis))
        max_slip  = max(max_slip,  len(c_slip))
        max_ins   = max(max_ins,   len(c_ins))

        all_clean_polis.append(c_polis)
        all_clean_slip.append(c_slip)
        all_clean_ins.append(c_ins)

        all_certificates.append(c_certificate)

    # Safety net: breakdown tidak boleh lebih dari MAX_BREAKDOWN kolom
    max_polis = min(max_polis, MAX_BREAKDOWN)
    max_slip  = min(max_slip,  MAX_BREAKDOWN)
    max_ins   = min(max_ins,   MAX_BREAKDOWN)

    print("[4/5] Menyusun kolom output ...")

    # Sisipkan kolom clean langsung setelah kolom _ori-nya
    new_columns = []

    for col in df.columns:

        new_columns.append(col)

        if col == "polis_ori":
            df["Certificate"] = all_certificates

            clean_polis_columns = _expand_clean_columns(
                df,
                all_clean_polis,
                "polis",
                max_polis
            )

            # Urutan:
            # polis_ori
            # clean polis 1
            # Certificate
            # clean polis 2
            # clean polis 3
            # dst.

            if clean_polis_columns:
                new_columns.append(clean_polis_columns[0])
                new_columns.append("Certificate")
                new_columns += clean_polis_columns[1:]
            else:
                new_columns.append("Certificate")

        elif col == "slip_ori":

            new_columns += _expand_clean_columns(
                df,
                all_clean_slip,
                "slip",
                max_slip
            )

        elif col == "insured_ori":

            new_columns += _expand_clean_columns(
                df,
                all_clean_ins,
                "insured",
                max_ins
            )

    df = df[new_columns]

    print(f"[5/5] Menyimpan hasil ke: {output_file} ...")
    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    df.to_excel(output_file, index=False)

    print(f"\n{'=' * 55}")
    print(f"  [OK] Selesai!")
    print(f"  Total baris output  : {len(df):,}")
    print(f"  Total kolom output  : {len(df.columns)}")
    print(f"  File disimpan di    : {output_file}")
    print(f"{'=' * 55}")


# ─────────────────────────────────────────────────────────────────────────────
# ENTRY POINT
# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    process_data(INPUT_FILE, OUTPUT_FILE)
    

# ─────────────────────────────────────────────────────────────────────────────
# CLEANING OSBAL
# ─────────────────────────────────────────────────────────────────────────────

INPUT_FILE = os.path.join("input", "2b. transaksi Osbal 01.01.23 - 17.08.26.xlsx")
OUTPUT_FILE = os.path.join("output", "wahana_output_osbalnew.xlsx")

CEDANT_COL   = "CCOS_COMP_NAME"
CEDANT_VALUE = "PT.ASURANSI WAHANA TATA"

POLIS_COL   = "FAC_POLICY_NO"
SLIP_COL    = "FAC_SLIP"
INSURED_COL = "FAC_INSURED"
CLASS_COL   = "CLASS_NAME" # Added for Marine Cargo check

CLSDT_POLIS_COL = "CLSDT_POLICY_NO"
CLSDT_SLIP_COL  = "CLSDT_SLIP_NO"
CLSDT_SERTF_COL = "CLSDT_SERTF_NO"

MAX_SPLIT_COLS = 5

# Polis: kata/prefix yang menyebabkan nilai dibiarkan apa adanya
POLIS_EXCEPTION_RE = re.compile(
    r"""
    MOP\s*MARINE
  | (?:LINE\s*SLIP|LINESLIP)
  | \b(?:P1|P2|P3|P73)\s*CANCEL
  | \b(?:P1|P2|P3|P73)\b
  | \bCANCEL\b
  | PENYELESAIAN
  | HUTANG\s*PIUTANG
    """,
    re.IGNORECASE | re.VERBOSE,
)

# Slip: kata yang menyebabkan nilai dibiarkan apa adanya
SLIP_EXCEPTION_RE = re.compile(
    r"""
    \bSUMMARY\b
  | \bBORDER[OA]\b
  | \bBORDRO\b
  | \bSINGGLESHIPMENT\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

MONTH_NAMES = frozenset({
    "JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI",
    "JULI", "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER",
    "JANUARY", "FEBRUARY", "MARCH", "MAY", "JUNE", "JULY", "AUGUST",
    "OCTOBER",
})

INSURED_JUNK_WORDS = frozenset({
    "PT", "CV", "TBK", "PERSERO", "LTD", "INC", "LLC",
    "AND", "OR", "THE", "OF", "AS",
    "NON FOOD", "DIV",
})

INSURED_SUFFIX_RE = re.compile(
    r"""
    ,?\s*\bTBK\b\s*(?:,?\s*PT\.?)?
  | ,?\s*\bPT\.?\s*$
  | ,?\s*\bCV\.?\s*$
    """,
    re.IGNORECASE | re.VERBOSE,
)

# Hapus token PT / QQ dimanapun posisinya (bukan hanya suffix)
INSURED_PT_QQ_RE = re.compile(
    r"""
    \bPT\.?\b
  | \bQQ\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

# Noise token dalam slip (bulan, tahun, mata uang, dsb.)
_SLIP_NOISE_RE = re.compile(
    r"""
    \b(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS
        |SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\b
  | \b(?:JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST
        |SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)\b
  | \b(?:IDR|USD|ENG)\b
  | \b(?:P1|P2|P3|P73)\b
  | \bNEW\b
  | \bVARIOUS\b
  | \b20[0-9]{2}\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

# Kandidat token nomor slip: alfanumerik ≥ 7 karakter
_SLIP_TOKEN_RE = re.compile(r"[A-Z0-9][A-Z0-9\-]{6,}", re.IGNORECASE)


# ─────────────────────────────────────────────────────────────────────────────
# SELEKSI SUMBER DATA (CLSDT vs FAC FALLBACK)
# ─────────────────────────────────────────────────────────────────────────────

def resolve_source_value(clsdt_val, fac_val):
    """
    Prioritaskan CLSDT jika nilainya valid.
    CLSDT tidak valid jika: kosong/NaN atau bernilai TBA, VAR, VARIOUS.
    Jika tidak valid, gunakan FAC sebagai fallback.
    """
    if pd.isna(clsdt_val):
        return fac_val

    clsdt_str = str(clsdt_val).strip()
    if not clsdt_str:
        return fac_val

    if clsdt_str.upper() in {"TBA", "VAR", "VARIOUS"}:
        return fac_val

    return clsdt_val


# ─────────────────────────────────────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _normalize_spaces(text: str) -> str:
    return re.sub(r"\s{2,}", " ", text).strip()


def _is_valid_polis_token(tok: str) -> bool:
    """Token polis valid: ≥5 char, punya digit, bukan TBA (murni/suffix)."""
    tok = tok.strip()
    return (
        len(tok) >= 5
        and not re.search(r"TBA$", tok, re.IGNORECASE)
        and bool(re.search(r"\d", tok))
    )


def _extract_polis_tokens(text: str) -> list:
    """Ekstrak semua token nomor polis valid dari teks."""
    tokens = []
    for block in re.split(r"\s{2,}", text.strip()):
        for tok in re.split(r"\+|\s+", block.strip()):
            tok = tok.strip().strip("-/")
            if _is_valid_polis_token(tok):
                tokens.append(tok)
    return tokens


def _is_valid_slip_token(tok: str) -> bool:
    """Token slip valid: ≥7 char, punya digit, bukan tahun 4-digit."""
    tok = tok.strip()
    return (
        len(tok) >= 7
        and bool(re.search(r"\d", tok))
        and not re.match(r"^\d{4}$", tok)
    )


def _strip_slip_suffix(tok: str) -> str:
    """Hapus suffix pendek numerik setelah dash (misal: -03, -000059)."""
    m = re.match(r"^(.+?)-(\d{1,6})$", tok)
    if m and len(m.group(1)) > len(m.group(2)):
        return m.group(1)
    return tok


def _extract_slip_tokens(text: str) -> list:
    """Ekstrak semua token nomor slip valid dari teks."""
    results = []
    for block in re.split(r"\s{2,}", text.strip()):
        clean = _SLIP_NOISE_RE.sub(" ", block)
        clean = re.sub(r"^[\s\-/+,]+|[\s\-/+,]+$", "", clean).strip()

        for cand in _SLIP_TOKEN_RE.findall(clean):
            cand = cand.strip("-")
            parts = cand.split("-")
            if len(parts) == 2:
                a, b = parts
                if _is_valid_slip_token(a) and _is_valid_slip_token(b) and abs(len(a) - len(b)) <= 2:
                    results.extend([a, b])
                    continue
            if _is_valid_slip_token(cand):
                results.append(_strip_slip_suffix(cand))

    return results


def _clean_insured_name(name: str) -> str:
    """Hapus suffix badan usaha (TBK, PT, CV) secara iteratif, lalu hapus
    token PT/QQ dimanapun posisinya dan tanda kurung (isi di dalamnya
    dipertahankan)."""
    name = _normalize_spaces(name)
    for _ in range(3):
        cleaned = _normalize_spaces(INSURED_SUFFIX_RE.sub("", name).strip().strip(","))
        if cleaned == name:
            break
        name = cleaned

    name = INSURED_PT_QQ_RE.sub(" ", name)
    name = name.replace("(", " ").replace(")", " ")
    name = _normalize_spaces(name)
    name = re.sub(r"^[\s.,\-]+|[\s.,\-]+$", "", name)

    return name


def _cap_or_join(items: list) -> list:
    """Jika jumlah item melebihi MAX_SPLIT_COLS, gabungkan dengan koma."""
    if len(items) > MAX_SPLIT_COLS:
        return [",".join(items)]
    return items

def parse_dash_chain(val):
    """
    Memisahkan nomor polis berdasarkan '-'.

    Hanya suffix yang merupakan breakdown polis yang
    akan dibuat menjadi clean polis.

    Contoh:
    09840502012020000115-110-114-109
    -> [
        09840502012020000115,
        09840502012020000110,
        09840502012020000114,
        09840502012020000109
    ]

    Sedangkan:
    02440502012023000394-000061-0003
    -> hanya base karena suffix bukan breakdown.
    """
    m = re.fullmatch(r"(\d+)((?:-\d+)+)", str(val).strip())

    if not m:
        return [], []

    base = m.group(1)
    suffixes = [x for x in m.group(2).split("-") if x]

    polis_list = [base]
    certificate_suffixes = []

    for suffix in suffixes:
        if _is_polis_breakdown_suffix(base, suffix):
            suffix_len = len(suffix)

            # Bentuk nomor polis lengkap berdasarkan suffix
            full_polis = (
                base[:-suffix_len] + suffix
                if len(base) > suffix_len
                else suffix
            )

            polis_list.append(full_polis)
        else:
            # Bukan breakdown -> biarkan certificate menangani
            certificate_suffixes.append(suffix)

    return polis_list, certificate_suffixes

# ─────────────────────────────────────────────────────────────────────────────
# CLEAN POLIS (WAHANA RULES - TIDAK DIUBAH)
# ─────────────────────────────────────────────────────────────────────────────
def _clean_polis_core(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).strip()

    if not val:
        return []

    if POLIS_EXCEPTION_RE.search(val):
        return [_normalize_spaces(val)]

    if re.search(r"\d+TBA\d+", val, re.IGNORECASE):
        return [_normalize_spaces(val)]

    val_upper = val.upper().strip()

    if "/" in val:
        val = val.split("/", 1)[0].strip()

    if re.fullmatch(r"VARIOUS", val_upper):
        return [val.strip()]

    if re.match(r"^VARIOUS\s*-\s*SEE\s+ATTACH", val_upper):
        return [val.strip()]

    if re.fullmatch(r"TBA", val_upper):
        return [val.strip()]

    if re.match(r"^\s*P1\s*/", val, re.IGNORECASE):
        after = re.sub(r"^\s*P1\s*/\s*", "", val, flags=re.IGNORECASE).strip()
        tokens = _extract_polis_tokens(after)
        if tokens:
            return [_normalize_spaces(tokens[0])]

    if re.fullmatch(r"\s*P1\s*(?:\+|\()\s*P2\s*\)?\s*", val, re.IGNORECASE):
        return [_normalize_spaces(val)]

    if re.search(r"[A-Za-z]", val):
        tokens = _extract_polis_tokens(val)
        if tokens:
            m = re.search(r"\d", val)
            if m:
                prefix = val[:m.start()].strip()
                if prefix:
                    return [_normalize_spaces(tokens[0])]

    if re.search(r"S/D", val, re.IGNORECASE):
        sd_parts = re.split(r"\s{2,}", val.strip())
        left = sd_parts[0] if sd_parts else val

        base_m = (re.match(r"^(\d[\d\-]+?)-\d+S/D\d+", left, re.IGNORECASE) or
                  re.match(r"^(\d[\d\-]+)S/D\d+", left, re.IGNORECASE))
        results = ([base_m.group(1).strip()] if base_m else [])

        for part in sd_parts[1:]:
            part = re.sub(r"\+?\s*\bTBA\b\s*", "", part, flags=re.IGNORECASE).strip().strip("+")
            results += [t for t in part.split() if _is_valid_polis_token(t.strip())]

        if results:
            return _cap_or_join(results)

    if re.match(r"^\s*TBA\s{2,}", val, re.IGNORECASE):
        rest = re.sub(r"^\s*TBA\s+", "", val, flags=re.IGNORECASE).strip()
        tokens = _extract_polis_tokens(rest)
        if tokens:
            return _cap_or_join(tokens)


    dash_match = re.fullmatch(r"(\d+)((?:-\d+)+)", val)

    if dash_match:
        base = dash_match.group(1)
        suffixes = [
            x for x in dash_match.group(2).split("-")
            if x
        ]

        results = [base]

        for suffix in suffixes:

            # =========================================================
            # 1. SUFFIX <= 6 DIGIT
            #    Cek apakah ini breakdown polis.
            # =========================================================
            if len(suffix) <= 6:
                if _is_polis_breakdown_suffix(base, suffix):
                    suffix_len = len(suffix)

                    if len(base) > suffix_len:
                        results.append(
                            base[:-suffix_len] + suffix
                        )

            # =========================================================
            # 2. SUFFIX > 6 DIGIT
            #    Tidak mungkin certificate.
            #    Untuk CLEAN POLIS dianggap breakdown / bagian polis.
            # =========================================================
            else:
                # Suffix > 6 digit = breakdown polis langsung
                # Tidak digabung dengan base
                results.append(suffix)

        return _cap_or_join(results)

    # ============================================================
    # DASH CHAIN:
    # BEDAKAN BREAKDOWN POLIS VS CERTIFICATE
    # ============================================================
    

    if re.search(r"\s{2,}", val):
        blocks = re.split(r"\s{2,}", val.strip())
        first_block = blocks[0].strip()

        dash_first = re.fullmatch(r"(\d+)((?:-\d+)+)", first_block)
        if dash_first:
            base = dash_first.group(1)
            suffixes = [x for x in dash_first.group(2).split("-") if x]
            if suffixes:
                lengths = {len(s) for s in suffixes}
                if len(lengths) == 1:
                    suffix_len = len(suffixes[0])
                    if 2 <= suffix_len <= 6:
                        if len(base) > suffix_len:
                            results = [base]
                            for suffix in suffixes:
                                results.append(base[:-suffix_len] + suffix)
                            for block in blocks[1:]:
                                tokens = _extract_polis_tokens(block)
                                for token in tokens:
                                    if _is_valid_polis_token(token.strip()):
                                        results.append(token.strip())
                            return _cap_or_join(results)

        results = []
        for block in blocks:
            tokens = _extract_polis_tokens(block)
            if tokens:
                results.append(_normalize_spaces(tokens[0]))
            if len(blocks) > 1:
                first = blocks[0].strip()
                m = re.fullmatch(r"(.+?)(\d{6})\.00", first)
                if m:
                    prefix = m.group(1)
                    results = [first]
                    ok = True
                    for b in blocks[1:]:
                        b = b.strip()
                        if re.fullmatch(r"\d{6}\.00", b):
                            results.append(prefix + b)
                        else:
                            ok = False
                            break
                    if ok:
                        return _cap_or_join(results)

        if results:
            return _cap_or_join(results)

    if "/" in val:
        slash_parts = re.split(r"\s*/\s*", val, maxsplit=1)
        left = slash_parts[0].strip()
        if left:
            if _is_valid_polis_token(left):
                return [_normalize_spaces(left)]
            tokens = _extract_polis_tokens(left)
            if tokens:
                return [_normalize_spaces(tokens[0])]

    dash_plus_match = re.fullmatch(r"([\d.]+)-(\d+(?:\+\d+)+)", val)

    if dash_plus_match:
        base = dash_plus_match.group(1)
        suffixes = [
            s for s in dash_plus_match.group(2).split("+")
            if s
        ]

        results = [base]

        for suffix in suffixes:
            if _is_polis_breakdown_suffix(base, suffix):
                suffix_len = len(suffix)

                if len(base) > suffix_len:
                    results.append(
                        base[:-suffix_len] + suffix
                    )

        if len(results) > 1:
            return _cap_or_join(results)

        return [_normalize_spaces(base)]

    val = re.sub(r"\s*\+\s*P[3-9]\d*\b", "", val, flags=re.IGNORECASE)
    val = _normalize_spaces(val)

    m = re.fullmatch(r"(.+?\.)(\d{6})((?:\+\d{3})+)", val.replace(" ", ""))
    if m:
        prefix = m.group(1)
        first = m.group(2)
        suffixes = [s for s in m.group(3).split("+") if s]
        results = [prefix + first + ".00"]
        for s in suffixes:
            results.append(prefix + first[:-3] + s + ".00")
        return _cap_or_join(results)

    val = re.sub(r"\b[^\s+]*[A-Za-z][^\s+]*\b", "", val)
    val = re.sub(r"\s*\+\s*", " + ", val)
    val = re.sub(r"^\s*\+\s*|\s*\+\s*$", "", val)
    val = re.sub(r"(?:\+\s*){2,}", "+ ", val)
    val = _normalize_spaces(val)

    if "+" in val:
        return [_normalize_spaces(val)]

    cleaned = _normalize_spaces(val)
    cleaned = cleaned.strip(" -/")

    if not cleaned:
        return []
    return [cleaned]



def clean_polis(val) -> list:
    return [r.replace(".", "") for r in _clean_polis_core(val)]


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN SLIP
# ─────────────────────────────────────────────────────────────────────────────
def _clean_slip_core(val) -> list:
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val:
        return []
    val = _normalize_spaces(val)
    
    if re.fullmatch(r"(P\d+(?:\s*\+\s*P\d+)*(\s+CANCEL)?|(JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\s+\d{4}(?:\s*-\s*(IDR|USD|EUR|GBP|SGD|JPY|AUD|CNY))?)", val, flags=re.IGNORECASE):
        return []

    val = re.sub(r"\s*/\s*END\b", "", val, flags=re.IGNORECASE).strip()
    val = re.sub(r"\+VAR\b", "", val, flags=re.IGNORECASE)
    val = _normalize_spaces(val)

    parts = re.split(r"\s{2,}", val)
    parts = [p.strip() for p in parts if p.strip()]
    if len(parts) >= 2:
        results = []
        for p in parts:
            results.append(_normalize_spaces(p))
        return _cap_or_join(results)
   
    if SLIP_EXCEPTION_RE.search(val):
        return [_normalize_spaces(val)]

    if re.fullmatch(r"\d+(?:\.\d+)+(?:-\d+){3,}", val):
        return [_normalize_spaces(val)]

    if "+" in val:
        parts = [p.strip() for p in re.split(r"\s*\+\s*", val) if p.strip()]
        parts = [p for p in parts if not re.fullmatch(r"P\d+", p, flags=re.IGNORECASE)]
        if len(parts) > 1:
            first = parts[0]
            rest = parts[1:]
            rest_is_digit = all(re.fullmatch(r"\d+", r) for r in rest)
            suffix_lengths = ({len(r) for r in rest} if rest_is_digit else set())
            if rest_is_digit and len(suffix_lengths) == 1:
                suffix_len = suffix_lengths.pop()
                m = re.match(rf"^(.*?)(\d{{{suffix_len}}})$", first)
                if m:
                    prefix = m.group(1)
                    results = [first]
                    for s in rest:
                        s = s.zfill(suffix_len)
                        results.append(prefix + s)
                    return _cap_or_join(results)
            if all(re.match(r"^\d+(?:\.\d+)+$", p) for p in parts):
                return _cap_or_join(parts)
            return [_normalize_spaces(val)]

    m = re.match(r"^(.*?)(\d+)-(\d+)$", val)
    if m:
        prefix = m.group(1)
        first = m.group(2)
        second = m.group(3)
        results = [prefix + first]
        if len(second) < len(first):
            second = second.zfill(len(first))
            results.append(prefix + second)
        else:
            results.append(second)
        return _cap_or_join(results)

    slips = re.findall(r"\d+(?:\.\d+)+", val)
    if len(slips) >= 2:
        return _cap_or_join(slips)

    m = re.search(r"\d+(?:\.\d+)+", val)
    if m:
        return [m.group(0)]

    return [_normalize_spaces(val)]     
    
def clean_slip(val) -> list:
    results = _clean_slip_core(val)
    cleaned = []
    for r in results:
        if pd.isna(r):
            continue
        r = str(r)
        r = _normalize_spaces(r)
        r = r.replace(".", "")
        r = r.strip()

        if not r:
            continue

        # RULE BARU:
        # Hanya slip yang diawali angka 6 yang ditambahkan "000"
        if r.startswith("6"):
            r = "000" + r

        if r not in cleaned:
            cleaned.append(r)

    return cleaned


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN INSURED
# ─────────────────────────────────────────────────────────────────────────────
def clean_insured(val) -> list:
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val:
        return []

    cleaned = []
    for p in re.split(r"/|,", val):
        p = _normalize_spaces(p.strip())
        if len(p) <= 2:
            continue
        if p.upper().strip() in INSURED_JUNK_WORDS:
            continue
        if re.match(r"^[^a-zA-Z0-9]+$", p):
            continue

        p_clean = _clean_insured_name(p)
        if not p_clean or len(p_clean) <= 2:
            continue
        if p_clean.upper().strip() in INSURED_JUNK_WORDS:
            continue
        cleaned.append(p_clean)

    if not cleaned:
        fallback = _clean_insured_name(_normalize_spaces(val))
        return [fallback] if fallback else []
    return _cap_or_join(cleaned)


# ─────────────────────────────────────────────────────────────────────────────
# NEW FORMAT CERT RANGE HELPER
# ─────────────────────────────────────────────────────────────────────────────
def format_cert_range(start_str, end_str):
    try:
        start_int = int(start_str)
        end_int = int(end_str)
        
        # JIKA TERBALIK (misal 61 lalu 3), JANGAN PAKAI SD! Pecah jadi koma.
        if end_int < start_int:
            return f"{str(start_int).zfill(6)}, {str(end_int).zfill(6)}"
            
        count = end_int - start_int + 1
        if count > 3:
            return f"{str(start_int).zfill(6)} SD {str(end_int).zfill(6)}"
        else:
            return ", ".join(str(i).zfill(6) for i in range(start_int, end_int + 1))
    except ValueError:
        return ""

# ─────────────────────────────────────────────────────────────────────────────
# CLEAN CERTIFICATE
# ─────────────────────────────────────────────────────────────────────────────
def clean_certificate(polis_ori, clsdt_sertf=None, class_name=None):
    """
    Membersihkan certificate.

    PRIORITAS:
    1. Jika CLSDT_SERTF_NO valid -> gunakan CLSDT_SERTF_NO
    2. Jika CLSDT_SERTF_NO tidak valid -> fallback ke POLIS ORI

    CLSDT tidak valid jika:
    - kosong / NaN
    - TBA
    - VAR
    - VARIOUS

    Certificate:
    - maksimal 3
    - numeric
    - maksimal 6 digit
    - jika < 6 digit -> zfill(6)
    - range menggunakan format SD jika > 3 nomor
    """

    # ============================================================
    # HELPER VALIDASI CLSDT CERTIFICATE
    # ============================================================
    def _valid_clsdt_certificate(value):
        if pd.isna(value):
            return False

        value = str(value).strip()

        if not value:
            return False

        if value.upper() in {"TBA", "VAR", "VARIOUS"}:
            return False

        return True

    # ============================================================
    # 1. PRIORITAS CLSDT_SERTF_NO
    # ============================================================
    if _valid_clsdt_certificate(clsdt_sertf):

        clsdt = str(clsdt_sertf).strip()

        # Normalisasi S/D
        clsdt = re.sub(
            r"\s*S\s*/?\s*D\s*",
            " SD ",
            clsdt,
            flags=re.IGNORECASE
        )

        clsdt = _normalize_spaces(clsdt)

        # --------------------------------------------------------
        # CLSDT berbentuk range:
        # 000001 S/D 000005
        # --------------------------------------------------------
        m_sd = re.fullmatch(
            r"(\d{1,6})(?:\.00)?\s*SD\s*(\d{1,6})(?:\.00)?",
            clsdt,
            re.IGNORECASE
        )

        if m_sd:
            result = format_cert_range(
                m_sd.group(1),
                m_sd.group(2)
            )

            return [result] if result else []

        # --------------------------------------------------------
        # CLSDT certificate tunggal
        # --------------------------------------------------------
        if re.fullmatch(r"\d{1,6}(?:\.00)?", clsdt):
            number = re.sub(r"\.00$", "", clsdt)

            return [number.zfill(6)]

        # --------------------------------------------------------
        # CLSDT beberapa certificate dengan koma
        # --------------------------------------------------------
        if "," in clsdt:
            certificates = []

            for part in clsdt.split(","):
                part = part.strip()

                if re.fullmatch(r"\d{1,6}(?:\.00)?", part):
                    part = re.sub(r"\.00$", "", part)
                    certificates.append(part.zfill(6))

                    if len(certificates) >= 3:
                        break

            if certificates:
                return certificates

    # ============================================================
    # 2. FALLBACK KE POLIS ORI
    # ============================================================
    if pd.isna(polis_ori):
        return []

    polis = str(polis_ori).strip()

    if not polis:
        return []

    # ============================================================
    # EXCEPTION
    # ============================================================
    if polis == "02210502012023000128-000129":
        return []

    if polis == "00940502012023001276-136-135-1275":
        return []

    # ============================================================
    # 3. HANDLE S/D DARI POLIS ORI
    # ============================================================
    m_sd = re.search(
        r"(\d{1,6})(?:\.00)?\s*S\s*/?\s*D\s*(\d{1,6})(?:\.00)?\b",
        polis,
        re.IGNORECASE
    )

    if m_sd:
        first = m_sd.group(1)
        last = m_sd.group(2)

        result = format_cert_range(first, last)

        return [result] if result else []

    # ============================================================
    # 4. HANDLE DASH DARI POLIS ORI
    # ============================================================
    if "-" not in polis:
        return []

    parts = [p.strip() for p in polis.split("-")]

    if len(parts) < 2:
        return []

    base = parts[0]
    suffixes = parts[1:]

    certificates = []

    for suffix in suffixes:

        # Harus angka murni
        if not re.fullmatch(r"\d+", suffix):
            continue

        # Lebih dari 6 digit bukan certificate
        if len(suffix) > 6:
            continue

        # ========================================================
        # CEK BREAKDOWN POLIS
        # ========================================================
        if _is_polis_breakdown_suffix(base, suffix):
            continue

        # ========================================================
        # BUKAN BREAKDOWN -> CERTIFICATE
        # ========================================================
        certificates.append(suffix.zfill(6))

        if len(certificates) >= 3:
            break

    return certificates

def _is_polis_breakdown_suffix(base, suffix):
    """
    Menentukan apakah suffix adalah lanjutan nomor polis,
    bukan certificate.

    Contoh:
    09810502012021000074-00073
    -> 00073 = breakdown polis

    09840502012020000115-110
    -> 110 = breakdown polis

    00940502012020001202-1201
    -> 1201 = breakdown polis

    Sedangkan:
    00940502012020001202-0074
    -> 0074 = certificate
    """

    if not re.fullmatch(r"\d+", base):
        return False

    if not re.fullmatch(r"\d+", suffix):
        return False

    suffix_len = len(suffix)

    # Suffix terlalu panjang → bukan breakdown polis
    if suffix_len > 6:
        return False

    # Minimal 2 digit untuk pola breakdown
    if suffix_len < 2:
        return False

    if len(base) <= suffix_len:
        return False

    base_end = base[-suffix_len:]

    try:
        base_num = int(base_end)
        suffix_num = int(suffix)

        diff = suffix_num - base_num

        # Breakdown polis biasanya nomor berdekatan.
        # Mendukung kasus turun seperti:
        # 74 -> 73
        # 1202 -> 1201
        #
        # dan naik seperti:
        # 128 -> 129
        return -10 <= diff <= 10

    except ValueError:
        return False

# ─────────────────────────────────────────────────────────────────────────────
# PROCESS DATA
# ─────────────────────────────────────────────────────────────────────────────
def _insert_clean_columns(df: pd.DataFrame, all_lists: list, prefix: str, max_cols: int) -> list:
    """Buat kolom clean_{prefix}_N dan kembalikan nama-nama kolom baru."""
    added = []
    for i in range(1, max_cols + 1):
        col_name = f"clean {prefix} {i}"
        df[col_name] = [lst[i - 1] if i - 1 < len(lst) else None for lst in all_lists]
        added.append(col_name)
    return added

def process_data(input_file: str, output_file: str) -> None:
    print(f"[1/5] Membaca data dari: {input_file} ...")
    df = pd.read_excel(input_file, header=0)
    print(f"      Total baris keseluruhan: {len(df):,}")

    if CEDANT_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{CEDANT_COL}' tidak ditemukan!")
        print(f"        Kolom tersedia: {list(df.columns)}")
        return

    required_cols = [
        POLIS_COL, SLIP_COL, INSURED_COL,
        CLSDT_POLIS_COL, CLSDT_SLIP_COL, CLSDT_SERTF_COL,
    ]

    if CLASS_COL not in df.columns:
        df[CLASS_COL] = ""

    for col in required_cols:
        if col not in df.columns:
            print(f"\n[ERROR] Kolom '{col}' tidak ditemukan di file!")
            print(f"        Kolom tersedia: {list(df.columns)}")
            return

    df[CEDANT_COL] = df[CEDANT_COL].astype(str).str.strip()
    is_wahana = df[CEDANT_COL] == CEDANT_VALUE
    print(
        f"[2/5] Filter cedant '{CEDANT_VALUE}': "
        f"{is_wahana.sum():,} baris WAHANA dari total {len(df):,} baris."
    )

    # Output cleaning hanya untuk Cedant WAHANA.
    df = df.loc[is_wahana].copy()
    is_wahana = pd.Series(True, index=df.index)
    print(f"      Output cleaning OSBAL dibatasi ke WAHANA: {len(df):,} baris.")

    # Simpan nama internal untuk kebutuhan cleaning/matching.
    df.rename(
        columns={
            POLIS_COL: "polis_ori",
            SLIP_COL: "slip_ori",
            INSURED_COL: "insured_ori",
        },
        inplace=True,
    )

    print("[3/5] Menjalankan proses selection sumber (CLSDT vs FAC) dan cleaning ...")

    all_clean_polis = []
    all_clean_slip = []
    all_clean_ins = []
    all_certificates = []

    max_polis = max_slip = max_ins = 1

    for idx, (_, row) in enumerate(df.iterrows(), 1):
        if idx % 50_000 == 0:
            print(f"      Progress: {idx:,} / {len(df):,} baris diproses...")

        raw_polis = resolve_source_value(
            row.get(CLSDT_POLIS_COL),
            row.get("polis_ori"),
        )
        raw_slip = resolve_source_value(
            row.get(CLSDT_SLIP_COL),
            row.get("slip_ori"),
        )
        raw_ins = row.get("insured_ori", "")
        class_nm = row.get(CLASS_COL, "")

        c_polis = clean_polis(raw_polis)
        c_slip = clean_slip(raw_slip)
        c_ins = clean_insured(raw_ins)
        c_certificate = clean_certificate(
            row.get("polis_ori"),
            row.get(CLSDT_SERTF_COL),
            class_name=class_nm,
        )

        max_polis = max(max_polis, len(c_polis))
        max_slip = max(max_slip, len(c_slip))
        max_ins = max(max_ins, len(c_ins))

        all_clean_polis.append(c_polis)
        all_clean_slip.append(c_slip)
        all_clean_ins.append(c_ins)

        if isinstance(c_certificate, list):
            all_certificates.append(c_certificate)
        elif c_certificate:
            all_certificates.append([str(c_certificate).strip()])
        else:
            all_certificates.append([])

    print("      Selesai diproses!")
    print(f"      -> Jumlah kolom clean polis  : {max_polis}")
    print(f"      -> Jumlah kolom clean slip   : {max_slip}")
    print(f"      -> Jumlah kolom clean insured: {max_ins}")

    print("[4/5] Menyusun kolom output sesuai template Excel ...")

    # Output Excel hanya memakai kolom template yang diberikan.
    # Kolom sumber/internal lain tetap tersedia di backend dataframe,
    # tetapi tidak dikeluarkan ke file cleaning.
    output_columns = [
        "CCOS_DOC_NO", "CCOS_DATE", "CCOS_REF_CODE", "CCOS_COMP",
        "CCOS_COMP_NAME", "CCOS_REF_COMP", "CCOS_REF_COMP_NAME",
        "FAC_INSURED", "INSURED_CLEAN_1", "INSURED_CLEAN_2",
        "INSURED_CLEAN_3", "INSURED_CLEAN_4", "INSURED_CLEAN_5",
        "CCOS_CURR", "CCOS_OR_BAL", "CCOS_BAL_DUE",
        "CCOS_OR_BAL_IN_IDR", "CCOS_BAL_DUE_IN_IDR",
        "FAC_COM_DATE", "FAC_EXP_DATE", "FAC_DUE_DATES",
        "FAC_SUB_CLASS", "FAC_POLICY_NO", "POLICY_CLEAN_1",
        "SERTIF_CLEAN_1", "POLICY_CLEAN_2", "SERTIF_CLEAN_2",
        "POLICY_CLEAN_3", "SERTIF_CLEAN_3", "POLICY_CLEAN_4",
        "POLICY_CLEAN_5", "FAC_SLIP", "SLIP_CLEAN_1",
        "SLIP_CLEAN_2", "SLIP_CLEAN_3", "SLIP_CLEAN_4",
        "SLIP_CLEAN_5", "CLASS_CODE", "CLASS_NAME",
        "CLSDT_POLICY_NO", "CLSDT_SLIP_NO", "CLSDT_SERTF_NO",
    ]

    def source_series(column_name, fallback=""):
        if column_name in df.columns:
            return df[column_name]
        return pd.Series([fallback] * len(df), index=df.index, dtype=object)

    output_df = pd.DataFrame(index=df.index)

    # Semua kolom template: kalau tidak tersedia, otomatis blank.
    for col in output_columns:
        output_df[col] = source_series(col)

    # Kolom raw yang sebelumnya dipakai sebagai alias internal.
    output_df["FAC_INSURED"] = source_series("insured_ori")
    output_df["FAC_POLICY_NO"] = source_series("polis_ori")
    output_df["FAC_SLIP"] = source_series("slip_ori")

    # Clean insured.
    for i in range(1, 6):
        output_df[f"INSURED_CLEAN_{i}"] = [
            vals[i - 1] if i - 1 < len(vals) else ""
            for vals in all_clean_ins
        ]

    # Clean policy + certificate.
    for i in range(1, 6):
        output_df[f"POLICY_CLEAN_{i}"] = [
            vals[i - 1] if i - 1 < len(vals) else ""
            for vals in all_clean_polis
        ]

    for i in range(1, 4):
        output_df[f"SERTIF_CLEAN_{i}"] = [
            vals[i - 1] if i - 1 < len(vals) else ""
            for vals in all_certificates
        ]

    # Clean slip.
    for i in range(1, 6):
        output_df[f"SLIP_CLEAN_{i}"] = [
            vals[i - 1] if i - 1 < len(vals) else ""
            for vals in all_clean_slip
        ]

    # Pastikan urutan persis sama dengan template.
    output_df = output_df[output_columns]

    print(f"[5/5] Menyimpan hasil ke: {output_file} ...")
    output_df.to_excel(output_file, index=False)

    print(f"\n{'=' * 55}")
    print("  [OK] Selesai!")
    print(f"  Total baris output  : {len(output_df):,}")
    print(f"  Total kolom output  : {len(output_df.columns)}")
    print(f"  File disimpan di    : {output_file}")
    print(f"{'=' * 55}")

# ─────────────────────────────────────────────────────────────────────────────
# ENTRY POINT (TEST POLA STRING)
# ─────────────────────────────────────────────────────────────────────────────
# ═════════════════════════════════════════════════════════════════════════════
# ENTRY POINT - TEST CERTIFICATE
# ═════════════════════════════════════════════════════════════════════════════

# if __name__ == "__main__":

#     TEST_DATA = [
#         # Test sebelumnya
#         "02210502012023000128-000129",
#         "02440502012023000394-000061-0003",
#         "00940502012023001276-136-135-1275",
#         "02210502012023000494-70320230003-002",

#         # Test tambahan
#         "0041050201202000183-0045",
#         "09810502012021000074-00073",
#         "01810502012020000259-0008",
#         "09840502012020000115-110-114-109",
#         "00940502012020001202-0074-1201-0073",
#         "02240502012020001376-000320",
#         "01540502012021000442-51481443442445452453457595861",
#         "09810502012020000566-00383-00045-00031",
#     ]

#     print("=" * 100)
#     print("TEST CLEAN POLIS + CERTIFICATE")
#     print("=" * 100)

#     for no, polis_ori in enumerate(TEST_DATA, 1):

#         clean_polis_result = clean_polis(polis_ori)
#         certificate_result = clean_certificate(polis_ori)

#         print(f"\n[{no}]")
#         print(f"POLIS ORI : {polis_ori}")

#         print("\nCLEAN POLIS:")
#         for i in range(5):
#             value = (
#                 clean_polis_result[i]
#                 if i < len(clean_polis_result)
#                 else ""
#             )
#             print(f"  clean polis {i + 1} : {value}")

#         print("\nCERTIFICATE:")
#         for i in range(3):
#             value = (
#                 certificate_result[i]
#                 if i < len(certificate_result)
#                 else ""
#             )
#             print(f"  Certificate {i + 1} : {value}")

#         print("-" * 100)

#     print("\n" + "=" * 100)
#     print("TEST SELESAI")
#     print("=" * 100)

if __name__ == "__main__":
    process_data(INPUT_FILE, OUTPUT_FILE)

# ─────────────────────────────────────────────────────────────────────────────
# CLEANING FACUL
# ─────────────────────────────────────────────────────────────────────────────    

INPUT_FILE  = os.path.join("input", "1b. Transaksi Facul 01.01.23 - 17.08.2026.xlsx")
OUTPUT_FILE = os.path.join("output", "wahana_output_facul.xlsx")

CEDANT_COL   = "COMP_NAME"
CEDANT_VALUE = "PT.ASURANSI WAHANA TATA"

POLIS_COL   = "FAC_POLICY_NO"
SLIP_COL    = "FAC_SLIP"
INSURED_COL = "FAC_INSURED"
BROKER_COL  = "COMP_NAME.1"

MAX_SPLIT_COLS = 5

# Polis: kata/prefix yang menyebabkan nilai dibiarkan apa adanya
POLIS_EXCEPTION_RE = re.compile(
    r"""
    MOP\s*MARINE
  | (?:LINE\s*SLIP|LINESLIP)
  | \b(?:P1|P2|P3|P73)\s*CANCEL
  | \b(?:P1|P2|P3|P73)\b
  | \bCANCEL\b
  | PENYELESAIAN
  | HUTANG\s*PIUTANG
    """,
    re.IGNORECASE | re.VERBOSE,
)

# Slip: kata yang menyebabkan nilai dibiarkan apa adanya
SLIP_EXCEPTION_RE = re.compile(
    r"""
    \bSUMMARY\b
  | \bBORDER[OA]\b
  | \bBORDRO\b
  | \bSINGGLESHIPMENT\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

MONTH_NAMES = frozenset({
    "JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI",
    "JULI", "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER",
    "JANUARY", "FEBRUARY", "MARCH", "MAY", "JUNE", "JULY", "AUGUST",
    "OCTOBER",
})

# =========================
# INSURED CONFIG
# =========================

INSURED_REMOVE_RE = re.compile(
    r"""
    \b(
        PT|CV|TBK|
        PERSERO|
        LTD|PTE|INC|LLC|
        MR|MRS|MS|
        BAPAK|BPK|IBU|NY|
        DR|DRS|DRA|IR|H|HJ
    )\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

INSURED_POLIS_RE = re.compile(
    r"""
    (POLIS\s*NO\.?.*)
    |
    (POLICY\s*NO\.?.*)
    |
    (SLIP\s*NO\.?.*)
    """,
    re.IGNORECASE | re.VERBOSE,
)

INSURED_SPLIT_RE = re.compile(
    r"""
    \s*,\s*
    |
    \s*/\s*
    |
    \s+QQ\s+
    |
    \s+AND/OR\s+
    |
    \s*&\s*
    |
    \s*\+\s*
    """,
    re.IGNORECASE | re.VERBOSE,
)

INSURED_JUNK_WORDS = {
    "",
    "AND",
    "OR",
    "THE",
    "OF",
    "AS"
}

# Noise token dalam slip
_SLIP_NOISE_RE = re.compile(
    r"""
    \b(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS
        |SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\b
  | \b(?:JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST
        |SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)\b
  | \b(?:IDR|USD|ENG)\b
  | \b(?:P1|P2|P3|P73)\b
  | \bNEW\b
  | \bVARIOUS\b
  | \b20[0-9]{2}\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

_SLIP_TOKEN_RE = re.compile(r"[A-Z0-9][A-Z0-9\-]{6,}", re.IGNORECASE)


# ─────────────────────────────────────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _normalize_spaces(text: str) -> str:
    return re.sub(r"\s{2,}", " ", text).strip()


def _is_valid_polis_token(tok: str) -> bool:
    """Token polis valid: ≥5 char, punya digit, bukan TBA (murni/suffix)."""
    tok = tok.strip()
    return (
        len(tok) >= 5
        and not re.search(r"TBA$", tok, re.IGNORECASE)
        and bool(re.search(r"\d", tok))
    )


def _extract_polis_tokens(text: str) -> list:
    """Ekstrak semua token nomor polis valid dari teks."""
    tokens = []
    for block in re.split(r"\s{2,}", text.strip()):
        for tok in re.split(r"\+|\s+", block.strip()):
            tok = tok.strip().strip("-/")
            if _is_valid_polis_token(tok):
                tokens.append(tok)
    return tokens


def _expand_dash_chain(segment: str) -> list:
    """
    Pecah satu segmen berisi base polis + suffix-suffix pendek dipisah dash.

    Contoh:
      '010115001310-1311-1313'
          → ['010115001310', '010115001311', '010115001313']
      '010115001310-010115300373'
          → ['010115001310', '010115300373']  (keduanya base)

    Aturan:
    - Token ≥8 digit → base baru.
    - Token pendek (semua digit, len < len(base)) → suffix: ganti N digit terakhir base.
    - Lainnya → tambahkan apa adanya jika valid.
    """
    results = []
    current_base = None

    for part in segment.split("-"):
        part = part.strip()
        if not part:
            continue

        if len(part) >= 8 and re.search(r"\d", part):
            current_base = part
            results.append(part)
        elif (
            current_base
            and len(part) >= 2
            and re.match(r"^\d+$", part)
            and len(part) < len(current_base)
        ):
            n = len(part)
            results.append(current_base[:-n] + part)
        else:
            if _is_valid_polis_token(part):
                results.append(part)

    return results or ([segment] if _is_valid_polis_token(segment) else [])


def _is_valid_slip_token(tok: str) -> bool:
    """Token slip valid: ≥7 char, punya digit, bukan tahun 4-digit."""
    tok = tok.strip()
    return (
        len(tok) >= 7
        and bool(re.search(r"\d", tok))
        and not re.match(r"^\d{4}$", tok)
    )


def _strip_slip_suffix(tok: str) -> str:
    """Hapus suffix pendek numerik setelah dash (misal: -03, -000059)."""
    m = re.match(r"^(.+?)-(\d{1,6})$", tok)
    if m and len(m.group(1)) > len(m.group(2)):
        return m.group(1)
    return tok


def _extract_slip_tokens(text: str) -> list:
    """Ekstrak semua token nomor slip valid dari teks."""
    results = []
    for block in re.split(r"\s{2,}", text.strip()):
        clean = _SLIP_NOISE_RE.sub(" ", block)
        clean = re.sub(r"^[\s\-/+,]+|[\s\-/+,]+$", "", clean).strip()

        for cand in _SLIP_TOKEN_RE.findall(clean):
            cand = cand.strip("-")
            parts = cand.split("-")
            if len(parts) == 2:
                a, b = parts
                if _is_valid_slip_token(a) and _is_valid_slip_token(b) and abs(len(a) - len(b)) <= 2:
                    results.extend([a, b])
                    continue
            if _is_valid_slip_token(cand):
                results.append(_strip_slip_suffix(cand))

    return results


def _clean_insured_name(name: str) -> str:

    if pd.isna(name):
        return ""

    name = str(name).upper()

    # hapus info polis/slip
    name = INSURED_POLIS_RE.sub("", name)

    # hapus badan usaha dan gelar
    name = INSURED_REMOVE_RE.sub(" ", name)

    # ganti karakter
    name = re.sub(r"[()]", " ", name)
    name = name.replace("/", " ")
    name = name.replace("-", " ")

    # rapikan spasi
    name = _normalize_spaces(name)

    return name

def split_insured(name):

    if pd.isna(name):
        return []

    name = str(name)

    parts = INSURED_SPLIT_RE.split(name)

    hasil = []

    for p in parts:

        p = _clean_insured_name(p)

        if not p:
            continue

        if p in INSURED_JUNK_WORDS:
            continue

        hasil.append(p)

    return hasil


def _cap_or_join(items: list) -> list:
    """Jika jumlah item melebihi MAX_SPLIT_COLS, gabungkan dengan koma."""
    if len(items) > MAX_SPLIT_COLS:
        return [",".join(items)]
    return items


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN POLIS  (facul memiliki tambahan: dot-suffix, ampersand, _expand_dash_chain)
# ─────────────────────────────────────────────────────────────────────────────

def clean_polis(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).strip()
    if not val:
        return []

    # exception
    if POLIS_EXCEPTION_RE.search(val):
        return [_normalize_spaces(val)]

    if re.search(r"\d+TBA\d+", val, re.IGNORECASE):
        return [_normalize_spaces(val)]

    val = _normalize_spaces(val)
    val_upper = val.upper()

    # ===============================
    # RULE 1
    # Hapus suffix /01 /02 /03 dst
    # Contoh:
    # 022117004330/02 -> 022117004330
    # ===============================
    val = re.sub(r"/\d{1,3}$", "", val)

    # ===============================
    # RULE 2
    # Jika hanya angka + titik
    # hapus semua titik
    # ===============================
    if re.fullmatch(r"[0-9.]+", val):
        return [val.replace(".", "")]

    # ===============================
    # RULE 3
    # Jika hanya angka + strip
    # hapus semua strip
    # ===============================
    if re.fullmatch(r"[0-9-]+", val):
        return [val.replace("-", "")]

    # ===============================
    # RULE 4
    # Polis yang ada huruf → tetap pertahankan isi,
    # tetapi hapus titik
    if re.search(r"[A-Z]", val_upper):
        return [val.replace(".", "")]

    # ===============================
    # RULE 5
    # Multi polis &
    # ===============================
    if "&" in val:
        hasil = []

        for p in val.split("&"):
            p = p.strip()

            p = re.sub(r"/\d{1,3}$", "", p)

            if re.fullmatch(r"[0-9.]+", p):
                p = p.replace(".", "")

            elif re.fullmatch(r"[0-9-]+", p):
                p = p.replace("-", "")

            if p:
                hasil.append(p)

        return _cap_or_join(hasil)

    # ===============================
    # HAPUS P1, P2, P3, P4, ...
    # ===============================

    val = re.sub(r"\bP\d+\b", "", val, flags=re.IGNORECASE)

    # Rapikan tanda + yang tersisa
    val = re.sub(r"\+\s*\+", "+", val)
    val = re.sub(r"^\s*\+\s*|\s*\+\s*$", "", val)

    val = _normalize_spaces(val)

    # ===============================
    # RULE 6
    # Multi polis +
    # ===============================
    if "+" in val:
        hasil = []

        for p in val.split("+"):
            p = p.strip()

            p = re.sub(r"/\d{1,3}$", "", p)

            if re.fullmatch(r"[0-9.]+", p):
                p = p.replace(".", "")

            elif re.fullmatch(r"[0-9-]+", p):
                p = p.replace("-", "")

            if p:
                hasil.append(p)

        return _cap_or_join(hasil)

    # ===============================
    # RULE 7
    # Polis biasa
    # ===============================
    return [val]
# ─────────────────────────────────────────────────────────────────────────────
# CLEAN SLIP
# ─────────────────────────────────────────────────────────────────────────────

def clean_slip(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).upper().strip()

    if not val:
        return []

    val = _normalize_spaces(val)

    # ==========================================
    # Rule 1 : Samakan delimiter multi slip
    # ==========================================
    val = re.sub(r"\s*&\s*", "/", val)
    val = re.sub(r"\s*\+\s*", "/", val)

    slips = []

    for s in val.split("/"):
        s = s.strip()

        if not s:
            continue

        # ==========================================
        # Rule 2 : Hapus keterangan di belakang
        # ==========================================
        s = re.sub(
            r"\b("
            r"USD|IDR|SGD|EUR|JPY|AUD|GBP|"
            r"ORI|ORI\.|ORIGINAL|COPY|"
            r"REALISASI|REALIZATION|"
            r"CANCEL|CANCELLED|"
            r"ENDT?|ENDORSEMENT|"
            r"SA|P1|P2|P3|VAR|REVISI|REV"
            r")\b.*$",
            "",
            s,
            flags=re.IGNORECASE,
        )

        # ==========================================
        # Rule 3 : Rapikan spasi
        # ==========================================
        s = _normalize_spaces(s)

        # ==========================================
        # Rule 4 : Hapus karakter di depan/belakang
        # ==========================================
        s = s.strip("-_,.; ")

        # ==========================================
        # Rule 4 : Hapus karakter di depan/belakang
        # ==========================================
        s = s.strip("-_,.; ")

        # Hapus titik di dalam slip
        s = s.replace(".", "")

        if not s:
            continue

        if not s:
            continue

        if s not in slips:
            slips.append(s)

    return _cap_or_join(slips)

# ─────────────────────────────────────────────────────────────────────────────
# CLEAN CERTIFICATE
# ─────────────────────────────────────────────────────────────────────────────

def clean_certificate(polis_ori):
    """
    Mengambil CERTIFICATE langsung dari POLIS ORI.

    RULE:
    - Polis harus mempunyai "-"
    - Satu suffix angka:
        ...-000156 -> 000156
        ...-00073  -> 000073
        ...-0045   -> 000045

    - Beberapa suffix:
        ...-110-114-109
        -> 000110 SD 000109

    - Suffix > 6 digit -> blank
    - Suffix bukan angka -> blank
    - Polis tanpa "-" -> blank
    """

    if pd.isna(polis_ori):
        return ""

    polis = str(polis_ori).strip()

    if not polis:
        return ""

    # ============================================================
    # HARUS ADA DASH
    # ============================================================

    if "-" not in polis:
        return ""

    # Ambil bagian setelah dash pertama
    parts = polis.split("-")

    if len(parts) < 2:
        return ""

    suffixes = [p.strip() for p in parts[1:]]

    # ============================================================
    # SEMUA SUFFIX HARUS ANGKA
    # ============================================================

    if not all(re.fullmatch(r"\d+", x) for x in suffixes):
        return ""

    # ============================================================
    # SINGLE CERTIFICATE
    # ============================================================

    if len(suffixes) == 1:

        cert = suffixes[0]

        if len(cert) > 6:
            return ""

        return cert.zfill(6)

    # ============================================================
    # RANGE CERTIFICATE
    # ============================================================

    first = suffixes[0]
    last = suffixes[-1]

    if len(first) > 6 or len(last) > 6:
        return ""

    return f"{first.zfill(6)} SD {last.zfill(6)}"

# ─────────────────────────────────────────────────────────────────────────────
# CLEAN INSURED
# ─────────────────────────────────────────────────────────────────────────────

def clean_insured(val) -> list:

    if pd.isna(val):
        return []

    val = str(val).strip()

    if not val:
        return []

    # Breakdown sesuai separator
    insureds = split_insured(val)

    cleaned = []

    for ins in insureds:

        ins = _normalize_spaces(ins)

        if not ins:
            continue

        if len(ins) <= 2:
            continue

        if ins.upper() in INSURED_JUNK_WORDS:
            continue

        # huruf kapital
        ins = ins.upper()

        if ins not in cleaned:
            cleaned.append(ins)

    # kalau tidak berhasil dibreakdown
    if not cleaned:

        fallback = _clean_insured_name(val)

        if fallback:
            cleaned.append(fallback.upper())

    return _cap_or_join(cleaned)
# ─────────────────────────────────────────────────────────────────────────────
# MITRA BISNIS
# ─────────────────────────────────────────────────────────────────────────────

def get_mitra_bisnis(comp_name2, comp_name) -> str:
    """
    Tentukan mitra_bisnis:
    - COMP_NAME2 kosong / 'DIRECT' → gunakan COMP_NAME (cedant)
    - Selain itu → gunakan COMP_NAME2 (broker)
    """
    broker = "" if pd.isna(comp_name2) else str(comp_name2).strip()
    if not broker or broker.upper() == "DIRECT":
        return "" if pd.isna(comp_name) else str(comp_name).strip()
    return broker

def clean_business_partners(val):

    if pd.isna(val):
        return ""

    val = str(val).strip().upper()

    # PT. -> PT
    val = re.sub(
        r"\bPT\.\s*",
        "PT ",
        val,
        flags=re.IGNORECASE
    )

    val = _normalize_spaces(val)

    return val

# ─────────────────────────────────────────────────────────────────────────────
# PROCESS DATA
# ─────────────────────────────────────────────────────────────────────────────

def _insert_clean_columns(
    df: pd.DataFrame,
    all_lists: list,
    prefix: str,
    max_cols: int,
) -> list:
    """Buat kolom clean_{prefix}_N dan kembalikan nama-nama kolom baru."""
    added = []
    for i in range(1, max_cols + 1):
        col_name = f"clean {prefix} {i}"
        df[col_name] = [lst[i - 1] if i - 1 < len(lst) else None for lst in all_lists]
        added.append(col_name)
    return added


def _fast_read_excel(path: str, sheet_name=None, header: int = 0) -> pd.DataFrame:
    """Baca file Excel besar lewat openpyxl dan tangani nama kolom duplikat."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    for _ in range(header):
        next(rows_iter)
    raw_cols = next(rows_iter)
    
    # Penanganan kolom duplikat (COMP_NAME kedua otomatis jadi COMP_NAME.1)
    cols = []
    counts = {}
    for c in raw_cols:
        if c in counts:
            counts[c] += 1
            cols.append(f"{c}.{counts[c]}")
        else:
            counts[c] = 0
            cols.append(c)

    data = list(rows_iter)
    wb.close()
    return pd.DataFrame(data, columns=cols)

def _fast_write_excel(df: pd.DataFrame, path: str) -> None:
    """Tulis DataFrame besar ke xlsx lebih cepat lewat openpyxl write_only mode."""
    import openpyxl
    os.makedirs(os.path.dirname(path), exist_ok=True)

    df2 = df.astype(object).where(pd.notnull(df), None)

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    ws.append(list(df2.columns))
    for row in df2.itertuples(index=False, name=None):
        ws.append(row)
    wb.save(path)


def process_data(input_file: str, output_file: str) -> None:
    print(f"[1/5] Membaca data dari: {input_file} ...")
    df = _fast_read_excel(input_file, header=0)
    print(f"      Total baris keseluruhan: {len(df):,}")

    if CEDANT_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{CEDANT_COL}' tidak ditemukan!")
        print(f"        Kolom tersedia: {list(df.columns)}")
        return

    required_cols = [POLIS_COL, SLIP_COL, INSURED_COL, BROKER_COL]
    for col in required_cols:
        if col not in df.columns:
            print(f"\n[ERROR] Kolom '{col}' tidak ditemukan di file!")
            print(f"        Kolom tersedia: {list(df.columns)}")
            return

    df[CEDANT_COL] = df[CEDANT_COL].astype(str).str.strip()
    is_wahana = df[CEDANT_COL] == CEDANT_VALUE
    print(
        f"[2/5] Filter cedant '{CEDANT_VALUE}': "
        f"{is_wahana.sum():,} baris WAHANA dari total {len(df):,} baris."
    )

    # Output cleaning hanya untuk Cedant WAHANA.
    df = df.loc[is_wahana].copy()
    is_wahana = pd.Series(True, index=df.index)
    print(f"      Output cleaning FACUL dibatasi ke WAHANA: {len(df):,} baris.")

    # Hitung mitra_bisnis sebelum rename.
    broker_s = df[BROKER_COL].fillna("").astype(str).str.strip()
    cedant_s = df[CEDANT_COL].fillna("").astype(str).str.strip()
    use_cedant = (broker_s == "") | (broker_s.str.upper() == "DIRECT")
    mitra_values = [
        clean_business_partners(x)
        for x in np.where(use_cedant, cedant_s, broker_s)
    ]

    df.rename(
        columns={
            POLIS_COL: "polis_ori",
            SLIP_COL: "slip_ori",
            INSURED_COL: "insured_ori",
        },
        inplace=True,
    )

    certificate_values = [
        clean_certificate(x)
        for x in df["polis_ori"]
    ]

    print("[3/5] Menjalankan proses cleaning hanya untuk baris WAHANA ...")

    n = len(df)
    all_clean_polis = [[] for _ in range(n)]
    all_clean_slip = [[] for _ in range(n)]
    all_clean_ins = [[] for _ in range(n)]

    wahana_idx = np.flatnonzero(is_wahana.to_numpy())
    polis_vals = df["polis_ori"].to_numpy()
    slip_vals = df["slip_ori"].to_numpy()
    insured_vals = df["insured_ori"].to_numpy()

    max_polis = max_slip = max_ins = 1
    total_w = len(wahana_idx)

    for n_done, pos in enumerate(wahana_idx, 1):
        if n_done % 5_000 == 0:
            print(f"      Progress: {n_done:,} / {total_w:,} baris WAHANA diproses...")

        c_polis = clean_polis(polis_vals[pos])
        c_slip = clean_slip(slip_vals[pos])
        c_ins = clean_insured(insured_vals[pos])

        max_polis = max(max_polis, len(c_polis))
        max_slip = max(max_slip, len(c_slip))
        max_ins = max(max_ins, len(c_ins))

        all_clean_polis[pos] = c_polis
        all_clean_slip[pos] = c_slip
        all_clean_ins[pos] = c_ins

    print("      Selesai diproses!")
    print(f"      -> Jumlah kolom clean polis  : {max_polis}")
    print(f"      -> Jumlah kolom clean slip   : {max_slip}")
    print(f"      -> Jumlah kolom clean insured: {max_ins}")

    print("[4/5] Menyusun kolom output sesuai template Excel ...")

    output_columns = [
        "FAC_CODE", "FAC_CEDANT", "COMP_NAME", "FAC_BROKER",
        "COMP_NAME2", "BUSSINESS_PATNER", "FAC_INSURED",
        "INSURED_CLEAN_1", "INSURED_CLEAN_2", "INSURED_CLEAN_3",
        "INSURED_CLEAN_4", "INSURED_CLEAN_5", "FAC_POLICY_NO",
        "POLICY_CLEAN_1", "SERTIF_CLEAN_1", "POLICY_CLEAN_2",
        "SERTIF_CLEAN_2", "POLICY_CLEAN_3", "SERTIF_CLEAN_3",
        "POLICY_CLEAN_4", "POLICY_CLEAN_5", "FAC_SLIP",
        "SLIP_CLEAN_1", "SLIP_CLEAN_2", "SLIP_CLEAN_3",
        "SLIP_CLEAN_4", "SLIP_CLEAN_5", "FAC_CURRENCY",
        "FAC_INP_DATE", "FAC_COM_DATE", "FAC_EXP_DATE",
        "FAC_SUB_CLASS", "FAC_RISK", "FAC_DESC", "FAC_ACC_STS",
        "FAC_STS_SLIP",
    ]

    def source_series(column_name, fallback=""):
        if column_name in df.columns:
            return df[column_name]
        return pd.Series([fallback] * len(df), index=df.index, dtype=object)

    output_df = pd.DataFrame(index=df.index)

    # Semua kolom template: jika sumber tidak tersedia, blank.
    for col in output_columns:
        output_df[col] = source_series(col)

    # Alias raw internal -> nama output template.
    output_df["FAC_INSURED"] = source_series("insured_ori")
    output_df["FAC_POLICY_NO"] = source_series("polis_ori")
    output_df["FAC_SLIP"] = source_series("slip_ori")

    # FAC_BROKER pada template dapat diisi dari source FAC_BROKER;
    # jika source tersebut tidak ada, gunakan kolom broker yang memang
    # dipakai script saat ini (COMP_NAME.1).
    if "FAC_BROKER" not in df.columns and BROKER_COL in df.columns:
        output_df["FAC_BROKER"] = df[BROKER_COL]

    # BUSINESS PARTNER menggunakan hasil cleaning yang sudah ada.
    output_df["BUSSINESS_PATNER"] = mitra_values

    # Clean insured.
    for i in range(1, 6):
        output_df[f"INSURED_CLEAN_{i}"] = [
            vals[i - 1] if i - 1 < len(vals) else ""
            for vals in all_clean_ins
        ]

    # Clean policy.
    for i in range(1, 6):
        output_df[f"POLICY_CLEAN_{i}"] = [
            vals[i - 1] if i - 1 < len(vals) else ""
            for vals in all_clean_polis
        ]

    # Certificate pada template FACUL hanya menyediakan SERTIF_CLEAN_1..3.
    for i in range(1, 4):
        output_df[f"SERTIF_CLEAN_{i}"] = [
            (
                cert[i - 1]
                if isinstance(cert, list) and i - 1 < len(cert)
                else (cert if i == 1 and cert else "")
            )
            for cert in certificate_values
        ]

    # Clean slip.
    for i in range(1, 6):
        output_df[f"SLIP_CLEAN_{i}"] = [
            vals[i - 1] if i - 1 < len(vals) else ""
            for vals in all_clean_slip
        ]

    output_df = output_df[output_columns]

    print(f"[5/5] Menyimpan hasil ke: {output_file} ...")
    _fast_write_excel(output_df, output_file)

    print(f"\n{'=' * 55}")
    print("  [OK] Selesai!")
    print(f"  Total baris output  : {len(output_df):,}")
    print(f"  Total kolom output  : {len(output_df.columns)}")
    print(f"  File disimpan di    : {output_file}")
    print(f"{'=' * 55}")

# ─────────────────────────────────────────────────────────────────────────────
# ENTRY POINT
# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    process_data(INPUT_FILE, OUTPUT_FILE)

# ─────────────────────────────────────────────────────────────────────────────
# MATCHING WAHANA
# ─────────────────────────────────────────────────────────────────────────────
import os
import re
from collections import defaultdict

import pandas as pd


# ============================================================
# CONFIG
# ============================================================

OUTPUT_DIR = "output"

SUSPEND_FILE = os.path.join(
    OUTPUT_DIR,
    "rev_wahana_output_suspend.xlsx"
)

OSBAL_FILE = os.path.join(
    OUTPUT_DIR,
    "wahana_output_osbalnew.xlsx"
)

FACUL_FILE = os.path.join(
    OUTPUT_DIR,
    "wahana_output_facul.xlsx"
)

# DATA 5: Database Rekap RI Slip
DATA5_FILE = os.path.join(
    "input",
    "5. Database Rekap RI Slip.xlsx"
)

OUTPUT_FILE = os.path.join(
    OUTPUT_DIR,
    "matching_wahana_output.xlsx"
)


# ============================================================
# SUSPEND COLUMNS
# ============================================================

SUSPEND_SLIP_COL = "polis_ori"
SUSPEND_POLIS_COL = "slip_ori"
SUSPEND_INSURED_COL = "insured_ori"

SUSPEND_CURR_COL = "CURR ORI"
SUSPEND_AMOUNT_COL = "AMOUNT ORI"
SUSPEND_DATE_COL = "RECEIPT DATE"


# ============================================================
# OSBAL COLUMNS
# ============================================================

OSBAL_FACODE_COL = "CCOS_REF_CODE"
OSBAL_POLIS_COL = "FAC_POLICY_NO"
OSBAL_SLIP_COL = "FAC_SLIP"
OSBAL_INSURED_COL = "FAC_INSURED"
OSBAL_CURR_COL = "CCOS_CURR"

# NAMA YANG BENAR
OSBAL_DATE_COL = "FAC_COM_DATE"

OSBAL_BAL_DUE_COL = "CCOS_BAL_DUE"
OSBAL_OR_BAL_COL = "CCOS_OR_BAL"
OSBAL_DOC_COL = "CCOS_DOC_NO"


# ============================================================
# FACUL COLUMNS
# ============================================================

FACUL_FACODE_COL = "FAC_CODE"
FACUL_POLIS_COL = "FAC_POLICY_NO"
FACUL_SLIP_COL = "FAC_SLIP"
FACUL_INSURED_COL = "FAC_INSURED"


# ============================================================
# SETTINGS
# ============================================================

MAX_SPLIT_COLS = 5

# Toleransi perbandingan amount
AMOUNT_TOLERANCE = 0.01


# ============================================================
# GENERAL HELPERS
# ============================================================

def normalize_text(value):
    """
    Normalisasi teks untuk kebutuhan matching.

    Tidak mengubah data asli dataframe.
    """
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    text = str(value).strip().upper()

    # Normalisasi whitespace
    text = re.sub(r"\s+", " ", text)

    return text


def normalize_code(value):
    """
    Normalisasi FAC_CODE / kode.
    """
    text = normalize_text(value)

    # Jika Excel membaca angka sebagai 123.0
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]

    return text


def clean_values(values):
    """
    Ambil nilai non-blank dari beberapa clean column.
    """
    result = []

    for value in values:
        value = normalize_text(value)

        if value and value not in result:
            result.append(value)

    return result


def get_clean_cols(df, prefix):
    """
    Cari:
        clean slip 1
        clean slip 2
        ...
        clean slip 5

    atau:
        clean polis 1
        ...
    """

    pattern = re.compile(
        rf"^clean {re.escape(prefix)} (\d+)$",
        re.IGNORECASE
    )

    template_patterns = {
        "polis": re.compile(r"^POLICY_CLEAN_(\d+)$", re.IGNORECASE),
        "slip": re.compile(r"^SLIP_CLEAN_(\d+)$", re.IGNORECASE),
        "insured": re.compile(r"^INSURED_CLEAN_(\d+)$", re.IGNORECASE),
    }

    found = []

    for col in df.columns:
        col_text = str(col)
        match = pattern.match(col_text)

        if match:
            number = int(match.group(1))
            if 1 <= number <= MAX_SPLIT_COLS:
                found.append((number, col))
                continue

        template_pattern = template_patterns.get(prefix.lower())
        if template_pattern:
            match = template_pattern.match(col_text)
            if match:
                number = int(match.group(1))
                if 1 <= number <= MAX_SPLIT_COLS:
                    found.append((number, col))

    found.sort(key=lambda x: x[0])

    return [col for _, col in found]


def require_column(df, column_name, df_name):
    """
    Pastikan kolom wajib tersedia.
    """
    if column_name not in df.columns:
        raise KeyError(
            f"\nKolom '{column_name}' tidak ditemukan di {df_name}.\n"
            f"Kolom yang tersedia:\n"
            f"{list(df.columns)}"
        )


def require_any_clean_columns(df, prefix, df_name):
    cols = get_clean_cols(df, prefix)

    if not cols:
        raise KeyError(
            f"\nTidak ditemukan kolom clean {prefix} 1..5 "
            f"di {df_name}."
        )

    return cols


# ============================================================
# NUMBER PARSER
# ============================================================

def to_number(value):
    """
    Konversi amount menjadi float.

    Menangani format:
        123456
        123,456
        123.456
        123.456,78
        123,456.78
        (123456)
    """

    if value is None:
        return None

    if isinstance(value, (int, float)):
        if pd.isna(value):
            return None

        return float(value)

    text = str(value).strip()

    if not text:
        return None

    # Negative accounting format
    negative = False

    if text.startswith("(") and text.endswith(")"):
        negative = True
        text = text[1:-1]

    # Buang karakter selain angka separator
    text = re.sub(r"[^\d,.\-]", "", text)

    if not text:
        return None

    try:
        if "," in text and "." in text:

            # Separator terakhir dianggap decimal
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "")
                text = text.replace(",", ".")
            else:
                text = text.replace(",", "")

        elif "," in text:

            parts = text.split(",")

            # 123,45 -> decimal
            if len(parts[-1]) == 2:
                text = text.replace(",", ".")
            else:
                text = text.replace(",", "")

        elif "." in text:

            parts = text.split(".")

            # Jika lebih dari satu titik,
            # diasumsikan thousands separator
            if len(parts) > 2:
                text = text.replace(".", "")

        result = float(text)

        if negative:
            result = -result

        return result

    except Exception:
        return None


# ============================================================
# DATE HELPERS
# ============================================================

def to_date(value):
    """
    Konversi tanggal tanpa mengubah dataframe asli.
    """

    if value is None:
        return pd.NaT

    try:
        return pd.to_datetime(value, errors="coerce")
    except Exception:
        return pd.NaT


# ============================================================
# LINE SLIP
# ============================================================

LINE_SLIP_PATTERN = re.compile(
    r"\bLINE\s*SLIP\b|\bLINESLIP\b",
    re.IGNORECASE
)


def is_line_slip(raw_slip, suspend_clean_slips):
    """
    Contoh:

    Suspend:
        AF2-2025-FR

    OSBAL:
        LINE SLIP ACA FIRE SEPTEMBER 2025 IDR AF2-2025-FR

    Maka dianggap Line Slip.
    """

    raw = normalize_text(raw_slip)

    if not raw:
        return False

    if not LINE_SLIP_PATTERN.search(raw):
        return False

    for slip in suspend_clean_slips:
        slip = normalize_text(slip)

        if slip and slip in raw:
            return True

    return False


def check_period(
    receipt_date,
    fac_com_date,
    raw_osbal_slip,
    suspend_clean_slips
):
    """
    Aturan periode:

    Normal:
        RECEIPT DATE > FAC_COM_DATE

    Line Slip:
        RECEIPT DATE >= FAC_COM_DATE + 1 bulan
    """

    receipt = to_date(receipt_date)
    fac_date = to_date(fac_com_date)

    if pd.isna(receipt) or pd.isna(fac_date):
        return None, "Tanggal Tidak Lengkap"

    line_slip = is_line_slip(
        raw_osbal_slip,
        suspend_clean_slips
    )

    if line_slip:
        minimum_date = fac_date + pd.DateOffset(months=1)

        if receipt >= minimum_date:
            return True, "Line Slip - Valid"

        return False, "Beda Periode"

    # Aturan normal: STRICT >
    if receipt > fac_date:
        return True, "Valid"

    return False, "Beda Periode"


# ============================================================
# INDEX BUILDING
# ============================================================

def build_clean_index(df, clean_cols):
    """
    Membuat index:

        nilai clean -> row index dataframe

    Contoh:
        AF2-2025-FR -> [12, 35, 88]
    """

    index = defaultdict(list)

    if not clean_cols:
        return index

    values_matrix = (
        df[clean_cols]
        .fillna("")
        .to_numpy(dtype=object)
    )

    for row_idx, row_values in enumerate(values_matrix):

        seen = set()

        for value in row_values:

            value = normalize_text(value)

            if not value:
                continue

            if value in seen:
                continue

            seen.add(value)

            index[value].append(row_idx)

    return index


def build_facul_code_index(df, clean_cols, facode_col):
    """
    Membuat:

        clean value -> FAC_CODE

    dari FACUL.
    """

    index = defaultdict(set)

    values_matrix = (
        df[clean_cols]
        .fillna("")
        .to_numpy(dtype=object)
    )

    facode_values = (
        df[facode_col]
        .fillna("")
        .to_numpy(dtype=object)
    )

    for row_idx, row_values in enumerate(values_matrix):

        fac_code = normalize_code(
            facode_values[row_idx]
        )

        if not fac_code:
            continue

        seen = set()

        for value in row_values:

            value = normalize_text(value)

            if not value:
                continue

            if value in seen:
                continue

            seen.add(value)

            index[value].add(fac_code)

    return index


def build_osbal_facode_index(df, facode_col):
    """
    FAC_CODE -> row OSBAL.
    """

    index = defaultdict(list)

    values = (
        df[facode_col]
        .fillna("")
        .to_numpy(dtype=object)
    )

    for row_idx, value in enumerate(values):

        fac_code = normalize_code(value)

        if fac_code:
            index[fac_code].append(row_idx)

    return index


# ============================================================
# INSURED LIKE / SUBSTRING MATCH
# ============================================================

# Kata yang terlalu generik untuk dijadikan satu-satunya dasar LIKE.
# Tujuannya mencegah kasus seperti "BANK" mem-match banyak perusahaan.
INSURED_GENERIC_TOKENS = {
    "BANK", "PT", "CV", "TBK", "LTD", "LIMITED", "INC", "CORP",
    "CORPORATION", "COMPANY", "CO", "GROUP", "HOLDING", "INDONESIA",
    "INDONESIAN", "INDUSTRY", "INDUSTRIES", "INTERNATIONAL", "PERSERO",
    "PERSEROAN", "TERBATAS", "FINANCE", "FINANCIAL", "ASURANSI"
}


def _insured_tokens(value):
    text = normalize_text(value)
    if not text:
        return []
    return [t for t in re.findall(r"[A-Z0-9]+", text) if t]


def insured_like_match(query, reference):
    """LIKE-like matching untuk INSURED.

    Prioritas:
      1. exact normalized match
      2. full-string substring bila sisi yang dipakai >= 10 karakter
      3. shared meaningful token >= 6 karakter, bukan generic token

    Shared token dipakai untuk kasus group/subsidiary seperti
    "INDORAMA POLYCHEM INDONESIA" vs "INDORAMA GROUP", tetapi
    token umum seperti "BANK" tidak pernah cukup untuk match.
    """
    q = normalize_text(query)
    r = normalize_text(reference)
    if not q or not r:
        return False
    if q == r:
        return True

    if (len(q) >= 10 and q in r) or (len(r) >= 10 and r in q):
        return True

    qt = {t for t in _insured_tokens(q) if len(t) >= 6 and t not in INSURED_GENERIC_TOKENS}
    rt = {t for t in _insured_tokens(r) if len(t) >= 6 and t not in INSURED_GENERIC_TOKENS}
    return bool(qt & rt)


def insured_like_rows(values, df, clean_cols):
    """Return row indices yang cocok dengan INSURED exact/LIKE."""
    queries = clean_values(values)
    if not queries or not clean_cols:
        return []

    matrix = df[clean_cols].fillna("").to_numpy(dtype=object)
    rows = []
    for row_idx, row_values in enumerate(matrix):
        refs = clean_values(row_values)
        if any(insured_like_match(q, r) for q in queries for r in refs):
            rows.append(row_idx)
    return rows


# ============================================================
# CANDIDATE SEARCH
# ============================================================

def candidate_rows_for_values(
    values,
    osbal_index,
    facul_index,
    osbal_by_facode
):
    """
    Cari kandidat OSBAL dari:

    1. direct clean value OSBAL
    2. clean value FACUL -> FAC_CODE -> OSBAL
    """

    candidates = set()

    fac_codes_from_facul = set()

    for value in values:

        value = normalize_text(value)

        if not value:
            continue

        # ----------------------------------------------------
        # Direct OSBAL
        # ----------------------------------------------------

        for row_idx in osbal_index.get(value, []):
            candidates.add(row_idx)

        # ----------------------------------------------------
        # FACUL -> FAC_CODE
        # ----------------------------------------------------

        for fac_code in facul_index.get(value, set()):
            fac_codes_from_facul.add(fac_code)

    # --------------------------------------------------------
    # FAC_CODE -> OSBAL
    # --------------------------------------------------------

    for fac_code in fac_codes_from_facul:

        for row_idx in osbal_by_facode.get(
            fac_code,
            []
        ):
            candidates.add(row_idx)

    return sorted(candidates)


# ============================================================
# CURRENCY + PERIOD FILTER
# ============================================================

def filter_candidates(
    candidates,
    suspend_curr,
    suspend_date,
    suspend_clean_slips,
    osbal_df
):
    """
    Filter berdasarkan:

    1. Currency
    2. Period

    Currency harus sama.

    Period:
        RECEIPT DATE > FAC_COM_DATE

    Line Slip:
        RECEIPT DATE >= FAC_COM_DATE + 1 bulan
    """

    if not candidates:
        return [], "Tidak Ada Kandidat", "Tidak Dicek"

    # --------------------------------------------------------
    # CURRENCY
    # --------------------------------------------------------

    suspend_currency = normalize_text(
        suspend_curr
    )

    currency_candidates = []

    if suspend_currency:

        for row_idx in candidates:

            osbal_currency = normalize_text(
                osbal_df.iloc[row_idx][OSBAL_CURR_COL]
            )

            if (
                osbal_currency
                and osbal_currency == suspend_currency
            ):
                currency_candidates.append(row_idx)

        if not currency_candidates:
            return [], "Beda Currency", "Tidak Dicek"

        currency_status = "Valid"

    else:
        # Jika currency Suspend kosong,
        # currency tidak bisa digunakan sebagai filter.
        currency_candidates = list(candidates)
        currency_status = "Currency Suspend Kosong"

    # --------------------------------------------------------
    # PERIOD
    # --------------------------------------------------------

    period_candidates = []

    period_unknown = []

    for row_idx in currency_candidates:

        row = osbal_df.iloc[row_idx]

        ok, period_status = check_period(
            suspend_date,
            row[OSBAL_DATE_COL],
            row[OSBAL_SLIP_COL],
            suspend_clean_slips
        )

        if ok is True:
            period_candidates.append(row_idx)

        elif ok is None:
            period_unknown.append(row_idx)

    if period_candidates:
        return (
            period_candidates,
            currency_status,
            "Valid"
        )

    # Kalau tanggal tidak lengkap, jangan langsung
    # membuang kandidat.
    if period_unknown:
        return (
            period_unknown,
            currency_status,
            "Tanggal Tidak Lengkap"
        )

    return (
        [],
        currency_status,
        "Beda Periode"
    )



# ============================================================
# DATA 5 / AUXILIARY SOURCE HELPERS
# ============================================================

def _normalize_header_name(value):
    """Normalisasi nama header agar spasi/koma/tanda baca ringan tidak menghalangi matching."""
    text = str(value).strip().upper()
    text = re.sub(r"[\s,;:]+", " ", text)
    return text.strip()


def _find_first_existing_col(df, aliases):
    """Cari nama kolom pertama yang tersedia (case-insensitive, toleran spasi/koma)."""
    lookup = {_normalize_header_name(c): c for c in df.columns}
    for alias in aliases:
        col = lookup.get(_normalize_header_name(alias))
        if col is not None:
            return col
    return None


def _find_clean_or_raw_cols(df, prefixes, raw_aliases):
    """Cari kolom clean 1..5; bila tidak ada, fallback ke satu kolom raw.

    Untuk Data 5, header kadang memiliki koma/titik/spasi tambahan.
    Pencarian clean/raw dibuat toleran, tetapi nama kolom asli tetap
    dikembalikan agar isi DataFrame tidak berubah.
    """
    found = []
    for prefix in prefixes:
        prefix_norm = _normalize_header_name(prefix)
        for col in df.columns:
            col_norm = _normalize_header_name(col)
            m = re.fullmatch(
                rf"clean\s+{re.escape(prefix_norm)}\s+(\d+)",
                col_norm,
                re.IGNORECASE
            )
            if m and 1 <= int(m.group(1)) <= MAX_SPLIT_COLS:
                found.append((int(m.group(1)), col))
        if found:
            return [c for _, c in sorted(found)]
    raw = _find_first_existing_col(df, raw_aliases)
    return [raw] if raw else []


def _read_data5_with_header_detection(path):
    """Baca Data 5 dan otomatis mencari baris header bila header bukan row 1."""
    df = pd.read_excel(path, header=0, dtype=object)

    aliases = {
        "SLIP", "SLIP NO", "SLIP_NO", "FAC SLIP",
        "POLIS", "POLICY", "POLICY NO", "FAC POLICY NO",
        "INSURED", "INSURED NAME", "FAC INSURED",
        "FAC CODE", "FAC. CODE", "FAC_CODE"
    }

    def score(columns):
        normalized = {_normalize_header_name(c) for c in columns}
        return sum(1 for a in aliases if _normalize_header_name(a) in normalized)

    if score(df.columns) >= 2:
        return df

    # Fallback: cari header pada beberapa baris pertama workbook.
    preview = pd.read_excel(path, header=None, nrows=12, dtype=object)
    best_row = None
    best_score = score(df.columns)
    for row_idx in range(len(preview)):
        row_score = score(preview.iloc[row_idx].tolist())
        if row_score > best_score:
            best_score = row_score
            best_row = row_idx

    if best_row is not None and best_score >= 2:
        return pd.read_excel(path, header=best_row, dtype=object)

    return df


def build_aux_facode_index(df, value_cols, facode_col, osbal_df, osbal_indexes):
    """
    Data 5/Data 1 -> clean value -> FAC_CODE.

    Jika source memiliki FAC_CODE, gunakan langsung.
    Jika source tidak memiliki FAC_CODE, FAC_CODE diturunkan dari OSBAL
    melalui nilai matching yang sama. Ini penting karena Data 5 dan Data 1
    tidak selengkap Data 2.
    """
    index = defaultdict(set)

    if not value_cols:
        return index

    # FAC_CODE langsung dari source bila tersedia
    direct_codes = None
    if facode_col and facode_col in df.columns:
        direct_codes = df[facode_col].fillna("").to_numpy(dtype=object)

    matrix = df[value_cols].fillna("").to_numpy(dtype=object)

    for row_idx, row_values in enumerate(matrix):
        row_codes = set()

        if direct_codes is not None:
            fc = normalize_code(direct_codes[row_idx])
            if fc:
                row_codes.add(fc)

        # Jika FAC_CODE source kosong/tidak ada, turunkan dari OSBAL.
        if not row_codes:
            for value in row_values:
                value = normalize_text(value)
                if not value:
                    continue
                for osbal_idx in (
                    osbal_indexes.get(value, [])
                ):
                    fc = normalize_code(
                        osbal_df.iloc[osbal_idx][OSBAL_FACODE_COL]
                    )
                    if fc:
                        row_codes.add(fc)

        for value in row_values:
            value = normalize_text(value)
            if value:
                index[value].update(row_codes)

    return index


def build_aux_insured_like_index(df, value_cols, osbal_df, osbal_insured_clean):
    """Source auxiliary -> LIKE INSURED -> FAC_CODE OSBAL."""
    index = defaultdict(set)
    if not value_cols:
        return index

    token_rows = defaultdict(set)
    osbal_matrix = osbal_df[osbal_insured_clean].fillna("").to_numpy(dtype=object)
    for os_idx, ref_values in enumerate(osbal_matrix):
        for ref in clean_values(ref_values):
            for tok in _insured_tokens(ref):
                if len(tok) >= 6 and tok not in INSURED_GENERIC_TOKENS:
                    token_rows[tok].add(os_idx)

    matrix = df[value_cols].fillna("").to_numpy(dtype=object)
    for row_values in matrix:
        for q in clean_values(row_values):
            candidate_rows = set()
            for tok in _insured_tokens(q):
                if len(tok) >= 6 and tok not in INSURED_GENERIC_TOKENS:
                    candidate_rows.update(token_rows.get(tok, set()))
            for os_idx in candidate_rows:
                if any(insured_like_match(q, r) for r in clean_values(osbal_matrix[os_idx])):
                    fc = normalize_code(osbal_df.iloc[os_idx][OSBAL_FACODE_COL])
                    if fc:
                        index[q].add(fc)
    return index


def candidate_rows_from_aux_values(values, aux_index, osbal_by_facode):
    """Aux source value -> FAC_CODE -> seluruh row OSBAL."""
    fac_codes = set()
    for value in values:
        value = normalize_text(value)
        if value:
            fac_codes.update(aux_index.get(value, set()))

    rows = set()
    for fc in fac_codes:
        rows.update(osbal_by_facode.get(fc, []))

    return sorted(rows), sorted(fac_codes)


def resolve_stage_candidates(
    values,
    source_index,
    osbal_by_facode
):
    return candidate_rows_from_aux_values(
        values,
        source_index,
        osbal_by_facode
    )

# ============================================================
# RESOLVE ONE SUSPEND ROW
# ============================================================

def _fac_codes_from_rows(rows, osbal_df):
    """Return unique FAC_CODE values from OSBAL row indices."""
    return sorted({
        normalize_code(osbal_df.iloc[row_idx][OSBAL_FACODE_COL])
        for row_idx in (rows or [])
        if normalize_code(osbal_df.iloc[row_idx][OSBAL_FACODE_COL])
    })


def resolve_one_suspend(
    suspend_row_idx,
    suspend_slip_values,
    suspend_polis_values,
    suspend_insured_values,
    suspend_curr,
    suspend_date,
    osbal_df,
    osbal_slip_index,
    osbal_polis_index,
    osbal_insured_index,
    data5_slip_index,
    data5_polis_index,
    data5_insured_index,
    facul_slip_index,
    facul_polis_index,
    facul_insured_index,
    osbal_by_facode,
    osbal_insured_like_index=None,
    data5_insured_like_index=None,
    facul_insured_like_index=None
):
    """
    Final matching cascade.

    Source cascade:
        Data 2 / OSBAL -> Data 5 -> Data 1 / FACUL

    Key cascade:
        SLIP -> POLIS -> INSURED

    BUSINESS RULE:
        - Data 2 adalah sumber kandidat FAC_CODE awal.
        - Data 5 dan Data 1 tidak menambah FAC_CODE baru.
        - Untuk key yang sama, kandidat dari Data 5/Data 1 hanya dipakai
          jika FAC_CODE-nya overlap dengan kandidat yang sudah diperoleh
          dari source sebelumnya.
        - Jika suatu source tidak mempunyai FAC_CODE yang overlap, kandidat
          sebelumnya dipertahankan dan proses diteruskan ke source berikutnya.
        - Setelah seluruh source untuk key tersebut selesai, BARU lakukan
          narrowing Currency -> Period terhadap kandidat FAC_CODE terakhir.
        - Jika hasil narrowing masih >1 FAC_CODE -> final
          "Matching >1 fac code".
        - Jika narrowing menggugurkan seluruh candidate karena currency atau
          period, candidate tetap dianggap ditemukan pada key tersebut;
          statusnya Beda Currency/Beda Periode.
        - Jika Data 2 benar-benar tidak mempunyai candidate pada key tersebut,
          baru lanjut ke key berikutnya.

    Data 5 dan Data 1 tetap memakai OSBAL sebagai referensi FAC_CODE dan
    detail narrowing/financial karena source tersebut tidak selengkap OSBAL.
    """

    source_defs = [
        (
            "Data 2",
            osbal_slip_index,
            osbal_polis_index,
            osbal_insured_index,
            osbal_insured_like_index,
        ),
        (
            "Data 5",
            data5_slip_index,
            data5_polis_index,
            data5_insured_index,
            data5_insured_like_index,
        ),
        (
            "Data 1",
            facul_slip_index,
            facul_polis_index,
            facul_insured_index,
            facul_insured_like_index,
        ),
    ]

    methods = [
        ("SLIP", suspend_slip_values),
        ("POLIS", suspend_polis_values),
        ("INSURED", suspend_insured_values),
    ]

    for method, values in methods:
        if not values:
            continue

        current_candidates = None
        current_source = ""
        current_currency_status = "Tidak Dicek"
        current_period_status = "Tidak Dicek"

        # ------------------------------------------------------------
        # STEP 1: Data 2 harus menjadi sumber kandidat awal.
        # ------------------------------------------------------------
        _, slip_index, polis_index, insured_index, like_map = source_defs[0]
        index = {
            "SLIP": slip_index,
            "POLIS": polis_index,
            "INSURED": insured_index,
        }[method]

        if not index:
            continue

        current_candidates = candidate_rows_for_values(
            values,
            index,
            defaultdict(set),
            osbal_by_facode
        )

        if method == "INSURED" and like_map:
            for q in clean_values(values):
                current_candidates.extend(like_map.get(q, []))

        current_candidates = sorted(set(current_candidates or []))

        if not current_candidates:
            # Data 2 tidak menemukan key ini -> baru boleh mencoba
            # key berikutnya (SLIP -> POLIS -> INSURED).
            continue

        current_fac_codes = _fac_codes_from_rows(
            current_candidates,
            osbal_df
        )
        if not current_fac_codes:
            continue

        current_source = "Data 2"

        # ------------------------------------------------------------
        # STEP 2: Data 5 lalu Data 1 hanya mempersempit FAC_CODE
        # melalui INTERSECTION. Tidak boleh menambah FAC_CODE baru.
        # ------------------------------------------------------------
        for source_name, slip_index, polis_index, insured_index, like_map in source_defs[1:]:
            index = {
                "SLIP": slip_index,
                "POLIS": polis_index,
                "INSURED": insured_index,
            }[method]

            if not index:
                continue

            stage_candidates = resolve_stage_candidates(
                values,
                index,
                osbal_by_facode
            )[0]

            if method == "INSURED" and like_map:
                like_codes = set()
                for q in clean_values(values):
                    like_codes.update(like_map.get(q, set()))
                for fc in like_codes:
                    stage_candidates.extend(
                        osbal_by_facode.get(fc, [])
                    )

            stage_candidates = sorted(set(stage_candidates or []))
            if not stage_candidates:
                # Source ini tidak punya kandidat; jangan mengganti kandidat
                # sebelumnya. Lanjut ke source berikutnya.
                continue

            stage_fac_codes = set(
                _fac_codes_from_rows(stage_candidates, osbal_df)
            )
            overlap_fac_codes = set(current_fac_codes) & stage_fac_codes

            if not overlap_fac_codes:
                # Tidak ada FAC_CODE yang sama dengan kandidat sebelumnya.
                # FAC_CODE source ini TIDAK digunakan.
                continue

            # Hanya FAC_CODE yang overlap yang dipertahankan.
            current_candidates = [
                row_idx
                for row_idx in current_candidates
                if normalize_code(
                    osbal_df.iloc[row_idx].get(OSBAL_FACODE_COL, "")
                ) in overlap_fac_codes
            ]
            current_candidates = sorted(set(current_candidates))
            current_fac_codes = sorted(overlap_fac_codes)
            current_source = source_name

        # ------------------------------------------------------------
        # STEP 3: SETELAH seluruh cross-source filtering selesai,
        # baru lakukan GENERAL NARROWING Currency -> Period.
        # ------------------------------------------------------------
        filtered, currency_status, period_status = filter_candidates(
            candidates=current_candidates,
            suspend_curr=suspend_curr,
            suspend_date=suspend_date,
            suspend_clean_slips=suspend_slip_values,
            osbal_df=osbal_df
        )

        narrowed_fac_codes = _fac_codes_from_rows(
            filtered,
            osbal_df
        )

        # ------------------------------------------------------------
        # STEP 4: narrowing menggugurkan seluruh candidate.
        # ------------------------------------------------------------
        if not filtered:
            if currency_status == "Beda Currency":
                failure = "Beda Currency"
            elif period_status == "Beda Periode":
                failure = "Beda Periode"
            elif len(current_fac_codes) > 1:
                failure = "Matching >1 fac code"
            else:
                failure = "Unmatching"

            return {
                "source": current_source,
                "method": method,
                "rows": [],
                "candidate_rows": list(current_candidates),
                "fac_codes": sorted(current_fac_codes),
                "currency_status": currency_status,
                "period_status": period_status,
                "failure": failure,
            }

        # ------------------------------------------------------------
        # STEP 5: narrowing selesai -> tentukan FAC_CODE.
        # ------------------------------------------------------------
        if len(narrowed_fac_codes) == 1:
            return {
                "source": current_source,
                "method": method,
                "rows": list(filtered),
                "candidate_rows": list(current_candidates),
                "fac_codes": narrowed_fac_codes,
                "currency_status": currency_status,
                "period_status": period_status,
                "failure": None,
            }

        if len(narrowed_fac_codes) > 1:
            return {
                "source": current_source,
                "method": method,
                "rows": list(filtered),
                "candidate_rows": list(current_candidates),
                "fac_codes": narrowed_fac_codes,
                "currency_status": currency_status,
                "period_status": period_status,
                "failure": "Matching >1 fac code",
            }

    # Tidak ada candidate sama sekali pada Data 2 untuk seluruh key.
    return {
        "source": "",
        "method": "",
        "rows": [],
        "candidate_rows": [],
        "fac_codes": [],
        "currency_status": "Tidak Dicek",
        "period_status": "Tidak Dicek",
        "failure": "Unmatching",
    }


# ============================================================
# FINANCIAL FLAG
# ============================================================

def compare_amounts(
    amount_total,
    bal_due_total,
    suspend_count,
    osbal_count
):
    """
    Flagging:

    Adjustment total tanpa akumulasi
        amount = bal due
        1 suspend -> 1 OSBAL

    Adjustment total dengan akumulasi
        amount = bal due
        1->many / many->1 / many->many

    Adjustment sebagian tanpa akumulasi
        amount < bal due
        1 -> 1

    Adjustment sebagian dengan akumulasi
        amount < bal due
        1->many / many->1 / many->many

    New entry sebagian
        amount > bal due

    New entry total
        amount != 0 AND bal due = 0
    """

    if amount_total is None:
        return "Amount ORI tidak dapat dihitung"

    if bal_due_total is None:
        return "CCOS_BAL_DUE tidak dapat dihitung"

    try:
        amount = float(amount_total)
        bal_due = float(bal_due_total)
    except Exception:
        return "Amount tidak dapat dihitung"

    # --------------------------------------------------------
    # New Entry Total
    # --------------------------------------------------------

    if (
        abs(amount) > AMOUNT_TOLERANCE
        and abs(bal_due) <= AMOUNT_TOLERANCE
    ):
        return "New Entry total"

    # --------------------------------------------------------
    # Accumulation
    # --------------------------------------------------------

    accumulation = (
        suspend_count > 1
        or osbal_count > 1
    )

    # --------------------------------------------------------
    # Zero / zero
    # --------------------------------------------------------
    # Jika baris SUDAH MATCH (misalnya Slip only), kondisi amount=0
    # dan bal_due=0 bukan berarti Unmatching. Unmatching hanya boleh
    # dipakai ketika memang tidak ada kandidat matching.
    if (
        abs(amount) <= AMOUNT_TOLERANCE
        and abs(bal_due) <= AMOUNT_TOLERANCE
    ):
        if accumulation:
            return "Adjustment total dengan akumulasi"
        return "Adjustment total tanpa akumulasi"

    # --------------------------------------------------------
    # Equal
    # --------------------------------------------------------

    if abs(amount - bal_due) <= AMOUNT_TOLERANCE:

        if accumulation:
            return "Adjustment total dengan akumulasi"

        return "Adjustment total tanpa akumulasi"

    # --------------------------------------------------------
    # Amount < Bal Due
    # --------------------------------------------------------

    if amount < bal_due:

        if accumulation:
            return "Adjustment sebagian dengan akumulasi"

        return "Adjustment sebagian tanpa akumulasi"

    # --------------------------------------------------------
    # Amount > Bal Due
    # --------------------------------------------------------

    if amount > bal_due:
        return "New Entry sebagian"

    return "Unmatching"



# ============================================================
# OUTPUT SCENARIO HELPER
# ============================================================

def build_scenario(
    slip_values,
    polis_values,
    insured_values,
    rows,
    osbal_df,
    osbal_slip_clean,
    osbal_polis_clean,
    osbal_insured_clean,
    suffix="",
    insured_like=False,
    selected_method=""
):
    """Determine SKENARIO from the actual matching evidence.

    Currency/period differences do not replace the matching scenario;
    they are appended as a suffix. Multiple FAC_CODE is also appended
    as ``(> 1 Fac)``.
    """
    if not rows:
        base = "Unmatching"
        return f"{base}{suffix}" if suffix else base

    def row_values(row_idx, cols):
        return {
            normalize_text(osbal_df.iloc[row_idx][c])
            for c in cols
            if normalize_text(osbal_df.iloc[row_idx][c])
        }

    slip_set = set(clean_values(slip_values))
    polis_set = set(clean_values(polis_values))
    insured_set = set(clean_values(insured_values))

    slip_hit = any(slip_set & row_values(i, osbal_slip_clean) for i in rows)
    polis_hit = any(polis_set & row_values(i, osbal_polis_clean) for i in rows)
    insured_hit = any(insured_set & row_values(i, osbal_insured_clean) for i in rows)
    if not insured_hit and insured_like:
        insured_hit = any(
            insured_like_match(q, r)
            for q in insured_set
            for i in rows
            for r in row_values(i, osbal_insured_clean)
        )

    # SKENARIO mengikuti key pertama yang benar-benar menghasilkan
    # keputusan matching. Jika Slip sudah menghasilkan 1 FAC_CODE,
    # cascade berhenti sehingga tidak boleh berubah menjadi
    # "Slip only + Insured" hanya karena insured kebetulan juga cocok
    # pada row OSBAL yang sama.
    if selected_method == "SLIP":
        base = "Slip only"
    elif selected_method == "POLIS":
        base = "Polis only"
    elif selected_method == "INSURED":
        base = "Insured only"
    else:
        # Fallback untuk kasus lama/edge case ketika method tidak tersedia.
        if polis_hit and slip_hit and insured_hit:
            base = "Polis + Slip + Insured"
        elif polis_hit and slip_hit:
            base = "Slip + Polis"
        elif polis_hit and insured_hit:
            base = "Polis only + Insured"
        elif slip_hit and insured_hit:
            base = "Slip only + Insured"
        elif polis_hit:
            base = "Polis only"
        elif slip_hit:
            base = "Slip only"
        elif insured_hit:
            base = "Insured only"
        else:
            base = "Unmatching"

    return f"{base}{suffix}" if suffix else base

def safe_sum(values):
    nums = [to_number(v) for v in values]
    nums = [v for v in nums if v is not None]
    return sum(nums) if nums else None

# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 70)
    print("MATCHING WAHANA")
    print("=" * 70)

    # --------------------------------------------------------
    # Check files
    # --------------------------------------------------------

    for file_path in [
        SUSPEND_FILE,
        OSBAL_FILE,
        DATA5_FILE,
        FACUL_FILE
    ]:

        if not os.path.exists(file_path):
            raise FileNotFoundError(
                f"\nFile tidak ditemukan:\n{file_path}"
            )

    os.makedirs(
        OUTPUT_DIR,
        exist_ok=True
    )

    # --------------------------------------------------------
    # READ FILES
    # --------------------------------------------------------

    print("\n[1/7] Membaca Suspend...")

    suspend = pd.read_excel(
        SUSPEND_FILE,
        header=0,
        dtype=object
    )

    # Safety filter: matching output hanya untuk Cedant WAHANA,
    # mengikuti kolom dan nilai Cedant yang dipakai script cleaning.
    if "CEDANT NAME" in suspend.columns:
        suspend = suspend.loc[
            suspend["CEDANT NAME"]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.upper()
            .eq("PT.ASURANSI WAHANA TATA")
        ].copy()

    print(
        f"       Suspend rows : {len(suspend):,} (WAHANA only)"
    )

    print("\n[2/7] Membaca OSBAL...")

    osbal = pd.read_excel(
        OSBAL_FILE,
        header=0,
        dtype=object
    )

    if "CCOS_COMP_NAME" in osbal.columns:
        osbal = osbal.loc[
            osbal["CCOS_COMP_NAME"]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.upper()
            .eq("PT.ASURANSI WAHANA TATA")
        ].copy()

    print(
        f"       OSBAL rows   : {len(osbal):,} (WAHANA only)"
    )

    print("\n[3/7] Membaca Database Rekap RI Slip (Data 5)...")

    data5 = _read_data5_with_header_detection(DATA5_FILE)

    print(
        f"       DATA 5 rows  : {len(data5):,}"
    )

    print("\n[4/7] Membaca FACUL...")

    facul = pd.read_excel(
        FACUL_FILE,
        header=0,
        dtype=object
    )

    if "COMP_NAME" in facul.columns:
        facul = facul.loc[
            facul["COMP_NAME"]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.upper()
            .eq("PT.ASURANSI WAHANA TATA")
        ].copy()

    print(
        f"       FACUL rows   : {len(facul):,} (WAHANA only)"
    )

    # --------------------------------------------------------
    # VALIDATE REQUIRED COLUMNS
    # --------------------------------------------------------

    print("\n[5/7] Validasi nama kolom...")

    # Suspend
    for col in [
        SUSPEND_SLIP_COL,
        SUSPEND_POLIS_COL,
        SUSPEND_INSURED_COL,
        SUSPEND_CURR_COL,
        SUSPEND_AMOUNT_COL,
        SUSPEND_DATE_COL
    ]:
        require_column(
            suspend,
            col,
            "SUSPEND"
        )

    # OSBAL
    for col in [
        OSBAL_FACODE_COL,
        OSBAL_POLIS_COL,
        OSBAL_SLIP_COL,
        OSBAL_INSURED_COL,
        OSBAL_CURR_COL,
        OSBAL_DATE_COL,
        OSBAL_BAL_DUE_COL,
        OSBAL_OR_BAL_COL,
        OSBAL_DOC_COL
    ]:
        require_column(
            osbal,
            col,
            "OSBAL"
        )

    # FACUL/Data 1 tidak wajib memiliki FAC_CODE.
    # Jika FAC_CODE tidak ada, nanti diturunkan dari Data 2 (OSBAL)
    # berdasarkan key matching yang sama.
    for col in [
        FACUL_POLIS_COL,
        FACUL_SLIP_COL,
        FACUL_INSURED_COL
    ]:
        require_column(
            facul,
            col,
            "FACUL"
        )

    # --------------------------------------------------------
    # CLEAN COLUMN DISCOVERY
    # --------------------------------------------------------

    # PENTING: pada file SUSPEND, isi kolom SLIP dan POLIS terbalik.
    # Nama kolom Excel TIDAK diubah. Backend hanya membalik interpretasinya:
    #   polis_ori + clean polis 1..5 -> dipakai sebagai SLIP
    #   slip_ori  + clean slip 1..5  -> dipakai sebagai POLIS
    suspend_slip_clean = require_any_clean_columns(
        suspend,
        "polis",
        "SUSPEND"
    )

    suspend_polis_clean = require_any_clean_columns(
        suspend,
        "slip",
        "SUSPEND"
    )

    suspend_insured_clean = require_any_clean_columns(
        suspend,
        "insured",
        "SUSPEND"
    )

    osbal_slip_clean = require_any_clean_columns(
        osbal,
        "slip",
        "OSBAL"
    )

    osbal_polis_clean = require_any_clean_columns(
        osbal,
        "polis",
        "OSBAL"
    )

    osbal_insured_clean = require_any_clean_columns(
        osbal,
        "insured",
        "OSBAL"
    )

    # DATA 5 tidak diasumsikan selengkap Data 2. Cari kolom clean/raw
    # yang tersedia; yang tidak ada cukup dilewati.
    data5_slip_clean = _find_clean_or_raw_cols(
        data5,
        ["slip"],
        ["FAC_SLIP", "SLIP", "SLIP NO", "SLIP_NO", "slip_ori"]
    )
    data5_polis_clean = _find_clean_or_raw_cols(
        data5,
        ["polis", "policy"],
        ["FAC_POLICY_NO", "POLIS", "POLICY", "POLICY NO", "POLIS_NO", "polis_ori"]
    )
    data5_insured_clean = _find_clean_or_raw_cols(
        data5,
        ["insured"],
        ["FAC_INSURED", "INSURED", "INSURED NAME", "insured_ori"]
    )

    if not (data5_slip_clean or data5_polis_clean or data5_insured_clean):
        raise KeyError(
            "Data 5 tidak memiliki kolom matching SLIP/POLIS/INSURED "
            f"yang dapat digunakan. Kolom yang terbaca: {list(data5.columns)}"
        )

    data5_facode_col = _find_first_existing_col(
        data5,
        ["FAC_CODE", "FAC CODE", "FAC. CODE", "FACCODE", "CCOS_REF_CODE", "REF_CODE"]
    )

    facul_slip_clean = require_any_clean_columns(
        facul,
        "slip",
        "FACUL"
    )

    facul_polis_clean = require_any_clean_columns(
        facul,
        "polis",
        "FACUL"
    )

    facul_insured_clean = require_any_clean_columns(
        facul,
        "insured",
        "FACUL"
    )

    print(
        f"       Suspend clean slip   : {suspend_slip_clean}"
    )

    print(
        f"       Suspend clean polis  : {suspend_polis_clean}"
    )

    print(
        f"       Suspend clean insured: {suspend_insured_clean}"
    )

    print(
        f"       OSBAL clean slip     : {osbal_slip_clean}"
    )

    print(
        f"       OSBAL clean polis    : {osbal_polis_clean}"
    )

    print(
        f"       OSBAL clean insured  : {osbal_insured_clean}"
    )

    print(
        f"       DATA 5 matching slip   : {data5_slip_clean}"
    )
    print(
        f"       DATA 5 matching polis  : {data5_polis_clean}"
    )
    print(
        f"       DATA 5 matching insured: {data5_insured_clean}"
    )
    print(
        f"       DATA 5 FAC_CODE        : {data5_facode_col or '[diturunkan dari Data 2]'}"
    )

    print(
        f"       FACUL clean slip     : {facul_slip_clean}"
    )

    print(
        f"       FACUL clean polis    : {facul_polis_clean}"
    )

    print(
        f"       FACUL clean insured  : {facul_insured_clean}"
    )

    # --------------------------------------------------------
    # BUILD INDEX
    # --------------------------------------------------------

    print("\n[6/7] Membuat index matching...")

    print("       Index OSBAL slip...")
    osbal_slip_index = build_clean_index(
        osbal,
        osbal_slip_clean
    )

    print("       Index OSBAL polis...")
    osbal_polis_index = build_clean_index(
        osbal,
        osbal_polis_clean
    )

    print("       Index OSBAL insured...")
    osbal_insured_index = build_clean_index(
        osbal,
        osbal_insured_clean
    )

    facul_facode_col = (
        FACUL_FACODE_COL
        if FACUL_FACODE_COL in facul.columns
        else None
    )

    print(
        "       FACUL FAC_CODE : "
        f"{facul_facode_col or '[diturunkan dari Data 2]'}"
    )

    print("       Index FACUL slip -> FAC_CODE...")
    facul_slip_index = build_aux_facode_index(
        facul,
        facul_slip_clean,
        facul_facode_col,
        osbal,
        osbal_slip_index
    )

    print("       Index FACUL polis -> FAC_CODE...")
    facul_polis_index = build_aux_facode_index(
        facul,
        facul_polis_clean,
        facul_facode_col,
        osbal,
        osbal_polis_index
    )

    print("       Index FACUL insured -> FAC_CODE...")
    facul_insured_index = build_aux_facode_index(
        facul,
        facul_insured_clean,
        facul_facode_col,
        osbal,
        osbal_insured_index
    )

    print("       Index OSBAL FAC_CODE...")
    osbal_by_facode = build_osbal_facode_index(
        osbal,
        OSBAL_FACODE_COL
    )

    # Data 5 -> FAC_CODE. Jika Data 5 tidak punya FAC_CODE,
    # helper akan mengambil FAC_CODE dari Data 2 berdasarkan key yang sama.
    print("       Index DATA 5 slip -> FAC_CODE...")
    data5_slip_index = build_aux_facode_index(
        data5,
        data5_slip_clean,
        data5_facode_col,
        osbal,
        osbal_slip_index
    )

    print("       Index DATA 5 polis -> FAC_CODE...")
    data5_polis_index = build_aux_facode_index(
        data5,
        data5_polis_clean,
        data5_facode_col,
        osbal,
        osbal_polis_index
    )

    print("       Index DATA 5 insured -> FAC_CODE...")
    data5_insured_index = build_aux_facode_index(
        data5,
        data5_insured_clean,
        data5_facode_col,
        osbal,
        osbal_insured_index
    )

    print("       Index LIKE DATA 5 insured...")
    data5_insured_like_index = build_aux_insured_like_index(
        data5, data5_insured_clean, osbal, osbal_insured_clean
    )

    print("       Index LIKE FACUL insured...")
    facul_insured_like_index = build_aux_insured_like_index(
        facul, facul_insured_clean, osbal, osbal_insured_clean
    )

    # --------------------------------------------------------
    # PREPARE SUSPEND MATRICES
    # --------------------------------------------------------

    suspend_slip_matrix = (
        suspend[suspend_slip_clean]
        .fillna("")
        .to_numpy(dtype=object)
    )

    suspend_polis_matrix = (
        suspend[suspend_polis_clean]
        .fillna("")
        .to_numpy(dtype=object)
    )

    suspend_insured_matrix = (
        suspend[suspend_insured_clean]
        .fillna("")
        .to_numpy(dtype=object)
    )

    # LIKE index khusus query INSURED Suspend -> row OSBAL.
    # Gunakan token bermakna agar tidak full-scan OSBAL untuk setiap row Suspend.
    osbal_insured_like_index = defaultdict(list)
    osbal_token_rows = defaultdict(set)
    osbal_insured_matrix = osbal[osbal_insured_clean].fillna("").to_numpy(dtype=object)
    for row_idx, ref_values in enumerate(osbal_insured_matrix):
        for ref in clean_values(ref_values):
            for tok in _insured_tokens(ref):
                if len(tok) >= 6 and tok not in INSURED_GENERIC_TOKENS:
                    osbal_token_rows[tok].add(row_idx)
    for query_values in suspend_insured_matrix:
        for q in clean_values(query_values):
            candidate_rows = set()
            for tok in _insured_tokens(q):
                if len(tok) >= 6 and tok not in INSURED_GENERIC_TOKENS:
                    candidate_rows.update(osbal_token_rows.get(tok, set()))
            for row_idx in candidate_rows:
                if any(insured_like_match(q, r) for r in clean_values(osbal_insured_matrix[row_idx])):
                    osbal_insured_like_index[q].append(row_idx)

    suspend_curr_values = (
        suspend[SUSPEND_CURR_COL]
        .fillna("")
        .to_numpy(dtype=object)
    )

    suspend_amount_values = (
        suspend[SUSPEND_AMOUNT_COL]
        .fillna("")
        .to_numpy(dtype=object)
    )

    suspend_date_values = (
        suspend[SUSPEND_DATE_COL]
        .to_numpy(dtype=object)
    )

    # --------------------------------------------------------
    # RESULT ARRAYS
    # --------------------------------------------------------

    match_method = []
    match_status = []
    flagging = []

    matched_fac_codes = []
    matching_fac_code_count = []

    matched_osbal_rows = []
    matched_osbal_excel_rows = []

    currency_status_list = []
    period_status_list = []

    # Untuk financial calculation
    resolved_rows_by_suspend = {}
    resolved_fac_by_suspend = {}

    # --------------------------------------------------------
    # MATCHING LOOP
    # --------------------------------------------------------

    print("\n[7/7] Proses matching...")

    total_rows = len(suspend)

    for idx in range(total_rows):

        if idx % 10000 == 0:
            print(
                f"       Processing {idx:,}/{total_rows:,}"
            )

        slip_values = clean_values(
            suspend_slip_matrix[idx]
        )

        polis_values = clean_values(
            suspend_polis_matrix[idx]
        )

        insured_values = clean_values(
            suspend_insured_matrix[idx]
        )

        curr_value = suspend_curr_values[idx]
        amount_value = suspend_amount_values[idx]
        date_value = suspend_date_values[idx]

        resolved = resolve_one_suspend(
            suspend_row_idx=idx,

            suspend_slip_values=slip_values,
            suspend_polis_values=polis_values,
            suspend_insured_values=insured_values,

            suspend_curr=curr_value,
            suspend_date=date_value,

            osbal_df=osbal,

            osbal_slip_index=osbal_slip_index,
            osbal_polis_index=osbal_polis_index,
            osbal_insured_index=osbal_insured_index,

            data5_slip_index=data5_slip_index,
            data5_polis_index=data5_polis_index,
            data5_insured_index=data5_insured_index,

            facul_slip_index=facul_slip_index,
            facul_polis_index=facul_polis_index,
            facul_insured_index=facul_insured_index,

            osbal_by_facode=osbal_by_facode,
            osbal_insured_like_index=osbal_insured_like_index,
            data5_insured_like_index=data5_insured_like_index,
            facul_insured_like_index=facul_insured_like_index
        )

        rows = resolved["rows"]
        fac_codes = resolved["fac_codes"]

        method = resolved["method"]

        # ----------------------------------------------------
        # NO VALID MATCH
        # ----------------------------------------------------

        if not rows:

            fallback_rows = list(
                resolved.get("candidate_rows", [])
            )
            fallback_fac_codes = sorted(
                set(
                    fc for fc in resolved.get("fac_codes", [])
                    if fc
                )
            )

            match_method.append(method)

            is_narrowing_failure = (
                resolved.get("failure") in
                ("Beda Currency", "Beda Periode")
                or resolved.get("currency_status") == "Beda Currency"
                or resolved.get("period_status") == "Beda Periode"
            )

            if is_narrowing_failure or resolved.get("failure") == "Matching >1 fac code":
                match_status.append(
                    resolved.get("failure") == "Matching >1 fac code"
                    and "Matching >1 fac code"
                    or "Matching"
                )
                flagging.append(
                    resolved.get("failure") or
                    resolved.get("currency_status") or
                    resolved.get("period_status")
                )
            else:
                match_status.append("Unmatching")
                flagging.append(
                    resolved.get("failure") or "Unmatching"
                )

            matched_fac_codes.append(
                ", ".join(fallback_fac_codes)
            )
            matching_fac_code_count.append(
                len(fallback_fac_codes)
            )
            matched_osbal_rows.append(len(fallback_rows))
            matched_osbal_excel_rows.append(
                "; ".join(str(row_idx + 2) for row_idx in fallback_rows)
            )

            currency_status_list.append(
                resolved["currency_status"]
            )
            period_status_list.append(
                resolved["period_status"]
            )

            # Kandidat tetap disimpan walaupun narrowing (currency/period)
            # gagal, agar FAC_CODE hasil matching tetap tampil di output.
            if fallback_rows:
                resolved_rows_by_suspend[idx] = set(fallback_rows)

            # Akumulasi hanya memakai FAC_CODE tunggal yang sudah pasti.
            # Jika kandidat memiliki beberapa FAC_CODE, FAC_CODE tetap
            # ditampilkan melalui CCOS_REF_CODE dan diberi flag multiple.
            if len(fallback_fac_codes) == 1:
                resolved_fac_by_suspend[idx] = fallback_fac_codes[0]

            continue

        # ----------------------------------------------------
        # MULTIPLE FAC CODE
        # ----------------------------------------------------

        fac_codes = sorted(
            set(
                fc
                for fc in fac_codes
                if fc
            )
        )

        if len(fac_codes) > 1:

            match_method.append(method)

            match_status.append(
                "Matching >1 fac code"
            )

            flagging.append(
                "Matching >1 fac code"
            )

            matched_fac_codes.append(
                ", ".join(fac_codes)
            )

            matching_fac_code_count.append(
                len(fac_codes)
            )

            matched_osbal_rows.append(
                len(rows)
            )

            matched_osbal_excel_rows.append(
                "; ".join(
                    str(row_idx + 2)
                    for row_idx in rows
                )
            )

            # Simpan seluruh kandidat row untuk tahap BUILD OUTPUT.
            # Tanpa ini, fac code jamak hilang saat output dibentuk karena
            # resolved_rows_by_suspend hanya diisi untuk satu FAC_CODE.
            resolved_rows_by_suspend[idx] = set(rows)

            currency_status_list.append(
                resolved["currency_status"]
            )

            period_status_list.append(
                resolved["period_status"]
            )

            continue

        # ----------------------------------------------------
        # ONE FAC CODE
        # ----------------------------------------------------

        if len(fac_codes) == 1:

            fac_code = fac_codes[0]

            match_method.append(method)

            match_status.append(
                "Matching"
            )

            flagging.append(
                ""
            )

            matched_fac_codes.append(
                fac_code
            )

            matching_fac_code_count.append(
                1
            )

            matched_osbal_rows.append(
                len(rows)
            )

            matched_osbal_excel_rows.append(
                "; ".join(
                    str(row_idx + 2)
                    for row_idx in rows
                )
            )

            currency_status_list.append(
                resolved["currency_status"]
            )

            period_status_list.append(
                resolved["period_status"]
            )

            resolved_rows_by_suspend[idx] = set(rows)
            resolved_fac_by_suspend[idx] = fac_code

            continue

        # ----------------------------------------------------
        # FAC CODE KOSONG
        # ----------------------------------------------------

        match_method.append(method)

        match_status.append(
            "Unmatching"
        )

        flagging.append(
            "FAC_CODE tidak ditemukan"
        )

        matched_fac_codes.append("")
        matching_fac_code_count.append(0)

        matched_osbal_rows.append(
            len(rows)
        )

        matched_osbal_excel_rows.append(
            "; ".join(
                str(row_idx + 2)
                for row_idx in rows
            )
        )

        currency_status_list.append(
            resolved["currency_status"]
        )

        period_status_list.append(
            resolved["period_status"]
        )

    print(
        f"       Selesai matching {total_rows:,} row Suspend."
    )

    # ========================================================
    # GROUP BY FAC_CODE
    # ========================================================

    print("\n[8/8] Menghitung accumulation & financial flag...")

    # --------------------------------------------------------
    # FAC_CODE -> Suspend rows
    # --------------------------------------------------------

    suspend_group = defaultdict(list)

    for suspend_idx, fac_code in (
        resolved_fac_by_suspend.items()
    ):
        suspend_group[fac_code].append(
            suspend_idx
        )

    # --------------------------------------------------------
    # FAC_CODE -> OSBAL rows
    # --------------------------------------------------------
    # AKUMULASI HARUS BERDASARKAN SELURUH BARIS OSBAL
    # DALAM FAC_CODE YANG SUDAH BERHASIL DITENTUKAN.
    #
    # Jangan hanya mengambil baris yang kebetulan ditemukan
    # oleh satu proses matching Suspend, karena satu FAC_CODE
    # dapat memiliki lebih dari satu SLIP/POLIS/INSURED.
    # Semua CCOS_BAL_DUE pada FAC_CODE tersebut harus dijumlahkan.

    osbal_group = defaultdict(set)

    for suspend_idx, fac_code in (
        resolved_fac_by_suspend.items()
    ):

        if not fac_code:
            continue

        for osbal_idx in osbal_by_facode.get(
            fac_code,
            []
        ):
            osbal_group[fac_code].add(
                osbal_idx
            )

    # --------------------------------------------------------
    # Calculate group financial values
    # --------------------------------------------------------

    group_financial = {}

    for fac_code, suspend_indices in suspend_group.items():

        osbal_indices = sorted(
            osbal_group.get(
                fac_code,
                set()
            )
        )

        # ----------------------------------------------
        # Suspend amount total
        # ----------------------------------------------

        suspend_amounts = []

        for suspend_idx in suspend_indices:

            value = to_number(
                suspend_amount_values[suspend_idx]
            )

            if value is not None:
                # AMOUNT ORI dikali -1 hanya untuk perhitungan
                suspend_amounts.append(value * -1)

        if suspend_amounts:
            amount_total = sum(
                suspend_amounts
            )
        else:
            amount_total = None

        # ----------------------------------------------
        # OSBAL CCOS_BAL_DUE total
        # ----------------------------------------------

        bal_due_values = []

        for osbal_idx in osbal_indices:

            value = to_number(
                osbal.iloc[osbal_idx][
                    OSBAL_BAL_DUE_COL
                ]
            )

            if value is not None:
                bal_due_values.append(value)

        if bal_due_values:
            # Jika hasil terlihat berbeda dari nilai satu row OSBAL, itu
            # memang berarti ada lebih dari satu row yang masuk ke FAC_CODE
            # tersebut. Rule bisnis tetap: seluruh row OSBAL dalam FAC_CODE
            # harus diakumulasi.
            #
            # Gunakan Decimal agar agregasi CCOS_BAL_DUE tidak mengalami
            # artefak floating-point. Jika hanya ada 1 baris OSBAL, nilai
            # tersebut dipakai langsung; jika >1 baris, baru dijumlahkan.
            #
            # PENTING: jangan melakukan round() secara membabi-buta.
            # Nilai pecahan yang memang ada di OSBAL harus tetap dipertahankan.
            from decimal import Decimal

            bal_due_decimal = sum(
                (Decimal(str(v)) for v in bal_due_values),
                Decimal("0")
            )

            if bal_due_decimal == bal_due_decimal.to_integral_value():
                bal_due_total = int(bal_due_decimal)
            else:
                bal_due_total = float(bal_due_decimal)
        else:
            bal_due_total = None

        suspend_count = len(
            suspend_indices
        )

        osbal_count = len(
            osbal_indices
        )

        financial_flag = compare_amounts(
            amount_total=amount_total,
            bal_due_total=bal_due_total,
            suspend_count=suspend_count,
            osbal_count=osbal_count
        )

        group_financial[fac_code] = {
            "amount_total": amount_total,
            "bal_due_total": bal_due_total,
            "suspend_count": suspend_count,
            "osbal_count": osbal_count,
            "financial_flag": financial_flag
        }

    # ========================================================
    # BUILD OUTPUT - HANYA KOLOM FINAL, URUTAN SUSPEND DIPERTAHANKAN
    # ========================================================

    # Hitung financial group untuk semua FAC_CODE yang valid.
    amount_total_result = []
    bal_due_total_result = []
    difference_result = []
    financial_flag_result = []
    scenario_result = []

    # Kolom output akhir sesuai kolom_match.xlsx
    output_rows = []

    for idx in range(len(suspend)):
        sus = suspend.iloc[idx]
        fac_code = resolved_fac_by_suspend.get(idx)
        rows = sorted(resolved_rows_by_suspend.get(idx, set()))
        fac_codes = sorted({
            normalize_code(osbal.iloc[r][OSBAL_FACODE_COL])
            for r in rows
            if normalize_code(osbal.iloc[r][OSBAL_FACODE_COL])
        })

        # Fallback untuk narrowing failure: jika baris kandidat tidak
        # tersedia pada mapping output, gunakan FAC_CODE yang sudah
        # dicatat pada tahap matching agar tidak hilang dari output.
        if not fac_codes and idx < len(matched_fac_codes):
            fac_codes = sorted({
                normalize_code(fc)
                for fc in str(matched_fac_codes[idx]).split(',')
                if normalize_code(fc)
            })

        # Data matching yang diagregasi jika lebih dari satu baris.
        amount_calc = None
        bal_due_total = None
        or_bal_total = None
        financial_flag = ""
        if fac_code and fac_code in group_financial:
            group = group_financial[fac_code]
            amount_calc = group["amount_total"]
            bal_due_total = group["bal_due_total"]
            financial_flag = group["financial_flag"]
            or_bal_total = safe_sum(
                [osbal.iloc[r][OSBAL_OR_BAL_COL] for r in
                 sorted(osbal_group.get(fac_code, set()))]
            )

        # AMOUNT_ORI tetap menyimpan nilai sumber per row.
        # AMOUNT_ORI_MIN1 untuk row yang sudah memiliki satu FAC_CODE
        # mengikuti hasil akumulasi seluruh row Suspend pada FAC_CODE
        # tersebut. Dengan demikian hasil accumulation tidak hanya hidup
        # di backend/financial calculation, tetapi benar-benar terlihat
        # pada kolom output.
        amount_ori = sus.get(SUSPEND_AMOUNT_COL)
        amount_ori_min1_row = (
            to_number(amount_ori) * -1
            if to_number(amount_ori) is not None else 0
        )

        if amount_calc is None:
            amount_calc = amount_ori_min1_row
        if bal_due_total is None:
            bal_due_total = 0

        if fac_code and fac_code in group_financial:
            amount_ori_min1 = group_financial[fac_code]["amount_total"]
            if amount_ori_min1 is None:
                amount_ori_min1 = amount_ori_min1_row
        else:
            # Untuk multi-FAC_CODE tidak ada satu FAC_CODE yang dapat
            # dipilih sebagai basis accumulation. Tetap tampilkan nilai
            # row agar tidak membuat angka financial palsu.
            amount_ori_min1 = amount_ori_min1_row

        # Untuk Beda Currency/Periode, kandidat matching tetap dipakai
        # untuk menentukan scenario. Scenario menjelaskan MATCH-NYA,
        # sedangkan narrowing menjadi suffix.
        fallback_rows = sorted(resolved_rows_by_suspend.get(idx, set()))
        scenario_rows = rows if rows else fallback_rows

        # Jika hasil narrowing memiliki status Beda Currency/Beda Periode,
        # status tersebut tetap menjadi keterangan final meskipun kandidat
        # awal memiliki lebih dari satu FAC_CODE. Dalam kondisi ini scenario
        # juga harus menampilkan suffix yang sama. Hanya jika currency/period
        # valid dan hasil akhir masih >1 FAC_CODE, scenario tidak memakai suffix.
        is_multi_fac_final = len(fac_codes) > 1

        if currency_status_list[idx] == "Beda Currency":
            scenario_suffix = " (Beda Currency)"
        elif period_status_list[idx] == "Beda Periode":
            scenario_suffix = " (Beda Periode)"
        else:
            scenario_suffix = ""

        scenario = build_scenario(
            clean_values(suspend_slip_matrix[idx]),
            clean_values(suspend_polis_matrix[idx]),
            clean_values(suspend_insured_matrix[idx]),
            scenario_rows,
            osbal,
            osbal_slip_clean,
            osbal_polis_clean,
            osbal_insured_clean,
            suffix=scenario_suffix,
            insured_like=(match_method[idx] == "INSURED"),
            selected_method=match_method[idx]
        )

        # Status narrowing lebih spesifik daripada jumlah FAC_CODE.
        # Jadi jika kandidat memiliki >1 FAC_CODE tetapi gugur pada
        # narrowing karena currency/period, FLAG_PROD tetap mengikuti
        # kategori Beda Currency/Beda Periode. "Matching >1 fac code"
        # hanya dipakai bila narrowing valid dan masih menyisakan >1 FAC_CODE.
        if currency_status_list[idx] == "Beda Currency":
            flag_prod = "Beda Currency"
        elif period_status_list[idx] == "Beda Periode":
            flag_prod = "Beda Periode"
        elif len(fac_codes) > 1:
            flag_prod = "Matching >1 fac code"
        elif not scenario_rows:
            # Unmatching hanya jika benar-benar tidak ada kandidat match.
            flag_prod = "Unmatching"
        else:
            # Baris yang sudah match tidak boleh jatuh ke FLAG_PROD
            # "Unmatching" hanya karena financial_flag kosong/error.
            flag_prod = financial_flag if financial_flag else ""

        # Untuk hasil yang belum memiliki satu FAC_CODE yang valid untuk
        # financial checking, CCOS_BAL_DUE tidak boleh diisi dari OSBAL.
        # Gunakan blank agar tidak terlihat seolah-olah ada balance yang
        # sudah ter-resolve.
        non_financial_flags = {
            "Beda Currency",
            "Beda Periode",
            "Unmatching",
            "Matching >1 fac code",
        }

        if flag_prod in non_financial_flags:
            bal_due_output = ""
            difference_output = ""
        else:
            bal_due_output = bal_due_total
            difference_output = (
                amount_calc - bal_due_total
                if amount_calc is not None and bal_due_total is not None
                else ""
            )

        def first_osbal(col):
            return osbal.iloc[scenario_rows[0]][col] if scenario_rows and col in osbal.columns else ""

        # ====================================================
        # OUTPUT COLUMN TEMPLATE: New kolom.xlsx
        # ====================================================
        # Isi Suspend TIDAK ditukar. Backend saja yang membalik
        # interpretasi slip/polis untuk matching.
        def source_clean(prefix, number):
            value = sus.get(f"clean {prefix} {number}", "")
            return value if normalize_text(value) else ""

        def source_cert(number):
            # Cleaning Wahana dapat menghasilkan Certificate 1..3.
            # Fallback ke Certificate untuk kompatibilitas dengan hasil
            # cleaning yang hanya mempunyai satu kolom certificate.
            value = sus.get(f"Certificate {number}", "")
            if normalize_text(value):
                return value
            if number == 1:
                return sus.get("Certificate", "")
            return ""

        out = {
            "CCOS_DOC_NO": first_osbal(OSBAL_DOC_COL),
            "CCOS_REF_CODE": ", ".join(fac_codes),
            "RECEIPT NO": sus.get("RECEIPT NO", ""),
            "CREDIT NOTES": sus.get("CREDIT NOTES", ""),
            "DETAIL RINCIAN NO": sus.get("DETAIL RINCIAN NO", ""),
            "RECEIPT DATE": sus.get(SUSPEND_DATE_COL, ""),
            "CEDANT NAME": sus.get("CEDANT NAME", ""),
            "CEDANT SHRT NAME": sus.get("CEDANT SHRT NAME", ""),
            "INSURED_ORI": sus.get(SUSPEND_INSURED_COL, ""),
            "INSURED_1": source_clean("insured", 1),
            "INSURED_2": source_clean("insured", 2),
            "CURR ORI": sus.get(SUSPEND_CURR_COL, ""),
            "AMOUNT ORI": amount_ori,
            "AMOUNT_ORI_MIN1": amount_ori_min1,
            "CURR PAY": sus.get("CURR PAY", ""),
            "AMOUNT PAY": sus.get("AMOUNT PAY", ""),
            "CCOS_OR_BAL": or_bal_total,
            "CCOS_BAL_DUE": bal_due_output,
            "DIFERENCE": difference_output,
            "FLAG_PROD": flag_prod,

            # Tetap mempertahankan isi sumber. Tidak ada pertukaran isi
            # antara kolom polis dan slip pada output.
            "POLIS": sus.get("polis_ori", ""),
            "POLICY_CLEAN_1": source_clean("polis", 1),
            "SERTIF_CLEAN_1": source_cert(1),
            "POLICY_CLEAN_2": source_clean("polis", 2),
            "SERTIF_CLEAN_2": source_cert(2),
            "POLICY_CLEAN_3": source_clean("polis", 3),
            "SERTIF_CLEAN_3": source_cert(3),
            "POLICY_CLEAN_4": source_clean("polis", 4),
            "POLICY_CLEAN_5": source_clean("polis", 5),
            "SLIP_NO": sus.get("slip_ori", ""),
            "SLIP_NO_CLN": source_clean("slip", 1),

            "DESC 1": sus.get("DESC 1", ""),
            "DESC 2": sus.get("DESC 2", ""),
            "DESC 3": sus.get("DESC 3", ""),
            "DESC 4": sus.get("DESC 4", ""),
            "STATUS": sus.get("STATUS", ""),
            "REC_TYPE": sus.get("REC_TYPE", ""),
            "SKENARIO": scenario,
        }
        output_rows.append(out)

    # Schema fixed mengikuti New kolom.xlsx agar semua cedant mempunyai
    # susunan kolom yang identik. Kolom yang tidak ada pada data -> kosong.
    result = pd.DataFrame(output_rows, columns=[
        "CCOS_DOC_NO", "CCOS_REF_CODE", "RECEIPT NO", "CREDIT NOTES",
        "DETAIL RINCIAN NO", "RECEIPT DATE", "CEDANT NAME", "CEDANT SHRT NAME",
        "INSURED_ORI", "INSURED_1", "INSURED_2", "CURR ORI", "AMOUNT ORI",
        "AMOUNT_ORI_MIN1", "CURR PAY", "AMOUNT PAY", "CCOS_OR_BAL",
        "CCOS_BAL_DUE", "DIFERENCE", "FLAG_PROD", "POLIS",
        "POLICY_CLEAN_1", "SERTIF_CLEAN_1", "POLICY_CLEAN_2",
        "SERTIF_CLEAN_2", "POLICY_CLEAN_3", "SERTIF_CLEAN_3",
        "POLICY_CLEAN_4", "POLICY_CLEAN_5", "SLIP_NO", "SLIP_NO_CLN",
        "DESC 1", "DESC 2", "DESC 3", "DESC 4", "STATUS", "REC_TYPE",
        "SKENARIO"
    ])

    # ========================================================
    # SAVE
    # ========================================================

    result.to_excel(
        OUTPUT_FILE,
        index=False
    )

    # ========================================================
    # SUMMARY
    # ========================================================

    print("\n" + "=" * 70)
    print("MATCHING SELESAI")
    print("=" * 70)

    print(
        f"\nOutput : {OUTPUT_FILE}"
    )

    print(
        f"Jumlah row output : {len(result):,}"
    )

    print(
        f"Jumlah row Suspend: {len(suspend):,}"
    )

    # --------------------------------------------------------
    # Sanity check row count
    # --------------------------------------------------------

    if len(result) != len(suspend):

        raise RuntimeError(
            "\nFATAL: jumlah row output berbeda "
            "dengan jumlah row Suspend!"
        )

    print(
        "\n✓ Row count tetap sama."
    )

    print(
        "✓ Urutan Suspend tidak di-sort."
    )

    print(
        "✓ Tidak dilakukan merge terhadap Suspend."
    )

    print(
        "✓ Tidak ada row Suspend yang diduplikasi."
    )

    # --------------------------------------------------------
    # Summary status
    # --------------------------------------------------------

    # --------------------------------------------------------
    # SUMMARY KHUSUS WAHANA
    # Hanya baris dengan CEDANT NAME persis:
    # PT.ASURANSI WAHANA TATA
    # --------------------------------------------------------

    WAHANA_SUMMARY_CEDANT = "PT.ASURANSI WAHANA TATA"

    cedant_name_summary = (
        result["CEDANT NAME"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    wahana_result = result.loc[
        cedant_name_summary == WAHANA_SUMMARY_CEDANT
    ].copy()

    print("\n" + "=" * 70)
    print("SUMMARY KHUSUS WAHANA")
    print("=" * 70)
    print(f"CEDANT NAME : {WAHANA_SUMMARY_CEDANT}")
    print(f"Rows        : {len(wahana_result):,}")
    print(f"Cols        : {len(wahana_result.columns):,}")
    print(f"File        : {OUTPUT_FILE}")

    print("\nFLAG_PROD summary:")

    flag_counts = (
        wahana_result["FLAG_PROD"]
        .fillna("")
        .replace("", "(blank)")
        .value_counts(dropna=False)
    )

    for flag, count in flag_counts.items():
        print(f"  {flag:<45}: {count:,}")

    print("\nSKENARIO summary:")

    scenario_counts = (
        wahana_result["SKENARIO"]
        .fillna("")
        .replace("", "(blank)")
        .value_counts(dropna=False)
    )

    for scenario_name, count in scenario_counts.items():
        print(f"  {scenario_name:<45}: {count:,}")


# ============================================================
# ENTRY POINT MATCHING WAHANA
# ============================================================
if __name__ == "__main__":
    main()
