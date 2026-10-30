import difflib
import functools
import os
import re
import sys
import time

sys.modules["numexpr"] = None
sys.modules["bottleneck"] = None

import numpy as np
import pandas as pd

# Konfigurasi direktori dan file
_BASE_DIR = os.path.dirname(os.path.abspath(__file__))
_CLEAN_DIR = os.path.join(_BASE_DIR, "data_clean")
_RAW_DIR = os.path.join(_BASE_DIR, "data_raw")


def _find_file_in(directory: str, basenames: list) -> str:
    # Mencari file pertama yang tersedia dari daftar nama yang diberikan
    for name in basenames:
        p = os.path.join(directory, name)
        if os.path.exists(p):
            return p
    return os.path.join(directory, basenames[0])


SUSPEND_FILE = _find_file_in(_CLEAN_DIR, ["mitrautama_output_suspend.xlsx"])
OSBAL_FILE = _find_file_in(_CLEAN_DIR, ["mitrautama_output_osbal.xlsx"])
FACUL_FILE = _find_file_in(_CLEAN_DIR, ["mitrautama_output_facul.xlsx"])
SLIPDB_FILE = _find_file_in(
    _RAW_DIR,
    ["5. Database Rekap RI Slip.xlsx", "5.Database Rekap RI Slip.xlsx", "Database Rekap RI Slip.xlsx"],
)
SLIPDB_HEADER_ROW = 1
OUTPUT_FILE = os.path.join(_BASE_DIR, "matching_mitrautama_v1.xlsx")

# Definisi nama kolom referensi dan kunci fac code
OSBAL_FACODE_COL = "CCOS_REF_CODE"
FACUL_FACODE_COL = "FAC_CODE"
SLIPDB_FACODE_COL = "FAC. CODE"

# Definisi kolom data Suspend (Data 3)
SUSPEND_CURR_COL = "CURR ORI"
SUSPEND_DATE_COL = "RECEIPT DATE"
SUSPEND_SLIP_PREFIX = "SLIP_NO_CLEAN"
SUSPEND_INSURED_PREFIX = "INSURED_CLEAN"
SUSPEND_POLIS_PREFIX = "POLIS_CLEAN"
SUSPEND_SERTIF_PREFIX = "CERTIFICATE_"
SUSPEND_GROUP_COL = "DESC 2"

# Definisi kolom data Osbal (Data 2)
OSBAL_CURR_COL = "CCOS_CURR"
OSBAL_DATE_COL = "FAC_COM_DATE"
OSBAL_SLIP_COLS = ["SLIP_CLEAN_1", "SLIP_CLEAN_2", "SLIP_CLEAN_3", "SLIP_CLEAN_4", "CLSDT_SLIP_NO"]
OSBAL_POLIS_COLS = ["POLICY_CLEAN_1", "POLICY_CLEAN_2", "CLSDT_POLICY_NO"]
OSBAL_INSURED_COLS = [
    "FAC_INSURED_CLN_1", "FAC_INSURED_CLN_2", "FAC_INSURED_CLN_3",
    "FAC_INSURED_CLN_4", "FAC_INSURED_CLN_5"
]

# Definisi kolom data Facul (Data 1)
FACUL_SLIP_COLS = ["FAC_SLIP_CLEAN_1", "FAC_SLIP_CLEAN_2", "FAC_SLIP_CLEAN_3", "FAC_SLIP_CLEAN_4"]
FACUL_POLIS_COLS = ["FAC_POLICY_NO_CLEAN_1", "FAC_POLICY_NO_CLEAN_2"]
FACUL_INSURED_COLS = [
    "FAC_INSURED_CLN_1", "FAC_INSURED_CLN_2", "FAC_INSURED_CLN_3",
    "FAC_INSURED_CLN_4", "FAC_INSURED_CLN_5"
]

# Definisi kolom data Slip DB (Data 5)
SLIPDB_RAW_SLIP = "SLIP"
SLIPDB_RAW_POLIS = "POLICY NO"
SLIPDB_RAW_INSURED = "INSURED"
SLIPDB_SLIP_COLS = ["_SLIP_KEY"]
SLIPDB_POLIS_COLS = ["_POLIS_KEY"]
SLIPDB_INSURED_COLS = ["_INSURED_KEY"]

# Konstanta pemfilteran dan pencocokan teks
_EXCLUDED_REF_MARKERS = ["SUSPENSE", "HUTANG", "UTANG"]
_LINESLIP_MARKERS = ("LINESLIP", "LINE SLIP")
_PLACEHOLDERS = frozenset({"", "EMPTY", "NAN", "NONE", "-", "TBA", "NULL", "0"})
_MIN_KEY_LEN = 4
_MIN_TOKEN_LEN = 3
_TOKEN_RE = re.compile(r"[^A-Z0-9]+")
_INSURED_LIKE_THRESHOLD: float = 0.75
_BANK_RE = re.compile(r"\bBANK\b|\bBPR\b|\bBPD\b")
_STRONG_WORDS: frozenset = frozenset({"PODOMORO"})

# Daftar stopwords untuk pembersihan nama tertanggung
_INSURED_STOPWORDS: frozenset = frozenset({
    "PT", "CV", "TBK", "UD", "PD", "NV", "LLC", "LTD", "INC",
    "GROUP", "GROUPS", "INDONESIA", "INDO",
    "INDUSTRY", "INDUSTRIES", "INDUSTRI",
    "INTERNATIONAL", "INTERNASIONAL",
    "COMPANY", "COMPANIES", "CORP", "CORPORATION",
    "ENTERPRISE", "ENTERPRISES", "TRADING", "TRADE",
    "UTAMA", "MANDIRI", "MAKMUR", "SEJAHTERA",
})

# Kolom-kolom yang digabungkan saat ada akumulasi suspend
_SUSPEND_JOIN_COLS = [
    "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "RECEIPT DATE",
    "CEDANT NAME", "CEDANT SHRT NAME",
    "INSURED", "INSURED_CLEAN_1", "INSURED_CLEAN_2", "INSURED_CLEAN_3", "INSURED_CLEAN_4", "INSURED_CLEAN_5",
    "CURR ORI", "CURR PAY", "AMOUNT PAY",
    "POLIS", "SLIP NO", "SLIP_NO_CLEAN",
    "DESC 1", "DESC 2", "DESC 3", "DESC 4",
    "STATUS", "REC_TYPE",
]
_SUSPEND_SUM_COLS = ["AMOUNT ORI"]

# Susunan urutan kolom final output
FINAL_COLUMNS = [
    "CCOS_DOC_NO", "CCOS_REF_CODE",
    "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "RECEIPT DATE",
    "CEDANT NAME", "CEDANT SHRT NAME",
    "INSURED_ORI", "INSURED_1", "INSURED_2", "INSURED_3", "INSURED_4", "INSURED_5",
    "CURR ORI", "AMOUNT ORI", "AMOUNT_ORI_MIN1",
    "CURR PAY", "AMOUNT PAY",
    "CCOS_OR_BAL", "CCOS_BAL_DUE", "DIFERENCE",
    "FLAG_PROD",
    "POLIS_ORI",
    "POLICY_CLEAN_1", "SERTIF_CLEAN_1",
    "POLICY_CLEAN_2", "SERTIF_CLEAN_2",
    "POLICY_CLEAN_3", "SERTIF_CLEAN_3",
    "POLICY_CLEAN_4", "POLICY_CLEAN_5",
    "SLIP_NO_ORI", "SLIP_NO_CLN",
    "DESC 1", "DESC 2", "DESC 3", "DESC 4",
    "STATUS", "REC_TYPE",
    "SKENARIO",
]

# Kolom yang harus diformat sebagai text di file excel
_TEXT_COLS = [
    "POLIS_ORI",
    "POLICY_CLEAN_1", "SERTIF_CLEAN_1", "POLICY_CLEAN_2", "SERTIF_CLEAN_2",
    "POLICY_CLEAN_3", "SERTIF_CLEAN_3", "POLICY_CLEAN_4", "POLICY_CLEAN_5",
    "SLIP_NO_ORI", "SLIP_NO_CLN",
    "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
]


@functools.lru_cache(maxsize=131072)
def _norm_cached(s: str) -> str:
    # Normalisasi teks kapital dan menghapus spasi ganda dengan cache
    text = s.strip().upper().replace("S/D", "SD")
    return re.sub(r"\s+", " ", text) if "  " in text else text


def _normalize(value) -> str:
    # Mengembalikan teks bersih atau string kosong jika bernilai null
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return _norm_cached(str(value))


def _valid(value: str) -> bool:
    # Memeriksa apakah nilai bukan merupakan placeholder kosong
    return value not in _PLACEHOLDERS


def _alnum_key(value) -> str:
    # Mengambil hanya karakter alfanumerik untuk kunci pencocokan
    v = _normalize(value)
    if not v or v in _PLACEHOLDERS:
        return ""
    return re.sub(r"[^A-Z0-9]", "", v)


def _format_sertif(val) -> str:
    # Memformat nomor sertifikat menjadi 6 digit angka
    if val is None or (isinstance(val, float) and pd.isna(val)) or str(val).strip() == "":
        return ""
    if isinstance(val, (int, float)):
        return str(int(val)).zfill(6)
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s.zfill(6) if s.isdigit() else s


@functools.lru_cache(maxsize=200000)
def _tokenize(text: str) -> frozenset:
    # Memecah teks menjadi token kata unik dengan panjang minimal
    if not text:
        return frozenset()
    return frozenset(t for t in _TOKEN_RE.split(text) if len(t) >= _MIN_TOKEN_LEN)


def _get_cols(columns: list, prefix: str) -> list:
    # Mengambil daftar kolom yang memiliki awalan prefix tertentu
    return [c for c in columns if str(c).lower().startswith(prefix.lower())]


def _to_float(v) -> float:
    # Mengubah berbagai representasi string/angka menjadi tipe float
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return 0.0
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _is_excluded_ref_row(row: dict, cols: list) -> bool:
    # Mengecek apakah baris data referensi mengandung kata pengecualian
    for col in cols:
        val = _normalize(row.get(col, ""))
        if val and any(m in val for m in _EXCLUDED_REF_MARKERS):
            return True
    return False


def _is_lineslip_row(sus_row: dict) -> bool:
    # Mendeteksi apakah baris suspend merupakan transaksi lineslip
    for col in ("INSURED", "DESC 1", "DESC 2", "DESC 3"):
        val = _normalize(sus_row.get(col, ""))
        if val and any(m in val for m in _LINESLIP_MARKERS):
            return True
    return False


def _collect_values(row: dict, cols: list, key_fn=None) -> list:
    # Mengumpulkan nilai unik yang valid dari beberapa kolom pada baris
    seen, result = set(), []
    for col in cols:
        raw = row.get(col, "")
        val = key_fn(raw) if key_fn else _normalize(raw)
        if val and _valid(val) and val not in seen:
            seen.add(val)
            result.append(val)
    return result


def _row_values(row: dict, cols: list) -> set:
    # Mengembalikan set nilai normalisasi dari kolom-kolom yang diminta
    return set(_collect_values(row, cols))


def _build_exact_index(rows: list, cols: list, key_fn=None, excluded: set = None) -> dict:
    # Membangun dictionary index nilai eksak terhadap indeks baris
    index: dict = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        for val in _collect_values(row, cols, key_fn):
            index.setdefault(val, []).append(i)
    return index


def _build_facode_index(rows: list, col: str, excluded: set = None) -> dict:
    # Membangun dictionary index fac code terhadap nomor indeks baris
    index: dict = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        val = _normalize(row.get(col, ""))
        if val:
            index.setdefault(val, []).append(i)
    return index


def _build_token_index(rows: list, cols: list, excluded: set = None) -> dict:
    # Membangun inverted index token kata untuk pencocokan like insured
    token_index: dict = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        tokens: frozenset = frozenset()
        for c in cols:
            t = _normalize(row.get(c, ""))
            if t and _valid(t):
                tokens = tokens | _tokenize(t)
        for tok in tokens:
            token_index.setdefault(tok, set()).add(i)
    return token_index


def _build_lookup(rows: list, slip_cols: list, polis_cols: list, insured_cols: list,
                  facode_col: str = None, label: str = "") -> dict:
    # Membuat seluruh struktur lookup index terbalik untuk dataset target
    if not rows:
        if label:
            print(f"  {label:<6} done (0.0s)", flush=True)
        return {"slip": {}, "polis": {}, "insured": {}, "tok_insured": {}, "facode": {}}

    t0 = time.perf_counter()
    excluded = {i for i, r in enumerate(rows)
                if _is_excluded_ref_row(r, slip_cols + polis_cols + insured_cols)}
    lookup = {
        "slip": _build_exact_index(rows, slip_cols, _alnum_key, excluded),
        "polis": _build_exact_index(rows, polis_cols, _alnum_key, excluded),
        "insured": _build_exact_index(rows, insured_cols, None, excluded),
        "tok_insured": _build_token_index(rows, insured_cols, excluded),
        "facode": _build_facode_index(rows, facode_col, excluded) if facode_col else {},
    }
    if label:
        exc_info = f" - {len(excluded):,} baris dikecualikan" if excluded else ""
        print(f"  {label:<6} done ({time.perf_counter()-t0:.1f}s){exc_info}", flush=True)
    return lookup


def _exact_match(query_values: list, index: dict) -> set:
    # Mencari set indeks baris yang cocok eksak pada index dictionary
    return {i for v in query_values if v and v in index for i in index[v]}


def _like_match(query_values: list, token_index: dict, rows: list, ref_cols: list,
                stopwords: frozenset = _INSURED_STOPWORDS,
                overlap_threshold: float = _INSURED_LIKE_THRESHOLD) -> list:
    # Mencocokkan nama tertanggung berdasarkan token fuzzy overlap
    valid_queries = [v for v in query_values if v and len(v) >= 2]
    if not valid_queries:
        return []

    query_tokens: frozenset = frozenset()
    for v in valid_queries:
        query_tokens = query_tokens | _tokenize(v)

    candidates: set = set()
    for tok in query_tokens:
        if tok in token_index:
            candidates |= token_index[tok]
    if not candidates:
        return []

    matched = []
    for i in candidates:
        ref_values = [_normalize(rows[i].get(c, "")) for c in ref_cols]
        ref_values = [rv for rv in ref_values if rv and _valid(rv)]
        if any(q in rv for q in valid_queries for rv in ref_values):
            matched.append(i)
            continue

        hit = False
        for q in valid_queries:
            sig_q = _tokenize(q) - stopwords
            if not sig_q:
                continue
            for rv in ref_values:
                rv_toks = _tokenize(rv)
                sig_rv = rv_toks - stopwords
                if not sig_rv:
                    continue
                if sig_q.issubset(rv_toks):
                    hit = True
                    break
                used, inter = set(), 0
                for t1 in sig_q:
                    for t2 in sig_rv:
                        if t2 in used:
                            continue
                        if t1 == t2 or (len(t1) >= 4 and len(t2) >= 4
                                        and difflib.SequenceMatcher(None, t1, t2).ratio() >= 0.85):
                            inter += 1
                            used.add(t2)
                            break
                denom = min(len(sig_q), len(sig_rv))
                if denom > 0 and inter / denom >= overlap_threshold:
                    hit = True
                    break

                if inter > 0 and (sig_q & sig_rv & _STRONG_WORDS):
                    hit = True
                    break

                if len(sig_q) == 1 and len(sig_rv) > 1:
                    q_tok = list(sig_q)[0]
                    if 2 <= len(q_tok) <= 5:
                        rv_words = [w for w in _TOKEN_RE.split(rv.upper()) if w and w not in stopwords]
                        rv_initials = "".join([w[0] for w in rv_words]) if rv_words else ""
                        if rv_initials and (q_tok == rv_initials or q_tok == rv_initials[1:] or q_tok == rv_initials[:-1]):
                            hit = True
                            break
            if hit:
                break
        if hit:
            matched.append(i)
    return matched


def _row_has_insured(ref_row: dict, ref_cols: list, query_values: list) -> bool:
    # Memeriksa apakah suatu baris referensi memuat salah satu nama tertanggung
    ref_vals = _collect_values(ref_row, ref_cols)
    if not ref_vals or not query_values:
        return False
    if set(query_values) & set(ref_vals):
        return True
    for q in query_values:
        sig_q = _tokenize(q) - _INSURED_STOPWORDS
        if not sig_q:
            continue
        if any(sig_q.issubset(_tokenize(rv)) for rv in ref_vals):
            return True
    return False


def _split_insured_queries(sus: dict, insured_cols: list) -> list:
    # Memisahkan query pencarian tertanggung per nama atau gabungan entitas bank
    names = _collect_values(sus, insured_cols)
    if not names:
        return []
    ori = _normalize(sus.get("INSURED", ""))
    if _BANK_RE.search(ori) or any(_BANK_RE.search(n) for n in names):
        return [names]
    return [[n] for n in names]


def _group_name(sus: dict) -> str:
    # Mengambil nama group tertanggung dari kolom DESC 2
    g = _normalize(sus.get(SUSPEND_GROUP_COL, ""))
    if not g or not _valid(g) or len(g) < 3:
        return ""
    return g


def _resolve_facode(ref_rows_matched: list, facode_col: str,
                    facode_osbal_idx: dict) -> list:
    # Menghubungkan fac code dari Facul atau SlipDB ke baris Osbal
    fac_codes = {_normalize(r.get(facode_col, "")) for r in ref_rows_matched}
    fac_codes.discard("")
    out: set = set()
    for fac in fac_codes:
        if fac in facode_osbal_idx:
            out.update(facode_osbal_idx[fac])
    return list(out)


def _narrow_currency_periode(sus: dict, idxs: list, osbal_rows: list,
                             effective_date, is_lineslip: bool) -> tuple:
    # Menyaring kandidat kecocokan berdasarkan kesesuaian mata uang dan periode
    cand = list(idxs)
    sus_curr = _normalize(sus.get(SUSPEND_CURR_COL, ""))
    if sus_curr:
        f = [i for i in cand if _normalize(osbal_rows[i].get(OSBAL_CURR_COL, "")) == sus_curr]
        if not f:
            return cand, True, False
        cand = f

    if effective_date is not None and pd.notna(effective_date):
        f = []
        for i in cand:
            com = osbal_rows[i].get("_com_date_parsed")
            if com is None or pd.isna(com):
                f.append(i)
            elif is_lineslip:
                if com.year == effective_date.year and com.month == effective_date.month:
                    f.append(i)
            elif effective_date >= com:
                f.append(i)
        if not f:
            return cand, False, True
        cand = f
    return cand, False, False


def _aggregate_ccos(osbal_rows: list, all_fac_osbal_rows: list = None) -> dict:
    # Menghitung agregasi saldo balance due dan nomor dokumen dari baris Osbal
    fac_codes = {str(r.get("CCOS_REF_CODE", "")).strip() for r in osbal_rows
                 if str(r.get("CCOS_REF_CODE", "")).strip()}
    multi_fac = len(fac_codes) > 1
    bal_due_rows = all_fac_osbal_rows if (all_fac_osbal_rows is not None and not multi_fac) else osbal_rows

    def join_unique(key):
        seen, vals = set(), []
        for r in osbal_rows:
            v = str(r.get(key, "")).strip()
            if v and v not in seen and v.lower() != "nan":
                seen.add(v)
                vals.append(v)
        return ", ".join(vals)

    return {
        "CCOS_DOC_NO": join_unique("CCOS_DOC_NO"),
        "CCOS_REF_CODE": join_unique("CCOS_REF_CODE"),
        "CCOS_OR_BAL": np.nan if multi_fac else sum(_to_float(r.get("CCOS_OR_BAL")) for r in bal_due_rows),
        "CCOS_BAL_DUE": np.nan if multi_fac else sum(_to_float(r.get("CCOS_BAL_DUE")) for r in bal_due_rows),
    }


def _text_or_empty(v) -> str:
    # Mengembalikan teks bersih atau string kosong jika bernilai nan/none
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("nan", "none") else s


def _build_output_row(sus: dict, source: str, scenario: str, osbal_rows: list,
                      osbal_count: int = 0, suspend_count: int = 1,
                      all_fac_osbal_rows: list = None) -> dict:
    # Membentuk kamus data untuk satu baris laporan output hasil matching
    has_osbal = bool(source and osbal_rows and osbal_rows[0])
    ccos = (_aggregate_ccos(osbal_rows, all_fac_osbal_rows)
            if has_osbal
            else {"CCOS_DOC_NO": "", "CCOS_REF_CODE": "", "CCOS_OR_BAL": np.nan, "CCOS_BAL_DUE": np.nan})

    final_scenario = "Unmatching" if not source or scenario in ("Unmatching", "UNMATCHED") else scenario

    row = {
        "CCOS_DOC_NO": ccos["CCOS_DOC_NO"],
        "CCOS_REF_CODE": ccos["CCOS_REF_CODE"],
        "RECEIPT NO": sus.get("RECEIPT NO", ""),
        "CREDIT NOTES": sus.get("CREDIT NOTES", ""),
        "DETAIL RINCIAN NO": sus.get("DETAIL RINCIAN NO", ""),
        "RECEIPT DATE": sus.get("RECEIPT DATE", ""),
        "CEDANT NAME": sus.get("CEDANT NAME", ""),
        "CEDANT SHRT NAME": sus.get("CEDANT SHRT NAME", ""),
        "INSURED_ORI": sus.get("INSURED", ""),
        "INSURED_1": sus.get("INSURED_CLEAN_1", ""),
        "INSURED_2": sus.get("INSURED_CLEAN_2", ""),
        "INSURED_3": sus.get("INSURED_CLEAN_3", ""),
        "INSURED_4": sus.get("INSURED_CLEAN_4", ""),
        "INSURED_5": sus.get("INSURED_CLEAN_5", ""),
        "CURR ORI": sus.get("CURR ORI", ""),
        "AMOUNT ORI": sus.get("AMOUNT ORI", ""),
        "CURR PAY": sus.get("CURR PAY", ""),
        "AMOUNT PAY": sus.get("AMOUNT PAY", ""),
        "CCOS_OR_BAL": ccos["CCOS_OR_BAL"],
        "CCOS_BAL_DUE": ccos["CCOS_BAL_DUE"],
        "POLIS_ORI": sus.get("POLIS", ""),
        "SLIP_NO_ORI": sus.get("SLIP NO", ""),
        "SLIP_NO_CLN": sus.get("SLIP_NO_CLEAN", ""),
        "DESC 1": sus.get("DESC 1", ""), "DESC 2": sus.get("DESC 2", ""),
        "DESC 3": sus.get("DESC 3", ""), "DESC 4": sus.get("DESC 4", ""),
        "STATUS": sus.get("STATUS", ""),
        "REC_TYPE": sus.get("REC_TYPE", ""),
        "SKENARIO": final_scenario,
        "_OSBAL_ROW_COUNT": osbal_count,
        "_SUSPEND_ROW_COUNT": suspend_count,
    }
    for n in range(1, 6):
        pol = _text_or_empty(sus.get(f"POLIS_CLEAN_{n}", ""))
        if n == 1 and not pol:
            pol = _text_or_empty(sus.get("POLIS", ""))
        row[f"POLICY_CLEAN_{n}"] = pol
        if n <= 3:
            row[f"SERTIF_CLEAN_{n}"] = _format_sertif(sus.get(f"CERTIFICATE_{n}", ""))
    return row


def _merge_suspend_rows(rows: list) -> dict:
    # Menggabungkan data baris suspend yang memiliki fac code sama dalam akumulasi
    if len(rows) == 1:
        return rows[0]
    merged: dict = {}
    for col in _SUSPEND_SUM_COLS:
        merged[col] = sum(_to_float(r.get(col, 0)) for r in rows)
    for col in _SUSPEND_JOIN_COLS:
        seen, vals = set(), []
        for r in rows:
            v = str(r.get(col, "")).strip()
            if v and v not in seen:
                seen.add(v)
                vals.append(v)
        merged[col] = ", ".join(vals)
    return merged


def _compute_derived_cols(df: pd.DataFrame) -> pd.DataFrame:
    # Menghitung selisih nominal dan menentukan klasifikasi status FLAG_PROD
    amount_base = df.get("AMOUNT ORI", pd.Series([0] * len(df), index=df.index))
    if "_MERGED_AMOUNT_ORI" in df.columns:
        amount_raw = df["_MERGED_AMOUNT_ORI"].combine_first(amount_base)
    else:
        amount_raw = amount_base
    amount_numeric = pd.to_numeric(amount_raw.replace(r'^\s*$', np.nan, regex=True), errors="coerce")

    df["AMOUNT_ORI_MIN1"] = np.where(amount_numeric.notna(), amount_numeric * -1, np.nan)
    amount_ori = amount_numeric.fillna(0)
    amount_neg = df["AMOUNT_ORI_MIN1"].fillna(0)

    balance_due = pd.to_numeric(df.get("CCOS_BAL_DUE", 0), errors="coerce").fillna(0)
    osbal_count_col = pd.to_numeric(df.get("_OSBAL_ROW_COUNT", 0), errors="coerce").fillna(0)
    sus_count_col = pd.to_numeric(df.get("_SUSPEND_ROW_COUNT", 1), errors="coerce").fillna(1)
    ccos_ref = df.get("CCOS_REF_CODE", pd.Series("", index=df.index)).fillna("").astype(str)

    df["DIFERENCE"] = amount_neg - balance_due
    accumulated = (osbal_count_col > 1) | (sus_count_col > 1)
    is_equal = amount_neg.round(2) == balance_due.round(2)

    is_beda_curr = df["_IS_BEDA_CURR"].fillna(False).astype(bool) if "_IS_BEDA_CURR" in df.columns \
        else pd.Series(False, index=df.index)
    is_beda_per = df["_IS_BEDA_PERIODE"].fillna(False).astype(bool) if "_IS_BEDA_PERIODE" in df.columns \
        else pd.Series(False, index=df.index)

    is_gt1_fac = ccos_ref.str.contains(",", na=False)
    is_new_entry_total = (amount_ori != 0) & (balance_due == 0)
    is_adj_total = is_equal
    is_adj_sebagian = ~is_equal & (amount_neg < balance_due)
    is_new_entry_sebagian = ~is_equal & (amount_neg > balance_due)

    df["FLAG_PROD"] = np.select(
        [is_beda_curr, is_beda_per, is_gt1_fac,
         is_new_entry_total,
         is_adj_total & ~accumulated, is_adj_total & accumulated,
         is_adj_sebagian & ~accumulated, is_adj_sebagian & accumulated,
         is_new_entry_sebagian],
        ["Beda Currency", "Beda Periode", "Matching >1 fac code",
         "New Entry total",
         "Adjustment total tanpa akumulasi", "Adjustment total dengan akumulasi",
         "Adjustment sebagian tanpa akumulasi", "Adjustment sebagian dengan akumulasi",
         "New Entry sebagian"],
        default="Unmatching",
    )

    for col in ["_OSBAL_ROW_COUNT", "_SUSPEND_ROW_COUNT", "_MERGED_AMOUNT_ORI", "_IS_BEDA_CURR", "_IS_BEDA_PERIODE"]:
        if col in df.columns:
            df.drop(columns=[col], inplace=True)
    return df


def _format_skenario(row) -> str:
    # Memformat teks deskripsi skenario berdasarkan hasil status akhir
    sken, flag = str(row["SKENARIO"]), str(row["FLAG_PROD"])
    if sken in ("UNMATCHED", "Unmatching"):
        return "Unmatching"
    if flag == "Matching >1 fac code":
        return f"{sken} (> 1 Fac)"
    return f"{sken} ({flag})"


def _load_excel(path: str, label: str, optional: bool = False, header: int = 0) -> tuple:
    # Membaca data excel atau memuat dari pickle cache jika tersedia
    if not os.path.exists(path):
        print(f"  {label}: {path} ({'tidak ada, dilewati' if optional else 'TIDAK DITEMUKAN'})", flush=True)
        return [], []

    cache = path + (".cache.pkl" if header == 0 else f".h{header}.cache.pkl")
    if os.path.exists(cache) and os.path.getmtime(cache) >= os.path.getmtime(path):
        t = time.perf_counter()
        print(f"  {label}: loading dari cache ...", flush=True)
        try:
            data, hdr = pd.read_pickle(cache)
            print(f"  -> {len(data):,} rows, {len(hdr)} cols ({time.perf_counter()-t:.2f}s)", flush=True)
            return data, hdr
        except Exception:
            print(f"  {label}: cache rusak, baca ulang dari Excel ...", flush=True)

    t = time.perf_counter()
    print(f"  {label}: {path} ...", flush=True)
    try:
        df = pd.read_excel(path, header=header)
    except PermissionError:
        if os.path.exists(cache):
            print(f"  [WARN] {label} sedang dibuka di Excel, fallback membaca dari cache ...", flush=True)
            data, hdr = pd.read_pickle(cache)
            return data, hdr
        import shutil, tempfile
        tmp = os.path.join(tempfile.gettempdir(), os.path.basename(path))
        shutil.copy2(path, tmp)
        df = pd.read_excel(tmp, header=header)
        try:
            os.remove(tmp)
        except Exception:
            pass
    except Exception as e:
        print(f"  [RETRY] engine default gagal ({e}), coba openpyxl ...", flush=True)
        try:
            df = pd.read_excel(path, header=header, engine="openpyxl")
        except PermissionError:
            if os.path.exists(cache):
                print(f"  [WARN] {label} sedang dibuka di Excel, fallback membaca dari cache ...", flush=True)
                data, hdr = pd.read_pickle(cache)
                return data, hdr
            raise

    hdr = list(df.columns)
    data = df.to_dict("records")
    print(f"  -> {len(data):,} rows, {len(hdr)} cols ({time.perf_counter()-t:.1f}s)", flush=True)
    try:
        pd.to_pickle((data, hdr), cache)
        print(f"  -> cache disimpan: {cache}", flush=True)
    except Exception:
        pass
    return data, hdr


def _prepare_slipdb(rows: list) -> None:
    # Menyiapkan kolom kunci alfanumerik pada data referensi R/I Slip DB
    for r in rows:
        r["_SLIP_KEY"] = _alnum_key(r.get(SLIPDB_RAW_SLIP, ""))
        r["_POLIS_KEY"] = _alnum_key(r.get(SLIPDB_RAW_POLIS, ""))
        ins = _normalize(r.get(SLIPDB_RAW_INSURED, ""))
        ins = re.sub(r"^(PT|CV|UD|PD)\.?\s+|,\s*(PT|CV)\.?$", "", ins).strip()
        r["_INSURED_KEY"] = ins


def run() -> None:
    # Menjalankan seluruh alur proses pencocokan data suspend Mitrautama
    t_start = time.perf_counter()
    print("\n" + "=" * 60, flush=True)
    print("  SUSPEND MATCHING  (MITRAUTAMA) -- matching_mure", flush=True)
    print("=" * 60, flush=True)

    # 1. Memuat seluruh sumber data
    print("\n[1/5] Loading data ...", flush=True)
    suspend_rows, suspend_cols = _load_excel(SUSPEND_FILE, "SUSPEND")
    osbal_rows, osbal_cols = _load_excel(OSBAL_FILE, "OSBAL")
    facul_rows, facul_cols = _load_excel(FACUL_FILE, "FACUL")
    slipdb_rows, slipdb_cols = _load_excel(SLIPDB_FILE, "SLIPDB", optional=True, header=SLIPDB_HEADER_ROW)

    if not suspend_rows or not osbal_rows:
        print("  [ERROR] Data Suspend / Osbal kosong. Proses dihentikan.", flush=True)
        return

    # Validasi kelengkapan kolom kunci
    for label, cols, need in [
        ("SUSPEND", suspend_cols, [SUSPEND_SLIP_PREFIX, SUSPEND_INSURED_PREFIX, SUSPEND_GROUP_COL,
                                   SUSPEND_CURR_COL, SUSPEND_DATE_COL, "AMOUNT ORI"]),
        ("OSBAL", osbal_cols, [OSBAL_FACODE_COL, OSBAL_CURR_COL, OSBAL_DATE_COL,
                               "CCOS_DOC_NO", "CCOS_OR_BAL", "CCOS_BAL_DUE"]),
    ]:
        miss = [c for c in need if not any(str(x).startswith(c) for x in cols)]
        print(f"  Cek kolom {label}: {'OK' if not miss else '[MISSING] ' + str(miss)}", flush=True)

    if slipdb_rows:
        _prepare_slipdb(slipdb_rows)

    # Parsing tanggal pada data Osbal
    for r in osbal_rows:
        raw = r.get(OSBAL_DATE_COL, "")
        r["_com_date_parsed"] = pd.to_datetime(raw, errors="coerce") \
            if raw is not None and not (isinstance(raw, float) and pd.isna(raw)) and raw != "" else None

    # Parsing tanggal dan evaluasi status lineslip pada Suspend
    lineslip_count = 0
    for r in suspend_rows:
        raw = r.get(SUSPEND_DATE_COL, "")
        parsed = pd.to_datetime(raw, errors="coerce") \
            if raw is not None and not (isinstance(raw, float) and pd.isna(raw)) and raw != "" else None
        r["_sus_date_parsed"] = parsed
        r["_is_lineslip"] = _is_lineslip_row(r)
        if r["_is_lineslip"]:
            lineslip_count += 1
            r["_sus_date_lineslip"] = (parsed - pd.DateOffset(months=1)) \
                if parsed is not None and pd.notna(parsed) else None
        else:
            r["_sus_date_lineslip"] = None
    print(f"  -> {lineslip_count:,} baris LINESLIP terdeteksi", flush=True)

    # 2. Mendeteksi nama kolom pada dataset suspend
    print("\n[2/5] Detecting columns ...", flush=True)
    slip_sus = _get_cols(suspend_cols, SUSPEND_SLIP_PREFIX)
    insured_sus = _get_cols(suspend_cols, SUSPEND_INSURED_PREFIX)
    polis_sus = _get_cols(suspend_cols, SUSPEND_POLIS_PREFIX)
    print(f"  slip={slip_sus} insured={insured_sus} polis={polis_sus}", flush=True)

    # 3. Membangun indeks lookup pencarian
    print("\n[3/5] Building lookup indexes ...", flush=True)
    lk_osbal = _build_lookup(osbal_rows, OSBAL_SLIP_COLS, OSBAL_POLIS_COLS, OSBAL_INSURED_COLS,
                             OSBAL_FACODE_COL, "OSBAL")
    lk_facul = _build_lookup(facul_rows, FACUL_SLIP_COLS, FACUL_POLIS_COLS, FACUL_INSURED_COLS,
                             FACUL_FACODE_COL, "FACUL")
    lk_slipdb = _build_lookup(slipdb_rows, SLIPDB_SLIP_COLS, SLIPDB_POLIS_COLS, SLIPDB_INSURED_COLS,
                              SLIPDB_FACODE_COL, "SLIPDB")
    facode_osbal_idx = lk_osbal["facode"]

    # Cache lokal untuk pencarian referensi silang
    _via_cache: dict = {}

    def _via_refs(kind: str, queries: list, like: bool = False) -> tuple:
        # Melakukan pencarian data referensi dengan memanfaatkan cache lokal
        if not queries:
            return set(), None
        ckey = (kind, like, tuple(queries))
        if ckey in _via_cache:
            return set(_via_cache[ckey][0]), _via_cache[ckey][1]
        res = _via_refs_raw(kind, queries, like)
        _via_cache[ckey] = (frozenset(res[0]), res[1])
        return res

    def _via_refs_raw(kind: str, queries: list, like: bool) -> tuple:
        # Eksekusi pencarian kandidat ke Osbal, SlipDB, dan Facul
        if like:
            def finder(rows, lk, cols):
                if not lk["tok_insured"]:
                    return []
                return _like_match(queries, lk["tok_insured"], rows, cols)
        else:
            def finder(rows, lk, cols):
                return _exact_match(queries, lk[kind])

        hits = finder(osbal_rows, lk_osbal, OSBAL_INSURED_COLS)
        if hits:
            return set(hits), "OSBAL"
        for rows, lk, fcol, icols, src in [
            (slipdb_rows, lk_slipdb, SLIPDB_FACODE_COL, SLIPDB_INSURED_COLS, "SLIPDB"),
            (facul_rows, lk_facul, FACUL_FACODE_COL, FACUL_INSURED_COLS, "FACUL"),
        ]:
            if not rows:
                continue
            h = finder(rows, lk, icols)
            if h:
                res = _resolve_facode([rows[i] for i in h], fcol, facode_osbal_idx)
                if res:
                    return set(res), src
        return set(), None

    def _row_values_keyed(row: dict, cols: list) -> set:
        # Helper set nilai alfanumerik per baris
        return set(_collect_values(row, cols, _alnum_key))

    def _narrow_combo(cands: set, scenario: str, sus: dict, used: str, order: list) -> tuple:
        # Menyempitkan kandidat multi-fac dengan parameter sekunder
        def n_fac(idxs):
            return len({_normalize(osbal_rows[i].get(OSBAL_FACODE_COL, "")) for i in idxs})

        cur = list(cands)
        for crit in order:
            if n_fac(cur) <= 1:
                break
            if crit == "polis":
                vals = _collect_values(sus, polis_sus, _alnum_key)
                conf = [i for i in cur if set(vals) & _row_values_keyed(osbal_rows[i], OSBAL_POLIS_COLS)] if vals else []
                tag = "Polis"
            elif crit == "slip":
                vals = _collect_values(sus, slip_sus, _alnum_key)
                conf = [i for i in cur if set(vals) & _row_values_keyed(osbal_rows[i], OSBAL_SLIP_COLS)] if vals else []
                tag = "Slip"
            else:
                vals = _collect_values(sus, insured_sus)
                conf = [i for i in cur if _row_has_insured(osbal_rows[i], OSBAL_INSURED_COLS, vals)] if vals else []
                tag = "Insured"
            if conf and n_fac(conf) < n_fac(cur):
                cur = conf
                if tag.lower() not in scenario.lower():
                    scenario = f"{scenario.replace(' only', '')} + {tag}"
        return set(cur), scenario

    # 4. Melakukan matching iteratif per baris suspend
    total = len(suspend_rows)
    print(f"\n[4/5] Matching {total:,} suspend rows ...", flush=True)
    t4 = time.perf_counter()
    raw_results = []

    for n, sus in enumerate(suspend_rows, 1):
        if n % 500 == 0:
            el = time.perf_counter() - t4
            rate = n / el if el > 0 else 0
            print(f"    Memproses baris {n:,} / {total:,}  ({rate:.0f} rows/s, ETA {(total-n)/rate if rate else 0:.0f}s)", flush=True)

        if str(sus.get("STATUS", "")).strip().upper() == "MATCHING":
            continue

        is_ls = sus.get("_is_lineslip", False)
        eff_date = sus.get("_sus_date_lineslip") if is_ls else sus.get("_sus_date_parsed")

        # Pencocokan via slip
        def phase_slip():
            vals = [v for v in _collect_values(sus, slip_sus, _alnum_key) if len(v) >= _MIN_KEY_LEN]
            if not vals:
                return set(), None, ""
            idx, src = _via_refs("slip", vals)
            if idx:
                return idx, src, "Slip only"
            idx = _exact_match(vals, lk_osbal["polis"])
            if idx:
                return set(idx), "OSBAL", "Slip only"
            return set(), None, ""

        # Pencocokan via polis
        def phase_polis():
            vals = [v for v in _collect_values(sus, polis_sus, _alnum_key) if len(v) >= _MIN_KEY_LEN]
            if not vals:
                return set(), None, ""
            idx, src = _via_refs("polis", vals)
            if idx:
                return idx, src, "Polis only"
            idx = _exact_match(vals, lk_osbal["slip"])
            if idx:
                return set(idx), "OSBAL", "Polis to slip"
            return set(), None, ""

        # Pencocokan via insured
        def phase_insured():
            grp_tokens = _tokenize(_group_name(sus)) - _INSURED_STOPWORDS

            def _group_guard(idx_set: set) -> set:
                if len(grp_tokens) < 2 or not idx_set:
                    return idx_set
                filtered = set()
                for i in idx_set:
                    for col in OSBAL_INSURED_COLS:
                        rv = _normalize(osbal_rows[i].get(col, ""))
                        if not rv or not _valid(rv):
                            continue
                        if grp_tokens & (_tokenize(rv) - _INSURED_STOPWORDS):
                            filtered.add(i)
                            break
                return filtered

            for stage in _split_insured_queries(sus, insured_sus):
                idx, src = _via_refs("insured", stage)
                if not idx:
                    idx, src = _via_refs("insured", stage, like=True)
                    if idx:
                        idx = _group_guard(idx)
                if idx:
                    return idx, src, "Insured only"
            return set(), None, ""

        # Pencocokan via group name
        def phase_group():
            g = _group_name(sus)
            if not g:
                return set(), None, ""
            idx, src = _via_refs("insured", [g], like=True)
            return (idx, src, "Insured only") if idx else (set(), None, "")

        phases = [("slip", phase_slip), ("polis", phase_polis), ("insured", phase_insured), ("group", phase_group)]

        accepted = None
        fallback = None
        for used, fn in phases:
            cands, src, scen = fn()
            if not cands:
                continue
            order = {"slip": ["polis", "insured"],
                     "polis": ["slip", "insured"],
                     "insured": ["slip", "polis"],
                     "group": ["polis", "insured", "slip"]}[used]
            cands, scen = _narrow_combo(cands, scen, sus, used, order)
            sel, b_curr, b_per = _narrow_currency_periode(sus, list(cands), osbal_rows, eff_date, is_ls)
            if not b_curr and not b_per:
                accepted = (set(sel), src, scen, False, False)
                break
            if fallback is None:
                fallback = (set(sel), src, scen, b_curr, b_per)

        matched, source, scenario, is_beda_curr, is_beda_per = accepted or fallback or (set(), None, "Unmatching", False, False)
        if not matched:
            source, scenario = None, "Unmatching"

        osbal_ref = [osbal_rows[i] for i in matched] if matched else [{}]
        osbal_count = 0 if (is_beda_curr or is_beda_per or not matched) else len(osbal_ref)
        fac_codes = {_normalize(r.get(OSBAL_FACODE_COL, "")) for r in osbal_ref if r}
        fac_codes.discard("")

        raw_results.append({
            "suspend": sus, "source": source, "scenario": scenario,
            "osbal_idx": list(matched), "osbal_count": osbal_count,
            "fac_codes": fac_codes,
            "_is_beda_curr": is_beda_curr, "_is_beda_periode": is_beda_per,
        })

    print(f"  Matching selesai dalam {time.perf_counter()-t4:.1f}s", flush=True)

    # 4b. Pengelompokan fac code tunggal dan akumulasi nominal
    print("\n[4b/5] Fac code claim & akumulasi ...", flush=True)
    _EXCL = {"Unmatching", "UNMATCHED"}

    def _bad(r):
        return r["_is_beda_curr"] or r["_is_beda_periode"] or r["scenario"] in _EXCL or r["source"] is None

    groups: dict = {}
    final: list = []
    for r in raw_results:
        r.setdefault("suspend_count", 1)
        if len(r["fac_codes"]) == 1 and not _bad(r):
            groups.setdefault(next(iter(r["fac_codes"])), []).append(r)
        else:
            final.append(r)

    # Menggabungkan saldo suspend untuk fac code yang sama
    for fac, group in groups.items():
        if len(group) == 1:
            group[0]["suspend_count"] = 1
            final.append(group[0])
            continue
        merged_idx = set()
        for g in group:
            merged_idx.update(g["osbal_idx"])
        merged_sus = _merge_suspend_rows([g["suspend"] for g in group])
        for g in group:
            g["osbal_idx"] = list(merged_idx)
            g["osbal_count"] = len(merged_idx)
            g["suspend_count"] = len(group)
            g["_merged_amount_ori"] = merged_sus.get("AMOUNT ORI", 0)
            final.append(g)
    print(f"  {len(raw_results):,} suspend rows -> {len(final):,} output rows", flush=True)

    # 5. Membangun baris akhir output
    print(f"\n[5/5] Building output ({len(final):,} rows) ...", flush=True)
    output_rows = []
    for r in final:
        ref_osbal = [osbal_rows[i] for i in r["osbal_idx"]] if r["source"] else [{}]
        all_fac_rows = None
        osbal_cnt = r["osbal_count"]
        if r["source"] and len(r["fac_codes"]) == 1 and not (r["_is_beda_curr"] or r["_is_beda_periode"]):
            fac = next(iter(r["fac_codes"]))
            all_idx = facode_osbal_idx.get(fac, [])
            if all_idx:
                all_fac_rows = [osbal_rows[i] for i in all_idx]
                osbal_cnt = len(all_idx)

        out = _build_output_row(
            r["suspend"], r["source"], r["scenario"], ref_osbal,
            osbal_count=osbal_cnt, suspend_count=r.get("suspend_count", 1),
            all_fac_osbal_rows=all_fac_rows,
        )
        if r["_is_beda_curr"] or r["_is_beda_periode"]:
            out["CCOS_OR_BAL"] = np.nan
            out["CCOS_BAL_DUE"] = np.nan
        if "_merged_amount_ori" in r:
            out["_MERGED_AMOUNT_ORI"] = r["_merged_amount_ori"]
        out["_IS_BEDA_CURR"] = r["_is_beda_curr"]
        out["_IS_BEDA_PERIODE"] = r["_is_beda_periode"]
        output_rows.append(out)

    df = pd.DataFrame(output_rows)
    for col in ["AMOUNT ORI", "CCOS_OR_BAL", "CCOS_BAL_DUE"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    df = _compute_derived_cols(df)
    df.loc[df["SKENARIO"].isin(["UNMATCHED", "Unmatching"]), "FLAG_PROD"] = "Unmatching"

    # Penyesuaian nama skenario dengan format label akhir
    cond = df["FLAG_PROD"].isin(["Beda Currency", "Beda Periode", "Matching >1 fac code"]) \
        & ~df["SKENARIO"].isin(["UNMATCHED", "Unmatching"])
    if cond.any():
        df.loc[cond, "SKENARIO"] = df.loc[cond].apply(_format_skenario, axis=1)

    df["SKENARIO"] = df["SKENARIO"].str.replace(r"\bGroup\b", "Insured", regex=True)

    # Menjamin urutan dan keberadaan seluruh kolom final
    for col in FINAL_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    df = df[FINAL_COLUMNS]

    # Menyimpan file excel hasil matching
    print(f"\nSaving to: {OUTPUT_FILE} ...", flush=True)
    for c in df.select_dtypes(include=["object", "string"]).columns:
        df[c] = df[c].astype(str).str.replace(r'[\x00-\x08\x0B\x0C\x0E-\x1F]', '', regex=True)
        df[c] = df[c].replace("nan", "")

    save_path = OUTPUT_FILE
    try:
        df.to_excel(save_path, index=False)
    except PermissionError:
        save_path = os.path.join(_BASE_DIR, "matching_mitrautama_v1_new.xlsx")
        print(f"  [PERINGATAN] {OUTPUT_FILE} sedang dibuka di Excel! Menyimpan ke: {save_path} ...", flush=True)
        df.to_excel(save_path, index=False)

    # Format cell teks pada openpyxl
    try:
        from openpyxl import load_workbook
        wb = load_workbook(save_path)
        ws = wb.active
        idxs = [i for i, c in enumerate(df.columns, start=1) if c in _TEXT_COLS]
        for ci in idxs:
            for rw in range(2, ws.max_row + 1):
                cell = ws.cell(row=rw, column=ci)
                cell.number_format = "@"
                if cell.value is not None:
                    cell.value = str(cell.value)
        wb.save(save_path)
    except Exception as e:
        print(f"  [WARN] Format text gagal: {e}")

    elapsed = time.perf_counter() - t_start
    print(f"\n{'=' * 60}")
    print(f"  Selesai dalam {elapsed:.1f}s ({elapsed/60:.1f} menit)")
    print(f"  Rows  : {len(df):,}  |  Cols: {len(df.columns)}  |  File: {save_path}")
    print("\n  FLAG_PROD:")
    for flag, count in df["FLAG_PROD"].value_counts().items():
        print(f"    {flag:<45}: {count:,}")
    print("\n  SKENARIO:")
    for sce, count in df["SKENARIO"].value_counts().items():
        print(f"    {sce:<50}: {count:,}")
    print("=" * 60)


if __name__ == "__main__":
    run()