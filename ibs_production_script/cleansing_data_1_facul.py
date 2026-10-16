"""
================================================================================
 SCRIPT CLEANSING DATA 1 - FACULTATIVE (CEDANT : IBS)
================================================================================
"""

import sys
sys.modules['numexpr'] = None
sys.modules['bottleneck'] = None

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
_BASE_DIR   = os.path.dirname(os.path.abspath(__file__))  # folder ibs_production_script/
_RAW_DIR    = os.path.join(_BASE_DIR, "data_raw")
_CLEAN_DIR  = os.path.join(_BASE_DIR, "data_cleaned")

INPUT_FILE  = os.path.join(_RAW_DIR,   "1b. Transaksi Facul 01.01.23 - 17.08.2026.xlsx")
SHEET_NAME  = "Query result"
HEADER_ROW  = 0
OUTPUT_FILE = os.path.join(_CLEAN_DIR, "ibs_output_facul_V2.xlsx")
CEDANT_FILTER = "IBS"

CEDANT_COLUMN = "COMP_NAME"
BROKER_NAME_COLUMN_CANDIDATES = ["COMP_NAME2", "COMP_NAME.1", "COMP_NAME_1"]
DIRECT_MARKER = "DIRECT"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"

MAX_BREAKDOWN_CODES = 5
TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CERTIFICATE",
)

# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
_TRANSACTION_HEADER_RE = re.compile(r"^\s*(?:Pembayaran|Penagihan)?\s*(?:dan\s+Penagihan)?\s*Premi\s*(?:Fakultatif|SOA|Pensesian)?\s*(?:PAR\s*&?\s*EQ|BIT|GIT|CECR|QS)?\s*", flags=re.IGNORECASE)
_TRUNCATE_TRIGGER_RE = re.compile(r"\bsubsidiar|\b(?:and\s*/\s*or|and|its|their|all)\s+associat(?!ion)|&\s*/\s*or\s+associat(?!ion)|\baffiliat|\bfiliated\b|related\s+compan|respective\s+right|\bincluding\s+an[yd]|\bindu?ding\s+any|as\s+the\s+owner|as\s+owners?\b|as\s+main\s+contractor|\(as\s+mortgagee\s+and\s+loss\s+payee\)|\b(?:their|tehir|its)\s+client", flags=re.IGNORECASE)

def _truncate_from_first_trigger(text: str) -> str:
    match = _TRUNCATE_TRIGGER_RE.search(text)
    return text[: match.start()] if match else text

_REMOVE_ONLY_RE = re.compile(r"\bas\s+principal\b|\bsebagai\s+principal\b|\bsebagai\s+kontraktor\b", flags=re.IGNORECASE)
_GENERIC_BOILERPLATE_TAIL_RE = re.compile(r"\bsub\s*-?\s*kontraktor\b.*$|\banak\s+perusahaan\b.*$|\bafiliasi\s+perusahaan\b.*$|\bdan\s+semua\b\s*[–-]?\s*$", flags=re.IGNORECASE)
_TITLE_RE = re.compile(r"\b(?:S\.?H\.?|S\.?I\.?Kom\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b", flags=re.IGNORECASE)
_LEGAL_ENTITY_RE = re.compile(r"\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PP\.?|PERSERO|Perseroan\s+Terbatas|LTD\.?|PTE\.?|LIMITED)\b", flags=re.IGNORECASE)

_ENTITY_SPLIT_RE = re.compile(r"\b(?:QQ+|Q\.Q\.)\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|\bin\s+this\s+case\b|\bdalam\s+hal\s+ini\b|;|/|¿|:|,", flags=re.IGNORECASE)
_BARE_N_SEPARATOR_RE = re.compile(r"(?<=[A-Z])\s+N\s+(?=(?!QQ\b)[A-Z]{2,})", flags=re.IGNORECASE)

def _normalize_bare_n_separator(text: str) -> str:
    return _BARE_N_SEPARATOR_RE.sub(" / ", text)

_BARE_SUFFIX_FRAGMENT_RE = re.compile(r"^\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PP\.?|PERSERO|LTD\.?|PTE\.?|LIMITED)\s*$", flags=re.IGNORECASE)
_BARE_SUFFIX_WITH_PAREN_RE = re.compile(r"^\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PP\.?|PERSERO|LTD\.?|PTE\.?|LIMITED)?\s*(\uE100\d+\uE101)\s*$", flags=re.IGNORECASE)

def _merge_bare_suffix_fragments(parts: list) -> list:
    merged = []
    for p in parts:
        p_stripped = p.strip()
        if merged and (_BARE_SUFFIX_FRAGMENT_RE.match(p_stripped) or _BARE_SUFFIX_WITH_PAREN_RE.match(p_stripped)):
            merged[-1] = merged[-1] + " " + p_stripped
        else:
            merged.append(p)
    return merged

_HEADER_BEFORE_COLON_RE = re.compile(r"^[^:]*:\s*")

def _strip_list_header_before_colon(text: str) -> str:
    if ":" not in text:
        return text
    head, _, tail = text.partition(":")
    if tail.count(",") >= 1:
        return tail.strip()
    return text

_INSURED_MONTHS_RE = r"JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?"
_INSURED_TRAILING_MONTH_YEAR_RE = re.compile(rf"\s*/\s*(?:{_INSURED_MONTHS_RE})\s+\d{{4}}\s*$", re.IGNORECASE)

def _strip_trailing_month_year(text: str) -> str:
    return _INSURED_TRAILING_MONTH_YEAR_RE.sub("", text)

_INSURED_VARIOUS_STANDALONE_RE = re.compile(r"^\s*VARIOUS\s*$", flags=re.IGNORECASE)
_INSURED_BATCH_NOISE_RE = re.compile(r"\b(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)?\s*\bBATCH\.?\s*\d*(?:\s*(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)\b)?|\d+(?=\s*(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)\b)", flags=re.IGNORECASE)

def _strip_insured_batch_noise(text: str) -> str: 
    return _INSURED_BATCH_NOISE_RE.sub(" ", text)

_STRAY_SYMBOL_RE = re.compile(r"[¡‽]")
def _strip_stray_symbols(text: str) -> str: 
    return _STRAY_SYMBOL_RE.sub(" ", text)

_PAREN_GROUP_RE = re.compile(r"\([^()]*\)")
_PAREN_NOISE_RE = re.compile(r"\(\s*\d+\s*CURRENC(?:Y|IES)\s*\)", re.IGNORECASE)

def _strip_paren_noise(text: str) -> str:
    return _PAREN_NOISE_RE.sub(" ", text)

_PAREN_STASH: dict[str, str] = {}
_PAREN_TOKEN_RE = re.compile(r"\uE100(\d+)\uE101")

def _protect_parens_content(text: str) -> str:
    def _mask(match: "re.Match") -> str:
        idx, token = len(_PAREN_STASH), f"\uE100{len(_PAREN_STASH)}\uE101"
        _PAREN_STASH[token] = match.group(0)
        return token
    return _PAREN_GROUP_RE.sub(_mask, text)

def _restore_protected_chars(text: str) -> str:
    return text.replace("\uE0F0", "/").replace("\uE0F1", ",").replace("\uE0F2", ";")

def _restore_protected_parens(text: str) -> str:
    def _unmask(match: "re.Match") -> str:
        return _PAREN_STASH.get(f"\uE100{match.group(1)}\uE101", match.group(0))
    return _PAREN_TOKEN_RE.sub(_unmask, text)

_INSURED_MERGE_OVERRIDES = {}
_LIPPO_KNOWN_ENTITIES = [
    "BERLIAN JAYA", "GLOBAL PANGAN NUSANTARA", "ADARO INDONESIA", "ADARO ENERGY",
    "BANK MESTIKA DHARMA", "BANK MESTIKA", "TRAKINDO UTAMA", "TIARA MARGA TRAKINDO",
]
_LIPPO_KNOWN_ENTITIES_SORTED = sorted(_LIPPO_KNOWN_ENTITIES, key=len, reverse=True)
_KNOWN_ENTITY_RE = re.compile("|".join(re.escape(name) for name in _LIPPO_KNOWN_ENTITIES_SORTED))
_ENTITY_SLASH_NORMALIZE_PATTERNS = [(name, re.compile(r"\b" + r"(?:\s*/\s*|\s+)".join(re.escape(w) for w in name.split()) + r"\b", flags=re.IGNORECASE)) for name in _LIPPO_KNOWN_ENTITIES_SORTED if len(name.split()) > 1]

def _normalize_entity_slash_variants(text: str) -> str:
    for canonical, pattern in _ENTITY_SLASH_NORMALIZE_PATTERNS: text = pattern.sub(canonical, text)
    return text

def _split_by_known_entities(text: str) -> list | None:
    pos, n, found = 0, len(text), []
    while pos < n:
        while pos < n and text[pos] == " ": pos += 1
        if pos >= n: break
        match = _KNOWN_ENTITY_RE.match(text, pos)
        if not match: return None
        found.append(match.group(0))
        pos = match.end()
    return found if len(found) > 1 else None

_INSURED_DESCRIPTIVE_RE = re.compile(r"\bVARIOUS\s+INSUREDS?\b|\bACCEPTED\s+BY\b", flags=re.IGNORECASE)

def _normalize_dash_entity_list(text: str) -> str:
    return text

def _strip_transaction_header(text: str) -> str: return _TRANSACTION_HEADER_RE.sub("", text)
def _strip_titles(text: str) -> str: return _TITLE_RE.sub(" ", text)
def _strip_legal_entity(text: str) -> str: return _LEGAL_ENTITY_RE.sub(" ", text)

def _truncate_boilerplate_tail(text: str) -> str:
    text = _truncate_from_first_trigger(text)
    match_generic = _GENERIC_BOILERPLATE_TAIL_RE.search(text)
    if match_generic: text = text[: match_generic.start()]
    return _REMOVE_ONLY_RE.sub(" ", text)

def _final_polish(text: str) -> str:
    t = re.sub(r"[,;:\"'.]", " ", text)
    t = re.sub(r"[-/¿]", " ", t)
    t = re.sub(r"\(\s+", "(", t)
    t = re.sub(r"\s+\)", ")", t)
    t = re.sub(r"\s+", " ", t).strip()
    dangling_pattern = r"^\b(?:AS|AN|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*|\s*\b(?:AS|AN|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$"
    previous = None
    while previous != t:
        previous, t = t, re.sub(dangling_pattern, "", t, flags=re.IGNORECASE).strip()
    return t.upper()

_BARE_CLIENT_FRAGMENT_RE = re.compile(r"^\s*CLIENTS?\s*$", flags=re.IGNORECASE)

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return []

    original = text.strip()
    if original in _INSURED_MERGE_OVERRIDES: return _INSURED_MERGE_OVERRIDES[original]
    if _INSURED_DESCRIPTIVE_RE.search(original): return [original]

    _PAREN_STASH.clear()

    t = _strip_transaction_header(original)
    t = _strip_trailing_month_year(t)
    t = _strip_list_header_before_colon(t)
    t = _strip_paren_noise(t)
    t = _normalize_bare_n_separator(t)
    t = _strip_titles(t)
    t = _strip_stray_symbols(t)
    t = _strip_insured_batch_noise(t)
    t = _normalize_dash_entity_list(t)
    t = _normalize_entity_slash_variants(t)
    t = _truncate_boilerplate_tail(t)
    t = _protect_parens_content(t)

    raw_parts = _merge_bare_suffix_fragments(_ENTITY_SPLIT_RE.split(t))
    results = []
    for part in raw_parts:
        part = _restore_protected_chars(part)
        cleaned = _truncate_boilerplate_tail(part)
        cleaned = _strip_legal_entity(cleaned)
        cleaned = _final_polish(cleaned)
        cleaned = _restore_protected_parens(cleaned)
        cleaned = re.sub(r"\s+", " ", cleaned).strip()

        if not cleaned or len(cleaned) < 2 or not re.search(r"[A-Z]", cleaned): continue
        if _BARE_CLIENT_FRAGMENT_RE.match(cleaned): continue

        known_split = _split_by_known_entities(cleaned)
        if known_split is not None:
            for name in known_split:
                if name not in results: results.append(name)
            continue
        if cleaned not in results: results.append(cleaned)

    if len(results) > MAX_BREAKDOWN_CODES:
        return [original]
    if not results:
        return [original.upper()]
    return results

# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO
# ==============================================================================
_STANDALONE_KEEP_RE = re.compile(r"^\s*(VAR|VARIOUS|TBA|P\d{0,2})\s*$", flags=re.IGNORECASE)
_PENYELESAIAN_RE = re.compile(r"PENYELESAIAN\s+(SUSPENSE|\(?H\)?UTANG\s+PIUTANG)", flags=re.IGNORECASE)

_MONTHS_RE = r"JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY|ERUARI|RURI)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|NOV(?:EMBER)?|DES(?:EMBER|MBER)?|DEC(?:EMBER)?"
_MONTH_YEAR_BLOCK_RE = re.compile(rf"\b(?:{_MONTHS_RE})\b(?:\s*[,/&]\s*(?:{_MONTHS_RE})\b)*\s*\.?\s*\d{{4}}\b", flags=re.IGNORECASE)
_MONTH_ALONE_RE = re.compile(rf"\b(?:{_MONTHS_RE})\b", flags=re.IGNORECASE)
_CURRENCY_NOISE_RE = re.compile(r"\b(?:USD|IDR|SGD|EUR)\b", flags=re.IGNORECASE)

_ADMIN_LABEL_NOISE_RE = re.compile(r"\bBORDERO\b|\bINTER\s+ISLAND\b|\bSINGLE\b|\bIMPORT\b|\bEXPORT\b|\bFPG\b|\bISL\b|\bNO\s*COVER\s*NOTE\b|\bDEKL\.?\b", flags=re.IGNORECASE)

_N_POLIS_LABEL_RE = re.compile(r"\d+\s*POLIS\b", flags=re.IGNORECASE)
_END_SUFFIX_RE = re.compile(r"\bEND\.?\s*\d*\b", flags=re.IGNORECASE)
_PLUS_PCODE_RE = re.compile(r"\+\s*P\d{1,2}\b", flags=re.IGNORECASE)
_NO_LABEL_PREFIX_RE = re.compile(r"^[.\s]*NO\s*[:.]\s*", flags=re.IGNORECASE)
_PAREN_GROUP_NOISE_RE = re.compile(r"\(\s*[^()]*\s*\)")
_BARE_LABEL_SUFFIX_RE = re.compile(r"\b(?:EXPAT|NATIONAL)\b", flags=re.IGNORECASE)

_TBA_WITH_SEPARATOR_RE = re.compile(r"(?:^|(?<=[/.]))\s*TBA\s*(?=[/.]|$)", flags=re.IGNORECASE)

def _strip_tba_with_separator(text: str) -> str:
    if _STANDALONE_KEEP_RE.match(text):
        return text
    if not re.search(r"TBA\s*[/.]|[/.]\s*TBA", text, flags=re.IGNORECASE):
        return text
    t = _TBA_WITH_SEPARATOR_RE.sub("", text)
    t = re.sub(r"^[/.\s]+", "", t)
    t = re.sub(r"[/.\s]+$", "", t)
    t = re.sub(r"\s*/\s*/\s*", "/", t)
    return t.strip()

def _strip_noise(text: str) -> str:
    t = text
    t = _NO_LABEL_PREFIX_RE.sub("", t)
    t = _PAREN_GROUP_NOISE_RE.sub(" ", t)
    t = _PLUS_PCODE_RE.sub(" ", t)
    t = _END_SUFFIX_RE.sub(" ", t)
    t = _BARE_LABEL_SUFFIX_RE.sub(" ", t)
    t = _MONTH_YEAR_BLOCK_RE.sub(" ", t)
    t = _MONTH_ALONE_RE.sub(" ", t)
    t = _CURRENCY_NOISE_RE.sub(" ", t)
    t = _ADMIN_LABEL_NOISE_RE.sub(" ", t)
    t = _N_POLIS_LABEL_RE.sub(" ", t)
    t = re.sub(r"[/,]\s*(?=[/,]|$)", " ", t)
    t = re.sub(r"^[/,\s]+|[/,\s]+$", "", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t

_TRAILING_WORD_AFTER_SLASH_RE = re.compile(r"^(?P<code>.*\d.*?)/(?P<word>[A-Za-z]{3,})\s*$")

def _strip_trailing_word_after_slash(text: str) -> str:
    match = _TRAILING_WORD_AFTER_SLASH_RE.match(text)
    if match:
        return match.group("code").strip()
    return text

def _looks_irregular(text: str) -> bool:
    return not re.search(r"\d", text)

_SEGMENT_SPLIT_RE = re.compile(r"[,+]")
_TRAILING_NUM_RE = re.compile(r"(\d+)\s*$")

def _reuse_prefix_breakdown(text: str):
    if not re.search(r"[,+]", text) and not re.search(r"\d\s*-\s*\d", text):
        return None

    segments = [s.strip() for s in _SEGMENT_SPLIT_RE.split(text) if s.strip()]
    if not segments:
        return None

    first = segments[0]
    range_match = re.search(r"^(.*?)(\d+)\s*-\s*(\d+)\s*$", first)
    if range_match and len(segments) == 1:
        prefix = range_match.group(1).strip()
        start_n, end_n = range_match.group(2), range_match.group(3)
        if len(end_n) <= len(start_n) - 3:
            return None
        codes = [f"{prefix} {start_n}".strip(), f"{prefix} {end_n}".strip()]
        return codes

    anchor_match = _TRAILING_NUM_RE.search(first)
    if not anchor_match:
        return None
    
    anchor_num = anchor_match.group(1)
    prefix_raw = first[: anchor_match.start()]
    glued = bool(re.search(r"[\-.]\s*$", prefix_raw))
    prefix = prefix_raw.strip()
    prefix = re.sub(r"[/,]\s*$", "", prefix).strip()
    joiner = "" if glued else " "
    prefix_for_join = prefix_raw.rstrip() if glued else prefix

    first_result = first.strip()
    if prefix_for_join.endswith("-") and prefix_for_join.count("-") == 1 and len(segments) > 2:
        prefix_for_join = prefix_for_join[:-1]
        joiner = ""
        first_result = f"{prefix_for_join}{anchor_num}"

    codes = [first_result]
    for seg in segments[1:]:
        seg_range = re.match(r"^(\d+)\s*-\s*(\d+)$", seg)
        if seg_range:
            for n in (seg_range.group(1), seg_range.group(2)):
                if len(n) > len(anchor_num):
                    return None
                codes.append(f"{prefix_for_join}{joiner}{n}".strip() if prefix else n)
            continue
        if not re.fullmatch(r"\d+", seg) or len(seg) > len(anchor_num):
            return None
        codes.append(f"{prefix_for_join}{joiner}{seg}".strip() if prefix else seg)

    if len(codes) < 2:
        return None
    return codes

_LONG_NUM_RE = re.compile(r"\d{6,}")

def _reuse_numeric_suffix(text: str):
    parts = [p.strip() for p in re.split(r"[,+]", text) if p.strip()]
    if len(parts) < 2:
        return None

    base_match = re.search(r"(\d{6,})\s*$", parts[0])
    if not base_match:
        return None

    anchor_prefix_text = parts[0][: base_match.start()]
    anchor_num = base_match.group(1)
    dash_is_glue = (
        anchor_prefix_text.endswith("-")
        and anchor_prefix_text.count("-") == 1
        and len(parts) > 2
    )
    if dash_is_glue:
        anchor_prefix_text = anchor_prefix_text[:-1]
    
    codes = [f"{anchor_prefix_text}{anchor_num}".strip() if dash_is_glue else parts[0].strip()]

    for p in parts[1:]:
        full_match = re.search(r"(\d{6,})\s*$", p)
        has_own_prefix = bool(full_match and full_match.start() > 0)
        if full_match and has_own_prefix and len(full_match.group(1)) >= len(anchor_num):
            anchor_prefix_text = p[: full_match.start()]
            anchor_num = full_match.group(1)
            if anchor_prefix_text.endswith("-") and dash_is_glue:
                anchor_prefix_text = anchor_prefix_text[:-1]
            codes.append(f"{anchor_prefix_text}{anchor_num}".strip() if dash_is_glue else p.strip())
            continue

        if not re.fullmatch(r"\d{1,%d}" % len(anchor_num), p):
            return None
        reconstructed_num = anchor_num[: -len(p)] + p
        codes.append(f"{anchor_prefix_text}{reconstructed_num}".strip())

    seen = set()
    result = []
    for c in codes:
        if c not in seen:
            seen.add(c)
            result.append(c)
    return result if len(result) >= 2 else None

_IBSRE_CODE_EXTRACT_RE = re.compile(r"(?:(?P<noiseword>TBA|NO)\.?\s+)?(?P<classprefix>(?:[A-Z]{1,10}\s+)?[A-Z0-9&]*[A-Z][A-Z0-9&]{0,9}[\s.]*(?:IBS\s*RE|IBSRESG|IBS|BSRE)[\s.]*\d{4}[\s.:]*)(?P<nums>[\d][\d,+&\-\s]*\d|\d)", flags=re.IGNORECASE)
_AIGMAP_CODE_EXTRACT_RE = re.compile(r"(?P<classprefix>AIG[\s\-]*MAP[\s\-]*)(?P<nums>[\d][\d,+\-\s]*\d|\d)", flags=re.IGNORECASE)

def _codes_from_classprefix_and_nums(classprefix: str, nums_raw: str):
    classprefix = re.sub(r"\s+", " ", classprefix)
    classprefix = re.sub(r"\s+:", " :", classprefix)
    joiner = ""
    nums_raw = nums_raw.strip()
    if _DAN_SEPARATOR_RE.search(nums_raw):
        nums_raw = _DAN_SEPARATOR_RE.sub(",", nums_raw)

    range_match = re.fullmatch(r"(\d+)\s*-\s*(\d+)", nums_raw)
    if range_match:
        start_n, end_n = range_match.group(1), range_match.group(2)
        codes = [f"{classprefix}{joiner}{start_n}".strip(), f"{classprefix}{joiner}{end_n}".strip()]
        return codes if len(codes) <= MAX_BREAKDOWN_CODES else None

    if re.search(r"[,+&]", nums_raw):
        raw_tokens = [t.strip() for t in re.split(r"[,+&]", nums_raw) if t.strip()]
        year_match = re.search(r"\b(\d{4})\b", classprefix)
        
        if year_match:
            year_str = year_match.group(1)
            cleaned_tokens = []
            for tok in raw_tokens:
                dup_year_match = re.match(r"^%s\s+(\d+)$" % re.escape(year_str), tok)
                cleaned_tokens.append(dup_year_match.group(1) if dup_year_match else tok)
            raw_tokens = cleaned_tokens
            
        codes = []
        anchor_num = None
        for tok in raw_tokens:
            tok_range = re.match(r"^(\d+)\s*-\s*(\d+)$", tok)
            if tok_range:
                codes.append(f"{classprefix}{joiner}{tok_range.group(1)}".strip())
                codes.append(f"{classprefix}{joiner}{tok_range.group(2)}".strip())
                anchor_num = tok_range.group(2)
                continue
                
            if anchor_num and re.fullmatch(r"\d+", tok) and len(tok) < len(anchor_num):
                reconstructed = anchor_num[: -len(tok)] + tok
                codes.append(f"{classprefix}{joiner}{reconstructed}".strip())
                anchor_num = reconstructed
                continue
                
            codes.append(f"{classprefix}{joiner}{tok}".strip())
            if re.fullmatch(r"\d+", tok):
                anchor_num = tok
                
        if not codes:
            return None
        return codes if len(codes) <= MAX_BREAKDOWN_CODES else None

    return [f"{classprefix}{joiner}{nums_raw}".strip()]

_IBSRE_VARIOUS_RE = re.compile(r"(?P<classprefix>[A-Z]{1,10}[\s.]*(?:IBS\s*RE|IBSRESG|BSRE)[\s.]*\d{4}[\s.:]*VARIOUS)", flags=re.IGNORECASE)

def _extract_ibsre_breakdown(original: str):
    various_match = _IBSRE_VARIOUS_RE.search(original)
    if various_match:
        code = re.sub(r"\s+", " ", various_match.group("classprefix")).strip()
        code = re.sub(r"\bVARIOUS\b", "Various", code, flags=re.IGNORECASE)
        return [code]

    match = _IBSRE_CODE_EXTRACT_RE.search(original)
    if not match:
        return None

    remainder = original[match.end():]
    if re.search(r"IBS\s*RE|IBSRESG|BSRE", remainder, re.IGNORECASE):
        return None

    result = _codes_from_classprefix_and_nums(match.group("classprefix"), match.group("nums"))
    return result if result is not None else [original]

def _extract_aigmap_breakdown(original: str):
    normalized = _DAN_SEPARATOR_RE.sub(",", original)
    match = _AIGMAP_CODE_EXTRACT_RE.search(normalized)
    if not match:
        return None

    classprefix = re.sub(r"\s+", "", match.group("classprefix")).strip()
    nums_raw = match.group("nums").strip()

    parts = [p.strip() for p in re.split(r"[,+]", nums_raw) if p.strip()]
    if not parts:
        return None

    anchor_num = parts[0]
    codes = [f"{classprefix}{anchor_num}"]
    for p in parts[1:]:
        if re.fullmatch(r"\d+", p) and len(p) < len(anchor_num):
            reconstructed = anchor_num[: -len(p)] + p
            codes.append(f"{classprefix}{reconstructed}")
        else:
            codes.append(f"{classprefix}{p}")

    return codes if len(codes) <= MAX_BREAKDOWN_CODES else [original]

_REAL_CODE_PATTERN_RE = re.compile(r"IBS\s*RE|IBSRESG|BSRE|MCOC|AIG[\s\-]*MAP|MD[\s\-]*C\d|11-F\d|[A-Z]{1,10}\s*IBS\s*\d{4}", flags=re.IGNORECASE)

def _has_real_code(text: str) -> bool:
    return bool(_REAL_CODE_PATTERN_RE.search(text))

_DAN_SEPARATOR_RE = re.compile(r"\bDAN\b|\bAND\b", flags=re.IGNORECASE)

def _normalize_dan_separator(text: str) -> str:
    return _DAN_SEPARATOR_RE.sub(",", text)

_MULTI_RANGE_RE = re.compile(r"^(?P<prefix>.*?)(?P<ranges>(?:\d+\s*-\s*\d+)(?:\s*\+\s*\d+\s*-\s*\d+)+)\s*$")

def _expand_multi_range(text: str):
    match = _MULTI_RANGE_RE.match(text)
    if not match:
        return None
    prefix = match.group("prefix").strip()
    prefix = re.sub(r"[/,]\s*$", "", prefix).strip()
    range_tokens = re.findall(r"(\d+)\s*-\s*(\d+)", match.group("ranges"))
    
    codes = []
    for start_n, end_n in range_tokens:
        codes.append(f"{prefix} {start_n}".strip() if prefix else start_n)
        codes.append(f"{prefix} {end_n}".strip() if prefix else end_n)
        
    seen, result = set(), []
    for c in codes:
        if c not in seen:
            seen.add(c); result.append(c)
    return result if len(result) >= 2 else None

_CHARACTERISTIC_WHITELIST_RE = re.compile(r"IBS\s*RE|IBSRESG|BSRE|AIG[\s\-]*MAP|\bOC\.?\s*\d{2}\.?\s*\d{2}\.?\s*\d{2}\.?\s*\d+\b", flags=re.IGNORECASE)

def _looks_like_characteristic_code(code: str) -> bool:
    if not isinstance(code, str) or not code.strip():
        return False
    c = code.strip()
    if not _CHARACTERISTIC_WHITELIST_RE.search(c):
        return False
    if "," in c:
        return False
    noise_re = re.compile(rf"\b(?:{_MONTHS_RE})\b|\bINTER\s+ISLAND\b|\bIMPORT\b|\bEXPORT\b|\bBORDERO\b", flags=re.IGNORECASE)
    if noise_re.search(c):
        return False
    return True

def _dempet_codes(codes: list) -> list:
    return [re.sub(r"[^A-Za-z0-9]", "", c) if isinstance(c, str) else c for c in codes]

def clean_split_code(text: str) -> list:
    results, should_dempet = _clean_split_code_impl(text)
    results = [r.upper() if isinstance(r, str) else r for r in results]
    if should_dempet:
        results = _dempet_codes(results)
    return results

def _clean_split_code_impl(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return [], False
    original = text.strip()

    if _STANDALONE_KEEP_RE.match(original):
        return [original.upper()], False
    if _PENYELESAIAN_RE.search(original):
        return [original], False

    trailing_end_stripped = re.sub(r"(?:\s*\+\s*END\s*\d*)+\s*$", "", original, flags=re.IGNORECASE).strip()
    if trailing_end_stripped != original and trailing_end_stripped:
        month_in_remainder = re.search(
            r"\b(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|"
            r"JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|"
            r"NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)\b",
            trailing_end_stripped, re.IGNORECASE,
        )
        if not month_in_remainder and not _has_real_code(trailing_end_stripped):
            return [trailing_end_stripped], False

    if not _has_real_code(original) and not re.search(r"TBA\s*[/.]|[/.]\s*TBA|^\s*TBA\s*[/.]|/\s*TBA\s*$", original, re.IGNORECASE):
        months_check = re.search(
            r"\b(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|"
            r"JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|"
            r"NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)\b|\bEND\b|\bSINGLE\s+RATE\b|\bINTER\s+ISLAND\b",
            original, re.IGNORECASE,
        )
        if months_check:
            return [original], False

    if re.search(r"IBS\s*RE|IBSRESG|BSRE|\b[A-Z]{1,10}\s*IBS\s*\d{4}", original, re.IGNORECASE):
        extracted = _extract_ibsre_breakdown(original)
        if extracted is not None:
            return extracted, all(_looks_like_characteristic_code(c) for c in extracted)

    if re.search(r"AIG[\s\-]*MAP", original, re.IGNORECASE):
        extracted = _extract_aigmap_breakdown(original)
        if extracted is not None:
            return extracted, all(_looks_like_characteristic_code(c) for c in extracted)

    t = _strip_tba_with_separator(original)
    if not t:
        return [original], False

    if "/" in t and not re.search(r"\d", t):
        return [t], False

    t = _strip_noise(t)
    if not t:
        return [original], False
    if _looks_irregular(t):
        return [t], False

    t = _strip_trailing_word_after_slash(t)
    t = _normalize_dan_separator(t)

    if not _has_real_code(original):
        if re.search(r"\d+\s*-\s*\d+", t) and "+" in t:
            return [t], False
        multi_range_result = None
    else:
        multi_range_result = _expand_multi_range(t)
        if multi_range_result is not None:
            if len(multi_range_result) > MAX_BREAKDOWN_CODES:
                return [t], False
            return multi_range_result, all(_looks_like_characteristic_code(c) for c in multi_range_result)

    if "+" in t and not re.search(r"P\d{1,2}\s*$", t):
        plus_parts = [p.strip() for p in t.split("+") if p.strip()]
        if len(plus_parts) >= 2 and all(re.search(r"\d", p) for p in plus_parts):
            if all(len(re.sub(r"[^A-Za-z]", "", p)) >= 2 for p in plus_parts):
                if len(plus_parts) <= MAX_BREAKDOWN_CODES:
                    seen, deduped = set(), []
                    for p in plus_parts:
                        if p not in seen:
                            seen.add(p)
                            deduped.append(p)
                    return deduped, all(_looks_like_characteristic_code(p) for p in deduped)

    numeric_suffix_result = _reuse_numeric_suffix(t)
    if numeric_suffix_result is not None:
        if len(numeric_suffix_result) > MAX_BREAKDOWN_CODES:
            return [t], False
        return numeric_suffix_result, all(_looks_like_characteristic_code(c) for c in numeric_suffix_result)

    prefix_result = _reuse_prefix_breakdown(t)
    if prefix_result is not None:
        if len(prefix_result) > MAX_BREAKDOWN_CODES:
            return [t], False
        return prefix_result, all(_looks_like_characteristic_code(c) for c in prefix_result)

    t_final = re.sub(r"[.,]\s*$", "", t).strip()
    if not t_final:
        return [original], False
    return [t_final], _looks_like_characteristic_code(t_final)

# ==============================================================================
# 2c. KOLOM CERTIFICATE
# ==============================================================================
CERTIFICATE_SOURCE_COLUMN = "FAC_POLICY_NO"
_CERT_SD_NORMALIZE_RE = re.compile(r"\bs\s*/\s*d\b", flags=re.IGNORECASE)

_CERT_TAIL_RANGE_RE = re.compile(r"^(?P<prefix>.+?)[\s\-]+(?P<start>0\d{0,5})\s*(?:-|S\s*/\s*D|SD)\s*(?P<end>0?\d{1,5})\s*$", flags=re.IGNORECASE)
_CERT_TAIL_SINGLE_RE = re.compile(r"^(?P<prefix>.+?)-\s*(?P<num>0\d{5})\s*$", flags=re.IGNORECASE)

def _cert_pad6(num_str: str) -> str:
    digits = re.sub(r"\D", "", num_str)
    if not digits:
        return num_str
    return digits.zfill(6)

def _format_cert_range(start_raw: str, end_raw: str):
    if len(start_raw) > 6 or len(end_raw) > 6:
        return None
    start_num, end_num = int(start_raw), int(end_raw)
    if end_num < start_num:
        return None
    start_pad, end_pad = _cert_pad6(start_raw), _cert_pad6(end_raw)
    member_count = end_num - start_num + 1
    if member_count > 3:
        return f"{start_pad} SD {end_pad}"
    codes = [_cert_pad6(str(n)) for n in range(start_num, end_num + 1)]
    return ", ".join(codes)

def extract_certificate_and_remainder(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "", value
    text = str(value).strip()
    if not text or text.lower() in ("nan", "none", "-"):
        return "", value

    normalized = _CERT_SD_NORMALIZE_RE.sub("SD", text)

    range_match = _CERT_TAIL_RANGE_RE.match(normalized)
    if range_match:
        cert = _format_cert_range(range_match.group("start"), range_match.group("end"))
        if cert is not None:
            return cert, range_match.group("prefix").strip()

    single_match = _CERT_TAIL_SINGLE_RE.match(normalized)
    if single_match:
        digits = single_match.group("num")
        if len(digits) <= 6:
            return _cert_pad6(digits), single_match.group("prefix").strip()

    return "", value

def extract_polis_certificate_pairs(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    text = str(value).strip()
    if not text or text.lower() in ("nan", "none", "-"):
        return None

    if "/" not in text:
        return None

    segments = [s.strip() for s in text.split("/") if s.strip()]
    if len(segments) < 2:
        return None

    any_cert_matched = False
    pairs = []
    for seg in segments:
        cert_value, pol_remainder = extract_certificate_and_remainder(seg)
        if cert_value:
            any_cert_matched = True
            pairs.append(([str(pol_remainder).strip().upper()], cert_value))
        else:
            pol_cln = clean_split_code(seg)
            pairs.append((pol_cln, ""))

    if not any_cert_matched:
        return None
    return pairs

# ==============================================================================
# 3. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    try:
        return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row, engine="calamine")
    except Exception:
        return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw: pd.DataFrame, keyword: str, cedant_column: str) -> pd.DataFrame:
    if cedant_column not in df_raw.columns:
        raise ValueError(f"CEDANT_COLUMN='{cedant_column}' tidak ditemukan. Tersedia: {list(df_raw.columns)}")
    mask = df_raw[cedant_column].astype(str).str.contains(keyword, case=False, na=False)
    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant (kolom '{cedant_column}') mengandung '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered

def _resolve_broker_column(df: pd.DataFrame) -> str | None:
    for candidate in BROKER_NAME_COLUMN_CANDIDATES:
        if candidate in df.columns:
            return candidate
    return None

# ==============================================================================
# 4. BUILD OUTPUT
# ==============================================================================
def _breakdown_all_rows(df: pd.DataFrame, col_insured: str, col_polis: str, col_slip: str):
    ins_arr = df[col_insured].to_numpy()
    pol_arr = df[col_polis].to_numpy()
    slp_arr = df[col_slip].to_numpy()

    insured_cln_all, polis_cln_all, slip_cln_all, certificate_cln_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val in zip(ins_arr, pol_arr, slp_arr):
        ins_cln = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []

        pairs = extract_polis_certificate_pairs(pol_val) if pd.notna(pol_val) else None

        if pairs is not None:
            pol_cln = []
            cert_cln = []
            for seg_pol_list, seg_cert in pairs:
                for code in seg_pol_list:
                    pol_cln.append(code)
                    cert_cln.append(seg_cert)
        else:
            if pd.notna(pol_val):
                cert_value, pol_remainder = extract_certificate_and_remainder(str(pol_val))
            else:
                cert_value, pol_remainder = "", pol_val
            pol_cln = clean_split_code(str(pol_remainder)) if pd.notna(pol_remainder) else []
            cert_cln = [cert_value] + [""] * (len(pol_cln) - 1) if pol_cln else []

        slp_cln = clean_split_code(str(slp_val)) if pd.notna(slp_val) else []

        insured_cln_all.append(ins_cln)
        polis_cln_all.append(pol_cln)
        slip_cln_all.append(slp_cln)
        certificate_cln_all.append(cert_cln)
        
        max_ins = max(max_ins, len(ins_cln))
        max_pol = max(max_pol, len(pol_cln))
        max_slp = max(max_slp, len(slp_cln))
        
    return insured_cln_all, polis_cln_all, slip_cln_all, certificate_cln_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_insured, col_polis, col_slip = INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN
    for col, label in ((col_insured, "INSURED_COLUMN"), (col_polis, "POLIS_COLUMN"), (col_slip, "SLIP_COLUMN")):
        if col not in df_filtered.columns:
            raise ValueError(f"Kolom {label}='{col}' tidak ditemukan. Tersedia: {list(df_filtered.columns)}")

    broker_column = _resolve_broker_column(df_filtered)

    insured_cln_all, polis_cln_all, slip_cln_all, certificate_cln_all, max_ins, max_pol, max_slp = _breakdown_all_rows(df_filtered, col_insured, col_polis, col_slip)
    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)

    if broker_column is not None and CEDANT_COLUMN in df_filtered.columns:
        is_direct = df_filtered[broker_column].astype(str).str.strip().str.upper() == DIRECT_MARKER
        bp_series = df_filtered[CEDANT_COLUMN].where(is_direct, df_filtered[broker_column])
        business_partner_values = bp_series.astype(str).str.replace(r"^(\s*PT)\.\s*", r"\1 ", regex=True, flags=re.IGNORECASE).values
    else:
        business_partner_values = None

    output_cols_data = {}
    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == broker_column and business_partner_values is not None:
            output_cols_data["BUSINESS_PARTNER"] = business_partner_values

        if col == col_polis:
            for i in range(1, max_pol + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in polis_cln_all]
                output_cols_data[f"CERTIFICATE_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in certificate_cln_all]
        if col == col_insured:
            for i in range(1, max_ins + 1): output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in insured_cln_all]
        if col == col_slip:
            for i in range(1, max_slp + 1): output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slip_cln_all]

    df_master = pd.DataFrame(output_cols_data)
    drop_cols = [c for c in df_master.columns if ("_CLN_" in c or c.startswith("CERTIFICATE_")) and df_master[c].astype(str).str.strip().eq("").all()]
    return df_master.drop(columns=drop_cols)

# ==============================================================================
# 5. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook = load_workbook(output_path)
    worksheet = workbook.active
    text_col_indices = [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]

    for col_idx in text_col_indices:
        col_name = df.columns[col_idx - 1]
        col_values = df[col_name].values
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            src_val = col_values[row - 2]
            if src_val is None or (isinstance(src_val, float) and pd.isna(src_val)):
                cell.value = None
            else:
                cell.value = str(src_val)
    workbook.save(output_path)

# ==============================================================================
# 6. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70 + "\n CLEANSING DATA 1 - FACULTATIVE | CEDANT: IBS\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(_RAW_DIR) if f.endswith(".xlsx")]
        if candidates:
            input_path = os.path.join(_RAW_DIR, candidates[0])
            print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")
        else:
            raise FileNotFoundError(f"File Excel sumber tidak ditemukan di: {_RAW_DIR}")

    os.makedirs(_CLEAN_DIR, exist_ok=True)

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    print(f"[INFO] Total baris source: {len(df_raw)}")

    df_ibs = filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)
    if df_ibs.empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter cedant. Proses dihentikan.")
        return

    df_hasil = build_output(df_ibs)
    save_with_text_format(df_hasil, OUTPUT_FILE)

    print("=" * 70 + f"\n[SUCCESS] Selesai. Total baris output: {len(df_hasil)}\n[SUCCESS] File hasil: '{OUTPUT_FILE}'\n" + "=" * 70)

if __name__ == "__main__":
    main()