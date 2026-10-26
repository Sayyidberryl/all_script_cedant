"""Audit Insured TF-IDF mandiri untuk IBS.

Tidak mengubah data input atau keputusan matching. Jalankan dari folder project:
    python TFIDF-IBS.py

Dependensi: pandas, openpyxl, numpy, scikit-learn.

Sumber referensi nama Insured:
  - OSBAL  : data_cleaned/osbal_clean_ibs.xlsx          (kolom INSURED_CLN_1..5, FAC_INSURED)
  - FACUL  : data_cleaned/facul_clean_ibs.xlsx           (kolom INSURED_CLN_1..5, FAC_INSURED)
  - RI SLIP: data_raw/5. Database Rekap RI Slip.xlsx     (kolom INSURED, header baris 2)
  - BORDERO: data_raw/September 2024 - Agustus 2026 ...  (kolom INSURED, delimiter /)

Baris yang diaudit (dari matching_ibs_bordero.xlsx):
  - SKENARIO mengandung: Insured, Unmatching
  - FLAG_PROD = Unmatching
"""
from pathlib import Path
from collections import defaultdict
import re
import unicodedata

import numpy as np
import pandas as pd
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.neighbors import NearestNeighbors


# =============================================================================
# KONFIGURASI PATH
# =============================================================================

ROOT     = Path(__file__).resolve().parent
CLEAN    = ROOT / 'data_cleaned'
RAW      = ROOT / 'data_raw'

# File input utama (output matching IBS)
_MATCH_CANDIDATES = [
    ROOT / 'matching_ibs_bordero.xlsx',
    ROOT / 'final_output_ibs_v1.xlsx',
    ROOT / 'matching_ibs_v1.xlsx',
    ROOT / 'matching_ibs.xlsx',
]
MATCH_FILE = next((f for f in _MATCH_CANDIDATES if f.exists()), _MATCH_CANDIDATES[0])

# File referensi (4 sumber)
OSBAL_FILE   = CLEAN / 'osbal_clean_ibs.xlsx'
FACUL_FILE   = CLEAN / 'facul_clean_ibs.xlsx'
RISLIP_FILE  = RAW   / '5. Database Rekap RI Slip.xlsx'
BORDERO_FILE = RAW   / 'September 2024 - Agustus 2026 (COB PAR & Earthquake).xlsx'

# File output audit
DEST = ROOT / 'tfidf_ibs_audit.xlsx'

# =============================================================================
# PARAMETER TF-IDF
# =============================================================================

TOP_N     = 10    # Kandidat per nama Insured
MIN_SCORE = 0.0   # Skor minimum (0–100); 0 = tampilkan semua

# =============================================================================
# STOPWORD — dibuang sebelum normalisasi kunci TF-IDF
# Gabungan dari COMMON (WAHANA) + _INSURED_STOPWORDS (matching_ibs_bordero.py)
# =============================================================================

COMMON = frozenset({
    'PT', 'CV', 'UD', 'PD', 'NV', 'LLC', 'LTD', 'INC', 'TBK', 'PERSERO',
    'CO', 'CORP', 'CORPORATION', 'COMPANY', 'COMPANIES',
    'GROUP', 'GROUPS',
    'INDONESIA', 'INDO',
    'INDUSTRY', 'INDUSTRIES', 'INDUSTRI',
    'INTERNATIONAL', 'INTERNASIONAL',
    'ENTERPRISE', 'ENTERPRISES',
    'TRADING', 'TRADE',
    'UTAMA', 'MANDIRI', 'MAKMUR', 'SEJAHTERA',
})


# =============================================================================
# FUNGSI UTILITAS
# =============================================================================

def txt(value) -> str:
    """Konversi nilai ke string; kembalikan '' untuk null/nan."""
    if value is None:
        return ''
    if not isinstance(value, (list, tuple)):
        try:
            if pd.isna(value):
                return ''
        except (TypeError, ValueError):
            pass
    s = str(value).strip()
    return '' if s.upper() in {'NAN', 'NONE', 'NAT', '-'} else s


def norm(value) -> str:
    """Normalisasi string jadi kunci TF-IDF: NFKC → upper → hanya alfanumerik → buang stopword."""
    v = unicodedata.normalize('NFKC', txt(value)).upper()
    v = re.sub(r'[^A-Z0-9]+', ' ', v)
    return ' '.join(token for token in v.split() if token not in COMMON)


def col_find(df: pd.DataFrame, *names) -> str | None:
    """Temukan nama kolom DataFrame yang cocok (case/symbol-insensitive)."""
    lookup = {re.sub(r'[^A-Z0-9]+', '', str(k).upper()): k for k in df.columns}
    for name in names:
        result = lookup.get(re.sub(r'[^A-Z0-9]+', '', name.upper()))
        if result is not None:
            return result
    return None


def insured_clean_columns(df: pd.DataFrame) -> list:
    """Temukan kolom INSURED_CLN_1..5 atau INSURED_CLEAN_1..5 dalam DataFrame."""
    result = []
    for column in df.columns:
        name = re.sub(r'[^A-Z0-9]+', '', str(column).upper())
        if re.fullmatch(r'(?:INSUREDCLEAN|INSUREDCLN|CLEANINSURED|INSURED)[1-5]', name):
            result.append(column)
    return sorted(result, key=lambda c: int(re.search(r'[1-5]$', str(c)).group()))


def name_columns(df: pd.DataFrame, source: str) -> list:
    """Tentukan kolom nama Insured yang dipakai per sumber."""
    if source == 'BORDERO':
        c = col_find(df, 'INSURED')
        if c is None:
            raise ValueError(f'BORDERO: kolom INSURED tidak ditemukan. Kolom ada: {list(df.columns)}')
        return [c]
    if source == 'RISLIP':
        c = col_find(df, 'INSURED', 'INSURED NAME', 'INSURED_NAME')
        if c is None:
            raise ValueError(f'RI SLIP: kolom INSURED tidak ditemukan. Kolom ada: {list(df.columns)}')
        return [c]
    # OSBAL / FACUL: utamakan kolom clean, fallback ke raw
    clean = insured_clean_columns(df)
    if clean:
        return clean
    raw = col_find(df, 'FAC_INSURED', 'INSURED_ORI', 'INSURED', 'INSURED NAME', 'INSURED_NAME')
    if raw is None:
        raise ValueError(f'{source}: kolom Insured tidak ditemukan. Kolom ada: {list(df.columns)}')
    return [raw]


def variants(row, columns: list, source: str):
    """Generator (column, raw_str, normalized_key) per varian nama Insured satu baris."""
    for column in columns:
        raw = txt(row[column])
        if not raw:
            continue
        # Bordero: nama dipisah delimiter '/'
        parts = re.split(r'\s*/\s*', raw) if source == 'BORDERO' else [raw]
        for part in parts:
            part = part.strip()
            key = norm(part)
            if key:
                yield str(column), part, key


# =============================================================================
# BACA FILE
# =============================================================================

def read_required(path: Path, label: str) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f'{label} tidak ditemukan: {path}')
    print(f'  Membaca {label}: {path.name}')
    return pd.read_excel(path, dtype=object)


def read_rislip(path: Path) -> pd.DataFrame:
    """Baca RI SLIP; coba header baris 0–3 hingga menemukan kolom INSURED."""
    for header in range(4):
        frame = pd.read_excel(path, header=header, dtype=object)
        has_insured = (
            insured_clean_columns(frame) or
            col_find(frame, 'FAC_INSURED', 'INSURED_ORI', 'INSURED', 'INSURED NAME', 'INSURED_NAME') is not None
        )
        if has_insured:
            if header:
                print(f'    [RI SLIP] Header ditemukan di baris {header + 1}')
            return frame
    raise ValueError(f'RI SLIP: kolom INSURED tidak ditemukan di baris header 1–4: {path}')


# =============================================================================
# BANGUN REFERENSI
# =============================================================================

def references(source_df: pd.DataFrame, source: str) -> list:
    """
    Kembalikan list dict referensi unik per (source, normalized_key):
      { source, name, key, codes, currencies, count, columns }
    """
    namecols  = name_columns(source_df, source)
    fc_col    = col_find(source_df, 'CCOS_REF_CODE', 'FAC_CODE', 'FAC CODE', 'FAC. CODE', 'FACODE')
    curr_col  = col_find(source_df, 'CCOS_CURR', 'FAC_CURRENCY', 'CURR', 'CURR ORI', 'CURRENCY')

    records: dict = {}
    for _, row in source_df.iterrows():
        for column, raw, key in variants(row, namecols, source):
            rec = records.setdefault(key, {
                'source': source, 'name': raw, 'key': key,
                'codes': set(), 'currencies': set(),
                'count': 0, 'columns': set(),
            })
            rec['count'] += 1
            rec['columns'].add(column)
            if fc_col and txt(row[fc_col]):
                rec['codes'].add(txt(row[fc_col]))
            if curr_col and txt(row[curr_col]):
                rec['currencies'].add(txt(row[curr_col]).upper())
    return list(records.values())


# =============================================================================
# FILTER BARIS AUDIT DARI OUTPUT MATCHING IBS
# =============================================================================

def filter_audit_rows(result: pd.DataFrame) -> tuple:
    """
    Kembalikan (selected_df, skenario_col, flag_col).

    Baris yang masuk audit:
      - SKENARIO mengandung kata 'insured' atau 'unmatching'
      - ATAU FLAG_PROD = 'Unmatching'
    """
    skenario_col = col_find(result, 'SKENARIO')
    flag_col     = col_find(result, 'FLAG_PROD')

    if skenario_col is None and flag_col is None:
        raise ValueError('Output matching tidak punya kolom SKENARIO maupun FLAG_PROD')

    ske   = result[skenario_col].fillna('').astype(str) if skenario_col else pd.Series('', index=result.index)
    flags = result[flag_col].fillna('').astype(str)     if flag_col     else pd.Series('', index=result.index)

    mask = (
        ske.str.contains('insured', case=False, regex=False) |
        ske.str.contains('unmatching|unmatched', case=False, regex=True) |
        flags.str.contains('unmatching|unmatched', case=False, regex=True)
    )
    return result.loc[mask].copy(), skenario_col, flag_col


# =============================================================================
# RISK FLAGS
# =============================================================================

def compute_risk_flags(qt: set, rt: set, ref_codes: set) -> str:
    risks = []
    if (qt ^ rt) & {'SYARIAH', 'SHARIA'}:
        risks.append('SYARIAH_DIFFERENCE')
    if (qt ^ rt) & {'GROUP', 'GRUP'}:
        risks.append('GROUP_DIFFERENCE')
    if (qt ^ rt) & {'PALM', 'SAWIT', 'KEBUN', 'PERKEBUNAN'}:
        risks.append('BUSINESS_WORD_DIFFERENCE')
    if len(qt & rt - {'INDONESIA', 'GROUP', 'GRUP'}) <= 1:
        risks.append('WEAK_SHARED_TOKENS')
    if len(ref_codes) > 1:
        risks.append('REFERENCE_MULTI_FAC_CODE')
    return ' | '.join(risks)


# =============================================================================
# MAIN
# =============================================================================

def main():
    print('=' * 60)
    print('TFIDF-IBS Audit Tool')
    print('=' * 60)

    # ── 1. Baca output matching IBS ──────────────────────────────
    print('\n[1] Membaca output matching IBS ...')
    result = read_required(MATCH_FILE, 'Matching IBS')
    selected, skenario_col, flag_col = filter_audit_rows(result)
    print(f'    Total baris matching : {len(result):,}')
    print(f'    Baris terpilih audit : {len(selected):,}')

    # ── 2. Baca 4 sumber referensi ───────────────────────────────
    print('\n[2] Membaca sumber referensi ...')
    sources_data = [
        ('OSBAL',   read_required(OSBAL_FILE,   'OSBAL')),
        ('FACUL',   read_required(FACUL_FILE,   'FACUL')),
        ('RISLIP',  read_rislip(RISLIP_FILE)),
        ('BORDERO', read_required(BORDERO_FILE, 'BORDERO')),
    ]

    refs  = []
    stats = []
    for source, frame in sources_data:
        group = references(frame, source)
        stats.append({'METRIK': f'Referensi unik {source}', 'NILAI': len(group)})
        refs.extend(group)
        print(f'    {source:<8}: {len(group):,} nama unik')

    if not refs:
        raise ValueError('Keempat sumber tidak memiliki nama Insured yang valid.')

    print(f'    TOTAL   : {len(refs):,} referensi gabungan')

    # ── 3. Fit TF-IDF ────────────────────────────────────────────
    print('\n[3] Fit TF-IDF vectorizer ...')
    vectorizer = TfidfVectorizer(
        analyzer='char_wb',
        ngram_range=(2, 5),
        min_df=1,
        sublinear_tf=True,
    )
    matrix = vectorizer.fit_transform([r['key'] for r in refs])

    neighbors = NearestNeighbors(
        metric='cosine',
        algorithm='brute',
        n_neighbors=min(TOP_N, len(refs)),
    )
    neighbors.fit(matrix)
    print(f'    Matrix shape: {matrix.shape}')

    # ── 4. Bangun query dari baris audit ─────────────────────────
    print('\n[4] Menyusun query Insured dari baris audit ...')

    # Kolom Insured di output matching (INSURED_1, INSURED_2, atau INSURED_ORI)
    qcols = [c for c in ['INSURED_1', 'INSURED_2'] if c in result.columns]
    if not qcols:
        qcols = [c for c in ['INSURED_ORI', 'INSURED'] if c in result.columns]

    ori_col      = col_find(result, 'INSURED_ORI', 'INSURED')
    fac_col      = col_find(result, 'CCOS_REF_CODE', 'FAC_CODE', 'FAC CODE')
    curr_col     = col_find(result, 'CURR ORI', 'CURR')
    receipt_col  = col_find(result, 'RECEIPT NO', 'RECEIPT_NO')

    queries: dict = {}
    details: list = []

    for idx, row in selected.iterrows():
        row_variants = list(variants(row, qcols, 'SUSPEND'))
        # Fallback ke kolom ori jika kolom clean kosong
        if not row_variants and ori_col:
            row_variants = list(variants(row, [ori_col], 'SUSPEND'))

        for column, raw, key in row_variants:
            q = queries.setdefault(key, {
                'name': raw, 'key': key,
                'receipts': set(), 'scenarios': set(),
                'flags': set(), 'columns': set(),
            })
            q['receipts'].add(idx)
            q['columns'].add(column)
            q['scenarios'].add(txt(row[skenario_col]) if skenario_col else '')
            q['flags'].add(txt(row[flag_col])         if flag_col     else '')

            details.append({
                'EXCEL_ROW':     idx + 2,
                'INSURED_QUERY': raw,
                'INSURED_KOLOM': column,
                'INSURED_KEY':   key,
                'INSURED_ORI':   txt(row[ori_col])     if ori_col     else '',
                'SKENARIO':      txt(row[skenario_col]) if skenario_col else '',
                'FLAG_PROD':     txt(row[flag_col])     if flag_col    else '',
                'FAC_CODE_HASIL':txt(row[fac_col])      if fac_col     else '',
                'CURR_ORI':      txt(row[curr_col])     if curr_col    else '',
                'RECEIPT_NO':    txt(row[receipt_col])  if receipt_col else '',
            })

    print(f'    Nama Insured unik untuk diaudit: {len(queries):,}')

    # ── 5. Query TF-IDF → kandidat ───────────────────────────────
    print('\n[5] Menjalankan TF-IDF nearest neighbors ...')
    candidates: list = []

    if queries:
        all_queries = list(queries.values())
        q_matrix    = vectorizer.transform([q['key'] for q in all_queries])
        distances, indices = neighbors.kneighbors(q_matrix)

        for q, ds, ids in zip(all_queries, distances, indices):
            qt = set(q['key'].split())
            for rank, (distance, ri) in enumerate(zip(ds, ids), start=1):
                ref   = refs[int(ri)]
                score = round(max(0.0, 1.0 - float(distance)) * 100, 2)
                if score < MIN_SCORE:
                    continue

                rt     = set(ref['key'].split())
                common = qt & rt

                candidates.append({
                    'INSURED_SUSPEND':       q['name'],
                    'NORMALIZED_SUSPEND':    q['key'],
                    'KOLOM_QUERY':           ', '.join(sorted(q['columns'])),
                    'JUMLAH_RECEIPT':        len(q['receipts']),
                    'SKENARIO_ASAL':         ' | '.join(sorted(x for x in q['scenarios'] if x)),
                    'FLAG_ASAL':             ' | '.join(sorted(x for x in q['flags']     if x)),
                    'RANK':                  rank,
                    'TFIDF_SCORE':           score,
                    'KANDIDAT_INSURED':      ref['name'],
                    'KANDIDAT_NORMALIZED':   ref['key'],
                    'SUMBER':                ref['source'],
                    'SHARED_TOKENS':         ', '.join(sorted(common)),
                    'RISK_FLAGS':            compute_risk_flags(qt, rt, ref['codes']),
                    'FAC_CODE_REFERENSI':    ', '.join(sorted(ref['codes'])),
                    'JUMLAH_FAC_CODE_REF':   len(ref['codes']),
                    'CURRENCY_REFERENSI':    ', '.join(sorted(ref['currencies'])),
                    'JUMLAH_REFERENSI':      ref['count'],
                    'KOLOM_REFERENSI':       ', '.join(sorted(ref['columns'])),
                    # Kolom untuk diisi reviewer
                    'STATUS_VALIDASI':       'PENDING',
                    'CATATAN_REVIEW':        '',
                    'KEPUTUSAN_FAC_CODE':    '',
                })

    print(f'    Total kandidat dihasilkan: {len(candidates):,}')

    # ── 6. Susun summary ─────────────────────────────────────────
    # Breakdown skenario baris audit
    ske_breakdown = []
    if skenario_col:
        vc = selected[skenario_col].fillna('(kosong)').value_counts()
        for k, v in vc.items():
            ske_breakdown.append({'METRIK': f'  Skenario: {k}', 'NILAI': v})

    summary = [
        {'METRIK': 'Perusahaan',                         'NILAI': 'IBS'},
        {'METRIK': 'File matching input',                 'NILAI': MATCH_FILE.name},
        {'METRIK': 'Total baris matching',                'NILAI': len(result)},
        {'METRIK': 'Baris terpilih untuk audit',          'NILAI': len(selected)},
        *ske_breakdown,
        {'METRIK': 'Nama Insured unik diaudit',           'NILAI': len(queries)},
        *stats,
        {'METRIK': 'Total referensi gabungan (unik key)', 'NILAI': len(refs)},
        {'METRIK': 'Top N kandidat per nama Insured',     'NILAI': TOP_N},
        {'METRIK': 'Jumlah kandidat dihasilkan',          'NILAI': len(candidates)},
    ]

    guide = [
        {'PANDUAN': 'Sumber kandidat: OSBAL, FACUL, RISLIP, BORDERO.'},
        {'PANDUAN': 'OSBAL & FACUL: menggunakan INSURED_CLN_1..5 (clean); fallback FAC_INSURED.'},
        {'PANDUAN': 'RI SLIP: kolom INSURED (header baris 2 file asli).'},
        {'PANDUAN': 'BORDERO: kolom INSURED, nama dipecah berdasarkan delimiter /.'},
        {'PANDUAN': 'Baris diaudit: SKENARIO mengandung Insured/Unmatching, atau FLAG_PROD=Unmatching.'},
        {'PANDUAN': 'TFIDF_SCORE: 0–100, semakin tinggi semakin mirip (karakter n-gram 2–5).'},
        {'PANDUAN': 'RISK_FLAGS: WEAK_SHARED_TOKENS = hati-hati, mungkin false positive.'},
        {'PANDUAN': 'RISK_FLAGS: REFERENCE_MULTI_FAC_CODE = 1 nama referensi terhubung >1 FAC CODE.'},
        {'PANDUAN': 'TFIDF_SCORE hanya ranking kemiripan; TIDAK boleh otomatis mengubah FAC_CODE.'},
        {'PANDUAN': 'Isi kolom STATUS_VALIDASI (VALID/INVALID/UNCERTAIN) dan KEPUTUSAN_FAC_CODE.'},
        {'PANDUAN': 'Tool ini hanya membaca file; tidak mengubah matching maupun output asli.'},
    ]

    # ── 7. Tulis ke Excel ────────────────────────────────────────
    print(f'\n[6] Menulis output ke {DEST.name} ...')
    from openpyxl.styles import Font, PatternFill

    with pd.ExcelWriter(DEST, engine='openpyxl') as writer:
        sheets = [
            ('AUDIT_SUMMARY',    summary),
            ('TFIDF_CANDIDATES', candidates),
            ('RECEIPT_DETAILS',  details),
            ('CARA_PAKAI',       guide),
        ]
        for sheet_name, rows in sheets:
            frame = pd.DataFrame(rows)
            if frame.empty:
                frame = pd.DataFrame({'INFO': ['Tidak ada data yang memenuhi filter']})
            frame.to_excel(writer, sheet_name=sheet_name, index=False)

            ws = writer.sheets[sheet_name]
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions

            # Format header: putih bold di atas biru tua IBS
            for cell in ws[1]:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='1A3A5C')

            # Auto-width kolom (max 55 karakter)
            for column in ws.columns:
                max_len = max(
                    len(str(cell.value or ''))
                    for cell in list(column)[:200]
                )
                ws.column_dimensions[column[0].column_letter].width = min(55, max(14, max_len + 2))

    print(f'\nSelesai! [OK]')
    print(f'  Output : {DEST}')
    print(f'  Diaudit: {len(selected)} baris | Query: {len(queries)} nama | Kandidat: {len(candidates)}')
    print('=' * 60)


if __name__ == '__main__':
    main()
