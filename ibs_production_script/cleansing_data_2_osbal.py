"""
================================================================================
 SCRIPT CLEANSING DATA 2 - OSBAL (CEDANT : IBS )
================================================================================
"""

import sys
sys.modules['numexpr'] = None
sys.modules['bottleneck'] = None

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
_BASE_DIR   = os.path.dirname(os.path.abspath(__file__))  # folder ibs_production_script/
_RAW_DIR    = os.path.join(_BASE_DIR, "data_raw")
_CLEAN_DIR  = os.path.join(_BASE_DIR, "data_cleaned")

INPUT_FILE  = os.path.join(_RAW_DIR,   "2b. transaksi Osbal 01.01.23 - 17.08.26.xlsx")
SHEET_NAME  = "Query result"
HEADER_ROW  = 0
OUTPUT_FILE = os.path.join(_CLEAN_DIR, "ibs_output_osbal_V2.xlsx")
CEDANT_FILTER = "IBS"

CEDANT_COLUMN = "CCOS_COMP_NAME"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"

CLSDT_POLIS_COLUMN = "CLSDT_POLICY_NO"
CLSDT_SLIP_COLUMN = "CLSDT_SLIP_NO"
CLSDT_SERTF_COLUMN = "CLSDT_SERTF_NO"
CERTIFICATE_OUTPUT_COL = "CERTIFICATE"

MAX_BREAKDOWN_CODES = 5
TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CLSDT", "CERTIFICATE",
)

CERT_DIGIT_LEN = 6
CERT_RANGE_JOIN_THRESHOLD = 3

# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
_BYPASS_ALL_RE = re.compile(r"\bPENYELESAIAN\s+SUSPENSE\b|\bPENYELESAIAN\s+HUTANG\s+PIUTANG\b|\bPENYELESAIAN\s+UTANG\s+PIUTANG\b", flags=re.IGNORECASE)
_TITLE_RE = re.compile(r"\b(?:S\.?H\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b", flags=re.IGNORECASE)
_LEGAL_ENTITY_RE = re.compile(r"\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO|Perseroan\s+Terbatas)\b", flags=re.IGNORECASE)
_LEGAL_ENTITY_WITH_LEADING_PUNCT_RE = re.compile(r"[,.]?\s*\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO|Perseroan\s+Terbatas)\b", flags=re.IGNORECASE)

_ENTITY_SPLIT_RE = re.compile(r"\bQQ\b|\band\s*/\s*or\b|\bN\b", flags=re.IGNORECASE)
_ENTITY_SLASH_SPLIT_RE = re.compile(r"\s*/\s*")
_INSURED_DESCRIPTIVE_RE = re.compile(r"\bVARIOUS\s+INSUREDS?\b|\bTBA\b|\bVAR(?:IOUS)?\b", flags=re.IGNORECASE)
_INSURED_DASH_SPLIT_EXCEPTIONS = ("EBDI VAUNDRI - OSCAR OMEGA",)

def _strip_legal_entity_with_punct(text: str) -> str:
    t = _LEGAL_ENTITY_WITH_LEADING_PUNCT_RE.sub("", text)
    t = re.sub(r",\s*(?=[(\s]|$)", " ", t)
    return t

def _final_polish(text: str) -> str:
    t = re.sub(r"\s+", " ", re.sub(r"[().;:\"']", " ", text)).strip()
    t = re.sub(r"^[\s,.\-–]+|[\s,.\-–]+$", "", t)
    dangling_pattern = r"^\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*|\s*\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$"
    previous = None
    while previous != t:
        previous, t = t, re.sub(dangling_pattern, "", t, flags=re.IGNORECASE).strip()
    return t.upper()

def _split_insured_entities(text: str):
    stripped = text.strip()
    for exception in _INSURED_DASH_SPLIT_EXCEPTIONS:
        if stripped.upper() == exception.upper():
            return [p.strip() for p in re.split(r"\s*-\s*", stripped) if p.strip()]

    t = _ENTITY_SPLIT_RE.sub("/", text)
    parts = _ENTITY_SLASH_SPLIT_RE.split(t)
    return [p.strip() for p in parts if p.strip()]

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []
    original = text.strip()

    if _BYPASS_ALL_RE.search(original) or _INSURED_DESCRIPTIVE_RE.fullmatch(original.strip()):
        return [original]

    raw_parts = _split_insured_entities(original)
    if not raw_parts:
        return [original]

    if len(raw_parts) > MAX_BREAKDOWN_CODES:
        return [original]

    results = []
    for part in raw_parts:
        cleaned = _final_polish(_strip_titles(_strip_legal_entity_with_punct(part)))
        if not cleaned or len(cleaned) < 2 or not re.search(r"[A-Z]", cleaned):
            continue
        if cleaned not in results:
            results.append(cleaned)

    if not results:
        return [original]

    return [original] if len(results) > MAX_BREAKDOWN_CODES else results

def _strip_titles(text: str) -> str:
    return _TITLE_RE.sub(" ", text)

# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO
# ==============================================================================
_COMMA_ONLY_RE = re.compile(r",")
_REF_SUFFIX_RE = re.compile(r"^P\d*$|^END\s*\d*$", flags=re.IGNORECASE)
_TBA_ALONE_RE = re.compile(r"^\s*TBA\s*\.?\s*/?\s*$", flags=re.IGNORECASE)
_VAR_ALONE_RE = re.compile(r"^\s*VAR(?:IOUS)?\s*$", flags=re.IGNORECASE)
_P_ALONE_RE = re.compile(r"^\s*P\s*$", flags=re.IGNORECASE)

_BYPASS_KEEP_ASIS_RE = re.compile(r"\bPENYELESAIAN\s+SUSPENSE\b|\bPENYELESAIAN\s+HUTANG\s+PIUTANG\b|\bPENYELESAIAN\s+UTANG\s+PIUTANG\b", flags=re.IGNORECASE)
_POLICY_SLIP_SPLIT_RE = re.compile(r"\s*\+\s*")
_SD_RANGE_RE = re.compile(r"^(.*?)(\d+)\s+S\s*/\s*D\s+(\d+)\s*$", flags=re.IGNORECASE)
_DASH_RANGE_RE = re.compile(r"^(.*?)(\d+)\s*-\s*(\d+)\s*$")

_NARRATIVE_PREFIX_RE = re.compile(r"^\s*IBS\s*RE\s*NO\.?\s*", flags=re.IGNORECASE)
_LEADING_NO_PREFIX_RE = re.compile(r"^\s*No\s+(?=[A-Za-z0-9])", flags=re.IGNORECASE)
_PAREN_CONTENT_RE = re.compile(r"\s*\([^)]*\)\s*")
_MONTH_WORDS_ID = ("JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI", "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER")
_SLASH_MONTH_YEAR_RE = re.compile(r"\s*/\s*(?:" + "|".join(_MONTH_WORDS_ID) + r")\s+\d{4}\s*", flags=re.IGNORECASE)
_SLASH_END_SUFFIX_RE = re.compile(r"\s*/\s*END\s*$", flags=re.IGNORECASE)
_TRAILING_CURRENCY_RE = re.compile(r"(?<=\d)\s+(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\s*$", flags=re.IGNORECASE)
_PURE_NARRATIVE_RE = re.compile(r"^\s*(?:" + "|".join(_MONTH_WORDS_ID) + r")\b.*$|^\s*BORDERO\b.*$", flags=re.IGNORECASE)

def _is_complex_month_year_pattern(text: str) -> bool:
    m = _SLASH_MONTH_YEAR_RE.search(text)
    if not m:
        return False
    before = text[: m.start()]
    if re.search(r"\bIBS\s*RE\b", before, flags=re.IGNORECASE):
        return False
    segments = before.split("/")
    return len(segments) >= 4

def _strip_polis_slip_noise(text: str) -> str:
    t = _NARRATIVE_PREFIX_RE.sub("", text)
    t = _LEADING_NO_PREFIX_RE.sub("", t)
    t = _PAREN_CONTENT_RE.sub(" ", t)
    t = _SLASH_MONTH_YEAR_RE.sub("", t)
    t = _SLASH_END_SUFFIX_RE.sub("", t)
    t = _TRAILING_CURRENCY_RE.sub("", t)
    t = re.sub(r"\s{2,}", " ", t).strip()
    t = re.sub(r"\s*\+\s*$", "", t).strip()
    return t

def _is_kept_as_is(value: str) -> bool:
    v = value.strip()
    if _TBA_ALONE_RE.match(v) or _VAR_ALONE_RE.match(v) or _P_ALONE_RE.match(v):
        return True
    if _BYPASS_KEEP_ASIS_RE.search(v) or _PURE_NARRATIVE_RE.match(v):
        return True
    return False

def _clean_single_code(value: str) -> str:
    s = value.strip()
    s = re.sub(r",\s*P\d*\s*$", "", s, flags=re.IGNORECASE)
    s = _dot_to_space_if_applicable(s)
    s = _COMMA_ONLY_RE.sub("", s)
    s = re.sub(r"\s{2,}", " ", s).strip()
    return s

_ALL_NUMERIC_DOT_SEGMENTS_RE = re.compile(r"^OC(?:\.\d+)+(?:-\d+)?$", flags=re.IGNORECASE)

def _dot_to_space_if_applicable(text: str) -> str:
    if "." not in text:
        return text
    tokens = text.split(" ")
    result_tokens = []
    for tok in tokens:
        if "." in tok and _ALL_NUMERIC_DOT_SEGMENTS_RE.match(tok):
            result_tokens.append(tok)
        else:
            result_tokens.append(tok.replace(".", " "))
    return " ".join(result_tokens)

def _strip_trailing_ref_codes(parts: list) -> list:
    return [p for p in parts if not _REF_SUFFIX_RE.match(p.strip())]

_LONG_GAP_DUPLICATE_RE = re.compile(r"^(.+?)\s{4,}(?:No\s+)?(.+)$", flags=re.IGNORECASE)

def _try_resolve_long_gap_duplicate(text: str):
    match = _LONG_GAP_DUPLICATE_RE.match(text)
    if not match:
        return None
    left, right = match.group(1).strip(), match.group(2).strip()

    if _TBA_ALONE_RE.match(right):
        return left if left else None

    left_clean = _clean_single_code(left)
    right_clean = _clean_single_code(right)
    if left_clean == right_clean and left_clean:
        return left_clean
    return None

REBUILD_DIGIT_WIDTH = 4

def _rebuild_number_with_base(base: str, tail: str) -> str:
    base_digit_match = re.search(r"\d+\s*$", base)
    if not base_digit_match:
        return base + tail
    base_digits = base_digit_match.group(0).strip()
    base_prefix = base[: base_digit_match.start()]
    width = max(REBUILD_DIGIT_WIDTH, len(tail))
    if len(base_digits) < width:
        width = len(base_digits)
    if len(tail) >= width:
        replaced_tail = tail
        kept_prefix_digits = base_digits[: len(base_digits) - len(tail)] if len(tail) < len(base_digits) else ""
    else:
        replaced_tail = base_digits[-width:][: width - len(tail)] + tail
        kept_prefix_digits = base_digits[: len(base_digits) - width]
    return base_prefix + kept_prefix_digits + replaced_tail

def _dedupe_fuzzy_parts(parts: list) -> list:
    cleaned = [_clean_single_code(p) for p in parts]

    if len(set(cleaned)) == 1:
        return [parts[0]]

    def trailing_digits(s: str, n: int = 4):
        tail = s[-n:] if len(s) >= n else None
        return tail if tail and tail.isdigit() else None

    def leading_letters(s: str):
        m = re.match(r"^([^\d]*)", s)
        return m.group(1) if m else ""

    dropped = set()
    for i in range(len(cleaned)):
        if i in dropped:
            continue
        for j in range(len(cleaned)):
            if j == i or j in dropped:
                continue
            a, b = cleaned[i], cleaned[j]
            if a == b:
                dropped.add(j)
                continue
            td_a, td_b = trailing_digits(a), trailing_digits(b)
            if td_a and td_b and td_a == td_b and leading_letters(a) == leading_letters(b):
                shorter_idx = i if len(a) < len(b) else j
                dropped.add(shorter_idx)
                continue
            shorter, longer = (a, b) if len(a) < len(b) else (b, a)
            if shorter and (longer == shorter or longer.endswith(" " + shorter)):
                shorter_idx = i if len(a) < len(b) else j
                dropped.add(shorter_idx)

    return [parts[i] for i in range(len(parts)) if i not in dropped]

def _try_expand_sd_range(text: str):
    match = _SD_RANGE_RE.match(text)
    if not match:
        return None
    prefix, start_s, end_s = match.groups()
    if not (start_s.isdigit() and end_s.isdigit()):
        return "__KEEP_AS_IS__"
    try:
        start, end = int(start_s), int(end_s)
    except ValueError:
        return "__KEEP_AS_IS__"
    width = len(start_s)
    if start > end or (end - start + 1) > MAX_BREAKDOWN_CODES:
        return "__KEEP_AS_IS__"
    nums = [f"{i:0{width}d}" for i in range(start, end + 1)]
    prefix_clean = re.sub(r"\s+$", "", prefix)
    return [_clean_single_code(f"{prefix_clean}{n}") for n in nums]

def _try_expand_dash_pattern(text: str):
    match = _DASH_RANGE_RE.match(text)
    if not match:
        return None
    prefix, left_num, right_num = match.groups()

    prefix_clean_check = re.sub(r"\s+$", "", prefix)
    if not re.search(r"[A-Za-z]", prefix_clean_check):
        return "__KEEP_AS_IS__"

    has_space_before_range = bool(re.search(r"\s" + re.escape(left_num) + r"\s*-\s*" + re.escape(right_num) + r"\s*$", text))

    if not has_space_before_range:
        rebuilt = f"{prefix}{left_num}{right_num}"
        return [_clean_single_code(rebuilt)]

    prefix_clean = re.sub(r"\s+$", "", prefix)
    _trailing_digit_match = re.search(r"(\d+)$", prefix_clean)
    _trailing_digit_len = len(_trailing_digit_match.group(1)) if _trailing_digit_match else 0
    needs_space = bool(prefix_clean) and _trailing_digit_len <= 4
    sep = " " if needs_space else ""
    code_left = _clean_single_code(f"{prefix_clean}{sep}{left_num}")
    code_right = _clean_single_code(f"{prefix_clean}{sep}{right_num}")
    if code_left == code_right:
        return [code_left]
    return [code_left, code_right]

_LEADING_TBA_SLASH_RE = re.compile(r"^\s*TBA\s*\.?\s*/\s*", flags=re.IGNORECASE)
_TBA_TOKEN_RE = re.compile(r"^TBA\.?$", flags=re.IGNORECASE)

def _has_tba_alongside_others(text: str) -> bool:
    if "/" not in text:
        return False
    tokens = [t.strip() for t in text.split("/") if t.strip()]
    has_tba = any(_TBA_TOKEN_RE.match(t) for t in tokens)
    has_other = any(not _TBA_TOKEN_RE.match(t) for t in tokens)
    return has_tba and has_other

def _strip_tba_tokens_keep_rest(text: str) -> str:
    tokens = [t.strip() for t in text.split("/")]
    kept = [t for t in tokens if t and not _TBA_TOKEN_RE.match(t)]
    rejoined = "/".join(kept)
    return _clean_single_code_keep_slash(rejoined) if rejoined else ""

_MIN_DIGITS_FOR_VALID_NUMBER = 4

def _process_slash_pattern(text: str):
    if "/" not in text:
        return None

    if _has_tba_alongside_others(text):
        rest = _strip_tba_tokens_keep_rest(text)
        return [rest] if rest else ["__KEEP_AS_IS__"]

    parts = [p.strip() for p in text.split("/")]
    if len(parts) < 2 or not parts[0]:
        return None

    base = parts[0]
    tail_parts = parts[1:]

    if all(re.fullmatch(r"\d+", p) for p in tail_parts):
        if not re.search(r"[A-Za-z]", base):
            return "__KEEP_AS_IS__"
        if len(parts) > MAX_BREAKDOWN_CODES:
            return "__KEEP_AS_IS__"
        codes = [_clean_single_code(base)]
        for tail in tail_parts:
            rebuilt = _rebuild_number_with_base(base, tail)
            codes.append(_clean_single_code(rebuilt))
        deduped = []
        for c in codes:
            if c not in deduped:
                deduped.append(c)
        return deduped

    digit_count_in_base = len(re.findall(r"\d", base))
    if digit_count_in_base >= _MIN_DIGITS_FOR_VALID_NUMBER:
        return [_clean_single_code(base)]

    return "__KEEP_AS_IS__"

def _clean_single_code_keep_slash(value: str) -> str:
    s = value.strip()
    s = re.sub(r",\s*P\d*\s*$", "", s, flags=re.IGNORECASE)
    segments = s.split("/")
    segments = [_strip_legal_entity_with_punct(seg).strip() for seg in segments]
    s = "/".join(seg for seg in segments if seg)
    s = _COMMA_ONLY_RE.sub("", s)
    s = re.sub(r"\s{2,}", " ", s).strip()
    return s

_END_MARKER_RE = re.compile(r"\bEND\s*\d*\b", flags=re.IGNORECASE)

def _try_resolve_numeric_base_with_end_markers(text: str):
    if not re.search(r"\bEND\s*\d*\b", text, flags=re.IGNORECASE):
        return None

    parts = [p.strip() for p in _POLICY_SLIP_SPLIT_RE.split(text) if p.strip()]
    if len(parts) < 2:
        return None

    end_parts = [p for p in parts if _END_MARKER_RE.fullmatch(p.strip())]
    non_end_parts = [p for p in parts if not _END_MARKER_RE.fullmatch(p.strip())]

    if not end_parts or not non_end_parts:
        return None

    combined_non_end = " ".join(non_end_parts)
    if re.search(r"[A-Za-z]", combined_non_end):
        return None
    if not re.search(r"\d", combined_non_end):
        return None

    base_rejoined = " + ".join(non_end_parts)
    base_rejoined = re.sub(r",\s*P\d*\s*$", "", base_rejoined, flags=re.IGNORECASE)
    return base_rejoined.strip()

_NUMERIC_MULTI_GROUP_PLUS_RE = re.compile(r"^(\d{4})\s+(\d{1,4})\s+(\d{1,4})\s*\+\s*(.+)$")

def _try_expand_numeric_multi_group(text: str):
    match = _NUMERIC_MULTI_GROUP_PLUS_RE.match(text)
    if not match:
        return None
    year, group1, group2, rest = match.groups()

    plus_segments = [s.strip() for s in rest.split("+") if s.strip()]
    all_tail_numbers = []
    for seg in plus_segments:
        seg_numbers = seg.split()
        if not all(re.fullmatch(r"\d{1,4}", n) for n in seg_numbers) or not seg_numbers:
            return None
        all_tail_numbers.extend(seg_numbers)

    if not all_tail_numbers:
        return None

    codes = [_clean_single_code(f"{year} {group1} {group2}")]
    for num in all_tail_numbers:
        codes.append(_clean_single_code(f"{year} {group1} {num}"))

    deduped = []
    for c in codes:
        if c not in deduped:
            deduped.append(c)
    if len(deduped) > MAX_BREAKDOWN_CODES:
        return None
    return deduped

def _split_and_rebuild_plus(cleaned_noise: str):
    parts = _POLICY_SLIP_SPLIT_RE.split(cleaned_noise)
    parts = [p.strip() for p in parts if p.strip()]

    if len(parts) <= 1:
        return [_clean_single_code(cleaned_noise)]

    parts = _strip_trailing_ref_codes(parts)
    if len(parts) == 0:
        return [_clean_single_code(cleaned_noise)]
    if len(parts) == 1:
        return [_clean_single_code(parts[0])]

    parts = _dedupe_fuzzy_parts(parts)
    if len(parts) == 1:
        return [_clean_single_code(parts[0])]

    if len(parts) > MAX_BREAKDOWN_CODES:
        return [_clean_single_code(cleaned_noise)]

    base = parts[0]
    codes = [_clean_single_code(base)]
    for tail in parts[1:]:
        if re.fullmatch(r"\d+", tail):
            rebuilt = _rebuild_number_with_base(base, tail)
            codes.append(_clean_single_code(rebuilt))
        else:
            codes.append(_clean_single_code(tail))

    deduped = []
    for c in codes:
        if c not in deduped:
            deduped.append(c)
    return deduped

def clean_split_code_impl(text: str, is_cert_check=False):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return [], False
    original = text.strip()

    if _is_kept_as_is(original):
        return [original], True

    if _is_complex_month_year_pattern(original):
        return [original], True

    long_gap_resolved = _try_resolve_long_gap_duplicate(original)
    if long_gap_resolved is not None:
        original = long_gap_resolved

    cleaned_noise = _strip_polis_slip_noise(original)
    if not cleaned_noise:
        return [original], True

    if _is_kept_as_is(cleaned_noise):
        return [original], True

    sd_expanded = _try_expand_sd_range(cleaned_noise)
    if sd_expanded == "__KEEP_AS_IS__":
        return [original], True
    if sd_expanded is not None:
        return sd_expanded, False

    end_marker_resolved = _try_resolve_numeric_base_with_end_markers(original)
    if end_marker_resolved is not None:
        return [end_marker_resolved], False

    if "+" in cleaned_noise:
        _plus_parts_pcode_check = [p.strip() for p in _POLICY_SLIP_SPLIT_RE.split(cleaned_noise) if p.strip()]
        if len(_plus_parts_pcode_check) >= 2 and all(re.fullmatch(r"P\d+", p, flags=re.IGNORECASE) for p in _plus_parts_pcode_check):
            return [original], True

    if "+" in cleaned_noise:
        _plus_parts_long_check = [p.strip() for p in _POLICY_SLIP_SPLIT_RE.split(cleaned_noise) if p.strip()]
        if len(_plus_parts_long_check) >= 2 and all(re.fullmatch(r"\d{10,}", p) for p in _plus_parts_long_check):
            if len(_plus_parts_long_check) > MAX_BREAKDOWN_CODES:
                return [original], True
            codes = [_clean_single_code(p) for p in _plus_parts_long_check]
            deduped = []
            for c in codes:
                if c not in deduped:
                    deduped.append(c)
            return deduped, False

    if "+" in cleaned_noise:
        _plus_parts_check = [p.strip() for p in _POLICY_SLIP_SPLIT_RE.split(cleaned_noise) if p.strip()]
        if len(_plus_parts_check) >= 2 and all(not re.search(r"[A-Za-z]", p) and re.search(r"\d", p) for p in _plus_parts_check):
            _cleaned_check = [_clean_single_code(p) for p in _plus_parts_check]
            if len(set(_cleaned_check)) > 1:
                return [original], True

    _first_segment = cleaned_noise.split("+", 1)[0]
    _dash_in_first = re.search(r"(.*?)(\d+)\s*-\s*(\d+)\s*$", _first_segment.strip())
    if _dash_in_first and not re.search(r"[A-Za-z]", _dash_in_first.group(1)):
        return [original], True

    _dash_group_pat = re.compile(r"\d+\s*-\s*\d+")
    _dash_comma_pat = re.compile(r"\d+\s*-\s*\d+\s*,\s*\d+")
    if len(_dash_group_pat.findall(cleaned_noise)) >= 2 or _dash_comma_pat.search(cleaned_noise):
        return [original], True

    if "+" not in cleaned_noise and re.search(r"\d+\s*-\s*\d+", cleaned_noise):
        dash_result = _try_expand_dash_pattern(cleaned_noise)
        if dash_result == "__KEEP_AS_IS__":
            return [original], True
        if dash_result is not None:
            return dash_result, False

    if "/" in cleaned_noise:
        if _has_tba_alongside_others(cleaned_noise):
            rest = _strip_tba_tokens_keep_rest(cleaned_noise)
            if not rest:
                return [original], True
            return [rest], True
        elif "+" not in cleaned_noise:
            slash_result = _process_slash_pattern(cleaned_noise)
            if slash_result == "__KEEP_AS_IS__":
                return [original], True
            if slash_result is not None:
                return slash_result, False

    return _split_and_rebuild_plus(cleaned_noise), False

def clean_split_code(text: str) -> list:
    results, _ = clean_split_code_impl(text)
    return results

# ==============================================================================
# 2B. CLEANSING CERTIFICATE
# ==============================================================================
_CERT_SD_NORMALIZE_RE = re.compile(r"\bS\s*/\s*D\b", flags=re.IGNORECASE)
_CERT_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9])(\d{5,6})(?![A-Za-z0-9])")
_CERT_SD_PAIR_RE = re.compile(r"(\d{5,6})\s+SD\s+(\d{5,6})", flags=re.IGNORECASE)

def _pad_cert_digits(digits: str) -> str:
    return digits.zfill(CERT_DIGIT_LEN)

def _extract_cert_numbers_from_raw(text: str):
    if not text:
        return []

    found = []
    consumed_spans = []

    for m in _CERT_SD_PAIR_RE.finditer(text):
        start_s, end_s = m.group(1), m.group(2)
        before = text[: m.start(1)]
        has_prefix = bool(re.search(r"[A-Za-z0-9]", before))
        if not has_prefix:
            continue
        try:
            start_n, end_n = int(start_s), int(end_s)
        except ValueError:
            continue
        if start_n > end_n or (end_n - start_n + 1) > MAX_BREAKDOWN_CODES:
            continue
        for n in range(start_n, end_n + 1):
            found.append(_pad_cert_digits(str(n)))
        consumed_spans.append((m.start(), m.end()))

    if consumed_spans:
        masked_chars = list(text)
        for s, e in consumed_spans:
            for i in range(s, e):
                masked_chars[i] = "#"
        scan_text = "".join(masked_chars)
    else:
        scan_text = text

    for m in _CERT_TOKEN_RE.finditer(scan_text):
        before = scan_text[: m.start()]
        has_prefix = bool(re.search(r"[A-Za-z0-9]", before))
        if has_prefix:
            found.append(_pad_cert_digits(m.group(1)))

    return found

def _join_certificates(cert_list: list) -> str:
    if not cert_list:
        return ""

    seen = []
    for c in cert_list:
        if c not in seen:
            seen.append(c)

    if len(seen) == 1:
        return seen[0]

    is_consecutive = False
    try:
        nums = [int(c) for c in seen]
        is_consecutive = all(nums[i + 1] - nums[i] == 1 for i in range(len(nums) - 1))
    except ValueError:
        is_consecutive = False

    if len(seen) > CERT_RANGE_JOIN_THRESHOLD and is_consecutive:
        return f"{seen[0]} SD {seen[-1]}"
    return ", ".join(seen)

def build_certificate_value(raw_sertf_value) -> str:
    if raw_sertf_value is None or (isinstance(raw_sertf_value, float) and pd.isna(raw_sertf_value)):
        return ""

    text = str(raw_sertf_value).strip()
    if not text or text.lower() in ("nan", "none", "-"):
        return ""

    text = _CERT_SD_NORMALIZE_RE.sub("SD", text)

    _, is_kept_as_is = clean_split_code_impl(text, is_cert_check=True)
    if is_kept_as_is:
        return ""

    cert_candidates = _extract_cert_numbers_from_raw(text)
    return _join_certificates(cert_candidates)

# ==============================================================================
# 3. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    try:
        return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row, engine="calamine")
    except Exception:
        return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw: pd.DataFrame, keyword: str, cedant_column: str = "") -> pd.DataFrame:
    if cedant_column:
        if cedant_column not in df_raw.columns:
            raise ValueError(f"Kolom {cedant_column} tidak ada.")
        cedant_cols = [cedant_column]
    else:
        if not (cedant_cols := [c for c in df_raw.columns if "CEDANT" in str(c).upper()]):
            raise ValueError("Kolom CEDANT tidak ditemukan lewat auto-detect.")

    mask = pd.Series(False, index=df_raw.index)
    for col in cedant_cols:
        mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)

    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant '{keyword}': {len(df_filtered)} dari total {len(df_raw)} baris.")
    return df_filtered

# ==============================================================================
# 4. BUILD OUTPUT
# ==============================================================================
def _detect_key_columns(df: pd.DataFrame, insured_col: str, polis_col: str, slip_col: str):
    def _res(col: str, kw_or_fn, lbl: str):
        if col:
            if col not in df.columns:
                raise ValueError(f"{lbl}='{col}' tidak ada.")
            return col
        return next((c for c in df.columns if (kw_or_fn(str(c).upper()) if callable(kw_or_fn) else kw_or_fn in str(c).upper())), None)

    col_ins = _res(insured_col, "INSURED", "INSURED_COLUMN")
    col_pol = _res(polis_col, lambda c: c.strip() == "POLIS", "POLIS_COLUMN")
    col_slp = _res(slip_col, "SLIP", "SLIP_COLUMN")

    if None in (col_ins, col_pol, col_slp):
        raise ValueError("Kolom kunci tidak lengkap.")

    print(f"[INFO] Kolom terdeteksi -> INSURED='{col_ins}', POLIS='{col_pol}', SLIP='{col_slp}'")
    return col_ins, col_pol, col_slp

_GENERIC_VALUE_RE = re.compile(r"^\s*(?:TBA\.?|VAR(?:IOUS)?|P|-|NAN|NONE)\s*$", flags=re.IGNORECASE)

def _is_blank_value(value) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    s = str(value).strip()
    return s == "" or s.lower() in ("nan", "none")

def _informativeness_score(value: str) -> int:
    if _is_blank_value(value):
        return -1
    s = str(value).strip()
    if _GENERIC_VALUE_RE.match(s):
        return 0
    digit_count = len(re.findall(r"\d", s))
    return 1 + digit_count

def _pick_best_source(clsdt_value, fac_value) -> str:
    clsdt_score = _informativeness_score(clsdt_value)
    fac_score = _informativeness_score(fac_value)

    if clsdt_score < 0 and fac_score < 0:
        return ""
    if clsdt_score < 0:
        return str(fac_value).strip()
    if fac_score < 0:
        return str(clsdt_value).strip()
    if fac_score > clsdt_score:
        return str(fac_value).strip()
    return str(clsdt_value).strip()

_NON_ALNUM_RE = re.compile(r"[^A-Za-z0-9]")

def _compact_value(value: str) -> str:
    if not isinstance(value, str):
        return value
    return _NON_ALNUM_RE.sub("", value)

def _compact_list(values: list) -> list:
    return [_compact_value(v) for v in values]

def _compact_split_code_results(results: list, is_kept_as_is: bool) -> list:
    if is_kept_as_is:
        return results
    return [_compact_value(v) for v in results]

def _breakdown_all_rows(
    df: pd.DataFrame,
    col_ins: str,
    col_pol: str,
    col_slp: str,
    col_clsdt_pol: str = None,
    col_clsdt_slp: str = None,
    col_clsdt_sertf: str = None,
):
    ins_arr = df[col_ins].to_numpy()
    pol_arr = df[col_pol].to_numpy()
    slp_arr = df[col_slp].to_numpy()

    clsdt_pol_arr = df[col_clsdt_pol].to_numpy() if col_clsdt_pol and col_clsdt_pol in df.columns else [None] * len(df)
    clsdt_slp_arr = df[col_clsdt_slp].to_numpy() if col_clsdt_slp and col_clsdt_slp in df.columns else [None] * len(df)
    clsdt_sertf_arr = df[col_clsdt_sertf].to_numpy() if col_clsdt_sertf and col_clsdt_sertf in df.columns else [None] * len(df)

    ins_all, pol_all, slp_all, cert_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val, clsdt_pol_val, clsdt_slp_val, clsdt_sertf_val in zip(
        ins_arr, pol_arr, slp_arr, clsdt_pol_arr, clsdt_slp_arr, clsdt_sertf_arr
    ):
        ins = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []

        raw_pol = pol_val if pd.notna(pol_val) else ""
        raw_slp = slp_val if pd.notna(slp_val) else ""

        if col_clsdt_pol and pd.notna(clsdt_pol_val):
            source_pol = _pick_best_source(clsdt_pol_val, raw_pol)
        else:
            source_pol = str(raw_pol).strip()

        if col_clsdt_slp and pd.notna(clsdt_slp_val):
            source_slp = _pick_best_source(clsdt_slp_val, raw_slp)
        else:
            source_slp = str(raw_slp).strip()

        if source_pol:
            pol_results, pol_kept_as_is = clean_split_code_impl(source_pol)
            pol = _compact_split_code_results(pol_results, pol_kept_as_is)
        else:
            pol = []

        if source_slp:
            slp_results, slp_kept_as_is = clean_split_code_impl(source_slp)
            slp = _compact_split_code_results(slp_results, slp_kept_as_is)
        else:
            slp = []

        cert_str = build_certificate_value(clsdt_sertf_val)
        cert = [cert_str] if cert_str else []

        ins_all.append(ins)
        pol_all.append(pol)
        slp_all.append(slp)
        cert_all.append(cert)

        max_ins = max(max_ins, len(ins))
        max_pol = max(max_pol, len(pol))
        max_slp = max(max_slp, len(slp))

    return ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_ins, col_pol, col_slp = _detect_key_columns(df_filtered, INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN)

    col_clsdt_pol = CLSDT_POLIS_COLUMN if CLSDT_POLIS_COLUMN in df_filtered.columns else None
    col_clsdt_slp = CLSDT_SLIP_COLUMN if CLSDT_SLIP_COLUMN in df_filtered.columns else None
    col_clsdt_sertf = CLSDT_SERTF_COLUMN if CLSDT_SERTF_COLUMN in df_filtered.columns else None

    ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp = _breakdown_all_rows(
        df_filtered, col_ins, col_pol, col_slp, col_clsdt_pol, col_clsdt_slp, col_clsdt_sertf
    )

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)
    output_cols_data = {}

    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == col_ins:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in ins_all]
        if col == col_pol:
            max_cert = max((len(c) for c in cert_all), default=0)
            max_iter = max(1, max_pol, max_cert)
            
            for i in range(1, max_iter + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in pol_all]
                output_cols_data[f"CERTIFICATE_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_all]
        if col == col_slp:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slp_all]

    df_master = pd.DataFrame(output_cols_data)
    
    drop_cols = [c for c in df_master.columns if "_CLN_" in c and df_master[c].astype(str).str.strip().eq("").all()]
    drop_cols.extend([c for c in df_master.columns if "CERTIFICATE_" in c and c != "CERTIFICATE_1" and df_master[c].astype(str).str.strip().eq("").all()])
    
    return df_master.drop(columns=drop_cols)

# ==============================================================================
# 5. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook = load_workbook(output_path)
    worksheet = workbook.active
    text_col_indices = [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)
    workbook.save(output_path)

# ==============================================================================
# 6. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70 + "\n CLEANSING DATA 2 - OSBAL | CEDANT: IBS\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(_RAW_DIR) if f.endswith(".xlsx")]
        if not candidates:
            raise FileNotFoundError(f"File Excel sumber tidak ditemukan di: {_RAW_DIR}")
        input_path = os.path.join(_RAW_DIR, candidates[0])
        print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")

    os.makedirs(_CLEAN_DIR, exist_ok=True)

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    df_ibs = filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)

    if not df_ibs.empty:
        df_hasil = build_output(df_ibs)
        save_with_text_format(df_hasil, OUTPUT_FILE)
        print("=" * 70 + f"\n[SUCCESS] Selesai. Output: '{OUTPUT_FILE}' ({len(df_hasil)} baris)\n" + "=" * 70)
    else:
        print("[WARNING] Tidak ada baris yang cocok dengan filter.")

if __name__ == "__main__":
    main()