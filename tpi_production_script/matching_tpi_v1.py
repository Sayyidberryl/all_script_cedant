import os
import re
import sys
import time
import functools

sys.modules['numexpr'] = None
sys.modules['bottleneck'] = None

import numpy as np
import pandas as pd
from dotenv import load_dotenv
from sqlalchemy import create_engine


# =============================================================================
# KONFIGURASI  (CEDANT: TPI — PT Asuransi Tugu Pratama Indonesia)
#
# Sumber arsitektur : prod_sus_1_new.py  (ACA reference)
# Referensi         : matching_IBS.md, teknis_matching_aca.md
# Dokumen workflow  : Workflow Teknis Suspend 2026.docx
#
# Delta utama vs IBS / ACA:
#   1. BORDERO / Open Cover Marine Cargo — TIDAK ADA untuk TPI.
#      Seluruh kode bordero dihapus bersih dari pipeline.
#   2. FASE 1 IBS (DESC_3_CLEAN) — TIDAK DITERAPKAN untuk TPI.
#      TPI tidak memiliki rules tambahan untuk pengecekan ke DESC 3
#      pada data suspend, sehingga FASE 1 DESC_3_CLEAN dihapus.
#      Pipeline TPI mengikuti urutan ACA standar:
#      Polis -> Slip -> Insured -> Currency & Periode -> Narrowing Non-Destruktif
#   3. Naming kolom TPI:
#      Suspend : INSURED_ORI, INSURED_CLEAN_N, POLIS_ORI, POLIS_CLEAN_1,
#                CERTIFICATE_1, SLIP_NO_ORI, SLIP_CLEAN_1
#      OSBAL   : FAC_POLICY_NO_ORI, POLIS_CLEAN_N, CERTIFICATE_N,
#                FAC_SLIP_ORI, SLIP_CLEAN_N, FAC_INSURED_ORI, INSURED_CLEAN_N,
#                CCOS_REF_CODE, CCOS_CURR, FAC_COM_DATE
#      FACUL   : FAC_CODE (kolom FAC CODE), pola kolom bersih sama dg OSBAL
# =============================================================================

_BASE_DIR  = os.path.dirname(os.path.abspath(__file__))   # folder TPI/
_CLEAN_DIR = os.path.join(_BASE_DIR, "data_clean")         # hasil cleansing
_RAW_DIR   = os.path.join(_BASE_DIR, "data_mentah")        # data mentah


def _find_file_in(directory: str, basenames: list) -> str:
    """Cari file pertama yang ada dari daftar kandidat nama dalam direktori tertentu."""
    for name in basenames:
        p = os.path.join(directory, name)
        if os.path.exists(p):
            return p
    return os.path.join(directory, basenames[0])


# [TPI] Path file menggunakan suffix/nama TPI
SUSPEND_FILE = _find_file_in(_CLEAN_DIR, ["tpi_output_suspend.xlsx"])
OSBAL_FILE   = _find_file_in(_CLEAN_DIR, ["tpi_output_osbal.xlsx"])
FACUL_FILE   = _find_file_in(_CLEAN_DIR, ["tpi_output_facul.xlsx"])
SLIPDB_FILE  = _find_file_in(_RAW_DIR, [
    "5. Database Rekap RI Slip.xlsx",
    "5.Database Rekap RI Slip.xlsx",
    "Database Rekap RI Slip.xlsx",
])
OUTPUT_FILE  = os.path.join(_BASE_DIR, "matching_tpi.xlsx")

# [TPI] Tidak ada BORDERO_FILE.
# TPI tidak punya data Open Cover Marine Cargo. Variabel ini tidak
# dideklarasikan agar jalur bordero tidak pernah dieksekusi secara struktural.

# Kolom FAC CODE di masing-masing tabel referensi
OSBAL_FACODE_COL  = "CCOS_REF_CODE"
FACUL_FACODE_COL  = "FAC_CODE"
SLIPDB_FACODE_COL = "FAC_CODE"   # dapat di-override runtime sesuai header aktual

# [TPI] Kolom polis/slip/insured SUSPEND TPI (dari script cleaning-suspend.py)
# Tidak ada SUSPEND_SLIP_SOURCE_COL (DESC_3_CLEAN) karena TPI tidak punya FASE 1 khusus
CLSDT_POLIS_COL = "POLIS_ORI"       # kolom polis original di suspend TPI
CLSDT_SLIP_COL  = "SLIP_NO_ORI"     # kolom slip original di suspend TPI
FAC_POLIS_COL   = "FAC_POLICY_NO"   # kolom wildcard polis (jika ada di suspend)
FAC_SLIP_COL    = "FAC_SLIP"        # kolom wildcard slip (jika ada di suspend)

# [TPI] Nama kolom currency/periode sesuai data TPI
SUSPEND_CURR_COL = "CURR ORI"    # kolom currency di tabel Suspend TPI
OSBAL_CURR_COL   = "CCOS_CURR"  # kolom currency di tabel OSBAL TPI
SUSPEND_DATE_COL = "RECEIPT DATE"
OSBAL_DATE_COL   = "FAC_COM_DATE"

# Marker baris administratif yang dikecualikan dari matching
_EXCLUDED_REF_MARKERS = ["HUTANG PIUTANG", "DATA SUSPENSE"]

# Marker LINESLIP - pakai RECEIPT DATE - 1 bulan sebagai tanggal efektif
_LINESLIP_MARKERS = ("LINESLIP", "LINE SLIP")

_MIN_TOKEN_LEN = 3
_TOKEN_RE      = re.compile(r"[^A-Z0-9]+")
_DELIMITER_RE  = re.compile(r"[,;|+\s]+")

# Token generik yang dikecualikan saat menghitung overlap INSURED (kurangi false positive).
_INSURED_STOPWORDS: frozenset = frozenset({
    "PT", "CV", "TBK", "UD", "PD", "NV", "LLC", "LTD", "INC",
    "GROUP", "GROUPS",
    "INDONESIA", "INDO",
    "INDUSTRY", "INDUSTRIES", "INDUSTRI",
    "INTERNATIONAL", "INTERNASIONAL",
    "COMPANY", "COMPANIES", "CORP", "CORPORATION",
    "ENTERPRISE", "ENTERPRISES",
    "TRADING", "TRADE",
    "UTAMA", "MANDIRI", "MAKMUR", "SEJAHTERA",
})

# Overlap coefficient = |A intersect B| / min(|A|, |B|) setelah exclude stopword.
_INSURED_LIKE_THRESHOLD: float = 0.65

# Threshold ketat untuk POLIS/SLIP (identifier dokumen — harus lebih presisi).
_POLIS_LIKE_THRESHOLD: float  = 0.80

# [TPI] Kolom join saat akumulasi suspend — nama mengikuti pola ACA
_SUSPEND_JOIN_COLS = [
    "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "RECEIPT DATE",
    "CEDANT NAME", "CEDANT SHRT NAME",
    "INSURED_ORI", "INSURED_CLEAN_1", "INSURED_CLEAN_2",
    "CURR ORI", "CURR PAY", "AMOUNT PAY",
    "POLIS_ORI", "POLIS_CLEAN_1",
    "SLIP_NO_ORI", "SLIP_CLEAN_1",
    "DESC 1", "DESC 2", "DESC 3", "DESC 4",
    "STATUS", "REC_TYPE",
]
_SUSPEND_SUM_COLS = ["AMOUNT ORI"]

# [TPI] FINAL_COLUMNS: 38 kolom standar output (mengikuti pola ACA)
FINAL_COLUMNS = [
    "CCOS_DOC_NO", "CCOS_REF_CODE",
    "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "RECEIPT DATE",
    "CEDANT NAME", "CEDANT SHRT NAME",
    "INSURED_ORI", "INSURED_1", "INSURED_2",
    "CURR ORI", "AMOUNT ORI", "AMOUNT_ORI_MIN1",
    "CURR PAY", "AMOUNT PAY",
    "CCOS_OR_BAL", "CCOS_BAL_DUE", "DIFERENCE",
    "FLAG_PROD",
    "POLIS",
    "POLICY_CLEAN_1", "SERTIF_CLEAN_1",
    "POLICY_CLEAN_2", "SERTIF_CLEAN_2",
    "POLICY_CLEAN_3", "SERTIF_CLEAN_3",
    "POLICY_CLEAN_4", "POLICY_CLEAN_5",
    "SLIP_NO", "SLIP_NO_CLN",
    "DESC 1", "DESC 2", "DESC 3", "DESC 4",
    "STATUS", "REC_TYPE",
    "SKENARIO",
]


# =============================================================================
# STRING UTILITIES
# =============================================================================

@functools.lru_cache(maxsize=131072)
def _norm_cached(s: str) -> str:
    # Normalisasi string - cached karena dipanggil jutaan kali
    text = s.strip().upper().replace("S/D", "SD")
    return re.sub(r"\s+", " ", text) if "  " in text else text


def _normalize(value) -> str:
    # Uppercase + strip + collapse spasi, return '' untuk null
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return _norm_cached(str(value))


def _format_sertif(val) -> str:
    # Format nilai sertif jadi string 6-digit (contoh: 7 -> '000007')
    if pd.isna(val) or str(val).strip() == "":
        return ""
    if isinstance(val, (int, float)):
        return str(int(val)).zfill(6)
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s.zfill(6) if s.isdigit() else s


def _tokenize(text: str) -> frozenset:
    # Pecah string jadi frozenset token alfanumerik (min 3 karakter)
    if not text:
        return frozenset()
    return frozenset(t for t in _TOKEN_RE.split(text) if len(t) >= _MIN_TOKEN_LEN)


@functools.lru_cache(maxsize=100000)
def _split_composite(text: str) -> frozenset:
    # Pecah string berdelimiter (koma/plus/spasi/titik koma/pipe) jadi token
    if not text:
        return frozenset()
    return frozenset(tok.strip() for tok in _DELIMITER_RE.split(text) if tok.strip())


def _get_clean_cols(columns: list, prefix: str) -> list:
    # Ambil nama kolom yang berawalan prefix (case-insensitive)
    if not columns:
        return []
    return [c for c in columns if str(c).lower().startswith(prefix.lower())]


def _map_scenario(label: str) -> str:
    # Petakan label internal ke nama skenario resmi
    if not label or label in ("Unmatching", "UNMATCHED"):
        return "Unmatching"
    if "INSURED" in label:
        return "Insured only"
    if "SLIP" in label:
        return "Slip only"
    if "POLIS" in label:
        return "Polis only"
    return label


def _clean_base_match(scenario: str) -> str:
    """Bersihkan suffix tambahan dari skenario dasar."""
    s = str(scenario or "").strip()
    for to_remove in [" (Beda Currency)", " (Beda Periode)", " (Matching >1 fac code)"]:
        s = s.replace(to_remove, "")
    s = s.strip()
    if s.startswith("Slip + Polis"):
        s = s.replace("Slip + Polis", "Polis + Slip")
    if " only + " in s:
        s = s.replace(" only + ", " + ")
    return s


def _format_final_scenario(
    base_scenario: str,
    source: str,
    is_beda_curr: bool,
    is_beda_per: bool,
    num_fac: int,
) -> str:
    """Format skenario final sesuai pola ACA (tanpa Bordero untuk TPI).

    Daftar skenario yang dihasilkan:
      Insured only
      Insured only (Matching >1 fac code)
      Slip only
      Unmatching
      Slip + Insured
      Insured only (Beda Currency)
      Slip only (Matching >1 fac code)
      Slip + Insured (Beda Currency)
      Slip only (Beda Currency)
      ... dst.
    """
    if not source or not base_scenario or str(base_scenario).strip().upper() in ("UNMATCHING", "UNMATCHED"):
        return "Unmatching"

    base = _clean_base_match(base_scenario)
    if not base or base.upper() in ("UNMATCHING", "UNMATCHED"):
        return "Unmatching"

    if is_beda_curr:
        return f"{base} (Beda Currency)"
    elif is_beda_per:
        return f"{base} (Beda Periode)"
    elif num_fac > 1:
        return f"{base} (Matching >1 fac code)"
    else:
        return base


# =============================================================================
# ROW DICT HELPERS
# =============================================================================

def _is_excluded_ref_row(row: dict, polis_cols: list, slip_cols: list) -> bool:
    # True jika baris referensi adalah baris administratif
    for col in list(polis_cols) + ["polis_ori", "FAC_POLICY_NO_ORI"] + list(slip_cols) + ["slip_ori", "FAC_SLIP_ORI"]:
        val = _normalize(row.get(col, ""))
        if val and any(m in val for m in _EXCLUDED_REF_MARKERS):
            return True
    return False


def _is_lineslip_row(sus_row: dict, insured_cols: list) -> bool:
    # True jika baris suspend mengandung marker LINESLIP di kolom insured
    # [TPI] Kolom ori adalah "INSURED_ORI"
    for col in list(insured_cols) + ["INSURED_ORI"]:
        val = _normalize(sus_row.get(col, ""))
        if val and any(m in val for m in _LINESLIP_MARKERS):
            return True
    return False


def _collect_clean_values(row: dict, cols: list) -> list:
    # Kumpulkan nilai bersih unik (non-empty) dari kolom yang diberikan
    seen, result = set(), []
    for col in cols:
        val = _normalize(row.get(col, ""))
        if val and val not in seen:
            seen.add(val)
            result.append(val)
    return result


def _collect_all_values(row: dict, clean_cols: list, ori_col: str) -> list:
    # Gabung nilai clean menjadi list unik
    return list(set(_collect_clean_values(row, clean_cols)))


def _get_effective_sus_cols(sus_row: dict, base_polis_cols: list, base_slip_cols: list) -> tuple:
    """Tentukan kolom efektif polis/slip per baris.

    [TPI] Mengikuti logika ACA standar:
    - eff_polis: POLIS_ORI (jika terisi) + POLIS_CLEAN_N
    - eff_slip:  SLIP_NO_ORI (jika terisi) + SLIP_CLEAN_N
    Urutan fase: Polis -> Slip -> Insured (tidak ada FASE 1 DESC_3_CLEAN)
    """
    clsdt_polis = _normalize(sus_row.get(CLSDT_POLIS_COL, ""))
    clsdt_slip  = _normalize(sus_row.get(CLSDT_SLIP_COL, ""))
    fac_polis   = _normalize(sus_row.get(FAC_POLIS_COL,  ""))
    fac_slip    = _normalize(sus_row.get(FAC_SLIP_COL,   ""))

    if clsdt_polis:
        eff_polis  = list(dict.fromkeys([CLSDT_POLIS_COL] + list(base_polis_cols)))
        like_polis = list(dict.fromkeys([CLSDT_POLIS_COL] + list(base_polis_cols)))
    else:
        eff_polis  = list(dict.fromkeys(list(base_polis_cols)))
        like_polis = list(dict.fromkeys(list(base_polis_cols) + ([FAC_POLIS_COL] if fac_polis else [])))

    if clsdt_slip:
        eff_slip  = list(dict.fromkeys([CLSDT_SLIP_COL] + list(base_slip_cols)))
        like_slip = list(dict.fromkeys([CLSDT_SLIP_COL] + list(base_slip_cols)))
    else:
        eff_slip  = base_slip_cols
        like_slip = list(base_slip_cols) + ([FAC_SLIP_COL] if fac_slip else [])

    return eff_polis, eff_slip, like_polis, like_slip


def _ref_has_value(ref_row: dict, clean_cols: list, ori_col: str, query_values: list) -> bool:
    if not query_values:
        return False
    ref_vals = set(_collect_clean_values(ref_row, clean_cols))
    if set(query_values) & ref_vals:
        return True
    for rv in ref_vals:
        if any(v in _split_composite(rv) for v in query_values if v and len(v) >= 2):
            return True
    return any(
        v in rv
        for v in query_values if v and len(v) >= 10
        for rv in ref_vals if rv
    )


def _slip_in_range(sus_slip: str, osbal_row: dict) -> bool:
    """Cek apakah nomor slip suspend berada di dalam range slip OSBAL."""
    if not sus_slip:
        return False
    s_clean = re.sub(r'\D', '', str(sus_slip))
    if not s_clean:
        return False
    sl1 = re.sub(r'\D', '', str(osbal_row.get("SLIP_CLEAN_1", "") or osbal_row.get("clean slip 1", "")))
    sl2 = re.sub(r'\D', '', str(osbal_row.get("SLIP_CLEAN_2", "") or osbal_row.get("clean slip 2", "")))
    if sl1 and sl2 and len(sl1) == len(s_clean) and len(sl2) == len(s_clean):
        try:
            n_sus = int(s_clean)
            n_sl1 = int(sl1)
            n_sl2 = int(sl2)
            if min(n_sl1, n_sl2) <= n_sus <= max(n_sl1, n_sl2):
                return True
        except ValueError:
            pass
    return False


# =============================================================================
# INDEX BUILDERS
# =============================================================================

def _build_exact_index(rows: list, cols: list, excluded: set = None) -> dict:
    # Bangun inverted index: value -> [indeks baris] untuk exact match
    index: dict = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        for col in cols:
            val = _normalize(row.get(col, ""))
            if val:
                index.setdefault(val, []).append(i)
    return index


def _build_facode_index(rows: list, col: str, excluded: set = None) -> dict:
    # Bangun inverted index: FAC_CODE -> [indeks baris]
    index: dict = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        val = _normalize(row.get(col, ""))
        if val:
            index.setdefault(val, []).append(i)
    return index


def _build_token_index(rows: list, clean_cols: list, ori_col: str, excluded: set = None) -> tuple:
    # Bangun token inverted index + cache token per baris untuk fuzzy matching
    token_index: dict = {}
    token_cache: dict = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        tokens: frozenset = frozenset()
        for c in clean_cols:
            t = _normalize(row.get(c, ""))
            if t:
                tokens = tokens | _tokenize(t)
        token_cache[i] = tokens
        for tok in tokens:
            token_index.setdefault(tok, set()).add(i)
    return token_index, token_cache


def _build_lookup(rows: list, polis_cols: list, slip_cols: list,
                  insured_cols: list, facode_col: str = None, label: str = "") -> tuple:
    """Bangun semua struktur lookup untuk satu tabel referensi (10-tuple).

    Index tuple identik dengan prod_sus_1_new.py:
    [0] slip exact  [1] polis exact  [2] insured exact
    [3][4][5] unused
    [6] tok_slip  [7] tok_polis  [8] tok_insured  [9] facode_index
    """
    if not rows:
        if label:
            print(f"  {label:<6} done (0.0s)", flush=True)
        return ({}, {}, {}, {}, {}, {}, ({}, {}), ({}, {}), ({}, {}), {})

    t0 = time.perf_counter()
    excluded = {i for i, r in enumerate(rows) if _is_excluded_ref_row(r, polis_cols, slip_cols)}

    facode_index = _build_facode_index(rows, facode_col, excluded=excluded) if facode_col else {}
    lookup = (
        _build_exact_index(rows, slip_cols,    excluded=excluded),   # [0] slip
        _build_exact_index(rows, polis_cols,   excluded=excluded),   # [1] polis
        _build_exact_index(rows, insured_cols, excluded=excluded),   # [2] insured
        {}, {}, {},                                                   # [3][4][5] unused
        _build_token_index(rows, slip_cols,    "slip_ori",    excluded=excluded),  # [6]
        _build_token_index(rows, polis_cols,   "polis_ori",   excluded=excluded),  # [7]
        _build_token_index(rows, insured_cols, "insured_ori", excluded=excluded),  # [8]
        facode_index,                                                 # [9]
    )

    elapsed = time.perf_counter() - t0
    if label:
        exc_info = f" - {len(excluded):,} baris dikecualikan" if excluded else ""
        print(f"  {label:<6} done ({elapsed:.1f}s){exc_info}", flush=True)
    return lookup


# =============================================================================
# MATCH ENGINE
# =============================================================================




def _ref_matches_insured(ref_row: dict, insured_ref_cols: list, sus_ins_vals: list) -> bool:
    """Cek apakah insured suspend cocok dengan referensi (INSURED_CLEAN atau FAC_INSURED_ORI / QQ / /)."""
    if not sus_ins_vals:
        return False
    # 1. Cek exact lewat kolom clean standar
    if _ref_has_value(ref_row, insured_ref_cols, "insured_ori", sus_ins_vals):
        return True

    # 2. [TPI] Cek token match lewat kolom-kolom INSURED_CLEAN
    for c in insured_ref_cols:
        val = ref_row.get(c, "")
        if val and str(val) != "nan":
            toks = _tokenize(str(val)) - _INSURED_STOPWORDS
            for sus_v in sus_ins_vals:
                sus_toks = _tokenize(sus_v) - _INSURED_STOPWORDS
                if toks and sus_toks:
                    if toks.issubset(sus_toks) or sus_toks.issubset(toks):
                        return True
                    common = len(toks & sus_toks)
                    if (common / min(len(toks), len(sus_toks))) >= 0.6:
                        return True

    # 3. [TPI] Cek FAC_INSURED_ORI (termasuk subsidiary setelah 'QQ' / 'C/O' / '/')
    fac_ins_ori = _normalize(ref_row.get("FAC_INSURED_ORI", ""))
    if fac_ins_ori:
        subsidiary_parts = [p.strip() for p in fac_ins_ori.replace('/', ' QQ ').replace('C/O', ' QQ ').split('QQ') if p.strip()]
        subsidiary_parts.append(fac_ins_ori)

        for ins_v in sus_ins_vals:
            if not ins_v:
                continue
            sus_tokens = _tokenize(ins_v) - _INSURED_STOPWORDS
            for sp in subsidiary_parts:
                sp_tokens = _tokenize(sp) - _INSURED_STOPWORDS
                if sus_tokens and sp_tokens:
                    if sp_tokens.issubset(sus_tokens) or sus_tokens.issubset(sp_tokens):
                        return True
                    common = len(sus_tokens & sp_tokens)
                    if common > 0 and (common / min(len(sus_tokens), len(sp_tokens))) >= 0.6:
                        return True
    return False


def _exact_match(query_values: list, index: dict) -> set:
    # Kembalikan indeks baris yang cocok exact dengan salah satu nilai query
    return {i for v in query_values if v and v in index for i in index[v]}


def _like_match(
    query_values: list,
    token_index: dict,
    rows: list,
    ref_cols: list,
    ori_col: str = None,
    stopwords: frozenset = frozenset(),
    overlap_threshold: float = _INSURED_LIKE_THRESHOLD,
    match_log: list = None,
) -> list:
    """Kembalikan indeks baris yang mirip dengan query, tiga kriteria berurutan.

    Kriteria kecocokan (ANY -> match):
      SUBSTR  - seluruh string query ada sebagai substring di nilai referensi.
      SUBSET  - semua token signifikan query adalah subset dari token referensi.
      OVERLAP - overlap coefficient |A intersect B| / min(|A|, |B|) >= threshold.
    """
    valid_queries = [v for v in query_values if v and len(v) >= 2]
    if not valid_queries:
        return []

    query_tokens: frozenset = frozenset()
    for v in valid_queries:
        query_tokens = query_tokens | _tokenize(v)

    sig_query_tokens = query_tokens - stopwords

    candidates: set = set()
    for tok in query_tokens:
        if tok in token_index:
            candidates |= token_index[tok]
    if not candidates:
        return []

    matched_set: dict = {}
    for i in candidates:
        ref_values = [_normalize(rows[i].get(c, "")) for c in ref_cols if rows[i].get(c)]

        substr_ok = any(q in rv for q in valid_queries for rv in ref_values if rv)
        if substr_ok:
            if i not in matched_set:
                matched_set[i] = "SUBSTR"
            continue

        if isinstance(token_index, dict):
            ref_tokens: frozenset = frozenset()
            for rv in ref_values:
                ref_tokens = ref_tokens | _tokenize(rv)
        else:
            ref_tokens = frozenset()

        if not ref_tokens:
            continue

        sig_ref_tokens = ref_tokens - stopwords

        if sig_query_tokens and sig_query_tokens.issubset(ref_tokens):
            if i not in matched_set:
                matched_set[i] = "SUBSET"
            continue

        if sig_query_tokens and sig_ref_tokens:
            intersection = len(sig_query_tokens & sig_ref_tokens)
            denom = min(len(sig_query_tokens), len(sig_ref_tokens))
            if denom > 0 and (intersection / denom) >= overlap_threshold:
                if i not in matched_set:
                    matched_set[i] = "OVERLAP"

    if not matched_set:
        return []

    if match_log is not None:
        for idx, method in matched_set.items():
            match_log.append((idx, method))

    return list(matched_set.keys())


# =============================================================================
# NARROWING
# =============================================================================

def _rematch_strict(
    suspend_row: dict, osbal_rows: list, lookup_osbal: tuple,
    polis_sus_cols: list, slip_sus_cols: list, insured_sus_cols: list,
    polis_ref_cols: list, slip_ref_cols: list, insured_ref_cols: list,
    facode_osbal_idx: dict,
    slipdb_rows: list = None, lookup_slipdb: tuple = None,
    facul_rows: list = None, lookup_facul: tuple = None,
    polis_sus_like_cols: list = None,
    effective_sus_date=None, is_lineslip: bool = False
) -> tuple:
    """Rematch ketat dengan filter currency + periode (R1-R5 standar ACA).

    [TPI] Tidak ada jalur R0 berbasis DESC_3_CLEAN (TPI tidak punya FASE 1 ini).
    Implementasi standar R1-R5 ACA dipakai apa adanya.
    """
    sus_curr = _normalize(suspend_row.get(SUSPEND_CURR_COL, ""))

    def _strict_filter(idxs):
        res = []
        for i in idxs:
            if sus_curr and _normalize(osbal_rows[i].get(OSBAL_CURR_COL, "")) != sus_curr:
                continue
            if pd.notna(effective_sus_date):
                com_date = osbal_rows[i].get("_com_date_parsed")
                if pd.notna(com_date):
                    if is_lineslip:
                        if com_date.year != effective_sus_date.year or com_date.month != effective_sus_date.month:
                            continue
                    else:
                        if effective_sus_date < com_date:
                            continue
            res.append(i)
        return res

    polis_vals = _collect_clean_values(suspend_row, polis_sus_cols)

    # R1: OSBAL polis exact
    curr_hits = _strict_filter(_exact_match(polis_vals, lookup_osbal[1]))
    if curr_hits:
        slip_vals = _collect_clean_values(suspend_row, slip_sus_cols)
        confirmed = [i for i in curr_hits if any(
            _normalize(osbal_rows[i].get(c, "")) in slip_vals
            for c in slip_ref_cols if _normalize(osbal_rows[i].get(c, ""))
        )]
        chosen = confirmed if confirmed else curr_hits
        return "OSBAL", ("Polis + Slip" if confirmed else "Polis only"), chosen, True

    # R2: OSBAL polis LIKE
    like_cols = polis_sus_like_cols if polis_sus_like_cols is not None else polis_sus_cols
    if lookup_osbal[7] and lookup_osbal[7][0] and like_cols:
        like_vals = _collect_clean_values(suspend_row, like_cols)
        if like_vals:
            curr_hits = _strict_filter(_like_match(
                like_vals, lookup_osbal[7][0], osbal_rows, polis_ref_cols,
                stopwords=frozenset(),
                overlap_threshold=_POLIS_LIKE_THRESHOLD,
            ))
            if curr_hits:
                return "OSBAL", "POLIS_LIKE", curr_hits, True

    # R3: SLIPDB
    if slipdb_rows and lookup_slipdb:
        hits = _exact_match(polis_vals, lookup_slipdb[1])
        if hits:
            res_idx, ok = _resolve_facode([slipdb_rows[i] for i in hits], SLIPDB_FACODE_COL, facode_osbal_idx, osbal_rows)
            if ok:
                curr_res = _strict_filter(res_idx)
                if curr_res:
                    return "SLIPDB", "Polis only", curr_res, True

    # R4: FACUL
    if facul_rows and lookup_facul:
        hits = _exact_match(polis_vals, lookup_facul[1])
        if hits:
            res_idx, ok = _resolve_facode([facul_rows[i] for i in hits], FACUL_FACODE_COL, facode_osbal_idx, osbal_rows)
            if ok:
                curr_res = _strict_filter(res_idx)
                if curr_res:
                    return "FACUL", "Polis only", curr_res, True

    # R5: OSBAL insured fallback
    # [TPI] Jika slip suspend ada di baris matched OSBAL, label "Slip + Insured" bukan "Insured only".
    insured_vals_r5 = _collect_clean_values(suspend_row, insured_sus_cols)
    curr_hits = _strict_filter(_exact_match(insured_vals_r5, lookup_osbal[2]))
    if curr_hits:
        slip_check_vals = _collect_clean_values(suspend_row, slip_sus_cols)
        r5_slip_confirmed = any(
            any(_normalize(osbal_rows[i].get(c, "")) in slip_check_vals
                for c in slip_ref_cols if _normalize(osbal_rows[i].get(c, "")))
            for i in curr_hits
        ) if slip_check_vals else False
        scen_r5 = "Slip + Insured" if r5_slip_confirmed else "Insured only"
        return "OSBAL", scen_r5, curr_hits, True

    if lookup_osbal[8] and lookup_osbal[8][0] and insured_vals_r5:
        curr_hits = _strict_filter(_like_match(
            insured_vals_r5, lookup_osbal[8][0], osbal_rows, insured_ref_cols,
            stopwords=_INSURED_STOPWORDS,
            overlap_threshold=_INSURED_LIKE_THRESHOLD,
        ))
        if curr_hits:
            slip_check_vals = _collect_clean_values(suspend_row, slip_sus_cols)
            r5_slip_confirmed = any(
                any(_normalize(osbal_rows[i].get(c, "")) in slip_check_vals
                    for c in slip_ref_cols if _normalize(osbal_rows[i].get(c, "")))
                for i in curr_hits
            ) if slip_check_vals else False
            scen_like = "Slip + Insured" if r5_slip_confirmed else "INSURED_LIKE"
            return "OSBAL", scen_like, curr_hits, True

    return None


# =============================================================================
# FAC CODE RESOLUTION
# =============================================================================

def _resolve_facode(ref_rows_matched: list, facode_col: str,
                    facode_osbal_idx: dict, osbal_rows: list) -> tuple:
    # Resolve FAC_CODE dari SLIPDB/FACUL ke indeks baris OSBAL
    fac_codes = {_normalize(r.get(facode_col, "")) for r in ref_rows_matched}
    fac_codes.discard("")
    if not fac_codes:
        return [], False
    osbal_indices: set = set()
    for fac in fac_codes:
        if fac in facode_osbal_idx:
            osbal_indices.update(facode_osbal_idx[fac])
    result = list(osbal_indices)
    return result, bool(result)


# =============================================================================
# OUTPUT BUILDERS
# =============================================================================

def _aggregate_ccos(osbal_rows: list, all_fac_osbal_rows: list = None) -> dict:
    # Agregasi nilai finansial OSBAL
    fac_codes = {str(r.get("CCOS_REF_CODE", "")).strip() for r in osbal_rows
                 if str(r.get("CCOS_REF_CODE", "")).strip()}
    multi_fac = len(fac_codes) > 1
    bal_due_rows = all_fac_osbal_rows if (all_fac_osbal_rows is not None and not multi_fac) else osbal_rows

    def join_unique(key):
        seen, vals = set(), []
        for r in osbal_rows:
            v = str(r.get(key, "")).strip()
            if v and v not in seen:
                seen.add(v)
                vals.append(v)
        return ", ".join(vals)

    def sum_col(key, rows):
        if multi_fac:
            return np.nan
        total = 0.0
        for r in rows:
            v = r.get(key, 0)
            if pd.isna(v):
                continue
            if isinstance(v, (int, float)):
                total += float(v)
            else:
                s = str(v).strip()
                if not s:
                    continue
                if ',' in s and '.' in s:
                    s = s.replace('.', '').replace(',', '.') if s.rfind(',') > s.rfind('.') else s.replace(',', '')
                elif ',' in s:
                    s = s.replace(',', '.')
                try:
                    total += float(s)
                except ValueError:
                    pass
        return total

    return {
        "CCOS_DOC_NO":   join_unique("CCOS_DOC_NO"),
        "CCOS_REF_CODE": join_unique("CCOS_REF_CODE"),
        "CCOS_OR_BAL":   np.nan if multi_fac else join_unique("CCOS_OR_BAL"),
        "CCOS_BAL_DUE":  sum_col("CCOS_BAL_DUE", bal_due_rows),
    }


def _facode_label(osbal_rows: list, source: str) -> str:
    # Kembalikan FAC_CODE tunggal, atau 'facode lebih dari 1' jika multiple
    if not osbal_rows or not osbal_rows[0]:
        return ""
    col = OSBAL_FACODE_COL if source == "OSBAL" else (
          SLIPDB_FACODE_COL if source == "SLIPDB" else FACUL_FACODE_COL)
    codes = {str(r.get(col, "")).strip() for r in osbal_rows if str(r.get(col, "")).strip()}
    return "facode lebih dari 1" if len(codes) > 1 else next(iter(codes), "")


def _build_output_row(
    suspend_row: dict, source: str, scenario: str, osbal_rows: list,
    osbal_count: int = 0, resolved: bool = True,
    suspend_count: int = 1, all_fac_osbal_rows: list = None,
    is_beda_curr: bool = False, is_beda_per: bool = False,
) -> dict:
    """Bangun satu baris output (38 kolom standar mengikuti pola ACA).

    [TPI] Nama kolom insured mengikuti TPI: INSURED_ORI, INSURED_CLEAN_1/2.
    Kolom polis/sertif mengikuti pola ACA: POLIS, POLICY_CLEAN_1..5, SERTIF_CLEAN_1..3.
    Kolom slip: SLIP_NO, SLIP_NO_CLN.
    """
    has_match = bool(source and osbal_rows and osbal_rows[0])
    has_osbal = has_match and (source in ("OSBAL", "SLIPDB", "FACUL") or resolved)

    ccos = (
        _aggregate_ccos(osbal_rows, all_fac_osbal_rows=all_fac_osbal_rows)
        if has_osbal
        else {"CCOS_DOC_NO": "", "CCOS_REF_CODE": "", "CCOS_OR_BAL": np.nan, "CCOS_BAL_DUE": np.nan}
    )
    if is_beda_curr or is_beda_per or "Beda Currency" in str(scenario) or "Beda Periode" in str(scenario):
        ccos["CCOS_OR_BAL"] = np.nan
        ccos["CCOS_BAL_DUE"] = np.nan

    final_scenario = "Unmatching" if not source or scenario in ("Unmatching", "UNMATCHED") else scenario

    return {
        "CCOS_DOC_NO":       ccos["CCOS_DOC_NO"] if has_osbal else "",
        "CCOS_REF_CODE":     ccos["CCOS_REF_CODE"],
        "RECEIPT NO":        suspend_row.get("RECEIPT NO",        ""),
        "CREDIT NOTES":      suspend_row.get("CREDIT NOTES",      ""),
        "DETAIL RINCIAN NO": suspend_row.get("DETAIL RINCIAN NO", ""),
        "RECEIPT DATE":      suspend_row.get("RECEIPT DATE",      ""),
        "CEDANT NAME":       suspend_row.get("CEDANT NAME",       ""),
        "CEDANT SHRT NAME":  suspend_row.get("CEDANT SHRT NAME",  ""),

        # [TPI] Nama kolom insured dari file riil TPI
        "INSURED_ORI": suspend_row.get("INSURED_ORI",     ""),
        "INSURED_1":   suspend_row.get("INSURED_CLEAN_1", ""),
        "INSURED_2":   suspend_row.get("INSURED_CLEAN_2", ""),

        "CURR ORI":   suspend_row.get("CURR ORI",   ""),
        "AMOUNT ORI": suspend_row.get("AMOUNT ORI", ""),
        "CURR PAY":   suspend_row.get("CURR PAY",   ""),
        "AMOUNT PAY": suspend_row.get("AMOUNT PAY", ""),

        "CCOS_OR_BAL":  ccos["CCOS_OR_BAL"],
        "CCOS_BAL_DUE": ccos["CCOS_BAL_DUE"],

        # [TPI] Kolom polis/sertif mengikuti pola ACA (alternating Polis & Sertif)
        "POLIS":          suspend_row.get("POLIS_ORI",      ""),
        "POLICY_CLEAN_1": suspend_row.get("POLIS_CLEAN_1",  ""),
        "SERTIF_CLEAN_1": _format_sertif(suspend_row.get("CERTIFICATE_1",  "")),
        "POLICY_CLEAN_2": suspend_row.get("POLIS_CLEAN_2",  ""),
        "SERTIF_CLEAN_2": _format_sertif(suspend_row.get("CERTIFICATE_2",  "")),
        "POLICY_CLEAN_3": suspend_row.get("POLIS_CLEAN_3",  ""),
        "SERTIF_CLEAN_3": _format_sertif(suspend_row.get("CERTIFICATE_3",  "")),
        "POLICY_CLEAN_4": suspend_row.get("POLIS_CLEAN_4",  ""),
        "POLICY_CLEAN_5": suspend_row.get("POLIS_CLEAN_5",  ""),

        # [TPI] Kolom slip mengikuti pola ACA
        "SLIP_NO":     suspend_row.get("SLIP_NO_ORI",  ""),
        "SLIP_NO_CLN": suspend_row.get("SLIP_CLEAN_1", ""),

        "DESC 1": suspend_row.get("DESC 1", ""),
        "DESC 2": suspend_row.get("DESC 2", ""),
        "DESC 3": suspend_row.get("DESC 3", ""),
        "DESC 4": suspend_row.get("DESC 4", ""),

        "STATUS":   suspend_row.get("STATUS",   ""),
        "REC_TYPE": suspend_row.get("REC_TYPE", ""),

        "SKENARIO":            final_scenario,
        "_OSBAL_ROW_COUNT":    osbal_count,
        "_SUSPEND_ROW_COUNT":  suspend_count,
    }


def _merge_suspend_rows(rows: list) -> dict:
    # Merge beberapa baris suspend (fac code sama): sum AMOUNT ORI, join-unique sisanya
    if len(rows) == 1:
        return rows[0]
    merged: dict = {}
    for col in _SUSPEND_SUM_COLS:
        total = 0.0
        for r in rows:
            v = r.get(col, 0)
            if pd.isna(v) or v == "":
                continue
            if isinstance(v, (int, float)):
                total += float(v)
            else:
                s = str(v).strip()
                if not s:
                    continue
                if ',' in s and '.' in s:
                    s = s.replace('.', '').replace(',', '.') if s.rfind(',') > s.rfind('.') else s.replace(',', '')
                elif ',' in s:
                    s = s.replace(',', '.')
                try:
                    total += float(s)
                except ValueError:
                    pass
        merged[col] = total
    for col in _SUSPEND_JOIN_COLS:
        seen, vals = set(), []
        for r in rows:
            v = str(r.get(col, "")).strip()
            if v and v not in seen:
                seen.add(v)
                vals.append(v)
        merged[col] = ", ".join(vals)
    return merged


def _compute_derived_cols(df: pd.DataFrame) -> pd.DataFrame:
    """Hitung AMOUNT_ORI_MIN1, DIFERENCE, dan FLAG_PROD (10 kategori).

    [TPI] Matriks np.select identik dengan ACA/IBS.
    Kondisi is_gt1_ocmc selalu False (TPI tidak punya bordero).
    """
    amount_ori_base = df.get("AMOUNT ORI", pd.Series([0]*len(df), index=df.index))
    if "_MERGED_AMOUNT_ORI" in df.columns:
        amount_ori_raw = df["_MERGED_AMOUNT_ORI"].combine_first(amount_ori_base)
    else:
        amount_ori_raw = amount_ori_base

    amount_ori_clean = amount_ori_raw.replace(r'^\s*$', np.nan, regex=True)
    if amount_ori_clean.dtype == object:
        amount_ori_clean = amount_ori_clean.astype(str).str.strip().str.replace(r'\s+', '', regex=True)
        amount_ori_clean = amount_ori_clean.replace("nan", np.nan)
        has_comma = amount_ori_clean.str.contains(',', na=False)
        amount_ori_eu = (amount_ori_clean
                         .str.replace('.', '', regex=False)
                         .str.replace(',', '.', regex=False))
        amount_ori_clean = amount_ori_eu.where(has_comma, amount_ori_clean)
    amount_ori_numeric = pd.to_numeric(amount_ori_clean, errors="coerce")

    df["AMOUNT_ORI_MIN1"] = np.where(amount_ori_numeric.notna(), amount_ori_numeric * -1, np.nan)
    amount_ori  = amount_ori_numeric.fillna(0)
    amount_neg  = df["AMOUNT_ORI_MIN1"].fillna(0)

    balance_due     = pd.to_numeric(df.get("CCOS_BAL_DUE",      0), errors="coerce").fillna(0)
    osbal_count_col = pd.to_numeric(df.get("_OSBAL_ROW_COUNT",  0), errors="coerce").fillna(0)
    sus_count_col   = pd.to_numeric(df.get("_SUSPEND_ROW_COUNT",1), errors="coerce").fillna(1)
    ccos_ref        = df.get("CCOS_REF_CODE", pd.Series("", index=df.index)).fillna("").astype(str)
    skenario_col    = df.get("SKENARIO",      pd.Series("", index=df.index)).fillna("")

    df["DIFERENCE"] = amount_neg - balance_due
    accumulated = (osbal_count_col > 1) | (sus_count_col > 1)
    is_equal    = amount_neg.round(2) == balance_due.round(2)

    if "_IS_BEDA_PER" in df.columns:
        is_beda_periode = df["_IS_BEDA_PER"].fillna(False).astype(bool)
    else:
        is_beda_periode = skenario_col.str.contains("Beda Periode",  na=False, regex=False)

    if "_IS_BEDA_CURR" in df.columns:
        is_beda_currency = df["_IS_BEDA_CURR"].fillna(False).astype(bool)
    else:
        is_beda_currency = skenario_col.str.contains("Beda Currency", na=False, regex=False)

    # [TPI] is_gt1_ocmc: selalu False karena TPI tidak punya bordero
    # is_gt1_fac: deteksi dari skenario "(Matching >1 fac code)" sesuai pola ACA
    is_gt1_ocmc          = pd.Series(False, index=df.index)
    is_gt1_fac           = skenario_col.str.contains("Matching >1 fac code", na=False, regex=False) | (ccos_ref.str.contains(",", na=False) & ~is_gt1_ocmc)
    is_new_entry_total   = (amount_ori != 0) & (balance_due == 0)
    is_adj_total         = is_equal
    is_adj_sebagian      = ~is_equal & (amount_neg < balance_due)
    is_new_entry_sebagian= ~is_equal & (amount_neg > balance_due)

    df["FLAG_PROD"] = np.select(
        [is_beda_currency, is_beda_periode, is_gt1_ocmc, is_gt1_fac,
         is_new_entry_total,
         is_adj_total & ~accumulated, is_adj_total & accumulated,
         is_adj_sebagian & ~accumulated, is_adj_sebagian & accumulated,
         is_new_entry_sebagian],
        ["Beda Currency", "Beda Periode",
         "Matching >1 Fac code OC MC", "Matching >1 fac code",
         "New Entry total",
         "Adjustment total tanpa akumulasi", "Adjustment total dengan akumulasi",
         "Adjustment sebagian tanpa akumulasi", "Adjustment sebagian dengan akumulasi",
         "New Entry sebagian"],
        default="Unmatching",
    )

    for col in ["_OSBAL_ROW_COUNT", "_SUSPEND_ROW_COUNT", "_MERGED_AMOUNT_ORI", "_IS_BEDA_CURR", "_IS_BEDA_PER"]:
        if col in df.columns:
            df.drop(columns=[col], inplace=True)
    return df


# =============================================================================
# PIPELINE
# =============================================================================

def _load_excel(path: str, label: str, optional: bool = False) -> tuple:
    # Load Excel ke (list[dict], list[str]) dengan pickle cache otomatis
    if not os.path.exists(path):
        print(f"  {label}: {path} ({'tidak ada, dilewati' if optional else 'TIDAK DITEMUKAN'})", flush=True)
        return [], []

    cache = path + ".cache.pkl"
    if os.path.exists(cache) and os.path.getmtime(cache) >= os.path.getmtime(path):
        t = time.perf_counter()
        print(f"  {label}: loading dari cache ...", flush=True)
        try:
            data, header = pd.read_pickle(cache)
            print(f"  -> {len(data):,} rows, {len(header)} cols ({time.perf_counter()-t:.2f}s)", flush=True)
            return data, header
        except Exception:
            print(f"  {label}: cache rusak, baca ulang dari Excel ...", flush=True)

    t = time.perf_counter()
    print(f"  {label}: {path} ...", flush=True)
    try:
        df = pd.read_excel(path)
    except Exception as e:
        print(f"  [RETRY] engine default gagal ({e}), coba openpyxl ...", flush=True)
        df = pd.read_excel(path, engine="openpyxl")

    header = list(df.columns)
    data   = df.to_dict("records")
    print(f"  -> {len(data):,} rows, {len(header)} cols ({time.perf_counter()-t:.1f}s)", flush=True)

    try:
        pd.to_pickle((data, header), cache)
        print(f"  -> cache disimpan: {cache}", flush=True)
    except Exception:
        pass
    return data, header


def run() -> None:
    """Pipeline utama TPI: matching 5 fase (tanpa Bordero, tanpa FASE 1 DESC_3_CLEAN).

    [TPI] Urutan fase: Polis -> Slip -> Insured ->
                       Currency & Periode -> Narrowing Non-Destruktif
    Ref: teknis_matching_aca.md, matching_IBS.md (adaptasi TPI tanpa FASE 1 DESC_3_CLEAN)
    """
    t_start = time.perf_counter()
    print("\n" + "=" * 60, flush=True)
    print("  PRODUCTION SCRIPT -- SUSPEND MATCHING  (TPI) -- V1", flush=True)
    print("  Cedant: PT Asuransi Tugu Pratama Indonesia", flush=True)
    print("=" * 60, flush=True)

    # [1] Load data
    print("\n[1/6] Loading data ...", flush=True)
    suspend_rows, suspend_cols = _load_excel(SUSPEND_FILE, "SUSPEND")
    osbal_rows,   osbal_cols   = _load_excel(OSBAL_FILE,   "OSBAL")
    facul_rows,   facul_cols   = _load_excel(FACUL_FILE,   "FACUL")
    slipdb_rows,  slipdb_cols  = _load_excel(SLIPDB_FILE,  "SLIPDB", optional=True)
    print("  Bordero: TIDAK DIGUNAKAN untuk TPI (fase dihapus dari pipeline)", flush=True)
    print("  FASE 1 DESC_3_CLEAN: TIDAK DITERAPKAN untuk TPI", flush=True)

    # Diagnostik kolom kunci
    print("\n  --- Cek kolom SUSPEND ---", flush=True)
    print(f"  Semua kolom SUSPEND: {suspend_cols}", flush=True)
    for col in ["INSURED_ORI", "INSURED_CLEAN_1", "INSURED_CLEAN_2",
                "POLIS_ORI", "POLIS_CLEAN_1", "CERTIFICATE_1",
                "SLIP_NO_ORI", "SLIP_CLEAN_1",
                SUSPEND_CURR_COL, SUSPEND_DATE_COL]:
        status = "OK" if col in suspend_cols else "[MISSING]"
        print(f"    {col:<22}: {status}", flush=True)

    print("\n  --- Cek kolom OSBAL ---", flush=True)
    print(f"  Semua kolom OSBAL: {osbal_cols}", flush=True)
    for col in [OSBAL_FACODE_COL, OSBAL_CURR_COL, OSBAL_DATE_COL,
                "CCOS_DOC_NO", "CCOS_OR_BAL", "CCOS_BAL_DUE"]:
        status = "OK" if col in osbal_cols else "[MISSING]"
        print(f"    {col:<22}: {status}", flush=True)

    # Pre-parse tanggal OSBAL FAC_COM_DATE
    print("  Pre-parsing OSBAL FAC_COM_DATE ...", flush=True)
    for r in osbal_rows:
        raw = r.get(OSBAL_DATE_COL, "")
        r["_com_date_parsed"] = pd.to_datetime(raw, errors="coerce") \
            if raw and not (isinstance(raw, float) and pd.isna(raw)) else None

    # Pre-parse tanggal Suspend RECEIPT DATE
    print("  Pre-parsing Suspend RECEIPT DATE ...", flush=True)
    for r in suspend_rows:
        raw = r.get(SUSPEND_DATE_COL, "")
        r["_sus_date_parsed"] = pd.to_datetime(raw, errors="coerce") \
            if raw and not (isinstance(raw, float) and pd.isna(raw)) else None

    # [2] Deteksi kolom bersih
    print("\n[2/6] Detecting clean columns ...", flush=True)

    # [TPI] Kolom polis/slip/insured suspend (dari output cleaning-suspend.py)
    polis_sus   = _get_clean_cols(suspend_cols, "POLIS_CLEAN")   or _get_clean_cols(suspend_cols, "clean polis")   or []
    slip_sus    = _get_clean_cols(suspend_cols, "SLIP_CLEAN")    or _get_clean_cols(suspend_cols, "clean slip")    or []
    insured_sus = _get_clean_cols(suspend_cols, "INSURED_CLEAN") or _get_clean_cols(suspend_cols, "clean insured") or []

    print(f"  Kolom polis suspend   : {polis_sus}", flush=True)
    print(f"  Kolom slip suspend    : {slip_sus}", flush=True)
    print(f"  Kolom insured suspend : {insured_sus}", flush=True)

    # Pre-parse LINESLIP flags
    print("  Pre-parsing LINESLIP flags ...", flush=True)
    lineslip_count = 0
    for r in suspend_rows:
        r["_is_lineslip"] = _is_lineslip_row(r, insured_sus)
        if r["_is_lineslip"]:
            lineslip_count += 1
            parsed = r.get("_sus_date_parsed")
            try:
                r["_sus_date_lineslip"] = (parsed - pd.DateOffset(months=1)) if parsed is not None else None
            except Exception:
                r["_sus_date_lineslip"] = None
        else:
            r["_sus_date_lineslip"] = None
    print(f"  -> {lineslip_count:,} baris LINESLIP terdeteksi", flush=True)

    # [TPI] Kolom referensi OSBAL (dari output cleaning-osbal.py)
    polis_osbal   = _get_clean_cols(osbal_cols, "POLIS_CLEAN")   or _get_clean_cols(osbal_cols, "clean polis")   or []
    slip_osbal    = _get_clean_cols(osbal_cols, "SLIP_CLEAN")    or _get_clean_cols(osbal_cols, "clean slip")    or []
    insured_osbal = _get_clean_cols(osbal_cols, "INSURED_CLEAN") or _get_clean_cols(osbal_cols, "clean insured") or []
    sertif_osbal  = _get_clean_cols(osbal_cols, "CERTIFICATE")   or _get_clean_cols(osbal_cols, "clean sertif") or []

    polis_facul   = _get_clean_cols(facul_cols, "POLIS_CLEAN")   or _get_clean_cols(facul_cols, "clean polis")   or []
    slip_facul    = _get_clean_cols(facul_cols, "SLIP_CLEAN")    or _get_clean_cols(facul_cols, "clean slip")    or []
    insured_facul = _get_clean_cols(facul_cols, "INSURED_CLEAN") or _get_clean_cols(facul_cols, "clean insured") or []

    polis_slipdb   = _get_clean_cols(slipdb_cols, "POLIS_CLEAN")   or _get_clean_cols(slipdb_cols, "clean polis")   or []
    slip_slipdb    = _get_clean_cols(slipdb_cols, "SLIP_CLEAN")    or _get_clean_cols(slipdb_cols, "clean slip")    or []
    insured_slipdb = _get_clean_cols(slipdb_cols, "INSURED_CLEAN") or _get_clean_cols(slipdb_cols, "clean insured") or []

    if not polis_slipdb:
        polis_slipdb = [c for c in ["CLSDT_POLICY_NO", "FAC_POLICY_NO", "POLIS", "POLICY_NO", "POLICY NO"] if c in slipdb_cols]
    if not slip_slipdb:
        slip_slipdb  = [c for c in ["CLSDT_SLIP_NO", "FAC_SLIP", "SLIP", "SLIP_NO", "SLIP NO"] if c in slipdb_cols]
    if not insured_slipdb:
        insured_slipdb = [c for c in ["FAC_INSURED", "INSURED", "CEDANT_NAME", "CEDANT NAME"] if c in slipdb_cols]

    global SLIPDB_FACODE_COL
    if   "FAC_CODE"      in slipdb_cols: SLIPDB_FACODE_COL = "FAC_CODE"
    elif "FAC CODE"      in slipdb_cols: SLIPDB_FACODE_COL = "FAC CODE"
    elif "CCOS_REF_CODE" in slipdb_cols: SLIPDB_FACODE_COL = "CCOS_REF_CODE"

    def _prepend(col, cols_list, header):
        return ([col] + cols_list) if col in header and col not in cols_list else cols_list

    # [TPI] Prepend kolom ORI ke daftar kolom bersih untuk referensi
    polis_osbal  = _prepend("FAC_POLICY_NO_ORI", polis_osbal,  osbal_cols)
    slip_osbal   = _prepend("FAC_SLIP_ORI",       slip_osbal,   osbal_cols)
    polis_facul  = _prepend("FAC_POLICY_NO_ORI", polis_facul,  facul_cols)
    slip_facul   = _prepend("FAC_SLIP_ORI",       slip_facul,   facul_cols)
    polis_slipdb = _prepend("CLSDT_POLICY_NO",   polis_slipdb, slipdb_cols)
    slip_slipdb  = _prepend("CLSDT_SLIP_NO",      slip_slipdb,  slipdb_cols)

    print(f"  Kolom polis OSBAL     : {polis_osbal}", flush=True)
    print(f"  Kolom slip OSBAL      : {slip_osbal}", flush=True)
    print(f"  Kolom insured OSBAL   : {insured_osbal}", flush=True)

    # [3] Bangun lookup index
    print("\n[3/6] Building lookup indexes ...", flush=True)
    lookup_osbal  = _build_lookup(osbal_rows,  polis_osbal,  slip_osbal,  insured_osbal,  OSBAL_FACODE_COL,  label="OSBAL")
    lookup_facul  = _build_lookup(facul_rows,  polis_facul,  slip_facul,  insured_facul,  FACUL_FACODE_COL,  label="FACUL")
    lookup_slipdb = _build_lookup(slipdb_rows, polis_slipdb, slip_slipdb, insured_slipdb, SLIPDB_FACODE_COL, label="SLIPDB")

    facode_osbal_idx = lookup_osbal[9]

    # [4] Matching per baris
    total = len(suspend_rows)
    print(f"\n[4/6] Matching {total:,} suspend rows ...", flush=True)
    t4 = time.perf_counter()
    raw_results = []
    stat_like_substr  = 0
    stat_like_subset  = 0
    stat_like_overlap = 0

    for n, sus in enumerate(suspend_rows, 1):
        if n % 1000 == 0:
            elapsed_so_far = time.perf_counter() - t4
            rate = n / elapsed_so_far if elapsed_so_far > 0 else 0
            eta  = (total - n) / rate if rate > 0 else 0
            print(f"    Memproses baris {n:,} / {total:,}  ({rate:.0f} rows/s, ETA {eta:.0f}s)", flush=True)

        if str(sus.get("STATUS", "")).strip().upper() == "MATCHING":
            continue

        eff_polis, eff_slip, like_polis, like_slip = _get_effective_sus_cols(sus, polis_sus, slip_sus)
        _is_ls = sus.get("_is_lineslip", False)
        effective_date = sus.get("_sus_date_lineslip") if _is_ls else sus.get("_sus_date_parsed")

        matched_osbal_idx = set()
        source = None
        scenario = ""
        resolved = False
        _like_log: list = []

        # =================================================================
        # [TPI] FASE 1 DESC_3_CLEAN: TIDAK DITERAPKAN
        # TPI tidak memiliki rules tambahan untuk pengecekan ke DESC 3.
        # Pipeline langsung ke FASE 1 ACA standar (Polis).
        # =================================================================

        # =================================================================
        # FASE 1 (TPI = ACA FASE 1): Pencarian berdasarkan POLIS
        # Cross-check: OSBAL -> SLIPDB -> FACUL
        # Jika exact miss, lanjut ke LIKE match polis.
        # Ref: teknis_matching_aca.md §2 FASE 1
        # =================================================================
        found_by_polis = False
        polis_vals = _collect_clean_values(sus, eff_polis)
        if polis_vals:
            # OSBAL exact
            osb_hits = _exact_match(polis_vals, lookup_osbal[1])
            if osb_hits:
                matched_osbal_idx.update(osb_hits)
                found_by_polis = True

            # SLIPDB exact
            if slipdb_rows:
                sld_hits = _exact_match(polis_vals, lookup_slipdb[1])
                if sld_hits:
                    res_idx, _ = _resolve_facode([slipdb_rows[idx] for idx in sld_hits], SLIPDB_FACODE_COL, facode_osbal_idx, osbal_rows)
                    if res_idx:
                        matched_osbal_idx.update(res_idx)
                        found_by_polis = True

            # FACUL exact
            if facul_rows:
                fac_hits = _exact_match(polis_vals, lookup_facul[1])
                if fac_hits:
                    res_idx, _ = _resolve_facode([facul_rows[idx] for idx in fac_hits], FACUL_FACODE_COL, facode_osbal_idx, osbal_rows)
                    if res_idx:
                        matched_osbal_idx.update(res_idx)
                        found_by_polis = True

            # LIKE match polis (jika exact miss)
            if not found_by_polis and lookup_osbal[7] and lookup_osbal[7][0] and like_polis:
                like_vals = _collect_clean_values(sus, like_polis)
                if like_vals:
                    like_hits = _like_match(
                        like_vals, lookup_osbal[7][0], osbal_rows, polis_osbal,
                        stopwords=frozenset(),
                        overlap_threshold=_POLIS_LIKE_THRESHOLD,
                        match_log=_like_log,
                    )
                    if like_hits:
                        matched_osbal_idx.update(like_hits)
                        found_by_polis = True

        if found_by_polis:
            source = "OSBAL"
            scenario = "Polis only"

        # =================================================================
        # FASE 2 (TPI = ACA FASE 2): Pencarian berdasarkan SLIP (fallback)
        # Hanya dijalankan jika FASE 1 (Polis) tidak menghasilkan match.
        # Cross-check: OSBAL -> SLIPDB -> FACUL
        # Ref: teknis_matching_aca.md §2 FASE 2
        # =================================================================
        fallback_slip_hits = set()
        fallback_slip_source = None
        fallback_slip_resolved = False

        if not found_by_polis:
            slip_vals = _collect_clean_values(sus, eff_slip)
            if slip_vals:
                # 1. OSBAL exact
                osb_hits = _exact_match(slip_vals, lookup_osbal[0])
                if osb_hits:
                    sub_facs = {str(osbal_rows[idx].get(OSBAL_FACODE_COL, "")).strip() for idx in osb_hits}
                    sus_ins_curr = _collect_all_values(sus, insured_sus, "insured_ori")
                    sub_ins_confirmed = [idx for idx in osb_hits if _ref_matches_insured(osbal_rows[idx], insured_osbal, sus_ins_curr)]
                    if sub_ins_confirmed:
                        matched_osbal_idx.update(sub_ins_confirmed)
                        source = "OSBAL"
                        scenario = "Slip + Insured"
                    elif len(sub_facs) == 1:
                        matched_osbal_idx.update(osb_hits)
                        source = "OSBAL"
                        scenario = "Slip only"
                    else:
                        # [TPI] Simpan sebagai fallback jika FASE 3 (Insured) miss
                        fallback_slip_hits = set(osb_hits)
                        fallback_slip_source = "OSBAL"
                        fallback_slip_resolved = False

                # 2. SLIPDB exact
                if not matched_osbal_idx and not fallback_slip_hits and slipdb_rows:
                    sld_hits = _exact_match(slip_vals, lookup_slipdb[0])
                    if sld_hits:
                        sus_ins_curr = _collect_all_values(sus, insured_sus, "insured_ori")
                        sld_ins_confirmed = [idx for idx in sld_hits if _ref_matches_insured(slipdb_rows[idx], insured_slipdb, sus_ins_curr)]
                        if sld_ins_confirmed:
                            res_idx, ok = _resolve_facode([slipdb_rows[idx] for idx in sld_ins_confirmed], SLIPDB_FACODE_COL, facode_osbal_idx, osbal_rows)
                            if ok:
                                matched_osbal_idx.update(res_idx)
                                source = "SLIPDB"
                                scenario = "Slip + Insured"
                 
                                resolved = True
                        else:
                            res_idx, ok = _resolve_facode([slipdb_rows[idx] for idx in sld_hits], SLIPDB_FACODE_COL, facode_osbal_idx, osbal_rows)
                            if ok:
                                sub_facs = {str(osbal_rows[idx].get(OSBAL_FACODE_COL, "")).strip() for idx in res_idx}
                                if len(sub_facs) == 1:
                                    matched_osbal_idx.update(res_idx)
                                    source = "SLIPDB"
                                    scenario = "Slip only"
                                    resolved = True
                                else:
                                    fallback_slip_hits = set(res_idx)
                                    fallback_slip_source = "SLIPDB"
                                    fallback_slip_resolved = True

                # 3. FACUL exact
                if not matched_osbal_idx and not fallback_slip_hits and facul_rows:
                    fac_hits = _exact_match(slip_vals, lookup_facul[0])
                    if fac_hits:
                        sus_ins_curr = _collect_all_values(sus, insured_sus, "insured_ori")
                        fac_ins_confirmed = [idx for idx in fac_hits if _ref_matches_insured(facul_rows[idx], insured_facul, sus_ins_curr)]
                        if fac_ins_confirmed:
                            res_idx, ok = _resolve_facode([facul_rows[idx] for idx in fac_ins_confirmed], FACUL_FACODE_COL, facode_osbal_idx, osbal_rows)
                            if ok:
                                matched_osbal_idx.update(res_idx)
                                source = "FACUL"
                                scenario = "Slip + Insured"
                                resolved = True
                        else:
                            res_idx, ok = _resolve_facode([facul_rows[idx] for idx in fac_hits], FACUL_FACODE_COL, facode_osbal_idx, osbal_rows)
                            if ok:
                                sub_facs = {str(osbal_rows[idx].get(OSBAL_FACODE_COL, "")).strip() for idx in res_idx}
                                if len(sub_facs) == 1:
                                    matched_osbal_idx.update(res_idx)
                                    source = "FACUL"
                                    scenario = "Slip only"
                                    resolved = True
                                else:
                                    fallback_slip_hits = set(res_idx)
                                    fallback_slip_source = "FACUL"
                                    fallback_slip_resolved = True

                # 4. LIKE match slip (hanya jika exact miss, tetap verifikasi insured jika bisa)
                if not matched_osbal_idx and not fallback_slip_hits and not osb_hits and lookup_osbal[6] and lookup_osbal[6][0] and like_slip:
                    like_slip_vals = _collect_clean_values(sus, like_slip)
                    if like_slip_vals:
                        like_hits = _like_match(
                            like_slip_vals, lookup_osbal[6][0], osbal_rows, slip_osbal,
                            stopwords=frozenset(),
                            overlap_threshold=_POLIS_LIKE_THRESHOLD,
                            match_log=_like_log,
                        )
                        if like_hits:
                            sus_ins_curr = _collect_all_values(sus, insured_sus, "insured_ori")
                            sub_ins_confirmed = [idx for idx in like_hits if _ref_matches_insured(osbal_rows[idx], insured_osbal, sus_ins_curr)]
                            if sub_ins_confirmed:
                                matched_osbal_idx.update(sub_ins_confirmed)
                                source = "OSBAL"
                                scenario = "Slip + Insured"
                            else:
                                sub_facs = {str(osbal_rows[idx].get(OSBAL_FACODE_COL, "")).strip() for idx in like_hits}
                                if len(sub_facs) == 1:
                                    matched_osbal_idx.update(like_hits)
                                    source = "OSBAL"
                                    scenario = "Slip only"
                                else:
                                    fallback_slip_hits = set(like_hits)
                                    fallback_slip_source = "OSBAL"
                                    fallback_slip_resolved = False

        # =================================================================
        # FASE 3 (TPI = ACA FASE 3): Pencarian berdasarkan INSURED (fallback)
        # Hanya dijalankan jika FASE 1 & 2 tidak menghasilkan match.
        # Ref: teknis_matching_aca.md §2 FASE 3
        # =================================================================
        if not matched_osbal_idx:
            ins_vals = _collect_clean_values(sus, insured_sus)
            if ins_vals:
                slip_check_vals = _collect_clean_values(sus, eff_slip)

                # 1. Coba OSBAL Exact & Like
                osb_hits = _exact_match(ins_vals, lookup_osbal[2])
                if not osb_hits and lookup_osbal[8] and lookup_osbal[8][0]:
                    osb_hits = _like_match(
                        ins_vals, lookup_osbal[8][0], osbal_rows, insured_osbal,
                        stopwords=_INSURED_STOPWORDS,
                        overlap_threshold=_INSURED_LIKE_THRESHOLD,
                        match_log=_like_log,
                    )
                if osb_hits:
                    matched_osbal_idx.update(osb_hits)
                    source = "OSBAL"
                    osb_slip_conf = any(
                        any(_normalize(osbal_rows[i].get(c, "")) in slip_check_vals
                            for c in slip_osbal if _normalize(osbal_rows[i].get(c, "")))
                        for i in osb_hits
                    ) if slip_check_vals else False
                    scenario = "Slip + Insured" if osb_slip_conf else "Insured only"

                # 2. Coba SLIPDB Exact & Like
                elif slipdb_rows and not matched_osbal_idx:
                    sld_hits = _exact_match(ins_vals, lookup_slipdb[2])
                    if not sld_hits and lookup_slipdb[8] and lookup_slipdb[8][0]:
                        sld_hits = _like_match(
                            ins_vals, lookup_slipdb[8][0], slipdb_rows, insured_slipdb,
                            stopwords=_INSURED_STOPWORDS,
                            overlap_threshold=_INSURED_LIKE_THRESHOLD,
                            match_log=_like_log,
                        )
                    if sld_hits:
                        res_idx, _ = _resolve_facode([slipdb_rows[idx] for idx in sld_hits], SLIPDB_FACODE_COL, facode_osbal_idx, osbal_rows)
                        if res_idx:
                            matched_osbal_idx.update(res_idx)
                            source = "SLIPDB"
                            sld_slip_conf = any(
                                any(_normalize(slipdb_rows[idx].get(c, "")) in slip_check_vals
                                    for c in slip_slipdb if _normalize(slipdb_rows[idx].get(c, "")))
                                for idx in sld_hits
                            ) if slip_check_vals else False
                            scenario = "Slip + Insured" if sld_slip_conf else "Insured only"
                            resolved = True

                # 3. Coba FACUL Exact & Like
                elif facul_rows and not matched_osbal_idx:
                    fac_hits = _exact_match(ins_vals, lookup_facul[2])
                    if not fac_hits and lookup_facul[8] and lookup_facul[8][0]:
                        fac_hits = _like_match(
                            ins_vals, lookup_facul[8][0], facul_rows, insured_facul,
                            stopwords=_INSURED_STOPWORDS,
                            overlap_threshold=_INSURED_LIKE_THRESHOLD,
                            match_log=_like_log,
                        )
                    if fac_hits:
                        res_idx, _ = _resolve_facode([facul_rows[idx] for idx in fac_hits], FACUL_FACODE_COL, facode_osbal_idx, osbal_rows)
                        if res_idx:
                            matched_osbal_idx.update(res_idx)
                            source = "FACUL"
                            fac_slip_conf = any(
                                any(_normalize(facul_rows[idx].get(c, "")) in slip_check_vals
                                    for c in slip_facul if _normalize(facul_rows[idx].get(c, "")))
                                for idx in fac_hits
                            ) if slip_check_vals else False
                            scenario = "Slip + Insured" if fac_slip_conf else "Insured only"
                            resolved = True

        # Fallback Slip: jika FASE 3 (Insured) tidak menghasilkan match,
        # gunakan kandidat slip FASE 2 agar tidak menjadi Unmatching (termasuk slip payung seperti 1100000030)
        if not matched_osbal_idx and fallback_slip_hits:
            matched_osbal_idx.update(fallback_slip_hits)
            source = fallback_slip_source
            scenario = "Slip only"
            resolved = fallback_slip_resolved

        # =================================================================
        # FASE 4: Narrowing Wajib (Currency & Periode) + Strict Rematch
        # Identik dengan ACA
        # Ref: teknis_matching_aca.md §2 FASE 4
        # =================================================================
        is_beda_curr = False
        is_beda_per  = False
        if matched_osbal_idx:
            matched_list = list(matched_osbal_idx)

            sus_curr = _normalize(sus.get(SUSPEND_CURR_COL, ""))
            if sus_curr:
                filtered_curr = [i for i in matched_list if _normalize(osbal_rows[i].get(OSBAL_CURR_COL, "")) == sus_curr]
                if not filtered_curr:
                    is_beda_curr = True
                elif len(filtered_curr) < len(matched_list):
                    matched_list = filtered_curr

            if not is_beda_curr:
                sus_date = effective_date
                if pd.notna(sus_date):
                    filtered_per = []
                    for i in matched_list:
                        com_date = osbal_rows[i].get("_com_date_parsed")
                        if pd.isna(com_date):
                            filtered_per.append(i)
                        else:
                            if _is_ls:
                                if com_date.year == sus_date.year and com_date.month == sus_date.month:
                                    filtered_per.append(i)
                            else:
                                if sus_date >= com_date:
                                    filtered_per.append(i)
                    if not filtered_per:
                        is_beda_per = True
                    elif len(filtered_per) < len(matched_list):
                        matched_list = filtered_per

            if is_beda_curr or is_beda_per:
                rm = _rematch_strict(
                    sus, osbal_rows, lookup_osbal,
                    eff_polis, eff_slip, insured_sus,
                    polis_osbal, slip_osbal, insured_osbal,
                    facode_osbal_idx,
                    slipdb_rows, lookup_slipdb,
                    facul_rows, lookup_facul,
                    like_polis,
                    effective_sus_date=effective_date, is_lineslip=_is_ls
                )
                if rm:
                    src, scen, hits, _ = rm
                    source = src
                    scenario = scen
                    matched_list = hits
                    is_beda_curr = False
                    is_beda_per  = False

            matched_osbal_idx = set(matched_list)

        # =================================================================
        # FASE 5: Narrowing Utama Bertingkat Non-Destruktif
        # Identik dengan ACA
        # Ref: teknis_matching_aca.md §2 FASE 5
        # =================================================================
        if len(matched_osbal_idx) > 1:
            matched_list = list(matched_osbal_idx)

            # Step 1a: Konfirmasi Slip (jika match via Polis)
            if "Polis" in scenario:
                slip_confirm_vals = _collect_clean_values(sus, eff_slip)
                if slip_confirm_vals:
                    confirmed = [idx for idx in matched_list if _ref_has_value(osbal_rows[idx], slip_osbal, "slip_ori", slip_confirm_vals)]
                    if confirmed:
                        matched_list = confirmed
                        scenario = "Polis + Slip"

            # Step 1b: Slip Range Numerik (hanya jika belum terkonfirmasi Slip)
            if len(matched_list) > 1 and "Slip" not in scenario:
                slip_range_vals = _collect_clean_values(sus, eff_slip if eff_slip else slip_sus)
                if slip_range_vals:
                    confirmed_slip = [
                        idx for idx in matched_list
                        if any(_slip_in_range(sv, osbal_rows[idx]) for sv in slip_range_vals)
                    ]
                    if confirmed_slip:
                        matched_list = confirmed_slip
                        scenario = "Polis + Slip" if "Polis" in scenario else scenario

            # Step 2: Konfirmasi Polis (jika match via Slip)
            if len(matched_list) > 1 and "Slip" in scenario and "Polis" not in scenario:
                polis_confirm_vals = _collect_all_values(sus, eff_polis, "polis_ori")
                if polis_confirm_vals:
                    confirmed = [idx for idx in matched_list if _ref_has_value(osbal_rows[idx], polis_osbal, "polis_ori", polis_confirm_vals)]
                    if confirmed:
                        matched_list = confirmed
                        scenario = "Slip + Polis"



            # Step 3: Konfirmasi Insured
            # [TPI] Selain INSURED_CLEAN_N, cek juga FAC_INSURED_ORI (termasuk nama subsidiary setelah 'QQ'
            # seperti 'LION AIR FLEET GROUP QQ LION MENTARI') -> mempersempit kandidat ke 1 fac code spesifik.
            ins_vals = _collect_all_values(sus, insured_sus, "insured_ori")
            if ins_vals and len(matched_list) > 1:
                confirmed = [idx for idx in matched_list if _ref_matches_insured(osbal_rows[idx], insured_osbal, ins_vals)]
                if confirmed:
                    matched_list = confirmed
                    if "Insured" not in scenario:
                        insured_vals_before: set = set()
                        for _idx in list(matched_osbal_idx):
                            for _c in insured_osbal:
                                _v = _normalize(osbal_rows[_idx].get(_c, ""))
                                if _v:
                                    insured_vals_before.add(_v)
                        if len(insured_vals_before) > 1:
                            base = scenario.replace(" only", "") if not scenario.endswith("only") else scenario
                            scenario = base + " + Insured"

            matched_osbal_idx = set(matched_list)

        # =================================================================
        # [TPI] FASE Bordero Open Cover Override -- DIHAPUS UNTUK TPI
        # TPI tidak memiliki data Open Cover Marine Cargo.
        # =================================================================

        # Post-matching normalisasi label
        if is_beda_curr or is_beda_per:
            osbal_ref   = [osbal_rows[i] for i in list(matched_osbal_idx)] if matched_osbal_idx else [{}]
            osbal_count = 0
            resolved    = False
        elif not matched_osbal_idx:
            source      = None
            scenario    = "Unmatching"
            osbal_ref   = [{}]
            osbal_count = 0
            resolved    = False
        else:
            osbal_ref   = [osbal_rows[i] for i in matched_osbal_idx]
            osbal_count = len(osbal_ref)

        if scenario == "INSURED_LIKE":
            scenario = "Insured only"
        elif scenario and " only + " in scenario:
            scenario = scenario.replace(" only + ", " + ")

        fac_codes = {_normalize(r.get(OSBAL_FACODE_COL, "")) for r in osbal_ref if r}
        fac_codes.discard("")

        raw_results.append({
            "suspend":        sus,
            "source":         source,
            "base_scenario":  scenario,
            "is_bordero":     False,   # [TPI] tidak ada bordero
            "is_beda_curr":   is_beda_curr,
            "is_beda_per":    is_beda_per,
            "osbal_idx":      list(matched_osbal_idx),
            "osbal_count":    osbal_count,
            "resolved":       resolved,
            "fac_codes":      fac_codes,
            "_like_log":      [m for _, m in _like_log],
        })

    print(f"  Matching selesai dalam {time.perf_counter()-t4:.1f}s", flush=True)
    for _row in raw_results:
        for _method in _row.get("_like_log", []):
            if _method == "SUBSTR":    stat_like_substr  += 1
            elif _method == "SUBSET":  stat_like_subset  += 1
            elif _method == "OVERLAP": stat_like_overlap += 1

    # [4c] Fac code claim & osbal accumulation (identik dengan ACA)
    print("\n[4c/6] Fac code claim & osbal accumulation (row asli dipertahankan) ...", flush=True)

    claimed: set = set()

    def _is_invalid_match(r):
        return (
            r.get("is_beda_curr", False)
            or r.get("is_beda_per", False)
            or r.get("source") is None
            or r.get("base_scenario", "") in ("Unmatching", "UNMATCHED")
        )

    for r in raw_results:
        if len(r["fac_codes"]) == 1 and not _is_invalid_match(r):
            claimed |= r["fac_codes"]

    for r in raw_results:
        if len(r["fac_codes"]) > 1:
            remaining = r["fac_codes"] - claimed
            if len(remaining) == 1 and not _is_invalid_match(r):
                r["_originally_multi_fac"] = True
                claimed |= remaining

    groups: dict = {}
    final:  list = []

    for r in raw_results:
        r.setdefault("suspend_count", 1)
        if len(r["fac_codes"]) == 1 and not _is_invalid_match(r):
            groups.setdefault(next(iter(r["fac_codes"])), []).append(r)
        else:
            final.append(r)

    for fac, group in groups.items():
        if len(group) == 1:
            group[0]["suspend_count"] = 1
            final.append(group[0])
            continue
        merged_osbal_idx = set()
        for g in group:
            merged_osbal_idx.update(g["osbal_idx"])
        merged_sus = _merge_suspend_rows([g["suspend"] for g in group])
        for g in group:
            g["osbal_idx"]          = list(merged_osbal_idx)
            g["osbal_count"]        = len(merged_osbal_idx)
            g["suspend_count"]      = len(group)
            g["_merged_amount_ori"] = merged_sus.get("AMOUNT ORI", 0)
            final.append(g)

    print(f"  {len(raw_results):,} suspend rows -> {len(final):,} output rows (row asli dipertahankan)", flush=True)
    total_like = stat_like_substr + stat_like_subset + stat_like_overlap
    print(f"  [DEBUG] _like_match total hits     : {total_like}", flush=True)
    print(f"    SUBSTR  (substring lama)          : {stat_like_substr}", flush=True)
    print(f"    SUBSET  (token subset)            : {stat_like_subset}", flush=True)
    print(f"    OVERLAP (overlap coeff >={_INSURED_LIKE_THRESHOLD:.0%}) : {stat_like_overlap}", flush=True)

    # [5] Build output DataFrame
    print(f"\n[5/6] Building output ({len(final):,} rows) ...", flush=True)
    output_rows = []
    for r in final:
        ref_osbal = [osbal_rows[i] for i in r["osbal_idx"]] if r["source"] else [{}]
        all_fac_osbal_rows = None
        osbal_cnt = r["osbal_count"]
        if r["source"] and len(r["fac_codes"]) == 1:
            fac = next(iter(r["fac_codes"]))
            all_fac_idx = facode_osbal_idx.get(fac, [])
            if len(all_fac_idx) > 0:
                all_fac_osbal_rows = [osbal_rows[i] for i in all_fac_idx]
                osbal_cnt = len(all_fac_idx)

        # Format skenario final sesuai pola ACA
        final_scenario = _format_final_scenario(
            base_scenario=r.get("base_scenario", ""),
            source=r["source"],
            is_beda_curr=r.get("is_beda_curr", False),
            is_beda_per=r.get("is_beda_per", False),
            num_fac=len(r["fac_codes"]),
        )

        out = _build_output_row(
            r["suspend"], source=r["source"], scenario=final_scenario,
            osbal_rows=ref_osbal, osbal_count=osbal_cnt,
            resolved=r["resolved"], suspend_count=r.get("suspend_count", 1),
            all_fac_osbal_rows=all_fac_osbal_rows,
            is_beda_curr=r.get("is_beda_curr", False),
            is_beda_per=r.get("is_beda_per", False),
        )
        if "_merged_amount_ori" in r:
            out["_MERGED_AMOUNT_ORI"] = r["_merged_amount_ori"]
        out["_IS_BEDA_CURR"] = r.get("is_beda_curr", False)
        out["_IS_BEDA_PER"]  = r.get("is_beda_per", False)
        output_rows.append(out)

    df = pd.DataFrame(output_rows)

    # Fix format angka koma-desimal
    for col in ["AMOUNT ORI", "CCOS_OR_BAL", "CCOS_BAL_DUE"]:
        if col in df.columns:
            s = df[col]
            if s.dtype == object:
                s = s.astype(str).str.strip().str.replace(r'\s+', '', regex=True)
                has_comma = s.str.contains(',', na=False)
                s_eu = s.str.replace('.', '', regex=False).str.replace(',', '.', regex=False)
                s = s_eu.where(has_comma, s)
            df[col] = pd.to_numeric(s, errors="coerce")

    df = _compute_derived_cols(df)
    df["SKENARIO"] = df["SKENARIO"].replace({"INSURED_LIKE": "Insured only"}).astype(str).str.strip()
    df.loc[df["SKENARIO"].isin(["UNMATCHED", "Unmatching"]), "FLAG_PROD"] = "Unmatching"
    df.loc[df["SKENARIO"] == "UNMATCHED", "SKENARIO"] = "Unmatching"

    for col in FINAL_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    df = df[FINAL_COLUMNS]

    # [6] Export ke Excel
    _out_file = OUTPUT_FILE
    print(f"\n[6/6] Saving to: {_out_file} ...", flush=True)
    for c in df.select_dtypes(include=['object', 'string']).columns:
        df[c] = df[c].astype(str).str.replace(r'[\x00-\x08\x0B\x0C\x0E-\x1F]', '', regex=True)
    try:
        df.to_excel(_out_file, index=False)
    except PermissionError:
        _out_file = OUTPUT_FILE.replace(".xlsx", "_new.xlsx")
        print(f"  [WARN] File sedang terbuka di Excel — menyimpan ke: {_out_file}", flush=True)
        df.to_excel(_out_file, index=False)

    try:
        from openpyxl import load_workbook
        workbook = load_workbook(_out_file)
        worksheet = workbook.active
        text_cols = [
            "POLIS", "POLICY_CLEAN_1", "POLICY_CLEAN_2", "POLICY_CLEAN_3",
            "POLICY_CLEAN_4", "POLICY_CLEAN_5",
            "SERTIF_CLEAN_1", "SERTIF_CLEAN_2", "SERTIF_CLEAN_3",
            "SLIP_NO", "SLIP_NO_CLN",
            "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
        ]
        text_col_indices = [idx for idx, col in enumerate(df.columns, start=1) if col in text_cols]
        for col_idx in text_col_indices:
            for row in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=col_idx)
                cell.number_format = "@"
                if cell.value is not None:
                    cell.value = str(cell.value)
        workbook.save(_out_file)
    except Exception as e:
        print(f"  [WARN] Format text gagal: {e}")

    try:
        # pyrefly: ignore [missing-import]
        from excel_styler import apply_purple_column_style
        apply_purple_column_style(_out_file, "matching")
    except Exception as e:
        print(f"  [WARN] Styling gagal: {e}")

    # Export ke PostgreSQL
    load_dotenv()
    print("  Exporting ke PostgreSQL ...")
    conn_str = "postgresql+psycopg2://{}:{}@{}:{}/{}".format(
        os.environ.get("DB_USER", "postgres"), os.environ.get("DB_PASS", "postgres"),
        os.environ.get("DB_HOST", "localhost"), os.environ.get("DB_PORT", "5432"),
        os.environ.get("DB_NAME", "postgres"),
    )
    try:
        engine = create_engine(conn_str)
        df.to_sql("SUSPENSE_DATA_SUSPENSE_TPI_V1", con=engine, if_exists="append", index=False)
        print(f"      -> {len(df):,} rows berhasil di-export ke 'SUSPENSE_DATA_SUSPENSE_TPI_V1'")
    except Exception as e:
        print(f"  [WARN] PostgreSQL export gagal: {e}")

    elapsed = time.perf_counter() - t_start
    print(f"\n{'=' * 60}")
    print(f"  Selesai dalam {elapsed:.1f}s ({elapsed/60:.1f} menit)")
    print(f"  Rows  : {len(df):,}  |  Cols: {len(df.columns)}  |  File: {OUTPUT_FILE}")
    print(f"\n  FLAG_PROD:")
    flag_counts = df["FLAG_PROD"].astype(str).str.strip().value_counts()
    for flag, count in flag_counts.items():
        print(f"    {str(flag):<45}: {count:,}")
    print(f"\n  SKENARIO:")
    scen_counts = df["SKENARIO"].astype(str).str.strip().value_counts()
    for sce, count in scen_counts.items():
        print(f"    {str(sce):<50}: {count:,}")
    print("=" * 60)


if __name__ == "__main__":
    run()
