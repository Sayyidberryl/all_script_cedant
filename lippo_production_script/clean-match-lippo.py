"""
================================================================================
 SCRIPT CLEANSING DATA 3 - SUSPENSE (KHUSUS CEDANT: PT LIPPO GENERAL INSURANCE)
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_DIR = os.path.join(SCRIPT_DIR, "input")
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "output")

INPUT_FILE = os.path.join(INPUT_DIR, "3b. Database Suspense 150826.xlsx")
SHEET_NAME = "Sheet1"
HEADER_ROW = 2
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "lippo_output_suspense_V3.xlsx")
CEDANT_FILTER = "LIPPO"
STATUS_FILTER = "SUSPENSE"

MAX_BREAKDOWN_CODES = 5
MAX_BREAKDOWN_INSURED = 5
TEXT_FORMAT_COLUMN_KEYWORDS = ("POLIS", "POLICY", "SLIP", "SERTIF", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "CERTIFICATE")

# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
_ASTERISK_RE = re.compile(r"\*")
_TRANSACTION_HEADER_RE = re.compile(r"(?i)^\s*(?:Pembayaran|Penagihan)?\s*(?:dan\s+Penagihan)?\s*Premi\s*(?:Fakultatif|SOA|Pensesian)?\s*(?:PAR\s*&?\s*EQ|BIT|GIT|CECR|QS)?\s*")
_TRUNCATE_TRIGGER_RE = re.compile(r"(?i)\bsubsidiar|\bassociat|\baffiliat|\bfiliated\b|related\s+compan|respective\s+right|\bincluding\s+an[yd]|\bindu?ding\s+any|as\s+the\s+owner|as\s+owners?\b|as\s+main\s+contractor|\(as\s+mortgagee\s+and\s+loss\s+payee\)")
_REMOVE_ONLY_RE = re.compile(r"(?i)\bas\s+principal\b|\bsebagai\s+principal\b|\bsebagai\s+kontraktor\b|\bas\s+(?:co\s+)?propert(?:y)?\s+owner\b|\bas\s+operating\s+company\b|\bas\s+property\s+manager\b|\bas\s+contractor\b|\bas\s+co\s+ben[ie]ficiary\b|\bas\s+event\s+project\s+owner\b")
_GENERIC_BOILERPLATE_TAIL_RE = re.compile(r"(?i)\bsub\s*-?\s*kontraktor\b.*$|\banak\s+perusahaan\b.*$|\bafiliasi\s+perusahaan\b.*$|\bsemua\b\s*[–-]?\s*$")
_TITLE_RE = re.compile(r"(?i)\b(?:S\.?H\.?|S\.?I\.?Kom\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b")
_LEGAL_ENTITY_RE = re.compile(r"(?i)\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO|Perseroan\s+Terbatas)\b")
_ENTITY_SPLIT_RE = re.compile(r"(?i)\b(?:QQ|Q\.Q\.)\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|\bin\s+this\s+case\b|\bdalam\s+hal\s+ini\b|\band\b|\bor\b|,")
_LEGAL_ENTITY_DASH_SIGNAL_RE = re.compile(r"(?i),\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO)\s*[\r\n]*\s*-\s*")
_GENERIC_DASH_SEPARATOR_RE = re.compile(r"\s*[\r\n]+\s*-\s*|\s+-\s+")
_MIDDLE_BOILERPLATE_CLAUSE_RE = re.compile(r"(?i)(?:\band\s*/\s*or\s+)?(?:associat\w*|subsidiar\w*|affiliat\w*|related\s+compan\w*)(?:\s+and\s*/\s*or\s+(?:associat\w*|subsidiar\w*|affiliat\w*|related\s+compan\w*))*\s+for\s+(?:their|its)\s+respective\s+rights?\s+and\s+interests?\b")

_COMPRISING_OF_RE = re.compile(r"(?is)^(?P<head>.*?\bgroup)\s+of\s+companies\s+comprising\s+of\s*:\s*(?P<tail>.+)$")
_CONSISTS_OF_RE = re.compile(r"(?i)consists\s+of\s*:")
_NUMBERED_ITEM_RE = re.compile(r"\d+\s*\.\s*")
_AS_PER_LIST_RE = re.compile(r"(?i)as\s+per\s+list\s+attached\s*\((?P<inner>[^)]+)\)")
_GROUP_OF_COMPANIES_SUFFIX_RE = re.compile(r"(?i)\s+group\s+of\s+companies\b")
_BARE_AND_OR_RE = re.compile(r"(?i)\band\b|\bor\b")
_PAREN_PT_RE = re.compile(r"(?i)\(\s*(?P<inner>(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?)\s+[^)]+)\)")
_PAREN_STRIP_PREFIX_RE = re.compile(r"(?i)^(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?)\s*")
_DANGLING_WORD_RE = re.compile(r"(?i)^\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*|\s*\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$")

_SPECIFIC_SENTENCE_WHITELIST = {
    "PT Palladium Megah Lestari and/ or PERPETUAL (ASIA) LIMITED (in its capacity as Trustee of Lippo Malls Indonesia Retail Trust) Lippo Malls Indonesia and all subsidiary or controlled companies, for their respective rights and": [
        "PALLADIUM MEGAH LESTARI",
        "PERPETUAL ASIA LIMITED",
        "LIPPO MALLS INDONESIA",
    ],
}
_SPECIFIC_SENTENCE_WHITELIST_NORMALIZED = {re.sub(r"\s+", " ", k.strip()).upper(): v for k, v in _SPECIFIC_SENTENCE_WHITELIST.items()}

_AND_NAME_WHITELIST = [
    "MITSUBISHI HC CAPITAL AND FINANCE INDONESIA",
    "MITSUBISHI UFJ LEASE AND FINANCE INDONESIA",
]

def _strip_asterisk(text: str) -> str:
    return _ASTERISK_RE.sub("", text)

def _truncate_from_first_trigger(text: str) -> str:
    match = _TRUNCATE_TRIGGER_RE.search(text)
    return text[: match.start()] if match else text

def _strip_middle_boilerplate_clause(text: str) -> str:
    match = _MIDDLE_BOILERPLATE_CLAUSE_RE.search(text)
    if not match:
        return text
        
    trailing = text[match.end():].strip(" .,;:-")
    if len(trailing) >= 3 and re.search(r"[A-Za-z]{3,}", trailing):
        return f"{text[: match.start()].rstrip()}, {trailing}"
    return text

def _normalize_dash_entity_list(text: str) -> str:
    if _LEGAL_ENTITY_DASH_SIGNAL_RE.search(text):
        return _GENERIC_DASH_SEPARATOR_RE.sub(", ", text)
    return text

def _strip_transaction_header(text: str) -> str:
    return _TRANSACTION_HEADER_RE.sub("", text)

def _try_pattern_h_specific_sentence(text: str):
    key = re.sub(r"\s+", " ", text.strip()).upper()
    return _SPECIFIC_SENTENCE_WHITELIST_NORMALIZED.get(key)

def _is_real_company_name(inner: str) -> bool:
    name = _PAREN_STRIP_PREFIX_RE.sub("", inner).strip()
    words = [w for w in re.split(r"[\s\-]+", name) if w]
    
    if len(words) < 2:
        return False
        
    has_lowercase = any(c.islower() for c in name)
    all_short_upper = all(len(w) <= 3 and w.isupper() for w in words)
    return not (all_short_upper and not has_lowercase)

def _try_pattern_f_paren_pt(text: str):
    m = _PAREN_PT_RE.search(text)
    if not m:
        return None
        
    inner = m.group("inner").strip()
    if not _is_real_company_name(inner):
        return None
        
    head = text[: m.start()].strip()
    tail_after = text[m.end():].strip()
    
    if tail_after:
        head = f"{head} {tail_after}".strip()
    return [head, inner]

def _try_pattern_a_comprising_of(text: str):
    m = _COMPRISING_OF_RE.match(text.strip())
    if m:
        return [m.group("head").strip(), m.group("tail").strip()]
    return None

def _try_pattern_b_consists_of(text: str):
    m = _CONSISTS_OF_RE.search(text)
    if not m:
        return None
        
    tail = _truncate_from_first_trigger(text[m.end():])
    items = [seg.strip() for seg in _NUMBERED_ITEM_RE.split(tail) if seg.strip()]
    return items or None

def _try_pattern_c_as_per_list(text: str):
    m = _AS_PER_LIST_RE.search(text)
    if not m:
        return None
        
    parts = [p.strip() for p in re.split(r"\s*[\u2013\u2014-]\s*", m.group("inner")) if p.strip()]
    return parts or None

def _strip_group_of_companies_suffix(text: str) -> str:
    return _GROUP_OF_COMPANIES_SUFFIX_RE.sub("", text)

def _truncate_boilerplate_tail(text: str) -> str:
    text = _truncate_from_first_trigger(text)
    match_generic = _GENERIC_BOILERPLATE_TAIL_RE.search(text)
    if match_generic:
        text = text[: match_generic.start()]
    return _REMOVE_ONLY_RE.sub(" ", text)

def _strip_titles(text: str) -> str:
    return _TITLE_RE.sub(" ", text)

def _strip_legal_entity(text: str) -> str:
    return _LEGAL_ENTITY_RE.sub(" ", text)

def _protect_and_whitelist(text: str) -> tuple[str, dict]:
    placeholders = {}
    for i, phrase in enumerate(_AND_NAME_WHITELIST):
        amp_variant = re.sub(r"(?i)\bAND\b", "&", phrase)
        for j, candidate in enumerate([phrase, amp_variant]):
            pattern = re.compile(re.escape(candidate).replace(r"\ ", r"\s+"), re.IGNORECASE)
            match = pattern.search(text)
            if match:
                key = f"ANDWL{i}{j}PLACEHOLDER"
                placeholders[key] = phrase
                text = pattern.sub(key, text)
    return text, placeholders

def _restore_and_whitelist(text: str, placeholders: dict[str, str]) -> str:
    for key, phrase in placeholders.items():
        text = re.sub(key, phrase, text, flags=re.IGNORECASE)
    return text

def _final_polish(text: str) -> str:
    t = re.sub(r"[().;:\"']", " ", text)
    t = re.sub(r"[-/]", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    
    previous = None
    while previous != t:
        previous = t
        t = _DANGLING_WORD_RE.sub("", t).strip()
    return t.upper()

def _finalize_segments(raw_parts: list) -> list:
    results = []
    for part in raw_parts:
        part, placeholders = _protect_and_whitelist(part)
        for sub in _ENTITY_SPLIT_RE.split(part):
            cleaned = _truncate_boilerplate_tail(sub)
            cleaned = _strip_group_of_companies_suffix(cleaned)
            cleaned = _strip_legal_entity(cleaned)
            cleaned = _final_polish(cleaned)
            cleaned = _restore_and_whitelist(cleaned, placeholders)
            
            if cleaned and len(cleaned) >= 2 and re.search(r"[A-Z]", cleaned) and cleaned not in results:
                results.append(cleaned)
    return results

def _clean_insured_name_breakdown(text: str) -> list:
    specific_result = _try_pattern_h_specific_sentence(text)
    if specific_result is not None:
        return specific_result

    t = text.strip()
    t = _strip_transaction_header(t)
    t = _strip_titles(t)
    t = _normalize_dash_entity_list(t)
    t = _strip_middle_boilerplate_clause(t)

    special_segments = _try_pattern_c_as_per_list(t) or _try_pattern_b_consists_of(t)
    if special_segments is not None:
        return _finalize_segments(special_segments)

    special_segments = _try_pattern_a_comprising_of(t)
    if special_segments is not None:
        return _finalize_segments(special_segments)

    special_segments = _try_pattern_f_paren_pt(t)
    if special_segments is not None:
        return _finalize_segments(special_segments)

    return _finalize_segments([_truncate_boilerplate_tail(_strip_group_of_companies_suffix(t))])

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []
        
    original = text.strip()
    t = _strip_asterisk(original)
    results = _clean_insured_name_breakdown(t)
    
    if len(results) > MAX_BREAKDOWN_INSURED:
        return [original]
    return results


# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO
# ==============================================================================
_TOK_MAIN_RE = re.compile(r"^\d{13}$")
_TOK_CERT_RE = re.compile(r"^\d{6}$")
_TOK_ENDORSE_RE = re.compile(r"^\d{1,2}/\d{1,2}$")
_TOK_STATUS_RE = re.compile(r"(?i)^(?:New|Endorsement|Cancel\s*All|Adjustment|Cancellation|Reinstatement)$")

def _classify_token(tok: str) -> str:
    if _TOK_MAIN_RE.match(tok): return "MAIN"
    if _TOK_CERT_RE.match(tok): return "CERT"
    if _TOK_ENDORSE_RE.match(tok): return "ENDORSE"
    if _TOK_STATUS_RE.match(tok): return "STATUS"
    return "OTHER"

def clean_split_code(text: str, skip_combined_if_cert: bool = False) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []

    t = text.strip()
    t = re.sub(r"[.,]", "", t)
    t = re.sub(r"\s+", " ", t).strip()
    
    classified = []
    for tok in t.split("-"):
        tok = tok.strip()
        if tok:
            classified.append((_classify_token(tok), tok))

    codes = []
    i = 0
    n = len(classified)
    
    while i < n:
        kind, tok = classified[i]
        
        if kind == "MAIN":
            suffix_parts = []
            j = i + 1
            
            while j < n and classified[j][0] in ("CERT", "ENDORSE"):
                suffix_parts.append(classified[j][1])
                j += 1

            followed_by_status = j < n and classified[j][0] == "STATUS"
            has_cert_suffix = False
            for k in range(i + 1, j):
                if classified[k][0] == "CERT":
                    has_cert_suffix = True
                    break

            if suffix_parts and not followed_by_status:
                if not (skip_combined_if_cert and has_cert_suffix):
                    combined = f"{tok}-" + "-".join(suffix_parts)
                    if combined not in codes:
                        codes.append(combined)
                
                if tok not in codes:
                    codes.append(tok)
            elif tok not in codes:
                codes.append(tok)

            if followed_by_status:
                j += 1
                
            i = j
            continue
            
        i += 1

    if not codes:
        return [t] if t else []
        
    if len(codes) > MAX_BREAKDOWN_CODES:
        return codes[:1]
    return codes


# ==============================================================================
# 3. EKSTRAKSI CERTIFICATE
# ==============================================================================
_SD_NORMALIZE_RE = re.compile(r"(?i)\bs\s*/\s*d\b")
_CERT_SEGMENT_RE = re.compile(r"(?i)(?P<base>\d{13})-(?P<start>\d{5,6})(?!\d)(?:\s*(?:SD|-)\s*(?:(?P=base)-)?(?P<end>\d{5,6})(?!\d))?")

def _normalize_sd(text: str) -> str:
    return _SD_NORMALIZE_RE.sub("SD", text)

def _to_six_digit(num_str: str) -> str | None:
    if len(num_str) == 5:
        return "0" + num_str
    if len(num_str) == 6:
        return num_str
    return None

def _format_certificate_range(start_six: str, end_six: str) -> str:
    start_num = int(start_six)
    end_num = int(end_six)
    
    if end_num < start_num:
        return ""
        
    count = end_num - start_num + 1
    if count > 3:
        return f"{start_six} SD {end_six}"
        
    return ", ".join(f"{n:06d}" for n in range(start_num, end_num + 1))

def extract_certificate(text: str) -> str:
    if not isinstance(text, str):
        return ""
        
    t = text.strip()
    if not t or t.lower() in ("nan", "none", "-"):
        return ""

    t = _normalize_sd(t)
    results = []
    
    for match in _CERT_SEGMENT_RE.finditer(t):
        start_six = _to_six_digit(match.group("start"))
        if start_six is None:
            continue
            
        end_digits = match.group("end")
        if end_digits:
            end_six = _to_six_digit(end_digits)
            if end_six is None:
                continue
            cert = _format_certificate_range(start_six, end_six)
        else:
            cert = start_six
            
        if cert and cert not in results:
            results.append(cert)

    return ", ".join(results)


# ==============================================================================
# 4. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_status_suspense(df_raw: pd.DataFrame, keyword: str) -> pd.DataFrame:
    status_col = None
    for c in df_raw.columns:
        if str(c).upper().strip() == "STATUS":
            status_col = c
            break
            
    if not status_col:
        raise ValueError("Kolom STATUS tidak ditemukan di source file.")

    mask = df_raw[status_col].astype(str).str.strip().str.upper() == keyword.upper()
    df_filtered = df_raw[mask].copy()
    
    print(f"[INFO] Filter STATUS = '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered

def filter_cedant(df_raw: pd.DataFrame, keyword: str) -> pd.DataFrame:
    cedant_cols = [c for c in df_raw.columns if "CEDANT" in str(c).upper()]
    if not cedant_cols:
        raise ValueError("Kolom CEDANT tidak ditemukan di source file.")

    mask = pd.Series(False, index=df_raw.index)
    for col in cedant_cols:
        mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)

    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant mengandung '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered


# ==============================================================================
# 5. BUILD OUTPUT
# ==============================================================================
def _detect_key_columns(df: pd.DataFrame) -> tuple[str, str, str]:
    col_insured = None
    col_polis = None
    col_slip = None
    
    for c in df.columns:
        c_upper = str(c).upper().strip()
        if "INSURED" in c_upper and col_insured is None:
            col_insured = c
        elif c_upper == "POLIS" and col_polis is None:
            col_polis = c
        elif "SLIP" in c_upper and col_slip is None:
            col_slip = c

    if None in (col_insured, col_polis, col_slip):
        raise ValueError(f"Kolom kunci tidak lengkap. INSURED={col_insured}, POLIS={col_polis}, SLIP={col_slip}")
        
    print(f"[INFO] Kolom terdeteksi -> INSURED='{col_insured}', POLIS='{col_polis}', SLIP='{col_slip}'")
    return col_insured, col_polis, col_slip

def _breakdown_all_rows(df: pd.DataFrame, col_insured: str, col_polis: str, col_slip: str) -> tuple[list, list, list, list, int, int, int]:
    insured_cln_all = []
    polis_cln_all = []
    slip_cln_all = []
    certificate_all = []
    
    max_ins = 0
    max_pol = 0
    max_slp = 0

    for _, row in df.iterrows():
        ins_val = row[col_insured]
        pol_val = row[col_polis]
        slp_val = row[col_slip]
        
        ins_cln = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []
        cert = extract_certificate(str(pol_val)) if pd.notna(pol_val) else ""
        pol_cln = clean_split_code(str(pol_val), skip_combined_if_cert=bool(cert)) if pd.notna(pol_val) else []
        slp_cln = clean_split_code(str(slp_val)) if pd.notna(slp_val) else []

        insured_cln_all.append(ins_cln)
        polis_cln_all.append(pol_cln)
        slip_cln_all.append(slp_cln)
        certificate_all.append(cert)
        
        max_ins = max(max_ins, len(ins_cln))
        max_pol = max(max_pol, len(pol_cln))
        max_slp = max(max_slp, len(slp_cln))

    return insured_cln_all, polis_cln_all, slip_cln_all, certificate_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_insured, col_polis, col_slip = _detect_key_columns(df_filtered)
    breakdown_results = _breakdown_all_rows(df_filtered, col_insured, col_polis, col_slip)
    insured_cln_all, polis_cln_all, slip_cln_all, certificate_all, max_ins, max_pol, max_slp = breakdown_results

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)
    output_cols_data = {}
    
    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        
        if col == col_insured:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLEAN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in insured_cln_all]
                
        if col == col_polis:
            if max_pol >= 1:
                output_cols_data["POLICY_CLEAN_1"] = [lst[0] if len(lst) >= 1 else "" for lst in polis_cln_all]
                
            output_cols_data["SERTIF_CLEAN_1"] = certificate_all
            
            for i in range(2, max_pol + 1):
                output_cols_data[f"POLICY_CLEAN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in polis_cln_all]
                
        if col == col_slip:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_CLEAN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slip_cln_all]

    df_master = pd.DataFrame(output_cols_data)
    
    drop_cols = []
    for c in df_master.columns:
        if "_CLN_" in c and df_master[c].astype(str).str.strip().eq("").all():
            drop_cols.append(c)
            
    return df_master.drop(columns=drop_cols)


# ==============================================================================
# 6. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    
    workbook = load_workbook(output_path)
    worksheet = workbook.active
    text_col_indices = []
    
    for idx, col in enumerate(df.columns, start=1):
        for kw in TEXT_FORMAT_COLUMN_KEYWORDS:
            if kw in str(col).upper():
                text_col_indices.append(idx)
                break

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)
                
    workbook.save(output_path)


# ==============================================================================
# 7. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70)
    print(" CLEANSING DATA 3 - SUSPENSE | CEDANT: PT LIPPO GENERAL INSURANCE")
    print("=" * 70)
    
    input_path = INPUT_FILE
    if not os.path.exists(input_path):
        candidates = []
        for f in os.listdir("."):
            if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_"):
                candidates.append(f)
                
        if not candidates:
            raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")
            
        input_path = candidates[0]
        print(f"[AUTO-DETECT] File input tidak ditemukan di path default, memakai: '{input_path}'")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    print(f"[INFO] Total baris source: {len(df_raw)}")

    df_suspense = filter_status_suspense(df_raw, STATUS_FILTER)
    if df_suspense.empty:
        print(f"[WARNING] Tidak ada baris dengan STATUS = '{STATUS_FILTER}'. Proses dihentikan.")
        return

    df_lippo = filter_cedant(df_suspense, CEDANT_FILTER)
    if df_lippo.empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter cedant. Proses dihentikan.")
        return

    df_hasil = build_output(df_lippo)
    save_with_text_format(df_hasil, OUTPUT_FILE)
    
    print("=" * 70)
    print(f"[SUCCESS] Selesai. Total baris output: {len(df_hasil)}")
    print(f"[SUCCESS] File hasil: '{OUTPUT_FILE}'")
    print("=" * 70)

if __name__ == "__main__":
    main()

"""
================================================================================
 SCRIPT CLEANSING DATA 2 - OSBAL (CEDANT: PT LIPPO GENERAL INSURANCE)
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_DIR = os.path.join(SCRIPT_DIR, "input")
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "output")

INPUT_FILE = os.path.join(INPUT_DIR, "2b. transaksi Osbal 01.01.23 - 17.08.26.xlsx")
SHEET_NAME = "Query result"
HEADER_ROW = 0
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "lippo_output_osbal_V2.xlsx")
CEDANT_FILTER = "LIPPO"

CEDANT_COLUMN = "CCOS_COMP_NAME"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"

CLSDT_POLIS_COLUMN = "CLSDT_POLICY_NO"
CLSDT_SLIP_COLUMN = "CLSDT_SLIP_NO"
CLSDT_SERTF_COLUMN = "CLSDT_SERTF_NO"
CERTIFICATE_OUTPUT_COL = "SERTIF"

MAX_BREAKDOWN_CODES = 5
CERT_DIGIT_LEN = 6
CERT_RANGE_JOIN_THRESHOLD = 3

TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CLSDT", "CERTIFICATE",
)

# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
_BYPASS_ALL_RE = re.compile(r"(?i)\bPENYELESAIAN\s+HUTANG\s+PIUTANG\b")
_TRANSACTION_HEADER_RE = re.compile(r"(?i)^\s*(?:Pembayaran|Penagihan)?\s*(?:dan\s+Penagihan)?\s*Premi\s*(?:Fakultatif|SOA|Pensesian)?\s*(?:PAR\s*&?\s*EQ|BIT|GIT|CECR|QS)?\s*")
_TRUNCATE_TRIGGER_RE = re.compile(r"(?i)\bsubsidiar|\bassociat|\baffiliat|\bfiliated\b|related\s+compan|respective\s+right|\bincluding\s+an[yd]|\bindu?ding\s+any|as\s+the\s+owner|as\s+owners?\b|as\s+main\s+contractor|\(as\s+mortgagee\s+and\s+loss\s+payee\)")
_REMOVE_ONLY_RE = re.compile(r"(?i)\bas\s+principal\b|\bsebagai\s+principal\b|\bsebagai\s+kontraktor\b")
_ONLY_WORD_RE = re.compile(r"(?i)\bONLY\b")
_GENERIC_BOILERPLATE_TAIL_RE = re.compile(r"(?i)\bsub\s*-?\s*kontraktor\b.*$|\banak\s+perusahaan\b.*$|\bafiliasi\s+perusahaan\b.*$|\bdan\s+semua\b\s*[–-]?\s*$")
_TITLE_RE = re.compile(r"(?i)\b(?:S\.?H\.?|S\.?I\.?Kom\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b")
_LEGAL_ENTITY_RE = re.compile(r"(?i)\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO|Perseroan\s+Terbatas)\b")
_ENTITY_SPLIT_RE = re.compile(r"(?i)\b(?:QQ|Q\.Q\.)\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|\bin\s+this\s+case\b|\bdalam\s+hal\s+ini\b|,|/")
_LEGAL_ENTITY_DASH_SIGNAL_RE = re.compile(r"(?i),\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO)\s*[\r\n]*\s*-\s*")
_GENERIC_DASH_SEPARATOR_RE = re.compile(r"\s*[\r\n]+\s*-\s*|\s+-\s+")
_TRAILING_PAREN_MERGE_RE = re.compile(r"(?i),?\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?)?\s*\(([^)]+)\)\s*$")
_INLINE_PAREN_MERGE_RE = re.compile(r"(?i),?\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?)?\s*\(([^)]+)\)")
_QQ_PAREN_UNWRAP_RE = re.compile(r"(?i)\bQQ\s*\(\s*([^)]*)\)")
_INSURED_DESCRIPTIVE_RE = re.compile(r"(?i)\bVARIOUS\s+INSUREDS?\b|\bACCEPTED\s+BY\b")
_SUSPENSE_WORD_RE = re.compile(r"(?i)\bSUSPENSE\b")

_MONTH_WORDS = ("JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI", "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER", "DES", "JANUARY", "FEBRUARY", "MARCH", "MAY", "JUNE", "JULY", "AUGUST", "OCTOBER", "DECEMBER", "JAN", "FEB", "MAR", "APR", "JUN", "JUL", "AUG", "SEP", "SEPT", "OCT", "NOV", "DEC")
_DATE_BATCH_TOKEN_RE = re.compile(r"(?i)\b(?:" + "|".join(_MONTH_WORDS) + r")\b|\b(?:19|20)\d{2}\b|\bBATCH\s*\d*\b")

_LIPPO_KNOWN_ENTITIES = ["MATAHARI PUTRA PRIMA", "MATAHARI BOSTON DRIGSTORE", "MATAHARI PUSAKA TAMA", "LIPPO GROUP"]
_LIPPO_KNOWN_ENTITIES_SORTED = sorted(_LIPPO_KNOWN_ENTITIES, key=len, reverse=True)
_KNOWN_ENTITY_RE = re.compile("|".join(re.escape(name) for name in _LIPPO_KNOWN_ENTITIES_SORTED))
_ENTITY_SLASH_NORMALIZE_PATTERNS = [(name, re.compile(r"(?i)\b" + r"(?:\s*/\s*|\s+)".join(re.escape(w) for w in name.split()) + r"\b")) for name in _LIPPO_KNOWN_ENTITIES_SORTED if len(name.split()) > 1]

_INSURED_DASH_SPLIT_EXCEPTIONS = ("EBDI VAUNDRI - OSCAR OMEGA", "ADHI - HUTAMA - NINDYA - ABIPRAYA KSO", "PEMERINTAH PROVINSI SULAWESI SELATAN - DINAS BINA MARGA DAN BINA KONSTRUKSI")
_INSURED_DASH_MERGE_SUBSTRINGS = ("PP - WASKITA - WIJAYA KARYA",)

def _unwrap_qq_paren(text: str) -> str:
    def _replace(match):
        inner = re.sub(r"\s+", " ", re.sub(r"[:;]", " ", match.group(1).strip())).strip()
        return f"QQ , {inner}" if inner else "QQ"
    return _QQ_PAREN_UNWRAP_RE.sub(_replace, text)

def _merge_trailing_paren(text: str) -> str:
    match = _TRAILING_PAREN_MERGE_RE.search(text)
    if not match: return text
    inner = match.group(1).strip()
    return text[: match.start()] + text[match.end():] if not inner or len(inner.split()) > 4 or _TRUNCATE_TRIGGER_RE.search(inner) else text[: match.start()] + " " + inner

def _merge_inline_paren(text: str) -> str:
    def _replace(match):
        inner = match.group(1).strip()
        return "" if not inner or len(inner.split()) > 4 or _TRUNCATE_TRIGGER_RE.search(inner) else f" {inner}"
    return _INLINE_PAREN_MERGE_RE.sub(_replace, text)

def _strip_date_batch_info(text: str) -> str: return _DATE_BATCH_TOKEN_RE.sub(" ", text)

def _truncate_boilerplate_tail(text: str) -> str:
    if match := _TRUNCATE_TRIGGER_RE.search(text): text = text[: match.start()]
    if match := _GENERIC_BOILERPLATE_TAIL_RE.search(text): text = text[: match.start()]
    return _REMOVE_ONLY_RE.sub(" ", text)

def _final_polish(text: str) -> str:
    t = re.sub(r"\s+", " ", re.sub(r"[-/().;:\"']", " ", text)).strip()
    prev = None
    while prev != t:
        prev, t = t, re.sub(r"(?i)^\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*|\s*\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$", "", t).strip()
    return t.upper()

def _apply_dash_merge_substrings(text: str) -> str:
    for sub in _INSURED_DASH_MERGE_SUBSTRINGS:
        if sub.upper() in text.upper():
            merged = re.sub(r"\s*-\s*", " ", sub)
            idx = text.upper().find(sub.upper())
            if idx != -1: text = text[:idx] + merged + text[idx + len(sub):]
    return text

def _try_split_insured_by_dash(text: str):
    stripped = text.strip()
    if _SUSPENSE_WORD_RE.search(stripped): return None
    for exception in _INSURED_DASH_SPLIT_EXCEPTIONS:
        if stripped.upper() == exception.upper():
            return [p.strip() for p in re.split(r"\s*-\s*", stripped) if p.strip()]
    return None

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return []
    original = text.strip()

    if _BYPASS_ALL_RE.search(original) or _INSURED_DESCRIPTIVE_RE.search(original) or _SUSPENSE_WORD_RE.search(original):
        return [original]

    if dash_split := _try_split_insured_by_dash(original):
        results = []
        for part in dash_split:
            cleaned = _final_polish(_strip_date_batch_info(_LEGAL_ENTITY_RE.sub(" ", _truncate_boilerplate_tail(part))))
            if cleaned and cleaned not in results: results.append(cleaned)
        return [original] if len(results) > MAX_BREAKDOWN_CODES or not results else results

    t = _apply_dash_merge_substrings(original)
    t = _unwrap_qq_paren(_ONLY_WORD_RE.sub(" ", t))
    t = _merge_trailing_paren(t)
    t = _merge_inline_paren(t)
    t = _TRANSACTION_HEADER_RE.sub("", t)
    t = _TITLE_RE.sub(" ", t)
    if _LEGAL_ENTITY_DASH_SIGNAL_RE.search(t): t = _GENERIC_DASH_SEPARATOR_RE.sub(", ", t)
    for canonical, pat in _ENTITY_SLASH_NORMALIZE_PATTERNS: t = pat.sub(canonical, t)
    t = _truncate_boilerplate_tail(t)

    results = []
    for part in _ENTITY_SPLIT_RE.split(t):
        cleaned = _final_polish(_strip_date_batch_info(_LEGAL_ENTITY_RE.sub(" ", _truncate_boilerplate_tail(part))))
        if not cleaned or len(cleaned) < 2 or not re.search(r"[A-Z]", cleaned): continue

        pos, n, known_found = 0, len(cleaned), []
        while pos < n:
            while pos < n and cleaned[pos] == " ": pos += 1
            if pos >= n or not (match := _KNOWN_ENTITY_RE.match(cleaned, pos)): break
            known_found.append(match.group(0))
            pos = match.end()

        if len(known_found) > 1:
            for name in known_found:
                if name not in results: results.append(name)
            continue
            
        if cleaned not in results: results.append(cleaned)

    return [original] if len(results) > MAX_BREAKDOWN_CODES else results

# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO
# ==============================================================================
_MAIN_DIGITS_PATTERN = r"\d{12,13}"
_TOK_MAIN_RE = re.compile(rf"^{_MAIN_DIGITS_PATTERN}$")
_TOK_CERT_RE = re.compile(r"^\d{6}$")
_TOK_SUFFIX_REPLACE_RE = re.compile(r"^\d{1,11}$")
_TOK_ENDORSE_RE = re.compile(r"^\d{1,2}[/\x00]\d{1,2}$")
_TOK_PLACEHOLDER_RE = re.compile(r"(?i)^P\d{1,2}$")
_TOK_STATUS_RE = re.compile(r"(?i)^(?:New|End(?:orsement)?|Cancel\s*All|Adjustment|Cancellation|Reinstatement|TBA\.?|Attachment|DN)$")
_VARIOUS_WORD_RE = re.compile(r"(?i)\bVAR(?:IOUS)?\b")
_SD_WORD_RE = re.compile(r"(?i)\bSD\b")
_AMPERSAND_SUFFIX_AMBIGUOUS_RE = re.compile(r"\d+\s*&\s*\d+")
_CURRENCY_RE = re.compile(r"(?i)\b(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\b")
_SEE_ATTACHMENT_RE = re.compile(r"(?i)\bSEE\s+ATTACHMENT\b")
_PENDING_SLIP_RE = re.compile(r"(?i)\bPENDING\s+SLIP\b")
_NO_LABEL_RE = re.compile(r"(?i)\bNO\s*[:.]")
_LETTER_REF_CODE_RE = re.compile(r"(?i)\b(?:\d{1,6}\s*/\s*)?[A-Z]{1,6}-[A-Z]{1,6}\s*/\s*[IVXLCM]{1,6}\s*/\s*\d{2,4}\b")
_BATCH_TOKEN_RE = re.compile(r"(?i)\bBATCH\s*\d*\b")
_BORDERO_WORD_RE = re.compile(r"(?i)\bBORDERO\b|\bBORD\b")
_POLIS_MONTH_RE = re.compile(r"(?i)\b(?:" + "|".join(_MONTH_WORDS) + r")\.?\s*(?:\d{2,4})?\b")
_POLIS_DATE_SLASH_RE = re.compile(r"\b\d{1,2}/\d{1,2}/\d{2,4}\b")
_POLIS_YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")
_POLIS_CN_RE = re.compile(r"(?i)\bCN\b")
_POLIS_NARRATIVE_WORDS_RE = re.compile(r"(?i)\bSANY\b|\bNO\s+SLIP\b|\bNO\s+POLIS\b|\bPOLIS\s+VARIOUS\b")
_POLIS_STATUS_WORDS_RE = re.compile(r"(?i)\bCancel\s*All\b|\bNew\b|\bEnd(?:orsement)?\b|\bAdjustment\b|\bCancellation\b|\bReinstatement\b")
_EDGE_SEPARATOR_RE = re.compile(r"^[\s\-/,]+|[\s\-/,]+$")
_NARRATIVE_PREFIX_RE = re.compile(r"(?i)^\s*(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\s+\d{4}\s*-\s*(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)?\s*/?\s*")
_MAIN_CODE_FIRST_RE = re.compile(_MAIN_DIGITS_PATTERN)
_STRONG_SPLIT_CAPTURE_RE = re.compile(r"(\+|-|[,/]|\s+)")
_ALPHA_WORD_RE = re.compile(r"[A-Za-z][A-Za-z0-9]*")
_ALLOWED_ALPHA_TOKENS = {"NEW", "END", "ENDORSEMENT", "CANCEL", "ALL", "ADJUSTMENT", "CANCELLATION", "REINSTATEMENT", "VARIOUS", "VAR", "SD", "ATTACHMENT", "TBA", "DN"}
_FALLBACK_NOISE_WORDS_RE = re.compile(r"(?i)\bTBA\.?\b|\bATTACHMENT\b|\bDN\b")
_BASE_SUFFIX_BASE_RE = re.compile(rf"^({_MAIN_DIGITS_PATTERN})-(\d{{1,11}})\s*-\s*({_MAIN_DIGITS_PATTERN})$")
_E_HASH_SUFFIX_RE = re.compile(rf"(?i)^({_MAIN_DIGITS_PATTERN})\s*-\s*E#\s*\d+\s*$")
_SHORT_BASE_MULTI_SUFFIX_RE = re.compile(r"^(\d{9,11})-(\d{1,6}(?:\s*[,+]\s*\d{1,6})+)\s*$")
_HAS_VALID_MAIN_DIGITS_RE = re.compile(_MAIN_DIGITS_PATTERN)

def _strip_polis_slip_noise(text: str) -> str:
    for pat in (_NARRATIVE_PREFIX_RE, _LETTER_REF_CODE_RE, _NO_LABEL_RE, _PENDING_SLIP_RE, _SEE_ATTACHMENT_RE, _POLIS_NARRATIVE_WORDS_RE, _POLIS_STATUS_WORDS_RE, _BATCH_TOKEN_RE, _BORDERO_WORD_RE, _POLIS_DATE_SLASH_RE, _POLIS_MONTH_RE, _POLIS_YEAR_RE, _POLIS_CN_RE, _CURRENCY_RE):
        text = pat.sub(" ", text)
    match = _MAIN_CODE_FIRST_RE.search(text)
    t = text[match.start():] if match and re.search(r"[A-Za-z]", text[:match.start()]) else text
    return _EDGE_SEPARATOR_RE.sub("", re.sub(r"\s+", " ", t).strip()).strip()

def _split_preserving_continuity(t: str):
    segments, is_soft_list, pending_soft = [], [], False
    for part in _STRONG_SPLIT_CAPTURE_RE.split(t):
        if part in ("+", "-"): pending_soft = True
        elif part is not None and part.strip() == "" and _STRONG_SPLIT_CAPTURE_RE.fullmatch(part or ""): continue
        elif part is not None and _STRONG_SPLIT_CAPTURE_RE.fullmatch(part): pending_soft = False
        elif (seg := (part or "").strip()):
            segments.append(seg)
            is_soft_list.append(pending_soft)
            pending_soft = False
    return segments, is_soft_list

def _clean_fallback_original(original: str) -> str:
    t = _FALLBACK_NOISE_WORDS_RE.sub("", original)
    return re.sub(r"\s{2,}", " ", re.sub(r"[+\-/,]\s*$", "", t.strip())).strip()

def _has_descriptive_narrative(text: str) -> bool:
    return any(not (_TOK_PLACEHOLDER_RE.match(w) or w.upper() in _ALLOWED_ALPHA_TOKENS) for w in _ALPHA_WORD_RE.findall(text))

def _expand_numeric_range(main_a: str, main_b: str):
    if len(main_a) != len(main_b) or not (main_a.isdigit() and main_b.isdigit()): return None
    diff_len = next((i for i in range(1, len(main_a) + 1) if main_a[-i] != main_b[-i]), 0)
    if diff_len == 0 or main_a[:-diff_len] != main_b[:-diff_len]: return None
    try: start, end = int(main_a[-diff_len:]), int(main_b[-diff_len:])
    except ValueError: return None
    return [f"{main_a[:-diff_len]}{str(n).zfill(diff_len)}" for n in range(start, end + 1)] if start <= end and (end - start + 1) <= MAX_BREAKDOWN_CODES else None

def _split_sd_range(text: str):
    if not _SD_WORD_RE.search(text): return None
    parts = _SD_WORD_RE.split(text, maxsplit=1)
    if len(parts) != 2: return "__FALLBACK__"

    left_all, right_all = re.findall(r"\d+", parts[0]), re.findall(r"\d+", parts[1])
    if not left_all or not right_all or not _TOK_MAIN_RE.match(left_all[-1]) or not _TOK_MAIN_RE.match(right_all[0]): return "__FALLBACK__"

    expanded = _expand_numeric_range(left_all[-1], right_all[0])
    return expanded + [t for t in right_all[1:] if _TOK_MAIN_RE.match(t) and t not in expanded] if expanded else "__FALLBACK__"

def _try_reconstruct_suffix23(tokens: list, tokens_is_soft: list):
    if not tokens or not _TOK_MAIN_RE.match(tokens[0]) or not all(_TOK_MAIN_RE.match(t) or _TOK_SUFFIX_REPLACE_RE.match(t) for t in tokens): return None
    if any(_TOK_SUFFIX_REPLACE_RE.match(tok) and not _TOK_MAIN_RE.match(tok) and not soft for tok, soft in zip(tokens, tokens_is_soft)): return None
    if not any(_TOK_SUFFIX_REPLACE_RE.match(tok) and not _TOK_MAIN_RE.match(tok) for tok in tokens): return None

    codes, current_base = [], None
    for tok in tokens:
        if _TOK_MAIN_RE.match(tok):
            current_base = tok
            if current_base not in codes: codes.append(current_base)
        elif current_base and (rec := current_base[:-len(tok)] + tok) not in codes:
            codes.append(rec)
    return codes

def clean_split_code(text: str) -> list:
    codes, _ = _clean_split_code_impl(text)
    return codes

def _clean_split_code_impl(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return [], False
    original = text.strip()

    if _BYPASS_ALL_RE.search(original): return [original], False

    if e_hash_match := _E_HASH_SUFFIX_RE.match(original): return [e_hash_match.group(1)], True

    if short_base_match := _SHORT_BASE_MULTI_SUFFIX_RE.match(original):
        base, suffix_blob = short_base_match.groups()
        suffixes = [s.strip() for s in re.split(r"[,+]", suffix_blob) if s.strip()]
        if suffixes:
            codes = [base]
            for suf in suffixes:
                if len(suf) < len(base): codes.append(base[: -len(suf)] + suf)
            deduped = []
            for c in codes:
                if c not in deduped: deduped.append(c)
            if len(deduped) <= MAX_BREAKDOWN_CODES: return deduped, True

    if not _HAS_VALID_MAIN_DIGITS_RE.search(original): return [original], False

    t = _strip_polis_slip_noise(original)
    if not t or _has_descriptive_narrative(t) or _AMPERSAND_SUFFIX_AMBIGUOUS_RE.search(t): return [original], False

    if (sd_res := _split_sd_range(t)) is not None: return ([original], False) if sd_res == "__FALLBACK__" else (sd_res, True)

    t = re.sub(r"(-\s*\d{1,2})\s*/\s*(\d{1,2})\b", lambda m: f"{m.group(1)}\x00{m.group(2)}", re.sub(r"[\u2012\u2013\u2014\u2015]", "-", _VARIOUS_WORD_RE.sub(" ", t)))
    strong_segments, segment_is_soft = _split_preserving_continuity(t)

    filtered_pairs = [(re.sub(r"[.]", "", seg).strip(), soft) for seg, soft in zip(strong_segments, segment_is_soft) if seg and not bool(_TOK_PLACEHOLDER_RE.match(re.sub(r"[.]", "", seg).strip()) or _TOK_STATUS_RE.match(re.sub(r"[.]", "", seg).strip()))]
    if not filtered_pairs: return [original], False

    tokens, tokens_is_soft = [p[0] for p in filtered_pairs], [p[1] for p in filtered_pairs]
    if (recon := _try_reconstruct_suffix23(tokens, tokens_is_soft)) is not None:
        return ([_clean_fallback_original(original)], False) if len(recon) > MAX_BREAKDOWN_CODES else (recon, True)

    codes, current_base = [], None
    for tok, soft in zip(tokens, tokens_is_soft):
        if current_base and soft and re.match(r"^\d{1,2}\x00\d{1,2}$", tok): continue
        if current_base and soft and _TOK_SUFFIX_REPLACE_RE.match(tok) and not _TOK_MAIN_RE.match(tok) and len(tok) < len(current_base):
            if (rec := current_base[: -len(tok)] + tok) not in codes: codes.append(rec)
            continue

        if "\x00" in tok and not _TOK_MAIN_RE.match(tok):
            if base_match := re.match(rf"^({_MAIN_DIGITS_PATTERN})-?(\d{{1,2}})\x00(\d{{1,2}})$", tok):
                current_base = base_match.group(1)
                if current_base not in codes: codes.append(current_base)
                continue

        anchor, *suffix_parts = [p.replace("\x00", "/") for p in tok.split("-")]
        if not _TOK_MAIN_RE.match(anchor): continue
        current_base = anchor

        while suffix_parts and _TOK_STATUS_RE.match(suffix_parts[-1]): suffix_parts.pop()
        for p in suffix_parts:
            if _TOK_ENDORSE_RE.match(p.replace("/", "\x00")) or re.match(r"^\d{1,2}/\d{1,2}$", p): continue
            if _TOK_SUFFIX_REPLACE_RE.match(p) and len(p) < len(anchor) and (rec := anchor[: -len(p)] + p) not in codes:
                codes.append(rec)
        if anchor not in codes: codes.append(anchor)

    return ([_clean_fallback_original(original)], False) if not codes or len(codes) > MAX_BREAKDOWN_CODES else (codes, True)

# ==============================================================================
# 2B. CLEANSING CERTIFICATE
# ==============================================================================
_CERT_SD_NORMALIZE_RE = re.compile(r"(?i)\bS\s*/\s*D\b")
_CERT_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9])(\d{5,6})(?![A-Za-z0-9])")
_CERT_SD_PAIR_RE = re.compile(r"(?i)(\d{5,6})\s+SD\s+(\d{5,6})")

def _pad_cert_digits(digits: str) -> str: return digits.zfill(CERT_DIGIT_LEN)

def detect_polis_cert_groups(text: str):
    if not text: return None

    tokens = re.findall(r"(?i)\d+|SD", text)
    groups, current_polis, current_certs, current_base_cert, pending_sd = [], None, [], None, False

    for tok in tokens:
        if tok.upper() == "SD":
            pending_sd = True; continue

        n = len(tok)
        if 10 <= n <= 13:
            if current_polis is not None: groups.append((current_polis, current_certs))
            current_polis, current_certs, current_base_cert, pending_sd = tok, [], None, False
            continue

        if current_polis is None: continue

        if pending_sd:
            if current_base_cert is None: pending_sd = False; continue
            endpoint = _pad_cert_digits(tok) if n in (5, 6) else (current_base_cert[:-n] + tok if 1 <= n <= 4 else None)
            
            if endpoint is not None:
                try: start_n, end_n = int(current_base_cert), int(endpoint)
                except ValueError: start_n = end_n = None
                
                if start_n is not None and start_n <= end_n:
                    for x in range(start_n, end_n + 1):
                        c = _pad_cert_digits(str(x))
                        if c not in current_certs: current_certs.append(c)
                current_base_cert = endpoint
            pending_sd = False
        elif n in (5, 6):
            padded = _pad_cert_digits(tok)
            current_certs.append(padded)
            current_base_cert = padded
        elif 1 <= n <= 4 and current_base_cert is not None:
            rebuilt = current_base_cert[:-n] + tok
            current_certs.append(rebuilt)
            current_base_cert = rebuilt

    if current_polis is not None: groups.append((current_polis, current_certs))
    return groups if groups and any(certs for _, certs in groups) else None

def _extract_cert_numbers_from_raw(text: str):
    if not text: return []
    found, consumed_spans = [], []

    for m in _CERT_SD_PAIR_RE.finditer(text):
        if not re.search(r"[A-Za-z0-9]", text[:m.start(1)]): continue
        try: start_n, end_n = int(m.group(1)), int(m.group(2))
        except ValueError: continue
        
        if start_n <= end_n and (end_n - start_n + 1) <= MAX_BREAKDOWN_CODES:
            for n in range(start_n, end_n + 1): found.append(_pad_cert_digits(str(n)))
            consumed_spans.append((m.start(), m.end()))

    scan_text = text
    if consumed_spans:
        chars = list(text)
        for s, e in consumed_spans:
            for i in range(s, e): chars[i] = "#"
        scan_text = "".join(chars)

    for m in _CERT_TOKEN_RE.finditer(scan_text):
        if re.search(r"[A-Za-z0-9]", scan_text[:m.start()]):
            found.append(_pad_cert_digits(m.group(1)))

    return found

def _single_certificate_from_clsdt(clsdt_sertf_value) -> str:
    if clsdt_sertf_value is None or (isinstance(clsdt_sertf_value, float) and pd.isna(clsdt_sertf_value)): return ""
    text = str(clsdt_sertf_value).strip()
    if not text or text.lower() in ("nan", "none", "-"): return ""

    text = _CERT_SD_NORMALIZE_RE.sub("SD", text)
    _, is_polis_pattern = _clean_split_code_impl(text)
    if is_polis_pattern: return ""
    if re.fullmatch(r"\d{5,6}", text): return _pad_cert_digits(text)

    candidates = _extract_cert_numbers_from_raw(text)
    return candidates[0] if candidates else ""

def _join_certificates(cert_list: list) -> str:
    if not cert_list: return ""
    
    seen = []
    for c in cert_list:
        if c not in seen: seen.append(c)
    if len(seen) == 1: return seen[0]

    try: nums = [int(c) for c in seen]
    except ValueError: return ", ".join(seen)

    segments, current_seg = [], [seen[0]]
    for i in range(1, len(seen)):
        if nums[i] - nums[i - 1] == 1: current_seg.append(seen[i])
        else:
            segments.append(current_seg)
            current_seg = [seen[i]]
    segments.append(current_seg)

    parts = []
    for seg in segments:
        if len(seg) > CERT_RANGE_JOIN_THRESHOLD: parts.append(f"{seg[0]} SD {seg[-1]}")
        else: parts.extend(seg)
    return ", ".join(parts)

def build_certificates_for_row(clsdt_sertf_value, fac_policy_value, num_polis: int):
    slots = [""] * max(num_polis, 1)
    if clsdt_cert := _single_certificate_from_clsdt(clsdt_sertf_value):
        slots[0] = clsdt_cert
        return slots

    if _is_blank_value(fac_policy_value): return slots
    fac_text = _CERT_SD_NORMALIZE_RE.sub("SD", str(fac_policy_value).strip())

    if groups := detect_polis_cert_groups(fac_text):
        for i, (_, certs) in enumerate(groups):
            if i < len(slots): slots[i] = _join_certificates(certs)
        return slots

    if single_candidates := _extract_cert_numbers_from_raw(fac_text):
        slots[0] = _join_certificates(single_candidates)

    return slots

# ==============================================================================
# 3. PEMILIHAN SUMBER TERBAIK: CLSDT vs FAC (POLIS & SLIP)
# ==============================================================================
_PENYELESAIAN_RE = re.compile(r"(?i)\bPENYELESAIAN\s+(?:HUTANG|UTANG)\s+PIUTANG\b")
_TRAILING_CURRENCY_WORD_RE = re.compile(r"(?i)\s+(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\s*$")

def _is_blank_value(value) -> bool:
    if value is None or (isinstance(value, float) and pd.isna(value)): return True
    return str(value).strip().lower() in ("", "nan", "none")

def _is_valid_clean_result(codes: list) -> bool:
    return all(re.fullmatch(r"\d{9,13}", c) for c in codes) if codes else False

def _pick_best_source_codes(clsdt_value, fac_value):
    clsdt_codes = clean_split_code(str(clsdt_value)) if not _is_blank_value(clsdt_value) else []
    if _is_valid_clean_result(clsdt_codes): return clsdt_codes, "CLSDT"
    fac_codes = clean_split_code(str(fac_value)) if not _is_blank_value(fac_value) else []
    return (fac_codes, "FAC") if fac_codes else (clsdt_codes, "CLSDT")

def _resolve_polis_or_slip(clsdt_value, fac_value):
    fac_str = "" if _is_blank_value(fac_value) else str(fac_value).strip()
    if _PENYELESAIAN_RE.search(fac_str):
        clsdt_codes = clean_split_code(str(clsdt_value)) if not _is_blank_value(clsdt_value) else []
        if _is_valid_clean_result(clsdt_codes):
            keterangan = re.sub(r"\s{2,}", " ", _TRAILING_CURRENCY_WORD_RE.sub("", fac_str).strip())
            return [f"{code} {keterangan}" for code in clsdt_codes], False
        fac_codes = clean_split_code(fac_str)
        return fac_codes, _is_valid_clean_result(fac_codes)

    codes, _ = _pick_best_source_codes(clsdt_value, fac_value)
    return codes, _is_valid_clean_result(codes)

# ==============================================================================
# 4. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw: pd.DataFrame, keyword: str, cedant_column: str = "") -> pd.DataFrame:
    if cedant_column:
        if cedant_column not in df_raw.columns: raise ValueError(f"Kolom {cedant_column} tidak ada.")
        cedant_cols = [cedant_column]
    else:
        if not (cedant_cols := [c for c in df_raw.columns if "CEDANT" in str(c).upper()]): raise ValueError("Kolom CEDANT tidak ditemukan lewat auto-detect.")

    mask = pd.Series(False, index=df_raw.index)
    for col in cedant_cols: mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)
    
    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant '{keyword}': {len(df_filtered)} dari total {len(df_raw)} baris.")
    return df_filtered

# ==============================================================================
# 5. BUILD OUTPUT
# ==============================================================================
def _detect_key_columns(df: pd.DataFrame, insured_col: str, polis_col: str, slip_col: str):
    def _res(col: str, kw_or_fn, lbl: str):
        if col:
            if col not in df.columns: raise ValueError(f"{lbl}='{col}' tidak ada.")
            return col
        return next((c for c in df.columns if (kw_or_fn(str(c).upper()) if callable(kw_or_fn) else kw_or_fn in str(c).upper())), None)

    col_ins = _res(insured_col, "INSURED", "INSURED_COLUMN")
    col_pol = _res(polis_col, lambda c: c.strip() == "POLIS", "POLIS_COLUMN")
    col_slp = _res(slip_col, "SLIP", "SLIP_COLUMN")
    
    if None in (col_ins, col_pol, col_slp): raise ValueError("Kolom kunci tidak lengkap.")
    print(f"[INFO] Kolom terdeteksi -> INSURED='{col_ins}', POLIS='{col_pol}', SLIP='{col_slp}'")
    return col_ins, col_pol, col_slp

_NON_ALNUM_RE = re.compile(r"[^A-Za-z0-9]")
def _compact_list(values: list) -> list: return [_NON_ALNUM_RE.sub("", v) if isinstance(v, str) else v for v in values]

def _breakdown_all_rows(df: pd.DataFrame, col_ins: str, col_pol: str, col_slp: str, col_clsdt_pol: str = None, col_clsdt_slp: str = None, col_clsdt_sertf: str = None):
    ins_arr, pol_arr, slp_arr = df[col_ins].to_numpy(), df[col_pol].to_numpy(), df[col_slp].to_numpy()
    clsdt_pol_arr = df[col_clsdt_pol].to_numpy() if col_clsdt_pol and col_clsdt_pol in df.columns else [None] * len(df)
    clsdt_slp_arr = df[col_clsdt_slp].to_numpy() if col_clsdt_slp and col_clsdt_slp in df.columns else [None] * len(df)
    clsdt_sertf_arr = df[col_clsdt_sertf].to_numpy() if col_clsdt_sertf and col_clsdt_sertf in df.columns else [None] * len(df)

    ins_all, pol_all, slp_all, cert_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val, clsdt_pol_val, clsdt_slp_val, clsdt_sertf_val in zip(ins_arr, pol_arr, slp_arr, clsdt_pol_arr, clsdt_slp_arr, clsdt_sertf_arr):
        ins = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []
        raw_pol, raw_slp = pol_val if pd.notna(pol_val) else "", slp_val if pd.notna(slp_val) else ""

        clsdt_pol_is_usable = col_clsdt_pol and pd.notna(clsdt_pol_val) and _is_valid_clean_result(clean_split_code(str(clsdt_pol_val)) if not _is_blank_value(clsdt_pol_val) else [])
        polis_cert_groups = detect_polis_cert_groups(_CERT_SD_NORMALIZE_RE.sub("SD", str(raw_pol).strip())) if not clsdt_pol_is_usable and not _is_blank_value(raw_pol) else None

        if polis_cert_groups:
            pol = [p for p, _ in polis_cert_groups]
            pol_should_compact = True
            cert = [_join_certificates(certs) for _, certs in polis_cert_groups]
        else:
            if col_clsdt_pol and pd.notna(clsdt_pol_val):
                pol, pol_should_compact = _resolve_polis_or_slip(clsdt_pol_val, raw_pol)
            else:
                pol = clean_split_code(str(raw_pol).strip()) if not _is_blank_value(raw_pol) else []
                pol_should_compact = _is_valid_clean_result(pol)
            cert = None

        if col_clsdt_slp and pd.notna(clsdt_slp_val): slp, slp_should_compact = _resolve_polis_or_slip(clsdt_slp_val, raw_slp)
        else:
            slp = clean_split_code(str(raw_slp).strip()) if not _is_blank_value(raw_slp) else []
            slp_should_compact = _is_valid_clean_result(slp)

        if pol_should_compact: pol = _compact_list(pol)
        if slp_should_compact: slp = _compact_list(slp)

        if cert is None: cert = build_certificates_for_row(clsdt_sertf_val, raw_pol, num_polis=len(pol) if pol else 1)
        if len(cert) < len(pol): cert = cert + [""] * (len(pol) - len(cert))
        elif len(cert) > len(pol): cert = cert[: len(pol)] if pol else cert

        ins_all.append(ins)
        pol_all.append(pol)
        slp_all.append(slp)
        cert_all.append(cert)
        
        max_ins, max_pol, max_slp = max(max_ins, len(ins)), max(max_pol, len(pol)), max(max_slp, len(slp))
        
    return ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_ins, col_pol, col_slp = _detect_key_columns(df_filtered, INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN)
    col_clsdt_pol = CLSDT_POLIS_COLUMN if CLSDT_POLIS_COLUMN in df_filtered.columns else None
    col_clsdt_slp = CLSDT_SLIP_COLUMN if CLSDT_SLIP_COLUMN in df_filtered.columns else None
    col_clsdt_sertf = CLSDT_SERTF_COLUMN if CLSDT_SERTF_COLUMN in df_filtered.columns else None

    ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp = _breakdown_all_rows(df_filtered, col_ins, col_pol, col_slp, col_clsdt_pol, col_clsdt_slp, col_clsdt_sertf)

    df_filtered = df_filtered.reset_index(drop=True)
    output_cols_data = {col: df_filtered[col].values for col in df_filtered.columns}

    for col in df_filtered.columns:
        if col == col_ins:
            for i in range(1, max_ins + 1): output_cols_data[f"INSURED_CLEAN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in ins_all]
        if col == col_pol:
            for i in range(1, max_pol + 1):
                output_cols_data[f"POLICY_CLEAN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in pol_all]
                output_cols_data[f"{CERTIFICATE_OUTPUT_COL}_CLEAN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_all]
        if col == col_slp:
            for i in range(1, max_slp + 1): output_cols_data[f"SLIP_CLEAN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slp_all]

    df_master = pd.DataFrame(output_cols_data)

    # Hanya mengatur urutan kolom output; tidak mengubah proses cleansing.
    ordered_cols = []
    for col in df_filtered.columns:
        ordered_cols.append(col)
        if col == col_ins:
            ordered_cols.extend(
                c for c in (f"INSURED_CLEAN_{i}" for i in range(1, max_ins + 1))
                if c in df_master.columns
            )
        elif col == col_pol:
            for i in range(1, max_pol + 1):
                for c in (f"POLICY_CLEAN_{i}", f"{CERTIFICATE_OUTPUT_COL}_CLEAN_{i}"):
                    if c in df_master.columns:
                        ordered_cols.append(c)
        elif col == col_slp:
            ordered_cols.extend(
                c for c in (f"SLIP_CLEAN_{i}" for i in range(1, max_slp + 1))
                if c in df_master.columns
            )

    ordered_cols.extend(c for c in df_master.columns if c not in ordered_cols)
    df_master = df_master[ordered_cols]

    # OUTPUT ONLY: nama, jumlah, dan urutan kolom mengikuti template OSBAL terbaru.
    visible_cols = [
        "CCOS_DOC_NO", "CCOS_DATE", "CCOS_REF_CODE", "CCOS_COMP", "CCOS_COMP_NAME",
        "CCOS_REF_COMP", "CCOS_REF_COMP_NAME", "FAC_INSURED",
        "INSURED_CLEAN_1", "INSURED_CLEAN_2", "INSURED_CLEAN_3", "INSURED_CLEAN_4", "INSURED_CLEAN_5",
        "CCOS_CURR", "CCOS_OR_BAL", "CCOS_BAL_DUE", "CCOS_OR_BAL_IN_IDR", "CCOS_BAL_DUE_IN_IDR",
        "FAC_COM_DATE", "FAC_EXP_DATE", "FAC_DUE_DATES", "FAC_SUB_CLASS", "FAC_POLICY_NO",
        "POLICY_CLEAN_1", "SERTIF_CLEAN_1", "POLICY_CLEAN_2", "SERTIF_CLEAN_2",
        "POLICY_CLEAN_3", "SERTIF_CLEAN_3", "POLICY_CLEAN_4", "POLICY_CLEAN_5",
        "FAC_SLIP", "SLIP_CLEAN_1", "SLIP_CLEAN_2", "SLIP_CLEAN_3", "SLIP_CLEAN_4", "SLIP_CLEAN_5",
        "CLASS_CODE", "CLASS_NAME", "CLSDT_POLICY_NO", "CLSDT_SLIP_NO", "CLSDT_SERTF_NO",
    ]
    for c in visible_cols:
        if c not in df_master.columns:
            df_master[c] = ""
    return df_master[visible_cols]

# ==============================================================================
# 6. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook = load_workbook(output_path)
    worksheet = workbook.active
    text_col_indices = [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None: cell.value = str(cell.value)
    workbook.save(output_path)

# ==============================================================================
# 7. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print(f"{'='*70}\n CLEANSING DATA 2 - OSBAL | CEDANT: PT LIPPO GENERAL INSURANCE\n{'='*70}")
    
    input_path = INPUT_FILE
    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
        if not candidates: raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")
        input_path = candidates[0]
        print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    df_lippo = filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)

    if not df_lippo.empty:
        df_hasil = build_output(df_lippo)
        save_with_text_format(df_hasil, OUTPUT_FILE)
        print(f"{'='*70}\n[SUCCESS] Selesai. Output: '{OUTPUT_FILE}' ({len(df_hasil)} baris)\n{'='*70}")
    else:
        print("[WARNING] Tidak ada baris yang cocok dengan filter.")

if __name__ == "__main__":
    main()


"""
================================================================================
 SCRIPT CLEANSING DATA 1 - FACULTATIVE (KHUSUS CEDANT: PT LIPPO GENERAL INSURANCE)
================================================================================
Dipakai untuk breakdown & cleansing kolom INSURED, POLIS, dan SLIP NO supaya
bisa dipakai sebagai key matching antar database.
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_DIR = os.path.join(SCRIPT_DIR, "input")
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "output")

INPUT_FILE = os.path.join(INPUT_DIR, "1b. Transaksi Facul 01.01.23 - 17.08.2026.xlsx")
SHEET_NAME = "Query result"
HEADER_ROW = 0
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "lippo_output_facul_V3.xlsx")
CEDANT_FILTER = "LIPPO"

CEDANT_COLUMN = "COMP_NAME"
BROKER_NAME_COLUMN = "COMP_NAME.1"
DIRECT_MARKER = "DIRECT"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"

MAX_BREAKDOWN_CODES = 5
TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CERTIFICATE",
)

RE_I = re.IGNORECASE

# ==============================================================================
# PRE-COMPILED REGEXES & CONSTANTS (OPTIMIZATION)
# ==============================================================================
_TRANSACTION_HEADER_RE = re.compile(r"^\s*(?:Pembayaran|Penagihan)?\s*(?:dan\s+Penagihan)?\s*Premi\s*(?:Fakultatif|SOA|Pensesian)?\s*(?:PAR\s*&?\s*EQ|BIT|GIT|CECR|QS)?\s*", RE_I)
_TRUNCATE_TRIGGER_RE = re.compile(r"\bsubsidiar|\b(?:and\s*/\s*or|and|its|their|all)\s+associat(?!ion)|&\s*/\s*or\s+associat(?!ion)|\baffiliat|\bfiliated\b|related\s+compan|respective\s+right|\bincluding\s+an[yd]|\bindu?ding\s+any|as\s+the\s+owner|as\s+owners?\b|as\s+main\s+contractor|\(as\s+mortgagee\s+and\s+loss\s+payee\)", RE_I)
_REMOVE_ONLY_RE = re.compile(r"\bas\s+principal\b|\bsebagai\s+principal\b|\bsebagai\s+kontraktor\b", RE_I)
_GENERIC_BOILERPLATE_TAIL_RE = re.compile(r"\bsub\s*-?\s*kontraktor\b.*$|\banak\s+perusahaan\b.*$|\bafiliasi\s+perusahaan\b.*$|\bdan\s+semua\b\s*[–-]?\s*$", RE_I)
_TITLE_RE = re.compile(r"\b(?:S\.?H\.?|S\.?I\.?Kom\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b", RE_I)
_LEGAL_ENTITY_RE = re.compile(r"\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PP\.?|PERSERO|Perseroan\s+Terbatas|LTD\.?|PTE\.?)\b", RE_I)
_ENTITY_SPLIT_RE = re.compile(r"\b(?:QQ|Q\.Q\.)\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|\bin\s+this\s+case\b|\bdalam\s+hal\s+ini\b|;|/|,(?!\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO|LTD\.?|PTE\.?)\b)", RE_I)
_INSURED_VARIOUS_RE = re.compile(r"\bVARIOUS\b", RE_I)
_INSURED_BATCH_NOISE_RE = re.compile(
    r"\b(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|"
    r"JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|"
    r"NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)?\s*\bBATCH\.?\s*\d*"
    r"(?:\s*(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|"
    r"JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|"
    r"NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)\b)?"
    r"|\d+(?=\s*(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|"
    r"JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|"
    r"NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)\b)", RE_I)
_STRAY_SYMBOL_RE = re.compile(r"[¿¡‽]")
_PAREN_GROUP_RE = re.compile(r"\([^()]*\)")
_PAREN_TOKEN_RE = re.compile(r"\uE100(\d+)\uE101")
_LEGAL_ENTITY_DASH_SIGNAL_RE = re.compile(r",\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO)\s*[\r\n]*\s*-\s*", RE_I)
_GENERIC_DASH_SEPARATOR_RE = re.compile(r"\s*[\r\n]+\s*-\s*|\s+-\s+")
_INSURED_DESCRIPTIVE_RE = re.compile(r"\bVARIOUS\s+INSUREDS?\b|\bACCEPTED\s+BY\b", RE_I)
_SUSPENSE_KEEP_AS_IS_RE = re.compile(r"\bsuspense\b|\bpenyelesaian\s+suspense\b|\bpenyelesaian\s+(?:h)?utang\s+piutang\b", RE_I)
_DANGLING_PATTERN_RE = re.compile(r"^\b(?:AS|AN|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*|\s*\b(?:AS|AN|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$", RE_I)

_LIPPO_KNOWN_ENTITIES = ["MATAHARI PUTRA PRIMA", "MATAHARI BOSTON DRIGSTORE", "MATAHARI PUSAKA TAMA", "MPP LIPPO GROUP", "LIPPO GROUP"]
_LIPPO_KNOWN_ENTITIES_SORTED = sorted(_LIPPO_KNOWN_ENTITIES, key=len, reverse=True)
_KNOWN_ENTITY_RE = re.compile("|".join(re.escape(name) for name in _LIPPO_KNOWN_ENTITIES_SORTED))
_ENTITY_SLASH_NORMALIZE_PATTERNS = [(name, re.compile(r"\b" + r"(?:\s*/\s*|\s+)".join(re.escape(w) for w in name.split()) + r"\b", RE_I)) for name in _LIPPO_KNOWN_ENTITIES_SORTED if len(name.split()) > 1]
_INSURED_MERGE_OVERRIDES = {"APARTEMEN EKSEKUTIF MENTENG,PERHIM.PENGH": ["APARTEMEN EKSEKUTIF MENTENG PERHIM PENGH"]}

_MAIN_LEN_MIN, _MAIN_LEN_MAX = 11, 16
_MAIN_DIGITS_PATTERN = rf"\d{{{_MAIN_LEN_MIN},{_MAIN_LEN_MAX}}}"
_TOK_MAIN_RE = re.compile(rf"^{_MAIN_DIGITS_PATTERN}$")
_TOK_INDEPENDENT_RE = re.compile(r"^\d{7,16}$")
_TOK_CERT_RE = re.compile(r"^\d{6}$")
_TOK_SUFFIX_REPLACE_RE = re.compile(r"^\d{1,11}$")
_TOK_PLACEHOLDER_RE = re.compile(r"^P\d{1,2}$", RE_I)
_TOK_STATUS_RE = re.compile(r"^(?:New|End(?:orsement)?|Cancel\s*All|Adjustment|Cancellation|Reinstatement)$", RE_I)
_VARIOUS_WORD_RE = re.compile(r"\bVAR(?:I(?:O(?:U(?:S)?)?)?)?\b", RE_I)
_SD_MARKER_RE = re.compile(r"\bS\s*/\s*D\b|\bSD\b", RE_I)
_PURE_DECIMAL_RE = re.compile(r"^\d+(?:\.\d+)+$")
_DASH_AMP_AMBIGUOUS_RE = re.compile(rf"{_MAIN_DIGITS_PATTERN}\s*-\s*\d+\s*&\s*\d+")
_ALPHA_WORD_RE = re.compile(r"[A-Za-z][A-Za-z0-9]*")
_ALLOWED_ALPHA_TOKENS = {"NEW", "END", "ENDORSEMENT", "CANCEL", "ALL", "ADJUSTMENT", "CANCELLATION", "REINSTATEMENT", "VARIOUS", "VAR", "SD", "S", "D"}

_ATTACHMENT_NOISE_RE = re.compile(r"\bSEE\s+ATTACHMENT\b|\bATTACHMENT\b", RE_I)
_MONTH_YEAR_NOISE_RE = re.compile(
    r"\b(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|"
    r"JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:T(?:E(?:M(?:B(?:E?R)?)?)?)?)?|SEPTEMEBR|OKT(?:OBER)?|OCT(?:OBER)?|"
    r"NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)\s*\.?\s*\d{2,4}\b", RE_I)
_CURRENCY_NOISE_RE = re.compile(r"\b(?:USD|IDR|SGD|EUR|CNY|SGD|GBP|JPY|AUD|HKD)\b", RE_I)
_UNIT_NOISE_RE = re.compile(r"\d*\s*UNIT\b", RE_I)
_STRAY_QUOTE_NOISE_RE = re.compile(r"['\u2018\u2019\u201c\u201d]")
_ADMIN_LABEL_NOISE_RE = re.compile(
    r"\bDEKL\.?\b|\b\d+\s*SLIP\s+DIJADIKAN\s+\d+\b|\bID\s*NO\.?\s*[A-Z0-9]*\b"
    r"|(?:\b\d{1,6}\s*/?\s*)?\bSHIPMENT\s*\d*\b|(?:\b\d{1,6}\s*/?\s*)?\bBORDERO\b"
    r"|(?:\b\d{1,6}\s*/?\s*)?\bPENDING\s+SLIP\b|(?:\b\d{1,6}\s*/?\s*)?\bBATCH\s*\d*\b", RE_I)
_DOC_REF_DATE_NOISE_RE = re.compile(r"(?:\b\d{1,6}\s*/\s*)?\b(?:CN|DN)\s*/\s*\d{1,2}\s*/\s*\d{1,2}\s*/\s*\d{2,4}(?:\s*/\s*\d{1,6}\b)?", RE_I)
_ENDORSE_HASH_NOISE_RE = re.compile(r"\bE\s*#\s*\d+\b", RE_I)
_HASH_CODE_NOISE_RE = re.compile(r"#\s*[A-Z]{1,3}\d{0,3}\b", RE_I)
_TREATY_REF_CODE_NOISE_RE = re.compile(
    rf"\b\d{{1,6}}\s*/\s*[A-Z]+-[A-Z]+(?:\s*/\s*(?:X{{0,3}}(?:IX|IV|V?I{{0,3}}))\b)?(?:\s*/\s*\d{{4}}\b)?"
    rf"|\b[A-Z]+-[A-Z]+(?:\s*/\s*(?:X{{0,3}}(?:IX|IV|V?I{{0,3}}))\b)?(?:\s*/\s*\d{{4}}\b)?", RE_I)
_MAIN_SPACE_SUFFIX_RE = re.compile(rf"^({_MAIN_DIGITS_PATTERN})\s+(\d{{1,6}})\s*(?:/\s*VAR[A-Z]*\s*)?$", RE_I)
_TBA_NAME_ONLY_RE = re.compile(r"^TBA\s*/\s*([A-Za-z][A-Za-z ,.]*[A-Za-z.])$", RE_I)
_TBA_RESULT_PT_SUFFIX_RE = re.compile(r"\s*,?\s*P\.?T\.?$", RE_I)

_PP_MAIN_TOKEN_RE = re.compile(rf"^{_MAIN_DIGITS_PATTERN}$")
_PP_NONCERT_SUFFIX_RE = re.compile(r"^\d{1,11}$")
_CERT_SD_NORMALIZE_RE = re.compile(r"S\s*/\s*D", RE_I)

_POLIS_SLIP_OVERRIDES = {
    "APRIL 2026 - 0015/LI-RBU/V/2026 / 1112302600206,1112302600215,1112302600218,1112302600221.": [
        "1112302600206", "1112302600215", "1112302600218", "1112302600221",
    ],
}

_POLICY_CERT_HARDCASE_FALLBACK = {
    "1303121900052-182 S/D 189 / 1303152000009 01S/D10",
    "1303121900052-198S/D213 / 1303152000009-19S/D35",
    "1303121900052 -191 S/D197 / 1303152000009 -22S/D30",
    "1303121900052-190/'1303152000009-11S/D21",
    "1903021900688-013S/D030 1903021900689-622S/D643",
    "1903021900688-015s/d050 /1903021900689-644s/d698",
    "1903021900688-001S/D0012 / 1903021900689-596S/D620",
    "120110180002-001-006-003-003-005004-006",
    "15011011004041-018135-026215-003072-017133-019134",
    "15011015003-029002028002010003011",
    "150110150026-215-004041018135019134017133003072",
}

_POLICY_CERT_HARDCASE_RESULT = {
    "1102221600017,1102281000725,1202281700127,13022217": (
        ["1102221600017", "1102281000725", "1302221700127"], ["", "", ""]
    ),
    "1102211500160,1102281800059,1102281900090,12022119": (
        ["1102211500160", "1102281800059", "1202211900090"], ["", "", ""]
    ),
    "1801092400272-1801052400198-1801092400277-18010524": (
        ["1801092400272", "1801052400198", "1801052400277"], ["", "", ""]
    ),
    "15011011007-044005042-003040-": (
        ["15011011007", "15044005042", "15044003040"], ["", "", ""]
    ),
}

_PAREN_STASH = {}


# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
def _truncate_from_first_trigger(text: str) -> str:
    return text[: match.start()] if (match := _TRUNCATE_TRIGGER_RE.search(text)) else text

def _strip_insured_batch_noise(text: str) -> str: return _INSURED_BATCH_NOISE_RE.sub(" ", text)
def _strip_stray_symbols(text: str) -> str: return _STRAY_SYMBOL_RE.sub(" ", text)
def _normalize_dash_entity_list(text: str) -> str: return _GENERIC_DASH_SEPARATOR_RE.sub(", ", text) if _LEGAL_ENTITY_DASH_SIGNAL_RE.search(text) else text
def _strip_transaction_header(text: str) -> str: return _TRANSACTION_HEADER_RE.sub("", text)
def _strip_titles(text: str) -> str: return _TITLE_RE.sub(" ", text)
def _strip_legal_entity(text: str) -> str: return _LEGAL_ENTITY_RE.sub(" ", text)

def _protect_parens_content(text: str) -> str:
    def _mask(match):
        token = f"\uE100{len(_PAREN_STASH)}\uE101"
        _PAREN_STASH[token] = match.group(0)
        return token
    return _PAREN_GROUP_RE.sub(_mask, text)

def _restore_protected_chars(text: str) -> str:
    return text.replace("\uE0F0", "/").replace("\uE0F1", ",").replace("\uE0F2", ";")

def _restore_protected_parens(text: str) -> str:
    def _unmask(match): return _PAREN_STASH.get(f"\uE100{match.group(1)}\uE101", match.group(0))
    return _PAREN_TOKEN_RE.sub(_unmask, text)

def _truncate_boilerplate_tail(text: str) -> str:
    text = _truncate_from_first_trigger(text)
    if match_generic := _GENERIC_BOILERPLATE_TAIL_RE.search(text): text = text[: match_generic.start()]
    return _REMOVE_ONLY_RE.sub(" ", text)

def _final_polish(text: str) -> str:
    t = re.sub(r"[-/]", " ", re.sub(r"[,;:\"'.]", " ", text))
    t = re.sub(r"\s+", " ", re.sub(r"\s+\)", ")", re.sub(r"\(\s+", "(", t))).strip()
    previous = None
    while previous != t:
        previous, t = t, _DANGLING_PATTERN_RE.sub("", t).strip()
    return t.upper()

def _normalize_entity_slash_variants(text: str) -> str:
    for canonical, pattern in _ENTITY_SLASH_NORMALIZE_PATTERNS: text = pattern.sub(canonical, text)
    return text

def _split_by_known_entities(text: str):
    pos, n, found = 0, len(text), []
    while pos < n:
        while pos < n and text[pos] == " ": pos += 1
        if pos >= n: break
        if not (match := _KNOWN_ENTITY_RE.match(text, pos)): return None
        found.append(match.group(0))
        pos = match.end()
    return found if len(found) > 1 else None

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return []
    original = text.strip()

    if original in _INSURED_MERGE_OVERRIDES: return _INSURED_MERGE_OVERRIDES[original]
    if _INSURED_DESCRIPTIVE_RE.search(original) or _SUSPENSE_KEEP_AS_IS_RE.search(original): return [original]

    _PAREN_STASH.clear()
    t = _protect_parens_content(_truncate_boilerplate_tail(_normalize_entity_slash_variants(
        _normalize_dash_entity_list(_INSURED_VARIOUS_RE.sub(" ", _strip_insured_batch_noise(
            _strip_stray_symbols(_strip_titles(_strip_transaction_header(original)))
        )))
    )))

    results = []
    for part in _ENTITY_SPLIT_RE.split(t):
        cleaned = re.sub(r"\s+", " ", _restore_protected_parens(_final_polish(_strip_legal_entity(_truncate_boilerplate_tail(_restore_protected_chars(part)))))).strip()
        if not cleaned or len(cleaned) < 2 or not re.search(r"[A-Z]", cleaned): continue

        if known_split := _split_by_known_entities(cleaned):
            for name in known_split:
                if name not in results: results.append(name)
            continue
        if cleaned not in results: results.append(cleaned)

    return [original] if len(results) > MAX_BREAKDOWN_CODES else results


# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO (FACULTATIVE)
# ==============================================================================
def _strip_known_noise_phrases(text: str) -> str:
    text = _STRAY_QUOTE_NOISE_RE.sub(" ", text)
    text = _PAREN_GROUP_RE.sub(" ", text)
    text = _ATTACHMENT_NOISE_RE.sub(" ", text)
    text = _DOC_REF_DATE_NOISE_RE.sub(" ", text)
    text = _TREATY_REF_CODE_NOISE_RE.sub(" ", text)
    text = _MONTH_YEAR_NOISE_RE.sub(" ", text)
    text = _CURRENCY_NOISE_RE.sub(" ", text)
    text = _UNIT_NOISE_RE.sub(" ", text)
    text = _ADMIN_LABEL_NOISE_RE.sub(" ", text)
    text = _ENDORSE_HASH_NOISE_RE.sub(" ", text)
    text = _HASH_CODE_NOISE_RE.sub(" ", text)
    return re.sub(r"\s+", " ", re.sub(r"[/\-]\s*(?=[/\-]|$)", " ", text)).strip()

def _strip_tba_prefix_if_name(text: str) -> str:
    return _TBA_RESULT_PT_SUFFIX_RE.sub("", match.group(1)).strip() if (match := _TBA_NAME_ONLY_RE.match(text)) else text

def _has_descriptive_narrative(text: str) -> bool:
    return any(not (_TOK_PLACEHOLDER_RE.match(word) or word.upper() in _ALLOWED_ALPHA_TOKENS) for word in _ALPHA_WORD_RE.findall(text))

def _strip_various(text: str) -> str: return _VARIOUS_WORD_RE.sub(" ", text)

def _expand_numeric_range(start_str: str, end_str: str, max_count: int = MAX_BREAKDOWN_CODES):
    if len(start_str) != len(end_str) or not (start_str.isdigit() and end_str.isdigit()): return None
    diff_len = next((i for i in range(1, len(start_str) + 1) if start_str[-i] != end_str[-i]), 0)
    if diff_len == 0 or start_str[:-diff_len] != end_str[:-diff_len]: return None
    try: start_n, end_n = int(start_str[-diff_len:]), int(end_str[-diff_len:])
    except ValueError: return None
    if start_n > end_n or (end_n - start_n + 1) > max_count: return None
    return [f"{start_str[:-diff_len]}{str(n).zfill(diff_len)}" for n in range(start_n, end_n + 1)]

def _expand_suffix_range(anchor: str, suffix_a: str, suffix_b: str):
    width = max(len(suffix_a), len(suffix_b))
    if width >= len(anchor) or not (suffix_a.isdigit() and suffix_b.isdigit()): return None
    start_n, end_n = int(suffix_a), int(suffix_b)
    return [f"{anchor[:-width]}{str(n).zfill(width)}" for n in range(start_n, end_n + 1)] if start_n <= end_n and (end_n - start_n + 1) <= MAX_BREAKDOWN_CODES else None

def _split_sd_range(text: str):
    if not (markers := list(_SD_MARKER_RE.finditer(text))) or len(markers) > 1: return "__FALLBACK__" if markers else None
    runs_left, runs_right = re.findall(r"\d+", text[: markers[0].start()]), re.findall(r"\d+", text[markers[0].end():])
    if not runs_left or not runs_right: return "__FALLBACK__"
    last_left, first_right = runs_left[-1], runs_right[0]
    
    if not (last_left_is_main := _TOK_MAIN_RE.match(last_left)) and len(runs_left) >= 2 and _TOK_MAIN_RE.match(runs_left[-2]):
        expanded = _expand_suffix_range(runs_left[-2], last_left, first_right)
    elif last_left_is_main and _TOK_MAIN_RE.match(first_right):
        expanded = _expand_numeric_range(last_left, first_right)
    else: return "__FALLBACK__"
    return expanded if expanded is not None else "__FALLBACK__"

def clean_split_code(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return []
    original = text.strip()

    if _SUSPENSE_KEEP_AS_IS_RE.search(original): return [original]
    if (tba_stripped := _strip_tba_prefix_if_name(original)) != original: return [tba_stripped]
    if space_concat_match := _MAIN_SPACE_SUFFIX_RE.match(original): return [space_concat_match.group(1) + space_concat_match.group(2)]
    if original in _POLIS_SLIP_OVERRIDES: return _POLIS_SLIP_OVERRIDES[original]

    t = _strip_known_noise_phrases(original)
    if _has_descriptive_narrative(t) or _PURE_DECIMAL_RE.match(t) or _DASH_AMP_AMBIGUOUS_RE.search(t): return [original]
    if (sd_result := _split_sd_range(t)) == "__FALLBACK__": return [original]
    if sd_result is not None: return sd_result

    t = re.sub(r"[\u2012\u2013\u2014\u2015]", "-", _strip_various(t))
    _had_slash_context, strong_segments, pending_delim = "/" in t, [], None

    for chunk in [c for c in re.split(r"(\s{2,}|[,+/;&])", t) if c != ""]:
        if re.fullmatch(r"\s{2,}", chunk) or chunk in (",", "+", "/", ";", "&"): pending_delim = "," if re.fullmatch(r"\s{2,}", chunk) else chunk; continue
        if seg := chunk.strip(): strong_segments.append((seg, pending_delim))
        pending_delim = None

    raw_tokens = []
    for seg, delim in strong_segments:
        sub_parts = [p.strip() for p in re.sub(r"[.]", "", seg).strip().split("-") if p.strip()]
        if len(sub_parts) <= 1:
            if sub_parts: raw_tokens.append((sub_parts[0], delim))
            continue
        buffer, buffer_delim = [sub_parts[0]], delim
        for p in sub_parts[1:]:
            if _TOK_MAIN_RE.match(p) or _TOK_PLACEHOLDER_RE.match(p):
                raw_tokens.append(("-".join(buffer), buffer_delim))
                buffer, buffer_delim = [p], "-"
            else: buffer.append(p)
        raw_tokens.append(("-".join(buffer), buffer_delim))

    tokens = [(tok, d) for tok, d in raw_tokens if not _TOK_PLACEHOLDER_RE.match(tok)]
    if not tokens: return [original]

    codes, current_base, current_base_is_strict, n_tokens = [], None, False, len(tokens)
    for i, (tok, delim) in enumerate(tokens):
        effective_delim = delim if delim is not None else (tokens[i + 1][1] if i + 1 < n_tokens else ("/" if _had_slash_context else None))
        if current_base is not None and current_base_is_strict and delim in ("+", "&", "/", ",") and _TOK_SUFFIX_REPLACE_RE.match(tok) and not _TOK_MAIN_RE.match(tok) and len(tok) < len(current_base):
            if (reconstructed := current_base[: -len(tok)] + tok) not in codes: codes.append(reconstructed)
            continue
        
        parts = tok.split("-")
        anchor, suffix_parts = parts[0], parts[1:]
        if _TOK_MAIN_RE.match(anchor): is_strict = True
        elif not suffix_parts and effective_delim == "/" and _TOK_INDEPENDENT_RE.match(anchor): is_strict = False
        else:
            if anchor.isdigit(): return [original]
            continue

        current_base, current_base_is_strict = anchor, is_strict
        while suffix_parts and _TOK_STATUS_RE.match(suffix_parts[-1]): suffix_parts.pop()
        append_suffix_parts = []

        for p in suffix_parts:
            if _TOK_CERT_RE.match(p): append_suffix_parts.append(p)
            elif _TOK_SUFFIX_REPLACE_RE.match(p) and len(p) < len(anchor):
                if (reconstructed := anchor[: -len(p)] + p) not in codes: codes.append(reconstructed)
        
        if append_suffix_parts and (combined := f"{anchor}-" + "-".join(append_suffix_parts)) not in codes: codes.append(combined)
        if anchor not in codes: codes.append(anchor)

    return [original] if not codes or len(codes) > MAX_BREAKDOWN_CODES else codes


# ==============================================================================
# 2B. PROSES TERPADU POLIS + CERTIFICATE (REVISI V4)
# ==============================================================================
def _cert_pad6(raw_suffix: str) -> str: return raw_suffix[-6:] if len(raw_suffix) > 6 else raw_suffix.zfill(6)
def _is_raw_cert_suffix(raw_suffix: str) -> bool: return len(raw_suffix) in (5, 6) and raw_suffix.isdigit() and raw_suffix[0] == "0"

def _format_cert_list(cert_nums: list, is_sd: bool) -> str:
    if not cert_nums: return ""
    seen = []
    for c in cert_nums:
        if c not in seen: seen.append(c)
    if len(seen) == 1: return seen[0]

    nums_int = [int(c) for c in seen]
    is_sequential = all(nums_int[i + 1] - nums_int[i] == (1 if nums_int[-1] >= nums_int[0] else -1) for i in range(len(nums_int) - 1))
    return f"{seen[0]} SD {seen[-1]}" if is_sd or (is_sequential and len(seen) > 3) else ", ".join(seen)

def process_policy_certificate(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return [], []
    original = text.strip()

    if original in _POLICY_CERT_HARDCASE_FALLBACK or _SUSPENSE_KEEP_AS_IS_RE.search(original): return [""], [original]
    if original in _POLICY_CERT_HARDCASE_RESULT:
        pol, cert = _POLICY_CERT_HARDCASE_RESULT[original]
        return cert, pol
    if (tba_stripped := _strip_tba_prefix_if_name(original)) != original: return [""], [tba_stripped]
    if original in _POLIS_SLIP_OVERRIDES: return [""] * len(_POLIS_SLIP_OVERRIDES[original]), list(_POLIS_SLIP_OVERRIDES[original])

    t = re.sub(r"\bCERTIFICATE\b\s*:?\s*", " ", re.sub(r"\.{2,}", " ", re.sub(r"\bDLL\b\.?", " ", re.sub(r"\bVARIUOS\b", " VARIOUS ", re.sub(r"\bSD\b", " SD ", _CERT_SD_NORMALIZE_RE.sub(" SD ", original), flags=RE_I), flags=RE_I), flags=RE_I)))
    t_noise = _strip_known_noise_phrases(t)
    if _has_descriptive_narrative(t_noise) or _PURE_DECIMAL_RE.match(t_noise): return [""], [original]
    t = re.sub(r"[\u2012\u2013\u2014\u2015]", "-", _strip_various(t_noise))

    if "+" in t and "," not in t:
        plus_chunks = [c.strip() for c in t.split("+") if c.strip()]
        if plus_chunks and _PP_MAIN_TOKEN_RE.match(plus_chunks[0]) and len(plus_chunks[1:]) > 1 and all(re.fullmatch(r"\d+", c) for c in plus_chunks[1:]) and not any(_PP_MAIN_TOKEN_RE.match(c) for c in plus_chunks[1:]): return [""], [original]

    if not (top_chunks := [c.strip() for c in re.split(r"[+/:]", t) if c.strip()]): return [""], [original]

    pairs, current_main, fallback = [], None, False
    for chunk in top_chunks:
        expanded = []
        for p in re.split(r"(\bSD\b|&|-|,)", chunk.strip().strip("/").strip()):
            expanded.append(p) if p in ("SD", "&", "-", ",") else expanded.extend(p.split())
        
        if not (pieces := [p.strip() for p in expanded if p.strip()]): continue

        i, local_main, pending_op, active_idx, cert_target_idx = 0, None, None, None, None
        while i < len(pieces):
            piece = pieces[i]
            if piece in ("-", "&", "SD", ","): pending_op = piece; i += 1; continue
            
            if _PP_MAIN_TOKEN_RE.match(piece):
                local_main, current_main = piece, piece
                if (existing_idx := next((idx for idx, p in enumerate(pairs) if p[0] == piece), None)) is not None: active_idx = existing_idx
                else: pairs.append([piece, [], False]); active_idx = len(pairs) - 1
                cert_target_idx, pending_op = active_idx, None
                i += 1
                continue

            if piece.isdigit():
                anchor = local_main or current_main
                if active_idx is None and pairs: active_idx = len(pairs) - 1
                if cert_target_idx is None and pairs: cert_target_idx = active_idx if active_idx is not None else len(pairs) - 1

                if pending_op == "SD":
                    if cert_target_idx is not None: pairs[cert_target_idx][1].append(_cert_pad6(piece)); pairs[cert_target_idx][2] = True
                    pending_op = None; i += 1; continue

                if pending_op in ("-", "&", ",", None) and anchor:
                    padded, already_has_cert = _cert_pad6(piece), cert_target_idx is not None and bool(pairs[cert_target_idx][1])
                    if _is_raw_cert_suffix(piece) or (len(piece) == 7 and padded[0] == "0") or (already_has_cert and pending_op in ("-", "&", ",")):
                        if cert_target_idx is not None: pairs[cert_target_idx][1].append(padded)
                        pending_op = None; i += 1; continue

                    if pending_op == "-" and anchor and _PP_NONCERT_SUFFIX_RE.match(piece) and len(piece) < len(anchor) and not already_has_cert:
                        reconstructed = anchor[: -len(piece)] + piece
                        if (existing_idx := next((idx for idx, p in enumerate(pairs) if p[0] == reconstructed), None)) is None:
                            pairs.append([reconstructed, [], False]); active_idx = len(pairs) - 1
                        else: active_idx = existing_idx
                        pending_op = None; i += 1; continue

            fallback = True; break
        if fallback: break

    polis_codes = [p[0] for p in pairs]
    if fallback or not polis_codes or len(polis_codes) > MAX_BREAKDOWN_CODES: return [""], [original]
    return [_format_cert_list(cert_nums, is_sd) for _, cert_nums, is_sd in pairs], polis_codes


# ==============================================================================
# 3. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path, sheet_name, header_row): return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw, keyword, cedant_column):
    if cedant_column not in df_raw.columns: raise ValueError(f"CEDANT_COLUMN='{cedant_column}' tidak ditemukan. Tersedia: {list(df_raw.columns)}")
    df_filtered = df_raw[df_raw[cedant_column].astype(str).str.contains(keyword, case=False, na=False)].copy()
    print(f"[INFO] Filter cedant (kolom '{cedant_column}') mengandung '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered


# ==============================================================================
# 4. BUILD OUTPUT
# ==============================================================================
def _breakdown_all_rows(df, col_insured, col_polis, col_slip):
    ins_arr, pol_arr, slp_arr = df[col_insured].to_numpy(), df[col_polis].to_numpy(), df[col_slip].to_numpy()
    insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val in zip(ins_arr, pol_arr, slp_arr):
        ins_cln = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []
        slp_cln = clean_split_code(str(slp_val)) if pd.notna(slp_val) else []
        cert_cln, pol_cln = process_policy_certificate(str(pol_val)) if pd.notna(pol_val) else ([], [])

        insured_cln_all.append(ins_cln)
        polis_cln_all.append(pol_cln)
        slip_cln_all.append(slp_cln)
        cert_cln_all.append(cert_cln)
        max_ins, max_pol, max_slp = max(max_ins, len(ins_cln)), max(max_pol, len(pol_cln)), max(max_slp, len(slp_cln))

    return insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all, max_ins, max_pol, max_slp

def build_output(df_filtered):
    for col, label in ((INSURED_COLUMN, "INSURED_COLUMN"), (POLIS_COLUMN, "POLIS_COLUMN"), (SLIP_COLUMN, "SLIP_COLUMN")):
        if col not in df_filtered.columns: raise ValueError(f"Kolom {label}='{col}' tidak ditemukan. Tersedia: {list(df_filtered.columns)}")

    insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all, max_ins, max_pol, max_slp = _breakdown_all_rows(df_filtered, INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN)
    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)

    if BROKER_NAME_COLUMN in df_filtered.columns and CEDANT_COLUMN in df_filtered.columns:
        is_direct = df_filtered[BROKER_NAME_COLUMN].astype(str).str.strip().str.upper() == DIRECT_MARKER
        bp_series = df_filtered[CEDANT_COLUMN].where(is_direct, df_filtered[BROKER_NAME_COLUMN])
        business_partner_values = bp_series.astype(str).str.replace(r"^(\s*PT)\.\s*", r"\1 ", regex=True, flags=RE_I).values
    else: business_partner_values = None

    output_cols_data = {}
    for col in original_cols:
        output_cols_data["COMP_NAME2" if col == BROKER_NAME_COLUMN else col] = df_filtered[col].values
        if col == BROKER_NAME_COLUMN and business_partner_values is not None: output_cols_data["BUSSINESS_PATNER"] = business_partner_values
        
        if col == INSURED_COLUMN:
            for i in range(1, max_ins + 1): output_cols_data[f"INSURED_CLEAN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in insured_cln_all]
        if col == POLIS_COLUMN:
            for i in range(1, max_pol + 1):
                output_cols_data[f"POLICY_CLEAN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in polis_cln_all]
                output_cols_data[f"SERTIF_CLEAN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_cln_all]
        if col == SLIP_COLUMN:
            for i in range(1, max_slp + 1): output_cols_data[f"SLIP_CLEAN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slip_cln_all]

    df_master = pd.DataFrame(output_cols_data)

    # Hanya mengatur urutan kolom output; tidak mengubah proses cleansing.
    ordered_cols = []
    for col in df_filtered.columns:
        ordered_cols.append("COMP_NAME2" if col == BROKER_NAME_COLUMN else col)
        if col == INSURED_COLUMN:
            ordered_cols.extend(
                c for c in (f"INSURED_CLEAN_{i}" for i in range(1, max_ins + 1))
                if c in df_master.columns
            )
        elif col == POLIS_COLUMN:
            for i in range(1, max_pol + 1):
                for c in (f"POLICY_CLEAN_{i}", f"SERTIF_CLEAN_{i}"):
                    if c in df_master.columns:
                        ordered_cols.append(c)
        elif col == SLIP_COLUMN:
            ordered_cols.extend(
                c for c in (f"SLIP_CLEAN_{i}" for i in range(1, max_slp + 1))
                if c in df_master.columns
            )

    ordered_cols.extend(c for c in df_master.columns if c not in ordered_cols)
    df_master = df_master[ordered_cols]

    # OUTPUT ONLY: nama, jumlah, dan urutan kolom mengikuti template FACUL terbaru.
    visible_cols = [
        "FAC_CODE", "FAC_CEDANT", "COMP_NAME", "FAC_BROKER", "COMP_NAME2", "BUSSINESS_PATNER",
        "FAC_INSURED", "INSURED_CLEAN_1", "INSURED_CLEAN_2", "INSURED_CLEAN_3", "INSURED_CLEAN_4", "INSURED_CLEAN_5",
        "FAC_POLICY_NO", "POLICY_CLEAN_1", "SERTIF_CLEAN_1", "POLICY_CLEAN_2", "SERTIF_CLEAN_2",
        "POLICY_CLEAN_3", "SERTIF_CLEAN_3", "POLICY_CLEAN_4", "POLICY_CLEAN_5",
        "FAC_SLIP", "SLIP_CLEAN_1", "SLIP_CLEAN_2", "SLIP_CLEAN_3", "SLIP_CLEAN_4", "SLIP_CLEAN_5",
        "FAC_CURRENCY", "FAC_INP_DATE", "FAC_COM_DATE", "FAC_EXP_DATE", "FAC_SUB_CLASS", "FAC_RISK",
        "FAC_DESC", "FAC_ACC_STS", "FAC_STS_SLIP",
    ]
    for c in visible_cols:
        if c not in df_master.columns:
            df_master[c] = ""
    return df_master[visible_cols]


# ==============================================================================
# 5. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df, output_path):
    df.to_excel(output_path, index=False)
    workbook, text_col_indices = load_workbook(output_path), [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]
    worksheet = workbook.active

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None: cell.value = str(cell.value)
    workbook.save(output_path)


# ==============================================================================
# 6. EXECUTION RUNNER
# ==============================================================================
def main():
    print("=" * 70 + "\n CLEANSING DATA 1 - FACULTATIVE | CEDANT: PT LIPPO GENERAL INSURANCE\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        if candidates := [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]:
            input_path = candidates[0]
            print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")
        else: raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    print(f"[INFO] Total baris source: {len(df_raw)}")

    if (df_lippo := filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)).empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter cedant. Proses dihentikan.")
        return

    df_hasil = build_output(df_lippo)
    save_with_text_format(df_hasil, OUTPUT_FILE)
    print("=" * 70 + f"\n[SUCCESS] Selesai. Total baris output: {len(df_hasil)}\n[SUCCESS] File hasil: '{OUTPUT_FILE}'\n" + "=" * 70)

if __name__ == "__main__":
    main()

"""
================================================================================
 MATCHING LIPPO
================================================================================
"""

import os
import re
import sys
import time
import functools

sys.modules['numexpr'] = None
sys.modules['bottleneck'] = None

import numpy as np
import pandas as pd
from dotenv import load_dotenv
from sqlalchemy import create_engine, text


# =============================================================================
# KONFIGURASI LIPPO (V1)
# =============================================================================

SCRIPT_DIR   = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR   = os.path.join(SCRIPT_DIR, "output")

SUSPEND_FILE = os.path.join(OUTPUT_DIR, "lippo_output_suspense_V3.xlsx")
OSBAL_FILE   = os.path.join(OUTPUT_DIR, "lippo_output_osbal_V2.xlsx")
FACUL_FILE   = os.path.join(OUTPUT_DIR, "lippo_output_facul_V3.xlsx")

# Data 5 / Database Rekap RI Slip — dipakai hanya untuk narrowing FAC_CODE.
def _find_data5_file() -> str:
    candidates = [
        os.path.join(SCRIPT_DIR, "input", "5. Database Rekap RI Slip.xlsx"),
        os.path.join(SCRIPT_DIR, "input", "ri slip.xlsx"),
        os.path.join(SCRIPT_DIR, "input", "ri_slip.xlsx"),
        os.path.join(SCRIPT_DIR, "data", "5. Database Rekap RI Slip.xlsx"),
    ]
    for path in candidates:
        if os.path.exists(path):
            return path
    return candidates[0]

DATA5_FILE   = _find_data5_file()
OUTPUT_FILE  = os.path.join(OUTPUT_DIR, "final_output_v1.xlsx")

# Kolom FAC CODE
OSBAL_FACODE_COL = "CCOS_REF_CODE"
FACUL_FACODE_COL = "FAC_CODE"

# Kolom Currency & Periode
SUSPEND_CURR_COL = "CURR ORI"
OSBAL_CURR_COL   = "CCOS_CURR"
SUSPEND_DATE_COL = "RECEIPT DATE"
OSBAL_DATE_COL   = "FAC_COM_DATE"

# Marker baris administratif
_EXCLUDED_REF_MARKERS = ["HUTANG PIUTANG", "DATA SUSPENSE"]
_MIN_TOKEN_LEN = 3
_TOKEN_RE      = re.compile(r"[^A-Z0-9]+")
_DELIMITER_RE  = re.compile(r"[,;|+\s]+")

# Kolom join untuk akumulasi V1 (shared amount)
_SUSPEND_JOIN_COLS = [
    "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "RECEIPT DATE",
    "CEDANT NAME", "CEDANT SHRT NAME",
    "INSURED", "INSURED_CLEAN_1", "INSURED_CLEAN_2",
    "CURR ORI", "CURR PAY", "AMOUNT PAY",
    "POLIS", "POLICY_CLEAN_1", "POLICY_CLEAN_2",
    "SLIP NO", "SLIP_CLEAN_1",
    "DESC 1", "DESC 2", "DESC 3", "DESC 4",
    "STATUS", "REC_TYPE",
]
_SUSPEND_SUM_COLS = ["AMOUNT ORI"]

# Output 38 kolom standar V1
FINAL_COLUMNS = [
    "CCOS_DOC_NO", "CCOS_REF_CODE",
    "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "RECEIPT DATE",
    "CEDANT NAME", "CEDANT SHRT NAME",
    "INSURED_ORI", "INSURED_CLEAN_1", "INSURED_CLEAN_2",
    "CURR ORI", "AMOUNT ORI", "AMOUNT_ORI_MIN1",
    "CURR PAY", "AMOUNT PAY",
    "CCOS_OR_BAL", "CCOS_BAL_DUE", "DIFERENCE",
    "FLAG_PROD",
    "POLIS",
    "POLICY_CLEAN_1", "SERTIF_CLEAN_1",
    "POLICY_CLEAN_2", "SERTIF_CLEAN_2",
    "POLICY_CLEAN_3", "SERTIF_CLEAN_3",
    "POLICY_CLEAN_4", "POLICY_CLEAN_5",
    "SLIP_NO", "SLIP_NO_CLEAN",
    "DESC 1", "DESC 2", "DESC 3", "DESC 4",
    "STATUS", "REC_TYPE",
    "SKENARIO",
]


# =============================================================================
# STRING UTILITIES
# =============================================================================

@functools.lru_cache(maxsize=131072)
def _norm_cached(s: str) -> str:
    text = s.strip().upper().replace("S/D", "SD")
    if text.endswith(".0"):
        text = text[:-2]
    return re.sub(r"\s+", " ", text) if "  " in text else text


def _normalize(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return _norm_cached(str(value))


def _clean_cert_str(val) -> str:
    if val is None or (isinstance(value:=val, float) and pd.isna(value)):
        return ""
    s = str(val).strip()
    return re.sub(r'\.0+$', '', s)


def _expand_sertif_range(val: str) -> list:
    if not val:
        return []
    s = re.sub(r'\.0+$', '', str(val).strip())
    clean = re.sub(r'\s*SD\s*|\s*S/D\s*', '-', s, flags=re.IGNORECASE)
    clean = re.sub(r'\s+', '', clean)
    if clean.isdigit() and 1 <= len(clean) <= 6:
        return [clean.zfill(6)]
    m = re.match(r'^(\d{1,6})-(\d{1,6})$', clean)
    if m:
        start, end = int(m.group(1)), int(m.group(2))
        if start <= end and (end - start) <= 5000:
            return [str(i).zfill(6) for i in range(start, end + 1)]
    return []


def _cert_in_range(query_cert: str, ref_cert_str: str) -> bool:
    q = _clean_cert_str(query_cert)
    if not q or not q.isdigit():
        return False
    q_num = int(q)
    s = _clean_cert_str(ref_cert_str)
    clean = re.sub(r'\s*SD\s*|\s*S/D\s*', '-', s, flags=re.IGNORECASE)
    clean = re.sub(r'\s+', '', clean)
    if clean.isdigit():
        return q_num == int(clean)
    m = re.match(r'^(\d{1,6})-(\d{1,6})$', clean)
    if m:
        start, end = int(m.group(1)), int(m.group(2))
        if start <= end:
            return start <= q_num <= end
        return q_num >= start or q_num <= end
    return False


def _slip_in_range(sus_slip: str, osbal_row: dict) -> bool:
    if not sus_slip:
        return False
    s_clean = re.sub(r'\D', '', str(sus_slip))
    if not s_clean:
        return False
    sl1 = re.sub(r'\D', '', str(osbal_row.get("SLIP_CLEAN_1", osbal_row.get("clean slip 1", ""))))
    sl2 = re.sub(r'\D', '', str(osbal_row.get("SLIP_CLEAN_2", osbal_row.get("clean slip 2", ""))))
    if sl1 and sl2 and len(sl1) == len(s_clean) and len(sl2) == len(s_clean):
        try:
            n_sus = int(s_clean)
            n_sl1 = int(sl1)
            n_sl2 = int(sl2)
            if min(n_sl1, n_sl2) <= n_sus <= max(n_sl1, n_sl2):
                return True
        except ValueError:
            pass
    return False


def _tokenize(text: str) -> frozenset:
    if not text:
        return frozenset()
    return frozenset(t for t in _TOKEN_RE.split(text) if len(t) >= _MIN_TOKEN_LEN)


def _get_clean_cols(columns: list, prefix: str) -> list:
    if not columns:
        return []
    alt_map = {
        "clean polis": ["clean polis", "policy_clean", "polis_cln"],
        "clean slip": ["clean slip", "slip_clean", "slip_no_cln"],
        "clean insured": ["clean insured", "insured_clean", "insured_"],
        "clean sertif": ["clean sertif", "sertif_clean", "sertif_cln", "certificate", "clsdt_sertf"],
    }
    prefixes = alt_map.get(prefix.lower(), [prefix.lower()])
    res = []
    for c in columns:
        clow = str(c).lower()
        if any(clow.startswith(p) for p in prefixes):
            res.append(c)
    return res


def _format_sertif(val) -> str:
    if pd.isna(val) or str(val).strip() == "":
        return ""
    if isinstance(val, (int, float)):
        return str(int(val)).zfill(6)
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s.zfill(6) if s.isdigit() else s


def _is_excluded_ref_row(ref_row: dict, polis_cols: list, slip_cols: list) -> bool:
    for col in list(polis_cols) + list(slip_cols):
        val = _normalize(ref_row.get(col, ""))
        if val and any(m in val for m in _EXCLUDED_REF_MARKERS):
            return True
    return False


def _collect_clean_values(row: dict, cols: list) -> list:
    seen, result = set(), []
    for col in cols:
        val = _normalize(row.get(col, ""))
        if val and val not in seen:
            seen.add(val)
            result.append(val)
    return result


def _ref_has_value(ref_row: dict, clean_cols: list, query_values: list) -> bool:
    if not query_values:
        return False
    ref_vals = [_normalize(ref_row.get(c, "")) for c in clean_cols]
    ref_vals_nonempty = [rv for rv in ref_vals if rv]
    if any(qv in ref_vals_nonempty for qv in query_values):
        return True
    return any(v in rv for v in query_values if v and len(v) >= 10 for rv in ref_vals_nonempty)


def _ref_has_exact_value(ref_row: dict, clean_cols: list, query_values: list) -> bool:
    """Exact-only check untuk narrowing key yang memang berlabel Exact Matching."""
    if not query_values:
        return False
    ref_vals = {_normalize(ref_row.get(c, "")) for c in clean_cols}
    ref_vals.discard("")
    return any(_normalize(qv) in ref_vals for qv in query_values if _normalize(qv))


# =============================================================================
# DATA 5 / CROSS-SOURCE FAC_CODE HELPERS
# Hanya untuk narrowing FAC_CODE: Data 2 -> Data 5 -> Data 1.
# =============================================================================

def _header_norm(value) -> str:
    return re.sub(r"[\s,;:.]+", " ", str(value).strip().upper()).strip()


def _find_first_col(columns: list, aliases: list) -> str | None:
    lookup = {_header_norm(c): c for c in columns}
    for alias in aliases:
        hit = lookup.get(_header_norm(alias))
        if hit is not None:
            return hit
    return None


def _find_aux_cols(columns: list, kind: str) -> list:
    """Cari clean columns Data 5; bila tidak ada, fallback ke raw column."""
    clean = _get_clean_cols(columns, f"clean {kind}")
    if clean:
        return clean
    aliases = {
        "polis": ["POLIS", "POLICY", "POLICY NO", "FAC POLICY NO", "FAC_POLICY_NO"],
        "slip": ["SLIP", "SLIP NO", "SLIP_NO", "FAC SLIP", "FAC_SLIP"],
        "insured": ["INSURED", "INSURED NAME", "FAC INSURED", "FAC_INSURED"],
        "sertif": ["CERTIFICATE", "CERTIFICATE NO", "SERTIFIKAT", "SERTIF", "CLSDT SERTF", "CLSDT_SERTF"],
    }
    col = _find_first_col(columns, aliases.get(kind, []))
    return [col] if col else []


def _read_data5(path: str) -> tuple[list, list]:
    """Baca Data 5 dan cari header pada 12 baris pertama bila diperlukan."""
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"Data 5 / RI Slip tidak ditemukan: {path}\n"
            "Letakkan file '5. Database Rekap RI Slip.xlsx' di folder input."
        )
    df = pd.read_excel(path, header=0, dtype=object)
    aliases = {"SLIP", "SLIP NO", "POLIS", "POLICY", "INSURED", "FAC CODE", "FAC_CODE", "CERTIFICATE"}
    def score(cols):
        norms = {_header_norm(c) for c in cols}
        return sum(1 for a in aliases if _header_norm(a) in norms)
    if score(df.columns) < 2:
        preview = pd.read_excel(path, header=None, nrows=12, dtype=object)
        best_row, best_score = None, score(df.columns)
        for row_idx in range(len(preview)):
            sc = score(preview.iloc[row_idx].tolist())
            if sc > best_score:
                best_row, best_score = row_idx, sc
        if best_row is not None and best_score >= 2:
            df = pd.read_excel(path, header=best_row, dtype=object)
    return df.to_dict(orient="records"), list(df.columns)


def _fac_codes_from_indices(indices: list, rows: list, facode_col: str) -> set:
    return {
        _normalize(rows[i].get(facode_col, ""))
        for i in indices
        if _normalize(rows[i].get(facode_col, ""))
    }


def _build_aux_facode_map(rows: list, value_cols: list, facode_col: str | None,
                           osbal_value_index: dict, osbal_rows: list) -> dict:
    """Aux value -> FAC_CODE. Jika FAC_CODE kosong/tidak ada, turunkan dari OSBAL."""
    result = {}
    for row in rows:
        values = _collect_clean_values(row, value_cols)
        if not values:
            continue
        codes = set()
        if facode_col:
            fc = _normalize(row.get(facode_col, ""))
            if fc:
                codes.add(fc)
        if not codes:
            os_idx = set()
            for value in values:
                os_idx.update(osbal_value_index.get(value, []))
            codes.update(_fac_codes_from_indices(list(os_idx), osbal_rows, OSBAL_FACODE_COL))
        if codes:
            for value in values:
                result.setdefault(value, set()).update(codes)
    return result


def _build_aux_cert_facode_map(rows: list, cert_cols: list, facode_col: str | None,
                                sertif_osbal_idx: dict, osbal_rows: list) -> dict:
    """Certificate/range SD -> FAC_CODE, termasuk expansion SD/S-D di sisi auxiliary."""
    result = {}
    for row in rows:
        certs = set()
        for col in cert_cols:
            raw = row.get(col, "")
            expanded = _expand_sertif_range(raw)
            if expanded:
                certs.update(expanded)
            else:
                c = _clean_cert_str(raw)
                if c:
                    certs.add(c.zfill(6) if c.isdigit() else _normalize(c))
        if not certs:
            continue
        codes = set()
        if facode_col:
            fc = _normalize(row.get(facode_col, ""))
            if fc:
                codes.add(fc)
        if not codes:
            os_idx = set()
            for cert in certs:
                os_idx.update(sertif_osbal_idx.get(cert, []))
            codes.update(_fac_codes_from_indices(list(os_idx), osbal_rows, OSBAL_FACODE_COL))
        for cert in certs:
            result.setdefault(cert, set()).update(codes)
    return result


def _aux_codes_for_values(values: list, aux_map: dict, cert_mode: bool = False) -> set:
    codes = set()
    for value in values:
        keys = _expand_sertif_range(value) if cert_mode else [_normalize(value)]
        if cert_mode and not keys:
            c = _clean_cert_str(value)
            keys = [c.zfill(6) if c.isdigit() else _normalize(c)] if c else []
        for key in keys:
            if key:
                codes.update(aux_map.get(key, set()))
    return codes


def _cross_source_narrow_facodes(current_indices: set, primary_method: str,
                                  sus: dict, osbal_rows: list,
                                  data5_maps: dict, facul_maps: dict,
                                  polis_sus: list, slip_sus: list,
                                  sertif_sus: list, insured_sus: list) -> set:
    """
    Narrowing lintas source sesuai cascade Lippo:
      Data 2/OSBAL -> Data 5/RI Slip -> Data 1/FACUL.

    Di SETIAP source auxiliary, key dicoba berurutan:
      POLIS -> SLIP -> SERTIFIKAT -> INSURED.

    Source auxiliary hanya boleh mempersempit FAC_CODE kandidat dari source
    sebelumnya melalui intersection. Tidak boleh menambah/mengganti FAC_CODE.
    Jika overlap kosong ATAU overlap sama dengan seluruh kandidat saat ini,
    kandidat dipertahankan dan key berikutnya tetap dicoba.
    """
    current = set(current_indices)
    current_codes = _fac_codes_from_indices(list(current), osbal_rows, OSBAL_FACODE_COL)
    if len(current_codes) <= 1:
        return current

    method_cols = {
        "POLIS": polis_sus,
        "SLIP": slip_sus,
        "SERTIFIKAT": sertif_sus,
        "INSURED": insured_sus,
    }
    method_order = ("POLIS", "SLIP", "SERTIFIKAT", "INSURED")

    # Urutan source wajib: Data 5 / RI Slip -> Data 1 / FACUL.
    for source_maps in (data5_maps, facul_maps):
        if len(current_codes) <= 1:
            break

        for method in method_order:
            if len(current_codes) <= 1:
                break
            values = _collect_clean_values(sus, method_cols[method])
            if not values:
                continue

            source_codes = _aux_codes_for_values(
                values, source_maps.get(method, {}), cert_mode=(method == "SERTIFIKAT")
            )
            if not source_codes:
                continue

            overlap = current_codes & source_codes
            # Tidak overlap: source/key ini tidak boleh mengganti kandidat.
            if not overlap:
                continue
            # Semua kandidat overlap: belum ada narrowing; lanjut key berikutnya.
            if overlap == current_codes:
                continue

            # Proper subset: baru benar-benar mempersempit FAC_CODE.
            current_codes = set(overlap)
            current = {
                i for i in current
                if _normalize(osbal_rows[i].get(OSBAL_FACODE_COL, "")) in current_codes
            }

    return current


# =============================================================================
# INDEX BUILDERS
# =============================================================================

def _build_exact_index(rows: list, cols: list, excluded: set = None) -> dict:
    index = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        for col in cols:
            val = _normalize(row.get(col, ""))
            if val:
                index.setdefault(val, []).append(i)
    return index


def _build_facode_index(rows: list, col: str, excluded: set = None) -> dict:
    index = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        val = _normalize(row.get(col, ""))
        if val:
            index.setdefault(val, []).append(i)
    return index


def _build_token_index(rows: list, clean_cols: list, excluded: set = None) -> tuple:
    token_index = {}
    token_cache = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        tokens = frozenset()
        for c in clean_cols:
            t = _normalize(row.get(c, ""))
            if t:
                tokens = tokens | _tokenize(t)
        token_cache[i] = tokens
        for tok in tokens:
            token_index.setdefault(tok, set()).add(i)
    return token_index, token_cache


def _build_sertif_index(rows: list, sertif_cols: list, excluded: set = None) -> dict:
    index = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        for col in sertif_cols:
            val = _normalize(row.get(col, ""))
            for cert in _expand_sertif_range(val):
                index.setdefault(cert, []).append(i)
    return index


def _build_lookup(rows: list, polis_cols: list, slip_cols: list,
                  insured_cols: list, facode_col: str = None, label: str = "") -> tuple:
    if not rows:
        return ({}, {}, {}, ({}, {}), ({}, {}), ({}, {}), {})

    t0 = time.perf_counter()
    excluded = {i for i, r in enumerate(rows) if _is_excluded_ref_row(r, polis_cols, slip_cols)}
    facode_index = _build_facode_index(rows, facode_col, excluded=excluded) if facode_col else {}
    lookup = (
        _build_exact_index(rows, slip_cols,    excluded=excluded),
        _build_exact_index(rows, polis_cols,   excluded=excluded),
        _build_exact_index(rows, insured_cols, excluded=excluded),
        _build_token_index(rows, slip_cols,    excluded=excluded),
        _build_token_index(rows, polis_cols,   excluded=excluded),
        _build_token_index(rows, insured_cols, excluded=excluded),
        facode_index,
    )
    elapsed = time.perf_counter() - t0
    exc_info = f" - {len(excluded):,} baris dikecualikan" if excluded else ""
    print(f"  {label:<6} done ({elapsed:.1f}s){exc_info}", flush=True)
    return lookup


# =============================================================================
# MATCH ENGINE
# =============================================================================

def _exact_match(query_values: list, index: dict) -> set:
    return {i for v in query_values if v and v in index for i in index[v]}


def _like_match(query_values: list, token_index: dict, rows: list, ref_cols: list) -> list:
    valid_queries = [v for v in query_values if v and len(v) >= 2]
    if not valid_queries:
        return []
    query_tokens = frozenset()
    for v in valid_queries:
        query_tokens = query_tokens | _tokenize(v)
    candidates = set()
    for tok in query_tokens:
        if tok in token_index:
            candidates |= token_index[tok]
    if not candidates:
        return []
    matched = []
    for i in candidates:
        ref_values = [_normalize(rows[i].get(c, "")) for c in ref_cols]
        for q in valid_queries:
            if any(q in rv for rv in ref_values if rv):
                matched.append(i)
                break
    return matched


def _resolve_facode(matched_rows: list, facode_col: str, facode_osbal_idx: dict, osbal_rows: list) -> tuple:
    osbal_indices = set()
    found_fac_codes = set()
    for row in matched_rows:
        fac = _normalize(row.get(facode_col, ""))
        if fac:
            found_fac_codes.add(fac)
            if fac in facode_osbal_idx:
                osbal_indices.update(facode_osbal_idx[fac])
    return osbal_indices, found_fac_codes


def _aggregate_ccos(rows: list, all_fac_osbal_rows: list = None) -> dict:
    if not rows or not rows[0]:
        return {"CCOS_DOC_NO": "", "CCOS_REF_CODE": "", "CCOS_OR_BAL": np.nan, "CCOS_BAL_DUE": np.nan}
    docs, refs = [], []
    for r in rows:
        d = str(r.get("CCOS_DOC_NO", "")).strip()
        rf = str(r.get(OSBAL_FACODE_COL, "")).strip()
        if d and d not in docs:
            docs.append(d)
        if rf and rf not in refs:
            refs.append(rf)
    calc_rows = all_fac_osbal_rows if all_fac_osbal_rows is not None else rows
    seen_indices = set()
    unique_calc_rows = []
    for r in calc_rows:
        r_id = id(r)
        if r_id not in seen_indices:
            seen_indices.add(r_id)
            unique_calc_rows.append(r)
    or_bal_sum  = sum(float(r.get("CCOS_OR_BAL", 0) or 0) for r in unique_calc_rows)
    bal_due_sum = sum(float(r.get("CCOS_BAL_DUE", 0) or 0) for r in unique_calc_rows)
    return {
        "CCOS_DOC_NO":   ", ".join(docs),
        "CCOS_REF_CODE": ", ".join(refs),
        "CCOS_OR_BAL":   or_bal_sum,
        "CCOS_BAL_DUE":  bal_due_sum,
    }


def _format_final_scenario(base_scenario: str, source: str, is_beda_curr: bool = False, is_beda_per: bool = False, num_fac: int = 1) -> str:
    base = base_scenario or "Unmatching"
    for to_remove in [" (Beda Currency)", " (Beda Periode)", " (> 1 Fac)"]:
        base = base.replace(to_remove, "")
    if is_beda_curr and is_beda_per:
        return f"{base} (Beda Currency) (Beda Periode)"
    elif is_beda_curr:
        return f"{base} (Beda Currency)"
    elif is_beda_per:
        return f"{base} (Beda Periode)"
    elif num_fac > 1:
        return f"{base} (> 1 Fac)"
    return base


def _build_output_row(
    suspend_row: dict, source: str, scenario: str, osbal_rows: list,
    osbal_count: int = 0, resolved: bool = True,
    suspend_count: int = 1, all_fac_osbal_rows: list = None,
) -> dict:
    has_match = bool(source and osbal_rows and osbal_rows[0])
    has_osbal = has_match and (source in ("OSBAL", "FACUL") or resolved)

    ccos = (
        _aggregate_ccos(osbal_rows, all_fac_osbal_rows=all_fac_osbal_rows)
        if has_osbal
        else {"CCOS_DOC_NO": "", "CCOS_REF_CODE": "", "CCOS_OR_BAL": np.nan, "CCOS_BAL_DUE": np.nan}
    )
    if "Beda Currency" in str(scenario) or "Beda Periode" in str(scenario):
        ccos["CCOS_OR_BAL"] = np.nan
        ccos["CCOS_BAL_DUE"] = np.nan

    final_scenario = "Unmatching" if not source or scenario in ("Unmatching", "UNMATCHED") else scenario

    return {
        "CCOS_DOC_NO":       ccos["CCOS_DOC_NO"] if has_osbal else "",
        "CCOS_REF_CODE":     ccos["CCOS_REF_CODE"],
        "RECEIPT NO":        suspend_row.get("RECEIPT NO",        ""),
        "CREDIT NOTES":      suspend_row.get("CREDIT NOTES",      ""),
        "DETAIL RINCIAN NO": suspend_row.get("DETAIL RINCIAN NO", ""),
        "RECEIPT DATE":      suspend_row.get("RECEIPT DATE",      ""),
        "CEDANT NAME":       suspend_row.get("CEDANT NAME",       ""),
        "CEDANT SHRT NAME":  suspend_row.get("CEDANT SHRT NAME",  ""),
        "INSURED_ORI":       suspend_row.get("INSURED") or suspend_row.get("insured_ori") or "",
        "INSURED_CLEAN_1":   suspend_row.get("INSURED_CLEAN_1") or suspend_row.get("clean insured 1") or "",
        "INSURED_CLEAN_2":   suspend_row.get("INSURED_CLEAN_2") or suspend_row.get("clean insured 2") or "",
        "CURR ORI":          suspend_row.get("CURR ORI",          ""),
        "AMOUNT ORI":        suspend_row.get("AMOUNT ORI",        ""),
        "CURR PAY":          suspend_row.get("CURR PAY",          ""),
        "AMOUNT PAY":        suspend_row.get("AMOUNT PAY",        ""),
        "CCOS_OR_BAL":       ccos["CCOS_OR_BAL"],
        "CCOS_BAL_DUE":      ccos["CCOS_BAL_DUE"],
        "POLIS":             suspend_row.get("POLIS") or "",
        "POLICY_CLEAN_1":    suspend_row.get("POLICY_CLEAN_1") or "",
        "SERTIF_CLEAN_1":    _format_sertif(suspend_row.get("SERTIF_CLEAN_1") or suspend_row.get("SERTIF") or suspend_row.get("CERTIFICATE") or ""),
        "POLICY_CLEAN_2":    suspend_row.get("POLICY_CLEAN_2") or "",
        "SERTIF_CLEAN_2":    _format_sertif(suspend_row.get("SERTIF_CLEAN_2") or suspend_row.get("CERTIFICATE_2") or ""),
        "POLICY_CLEAN_3":    suspend_row.get("POLICY_CLEAN_3") or "",
        "SERTIF_CLEAN_3":    _format_sertif(suspend_row.get("SERTIF_CLEAN_3") or suspend_row.get("CERTIFICATE_3") or ""),
        "POLICY_CLEAN_4":    suspend_row.get("POLICY_CLEAN_4") or "",
        "POLICY_CLEAN_5":    suspend_row.get("POLICY_CLEAN_5") or "",
        "SLIP_NO":           suspend_row.get("SLIP NO") or suspend_row.get("SLIP_NO") or "",
        "SLIP_NO_CLEAN":      suspend_row.get("SLIP_CLEAN_1") or suspend_row.get("SLIP_NO_CLEAN") or suspend_row.get("SLIP_NO_CLN") or "",
        "DESC 1":            suspend_row.get("DESC 1",            ""),
        "DESC 2":            suspend_row.get("DESC 2",            ""),
        "DESC 3":            suspend_row.get("DESC 3",            ""),
        "DESC 4":            suspend_row.get("DESC 4",            ""),
        "STATUS":            suspend_row.get("STATUS",            ""),
        "REC_TYPE":          suspend_row.get("REC_TYPE",          ""),
        "SKENARIO":          final_scenario,
        "_OSBAL_ROW_COUNT":  osbal_count,
        "_SUSPEND_ROW_COUNT": suspend_count,
    }


def _merge_suspend_rows(rows: list) -> dict:
    if len(rows) == 1:
        return rows[0]
    merged = {}
    for col in _SUSPEND_SUM_COLS:
        total = 0.0
        for r in rows:
            v = r.get(col, 0)
            if pd.isna(v) or v == "":
                continue
            if isinstance(v, (int, float)):
                total += float(v)
            else:
                s = str(v).strip()
                if not s:
                    continue
                if ',' in s and '.' in s:
                    s = s.replace('.', '').replace(',', '.') if s.rfind(',') > s.rfind('.') else s.replace(',', '')
                elif ',' in s:
                    s = s.replace(',', '.')
                try:
                    total += float(s)
                except ValueError:
                    pass
        merged[col] = total
    for col in _SUSPEND_JOIN_COLS:
        seen, vals = set(), []
        for r in rows:
            v = str(r.get(col, "")).strip()
            if v and v not in seen:
                seen.add(v)
                vals.append(v)
        merged[col] = ", ".join(vals)
    return merged


def _compute_derived_cols(df: pd.DataFrame) -> pd.DataFrame:
    amount_ori_base = df.get("AMOUNT ORI", pd.Series([0]*len(df), index=df.index))
    if "_MERGED_AMOUNT_ORI" in df.columns:
        amount_ori_raw = df["_MERGED_AMOUNT_ORI"].combine_first(amount_ori_base)
    else:
        amount_ori_raw = amount_ori_base

    amount_ori_clean = amount_ori_raw.replace(r'^\s*$', np.nan, regex=True)
    if amount_ori_clean.dtype == object:
        amount_ori_clean = amount_ori_clean.astype(str).str.strip().str.replace(r'\s+', '', regex=True)
        amount_ori_clean = amount_ori_clean.replace("nan", np.nan)
        has_comma = amount_ori_clean.str.contains(',', na=False)
        amount_ori_eu = (amount_ori_clean
                         .str.replace('.', '', regex=False)
                         .str.replace(',', '.', regex=False))
        amount_ori_clean = amount_ori_eu.where(has_comma, amount_ori_clean)
    amount_ori_numeric = pd.to_numeric(amount_ori_clean, errors="coerce")

    df["AMOUNT_ORI_MIN1"] = np.where(amount_ori_numeric.notna(), amount_ori_numeric * -1, np.nan)
    amount_ori = amount_ori_numeric.fillna(0)
    amount_neg = df["AMOUNT_ORI_MIN1"].fillna(0)

    balance_due     = pd.to_numeric(df["CCOS_BAL_DUE"] if "CCOS_BAL_DUE" in df.columns else pd.Series(0, index=df.index), errors="coerce").fillna(0)
    osbal_count_col = pd.to_numeric(df["_OSBAL_ROW_COUNT"] if "_OSBAL_ROW_COUNT" in df.columns else pd.Series(0, index=df.index), errors="coerce").fillna(0)
    sus_count_col   = pd.to_numeric(df["_SUSPEND_ROW_COUNT"] if "_SUSPEND_ROW_COUNT" in df.columns else pd.Series(1, index=df.index), errors="coerce").fillna(1)
    ccos_ref        = df.get("CCOS_REF_CODE", pd.Series("", index=df.index)).fillna("").astype(str)
    skenario_col    = df.get("SKENARIO",      pd.Series("", index=df.index)).fillna("")

    df["DIFERENCE"] = amount_neg - balance_due
    accumulated = (osbal_count_col > 1) | (sus_count_col > 1)
    is_equal    = amount_neg.round(2) == balance_due.round(2)

    is_unmatched         = skenario_col.str.strip().str.upper().isin(["UNMATCHING", "UNMATCHED"]) | (skenario_col == "")
    is_beda_currency     = skenario_col.str.contains("Beda Currency", na=False, regex=False)
    is_beda_periode      = skenario_col.str.contains("Beda Periode",  na=False, regex=False)
    is_gt1_fac           = (skenario_col.str.contains("> 1 Fac", na=False, regex=False) | ccos_ref.str.contains(",", na=False))
    is_new_entry_total   = ~is_unmatched & (amount_ori != 0) & (balance_due == 0)
    is_adj_total         = ~is_unmatched & is_equal
    is_adj_sebagian      = ~is_unmatched & ~is_equal & (amount_neg < balance_due)
    is_new_entry_sebagian= ~is_unmatched & ~is_equal & (amount_neg > balance_due)

    df["FLAG_PROD"] = np.select(
        [is_beda_currency, is_beda_periode, is_gt1_fac,
         is_new_entry_total,
         is_adj_total & ~accumulated, is_adj_total & accumulated,
         is_adj_sebagian & ~accumulated, is_adj_sebagian & accumulated,
         is_new_entry_sebagian],
        ["Beda Currency", "Beda Periode",
         "Matching >1 fac code",
         "New Entry total",
         "Adjustment total tanpa akumulasi", "Adjustment total dengan akumulasi",
         "Adjustment sebagian tanpa akumulasi", "Adjustment sebagian dengan akumulasi",
         "New Entry sebagian"],
        default="Unmatching",
    )

    for col in ["_OSBAL_ROW_COUNT", "_SUSPEND_ROW_COUNT", "_MERGED_AMOUNT_ORI"]:
        if col in df.columns:
            df.drop(columns=[col], inplace=True)
    return df


# =============================================================================
# PIPELINE EXECUTION
# =============================================================================

def _load_excel(path: str, label: str) -> tuple:
    if not os.path.exists(path):
        print(f"  {label}: {path} (TIDAK DITEMUKAN)", flush=True)
        return [], []

    cache = path + ".cache.pkl"
    if os.path.exists(cache) and os.path.getmtime(cache) >= os.path.getmtime(path):
        t = time.perf_counter()
        print(f"  {label}: loading dari cache ...", flush=True)
        try:
            data, header = pd.read_pickle(cache)
            print(f"  -> {len(data):,} rows, {len(header)} cols ({time.perf_counter()-t:.2f}s)", flush=True)
            return data, header
        except Exception:
            pass

    t = time.perf_counter()
    print(f"  {label}: membaca Excel {path} ...", flush=True)
    df = pd.read_excel(path)
    data = df.to_dict(orient="records")
    header = list(df.columns)
    print(f"  -> {len(data):,} rows, {len(header)} cols ({time.perf_counter()-t:.2f}s)", flush=True)
    try:
        pd.to_pickle((data, header), cache)
    except Exception:
        pass
    return data, header


def run() -> None:
    t_start = time.perf_counter()
    print("\n" + "=" * 60, flush=True)
    print("  PRODUCTION SCRIPT — SUSPEND MATCHING (LIPPO) — V1", flush=True)
    print("=" * 60, flush=True)

    # [1] Load data
    print("\n[1/6] Loading data ...", flush=True)
    suspend_rows, suspend_cols = _load_excel(SUSPEND_FILE, "SUSPEND")
    osbal_rows,   osbal_cols   = _load_excel(OSBAL_FILE,   "OSBAL")
    facul_rows,   facul_cols   = _load_excel(FACUL_FILE,   "FACUL")
    print(f"  DATA 5: membaca RI Slip {DATA5_FILE} ...", flush=True)
    data5_rows,   data5_cols   = _read_data5(DATA5_FILE)
    print(f"  -> {len(data5_rows):,} rows, {len(data5_cols)} cols", flush=True)

    # Pre-parse tanggal OSBAL FAC_COM_DATE
    print("  Pre-parsing OSBAL FAC_COM_DATE ...", flush=True)
    for r in osbal_rows:
        raw = r.get(OSBAL_DATE_COL, "")
        r["_com_date_parsed"] = pd.to_datetime(raw, errors="coerce") \
            if raw and not (isinstance(raw, float) and pd.isna(raw)) else None

    # Pre-parse tanggal Suspend RECEIPT DATE
    print("  Pre-parsing Suspend RECEIPT DATE ...", flush=True)
    for r in suspend_rows:
        raw = r.get(SUSPEND_DATE_COL, "")
        r["_sus_date_parsed"] = pd.to_datetime(raw, errors="coerce") \
            if raw and not (isinstance(raw, float) and pd.isna(raw)) else None

    # [2] Deteksi kolom bersih
    print("\n[2/6] Detecting clean columns ...", flush=True)
    polis_sus   = _get_clean_cols(suspend_cols, "clean polis")
    slip_sus    = _get_clean_cols(suspend_cols, "clean slip")
    insured_sus = _get_clean_cols(suspend_cols, "clean insured")
    sertif_sus  = _get_clean_cols(suspend_cols, "clean sertif")

    polis_osbal   = _get_clean_cols(osbal_cols, "clean polis")
    slip_osbal    = _get_clean_cols(osbal_cols, "clean slip")
    insured_osbal = _get_clean_cols(osbal_cols, "clean insured")
    sertif_osbal  = _get_clean_cols(osbal_cols, "clean sertif")

    polis_facul   = _get_clean_cols(facul_cols, "clean polis")
    slip_facul    = _get_clean_cols(facul_cols, "clean slip")
    insured_facul = _get_clean_cols(facul_cols, "clean insured")
    sertif_facul  = _get_clean_cols(facul_cols, "clean sertif")

    # Data 5 tidak diasumsikan selengkap Data 2; gunakan clean/raw yang tersedia.
    polis_data5   = _find_aux_cols(data5_cols, "polis")
    slip_data5    = _find_aux_cols(data5_cols, "slip")
    insured_data5 = _find_aux_cols(data5_cols, "insured")
    sertif_data5  = _find_aux_cols(data5_cols, "sertif")
    data5_facode_col = _find_first_col(data5_cols, ["FAC_CODE", "FAC CODE", "FAC. CODE", "CCOS_REF_CODE"])

    def _prepend(col, cols_list, header):
        return ([col] + cols_list) if col in header and col not in cols_list else cols_list

    polis_osbal  = _prepend("CLSDT_POLICY_NO", polis_osbal,  osbal_cols)
    polis_osbal  = _prepend("FAC_POLICY_NO",   polis_osbal,  osbal_cols)
    slip_osbal   = _prepend("CLSDT_SLIP_NO",   slip_osbal,   osbal_cols)
    slip_osbal   = _prepend("FAC_SLIP",        slip_osbal,   osbal_cols)

    polis_facul  = _prepend("FAC_POLICY_NO",   polis_facul,  facul_cols)
    slip_facul   = _prepend("FAC_SLIP",        slip_facul,   facul_cols)

    # [3] Bangun lookup index
    print("\n[3/6] Building lookup indexes ...", flush=True)
    lookup_osbal = _build_lookup(osbal_rows, polis_osbal, slip_osbal, insured_osbal, OSBAL_FACODE_COL, label="OSBAL")
    lookup_facul = _build_lookup(facul_rows, polis_facul, slip_facul, insured_facul, FACUL_FACODE_COL, label="FACUL")

    facode_osbal_idx = lookup_osbal[6]
    excluded_osbal   = {i for i, r in enumerate(osbal_rows) if _is_excluded_ref_row(r, polis_osbal, slip_osbal)}
    sertif_osbal_idx = _build_sertif_index(osbal_rows, sertif_osbal, excluded=excluded_osbal)
    print(f"  Sertif OSBAL index: {len(sertif_osbal_idx):,} sertif unik", flush=True)

    # Index FAC_CODE auxiliary. Data 5 lalu Data 1 hanya dipakai untuk
    # intersection terhadap FAC_CODE kandidat Data 2.
    data5_maps = {
        "POLIS": _build_aux_facode_map(data5_rows, polis_data5, data5_facode_col, lookup_osbal[1], osbal_rows),
        "SLIP": _build_aux_facode_map(data5_rows, slip_data5, data5_facode_col, lookup_osbal[0], osbal_rows),
        "SERTIFIKAT": _build_aux_cert_facode_map(data5_rows, sertif_data5, data5_facode_col, sertif_osbal_idx, osbal_rows),
        "INSURED": _build_aux_facode_map(data5_rows, insured_data5, data5_facode_col, lookup_osbal[2], osbal_rows),
    }
    facul_maps = {
        "POLIS": _build_aux_facode_map(facul_rows, polis_facul, FACUL_FACODE_COL, lookup_osbal[1], osbal_rows),
        "SLIP": _build_aux_facode_map(facul_rows, slip_facul, FACUL_FACODE_COL, lookup_osbal[0], osbal_rows),
        "SERTIFIKAT": _build_aux_cert_facode_map(facul_rows, sertif_facul, FACUL_FACODE_COL, sertif_osbal_idx, osbal_rows),
        "INSURED": _build_aux_facode_map(facul_rows, insured_facul, FACUL_FACODE_COL, lookup_osbal[2], osbal_rows),
    }

    # [4] Matching per baris
    print(f"\n[4/6] Matching {len(suspend_rows):,} baris suspend (Basic Rules, No Bordero) ...", flush=True)
    t4 = time.perf_counter()
    raw_results = []

    for n, sus in enumerate(suspend_rows, 1):
        if n % 1000 == 0:
            print(f"    Memproses baris {n:,} ...", flush=True)

        eff_polis = polis_sus
        eff_slip  = slip_sus
        sus_date  = sus.get("_sus_date_parsed")

        matched_fac_codes = set()
        matched_osbal_idx = set()
        source = None
        scenario = ""
        resolved = False
        # Key terakhir yang benar-benar berhasil mempersempit kandidat Data 2.
        # Dipakai hanya untuk pencarian FAC_CODE ke Data 5 -> Data 1.
        fac_lookup_method = ""

        # FASE 0: PASS A (Polis + Sertif)
        sus_cert_raw = sus.get("SERTIF_CLEAN_1", sus.get("SERTIF", sus.get("CERTIFICATE", "")))
        sus_cert_val = _clean_cert_str(sus_cert_raw) if sus_cert_raw else ""
        if sus_cert_val and sus_cert_val in sertif_osbal_idx:
            c_idxs = sertif_osbal_idx[sus_cert_val]
            p_vals = _collect_clean_values(sus, eff_polis)
            if p_vals:
                confirmed_idxs = [i for i in c_idxs if _ref_has_value(osbal_rows[i], polis_osbal, p_vals)]
                if confirmed_idxs:
                    matched_osbal_idx.update(confirmed_idxs)
                    source = "OSBAL"
                    scenario = "Polis + Sertifikat"
                    fac_lookup_method = "SERTIFIKAT"
                    resolved = True

        # FASE 1: POLIS
        if not matched_osbal_idx:
            polis_vals = _collect_clean_values(sus, eff_polis)
            if polis_vals:
                # R1: OSBAL polis exact
                curr_hits = _exact_match(polis_vals, lookup_osbal[1])
                if curr_hits:
                    matched_osbal_idx.update(curr_hits)
                    source = "OSBAL"
                    scenario = "Polis only"
                    fac_lookup_method = "POLIS"
                    resolved = True
                
                # R2: FACUL polis exact -> OSBAL
                elif facul_rows:
                    fac_hits = _exact_match(polis_vals, lookup_facul[1])
                    if fac_hits:
                        res_idx, _ = _resolve_facode([facul_rows[i] for i in fac_hits], FACUL_FACODE_COL, facode_osbal_idx, osbal_rows)
                        if res_idx:
                            matched_osbal_idx.update(res_idx)
                            source = "FACUL"
                            scenario = "Polis only"
                            fac_lookup_method = "POLIS"
                            resolved = True

                # R3: OSBAL polis LIKE
                if not matched_osbal_idx:
                    like_hits = _like_match(polis_vals, lookup_osbal[4][0], osbal_rows, polis_osbal)
                    if like_hits:
                        matched_osbal_idx.update(like_hits)
                        source = "OSBAL"
                        scenario = "Polis only"
                        resolved = True

                # R4: FACUL polis LIKE -> OSBAL
                elif not matched_osbal_idx and facul_rows:
                    fac_like = _like_match(polis_vals, lookup_facul[4][0], facul_rows, polis_facul)
                    if fac_like:
                        res_idx, _ = _resolve_facode([facul_rows[i] for i in fac_like], FACUL_FACODE_COL, facode_osbal_idx, osbal_rows)
                        if res_idx:
                            matched_osbal_idx.update(res_idx)
                            source = "FACUL"
                            scenario = "Polis only"
                            fac_lookup_method = "POLIS"
                            resolved = True

        # FASE 2: SLIP
        if not matched_osbal_idx:
            slip_vals = _collect_clean_values(sus, eff_slip)
            # Guard Lippo: nilai yang identik dengan CLEAN POLIS bukan bukti SLIP.
            polis_vals_for_slip_guard = {_normalize(v) for v in _collect_clean_values(sus, eff_polis) if _normalize(v)}
            slip_vals = [v for v in slip_vals if _normalize(v) not in polis_vals_for_slip_guard]
            if slip_vals:
                # R5: OSBAL slip exact
                curr_hits = _exact_match(slip_vals, lookup_osbal[0])
                if curr_hits:
                    matched_osbal_idx.update(curr_hits)
                    source = "OSBAL"
                    scenario = "Slip only"
                    fac_lookup_method = "SLIP"
                    resolved = True

                # R6: FACUL slip exact -> OSBAL
                elif facul_rows:
                    fac_hits = _exact_match(slip_vals, lookup_facul[0])
                    if fac_hits:
                        res_idx, _ = _resolve_facode([facul_rows[i] for i in fac_hits], FACUL_FACODE_COL, facode_osbal_idx, osbal_rows)
                        if res_idx:
                            matched_osbal_idx.update(res_idx)
                            source = "FACUL"
                            scenario = "Slip only"
                            fac_lookup_method = "SLIP"
                            resolved = True

                # R7: OSBAL slip LIKE
                if not matched_osbal_idx:
                    like_hits = _like_match(slip_vals, lookup_osbal[3][0], osbal_rows, slip_osbal)
                    if like_hits:
                        matched_osbal_idx.update(like_hits)
                        source = "OSBAL"
                        scenario = "Slip only"
                        resolved = True

                # R8: FACUL slip LIKE -> OSBAL
                elif not matched_osbal_idx and facul_rows:
                    fac_like = _like_match(slip_vals, lookup_facul[3][0], facul_rows, slip_facul)
                    if fac_like:
                        res_idx, _ = _resolve_facode([facul_rows[i] for i in fac_like], FACUL_FACODE_COL, facode_osbal_idx, osbal_rows)
                        if res_idx:
                            matched_osbal_idx.update(res_idx)
                            source = "FACUL"
                            scenario = "Slip only"
                            fac_lookup_method = "SLIP"
                            resolved = True

        # FASE 3: INSURED FALLBACK
        if not matched_osbal_idx:
            ins_vals = _collect_clean_values(sus, insured_sus)
            if ins_vals:
                curr_hits = _exact_match(ins_vals, lookup_osbal[2])
                if curr_hits:
                    matched_osbal_idx.update(curr_hits)
                    source = "OSBAL"
                    scenario = "Insured only"
                    fac_lookup_method = "INSURED"
                    resolved = True
                elif facul_rows:
                    fac_hits = _exact_match(ins_vals, lookup_facul[2])
                    if fac_hits:
                        res_idx, _ = _resolve_facode([facul_rows[i] for i in fac_hits], FACUL_FACODE_COL, facode_osbal_idx, osbal_rows)
                        if res_idx:
                            matched_osbal_idx.update(res_idx)
                            source = "FACUL"
                            scenario = "Insured only"
                            fac_lookup_method = "INSURED"
                            resolved = True

        # FASE 5: NARROWING UTAMA (Non-Destructive)
        # Key berikutnya HANYA diperiksa bila kandidat masih memiliki >1 FAC_CODE.
        # Banyak baris OSBAL dengan FAC_CODE yang sama tetap dianggap sudah resolved;
        # karena itu Slip/Sertifikat/Insured tidak boleh ditambahkan ke skenario.
        if len(_fac_codes_from_indices(list(matched_osbal_idx), osbal_rows, OSBAL_FACODE_COL)) > 1:
            matched_list = list(matched_osbal_idx)

            # Step 1: Slip Exact Matching
            # Slip hanya dianggap berhasil untuk pencarian FAC_CODE bila benar-benar
            # MEMPERSEMPIT kumpulan FAC_CODE. Match yang tidak mengurangi FAC_CODE
            # tidak boleh menambah skenario "+ Slip" atau mengganti key cross-check.
            # STRICT FIELD-TO-FIELD: nilai dari kolom Slip Suspend hanya diuji
            # terhadap kolom Slip OSBAL. Walaupun nilainya sama dengan Clean Polis,
            # nilainya tetap boleh diuji sebagai Slip, tetapi TIDAK boleh dianggap
            # match hanya karena ditemukan pada kolom Polis.
            slip_vals = _collect_clean_values(sus, slip_sus)
            if slip_vals:
                before_codes = _fac_codes_from_indices(matched_list, osbal_rows, OSBAL_FACODE_COL)
                confirmed = [idx for idx in matched_list if _ref_has_exact_value(osbal_rows[idx], slip_osbal, slip_vals)]
                confirmed_codes = _fac_codes_from_indices(confirmed, osbal_rows, OSBAL_FACODE_COL) if confirmed else set()
                if confirmed and confirmed_codes and confirmed_codes < before_codes:
                    matched_list = confirmed
                    scenario = "Polis + Slip" if "Polis" in scenario else scenario
                    fac_lookup_method = "SLIP"

            # Step 1b: Slip Numeric Range Check
            # Sama: range Slip hanya dipakai bila mengurangi FAC_CODE kandidat.
            if len(matched_list) > 1 and slip_vals:
                before_codes = _fac_codes_from_indices(matched_list, osbal_rows, OSBAL_FACODE_COL)
                confirmed_slip = [
                    idx for idx in matched_list
                    if any(_slip_in_range(sv, osbal_rows[idx]) for sv in slip_vals)
                ]
                confirmed_codes = _fac_codes_from_indices(confirmed_slip, osbal_rows, OSBAL_FACODE_COL) if confirmed_slip else set()
                if confirmed_slip and confirmed_codes and confirmed_codes < before_codes:
                    matched_list = confirmed_slip
                    if "Slip" not in scenario:
                        scenario = "Polis + Slip" if "Polis" in scenario else scenario
                    fac_lookup_method = "SLIP"

            # Step 2: Sertifikat Exact & In-Range Matching
            # Sama seperti Slip/Insured: Sertifikat hanya dianggap narrowing bila
            # benar-benar mengurangi kumpulan FAC_CODE.
            sus_cert_vals = set(_clean_cert_str(v) for v in _collect_clean_values(sus, sertif_sus) if _clean_cert_str(v))
            if sus_cert_vals and len(_fac_codes_from_indices(matched_list, osbal_rows, OSBAL_FACODE_COL)) > 1:
                before_codes = _fac_codes_from_indices(matched_list, osbal_rows, OSBAL_FACODE_COL)
                cert_matches = []
                for idx in matched_list:
                    ref_r = osbal_rows[idx]
                    r_certs = [_clean_cert_str(ref_r.get(c, "")) for c in sertif_osbal if _clean_cert_str(ref_r.get(c, ""))]
                    if any(sc in r_certs for sc in sus_cert_vals):
                        cert_matches.append(idx)
                    elif any(_cert_in_range(sc, rc) for sc in sus_cert_vals for rc in r_certs):
                        cert_matches.append(idx)
                cert_codes = _fac_codes_from_indices(cert_matches, osbal_rows, OSBAL_FACODE_COL) if cert_matches else set()
                if cert_matches and cert_codes and cert_codes < before_codes:
                    matched_list = cert_matches
                    if "Sertifikat" not in scenario:
                        scenario = f"{scenario} + Sertifikat"
                    fac_lookup_method = "SERTIFIKAT"

            # Step 3: Insured Token Matching
            # Insured tidak perlu diperiksa lagi bila FAC_CODE sudah tinggal satu.
            # Insured hanya boleh masuk skenario jika benar-benar mengurangi FAC_CODE.
            if len(_fac_codes_from_indices(matched_list, osbal_rows, OSBAL_FACODE_COL)) > 1:
                ins_vals = _collect_clean_values(sus, insured_sus)
                if ins_vals:
                    before_codes = _fac_codes_from_indices(matched_list, osbal_rows, OSBAL_FACODE_COL)
                    confirmed_ins = [
                        idx for idx in matched_list
                        if _ref_has_value(osbal_rows[idx], insured_osbal, ins_vals)
                    ]
                    confirmed_codes = _fac_codes_from_indices(confirmed_ins, osbal_rows, OSBAL_FACODE_COL) if confirmed_ins else set()
                    if confirmed_ins and confirmed_codes and confirmed_codes < before_codes:
                        matched_list = confirmed_ins
                        if "Insured" not in scenario:
                            scenario = f"{scenario} + Insured"
                        fac_lookup_method = "INSURED"

            # Step 4: Amount TIDAK digunakan untuk memilih FAC_CODE.
            # Amount tetap dipakai pada perhitungan/flagging output, tetapi tidak boleh
            # mengubah kandidat >1 FAC_CODE menjadi satu FAC_CODE. Ini penting agar
            # kandidat yang sama-sama valid berdasarkan Currency dan Receipt Date/FAC_COM_DATE
            # tetap berakhir sebagai Matching >1 fac code.

            matched_osbal_idx = set(matched_list)

        # FASE 5b: CROSS-SOURCE FAC_CODE NARROWING
        # Setelah seluruh key Data 2 selesai, bila masih >1 FAC_CODE lanjut ke:
        # Data 5 (Polis -> Slip -> Cert -> Insured), lalu Data 1 dengan urutan sama.
        # Keduanya hanya boleh melakukan intersection terhadap kandidat FAC_CODE
        # yang tersisa; tidak boleh menambah/mengganti FAC_CODE.
        if len(_fac_codes_from_indices(list(matched_osbal_idx), osbal_rows, OSBAL_FACODE_COL)) > 1:
            matched_osbal_idx = _cross_source_narrow_facodes(
                matched_osbal_idx, fac_lookup_method, sus, osbal_rows,
                data5_maps, facul_maps, polis_sus, slip_sus, sertif_sus, insured_sus
            )

        # FASE 6: CURRENCY -> PERIODE SELALU PALING AKHIR
        # Currency/period tidak boleh memilih FAC_CODE sebelum pencarian Data 2 ->
        # Data 5 -> Data 1 selesai. Bahkan bila FAC_CODE sudah tinggal satu, kedua
        # validasi ini tetap dijalankan untuk memastikan hasil make sense.
        is_beda_curr = False
        is_beda_per = False
        if matched_osbal_idx:
            matched_list = list(matched_osbal_idx)
            sus_curr = _normalize(sus.get(SUSPEND_CURR_COL, ""))
            if sus_curr:
                filtered_curr = [
                    i for i in matched_list
                    if _normalize(osbal_rows[i].get(OSBAL_CURR_COL, "")) == sus_curr
                ]
                if not filtered_curr:
                    is_beda_curr = True
                elif len(filtered_curr) < len(matched_list):
                    matched_list = filtered_curr

            if not is_beda_curr and pd.notna(sus_date):
                filtered_per = []
                for i in matched_list:
                    com_date = osbal_rows[i].get("_com_date_parsed")
                    if pd.isna(com_date) or sus_date >= com_date:
                        filtered_per.append(i)
                if not filtered_per:
                    is_beda_per = True
                elif len(filtered_per) < len(matched_list):
                    matched_list = filtered_per

            matched_osbal_idx = set(matched_list)

        osbal_count = len(matched_osbal_idx)
        osbal_ref = [osbal_rows[i] for i in matched_osbal_idx] if matched_osbal_idx else [{}]

        if scenario and " only + " in scenario:
            scenario = scenario.replace(" only + ", " + ")

        fac_codes = {_normalize(r.get(OSBAL_FACODE_COL, "")) for r in osbal_ref if r}
        fac_codes.discard("")

        raw_results.append({
            "suspend": sus, "source": source,
            "base_scenario": scenario,
            "is_beda_curr": is_beda_curr,
            "is_beda_per": is_beda_per,
            "osbal_idx": list(matched_osbal_idx),
            "osbal_count": osbal_count, "resolved": resolved, "fac_codes": fac_codes,
        })

    print(f"  Matching selesai dalam {time.perf_counter()-t4:.1f}s", flush=True)

    # [4c/6] Fac code claim & accumulation (V1: row asli dipertahankan)
    print("\n[4c/6] Fac code claim & accumulation (V1: row asli dipertahankan) ...", flush=True)

    # FAC_CODE tidak lagi dieliminasi karena sudah "claimed" oleh suspend row lain.
    # Resolusi multi-FAC hanya boleh terjadi lewat cascade Data 2 -> Data 5 -> Data 1 di atas.

    # Group per fac untuk shared osbal
    groups = {}
    final = []

    for r in raw_results:
        r.setdefault("suspend_count", 1)
        if (len(r["fac_codes"]) == 1 and r["source"] is not None
                and not r["is_beda_curr"] and not r["is_beda_per"]):
            groups.setdefault(next(iter(r["fac_codes"])), []).append(r)
        else:
            final.append(r)

    for fac, group in groups.items():
        if len(group) == 1:
            group[0]["suspend_count"] = 1
            final.append(group[0])
            continue
        merged_osbal_idx = set()
        for g in group:
            merged_osbal_idx.update(g["osbal_idx"])
        merged_sus = _merge_suspend_rows([g["suspend"] for g in group])
        for g in group:
            g["osbal_idx"]        = list(merged_osbal_idx)
            g["osbal_count"]      = len(merged_osbal_idx)
            g["suspend_count"]    = len(group)
            g["_merged_amount_ori"] = merged_sus.get("AMOUNT ORI", 0)
            final.append(g)

    print(f"  {len(raw_results):,} suspend rows -> {len(final):,} output rows (row asli dipertahankan)", flush=True)

    # [5] Build output DataFrame
    print(f"\n[5/6] Building output ({len(final):,} rows) ...", flush=True)
    output_rows = []
    for r in final:
        ref_osbal = [osbal_rows[i] for i in r["osbal_idx"]] if r["source"] else [{}]
        all_fac_osbal_rows = None
        osbal_cnt = r["osbal_count"]
        if r["source"] and len(r["fac_codes"]) == 1:
            fac = next(iter(r["fac_codes"]))
            all_fac_idx = facode_osbal_idx.get(fac, [])
            if len(all_fac_idx) > 0:
                all_fac_osbal_rows = [osbal_rows[i] for i in all_fac_idx]
                osbal_cnt = len(all_fac_idx)

        final_scenario = _format_final_scenario(
            base_scenario=r.get("base_scenario", ""),
            source=r["source"],
            is_beda_curr=r.get("is_beda_curr", False),
            is_beda_per=r.get("is_beda_per", False),
            num_fac=len(r["fac_codes"]),
        )

        out = _build_output_row(
            r["suspend"], source=r["source"], scenario=final_scenario,
            osbal_rows=ref_osbal, osbal_count=osbal_cnt,
            resolved=r["resolved"], suspend_count=r.get("suspend_count", 1),
            all_fac_osbal_rows=all_fac_osbal_rows,
        )
        if "_merged_amount_ori" in r:
            out["_MERGED_AMOUNT_ORI"] = r["_merged_amount_ori"]
        output_rows.append(out)

    df = pd.DataFrame(output_rows)

    # Fix format angka koma-desimal
    for col in ["AMOUNT ORI", "CCOS_OR_BAL", "CCOS_BAL_DUE"]:
        if col in df.columns:
            s = df[col]
            if s.dtype == object:
                s = s.astype(str).str.strip().str.replace(r'\s+', '', regex=True)
                has_comma = s.str.contains(',', na=False)
                s_eu = s.str.replace('.', '', regex=False).str.replace(',', '.', regex=False)
                s = s_eu.where(has_comma, s)
            df[col] = pd.to_numeric(s, errors="coerce")

    df = _compute_derived_cols(df)
    # OUTPUT ONLY: istilah skenario "Sertifikat" ditampilkan sebagai "Cert".
    if "SKENARIO" in df.columns:
        df["SKENARIO"] = df["SKENARIO"].astype(str).str.replace("Sertifikat", "Cert", regex=False)
    df.loc[df["SKENARIO"].isin(["UNMATCHED", "Unmatching"]), "FLAG_PROD"] = "Unmatching"
    df.loc[df["SKENARIO"] == "UNMATCHED", "SKENARIO"] = "Unmatching"

    # OUTPUT ONLY: kedua saldo dibuat blank untuk flag non-finansial tertentu.
    # Tidak mengubah proses/logic matching maupun FLAG_PROD.
    _blank_balance_flags = {
        "Unmatching", "Matching >1 fac code", "Beda Periode", "Beda Currency"
    }
    if "FLAG_PROD" in df.columns:
        _blank_mask = df["FLAG_PROD"].isin(_blank_balance_flags)
        for _balance_col in ("CCOS_BAL_DUE", "CCOS_OR_BAL"):
            if _balance_col in df.columns:
                df.loc[_blank_mask, _balance_col] = np.nan

    # [6] Export
    print(f"\n[6/6] Saving to: {OUTPUT_FILE} ...", flush=True)
    for col in FINAL_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    df = df[FINAL_COLUMNS]

    for c in df.select_dtypes(include=['object']).columns:
        df[c] = df[c].astype(str).str.replace(r'[\x00-\x08\x0B\x0C\x0E-\x1F]', '', regex=True)
    df.to_excel(OUTPUT_FILE, index=False)

    # FIX OUTPUT ONLY: cegah tampilan scientific notation (E+) di Excel.
    # Nilai identifier ditulis sebagai text; kolom nominal memakai format angka biasa.
    try:
        from decimal import Decimal, InvalidOperation

        wb_fmt = load_workbook(OUTPUT_FILE)
        ws_fmt = wb_fmt.active
        header_map = {cell.value: cell.column for cell in ws_fmt[1]}

        text_identifier_cols = [
            "CCOS_DOC_NO", "CCOS_REF_CODE", "RECEIPT NO", "CREDIT NOTES",
            "DETAIL RINCIAN NO", "POLIS", "POLICY_CLEAN_1", "POLICY_CLEAN_2",
            "POLICY_CLEAN_3", "POLICY_CLEAN_4", "POLICY_CLEAN_5",
            "SERTIF_CLEAN_1", "SERTIF_CLEAN_2", "SERTIF_CLEAN_3",
            "SLIP_NO", "SLIP_NO_CLEAN"
        ]

        for col_name in text_identifier_cols:
            col_idx = header_map.get(col_name)
            if not col_idx:
                continue
            for row_idx in range(2, ws_fmt.max_row + 1):
                cell = ws_fmt.cell(row=row_idx, column=col_idx)
                value = cell.value
                if value is None:
                    continue
                if isinstance(value, float) and value.is_integer():
                    value = format(value, ".0f")
                else:
                    value = str(value)
                    if "e" in value.lower():
                        try:
                            value = format(Decimal(value), "f")
                            if "." in value:
                                value = value.rstrip("0").rstrip(".")
                        except InvalidOperation:
                            pass
                cell.value = str(value)
                cell.number_format = "@"

        for col_name in ["AMOUNT ORI", "AMOUNT_ORI_MIN1", "AMOUNT PAY",
                         "CCOS_OR_BAL", "CCOS_BAL_DUE", "DIFERENCE"]:
            col_idx = header_map.get(col_name)
            if not col_idx:
                continue
            for row_idx in range(2, ws_fmt.max_row + 1):
                ws_fmt.cell(row=row_idx, column=col_idx).number_format = "#,##0.00"

        wb_fmt.save(OUTPUT_FILE)
    except Exception as e:
        print(f"  [WARN] Format anti-E+ gagal: {e}")

    try:
        from excel_styler import apply_purple_column_style
        apply_purple_column_style(OUTPUT_FILE, "matching")
    except Exception as e:
        print(f"  [WARN] Styling gagal: {e}")

    # Export ke PostgreSQL
    env_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env")
    load_dotenv(env_path)
    load_dotenv()
    print("  Exporting ke PostgreSQL ...")
    db_user = os.environ.get("DB_USER", "viy4user")
    db_pass = os.environ.get("DB_PASS", "fild42op61nx")
    db_host = os.environ.get("DB_HOST", "10.10.125.91")
    db_port = os.environ.get("DB_PORT", "5432")
    db_name = os.environ.get("DB_NAME", "fsi_db")
    conn_str = f"postgresql+psycopg2://{db_user}:{db_pass}@{db_host}:{db_port}/{db_name}"

    try:
        df_db = df.copy()
        df_db["RECEIPT DATE"] = pd.to_datetime(df_db["RECEIPT DATE"], errors="coerce").dt.date
        df_db["RECEIPT DATE"] = df_db["RECEIPT DATE"].where(df_db["RECEIPT DATE"].notnull(), None)

        numeric_cols = ["AMOUNT ORI", "AMOUNT_ORI_MIN1", "AMOUNT PAY", "CCOS_OR_BAL", "CCOS_BAL_DUE", "DIFERENCE"]
        for c in numeric_cols:
            df_db[c] = pd.to_numeric(df_db[c], errors="coerce")

        text_cols = [c for c in FINAL_COLUMNS if c not in numeric_cols and c != "RECEIPT DATE"]
        for c in text_cols:
            df_db[c] = df_db[c].astype(str).str.strip()
            df_db[c] = df_db[c].replace({"nan": None, "None": None, "": None})

        engine = create_engine(conn_str)
        with engine.begin() as con:
            con.execute(text('DELETE FROM "SUSPENSE_DATA_SUSPENSE_V1" WHERE "CEDANT NAME" = \'PT LIPPO GENERAL INSURANCE\';'))
            con.execute(text('DELETE FROM "SUSPENSE_CLEAN_2026" WHERE "CEDANT NAME" = \'PT LIPPO GENERAL INSURANCE\';'))
        df_db.to_sql("SUSPENSE_DATA_SUSPENSE_V1", con=engine, if_exists="append", index=False)
        df_db.to_sql("SUSPENSE_CLEAN_2026", con=engine, if_exists="append", index=False)
        print(f"      -> {len(df_db):,} rows berhasil di-export ke 'SUSPENSE_DATA_SUSPENSE_V1' & 'SUSPENSE_CLEAN_2026'")
    except Exception as e:
        print(f"  [WARN] PostgreSQL export gagal: {e}")

    elapsed = time.perf_counter() - t_start
    print(f"\n{'=' * 60}")
    print(f"  Selesai dalam {elapsed:.1f}s ({elapsed/60:.1f} menit)")
    print(f"  Rows  : {len(df):,}  |  Cols: {len(df.columns)}  |  File: {OUTPUT_FILE}")
    print(f"\n  FLAG_PROD:")
    for flag, count in df["FLAG_PROD"].value_counts().items():
        print(f"    {flag:<45}: {count:,}")
    print(f"\n  SKENARIO:")
    for sce, count in df["SKENARIO"].value_counts().items():
        print(f"    {sce:<50}: {count:,}")
    print("=" * 60)


if __name__ == "__main__":
    run()