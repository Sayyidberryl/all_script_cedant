"""
================================================================================
 SCRIPT CLEANSING DATA 2 - OSBAL (CEDANT: BRINS GENERAL INSURANCE)
================================================================================
 Changelog:
 - Tambah rule insured BIO FARMA campuran badan hukum
     "BIO FARMA/PT BRI ASURANSI INDONESIA/BIO FARMA PASTEUR"
     -> BIO FARMA | BRI ASURANSI INDONESIA | BIO FARMA PASTEUR
 - Tambah rule insured "PELABUHAN INDONESIA (PELINDO) HOLDING"
     -> PELABUHAN INDONESIA HOLDING
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
INPUT_FILE = "2b. transaksi Osbal.xlsx"
SHEET_NAME = "Query result"
HEADER_ROW = 0
OUTPUT_FILE = "brins_output_osbal_V2.xlsx"
CEDANT_FILTER = "BRINS"

CEDANT_COLUMN = "CCOS_COMP_NAME"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"

CLSDT_POLIS_COLUMN = "CLSDT_POLICY_NO"
CLSDT_SLIP_COLUMN = "CLSDT_SLIP_NO"
CLSDT_SERTF_COLUMN = "CLSDT_SERTF_NO"
CERTIFICATE_OUTPUT_COL = "CERTIFICATE"

MAX_BREAKDOWN_CODES = 5
TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CLSDT", "CERTIFICATE",
)

CERT_DIGIT_LEN = 6
CERT_RANGE_JOIN_THRESHOLD = 3
_BRINS_BASE_LEN_MIN, _BRINS_BASE_LEN_MAX = 13, 21

# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
_BYPASS_ALL_RE = re.compile(r"\bPENYELESAIAN\s+HUTANG\s+PIUTANG\b", flags=re.IGNORECASE)
_TRANSACTION_HEADER_RE = re.compile(r"^\s*(?:Pembayaran|Penagihan)?\s*(?:dan\s+Penagihan)?\s*Premi\s*(?:Fakultatif|SOA|Pensesian)?\s*(?:PAR\s*&?\s*EQ|BIT|GIT|CECR|QS)?\s*", flags=re.IGNORECASE)
_TRUNCATE_TRIGGER_RE = re.compile(r"\bsubsidiar|\bassociat|\baffiliat|\bfiliated\b|related\s+compan|respective\s+right|\bincluding\s+an[yd]|\bindu?ding\s+any|as\s+the\s+owner|as\s+owners?\b|as\s+main\s+contractor|\(as\s+mortgagee\s+and\s+loss\s+payee\)", flags=re.IGNORECASE)
_REMOVE_ONLY_RE = re.compile(r"\bas\s+principal\b|\bsebagai\s+principal\b|\bsebagai\s+kontraktor\b", flags=re.IGNORECASE)
_ONLY_WORD_RE = re.compile(r"\bONLY\b", flags=re.IGNORECASE)
_GENERIC_BOILERPLATE_TAIL_RE = re.compile(r"\bsub\s*-?\s*kontraktor\b.*$|\banak\s+perusahaan\b.*$|\bafiliasi\s+perusahaan\b.*$|\bdan\s+semua\b\s*[–-]?\s*$", flags=re.IGNORECASE)
_TITLE_RE = re.compile(r"\b(?:S\.?H\.?|S\.?I\.?Kom\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b", flags=re.IGNORECASE)
_LEGAL_ENTITY_RE = re.compile(r"\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO|Perseroan\s+Terbatas)\b", flags=re.IGNORECASE)
_ENTITY_SPLIT_RE = re.compile(r"\b(?:QQ|Q\.Q\.|CQ|C\.Q\.?)\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|\bin\s+this\s+case\b|\bdalam\s+hal\s+ini\b|,|/", flags=re.IGNORECASE)
_LEGAL_ENTITY_DASH_SIGNAL_RE = re.compile(r",\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO)\s*[\r\n]*\s*-\s*", flags=re.IGNORECASE)
_GENERIC_DASH_SEPARATOR_RE = re.compile(r"\s*[\r\n]+\s*-\s*|\s+-\s+")
_TRAILING_PAREN_MERGE_RE = re.compile(r",?\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?)?\s*\(([^)]+)\)\s*$", flags=re.IGNORECASE)
_INLINE_PAREN_MERGE_RE = re.compile(r",?\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?)?\s*\(([^)]+)\)", flags=re.IGNORECASE)
_QQ_PAREN_UNWRAP_RE = re.compile(r"\bQQ\s*\(\s*([^)]*)\)", flags=re.IGNORECASE)

def _unwrap_qq_paren(text: str) -> str:
    def _replace(match):
        inner = re.sub(r"\s+", " ", re.sub(r"[:;]", " ", match.group(1).strip())).strip()
        return f"QQ , {inner}" if inner else "QQ"
    return _QQ_PAREN_UNWRAP_RE.sub(_replace, text)

def _merge_trailing_paren(text: str) -> str:
    match = _TRAILING_PAREN_MERGE_RE.search(text)
    if not match: return text
    inner = match.group(1).strip()
    if re.search(r"\d", inner) or inner.strip().upper() == "PERSERODA":
        return text[: match.start()] + text[match.end():]
    return text[: match.start()] + text[match.end():] if not inner or len(inner.split()) > 4 or _TRUNCATE_TRIGGER_RE.search(inner) else text[: match.start()] + " " + inner

def _merge_inline_paren(text: str) -> str:
    def _replace(match):
        inner = match.group(1).strip()
        if re.search(r"\d", inner) or inner.strip().upper() == "PERSERODA":
            return ""
        return "" if not inner or len(inner.split()) > 4 or _TRUNCATE_TRIGGER_RE.search(inner) else f" {inner}"
    return _INLINE_PAREN_MERGE_RE.sub(_replace, text)

_MONTH_WORDS = ("JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI", "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER", "DES", "JANUARY", "FEBRUARY", "MARCH", "MAY", "JUNE", "JULY", "AUGUST", "OCTOBER", "DECEMBER", "JAN", "FEB", "MAR", "APR", "JUN", "JUL", "AUG", "SEP", "SEPT", "OCT", "NOV", "DEC")
_DATE_BATCH_TOKEN_RE = re.compile(r"\b(?:" + "|".join(_MONTH_WORDS) + r")\b|\b(?:19|20)\d{2}\b|\bBATCH\s*\d*\b", flags=re.IGNORECASE)

def _strip_date_batch_info(text: str) -> str: return _DATE_BATCH_TOKEN_RE.sub(" ", text)

def _truncate_from_first_trigger(text: str) -> str:
    match = _TRUNCATE_TRIGGER_RE.search(text)
    return text[: match.start()] if match else text

def _normalize_dash_entity_list(text: str) -> str: return _GENERIC_DASH_SEPARATOR_RE.sub(", ", text) if _LEGAL_ENTITY_DASH_SIGNAL_RE.search(text) else text
def _strip_transaction_header(text: str) -> str: return _TRANSACTION_HEADER_RE.sub("", text)
def _strip_titles(text: str) -> str: return _TITLE_RE.sub(" ", text)
def _strip_legal_entity(text: str) -> str: return _LEGAL_ENTITY_RE.sub(" ", text)

def _truncate_boilerplate_tail(text: str) -> str:
    text = _truncate_from_first_trigger(text)
    match_generic = _GENERIC_BOILERPLATE_TAIL_RE.search(text)
    if match_generic: text = text[: match_generic.start()]
    return _REMOVE_ONLY_RE.sub(" ", text)

def _final_polish(text: str) -> str:
    t = re.sub(r"\s+", " ", re.sub(r"[-/]", " ", re.sub(r"[().;:\"']", " ", text))).strip()
    dangling_pattern = r"^\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*|\s*\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$"
    previous = None
    while previous != t:
        previous, t = t, re.sub(dangling_pattern, "", t, flags=re.IGNORECASE).strip()
    return t.upper()

_KNOWN_ENTITIES = []
_KNOWN_ENTITIES_SORTED = sorted(_KNOWN_ENTITIES, key=len, reverse=True)
_KNOWN_ENTITY_RE = re.compile("|".join(re.escape(name) for name in _KNOWN_ENTITIES_SORTED)) if _KNOWN_ENTITIES_SORTED else re.compile(r"(?!)")
_ENTITY_SLASH_NORMALIZE_PATTERNS = [(name, re.compile(r"\b" + r"(?:\s*/\s*|\s+)".join(re.escape(w) for w in name.split()) + r"\b", flags=re.IGNORECASE)) for name in _KNOWN_ENTITIES_SORTED if len(name.split()) > 1]

def _normalize_entity_slash_variants(text: str) -> str:
    for canonical, pattern in _ENTITY_SLASH_NORMALIZE_PATTERNS: text = pattern.sub(canonical, text)
    return text

def _split_by_known_entities(text: str):
    pos, n, found = 0, len(text), []
    while pos < n:
        while pos < n and text[pos] == " ": pos += 1
        if pos >= n: break
        match = _KNOWN_ENTITY_RE.match(text, pos)
        if not match: return None
        found.append(match.group(0))
        pos = match.end()
    return found if len(found) > 1 else None

_INSURED_DESCRIPTIVE_RE = re.compile(r"\bVARIOUS\s+INSUREDS?\b|\bACCEPTED\s+BY\b", flags=re.IGNORECASE)
_INSURED_DASH_SPLIT_EXCEPTIONS = (
    "EBDI VAUNDRI - OSCAR OMEGA",
    "ADHI - HUTAMA - NINDYA - ABIPRAYA KSO",
    "PEMERINTAH PROVINSI SULAWESI SELATAN - DINAS BINA MARGA DAN BINA KONSTRUKSI",
)
_INSURED_DASH_MERGE_SUBSTRINGS = (
    "PP - WASKITA - WIJAYA KARYA",
)
_SUSPENSE_WORD_RE = re.compile(r"\bSUSPENSE\b", flags=re.IGNORECASE)

def _apply_dash_merge_substrings(text: str) -> str:
    for original_sub in _INSURED_DASH_MERGE_SUBSTRINGS:
        if original_sub.upper() in text.upper():
            merged = re.sub(r"\s*-\s*", " ", original_sub)
            idx = text.upper().find(original_sub.upper())
            if idx != -1:
                text = text[:idx] + merged + text[idx + len(original_sub):]
    return text

def _try_split_insured_by_dash(text: str):
    stripped = text.strip()
    if _SUSPENSE_WORD_RE.search(stripped): return None
    for exception in _INSURED_DASH_SPLIT_EXCEPTIONS:
        if stripped.upper() == exception.upper():
            return [p.strip() for p in re.split(r"\s*-\s*", stripped) if p.strip()]
    return None

_BIOFARMA_RE = re.compile(r"^\s*BIO\s*FARMA\b", flags=re.IGNORECASE)
_INSURED_CONJUNCTION_RE = re.compile(r"\s*(?:&|DAN)\s*", flags=re.IGNORECASE)
_PERURI_SYNONYM_RE = re.compile(r"PERUSAHAAN\s+UMUM\s+PERCETAKAN\s+UANG\s+REPUBLIK\s+INDONESIA\s*/\s*PERUM\s+PERURI", flags=re.IGNORECASE)

def _try_handle_peruri_synonym(text: str):
    if _PERURI_SYNONYM_RE.search(text): return ["PERUSAHAAN UMUM PERCETAKAN UANG REPUBLIK INDONESIA"]
    return None

_ROMAN_TO_ARABIC = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10}
_PELINDO_QQ_RE = re.compile(r"PELABUHAN\s+INDONESIA\s*,?\s*PT\s*QQ\s*PELINDO\s+([IVX]+(?:\s*,\s*[IVX]+)*)", flags=re.IGNORECASE)

def _try_handle_pelindo(text: str):
    match = _PELINDO_QQ_RE.search(text)
    if not match: return None
    romans = [r.strip().upper() for r in match.group(1).split(",") if r.strip()]
    results = []
    for r in romans:
        if r in _ROMAN_TO_ARABIC:
            val = f"PELABUHAN INDONESIA {_ROMAN_TO_ARABIC[r]}"
            if val not in results: results.append(val)
    return results if results else None

# --- [BARU] PELABUHAN INDONESIA (PELINDO) HOLDING -> PELABUHAN INDONESIA HOLDING ---
# Kurung "(PELINDO)" adalah alias singkatan, bukan bagian nama -> dibuang.
# Pola sengaja di-anchor (^...$) supaya tidak menyentuh varian insured lain.
_PELINDO_HOLDING_RE = re.compile(
    r"^\s*(?:PT\.?\s+)?PELABUHAN\s+INDONESIA\s*\(\s*PELINDO\s*\)\s*HOLDING(?:\s*\(?\s*PERSERO\s*\)?)?\s*$",
    flags=re.IGNORECASE,
)

def _try_handle_pelindo_holding(text: str):
    if _PELINDO_HOLDING_RE.match(text):
        return ["PELABUHAN INDONESIA HOLDING"]
    return None

# --- [BARU] BIO FARMA campuran dengan entitas berbadan hukum (PT/CV) ---
# Segmen yang diawali PT/CV = entitas independen, bukan lanjutan nama BIO FARMA.
_LEGAL_ENTITY_START_RE = re.compile(r"^\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?)(?=\s)", flags=re.IGNORECASE)

def _try_split_biofarma_mixed_entities(segments: list):
    """
    BIO FARMA/PT BRI ASURANSI INDONESIA/BIO FARMA PASTEUR
    -> segmen baru dimulai saat ketemu 'PT/CV' atau 'BIO FARMA'; selain itu disambung ke segmen sebelumnya.
    Return None jika tidak ada segmen berawalan badan hukum (biar logic lama tetap jalan).
    """
    if not any(_LEGAL_ENTITY_START_RE.match(s) for s in segments[1:]):
        return None

    groups = []
    for i, seg in enumerate(segments):
        if i == 0 or _LEGAL_ENTITY_START_RE.match(seg) or _BIOFARMA_RE.match(seg):
            groups.append(seg)
        else:
            groups[-1] = f"{groups[-1]} {seg}"

    results = []
    for g in groups:
        cleaned = _final_polish(_strip_legal_entity(g))
        if cleaned and cleaned not in results:
            results.append(cleaned)
    return results

def _try_handle_biofarma(text: str):
    if not _BIOFARMA_RE.match(text): return None
    segments = [s.strip() for s in text.split("/") if s.strip()]
    if len(segments) < 2: return [text.strip()]

    # [BARU] Rule campuran BIO FARMA + PT/CV (dicek sebelum logic ekspansi lama)
    mixed = _try_split_biofarma_mixed_entities(segments)
    if mixed is not None:
        return mixed if 0 < len(mixed) <= MAX_BREAKDOWN_CODES else [text.strip()]

    first_base, rest_segments = segments[0], segments[1:]
    expand_idx = next((i for i, seg in enumerate(rest_segments) if _INSURED_CONJUNCTION_RE.search(seg)), None)

    if expand_idx is None:
        combined = " ".join([first_base] + rest_segments)
        return [re.sub(r"\s+", " ", combined).strip()]

    expand_seg = rest_segments[expand_idx]
    parts = [p.strip() for p in _INSURED_CONJUNCTION_RE.split(expand_seg) if p.strip()]
    if len(parts) < 2: return [text.strip()]

    prefix_words = parts[0].split()
    shared_prefix = " ".join(prefix_words[:-1]) if len(prefix_words) > 1 else ""
    last_variant = prefix_words[-1] if prefix_words else ""
    pre_expand_segments = rest_segments[:expand_idx]
    expand_seg_has_own_base = bool(_BIOFARMA_RE.match(expand_seg))

    if expand_seg_has_own_base:
        independent_insureds = [first_base]
        full_prefix = shared_prefix if shared_prefix else parts[0]
    elif pre_expand_segments:
        expand_base = " ".join(pre_expand_segments).strip()
        independent_insureds = [first_base]
        full_prefix = f"{expand_base} {shared_prefix}".strip() if shared_prefix else expand_base
    else:
        independent_insureds = []
        full_prefix = f"{first_base} {shared_prefix}".strip() if shared_prefix else first_base

    full_prefix = re.sub(r"\s+", " ", full_prefix).strip()
    expanded = []
    for idx, p in enumerate(parts):
        val = f"{full_prefix} {last_variant}".strip() if idx == 0 else f"{full_prefix} {p}".strip()
        val = re.sub(r"\s+", " ", val).strip()
        if val not in expanded: expanded.append(val)

    result = independent_insureds + expanded
    return [r for i, r in enumerate(result) if r not in result[:i]]

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return []
    original = text.strip()

    res = _try_handle_peruri_synonym(original)
    if res is not None: return res
    
    res = _try_handle_pelindo(original)
    if res is not None: return res

    res = _try_handle_pelindo_holding(original)  # [BARU]
    if res is not None: return res
    
    res = _try_handle_biofarma(original)
    if res is not None: return res
    
    if _BYPASS_ALL_RE.search(original) or _INSURED_DESCRIPTIVE_RE.search(original): return [original]
    if _SUSPENSE_WORD_RE.search(original): return [original]

    dash_split = _try_split_insured_by_dash(original)
    if dash_split is not None:
        results = []
        for part in dash_split:
            cleaned = _final_polish(_strip_date_batch_info(_strip_legal_entity(_truncate_boilerplate_tail(part))))
            if cleaned and cleaned not in results: results.append(cleaned)
        return [original] if len(results) > MAX_BREAKDOWN_CODES or not results else results

    original_for_processing = _apply_dash_merge_substrings(original)
    t = _truncate_boilerplate_tail(_normalize_entity_slash_variants(_normalize_dash_entity_list(_strip_titles(_strip_transaction_header(_merge_inline_paren(_merge_trailing_paren(_unwrap_qq_paren(_ONLY_WORD_RE.sub(" ", original_for_processing)))))))))

    raw_parts, results = _ENTITY_SPLIT_RE.split(t), []
    for part in raw_parts:
        cleaned = _final_polish(_strip_date_batch_info(_strip_legal_entity(_truncate_boilerplate_tail(part))))
        if not cleaned or len(cleaned) < 2 or not re.search(r"[A-Z]", cleaned): continue

        known_split = _split_by_known_entities(cleaned)
        if known_split is not None:
            for name in known_split:
                if name not in results: results.append(name)
            continue
        if cleaned not in results: results.append(cleaned)

    return [original] if len(results) > MAX_BREAKDOWN_CODES else results

# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO
# ==============================================================================
_MAIN_LEN_MIN, _MAIN_LEN_MAX = _BRINS_BASE_LEN_MIN, _BRINS_BASE_LEN_MAX
_MAIN_DIGITS_PATTERN = rf"\d{{{_MAIN_LEN_MIN},{_MAIN_LEN_MAX}}}"
_TOK_MAIN_RE = re.compile(rf"^{_MAIN_DIGITS_PATTERN}$")
_TOK_CERT_RE = re.compile(r"^\d{6}$")
_TOK_SUFFIX_REPLACE_RE = re.compile(r"^\d{1,11}$")
_TOK_ENDORSE_RE = re.compile(r"^\d{1,2}[/\x00]\d{1,2}$")
_TOK_PLACEHOLDER_RE = re.compile(r"^P\d{1,2}$", flags=re.IGNORECASE)
_TOK_STATUS_RE = re.compile(r"^(?:New|End(?:orsement)?|Cancel\s*All|Adjustment|Cancellation|Reinstatement|TBA\.?|Attachment|DN)$", flags=re.IGNORECASE)
_VARIOUS_WORD_RE = re.compile(r"\bVAR(?:I|IOUS)?\b", flags=re.IGNORECASE)
_SD_WORD_RE = re.compile(r"\bSD\b", flags=re.IGNORECASE)
_AMPERSAND_SUFFIX_AMBIGUOUS_RE = re.compile(r"\d+\s*&\s*\d+")
_HAS_ANOMALOUS_CHAR_RE = re.compile(r"[^\w\s+\-,/.&()]")
_CURRENCY_RE = re.compile(r"\b(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\b", flags=re.IGNORECASE)
_SEE_ATTACHMENT_RE = re.compile(r"\bSEE\s+ATTACHMENT\b", flags=re.IGNORECASE)
_PENDING_SLIP_RE = re.compile(r"\bPENDING\s+SLIP\b", flags=re.IGNORECASE)
_NO_LABEL_RE = re.compile(r"\bNO\s*[:.]", flags=re.IGNORECASE)
_LETTER_REF_CODE_RE = re.compile(r"\b(?:\d{1,6}\s*/\s*)?[A-Z]{1,6}-[A-Z]{1,6}\s*/\s*[IVXLCM]{1,6}\s*/\s*\d{2,4}\b", flags=re.IGNORECASE)
_BATCH_TOKEN_RE = re.compile(r"\bBATCH\s*\d*\b", flags=re.IGNORECASE)
_BORDERO_WORD_RE = re.compile(r"\bBORDERO\b|\bBORD\b", flags=re.IGNORECASE)
_POLIS_MONTH_RE = re.compile(r"\b(?:" + "|".join(_MONTH_WORDS) + r")\.?\s*(?:\d{2,4})?\b", flags=re.IGNORECASE)
_POLIS_DATE_SLASH_RE = re.compile(r"\b\d{1,2}/\d{1,2}/\d{2,4}\b")
_POLIS_YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")
_POLIS_CN_RE = re.compile(r"\bCN\b", flags=re.IGNORECASE)
_POLIS_NARRATIVE_WORDS_RE = re.compile(r"\bSANY\b|\bNO\s+SLIP\b|\bNO\s+POLIS\b|\bPOLIS\s+VARIOUS\b", flags=re.IGNORECASE)
_POLIS_STATUS_WORDS_RE = re.compile(r"\bCancel\s*All\b|\bNew\b|\bEnd(?:orsement)?\b|\bAdjustment\b|\bCancellation\b|\bReinstatement\b", flags=re.IGNORECASE)
_EDGE_SEPARATOR_RE = re.compile(r"^[\s\-/,]+|[\s\-/,]+$")
_NARRATIVE_PREFIX_RE = re.compile(r"^\s*(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\s+\d{4}\s*-\s*(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)?\s*/?\s*", flags=re.IGNORECASE)
_MAIN_CODE_FIRST_RE = re.compile(_MAIN_DIGITS_PATTERN)

def _strip_leading_narrative(text: str) -> str:
    match = _MAIN_CODE_FIRST_RE.search(text)
    return text[match.start():] if match and re.search(r"[A-Za-z]", text[: match.start()]) else text

def _strip_polis_slip_noise(text: str) -> str:
    for pat in (_NARRATIVE_PREFIX_RE, _LETTER_REF_CODE_RE, _NO_LABEL_RE, _PENDING_SLIP_RE, _SEE_ATTACHMENT_RE, _POLIS_NARRATIVE_WORDS_RE, _POLIS_STATUS_WORDS_RE, _BATCH_TOKEN_RE, _BORDERO_WORD_RE, _POLIS_DATE_SLASH_RE, _POLIS_MONTH_RE, _POLIS_YEAR_RE, _POLIS_CN_RE, _CURRENCY_RE):
        text = pat.sub(" ", text)
    return _strip_leading_narrative(_EDGE_SEPARATOR_RE.sub("", re.sub(r"\s+", " ", text).strip()).strip())

_STRONG_SPLIT_CAPTURE_RE = re.compile(r"(\+|-|[,/]|\s+)")

def _split_preserving_continuity(t: str):
    segments, is_soft_list, pending_soft = [], [], False
    for part in _STRONG_SPLIT_CAPTURE_RE.split(t):
        if part in ("+", "-", ","):
            pending_soft = True
        elif part is not None and part.strip() == "" and _STRONG_SPLIT_CAPTURE_RE.fullmatch(part or ""): continue
        elif part is not None and _STRONG_SPLIT_CAPTURE_RE.fullmatch(part):
            pending_soft = False
        else:
            seg = (part or "").strip()
            if seg:
                segments.append(seg)
                is_soft_list.append(pending_soft)
                pending_soft = False
    return segments, is_soft_list

_ALPHA_WORD_RE = re.compile(r"[A-Za-z][A-Za-z0-9]*")
_ALLOWED_ALPHA_TOKENS = {"NEW", "END", "ENDORSEMENT", "CANCEL", "ALL", "ADJUSTMENT", "CANCELLATION", "REINSTATEMENT", "VARIOUS", "VAR", "VARI", "SD", "ATTACHMENT", "TBA", "DN"}
_FALLBACK_NOISE_WORDS_RE = re.compile(r"\bTBA\.?\b|\bATTACHMENT\b|\bDN\b", flags=re.IGNORECASE)

_DASH_CONCAT_MAX_BASE_LEN = 14
_DASH_CONCAT_RE = re.compile(rf"^(\d{{1,{_DASH_CONCAT_MAX_BASE_LEN}}})-(\d{{1,4}})(?=[+\s]|$)")
_HASH_SUFFIX_NOISE_RE = re.compile(r"\s*#\s*\d+\s*")
_FRACTION_SUFFIX_RE = re.compile(r"-\s*\d{1,2}\s*/\s*\d{1,2}\s*$")
_SPLIT_BASE_REJOIN_RE = re.compile(rf"^(\d{{1,{_BRINS_BASE_LEN_MAX - 1}}})\s+(\d{{1,{_BRINS_BASE_LEN_MAX - 1}}})(?=-)")
_SD_SLASH_NORMALIZE_RE = re.compile(r"\bS\s*/\s*D\b", flags=re.IGNORECASE)
_ALL_SPACED_DIGITS_RE = re.compile(r"^\s*\d+(?:\s+\d+)+\s*$")

def _try_rejoin_all_spaced_digits(text: str) -> str:
    if not _ALL_SPACED_DIGITS_RE.match(text): return text
    combined = re.sub(r"\s+", "", text)
    if _BRINS_BASE_LEN_MIN <= len(combined) <= _BRINS_BASE_LEN_MAX: return combined
    return text

_BASE_SPACE_SUFFIX_RE = re.compile(rf"(\d{{{_BRINS_BASE_LEN_MIN},{_BRINS_BASE_LEN_MAX}}}(?:\+\d{{1,6}})*)\s+(\d{{1,6}})(?=[\s/+-]|$)")
_BASE_DASH_SUFFIX_RE = re.compile(rf"(\d{{{_DASH_CONCAT_MAX_BASE_LEN + 1},{_BRINS_BASE_LEN_MAX}}}(?:\+\d{{1,6}})*)-(\d{{1,6}})(?=[\s/+-]|$)")

def _normalize_base_space_suffix(text: str) -> str:
    prev = None
    while prev != text:
        prev = text
        text = _BASE_SPACE_SUFFIX_RE.sub(lambda m: f"{m.group(1)}+{m.group(2)}", text, count=1)
        text = _BASE_DASH_SUFFIX_RE.sub(lambda m: f"{m.group(1)}+{m.group(2)}", text, count=1)
    return text

def _strip_brins_specific_noise(text: str) -> str:
    t = _HASH_SUFFIX_NOISE_RE.sub("", text)
    t = _FRACTION_SUFFIX_RE.sub("", t).strip()
    t = _try_rejoin_all_spaced_digits(t)
    m = _SPLIT_BASE_REJOIN_RE.match(t)
    if m:
        part1, part2 = m.groups()
        combined = part1 + part2
        if _BRINS_BASE_LEN_MIN <= len(combined) <= _BRINS_BASE_LEN_MAX:
            t = combined + t[m.end():]
    return _normalize_base_space_suffix(t).strip()

def _try_dash_concat_normalize(text: str) -> str:
    m = _DASH_CONCAT_RE.match(text)
    if not m: return text
    base, first_suffix = m.groups()
    return base + first_suffix + text[m.end():]

def _clean_fallback_original(original: str) -> str:
    t = _FALLBACK_NOISE_WORDS_RE.sub("", original)
    t = re.sub(r"[+\-/,]\s*$", "", t.strip()).strip()
    return re.sub(r"\s{2,}", " ", t)

def _has_descriptive_narrative(text: str) -> bool:
    return any(not (_TOK_PLACEHOLDER_RE.match(w) or w.upper() in _ALLOWED_ALPHA_TOKENS) for w in _ALPHA_WORD_RE.findall(text))

def _split_sd_range(text: str):
    if not _SD_WORD_RE.search(text): return None
    parts = _SD_WORD_RE.split(text, maxsplit=1)
    if len(parts) != 2: return "__FALLBACK__"

    left_all, right_all = re.findall(r"\d+", parts[0]), re.findall(r"\d+", parts[1])
    if not left_all or not right_all: return "__FALLBACK__"

    known_suffix_len = None
    left_resolved = left_all[-1]
    prepend_base = None
    if not _TOK_MAIN_RE.match(left_resolved):
        base = next((tok for tok in left_all if _TOK_MAIN_RE.match(tok)), None)
        if base is None or len(left_resolved) >= len(base): return "__FALLBACK__"
        known_suffix_len = len(left_resolved)
        prepend_base = base
        left_resolved = base[: -len(left_resolved)] + left_resolved

    right_first = right_all[0]
    if not _TOK_MAIN_RE.match(right_first):
        if len(right_first) >= len(left_resolved): return "__FALLBACK__"
        known_suffix_len = len(right_first)
        right_first = left_resolved[: -len(right_first)] + right_first

    expanded = _expand_numeric_range(left_resolved, right_first, known_suffix_len=known_suffix_len)
    if not expanded: return "__FALLBACK__"
    result = expanded + [t for t in right_all[1:] if _TOK_MAIN_RE.match(t) and t not in expanded]
    if prepend_base and prepend_base not in result: result = [prepend_base] + result
    return result

def _expand_numeric_range(main_a: str, main_b: str, known_suffix_len: int = None):
    if len(main_a) != len(main_b) or not (main_a.isdigit() and main_b.isdigit()): return None
    diff_len = known_suffix_len if known_suffix_len is not None else next((i for i in range(1, len(main_a) + 1) if main_a[-i] != main_b[-i]), 0)
    if diff_len == 0 or main_a[:-diff_len] != main_b[:-diff_len]: return None
    try:
        start, end = int(main_a[-diff_len:]), int(main_b[-diff_len:])
    except ValueError: return None
    return [f"{main_a[:-diff_len]}{str(n).zfill(diff_len)}" for n in range(start, end + 1)] if start <= end and (end - start + 1) <= MAX_BREAKDOWN_CODES else None

def _try_reconstruct_suffix23(tokens: list, tokens_is_soft: list):
    if not tokens or not _TOK_MAIN_RE.match(tokens[0]) or not all(_TOK_MAIN_RE.match(t) or _TOK_SUFFIX_REPLACE_RE.match(t) for t in tokens): return None
    if any(_TOK_SUFFIX_REPLACE_RE.match(tok) and not _TOK_MAIN_RE.match(tok) and not soft for tok, soft in zip(tokens, tokens_is_soft)): return None
    if not any(_TOK_SUFFIX_REPLACE_RE.match(tok) and not _TOK_MAIN_RE.match(tok) for tok in tokens): return None

    codes, current_base = [], None
    for tok in tokens:
        if _TOK_MAIN_RE.match(tok):
            current_base = tok
            if current_base not in codes: codes.append(current_base)
        elif current_base:
            rec = current_base[:-len(tok)] + tok
            if rec not in codes:
                codes.append(rec)
    return codes

_BASE_SUFFIX_BASE_RE = re.compile(rf"^({_MAIN_DIGITS_PATTERN})-(\d{{1,11}})\s*-\s*({_MAIN_DIGITS_PATTERN})$")

def _try_skip_ambiguous_middle_suffix(text: str):
    match = _BASE_SUFFIX_BASE_RE.match(text.strip())
    if not match: return None
    base1, _suffix_ignored, base2 = match.groups()
    return [base1] if base1 == base2 else [base1, base2]

_E_HASH_SUFFIX_RE = re.compile(rf"^({_MAIN_DIGITS_PATTERN})\s*-\s*E#\s*\d+\s*$", flags=re.IGNORECASE)

def _try_strip_e_hash_suffix(text: str):
    match = _E_HASH_SUFFIX_RE.match(text.strip())
    if not match: return None
    return [match.group(1)]

_SHORT_BASE_MULTI_SUFFIX_RE = re.compile(r"^(\d{9,11})-(\d{1,6}(?:\s*[,+]\s*\d{1,6})+)\s*$")

def _try_expand_short_base_multi_suffix(text: str):
    match = _SHORT_BASE_MULTI_SUFFIX_RE.match(text.strip())
    if not match: return None
    base, suffix_blob = match.groups()
    suffixes = [s.strip() for s in re.split(r"[,+]", suffix_blob) if s.strip()]
    if not suffixes: return None
    codes = [base]
    for suf in suffixes:
        if len(suf) < len(base): codes.append(base[: -len(suf)] + suf)
    return [c for i, c in enumerate(codes) if c not in codes[:i]] if len(codes) <= MAX_BREAKDOWN_CODES else None

_HAS_VALID_MAIN_DIGITS_RE = re.compile(_MAIN_DIGITS_PATTERN)

def clean_split_code(text: str) -> list:
    codes, _is_valid = _clean_split_code_impl(text)
    return codes

def _clean_split_code_impl(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return [], False
    raw_original = text.strip()
    original = _SD_SLASH_NORMALIZE_RE.sub("SD", raw_original)

    if _BYPASS_ALL_RE.search(original): return [raw_original], False
    original = _strip_brins_specific_noise(original)
    if not original: return [raw_original], False

    original = _try_dash_concat_normalize(original)
    
    res = _try_strip_e_hash_suffix(original)
    if res is not None: return res, True
    
    res = _try_expand_short_base_multi_suffix(original)
    if res is not None: return res, True
    
    if not _HAS_VALID_MAIN_DIGITS_RE.search(original): return [raw_original], False

    t = _strip_polis_slip_noise(original)
    if not t or _has_descriptive_narrative(t) or _AMPERSAND_SUFFIX_AMBIGUOUS_RE.search(t) or _HAS_ANOMALOUS_CHAR_RE.search(t): return [raw_original], False

    sd_result = _split_sd_range(t)
    if sd_result == "__FALLBACK__": return [raw_original], False
    if sd_result is not None: return sd_result, True

    t = re.sub(r"(-\s*\d{1,2})\s*/\s*(\d{1,2})\b", lambda m: f"{m.group(1)}\x00{m.group(2)}", re.sub(r"[\u2012\u2013\u2014\u2015]", "-", _VARIOUS_WORD_RE.sub(" ", t)))
    strong_segments, segment_is_soft = _split_preserving_continuity(t)

    def _is_droppable_token(seg_clean: str) -> bool:
        return bool(_TOK_PLACEHOLDER_RE.match(seg_clean) or _TOK_STATUS_RE.match(seg_clean))

    filtered_pairs = [(re.sub(r"[.]", "", seg).strip(), soft) for seg, soft in zip(strong_segments, segment_is_soft) if seg and not _is_droppable_token(re.sub(r"[.]", "", seg).strip())]
    if not filtered_pairs: return [raw_original], False

    tokens, tokens_is_soft = [p[0] for p in filtered_pairs], [p[1] for p in filtered_pairs]
    reconstructed = _try_reconstruct_suffix23(tokens, tokens_is_soft)
    if reconstructed is not None:
        return ([_clean_fallback_original(raw_original)], False) if len(reconstructed) > MAX_BREAKDOWN_CODES else (reconstructed, True)

    codes, current_base = [], None
    for tok, soft in zip(tokens, tokens_is_soft):
        if current_base and soft and re.match(r"^\d{1,2}\x00\d{1,2}$", tok): continue
        if current_base and soft and _TOK_SUFFIX_REPLACE_RE.match(tok) and not _TOK_MAIN_RE.match(tok) and len(tok) < len(current_base):
            rec = current_base[: -len(tok)] + tok
            if rec not in codes: codes.append(rec)
            continue

        if "\x00" in tok and not _TOK_MAIN_RE.match(tok):
            base_match = re.match(rf"^({_MAIN_DIGITS_PATTERN})-?(\d{{1,2}})\x00(\d{{1,2}})$", tok)
            if base_match:
                current_base = base_match.group(1)
                if current_base not in codes: codes.append(current_base)
                continue

        anchor, *suffix_parts = [p.replace("\x00", "/") for p in tok.split("-")]
        if not _TOK_MAIN_RE.match(anchor): continue
        current_base = anchor

        while suffix_parts and _TOK_STATUS_RE.match(suffix_parts[-1]): suffix_parts.pop()
        for p in suffix_parts:
            if _TOK_ENDORSE_RE.match(p.replace("/", "\x00")) or re.match(r"^\d{1,2}/\d{1,2}$", p): continue
            if _TOK_SUFFIX_REPLACE_RE.match(p) and len(p) < len(anchor):
                rec = anchor[: -len(p)] + p
                if rec not in codes: codes.append(rec)
        if anchor not in codes: codes.append(anchor)

    if not codes or len(codes) > MAX_BREAKDOWN_CODES: return [_clean_fallback_original(raw_original)], False
    return codes, True

# ==============================================================================
# 2B. CLEANSING CERTIFICATE
# ==============================================================================
_CERT_SD_NORMALIZE_RE = re.compile(r"\bS\s*/\s*D\b", flags=re.IGNORECASE)
_CERT_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9])(0\d{4,5})(?![A-Za-z0-9])")
_CERT_SD_PAIR_RE = re.compile(r"(0\d{4,5})\s+SD\s+(0\d{4,5})", flags=re.IGNORECASE)

def _pad_cert_digits(digits: str) -> str:
    return digits.zfill(CERT_DIGIT_LEN)

_MANUAL_CERT_OVERRIDE_BASIS = "1114031022000018"
_MANUAL_CERT_OVERRIDE_RE = re.compile(rf"^\s*{_MANUAL_CERT_OVERRIDE_BASIS}\s+70\s+SD\s+85\s*$", flags=re.IGNORECASE)

def try_manual_cert_override(text: str):
    if not text: return None
    normalized = _SD_SLASH_NORMALIZE_RE.sub("SD", text.strip())
    if _MANUAL_CERT_OVERRIDE_RE.match(normalized): return _MANUAL_CERT_OVERRIDE_BASIS, "000070 SD 000085"
    return None

def detect_polis_cert_groups(text: str):
    if not text: return None
    tokens = re.findall(r"\d+|SD", text, flags=re.IGNORECASE)
    groups = []
    current_polis, current_certs, current_base_cert, pending_sd = None, [], None, False

    for tok in tokens:
        if tok.upper() == "SD":
            pending_sd = True
            continue

        n = len(tok)
        if _BRINS_BASE_LEN_MIN <= n <= _BRINS_BASE_LEN_MAX:
            if current_polis is not None: groups.append((current_polis, current_certs))
            current_polis, current_certs, current_base_cert, pending_sd = tok, [], None, False
            continue

        if current_polis is None: continue

        if pending_sd:
            if current_base_cert is None:
                pending_sd = False
                continue
            endpoint = _pad_cert_digits(tok) if n in (5, 6) and tok.startswith("0") else (current_base_cert[:-n] + tok if 1 <= n <= 4 else None)
            if endpoint is not None:
                try:
                    start_n, end_n = int(current_base_cert), int(endpoint)
                except ValueError:
                    start_n, end_n = None, None
                if start_n is not None and start_n <= end_n:
                    for x in range(start_n, end_n + 1):
                        c = _pad_cert_digits(str(x))
                        if c not in current_certs: current_certs.append(c)
                current_base_cert = endpoint
            pending_sd = False
        elif n in (5, 6) and tok.startswith("0"):
            padded = _pad_cert_digits(tok)
            current_certs.append(padded)
            current_base_cert = padded
        elif 1 <= n <= 4 and current_base_cert is not None:
            endpoint = current_base_cert[:-n] + tok
            try:
                start_n, end_n = int(current_base_cert), int(endpoint)
            except ValueError:
                start_n, end_n = None, None
            if start_n is not None and start_n <= end_n:
                for x in range(start_n, end_n + 1):
                    c = _pad_cert_digits(str(x))
                    if c not in current_certs: current_certs.append(c)
            else:
                if endpoint not in current_certs: current_certs.append(endpoint)
            current_base_cert = endpoint

    if current_polis is not None: groups.append((current_polis, current_certs))
    if not groups or not any(certs for _polis, certs in groups): return None
    return groups

def _extract_cert_numbers_from_raw(text: str):
    if not text: return []
    found, consumed_spans = [], []

    for m in _CERT_SD_PAIR_RE.finditer(text):
        start_s, end_s = m.group(1), m.group(2)
        if not re.search(r"[A-Za-z0-9]", text[: m.start(1)]): continue
        try:
            start_n, end_n = int(start_s), int(end_s)
        except ValueError: continue
        if start_n > end_n or (end_n - start_n + 1) > MAX_BREAKDOWN_CODES: continue
        for n in range(start_n, end_n + 1): found.append(_pad_cert_digits(str(n)))
        consumed_spans.append((m.start(), m.end()))

    if consumed_spans:
        masked_chars = list(text)
        for s, e in consumed_spans:
            for i in range(s, e):
                masked_chars[i] = "#"
        scan_text = "".join(masked_chars)
    else:
        scan_text = text
        
    for m in _CERT_TOKEN_RE.finditer(scan_text):
        if re.search(r"[A-Za-z0-9]", scan_text[: m.start()]):
            found.append(_pad_cert_digits(m.group(1)))
    return found

def _single_certificate_from_clsdt(clsdt_sertf_value) -> str:
    if clsdt_sertf_value is None or (isinstance(clsdt_sertf_value, float) and pd.isna(clsdt_sertf_value)): return ""
    text = str(clsdt_sertf_value).strip()
    if not text or text.lower() in ("nan", "none", "-"): return ""

    text = _CERT_SD_NORMALIZE_RE.sub("SD", text)
    _, is_valid_polis_pattern = _clean_split_code_impl(text)
    if is_valid_polis_pattern: return ""
    if re.fullmatch(r"0\d{4,5}", text): return _pad_cert_digits(text)

    cert_candidates = _extract_cert_numbers_from_raw(text)
    return cert_candidates[0] if cert_candidates else ""

def _join_certificates(cert_list: list) -> str:
    if not cert_list: return ""
    seen = [c for i, c in enumerate(cert_list) if c not in cert_list[:i]]
    if len(seen) == 1: return seen[0]

    try:
        nums = [int(c) for c in seen]
    except ValueError: return ", ".join(seen)

    segments, current_seg = [], [seen[0]]
    for i in range(1, len(seen)):
        if nums[i] - nums[i - 1] == 1: current_seg.append(seen[i])
        else: segments.append(current_seg); current_seg = [seen[i]]
    segments.append(current_seg)

    parts = []
    for seg in segments:
        if len(seg) > CERT_RANGE_JOIN_THRESHOLD: parts.append(f"{seg[0]} SD {seg[-1]}")
        else: parts.extend(seg)
    return ", ".join(parts)

def build_certificates_for_row(clsdt_sertf_value, fac_policy_value, num_polis: int):
    slots = [""] * max(num_polis, 1)
    
    clsdt_cert = _single_certificate_from_clsdt(clsdt_sertf_value)
    if clsdt_cert:
        slots[0] = clsdt_cert
        return slots

    if _is_blank_value(fac_policy_value): return slots
    fac_text = _CERT_SD_NORMALIZE_RE.sub("SD", _strip_brins_specific_noise(str(fac_policy_value).strip()))

    groups = detect_polis_cert_groups(fac_text)
    if groups is not None:
        for i, (_polis, certs) in enumerate(groups):
            if i < len(slots): slots[i] = _join_certificates(certs)
        return slots

    single_candidates = _extract_cert_numbers_from_raw(fac_text)
    if single_candidates:
        slots[0] = _join_certificates(single_candidates)
    return slots

# ==============================================================================
# 3. PEMILIHAN SUMBER TERBAIK: CLSDT vs FAC (POLIS & SLIP)
# ==============================================================================
_PENYELESAIAN_RE = re.compile(r"\bPENYELESAIAN\s+(?:HUTANG|UTANG)\s+PIUTANG\b", flags=re.IGNORECASE)
_TRAILING_CURRENCY_WORD_RE = re.compile(r"\s+(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\s*$", flags=re.IGNORECASE)

def _is_blank_value(value) -> bool:
    if value is None or (isinstance(value, float) and pd.isna(value)): return True
    s = str(value).strip()
    return s == "" or s.lower() in ("nan", "none")

def _is_valid_clean_result(codes: list) -> bool:
    if not codes: return False
    return all(re.fullmatch(rf"\d{{{_BRINS_BASE_LEN_MIN},{_BRINS_BASE_LEN_MAX}}}", c) for c in codes)

def _pick_best_source_codes(clsdt_value, fac_value):
    clsdt_codes = clean_split_code(str(clsdt_value)) if not _is_blank_value(clsdt_value) else []
    if _is_valid_clean_result(clsdt_codes): return clsdt_codes, "CLSDT"
    fac_codes = clean_split_code(str(fac_value)) if not _is_blank_value(fac_value) else []
    if fac_codes: return fac_codes, "FAC"
    return clsdt_codes, "CLSDT"

def _resolve_polis_or_slip(clsdt_value, fac_value):
    fac_str = "" if _is_blank_value(fac_value) else str(fac_value).strip()
    if _PENYELESAIAN_RE.search(fac_str):
        clsdt_codes = clean_split_code(str(clsdt_value)) if not _is_blank_value(clsdt_value) else []
        if _is_valid_clean_result(clsdt_codes):
            keterangan = _TRAILING_CURRENCY_WORD_RE.sub("", fac_str).strip()
            keterangan = re.sub(r"\s{2,}", " ", keterangan)
            return [f"{code} {keterangan}" for code in clsdt_codes], False
        fac_codes = clean_split_code(fac_str)
        return fac_codes, _is_valid_clean_result(fac_codes)

    codes, _source_label = _pick_best_source_codes(clsdt_value, fac_value)
    return codes, _is_valid_clean_result(codes)

# ==============================================================================
# 4. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw: pd.DataFrame, keyword: str, cedant_column: str = "") -> pd.DataFrame:
    if cedant_column:
        if cedant_column not in df_raw.columns: raise ValueError(f"Kolom {cedant_column} tidak ada.")
        cedant_cols = [cedant_column]
    else:
        cedant_cols = [c for c in df_raw.columns if "CEDANT" in str(c).upper()]
        if not cedant_cols: raise ValueError("Kolom CEDANT tidak ditemukan.")

    mask = pd.Series(False, index=df_raw.index)
    for col in cedant_cols: mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)

    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant '{keyword}': {len(df_filtered)} dari total {len(df_raw)} baris.")
    return df_filtered

# ==============================================================================
# 5. BUILD OUTPUT
# ==============================================================================
def _detect_key_columns(df: pd.DataFrame, insured_col: str, polis_col: str, slip_col: str):
    def _res(col: str, kw_or_fn, lbl: str):
        if col:
            if col not in df.columns: raise ValueError(f"{lbl}='{col}' tidak ada.")
            return col
        return next((c for c in df.columns if (kw_or_fn(str(c).upper()) if callable(kw_or_fn) else kw_or_fn in str(c).upper())), None)

    col_ins, col_pol, col_slp = _res(insured_col, "INSURED", "INSURED_COLUMN"), _res(polis_col, lambda c: c.strip() == "POLIS", "POLIS_COLUMN"), _res(slip_col, "SLIP", "SLIP_COLUMN")
    if None in (col_ins, col_pol, col_slp): raise ValueError("Kolom kunci tidak lengkap.")
    return col_ins, col_pol, col_slp

_NON_ALNUM_RE = re.compile(r"[^A-Za-z0-9]")

def _compact_value(value: str) -> str:
    return _NON_ALNUM_RE.sub("", value) if isinstance(value, str) else value

def _compact_list(values: list) -> list:
    return [_compact_value(v) for v in values]

def _breakdown_all_rows(
    df: pd.DataFrame, col_ins: str, col_pol: str, col_slp: str,
    col_clsdt_pol: str = None, col_clsdt_slp: str = None, col_clsdt_sertf: str = None,
):
    ins_arr = df[col_ins].to_numpy()
    pol_arr = df[col_pol].to_numpy()
    slp_arr = df[col_slp].to_numpy()

    clsdt_pol_arr = df[col_clsdt_pol].to_numpy() if col_clsdt_pol and col_clsdt_pol in df.columns else [None] * len(df)
    clsdt_slp_arr = df[col_clsdt_slp].to_numpy() if col_clsdt_slp and col_clsdt_slp in df.columns else [None] * len(df)
    clsdt_sertf_arr = df[col_clsdt_sertf].to_numpy() if col_clsdt_sertf and col_clsdt_sertf in df.columns else [None] * len(df)

    ins_all, pol_all, slp_all, cert_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val, clsdt_pol_val, clsdt_slp_val, clsdt_sertf_val in zip(
        ins_arr, pol_arr, slp_arr, clsdt_pol_arr, clsdt_slp_arr, clsdt_sertf_arr
    ):
        ins = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []
        raw_pol = pol_val if pd.notna(pol_val) else ""
        raw_slp = slp_val if pd.notna(slp_val) else ""

        manual_override = try_manual_cert_override(str(raw_pol)) if not _is_blank_value(raw_pol) else None
        if manual_override is not None:
            manual_polis, manual_cert = manual_override
            pol, pol_should_compact, cert = [manual_polis], False, [manual_cert]
        else:
            clsdt_pol_is_usable = col_clsdt_pol and pd.notna(clsdt_pol_val) and _is_valid_clean_result(
                clean_split_code(str(clsdt_pol_val)) if not _is_blank_value(clsdt_pol_val) else []
            )
            polis_cert_groups = None
            if not clsdt_pol_is_usable and not _is_blank_value(raw_pol):
                fac_pol_normalized = _CERT_SD_NORMALIZE_RE.sub("SD", _strip_brins_specific_noise(str(raw_pol).strip()))
                polis_cert_groups = detect_polis_cert_groups(fac_pol_normalized)

            if polis_cert_groups:
                pol = [p for p, _certs in polis_cert_groups]
                pol_should_compact = True
                cert = [_join_certificates(certs) for _p, certs in polis_cert_groups]
            else:
                if col_clsdt_pol and pd.notna(clsdt_pol_val):
                    pol, pol_should_compact = _resolve_polis_or_slip(clsdt_pol_val, raw_pol)
                else:
                    pol = clean_split_code(str(raw_pol).strip()) if not _is_blank_value(raw_pol) else []
                    pol_should_compact = _is_valid_clean_result(pol)
                cert = None

        if col_clsdt_slp and pd.notna(clsdt_slp_val):
            slp, slp_should_compact = _resolve_polis_or_slip(clsdt_slp_val, raw_slp)
        else:
            slp = clean_split_code(str(raw_slp).strip()) if not _is_blank_value(raw_slp) else []
            slp_should_compact = _is_valid_clean_result(slp)

        if pol_should_compact: pol = _compact_list(pol)
        if slp_should_compact: slp = _compact_list(slp)

        if cert is None:
            cert = build_certificates_for_row(clsdt_sertf_val, raw_pol, num_polis=len(pol) if pol else 1)
        
        if len(cert) < len(pol): cert = cert + [""] * (len(pol) - len(cert))
        elif len(cert) > len(pol): cert = cert[: len(pol)] if pol else cert

        ins_all.append(ins)
        pol_all.append(pol)
        slp_all.append(slp)
        cert_all.append(cert)
        
        max_ins, max_pol, max_slp = max(max_ins, len(ins)), max(max_pol, len(pol)), max(max_slp, len(slp))

    return ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_ins, col_pol, col_slp = _detect_key_columns(df_filtered, INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN)
    col_clsdt_pol = CLSDT_POLIS_COLUMN if CLSDT_POLIS_COLUMN in df_filtered.columns else None
    col_clsdt_slp = CLSDT_SLIP_COLUMN if CLSDT_SLIP_COLUMN in df_filtered.columns else None
    col_clsdt_sertf = CLSDT_SERTF_COLUMN if CLSDT_SERTF_COLUMN in df_filtered.columns else None

    ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp = _breakdown_all_rows(
        df_filtered, col_ins, col_pol, col_slp, col_clsdt_pol, col_clsdt_slp, col_clsdt_sertf
    )

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)
    output_cols_data = {}

    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == col_ins:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in ins_all]
        if col == col_pol:
            for i in range(1, max_pol + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in pol_all]
                output_cols_data[f"{CERTIFICATE_OUTPUT_COL}_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_all]
        if col == col_slp:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slp_all]

    df_master = pd.DataFrame(output_cols_data)
    drop_cols = [c for c in df_master.columns if "_CLN_" in c and df_master[c].astype(str).str.strip().eq("").all()]
    return df_master.drop(columns=drop_cols)

# ==============================================================================
# 6. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook, text_col_indices = load_workbook(output_path), [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]
    worksheet = workbook.active

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None: cell.value = str(cell.value)
    workbook.save(output_path)

# ==============================================================================
# 7. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70 + "\n CLEANSING DATA 2 - OSBAL | CEDANT: BRINS GENERAL INSURANCE\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
        if not candidates: raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")
        input_path, _ = candidates[0], print(f"[AUTO-DETECT] Memakai file terdeteksi: '{candidates[0]}'")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    df_brins = filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)

    if not df_brins.empty:
        df_hasil = build_output(df_brins)
        save_with_text_format(df_hasil, OUTPUT_FILE)
        print("=" * 70 + f"\n[SUCCESS] Selesai. Output: '{OUTPUT_FILE}' ({len(df_hasil)} baris)\n" + "=" * 70)
    else:
        print("[WARNING] Tidak ada baris yang cocok dengan filter.")

if __name__ == "__main__":
    main()