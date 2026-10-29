# """
# ================================================================================
#  SCRIPT CLEANSING DATA 2 - OSBAL (KHUSUS CEDANT: JASA TANIA)
# ================================================================================
# """

# import os
# import re
# import pandas as pd
# from openpyxl import load_workbook

# # ==============================================================================
# # 0. KONFIGURASI
# # ==============================================================================
# INPUT_FILE = "2b. transaksi Osbal.xlsx"
# SHEET_NAME = "Query result"
# HEADER_ROW = 0
# OUTPUT_FILE = "jasatania_output_osbal_V2.xlsx"
# CEDANT_FILTER = "JASA TANIA"

# CEDANT_COLUMN = "CCOS_COMP_NAME"
# INSURED_COLUMN = "FAC_INSURED"
# POLIS_COLUMN = "FAC_POLICY_NO"
# SLIP_COLUMN = "FAC_SLIP"
# CLSDT_POLIS_COLUMN = "CLSDT_POLICY_NO"
# CLSDT_SLIP_COLUMN = "CLSDT_SLIP_NO"
# CLSDT_SERTF_COLUMN = "CLSDT_SERTF_NO"

# MAX_BREAKDOWN_CODES = 5
# CERTIFICATE_DIGIT_WIDTH = 6
# CERTIFICATE_RANGE_SD_THRESHOLD = 3
# TEXT_FORMAT_COLUMN_KEYWORDS = (
#     "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
#     "FAC_POLICY_NO", "FAC_SLIP", "CLSDT", "CERTIFICATE",
# )

# # ==============================================================================
# # 1. CLEANSING INSURED
# # ==============================================================================
# _BYPASS_ALL_RE = re.compile(
#     r"\bPENYELESAIAN\s+(?:HUTANG|UTANG)\s+PIUTANG\b|\bSUSPENSE\b", flags=re.IGNORECASE
# )
# _INSURED_DESCRIPTIVE_RE = re.compile(r"\bVARIOUS\s+INSUREDS?\b|\bACCEPTED\s+BY\b", flags=re.IGNORECASE)

# _TITLE_RE = re.compile(
#     r"\b(?:S\.?H\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|"
#     r"Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b",
#     flags=re.IGNORECASE,
# )
# _LEGAL_ENTITY_RE = re.compile(
#     r"\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO|Perseroan\s+Terbatas)\b",
#     flags=re.IGNORECASE,
# )
# _LEGAL_ENTITY_WITH_LEADING_PUNCT_RE = re.compile(
#     r"[,.]?\s*\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO|Perseroan\s+Terbatas)\b",
#     flags=re.IGNORECASE,
# )

# _BANK_NAME_RE = re.compile(
#     r"^(?:BANK\s+)?(?:BRI|BNI|BCA|MANDIRI|BTN|CIMB\s*NIAGA|PANIN|DANAMON|PERMATA|"
#     r"OCBC\s*NISP|MAYBANK|UOB|HSBC|CITIBANK|STANDARD\s*CHARTERED|BUKOPIN|MEGA|"
#     r"RAKYAT\s+INDONESIA|NEGARA\s+INDONESIA|CENTRAL\s+ASIA)\b.*$",
#     flags=re.IGNORECASE,
# )

# def _strip_legal_entity_with_punct(text: str) -> str:
#     def _strip_outside_parens(segment_text):
#         return _LEGAL_ENTITY_WITH_LEADING_PUNCT_RE.sub("", segment_text)

#     def _handle_paren_segment(paren_text):
#         inner = paren_text[1:-1].strip()
#         if _LEGAL_ENTITY_RE.fullmatch(inner):
#             return ""
#         return paren_text

#     segments = re.split(r"(\([^)]*\))", text)
#     result_segments = []
#     for seg in segments:
#         if seg.startswith("(") and seg.endswith(")"):
#             result_segments.append(_handle_paren_segment(seg))
#         else:
#             result_segments.append(_strip_outside_parens(seg))
#     t = "".join(result_segments)
#     t = re.sub(r",\s*(?=[(\s]|$)", " ", t)
#     return t

# def _strip_titles(text: str) -> str:
#     return _TITLE_RE.sub(" ", text)

# def _final_polish_insured(text: str) -> str:
#     segments = re.split(r"(\([^)]*\))", text)
#     cleaned_segments = []
#     for seg in segments:
#         if seg.startswith("(") and seg.endswith(")"):
#             cleaned_segments.append(seg)
#         else:
#             cleaned_segments.append(re.sub(r"[().;:\"']", " ", seg))
#     t = "".join(cleaned_segments)
#     t = re.sub(r"\s+", " ", t).strip()
#     t = re.sub(r"^[\s,.\-–]+|[\s,.\-–]+$", "", t)
#     t = re.sub(r"\(\s+", "(", t)
#     t = re.sub(r"\s+\)", ")", t)
#     return t.upper()

# _DIVISION_WORDS_RE = re.compile(
#     r"\b(?:UNIT|DIVISI|CABANG|BAGIAN|KANTOR)\b", flags=re.IGNORECASE
# )

# def _protect_division_ampersand(text: str) -> str:
#     def _replacer(match):
#         before = text[: match.start()]
#         recent_context = before[-60:]
#         if _DIVISION_WORDS_RE.search(recent_context):
#             return " \x00AMP\x00 "
#         return match.group(0)

#     return re.sub(r"\s&\s", _replacer, text)

# _ENTITY_SPLIT_RE = re.compile(
#     r"\bQQ\b|\band\s*/\s*or\s*|\s&\s|\s[–]\s|\+", flags=re.IGNORECASE
# )
# _ENTITY_SLASH_SPLIT_RE = re.compile(r"\s*/\s*")

# def _split_insured_entities(text: str):
#     protected = _protect_division_ampersand(text)
#     t = _ENTITY_SPLIT_RE.sub("/", protected)
#     parts = _ENTITY_SLASH_SPLIT_RE.split(t)
#     parts = [p.replace("\x00AMP\x00", "&").strip() for p in parts]
#     return [p for p in parts if p]

# def clean_insured_name(text: str) -> list:
#     if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
#         return []
#     original = text.strip()

#     if _BYPASS_ALL_RE.search(original) or _INSURED_DESCRIPTIVE_RE.search(original):
#         return [original]

#     raw_parts = _split_insured_entities(original)
#     if not raw_parts:
#         return [original]

#     if len(raw_parts) > MAX_BREAKDOWN_CODES:
#         cleaned_parts = []
#         for part in raw_parts:
#             cleaned = _strip_legal_entity_with_punct(part).strip()
#             cleaned_parts.append(cleaned if cleaned else part.strip())
#         return ["/".join(cleaned_parts)]

#     results = []
#     for part in raw_parts:
#         if _BANK_NAME_RE.match(_strip_legal_entity_with_punct(part).strip()):
#             continue
#         cleaned = _final_polish_insured(_strip_titles(_strip_legal_entity_with_punct(part)))
#         if not cleaned or len(cleaned) < 2 or not re.search(r"[A-Z]", cleaned):
#             continue
#         if cleaned not in results:
#             results.append(cleaned)

#     if not results:
#         return [original]

#     return [original] if len(results) > MAX_BREAKDOWN_CODES else results

# # ==============================================================================
# # 2. CLEANSING POLIS & SLIP NO
# # ==============================================================================
# _COMMA_DOT_RE = re.compile(r"[,.]")
# _REF_SUFFIX_RE = re.compile(r"^P\d*$|^VAR(?:IOUS)?$|^TBA\.?$", flags=re.IGNORECASE)
# _TBA_ALONE_RE = re.compile(r"^\s*TBA\.?\s*$", flags=re.IGNORECASE)
# _VAR_ALONE_RE = re.compile(r"^\s*VAR(?:IOUS)?\s*$", flags=re.IGNORECASE)
# _P_ALONE_RE = re.compile(r"^\s*P\d*\s*$", flags=re.IGNORECASE)
# _SEE_ATTACHMENT_ALONE_RE = re.compile(r"^\s*SEE\s+ATTACHMENT\s*$", flags=re.IGNORECASE)
# _CURRENCY_ALONE_RE = re.compile(r"^\s*(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\s*$", flags=re.IGNORECASE)

# _BYPASS_KEEP_ASIS_RE = re.compile(
#     r"\bPENYELESAIAN\s+(?:HUTANG|UTANG)\s+PIUTANG\b|\bPENYELESAIAN\s+SUSPEND\b|\bSUSPENSE\b",
#     flags=re.IGNORECASE,
# )
# _DATE_ONLY_RE = re.compile(r"^\s*.*\bTAHUN\b.*\b(?:19|20)\d{2}\b.*$", flags=re.IGNORECASE)
# _POLICY_SLIP_SPLIT_RE = re.compile(r"\s*\+\s*")

# _MAIN_CODE_RE = re.compile(r"^[A-Za-z]{1,3}\d{9,13}$")
# _MAIN_DIGITS_ONLY_RE = re.compile(r"^\d{9,14}$")
# _MAIN_ALNUM_MIXED_RE = re.compile(r"^[A-Za-z0-9]{9,16}$")

# def _is_valid_main_code(s: str) -> bool:
#     s = s.strip()
#     if _MAIN_CODE_RE.match(s) or _MAIN_DIGITS_ONLY_RE.match(s):
#         return True
#     if _MAIN_ALNUM_MIXED_RE.match(s) and re.search(r"\d", s):
#         return True
#     return False

# def _is_kept_as_is(value: str) -> bool:
#     v = value.strip()
#     if (_TBA_ALONE_RE.match(v) or _VAR_ALONE_RE.match(v) or _P_ALONE_RE.match(v)
#             or _SEE_ATTACHMENT_ALONE_RE.match(v) or _CURRENCY_ALONE_RE.match(v)):
#         return True
#     if _BYPASS_KEEP_ASIS_RE.search(v) or _DATE_ONLY_RE.match(v):
#         return True
#     return False

# _TRAILING_TBA_VAR_SPACE_RE = re.compile(r"^(.+?)\s+(?:TBA\.?|VAR(?:IOUS)?)\s*$", flags=re.IGNORECASE)

# def _try_strip_trailing_tba_var_space(text: str):
#     match = _TRAILING_TBA_VAR_SPACE_RE.match(text.strip())
#     if not match:
#         return None
#     remainder = match.group(1).strip()
#     cleaned = _clean_single_code(remainder)
#     if _is_valid_main_code(cleaned):
#         return [cleaned]
#     return None

# def _clean_single_code(value: str) -> str:
#     s = value.strip()
#     s = re.sub(r",\s*P\d*\s*$", "", s, flags=re.IGNORECASE)
#     s = _COMMA_DOT_RE.sub("", s)
#     s = re.sub(r"\s{2,}", " ", s).strip()
#     return s

# def _strip_trailing_ref_codes(parts: list) -> list:
#     return [p for p in parts if not _REF_SUFFIX_RE.match(p.strip())]

# def _rebuild_number_with_base(base: str, tail: str) -> str:
#     base_digit_match = re.search(r"\d+\s*$", base)
#     if not base_digit_match:
#         return base + tail
#     base_digits = base_digit_match.group(0).strip()
#     base_prefix = base[: base_digit_match.start()]
#     if len(tail) >= len(base_digits):
#         return base_prefix + tail
#     return base_prefix + base_digits[: -len(tail)] + tail

# _CNR_SEP = r"[/.]"
# _CNR_MULTI_RE = re.compile(
#     rf"^\s*\d+\s*{_CNR_SEP}\s*CNR\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{2,4}}"
#     rf"(?:\s*\+\s*\d+\s*{_CNR_SEP}\s*CNR\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{2,4}})+\s*$",
#     flags=re.IGNORECASE,
# )

# def _try_split_multi_cnr_keep_asis(text: str):
#     stripped = text.strip()
#     if not _CNR_MULTI_RE.match(stripped):
#         return None
#     parts = [p.strip() for p in stripped.split("+")]
#     parts = [p for p in parts if p]
#     return parts if parts else None

# _LEADING_P_PREFIX_RE = re.compile(r"^\s*P\d*\s*/\s*", flags=re.IGNORECASE)

# def _strip_leading_p_prefix(text: str) -> str:
#     return _LEADING_P_PREFIX_RE.sub("", text)

# _SD_RANGE_RE = re.compile(r"^(.+?)\s+SD\s+(\d{1,6})\s*$", flags=re.IGNORECASE)

# def _try_expand_sd_range(text: str):
#     match = _SD_RANGE_RE.match(text.strip())
#     if not match:
#         return None
#     base_full, end_suffix = match.groups()
#     base_full = base_full.strip()
#     base_digit_match = re.search(r"(\d+)\s*$", base_full)
#     if not base_digit_match:
#         return None
#     base_digits = base_digit_match.group(1)
#     if len(end_suffix) > len(base_digits):
#         return None
#     start_num_str = base_digits[-len(end_suffix):]
#     try:
#         start, end = int(start_num_str), int(end_suffix)
#     except ValueError:
#         return None
#     if start > end or (end - start + 1) > MAX_BREAKDOWN_CODES:
#         return None
#     width = len(end_suffix)
#     prefix = base_full[: base_digit_match.start(1)] + base_digits[: -width]
#     codes = [_clean_single_code(f"{prefix}{n:0{width}d}") for n in range(start, end + 1)]
#     deduped = []
#     for c in codes:
#         if c not in deduped:
#             deduped.append(c)
#     return deduped

# def _try_strip_narrative_after_slash(text: str):
#     if "/" not in text:
#         return None
#     parts = [p.strip() for p in text.split("/")]
#     if len(parts) < 2:
#         return None
#     first = parts[0]
#     if not _is_valid_main_code(_clean_single_code(first)):
#         return None
#     rest_parts = parts[1:]
#     if any(_is_valid_main_code(_clean_single_code(p)) for p in rest_parts):
#         return None
#     return [_clean_single_code(first)]

# _DASH_BASE_RE = re.compile(r"^([A-Za-z]{0,3}\d{6,14})-(\d{1,6}(?:\s*\+\s*\d{1,6})*)\s*(.*)$")

# def _try_expand_dash_suffix(text: str):
#     stripped = text.strip()
#     match = _DASH_BASE_RE.match(stripped)
#     if not match:
#         return None
#     base, suffix_blob, rest = match.groups()

#     base_clean = _clean_single_code(base)
#     has_alpha_prefix = bool(re.match(r"^[A-Za-z]", base_clean))

#     suffixes = [s.strip() for s in suffix_blob.split("+") if s.strip()]
#     codes = [base_clean] if not has_alpha_prefix else []
#     for suf in suffixes:
#         if has_alpha_prefix:
#             codes.append(base_clean + suf)
#         else:
#             codes.append(_clean_single_code(_rebuild_number_with_base(base_clean, suf)))

#     rest = rest.strip()
#     if rest.startswith("+"):
#         rest = rest[1:].strip()
#     if rest:
#         extra_parts = [p.strip() for p in rest.split("+") if p.strip()]
#         for p in extra_parts:
#             codes.append(_clean_single_code(p))

#     deduped = []
#     for c in codes:
#         if c not in deduped:
#             deduped.append(c)
#     return deduped if len(deduped) <= MAX_BREAKDOWN_CODES else None

# def _split_and_rebuild_plus(cleaned_noise: str):
#     parts = _POLICY_SLIP_SPLIT_RE.split(cleaned_noise)
#     parts = [p.strip() for p in parts if p.strip()]

#     if len(parts) <= 1:
#         single = _clean_single_code(cleaned_noise)
#         return [single], not _is_valid_main_code(single)

#     parts = _strip_trailing_ref_codes(parts)
#     if len(parts) == 0:
#         single = _clean_single_code(cleaned_noise)
#         return [single], not _is_valid_main_code(single)
#     if len(parts) == 1:
#         single = _clean_single_code(parts[0])
#         return [single], not _is_valid_main_code(single)

#     total_pieces = 1
#     for tail in parts[1:]:
#         sub_suffixes = [s for s in tail.split("_") if s]
#         total_pieces += max(len(sub_suffixes), 1)

#     if total_pieces > MAX_BREAKDOWN_CODES:
#         return [cleaned_noise.strip()], True

#     base = parts[0]
#     codes = [_clean_single_code(base)]
#     for tail in parts[1:]:
#         sub_suffixes = [s for s in tail.split("_") if s] or [tail]
#         for suf in sub_suffixes:
#             if re.fullmatch(r"\d+", suf) and len(suf) <= 6 and len(suf) < len(base):
#                 rebuilt = _rebuild_number_with_base(base, suf)
#                 codes.append(_clean_single_code(rebuilt))
#             else:
#                 codes.append(_clean_single_code(suf))

#     deduped = []
#     for c in codes:
#         if c not in deduped:
#             deduped.append(c)
#     return deduped, False

# _TRAILING_CERT_SUFFIX_RE = re.compile(
#     r"^([A-Za-z]{1,3}\d{9,13}|\d{9,14})\s*[-\s]\s*(\d{5,6})\s*$"
# )

# def _try_strip_trailing_certificate_suffix(text: str):
#     match = _TRAILING_CERT_SUFFIX_RE.match(text.strip())
#     if not match:
#         return None
#     base, _suffix = match.groups()
#     base_clean = _clean_single_code(base)
#     if _is_valid_main_code(base_clean):
#         return [base_clean]
#     return None

# _DASH_VAR_SUFFIX_RE = re.compile(r"^(.+?)\s*-\s*VAR(?:IOUS)?\s*$", flags=re.IGNORECASE)

# def _try_strip_trailing_dash_var(text: str):
#     match = _DASH_VAR_SUFFIX_RE.match(text.strip())
#     if not match:
#         return None
#     remainder = match.group(1).strip()
#     cleaned = _clean_single_code(remainder)
#     if _is_valid_main_code(cleaned):
#         return [cleaned]
#     return None

# _COMMA_BREAKDOWN_RE = re.compile(r"^([A-Za-z]{1,3}\d{6,14})\s*,\s*(\d{1,6}(?:\s*,\s*\d{1,6})*)\s*$")

# def _try_expand_comma_suffix(text: str):
#     match = _COMMA_BREAKDOWN_RE.match(text.strip())
#     if not match:
#         return None
#     base, suffix_blob = match.groups()
#     suffixes = [s.strip() for s in suffix_blob.split(",") if s.strip()]
#     if not suffixes:
#         return None
#     total_pieces = 1 + len(suffixes)
#     if total_pieces > MAX_BREAKDOWN_CODES:
#         return None
#     codes = [_clean_single_code(base)]
#     for suf in suffixes:
#         rebuilt = _rebuild_number_with_base(base, suf)
#         codes.append(_clean_single_code(rebuilt))
#     deduped = []
#     for c in codes:
#         if c not in deduped:
#             deduped.append(c)
#     return deduped

# _CNR_SINGLE_TRAILING_SUFFIX_PLUS_RE = re.compile(
#     rf"^\s*(\d{{4,14}})((?:\s*\+\s*\d{{1,6}})+)\s*({_CNR_SEP}\s*CNR\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{2,4}})\s*$",
#     flags=re.IGNORECASE,
# )

# def _try_expand_plus_with_trailing_cnr(text: str):
#     match = _CNR_SINGLE_TRAILING_SUFFIX_PLUS_RE.match(text.strip())
#     if not match:
#         return None
#     base, suffix_blob, cnr_suffix = match.groups()
#     suffixes = [s.strip() for s in suffix_blob.split("+") if s.strip()]
#     if not suffixes:
#         return None
#     total_pieces = 1 + len(suffixes)
#     if total_pieces > MAX_BREAKDOWN_CODES:
#         return None
#     numbers = [base] + [_rebuild_number_with_base(base, suf) for suf in suffixes]
#     deduped_numbers = []
#     for n in numbers:
#         if n not in deduped_numbers:
#             deduped_numbers.append(n)
#     return [f"{n}{cnr_suffix.strip()}" for n in deduped_numbers]

# def clean_split_code_impl(text: str):
#     if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
#         return [], True
#     original = text.strip()

#     if _is_kept_as_is(original):
#         return [original], True

#     plus_cnr_reattach_result = _try_expand_plus_with_trailing_cnr(original)
#     if plus_cnr_reattach_result is not None:
#         return plus_cnr_reattach_result, True

#     multi_cnr_result = _try_split_multi_cnr_keep_asis(original)
#     if multi_cnr_result is not None:
#         return multi_cnr_result, True

#     trailing_tba_result = _try_strip_trailing_tba_var_space(original)
#     if trailing_tba_result is not None:
#         return trailing_tba_result, False

#     working = _strip_leading_p_prefix(original)
#     if working != original and not working.strip():
#         return [original], True

#     trailing_cert_result = _try_strip_trailing_certificate_suffix(working)
#     if trailing_cert_result is not None:
#         return trailing_cert_result, False

#     multi_pair_result = _try_split_polis_certificate_pairs(working)
#     if multi_pair_result is not None:
#         return [polis for polis, _cert in multi_pair_result], False

#     dash_var_result = _try_strip_trailing_dash_var(working)
#     if dash_var_result is not None:
#         return dash_var_result, False

#     if "CNR" in working.upper():
#         return [original], True

#     sd_result = _try_expand_sd_range(working)
#     if sd_result is not None:
#         return sd_result, False

#     comma_result = _try_expand_comma_suffix(working)
#     if comma_result is not None:
#         return comma_result, False

#     narrative_result = _try_strip_narrative_after_slash(working)
#     if narrative_result is not None:
#         return narrative_result, False

#     if "-" in working:
#         dash_result = _try_expand_dash_suffix(working)
#         if dash_result is not None:
#             return dash_result, False

#     if "/" in working and "+" not in working and "-" not in working:
#         slash_parts = [p.strip() for p in working.split("/") if p.strip()]
#         cleaned_slash_parts = [_clean_single_code(p) for p in slash_parts]
#         if len(slash_parts) >= 2 and all(_is_valid_main_code(p) for p in cleaned_slash_parts):
#             if len(cleaned_slash_parts) <= MAX_BREAKDOWN_CODES:
#                 deduped = []
#                 for c in cleaned_slash_parts:
#                     if c not in deduped:
#                         deduped.append(c)
#                 return deduped, False

#     return _split_and_rebuild_plus(working)

# def clean_split_code(text: str) -> list:
#     results, _ = clean_split_code_impl(text)
#     return results

# # ==============================================================================
# # 3. PEMILIHAN SUMBER TERBAIK: CLSDT vs FAC (POLIS & SLIP)
# # ==============================================================================
# def _is_blank_value(value) -> bool:
#     if value is None:
#         return True
#     if isinstance(value, float) and pd.isna(value):
#         return True
#     s = str(value).strip()
#     return s == "" or s.lower() in ("nan", "none")

# def _is_valid_clean_result(codes: list) -> bool:
#     if not codes:
#         return False
#     return all(_is_valid_main_code(c) for c in codes)

# def _fac_has_bypass_keyword(fac_value) -> bool:
#     if _is_blank_value(fac_value):
#         return False
#     return bool(_BYPASS_KEEP_ASIS_RE.search(str(fac_value).strip()))

# _TRAILING_CURRENCY_CODE_RE = re.compile(
#     r"\s+(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\s*$", flags=re.IGNORECASE
# )

# def _strip_trailing_currency(text: str) -> str:
#     return _TRAILING_CURRENCY_CODE_RE.sub("", text).strip()

# _GENERIC_CLSDT_VALUE_RE = re.compile(
#     r"^\s*(?:TBA\.?|VAR(?:IOUS)?|P\d*|-|SEE\s+ATTACHMENT)\s*$", flags=re.IGNORECASE
# )

# def _is_generic_clsdt_value(value) -> bool:
#     if _is_blank_value(value):
#         return True
#     return bool(_GENERIC_CLSDT_VALUE_RE.match(str(value).strip()))

# def _resolve_polis_or_slip(clsdt_value, fac_value):
#     clsdt_codes = clean_split_code(str(clsdt_value)) if not _is_blank_value(clsdt_value) else []

#     if _is_valid_clean_result(clsdt_codes) and _fac_has_bypass_keyword(fac_value):
#         clsdt_single = clsdt_codes[0] if len(clsdt_codes) == 1 else " ".join(clsdt_codes)
#         fac_text = _strip_trailing_currency(str(fac_value).strip())
#         combined = f"{clsdt_single} {fac_text}"
#         return [combined], True

#     if _is_valid_clean_result(clsdt_codes):
#         return clsdt_codes, False

#     if clsdt_codes and not _is_generic_clsdt_value(clsdt_value):
#         return clsdt_codes, True

#     fac_codes = clean_split_code(str(fac_value)) if not _is_blank_value(fac_value) else []
#     if fac_codes:
#         return fac_codes, not _is_valid_clean_result(fac_codes)

#     if clsdt_codes:
#         return clsdt_codes, True

#     return [], True

# # ==============================================================================
# # 3b. EKSTRAKSI CERTIFICATE
# # ==============================================================================
# _CERT_VALID_PREFIX_CODE_RE = re.compile(r"^[A-Za-z]{1,3}\d{9,13}$|^\d{9,14}$")
# _SD_NOTATION_RE = re.compile(r"\bS\s*/\s*D\b", flags=re.IGNORECASE)

# def _normalize_sd_notation(text: str) -> str:
#     return _SD_NOTATION_RE.sub("SD", text)

# def _pad_certificate_digits(digit_str: str):
#     d = digit_str.strip()
#     if not d.isdigit():
#         return None
#     if len(d) > CERTIFICATE_DIGIT_WIDTH:
#         return None
#     return d.zfill(CERTIFICATE_DIGIT_WIDTH)

# def _is_certificate_like_clsdt(value: str) -> bool:
#     v = value.strip()
#     if not v:
#         return False
#     v_norm = _normalize_sd_notation(v)
#     if re.search(r"\bSD\b", v_norm, flags=re.IGNORECASE) or "," in v_norm:
#         return True
#     cleaned = re.sub(r"[.\-\s]", "", v_norm)
#     return cleaned.isdigit() and len(cleaned) <= CERTIFICATE_DIGIT_WIDTH + 2

# def _format_certificate_range(codes: list) -> str:
#     deduped = []
#     for c in codes:
#         if c not in deduped:
#             deduped.append(c)
#     deduped.sort(key=lambda x: int(x) if x.isdigit() else x)
#     if len(deduped) == 1:
#         return deduped[0]
#     if len(deduped) > CERTIFICATE_RANGE_SD_THRESHOLD:
#         return f"{deduped[0]} SD {deduped[-1]}"
#     return ", ".join(deduped)

# def _extract_certificate_from_clsdt(clsdt_sertf_value) -> str:
#     if _is_blank_value(clsdt_sertf_value):
#         return None

#     raw = str(clsdt_sertf_value).strip()
#     normalized = _normalize_sd_notation(raw)

#     sd_match = re.match(r"^(\d{1,6})\s*SD\s*(\d{1,6})\s*$", normalized, flags=re.IGNORECASE)
#     if sd_match:
#         start_s, end_s = sd_match.groups()
#         start_p, end_p = _pad_certificate_digits(start_s), _pad_certificate_digits(end_s)
#         if start_p is None or end_p is None:
#             return ""
#         try:
#             start_n, end_n = int(start_p), int(end_p)
#         except ValueError:
#             return ""
#         if start_n > end_n:
#             return ""
#         codes = [str(n).zfill(CERTIFICATE_DIGIT_WIDTH) for n in range(start_n, end_n + 1)]
#         return _format_certificate_range(codes)

#     if "," in normalized:
#         parts = [p.strip() for p in normalized.split(",") if p.strip()]
#         padded = [_pad_certificate_digits(p) for p in parts]
#         if not padded or any(p is None for p in padded):
#             return ""
#         return _format_certificate_range(padded)

#     if not _is_certificate_like_clsdt(raw):
#         return ""
#     cleaned = re.sub(r"[.\-\s]", "", normalized)
#     padded = _pad_certificate_digits(cleaned)
#     if padded is None:
#         return ""
#     return padded

# _CERT_TRAILING_SD_RE = re.compile(
#     r"^([A-Za-z]{1,3}\d{9,13}|\d{9,14})\s*[-\s]\s*(\d{6})\s+SD\s+(\d{1,6})\s*$",
#     flags=re.IGNORECASE,
# )
# _CERT_TRAILING_COMMA_RE = re.compile(
#     r"^([A-Za-z]{1,3}\d{9,13}|\d{9,14})\s*[-\s]\s*(\d{5,6}(?:\s*,\s*\d{1,6})+)\s*$"
# )
# _CERT_TRAILING_SINGLE_RE = re.compile(
#     r"^([A-Za-z]{1,3}\d{9,13}|\d{9,14})\s*[-\s]\s*(\d{5,6})\s*$"
# )

# def _extract_certificate_from_fac(fac_value) -> str:
#     if _is_blank_value(fac_value):
#         return ""
#     raw = str(fac_value).strip()
#     normalized = _normalize_sd_notation(raw)
#     normalized = _strip_leading_p_prefix(normalized).strip()

#     if _BYPASS_KEEP_ASIS_RE.search(normalized) or _is_kept_as_is(normalized) or "+" in normalized:
#         return ""

#     sd_match = _CERT_TRAILING_SD_RE.match(normalized)
#     if sd_match:
#         _base, start_s, end_s = sd_match.groups()
#         start_p, end_p = _pad_certificate_digits(start_s), _pad_certificate_digits(end_s)
#         if start_p is None or end_p is None:
#             return ""
#         try:
#             start_n, end_n = int(start_p), int(end_p)
#         except ValueError:
#             return ""
#         if start_n > end_n or (end_n - start_n + 1) > MAX_BREAKDOWN_CODES:
#             return ""
#         codes = [str(n).zfill(CERTIFICATE_DIGIT_WIDTH) for n in range(start_n, end_n + 1)]
#         return _format_certificate_range(codes)

#     comma_match = _CERT_TRAILING_COMMA_RE.match(normalized)
#     if comma_match:
#         _base, suffix_blob = comma_match.groups()
#         parts = [p.strip() for p in suffix_blob.split(",") if p.strip()]
#         padded = [_pad_certificate_digits(p) for p in parts]
#         if not padded or any(p is None for p in padded) or len(padded) > MAX_BREAKDOWN_CODES:
#             return ""
#         return _format_certificate_range(padded)

#     single_match = _CERT_TRAILING_SINGLE_RE.match(normalized)
#     if single_match:
#         _base, cert_digits = single_match.groups()
#         padded = _pad_certificate_digits(cert_digits)
#         if padded is None:
#             return ""
#         return padded

#     return ""

# def extract_certificate(fac_policy_value, clsdt_sertf_value=None) -> str:
#     if clsdt_sertf_value is not None:
#         clsdt_result = _extract_certificate_from_clsdt(clsdt_sertf_value)
#         if clsdt_result is not None:
#             return clsdt_result

#     return _extract_certificate_from_fac(fac_policy_value)

# _CERT_PAIR_SLASH_SPLIT_RE = re.compile(r"\s*/\s*")
# _CERT_PAIR_SINGLE_RE = re.compile(
#     r"^([A-Za-z]{1,3}\d{9,13}|\d{9,14})\s*[-\s]\s*(\d{5,6})\s*$"
# )

# def _try_split_polis_certificate_pairs(text: str):
#     stripped = text.strip()
#     if "/" not in stripped:
#         return None
#     segments = [s.strip() for s in _CERT_PAIR_SLASH_SPLIT_RE.split(stripped) if s.strip()]
#     if len(segments) < 2:
#         return None

#     pairs = []
#     for seg in segments:
#         seg_norm = _normalize_sd_notation(seg)
#         match = _CERT_PAIR_SINGLE_RE.match(seg_norm)
#         if not match:
#             return None
#         polis_raw, cert_raw = match.groups()
#         polis_clean = _clean_single_code(polis_raw)
#         if not _is_valid_main_code(polis_clean):
#             return None
#         cert_padded = _pad_certificate_digits(cert_raw)
#         if cert_padded is None:
#             return None
#         pairs.append((polis_clean, cert_padded))

#     if len(pairs) > MAX_BREAKDOWN_CODES:
#         return None
#     return pairs

# def extract_certificate_pairs(fac_policy_value, clsdt_sertf_value=None, polis_cln_count=None):
#     def _pad_to_count(values: list) -> list:
#         if polis_cln_count is None:
#             return values
#         if len(values) >= polis_cln_count:
#             return values[:polis_cln_count] if polis_cln_count > 0 else values
#         return values + [""] * (polis_cln_count - len(values))

#     if clsdt_sertf_value is not None:
#         clsdt_result = _extract_certificate_from_clsdt(clsdt_sertf_value)
#         if clsdt_result is not None:
#             return _pad_to_count([clsdt_result])

#     if _is_blank_value(fac_policy_value):
#         return _pad_to_count([])

#     raw = str(fac_policy_value).strip()
#     working = _strip_leading_p_prefix(_normalize_sd_notation(raw)).strip()

#     multi_pairs = _try_split_polis_certificate_pairs(working)
#     if multi_pairs is not None:
#         return _pad_to_count([cert for _polis, cert in multi_pairs])

#     single_cert = _extract_certificate_from_fac(raw)
#     return _pad_to_count([single_cert] if single_cert else [""])

# # ==============================================================================
# # 4. LOAD & FILTER DATA
# # ==============================================================================
# def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
#     return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

# def filter_cedant(df_raw: pd.DataFrame, keyword: str, cedant_column: str = "") -> pd.DataFrame:
#     if cedant_column:
#         if cedant_column not in df_raw.columns:
#             raise ValueError(f"Kolom {cedant_column} tidak ada.")
#         cedant_cols = [cedant_column]
#     else:
#         if not (cedant_cols := [c for c in df_raw.columns if "CEDANT" in str(c).upper()]):
#             raise ValueError("Kolom CEDANT tidak ditemukan lewat auto-detect.")

#     mask = pd.Series(False, index=df_raw.index)
#     for col in cedant_cols:
#         mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)

#     df_filtered = df_raw[mask].copy()
#     print(f"[INFO] Filter cedant '{keyword}': {len(df_filtered)} dari total {len(df_raw)} baris.")
#     return df_filtered

# # ==============================================================================
# # 5. BUILD OUTPUT
# # ==============================================================================
# _NON_ALNUM_RE = re.compile(r"[^A-Za-z0-9]")

# def _compact_value(value: str) -> str:
#     if not isinstance(value, str):
#         return value
#     return _NON_ALNUM_RE.sub("", value)

# def _compact_list(values: list) -> list:
#     return [_compact_value(v) for v in values]

# def _compact_split_code_results(results: list, is_kept_as_is: bool) -> list:
#     if is_kept_as_is:
#         return results
#     return [_compact_value(v) for v in results]

# def _detect_key_columns(df: pd.DataFrame, insured_col: str, polis_col: str, slip_col: str):
#     def _res(col: str, kw_or_fn, lbl: str):
#         if col:
#             if col not in df.columns:
#                 raise ValueError(f"{lbl}='{col}' tidak ada.")
#             return col
#         return next(
#             (c for c in df.columns if (kw_or_fn(str(c).upper()) if callable(kw_or_fn) else kw_or_fn in str(c).upper())),
#             None,
#         )

#     col_ins = _res(insured_col, "INSURED", "INSURED_COLUMN")
#     col_pol = _res(polis_col, lambda c: c.strip() == "POLIS", "POLIS_COLUMN")
#     col_slp = _res(slip_col, "SLIP", "SLIP_COLUMN")

#     if None in (col_ins, col_pol, col_slp):
#         raise ValueError("Kolom kunci tidak lengkap.")

#     print(f"[INFO] Kolom terdeteksi -> INSURED='{col_ins}', POLIS='{col_pol}', SLIP='{col_slp}'")
#     return col_ins, col_pol, col_slp

# def _breakdown_all_rows(
#     df: pd.DataFrame,
#     col_ins: str,
#     col_pol: str,
#     col_slp: str,
#     col_clsdt_pol: str = None,
#     col_clsdt_slp: str = None,
#     col_clsdt_sertf: str = None,
# ):
#     ins_arr = df[col_ins].to_numpy()
#     pol_arr = df[col_pol].to_numpy()
#     slp_arr = df[col_slp].to_numpy()

#     clsdt_pol_arr = df[col_clsdt_pol].to_numpy() if col_clsdt_pol and col_clsdt_pol in df.columns else [None] * len(df)
#     clsdt_slp_arr = df[col_clsdt_slp].to_numpy() if col_clsdt_slp and col_clsdt_slp in df.columns else [None] * len(df)
#     clsdt_sertf_arr = df[col_clsdt_sertf].to_numpy() if col_clsdt_sertf and col_clsdt_sertf in df.columns else [None] * len(df)

#     ins_all, pol_all, slp_all, cert_all = [], [], [], []
#     max_ins = max_pol = max_slp = 0

#     for ins_val, pol_val, slp_val, clsdt_pol_val, clsdt_slp_val, clsdt_sertf_val in zip(
#         ins_arr, pol_arr, slp_arr, clsdt_pol_arr, clsdt_slp_arr, clsdt_sertf_arr
#     ):
#         ins = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []

#         raw_pol = pol_val if pd.notna(pol_val) else ""
#         raw_slp = slp_val if pd.notna(slp_val) else ""

#         if col_clsdt_pol and pd.notna(clsdt_pol_val):
#             pol_results, pol_kept_as_is = _resolve_polis_or_slip(clsdt_pol_val, raw_pol)
#         else:
#             pol_results, pol_kept_as_is = clean_split_code_impl(str(raw_pol).strip()) if not _is_blank_value(raw_pol) else ([], True)
#         pol = _compact_split_code_results(pol_results, pol_kept_as_is)

#         if col_clsdt_slp and pd.notna(clsdt_slp_val):
#             slp_results, slp_kept_as_is = _resolve_polis_or_slip(clsdt_slp_val, raw_slp)
#         else:
#             slp_results, slp_kept_as_is = clean_split_code_impl(str(raw_slp).strip()) if not _is_blank_value(raw_slp) else ([], True)
#         slp = _compact_split_code_results(slp_results, slp_kept_as_is)

#         clsdt_sertf_for_row = clsdt_sertf_val if (col_clsdt_sertf and pd.notna(clsdt_sertf_val)) else None
#         cert = extract_certificate_pairs(raw_pol, clsdt_sertf_for_row, polis_cln_count=len(pol))

#         ins_all.append(ins)
#         pol_all.append(pol)
#         slp_all.append(slp)
#         cert_all.append(cert)

#         if len(ins) > max_ins: max_ins = len(ins)
#         if len(pol) > max_pol: max_pol = len(pol)
#         if len(slp) > max_slp: max_slp = len(slp)

#     return ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp

# def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
#     col_ins, col_pol, col_slp = _detect_key_columns(df_filtered, INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN)

#     col_clsdt_pol = CLSDT_POLIS_COLUMN if CLSDT_POLIS_COLUMN in df_filtered.columns else None
#     col_clsdt_slp = CLSDT_SLIP_COLUMN if CLSDT_SLIP_COLUMN in df_filtered.columns else None
#     col_clsdt_sertf = CLSDT_SERTF_COLUMN if CLSDT_SERTF_COLUMN in df_filtered.columns else None
    
#     if col_clsdt_pol or col_clsdt_slp:
#         print(f"[INFO] Kolom CLSDT terdeteksi -> POLIS='{col_clsdt_pol}', SLIP='{col_clsdt_slp}' (prioritas utama)")
#     else:
#         print("[INFO] Kolom CLSDT tidak ditemukan di file ini -> pakai FAC_POLICY_NO/FAC_SLIP saja.")
    
#     if col_clsdt_sertf:
#         print(f"[INFO] Kolom CLSDT_SERTF_NO terdeteksi -> '{col_clsdt_sertf}' (prioritas utama utk CERTIFICATE)")
#     else:
#         print("[INFO] Kolom CLSDT_SERTF_NO tidak ditemukan -> CERTIFICATE diekstrak dari FAC_POLICY_NO saja.")

#     ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp = _breakdown_all_rows(
#         df_filtered, col_ins, col_pol, col_slp, col_clsdt_pol, col_clsdt_slp, col_clsdt_sertf
#     )

#     original_cols = list(df_filtered.columns)
#     df_filtered = df_filtered.reset_index(drop=True)
#     output_cols_data = {}

#     for col in original_cols:
#         output_cols_data[col] = df_filtered[col].values
#         if col == col_ins:
#             for i in range(1, max_ins + 1):
#                 output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in ins_all]
#         if col == col_pol:
#             for i in range(1, max(max_pol, 1) + 1):
#                 output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in pol_all]
#                 output_cols_data[f"CERTIFICATE_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_all]
#         if col == col_slp:
#             for i in range(1, max_slp + 1):
#                 output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slp_all]

#     df_master = pd.DataFrame(output_cols_data)
#     drop_cols = [
#         c for c in df_master.columns
#         if ("_CLN_" in c or (c.startswith("CERTIFICATE_") and c != "CERTIFICATE_1"))
#         and df_master[c].astype(str).str.strip().eq("").all()
#     ]
#     return df_master.drop(columns=drop_cols)

# # ==============================================================================
# # 6. SIMPAN OUTPUT
# # ==============================================================================
# def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
#     df.to_excel(output_path, index=False)
#     workbook, text_col_indices = load_workbook(output_path), [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]
#     worksheet = workbook.active

#     for col_idx in text_col_indices:
#         for row in range(2, worksheet.max_row + 1):
#             cell = worksheet.cell(row=row, column=col_idx)
#             cell.number_format = "@"
#             if cell.value is not None:
#                 cell.value = str(cell.value)
#     workbook.save(output_path)

# # ==============================================================================
# # 7. EXECUTION RUNNER
# # ==============================================================================
# def main() -> None:
#     print("=" * 70 + "\n CLEANSING DATA 2 - OSBAL | CEDANT: JASA TANIA\n" + "=" * 70)
#     input_path = INPUT_FILE

#     if not os.path.exists(input_path):
#         candidates = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
#         if not candidates:
#             raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")
#         input_path = candidates[0]
#         print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")

#     df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
#     df_jt = filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)

#     if not df_jt.empty:
#         df_hasil = build_output(df_jt)
#         save_with_text_format(df_hasil, OUTPUT_FILE)
#         print("=" * 70 + f"\n[SUCCESS] Selesai. Output: '{OUTPUT_FILE}' ({len(df_hasil)} baris)\n" + "=" * 70)
#     else:
#         print("[WARNING] Tidak ada baris yang cocok dengan filter.")

# if __name__ == "__main__":
#     main()

"""
================================================================================
 SCRIPT CLEANSING DATA 2 - OSBAL (KHUSUS CEDANT: JASA TANIA)
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
INPUT_FILE = "2b. transaksi Osbal.xlsx"
SHEET_NAME = "Query result"
HEADER_ROW = 0
OUTPUT_FILE = "jasatania_output_osbal_v2.xlsx"
CEDANT_FILTER = "JASA TANIA"

CEDANT_COLUMN = "CCOS_COMP_NAME"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"
CLSDT_POLIS_COLUMN = "CLSDT_POLICY_NO"
CLSDT_SLIP_COLUMN = "CLSDT_SLIP_NO"
CLSDT_SERTF_COLUMN = "CLSDT_SERTF_NO"

MAX_BREAKDOWN_CODES = 5
CERTIFICATE_DIGIT_WIDTH = 6
CERTIFICATE_RANGE_SD_THRESHOLD = 3
TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CLSDT", "CERTIFICATE",
)

# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
_BYPASS_ALL_RE = re.compile(
    r"\bPENYELESAIAN\s+(?:HUTANG|UTANG)\s+PIUTANG\b|\bSUSPENSE\b", flags=re.IGNORECASE
)
_INSURED_DESCRIPTIVE_RE = re.compile(r"\bVARIOUS\s+INSUREDS?\b|\bACCEPTED\s+BY\b", flags=re.IGNORECASE)

_TITLE_RE = re.compile(
    r"\b(?:S\.?H\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|"
    r"Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b",
    flags=re.IGNORECASE,
)
_LEGAL_ENTITY_RE = re.compile(
    r"\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO|Perseroan\s+Terbatas)\b",
    flags=re.IGNORECASE,
)
_LEGAL_ENTITY_WITH_LEADING_PUNCT_RE = re.compile(
    r"[,.]?\s*\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO|Perseroan\s+Terbatas)\b",
    flags=re.IGNORECASE,
)

_BANK_NAME_RE = re.compile(
    r"^(?:BANK\s+)?(?:BRI|BNI|BCA|MANDIRI|BTN|CIMB\s*NIAGA|PANIN|DANAMON|PERMATA|"
    r"OCBC\s*NISP|MAYBANK|UOB|HSBC|CITIBANK|STANDARD\s*CHARTERED|BUKOPIN|MEGA|"
    r"RAKYAT\s+INDONESIA|NEGARA\s+INDONESIA|CENTRAL\s+ASIA)\b.*$",
    flags=re.IGNORECASE,
)

def _strip_legal_entity_with_punct(text: str) -> str:
    def _strip_outside_parens(segment_text):
        return _LEGAL_ENTITY_WITH_LEADING_PUNCT_RE.sub("", segment_text)

    def _handle_paren_segment(paren_text):
        inner = paren_text[1:-1].strip()
        if _LEGAL_ENTITY_RE.fullmatch(inner):
            return ""
        return paren_text

    segments = re.split(r"(\([^)]*\))", text)
    result_segments = []
    for seg in segments:
        if seg.startswith("(") and seg.endswith(")"):
            result_segments.append(_handle_paren_segment(seg))
        else:
            result_segments.append(_strip_outside_parens(seg))
    t = "".join(result_segments)
    t = re.sub(r",\s*(?=[(\s]|$)", " ", t)
    return t

def _strip_titles(text: str) -> str:
    return _TITLE_RE.sub(" ", text)

def _final_polish_insured(text: str) -> str:
    segments = re.split(r"(\([^)]*\))", text)
    cleaned_segments = []
    for seg in segments:
        if seg.startswith("(") and seg.endswith(")"):
            cleaned_segments.append(seg)
        else:
            cleaned_segments.append(re.sub(r"[().;:\"']", " ", seg))
    t = "".join(cleaned_segments)
    t = re.sub(r"\s+", " ", t).strip()
    t = re.sub(r"^[\s,.\-–]+|[\s,.\-–]+$", "", t)
    t = re.sub(r"\(\s+", "(", t)
    t = re.sub(r"\s+\)", ")", t)
    return t.upper()

_DIVISION_WORDS_RE = re.compile(
    r"\b(?:UNIT|DIVISI|CABANG|BAGIAN|KANTOR)\b", flags=re.IGNORECASE
)

def _protect_division_ampersand(text: str) -> str:
    def _replacer(match):
        before = text[: match.start()]
        recent_context = before[-60:]
        if _DIVISION_WORDS_RE.search(recent_context):
            return " \x00AMP\x00 "
        return match.group(0)

    return re.sub(r"\s&\s", _replacer, text)

_ENTITY_SPLIT_RE = re.compile(
    r"\bQQ\b|\band\s*/\s*or\s*|\s&\s|\s[–]\s|\+", flags=re.IGNORECASE
)
_ENTITY_SLASH_SPLIT_RE = re.compile(r"\s*/\s*")

def _split_insured_entities(text: str):
    protected = _protect_division_ampersand(text)
    t = _ENTITY_SPLIT_RE.sub("/", protected)
    parts = _ENTITY_SLASH_SPLIT_RE.split(t)
    parts = [p.replace("\x00AMP\x00", "&").strip() for p in parts]
    return [p for p in parts if p]

_ROMAN_AFTER_NUSANTARA_RE = re.compile(
    r"(?<![A-Za-z0-9])(NUSANTARA\s+)(X{0,3}(?:IX|IV|V?I{0,3}))(?![A-Za-z0-9])"
)
_ROMAN_VALUES = {"I": 1, "V": 5, "X": 10}

def _roman_to_int(token: str) -> int:
    total = 0
    for i, ch in enumerate(token):
        val = _ROMAN_VALUES[ch]
        if i + 1 < len(token) and _ROMAN_VALUES[token[i + 1]] > val:
            total -= val
        else:
            total += val
    return total

def _convert_roman_numerals(text: str) -> str:
    """Angka romawi 1-39 diubah jadi angka HANYA bila tepat setelah kata NUSANTARA
    (mis. 'PERKEBUNAN NUSANTARA IV REGIONAL 2' -> '... NUSANTARA 4 REGIONAL 2')."""
    def _repl(m):
        token = m.group(2)
        return m.group(1) + str(_roman_to_int(token)) if token else m.group(0)
    return _ROMAN_AFTER_NUSANTARA_RE.sub(_repl, text)

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []
    original = text.strip()

    if _BYPASS_ALL_RE.search(original) or _INSURED_DESCRIPTIVE_RE.search(original):
        return [original]

    raw_parts = _split_insured_entities(original)
    if not raw_parts:
        return [original]

    if len(raw_parts) > MAX_BREAKDOWN_CODES:
        cleaned_parts = []
        for part in raw_parts:
            cleaned = _strip_legal_entity_with_punct(part).strip()
            cleaned_parts.append(cleaned if cleaned else part.strip())
        return [_convert_roman_numerals("/".join(cleaned_parts).upper())]

    results = []
    for part in raw_parts:
        if _BANK_NAME_RE.match(_strip_legal_entity_with_punct(part).strip()):
            continue
        cleaned = _final_polish_insured(_strip_titles(_strip_legal_entity_with_punct(part)))
        if not cleaned or len(cleaned) < 2 or not re.search(r"[A-Z]", cleaned):
            continue
        cleaned = _convert_roman_numerals(cleaned)
        if cleaned not in results:
            results.append(cleaned)

    if not results:
        return [original]

    return [original] if len(results) > MAX_BREAKDOWN_CODES else results

# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO
# ==============================================================================
_COMMA_DOT_RE = re.compile(r"[,.]")
_REF_SUFFIX_RE = re.compile(r"^P\d*$|^VAR(?:IOUS)?$|^TBA\.?$", flags=re.IGNORECASE)
_TBA_ALONE_RE = re.compile(r"^\s*TBA\.?\s*$", flags=re.IGNORECASE)
_VAR_ALONE_RE = re.compile(r"^\s*VAR(?:IOUS)?\s*$", flags=re.IGNORECASE)
_P_ALONE_RE = re.compile(r"^\s*P\d*\s*$", flags=re.IGNORECASE)
_SEE_ATTACHMENT_ALONE_RE = re.compile(r"^\s*SEE\s+ATTACHMENT\s*$", flags=re.IGNORECASE)
_CURRENCY_ALONE_RE = re.compile(r"^\s*(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\s*$", flags=re.IGNORECASE)

_BYPASS_KEEP_ASIS_RE = re.compile(
    r"\bPENYELESAIAN\s+(?:HUTANG|UTANG)\s+PIUTANG\b|\bPENYELESAIAN\s+SUSPEND\b|\bSUSPENSE\b",
    flags=re.IGNORECASE,
)
_DATE_ONLY_RE = re.compile(r"^\s*.*\bTAHUN\b.*\b(?:19|20)\d{2}\b.*$", flags=re.IGNORECASE)
_POLICY_SLIP_SPLIT_RE = re.compile(r"\s*\+\s*")

_MAIN_CODE_RE = re.compile(r"^[A-Za-z]{1,3}\d{9,13}$")
_MAIN_DIGITS_ONLY_RE = re.compile(r"^\d{9,14}$")
_MAIN_ALNUM_MIXED_RE = re.compile(r"^[A-Za-z0-9]{9,16}$")

def _is_valid_main_code(s: str) -> bool:
    s = s.strip()
    if _MAIN_CODE_RE.match(s) or _MAIN_DIGITS_ONLY_RE.match(s):
        return True
    if _MAIN_ALNUM_MIXED_RE.match(s) and re.search(r"\d", s):
        return True
    return False

def _is_kept_as_is(value: str) -> bool:
    v = value.strip()
    if (_TBA_ALONE_RE.match(v) or _VAR_ALONE_RE.match(v) or _P_ALONE_RE.match(v)
            or _SEE_ATTACHMENT_ALONE_RE.match(v) or _CURRENCY_ALONE_RE.match(v)):
        return True
    if _BYPASS_KEEP_ASIS_RE.search(v) or _DATE_ONLY_RE.match(v):
        return True
    return False

_TAHUN_WORD_RE = re.compile(r"\bTAHUN\b", flags=re.IGNORECASE)
_EMBEDDED_MAIN_CODE_RE = re.compile(r"(?<![A-Za-z0-9])([A-Za-z]{1,3}\d{9,13})(?![A-Za-z0-9])")

def _try_extract_code_from_tahun_noise(text: str):
    """Teks berisi kode polis yang sesuai karakteristik + keterangan 'TAHUN ...'
    (mis. 'IP071224000144 - TAHUN 2024 KE 2025') -> ambil kodenya saja.
    Hanya jika ada tepat 1 kode unik; kalau ambigu dibiarkan (return None)."""
    if not _TAHUN_WORD_RE.search(text) or _BYPASS_KEEP_ASIS_RE.search(text) or "+" in text:
        return None
    codes = []
    for c in _EMBEDDED_MAIN_CODE_RE.findall(text):
        c = c.strip()
        if c not in codes:
            codes.append(c)
    if len(codes) != 1:
        return None
    return [codes[0]]

_TRAILING_TBA_VAR_SPACE_RE = re.compile(r"^(.+?)\s+(?:TBA\.?|VAR(?:IOUS)?)\s*$", flags=re.IGNORECASE)

def _try_strip_trailing_tba_var_space(text: str):
    match = _TRAILING_TBA_VAR_SPACE_RE.match(text.strip())
    if not match:
        return None
    remainder = match.group(1).strip()
    cleaned = _clean_single_code(remainder)
    if _is_valid_main_code(cleaned):
        return [cleaned]
    return None

def _clean_single_code(value: str) -> str:
    s = value.strip()
    s = re.sub(r",\s*P\d*\s*$", "", s, flags=re.IGNORECASE)
    s = _COMMA_DOT_RE.sub("", s)
    s = re.sub(r"\s{2,}", " ", s).strip()
    return s

def _strip_trailing_ref_codes(parts: list) -> list:
    return [p for p in parts if not _REF_SUFFIX_RE.match(p.strip())]

def _rebuild_number_with_base(base: str, tail: str) -> str:
    base_digit_match = re.search(r"\d+\s*$", base)
    if not base_digit_match:
        return base + tail
    base_digits = base_digit_match.group(0).strip()
    base_prefix = base[: base_digit_match.start()]
    if len(tail) >= len(base_digits):
        return base_prefix + tail
    return base_prefix + base_digits[: -len(tail)] + tail

_CNR_SEP = r"[/.]"
_CNR_MULTI_RE = re.compile(
    rf"^\s*\d+\s*{_CNR_SEP}\s*CNR\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{2,4}}"
    rf"(?:\s*\+\s*\d+\s*{_CNR_SEP}\s*CNR\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{2,4}})+\s*$",
    flags=re.IGNORECASE,
)

def _try_split_multi_cnr_keep_asis(text: str):
    stripped = text.strip()
    if not _CNR_MULTI_RE.match(stripped):
        return None
    parts = [p.strip() for p in stripped.split("+")]
    parts = [p for p in parts if p]
    return parts if parts else None

_LEADING_P_PREFIX_RE = re.compile(
    r"^\s*P\d*\s*/\s*|^\s*P\d*\s+(?=[A-Za-z]{0,3}\d{6})", flags=re.IGNORECASE
)

def _strip_leading_p_prefix(text: str) -> str:
    return _LEADING_P_PREFIX_RE.sub("", text)

_SD_RANGE_RE = re.compile(r"^(.+?)\s+SD\s+(\d{1,6})\s*$", flags=re.IGNORECASE)

def _try_expand_sd_range(text: str):
    match = _SD_RANGE_RE.match(text.strip())
    if not match:
        return None
    base_full, end_suffix = match.groups()
    base_full = base_full.strip()
    base_digit_match = re.search(r"(\d+)\s*$", base_full)
    if not base_digit_match:
        return None
    base_digits = base_digit_match.group(1)
    if len(end_suffix) > len(base_digits):
        return None
    start_num_str = base_digits[-len(end_suffix):]
    try:
        start, end = int(start_num_str), int(end_suffix)
    except ValueError:
        return None
    if start > end or (end - start + 1) > MAX_BREAKDOWN_CODES:
        return None
    width = len(end_suffix)
    prefix = base_full[: base_digit_match.start(1)] + base_digits[: -width]
    codes = [_clean_single_code(f"{prefix}{n:0{width}d}") for n in range(start, end + 1)]
    deduped = []
    for c in codes:
        if c not in deduped:
            deduped.append(c)
    return deduped

def _try_strip_narrative_after_slash(text: str):
    if "/" not in text:
        return None
    parts = [p.strip() for p in text.split("/")]
    if len(parts) < 2:
        return None
    first = parts[0]
    if not _is_valid_main_code(_clean_single_code(first)):
        return None
    rest_parts = parts[1:]
    if any(_is_valid_main_code(_clean_single_code(p)) for p in rest_parts):
        return None
    return [_clean_single_code(first)]

_DASH_BASE_RE = re.compile(r"^([A-Za-z]{0,3}\d{6,14})-(\d{1,6}(?:\s*\+\s*\d{1,6})*)\s*(.*)$")

def _try_expand_dash_suffix(text: str):
    stripped = text.strip()
    match = _DASH_BASE_RE.match(stripped)
    if not match:
        return None
    base, suffix_blob, rest = match.groups()

    base_clean = _clean_single_code(base)
    has_alpha_prefix = bool(re.match(r"^[A-Za-z]", base_clean))

    suffixes = [s.strip() for s in suffix_blob.split("+") if s.strip()]
    codes = [base_clean] if not has_alpha_prefix else []
    for suf in suffixes:
        if has_alpha_prefix:
            codes.append(base_clean + suf)
        else:
            codes.append(_clean_single_code(_rebuild_number_with_base(base_clean, suf)))

    rest = rest.strip()
    if rest.startswith("+"):
        rest = rest[1:].strip()
    if rest:
        extra_parts = [p.strip() for p in rest.split("+") if p.strip()]
        for p in extra_parts:
            codes.append(_clean_single_code(p))

    deduped = []
    for c in codes:
        if c not in deduped:
            deduped.append(c)
    return deduped if len(deduped) <= MAX_BREAKDOWN_CODES else None

def _split_and_rebuild_plus(cleaned_noise: str):
    parts = _POLICY_SLIP_SPLIT_RE.split(cleaned_noise)
    parts = [p.strip() for p in parts if p.strip()]

    if len(parts) <= 1:
        single = _clean_single_code(cleaned_noise)
        return [single], not _is_valid_main_code(single)

    parts = _strip_trailing_ref_codes(parts)
    if len(parts) == 0:
        single = _clean_single_code(cleaned_noise)
        return [single], not _is_valid_main_code(single)
    if len(parts) == 1:
        single = _clean_single_code(parts[0])
        return [single], not _is_valid_main_code(single)

    total_pieces = 1
    for tail in parts[1:]:
        sub_suffixes = [s for s in tail.split("_") if s]
        total_pieces += max(len(sub_suffixes), 1)

    if total_pieces > MAX_BREAKDOWN_CODES:
        return [cleaned_noise.strip()], True

    base = parts[0]
    codes = [_clean_single_code(base)]
    for tail in parts[1:]:
        sub_suffixes = [s for s in tail.split("_") if s] or [tail]
        for suf in sub_suffixes:
            if re.fullmatch(r"\d+", suf) and len(suf) <= 6 and len(suf) < len(base):
                rebuilt = _rebuild_number_with_base(base, suf)
                codes.append(_clean_single_code(rebuilt))
            else:
                codes.append(_clean_single_code(suf))

    deduped = []
    for c in codes:
        if c not in deduped:
            deduped.append(c)
    return deduped, False

_TRAILING_CERT_SUFFIX_RE = re.compile(
    r"^([A-Za-z]{1,3}\d{9,13}|\d{9,14})\s*[-\s]\s*(\d{5,6})\s*$"
)

def _try_strip_trailing_certificate_suffix(text: str):
    match = _TRAILING_CERT_SUFFIX_RE.match(text.strip())
    if not match:
        return None
    base, _suffix = match.groups()
    base_clean = _clean_single_code(base)
    if _is_valid_main_code(base_clean):
        return [base_clean]
    return None

_DASH_VAR_SUFFIX_RE = re.compile(r"^(.+?)\s*-\s*VAR(?:IOUS)?\s*$", flags=re.IGNORECASE)

def _try_strip_trailing_dash_var(text: str):
    match = _DASH_VAR_SUFFIX_RE.match(text.strip())
    if not match:
        return None
    remainder = match.group(1).strip()
    cleaned = _clean_single_code(remainder)
    if _is_valid_main_code(cleaned):
        return [cleaned]
    return None

_COMMA_BREAKDOWN_RE = re.compile(r"^([A-Za-z]{1,3}\d{6,14})\s*,\s*(\d{1,6}(?:\s*,\s*\d{1,6})*)\s*$")

def _try_expand_comma_suffix(text: str):
    match = _COMMA_BREAKDOWN_RE.match(text.strip())
    if not match:
        return None
    base, suffix_blob = match.groups()
    suffixes = [s.strip() for s in suffix_blob.split(",") if s.strip()]
    if not suffixes:
        return None
    total_pieces = 1 + len(suffixes)
    if total_pieces > MAX_BREAKDOWN_CODES:
        return None
    codes = [_clean_single_code(base)]
    for suf in suffixes:
        rebuilt = _rebuild_number_with_base(base, suf)
        codes.append(_clean_single_code(rebuilt))
    deduped = []
    for c in codes:
        if c not in deduped:
            deduped.append(c)
    return deduped

_CNR_SINGLE_TRAILING_SUFFIX_PLUS_RE = re.compile(
    rf"^\s*(\d{{4,14}})((?:\s*\+\s*\d{{1,6}})+)\s*({_CNR_SEP}\s*CNR\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{1,2}}\s*{_CNR_SEP}\s*\d{{2,4}})\s*$",
    flags=re.IGNORECASE,
)

def _try_expand_plus_with_trailing_cnr(text: str):
    match = _CNR_SINGLE_TRAILING_SUFFIX_PLUS_RE.match(text.strip())
    if not match:
        return None
    base, suffix_blob, cnr_suffix = match.groups()
    suffixes = [s.strip() for s in suffix_blob.split("+") if s.strip()]
    if not suffixes:
        return None
    total_pieces = 1 + len(suffixes)
    if total_pieces > MAX_BREAKDOWN_CODES:
        return None
    numbers = [base] + [_rebuild_number_with_base(base, suf) for suf in suffixes]
    deduped_numbers = []
    for n in numbers:
        if n not in deduped_numbers:
            deduped_numbers.append(n)
    return [f"{n}{cnr_suffix.strip()}" for n in deduped_numbers]

def clean_split_code_impl(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return [], True
    original = text.strip()

    tahun_noise_result = _try_extract_code_from_tahun_noise(original)
    if tahun_noise_result is not None:
        return tahun_noise_result, False

    if _is_kept_as_is(original):
        return [original], True

    plus_cnr_reattach_result = _try_expand_plus_with_trailing_cnr(original)
    if plus_cnr_reattach_result is not None:
        return plus_cnr_reattach_result, True

    multi_cnr_result = _try_split_multi_cnr_keep_asis(original)
    if multi_cnr_result is not None:
        return multi_cnr_result, True

    trailing_tba_result = _try_strip_trailing_tba_var_space(original)
    if trailing_tba_result is not None:
        return trailing_tba_result, False

    working = _strip_leading_p_prefix(original)
    if working != original and not working.strip():
        return [original], True

    trailing_cert_result = _try_strip_trailing_certificate_suffix(working)
    if trailing_cert_result is not None:
        return trailing_cert_result, False

    multi_pair_result = _try_split_polis_certificate_pairs(working)
    if multi_pair_result is not None:
        return [polis for polis, _cert in multi_pair_result], False

    dash_var_result = _try_strip_trailing_dash_var(working)
    if dash_var_result is not None:
        return dash_var_result, False

    if "CNR" in working.upper():
        return [original], True

    sd_result = _try_expand_sd_range(working)
    if sd_result is not None:
        return sd_result, False

    comma_result = _try_expand_comma_suffix(working)
    if comma_result is not None:
        return comma_result, False

    narrative_result = _try_strip_narrative_after_slash(working)
    if narrative_result is not None:
        return narrative_result, False

    if "-" in working:
        dash_result = _try_expand_dash_suffix(working)
        if dash_result is not None:
            return dash_result, False

    if "/" in working and "+" not in working and "-" not in working:
        slash_parts = [p.strip() for p in working.split("/") if p.strip()]
        cleaned_slash_parts = [_clean_single_code(p) for p in slash_parts]
        if len(slash_parts) >= 2 and all(_is_valid_main_code(p) for p in cleaned_slash_parts):
            if len(cleaned_slash_parts) <= MAX_BREAKDOWN_CODES:
                deduped = []
                for c in cleaned_slash_parts:
                    if c not in deduped:
                        deduped.append(c)
                return deduped, False

    return _split_and_rebuild_plus(working)

def clean_split_code(text: str) -> list:
    results, _ = clean_split_code_impl(text)
    return results

# ==============================================================================
# 3. PEMILIHAN SUMBER TERBAIK: CLSDT vs FAC (POLIS & SLIP)
# ==============================================================================
def _is_blank_value(value) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    s = str(value).strip()
    return s == "" or s.lower() in ("nan", "none")

def _is_valid_clean_result(codes: list) -> bool:
    if not codes:
        return False
    return all(_is_valid_main_code(c) for c in codes)

def _fac_has_bypass_keyword(fac_value) -> bool:
    if _is_blank_value(fac_value):
        return False
    return bool(_BYPASS_KEEP_ASIS_RE.search(str(fac_value).strip()))

_TRAILING_CURRENCY_CODE_RE = re.compile(
    r"\s+(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\s*$", flags=re.IGNORECASE
)

def _strip_trailing_currency(text: str) -> str:
    return _TRAILING_CURRENCY_CODE_RE.sub("", text).strip()

_GENERIC_CLSDT_VALUE_RE = re.compile(
    r"^\s*(?:TBA\.?|VAR(?:IOUS)?|P\d*|-|SEE\s+ATTACHMENT)\s*$", flags=re.IGNORECASE
)

def _is_generic_clsdt_value(value) -> bool:
    if _is_blank_value(value):
        return True
    return bool(_GENERIC_CLSDT_VALUE_RE.match(str(value).strip()))

def _resolve_polis_or_slip(clsdt_value, fac_value):
    clsdt_codes = clean_split_code(str(clsdt_value)) if not _is_blank_value(clsdt_value) else []

    if _is_valid_clean_result(clsdt_codes) and _fac_has_bypass_keyword(fac_value):
        clsdt_single = clsdt_codes[0] if len(clsdt_codes) == 1 else " ".join(clsdt_codes)
        fac_text = _strip_trailing_currency(str(fac_value).strip())
        combined = f"{clsdt_single} {fac_text}"
        return [combined], True

    if _is_valid_clean_result(clsdt_codes):
        return clsdt_codes, False

    if clsdt_codes and not _is_generic_clsdt_value(clsdt_value):
        return clsdt_codes, True

    fac_codes = clean_split_code(str(fac_value)) if not _is_blank_value(fac_value) else []
    if fac_codes:
        return fac_codes, not _is_valid_clean_result(fac_codes)

    if clsdt_codes:
        return clsdt_codes, True

    return [], True

# ==============================================================================
# 3b. EKSTRAKSI CERTIFICATE
# ==============================================================================
_CERT_VALID_PREFIX_CODE_RE = re.compile(r"^[A-Za-z]{1,3}\d{9,13}$|^\d{9,14}$")
_SD_NOTATION_RE = re.compile(r"\bS\s*/\s*D\b", flags=re.IGNORECASE)

def _normalize_sd_notation(text: str) -> str:
    return _SD_NOTATION_RE.sub("SD", text)

def _pad_certificate_digits(digit_str: str):
    d = digit_str.strip()
    if not d.isdigit():
        return None
    if len(d) > CERTIFICATE_DIGIT_WIDTH:
        return None
    return d.zfill(CERTIFICATE_DIGIT_WIDTH)

def _is_certificate_like_clsdt(value: str) -> bool:
    v = value.strip()
    if not v:
        return False
    v_norm = _normalize_sd_notation(v)
    if re.search(r"\bSD\b", v_norm, flags=re.IGNORECASE) or "," in v_norm:
        return True
    cleaned = re.sub(r"[.\-\s]", "", v_norm)
    return cleaned.isdigit() and len(cleaned) <= CERTIFICATE_DIGIT_WIDTH + 2

def _format_certificate_range(codes: list) -> str:
    deduped = []
    for c in codes:
        if c not in deduped:
            deduped.append(c)
    deduped.sort(key=lambda x: int(x) if x.isdigit() else x)
    if len(deduped) == 1:
        return deduped[0]
    if len(deduped) > CERTIFICATE_RANGE_SD_THRESHOLD:
        return f"{deduped[0]} SD {deduped[-1]}"
    return ", ".join(deduped)

def _extract_certificate_from_clsdt(clsdt_sertf_value) -> str:
    if _is_blank_value(clsdt_sertf_value):
        return None

    raw = str(clsdt_sertf_value).strip()
    normalized = _normalize_sd_notation(raw)

    sd_match = re.match(r"^(\d{1,6})\s*SD\s*(\d{1,6})\s*$", normalized, flags=re.IGNORECASE)
    if sd_match:
        start_s, end_s = sd_match.groups()
        start_p, end_p = _pad_certificate_digits(start_s), _pad_certificate_digits(end_s)
        if start_p is None or end_p is None:
            return ""
        try:
            start_n, end_n = int(start_p), int(end_p)
        except ValueError:
            return ""
        if start_n > end_n:
            return ""
        codes = [str(n).zfill(CERTIFICATE_DIGIT_WIDTH) for n in range(start_n, end_n + 1)]
        return _format_certificate_range(codes)

    if "," in normalized:
        parts = [p.strip() for p in normalized.split(",") if p.strip()]
        padded = [_pad_certificate_digits(p) for p in parts]
        if not padded or any(p is None for p in padded):
            return ""
        return _format_certificate_range(padded)

    if not _is_certificate_like_clsdt(raw):
        return ""
    cleaned = re.sub(r"[.\-\s]", "", normalized)
    padded = _pad_certificate_digits(cleaned)
    if padded is None:
        return ""
    return padded

_CERT_TRAILING_SD_RE = re.compile(
    r"^([A-Za-z]{1,3}\d{9,13}|\d{9,14})\s*[-\s]\s*(\d{6})\s+SD\s+(\d{1,6})\s*$",
    flags=re.IGNORECASE,
)
_CERT_TRAILING_COMMA_RE = re.compile(
    r"^([A-Za-z]{1,3}\d{9,13}|\d{9,14})\s*[-\s]\s*(\d{5,6}(?:\s*,\s*\d{1,6})+)\s*$"
)
_CERT_TRAILING_SINGLE_RE = re.compile(
    r"^([A-Za-z]{1,3}\d{9,13}|\d{9,14})\s*[-\s]\s*(\d{5,6})\s*$"
)

def _extract_certificate_from_fac(fac_value) -> str:
    if _is_blank_value(fac_value):
        return ""
    raw = str(fac_value).strip()
    normalized = _normalize_sd_notation(raw)
    normalized = _strip_leading_p_prefix(normalized).strip()

    if _BYPASS_KEEP_ASIS_RE.search(normalized) or _is_kept_as_is(normalized) or "+" in normalized:
        return ""

    sd_match = _CERT_TRAILING_SD_RE.match(normalized)
    if sd_match:
        _base, start_s, end_s = sd_match.groups()
        start_p, end_p = _pad_certificate_digits(start_s), _pad_certificate_digits(end_s)
        if start_p is None or end_p is None:
            return ""
        try:
            start_n, end_n = int(start_p), int(end_p)
        except ValueError:
            return ""
        if start_n > end_n or (end_n - start_n + 1) > MAX_BREAKDOWN_CODES:
            return ""
        codes = [str(n).zfill(CERTIFICATE_DIGIT_WIDTH) for n in range(start_n, end_n + 1)]
        return _format_certificate_range(codes)

    comma_match = _CERT_TRAILING_COMMA_RE.match(normalized)
    if comma_match:
        _base, suffix_blob = comma_match.groups()
        parts = [p.strip() for p in suffix_blob.split(",") if p.strip()]
        padded = [_pad_certificate_digits(p) for p in parts]
        if not padded or any(p is None for p in padded) or len(padded) > MAX_BREAKDOWN_CODES:
            return ""
        return _format_certificate_range(padded)

    single_match = _CERT_TRAILING_SINGLE_RE.match(normalized)
    if single_match:
        _base, cert_digits = single_match.groups()
        padded = _pad_certificate_digits(cert_digits)
        if padded is None:
            return ""
        return padded

    return ""

def extract_certificate(fac_policy_value, clsdt_sertf_value=None) -> str:
    if clsdt_sertf_value is not None:
        clsdt_result = _extract_certificate_from_clsdt(clsdt_sertf_value)
        if clsdt_result is not None:
            return clsdt_result

    return _extract_certificate_from_fac(fac_policy_value)

_CERT_PAIR_SLASH_SPLIT_RE = re.compile(r"\s*/\s*")
_CERT_PAIR_SINGLE_RE = re.compile(
    r"^([A-Za-z]{1,3}\d{9,13}|\d{9,14})\s*[-\s]\s*(\d{5,6})\s*$"
)

def _try_split_polis_certificate_pairs(text: str):
    stripped = text.strip()
    if "/" not in stripped:
        return None
    segments = [s.strip() for s in _CERT_PAIR_SLASH_SPLIT_RE.split(stripped) if s.strip()]
    if len(segments) < 2:
        return None

    pairs = []
    for seg in segments:
        seg_norm = _normalize_sd_notation(seg)
        match = _CERT_PAIR_SINGLE_RE.match(seg_norm)
        if not match:
            return None
        polis_raw, cert_raw = match.groups()
        polis_clean = _clean_single_code(polis_raw)
        if not _is_valid_main_code(polis_clean):
            return None
        cert_padded = _pad_certificate_digits(cert_raw)
        if cert_padded is None:
            return None
        pairs.append((polis_clean, cert_padded))

    if len(pairs) > MAX_BREAKDOWN_CODES:
        return None
    return pairs

def extract_certificate_pairs(fac_policy_value, clsdt_sertf_value=None, polis_cln_count=None):
    def _pad_to_count(values: list) -> list:
        if polis_cln_count is None:
            return values
        if len(values) >= polis_cln_count:
            return values[:polis_cln_count] if polis_cln_count > 0 else values
        return values + [""] * (polis_cln_count - len(values))

    if clsdt_sertf_value is not None:
        clsdt_result = _extract_certificate_from_clsdt(clsdt_sertf_value)
        if clsdt_result is not None:
            return _pad_to_count([clsdt_result])

    if _is_blank_value(fac_policy_value):
        return _pad_to_count([])

    raw = str(fac_policy_value).strip()
    working = _strip_leading_p_prefix(_normalize_sd_notation(raw)).strip()

    multi_pairs = _try_split_polis_certificate_pairs(working)
    if multi_pairs is not None:
        return _pad_to_count([cert for _polis, cert in multi_pairs])

    single_cert = _extract_certificate_from_fac(raw)
    return _pad_to_count([single_cert] if single_cert else [""])

# ==============================================================================
# 4. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw: pd.DataFrame, keyword: str, cedant_column: str = "") -> pd.DataFrame:
    if cedant_column:
        if cedant_column not in df_raw.columns:
            raise ValueError(f"Kolom {cedant_column} tidak ada.")
        cedant_cols = [cedant_column]
    else:
        if not (cedant_cols := [c for c in df_raw.columns if "CEDANT" in str(c).upper()]):
            raise ValueError("Kolom CEDANT tidak ditemukan lewat auto-detect.")

    mask = pd.Series(False, index=df_raw.index)
    for col in cedant_cols:
        mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)

    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant '{keyword}': {len(df_filtered)} dari total {len(df_raw)} baris.")
    return df_filtered

# ==============================================================================
# 5. BUILD OUTPUT
# ==============================================================================
_NON_ALNUM_RE = re.compile(r"[^A-Za-z0-9]")

def _compact_value(value: str) -> str:
    if not isinstance(value, str):
        return value
    return _NON_ALNUM_RE.sub("", value)

def _compact_list(values: list) -> list:
    return [_compact_value(v) for v in values]

def _compact_split_code_results(results: list, is_kept_as_is: bool) -> list:
    if is_kept_as_is:
        return results
    return [_compact_value(v) for v in results]

def _detect_key_columns(df: pd.DataFrame, insured_col: str, polis_col: str, slip_col: str):
    def _res(col: str, kw_or_fn, lbl: str):
        if col:
            if col not in df.columns:
                raise ValueError(f"{lbl}='{col}' tidak ada.")
            return col
        return next(
            (c for c in df.columns if (kw_or_fn(str(c).upper()) if callable(kw_or_fn) else kw_or_fn in str(c).upper())),
            None,
        )

    col_ins = _res(insured_col, "INSURED", "INSURED_COLUMN")
    col_pol = _res(polis_col, lambda c: c.strip() == "POLIS", "POLIS_COLUMN")
    col_slp = _res(slip_col, "SLIP", "SLIP_COLUMN")

    if None in (col_ins, col_pol, col_slp):
        raise ValueError("Kolom kunci tidak lengkap.")

    print(f"[INFO] Kolom terdeteksi -> INSURED='{col_ins}', POLIS='{col_pol}', SLIP='{col_slp}'")
    return col_ins, col_pol, col_slp

def _breakdown_all_rows(
    df: pd.DataFrame,
    col_ins: str,
    col_pol: str,
    col_slp: str,
    col_clsdt_pol: str = None,
    col_clsdt_slp: str = None,
    col_clsdt_sertf: str = None,
):
    ins_arr = df[col_ins].to_numpy()
    pol_arr = df[col_pol].to_numpy()
    slp_arr = df[col_slp].to_numpy()

    clsdt_pol_arr = df[col_clsdt_pol].to_numpy() if col_clsdt_pol and col_clsdt_pol in df.columns else [None] * len(df)
    clsdt_slp_arr = df[col_clsdt_slp].to_numpy() if col_clsdt_slp and col_clsdt_slp in df.columns else [None] * len(df)
    clsdt_sertf_arr = df[col_clsdt_sertf].to_numpy() if col_clsdt_sertf and col_clsdt_sertf in df.columns else [None] * len(df)

    ins_all, pol_all, slp_all, cert_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val, clsdt_pol_val, clsdt_slp_val, clsdt_sertf_val in zip(
        ins_arr, pol_arr, slp_arr, clsdt_pol_arr, clsdt_slp_arr, clsdt_sertf_arr
    ):
        ins = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []

        raw_pol = pol_val if pd.notna(pol_val) else ""
        raw_slp = slp_val if pd.notna(slp_val) else ""

        if col_clsdt_pol and pd.notna(clsdt_pol_val):
            pol_results, pol_kept_as_is = _resolve_polis_or_slip(clsdt_pol_val, raw_pol)
        else:
            pol_results, pol_kept_as_is = clean_split_code_impl(str(raw_pol).strip()) if not _is_blank_value(raw_pol) else ([], True)
        pol = _compact_split_code_results(pol_results, pol_kept_as_is)

        if col_clsdt_slp and pd.notna(clsdt_slp_val):
            slp_results, slp_kept_as_is = _resolve_polis_or_slip(clsdt_slp_val, raw_slp)
        else:
            slp_results, slp_kept_as_is = clean_split_code_impl(str(raw_slp).strip()) if not _is_blank_value(raw_slp) else ([], True)
        slp = _compact_split_code_results(slp_results, slp_kept_as_is)

        clsdt_sertf_for_row = clsdt_sertf_val if (col_clsdt_sertf and pd.notna(clsdt_sertf_val)) else None
        cert = extract_certificate_pairs(raw_pol, clsdt_sertf_for_row, polis_cln_count=len(pol))

        ins_all.append(ins)
        pol_all.append(pol)
        slp_all.append(slp)
        cert_all.append(cert)

        if len(ins) > max_ins: max_ins = len(ins)
        if len(pol) > max_pol: max_pol = len(pol)
        if len(slp) > max_slp: max_slp = len(slp)

    return ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_ins, col_pol, col_slp = _detect_key_columns(df_filtered, INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN)

    col_clsdt_pol = CLSDT_POLIS_COLUMN if CLSDT_POLIS_COLUMN in df_filtered.columns else None
    col_clsdt_slp = CLSDT_SLIP_COLUMN if CLSDT_SLIP_COLUMN in df_filtered.columns else None
    col_clsdt_sertf = CLSDT_SERTF_COLUMN if CLSDT_SERTF_COLUMN in df_filtered.columns else None
    
    if col_clsdt_pol or col_clsdt_slp:
        print(f"[INFO] Kolom CLSDT terdeteksi -> POLIS='{col_clsdt_pol}', SLIP='{col_clsdt_slp}' (prioritas utama)")
    else:
        print("[INFO] Kolom CLSDT tidak ditemukan di file ini -> pakai FAC_POLICY_NO/FAC_SLIP saja.")
    
    if col_clsdt_sertf:
        print(f"[INFO] Kolom CLSDT_SERTF_NO terdeteksi -> '{col_clsdt_sertf}' (prioritas utama utk CERTIFICATE)")
    else:
        print("[INFO] Kolom CLSDT_SERTF_NO tidak ditemukan -> CERTIFICATE diekstrak dari FAC_POLICY_NO saja.")

    ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp = _breakdown_all_rows(
        df_filtered, col_ins, col_pol, col_slp, col_clsdt_pol, col_clsdt_slp, col_clsdt_sertf
    )

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)
    output_cols_data = {}

    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == col_ins:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in ins_all]
        if col == col_pol:
            for i in range(1, max(max_pol, 1) + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in pol_all]
                output_cols_data[f"CERTIFICATE_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_all]
        if col == col_slp:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slp_all]

    df_master = pd.DataFrame(output_cols_data)
    drop_cols = [
        c for c in df_master.columns
        if ("_CLN_" in c or (c.startswith("CERTIFICATE_") and c != "CERTIFICATE_1"))
        and df_master[c].astype(str).str.strip().eq("").all()
    ]
    return df_master.drop(columns=drop_cols)

# ==============================================================================
# 6. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook, text_col_indices = load_workbook(output_path), [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]
    worksheet = workbook.active

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)
    workbook.save(output_path)

# ==============================================================================
# 7. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70 + "\n CLEANSING DATA 2 - OSBAL | CEDANT: JASA TANIA\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
        if not candidates:
            raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")
        input_path = candidates[0]
        print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    df_jt = filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)

    if not df_jt.empty:
        df_hasil = build_output(df_jt)
        save_with_text_format(df_hasil, OUTPUT_FILE)
        print("=" * 70 + f"\n[SUCCESS] Selesai. Output: '{OUTPUT_FILE}' ({len(df_hasil)} baris)\n" + "=" * 70)
    else:
        print("[WARNING] Tidak ada baris yang cocok dengan filter.")

if __name__ == "__main__":
    main()