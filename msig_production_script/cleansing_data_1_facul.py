"""
================================================================================
 SCRIPT CLEANSING DATA 1 - FACUL (CEDANT: MSIG )
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI UTAMA
# ==============================================================================
INPUT_FILE = "1b. Transaksi Facul.xlsx"
SHEET_NAME = "Query result"
HEADER_ROW = 0
OUTPUT_FILE = "msig_output_facul.xlsx"
CEDANT_FILTER = "PT. ASURANSI MSIG INDONESIA"

CEDANT_COLUMN = "COMP_NAME"
BROKER_NAME_COLUMN = "COMP_NAME.1"
DIRECT_MARKER = "DIRECT"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"

MAX_BREAKDOWN_CODES = 5
TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CERTIFICATE"
)
RE_I = re.IGNORECASE


# ==============================================================================
# 1. CONSTANTS & PRE-COMPILED REGEXES
# ==============================================================================
_TITLE_RE = re.compile(
    r"\b(?:S\.?H\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|"
    r"Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu(?!\s+KOTA)|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b",
    RE_I
)
_LEGAL_ENTITY_RE = re.compile(
    r"\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|CO\.?|PERSERODA|PERSERO|Perseroan\s+Terbatas)\b",
    RE_I
)
_LEGAL_ENTITY_WITH_LEADING_PUNCT_RE = re.compile(
    r"[,.]?\s*\b(?:PT\.?(?!\w)|P\.T\.?|P\s+T(?!\w)|CV\.?(?!\w)|C\.V\.?|TBK\.?(?!\w)|Tbk\.?(?!\w)|LTD\.?(?!\w)|LIMITED|PTE\.?(?!\w)|INC\.?(?!\w)|CO\.?(?!\w)|PERSERODA(?!\w)|PERSERO(?!\w)|Perseroan\s+Terbatas)",
    RE_I
)
_ENTITY_SPLIT_RE = re.compile(r"\bQQ+\.?\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|\bAND\b|;|/|,|:", RE_I)
_BARE_SUFFIX_FRAGMENT_RE = re.compile(
    r"^\s*(?:PT\.?|P\.T\.?|P\s+T|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO|KSO)\s*$",
    RE_I
)
_BARE_SUFFIX_WITH_PAREN_RE = re.compile(
    r"^\s*(?:PT\.?|P\.T\.?|P\s+T|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO)?\s*"
    r"(\uE100\d+\uE101)\s*$",
    RE_I
)
_SUFFIX_THEN_DASH_NEW_ENTITY_RE = re.compile(
    r"^\s*(PT\.?|P\.T\.?|P\s+T|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO)"
    r"\s+-\s+(.+)$",
    RE_I
)
_PAREN_GROUP_RE = re.compile(r"\([^()]*\)")
_PAREN_TOKEN_RE = re.compile(r"\uE100(\d+)\uE101")
_FINAL_PUNCT_RE = re.compile(r"[,;:\"']")
_FINAL_DASH_RE = re.compile(r"(?<!\s)(?<!\d)-(?!\s)(?!\d)")
_FINAL_SLASH_RE = re.compile(r"/")
_FINAL_DOT_RE = re.compile(r"\.")
_FINAL_LPAREN_RE = re.compile(r"\(\s+")
_FINAL_RPAREN_RE = re.compile(r"\s+\)")
_MULTISPACE_RE = re.compile(r"\s+")
_FINAL_EDGE_RE = re.compile(r"^[.\s]+|[.\s]+$")
_FINAL_EMPTY_PAREN_RE = re.compile(r"\(\s*\)")
_DANGLING_PATTERN_RE = re.compile(
    r"^\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*"
    r"|\s*\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$",
    RE_I
)
_BARE_CLIENT_FRAGMENT_RE = re.compile(r"^\s*CLIENTS?\s*$", RE_I)
_INSURED_DESCRIPTIVE_RE = re.compile(r"\bVARIOUS\s+INSUREDS?\b", RE_I)
_LANJUTAN_NOISE_RE = re.compile(r"\bLANJUTAN\s+[A-Za-z0-9]+\b", RE_I)
_HAS_UPPER_RE = re.compile(r"[A-Z]")

_SUSPENSE_KEEP_AS_IS_RE = re.compile(
    r"\bPENYELESAIAN\s+SUSPENSE\s+MSIG\b|\bSUSPENSE\b|\bPENYELESAIAN\s+(?:H)?UTANG\s+PIUTANG\b",
    RE_I
)
_MSIG_GROUP_YEAR_RE = re.compile(r"\bMSIG\s+GROUP\b", RE_I)
_STANDALONE_KEEP_RE = re.compile(r"^\s*(VAR|VARIOUS|TBA|P\d{0,2})\s*$", RE_I)
_SEE_ATTACHMENT_RE = re.compile(r"^\s*SEE\s+ATTACHMENT\s*$", RE_I)

_PAREN_STASH: dict[str, str] = {}

_ACRONYM_UNWRAP_OVERRIDES = {
    ("INDONESIA ASAHAN ALUMINIUM", "INALUM"),
}

_INSURED_RESULT_OVERRIDES = {
    "DUTA INTIDAYA (A.S WATSON),PT": ["DUTA INTIDAYA", "A S WATSON"],
    "ASTRA DAIHATSU MOTOR,PT (HO) , I & IV": ["ASTRA DAIHATSU MOTOR (HO) 1", "ASTRA DAIHATSU MOTOR (HO) 4"],
    "AERONE ( DRONE )": ["AERONE DRONE"],
    "SAVINO DEL BENE ( INDONESIA )": ["SAVINO DEL BENE INDONESIA"],
    "PAKUWON GROUP / ELITE PRIMA HUTAMA / PPRS OFFICE 88 (TOWER A) / PPRS CASA GRANDE 1 / ANGELO BELLA": [
        "PAKUWON GROUP", "ELITE PRIMA HUTAMA", "PPRS OFFICE 88 TOWER A", "PPRS CASA GRANDE 1", "ANGELO BELLA"
    ],
    "BHUMI JATI POWER ( TANJUNG JATI B COAL UNIT 5 DAN 6 )": [
        "BHUMI JATI POWER TANJUNG JATI B COAL UNIT 5", "BHUMI JATI POWER TANJUNG JATI B COAL UNIT 6"
    ],
    "CENTRAL JAVA POWER ( SPC ) QQ SUMITOMO CORPORATION AND SUMMIT POWER DEV QQ PLN QQ EXPANSION LENDERS": [
        "CENTRAL JAVA POWER (SPC)", "SUMITOMO CORPORATION AND SUMMIT POWER DEV", "PLN", "EXPANSION LENDERS"
    ],
    "LUXOTTICA GROUP / ESSILORLUXOTTICA S.A / OPTICAL SUPPORT OF INDONESIA PLC, PT": [
        "LUXOTTICA GROUP", "ESSILORLUXOTTICA S A", "OPTICAL SUPPORT OF INDONESIA PLC"
    ],
    "JAYA REAL PROPERTY, PT / PUSAT GROSIR SENEN JAYA BLOK 1 AND 2": [
        "JAYA REAL PROPERTY PUSAT GROSIR SENEN JAYA BLOK 1", "JAYA REAL PROPERTY PUSAT GROSIR SENEN JAYA BLOK 2"
    ],
    "JAYA REAL PROPERTY, PT / PUSAT GROSIR SENEN JAYA BLOK BLOK IV & SENEN BLOK V": [
        "JAYA REAL PROPERTY PUSAT GROSIR SENEN JAYA BLOK BLOK 4", "JAYA REAL PROPERTY PUSAT GROSIR SENEN JAYA BLOK BLOK 5"
    ],
    "PEMBANGUNAN JAYA (SENEN JAYA 1 &2)": ["PEMBANGUNAN JAYA SENEN 1", "PEMBANGUNAN JAYA SENEN 2"],
    "PERFETTI VAN MELLE HOLDING B.V": ["PERFETTI VAN MELLE HOLDING B V"],
    "MSIG INDONESIA/MSIG FACILITY/GILAT SATELLITE NETWORKS (HOLLAND) B.V": [
        "MSIG INDONESIA", "MSIG FACILITY", "GILAT SATELLITE NETWORKS (HOLLAND) B V"
    ],
    "TANJUNG JATI A DAN B QQ CENTRAL JAVA POWER ( SPC ) QQ SUMITOMO CORPORATION AND SUMMIT POWER D QQ PLN": [
        "TANJUNG JATI A", "TANJUNG JATI B", "CENTRAL JAVA POWER ( SPC )", "SUMITOMO CORPORATION AND SUMMIT POWER D", "PLN"
    ],
    "PAITON ENERGY COMPANY UNIT 7-8": ["PAITON ENERGY COMPANY UNIT 7", "PAITON ENERGY COMPANY UNIT 8"],
    "MSIG INDONESIA/INDONESIA RE - IF P&C/SCANIA PARTS INDONESIA": [
        "MSIG INDONESIA", "INDONESIA RE", "IF P&C", "SCANIA PARTS INDONESIA"
    ],
    "PAITON ENERGY COMPANY,PT PAITON 7-8": ["PAITON ENERGY COMPANY PAITON 7", "PAITON ENERGY COMPANY PAITON 8"],
    "PAITON ENERGY UNIT 7-8, PT": ["PAITON ENERGY UNIT 7", "PAITON ENERGY UNIT 8"],
}


# ==============================================================================
# 2. CLEANSING INSURED
# ==============================================================================
def _strip_legal_entity_with_punct(text: str) -> str:
    t = _LEGAL_ENTITY_WITH_LEADING_PUNCT_RE.sub("", text)
    return re.sub(r",\s*(?=[(\s]|$)", " ", t)


def _strip_titles(text: str) -> str:
    return _TITLE_RE.sub(" ", text)


def _merge_bare_suffix_fragments(parts: list) -> list:
    merged = []
    for p in parts:
        p_stripped = p.strip()
        if merged and (_BARE_SUFFIX_FRAGMENT_RE.match(p_stripped) or _BARE_SUFFIX_WITH_PAREN_RE.match(p_stripped)):
            merged[-1] = merged[-1] + " " + p_stripped
            continue
        suffix_dash_match = _SUFFIX_THEN_DASH_NEW_ENTITY_RE.match(p_stripped) if merged else None
        if suffix_dash_match:
            merged[-1] = merged[-1] + " " + suffix_dash_match.group(1).strip()
            merged.append(suffix_dash_match.group(2).strip())
            continue
        merged.append(p)
    return merged


def _protect_parens_content(text: str) -> str:
    def _mask(match: re.Match) -> str:
        token = f"\uE100{len(_PAREN_STASH)}\uE101"
        _PAREN_STASH[token] = match.group(0)
        return token
    return _PAREN_GROUP_RE.sub(_mask, text)


def _restore_protected_parens(text: str) -> str:
    def _unmask(match: re.Match) -> str:
        return _PAREN_STASH.get(f"\uE100{match.group(1)}\uE101", match.group(0))
    return _PAREN_TOKEN_RE.sub(_unmask, text)


def _final_polish(text: str) -> str:
    t = _FINAL_PUNCT_RE.sub(" ", text)
    t = _FINAL_DASH_RE.sub(" ", t)
    t = _FINAL_SLASH_RE.sub(" ", t)
    t = _FINAL_DOT_RE.sub(" ", t)
    t = _FINAL_LPAREN_RE.sub("(", t)
    t = _FINAL_RPAREN_RE.sub(")", t)
    t = _MULTISPACE_RE.sub(" ", t).strip()
    t = _FINAL_EDGE_RE.sub("", t)
    t = _FINAL_EMPTY_PAREN_RE.sub("", t).strip()
    
    previous = None
    while previous != t:
        previous, t = t, _DANGLING_PATTERN_RE.sub("", t).strip()
        t = _FINAL_EDGE_RE.sub("", t).strip()
    return t.upper()


def _is_acronym_of(before_text: str, paren_text: str) -> bool:
    before_norm = _MULTISPACE_RE.sub(" ", before_text.upper()).strip()
    paren_norm = re.sub(r"[^A-Z]", "", paren_text.upper())
    if (before_norm, paren_norm) in _ACRONYM_UNWRAP_OVERRIDES:
        return True
    words = [w for w in before_norm.split() if w.isalpha()]
    if not words or not paren_norm:
        return False
    initials = "".join(w[0] for w in words)
    if paren_norm == initials:
        return True
    it = iter(initials)
    return all(ch in it for ch in paren_norm)


def _resolve_parens_split_or_unwrap(text: str) -> str:
    def _repl(match: re.Match) -> str:
        paren_content = match.group(1)
        before_text = text[: match.start()]
        before_text_clean = _LEGAL_ENTITY_WITH_LEADING_PUNCT_RE.sub("", before_text)
        before_segment = re.split(r"[,/;]", before_text_clean)[-1]
        if _is_acronym_of(before_segment, paren_content):
            return f"\uE200{paren_content}\uE201"
        return f" , {paren_content} , "
        
    resolved = re.sub(r"\(([^()]*)\)", _repl, text)
    resolved = re.sub(
        r"(?:,\s*(?:PT\.?|P\.T\.?|CV\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO)\s*)?\uE200([^\uE200\uE201]*)\uE201",
        r" (\1) ",
        resolved,
        flags=RE_I,
    )
    return resolved


def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []

    original = text.strip()
    if _INSURED_DESCRIPTIVE_RE.search(original):
        return [original]

    normalized_check = _MULTISPACE_RE.sub(" ", original.upper()).strip()
    if normalized_check in _INSURED_RESULT_OVERRIDES:
        return list(_INSURED_RESULT_OVERRIDES[normalized_check])

    if _MSIG_GROUP_YEAR_RE.search(original):
        t = _MULTISPACE_RE.sub(" ", original).strip()
        t = re.sub(r"\b(?:IDR|USD|SGD|EUR|CNY|GBP|JPY|AUD|HKD)\b", " ", t, flags=RE_I)
        t = _MULTISPACE_RE.sub(" ", t).strip()
        return [f"{t.upper()} PENYELESAIAN SUSPEND"]

    if _SUSPENSE_KEEP_AS_IS_RE.search(original):
        return [original]

    _PAREN_STASH.clear()
    t = _LANJUTAN_NOISE_RE.sub(" ", original)
    t = _strip_titles(t)
    t = _resolve_parens_split_or_unwrap(t)

    t = _protect_parens_content(t)
    raw_parts = _merge_bare_suffix_fragments(_ENTITY_SPLIT_RE.split(t))
    results = []

    for part in raw_parts:
        cleaned = _strip_legal_entity_with_punct(part)
        cleaned = _final_polish(cleaned)
        cleaned = _restore_protected_parens(cleaned)
        cleaned = _MULTISPACE_RE.sub(" ", cleaned).strip()

        if not cleaned or len(cleaned) < 2 or not _HAS_UPPER_RE.search(cleaned):
            continue
        if _BARE_CLIENT_FRAGMENT_RE.match(cleaned):
            continue
        if cleaned not in results:
            results.append(cleaned)

    if len(results) > MAX_BREAKDOWN_CODES:
        return [original]
    if not results:
        return [original.upper()]
    return results


# ==============================================================================
# 3. CLEANSING SLIP NO
# ==============================================================================
_MONTH_WORDS_ID = ("JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI",
                   "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER", "DESEBER")
_MONTH_WORDS_EN = ("JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY",
                   "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER")
_MONTH_ALONE_RE = re.compile(r"\b(?:" + "|".join(_MONTH_WORDS_ID + _MONTH_WORDS_EN) + r")\b", RE_I)
_MONTH_YEAR_GLUED_RE = re.compile(r"\b(?:" + "|".join(_MONTH_WORDS_ID + _MONTH_WORDS_EN) + r")(?:19|20)\d{2}\b", RE_I)
_OPEN_COVER_LABEL_RE = re.compile(r"\b(?:BORDERO|SUMMARY)\b", RE_I)
_CURRENCY_NOISE_RE = re.compile(r"\b(?:IDR|USD|SGD|EUR|CNY|GBP|JPY|AUD|HKD)\b", RE_I)
_FROM_UPTO_NOISE_RE = re.compile(r"\bFROM\s+\d{1,2}[./]\d{2,4}\s+UP\s+TO\s+\d{1,2}[./]\d{2,4}\b", RE_I)
_NUMERIC_MONTH_YEAR_RE = re.compile(r"\b(?:\d{1,2}\s*[./ ]\s*(?:19|20)\d{2}|(?:19|20)\d{2}\s*[./ ]?\s*\d{1,2})\b")
_END_REF_NOISE_RE = re.compile(r"\bEND\.\d+(?:-\d+)?\s*/\s*\d{1,2}[./]\d{2,4}\b", RE_I)
_IF_PC_NOISE_RE = re.compile(r"\bIF\s+P\s*&\s*C\b", RE_I)
_XGEN_NOISE_RE = re.compile(r"\bXGEN\b", RE_I)
_NEW_SYSTEM_NOISE_RE = re.compile(r"\bNEW\s+SYSTEM\b", RE_I)
_ADMIN_LABEL_NOISE_RE = re.compile(
    r"\bACC(?:OUNT|ONUT)\s+NUM[NB]ER\s*:\s*|\bACCOUNT\s+NO\s*:\s*|\bAC\s+NO\s*:\s*"
    r"|\bCONTRACT\s+NUMBER\s*:\s*|\bTREATY\s+NUMBER\s*:\s*|\bBORD(?:ERO)?\b",
    RE_I
)
_DOT_COMMA_TAIL_RE = re.compile(r"[.,]\s*$")
_NON_ALNUM_RE = re.compile(r"[^A-Za-z0-9]")
_DIGIT_SEARCH_RE = re.compile(r"\d")

_SLIP_KEEP_AS_IS_LITERALS = {
    "SOA MSIG RI 2022-08 - R I U-IF P&C UPDATED",
    "SOA MSIG RI 2022-08 - R I U-IF P&C UPDATED/ENDORSEMENT",
}


def _has_real_code(text: str) -> bool:
    return bool(re.search(r"[A-Za-z]{1,6}\d{5,}|\d{7,}", text))


def _looks_irregular(text: str) -> bool:
    return not _DIGIT_SEARCH_RE.search(text)


def _strip_slip_noise(t: str) -> str:
    t = _FROM_UPTO_NOISE_RE.sub(" ", t)
    t = _END_REF_NOISE_RE.sub(" ", t)
    t = _IF_PC_NOISE_RE.sub(" ", t)
    t = _XGEN_NOISE_RE.sub(" ", t)
    t = _NEW_SYSTEM_NOISE_RE.sub(" ", t)
    t = _ADMIN_LABEL_NOISE_RE.sub(" ", t)
    t = _CURRENCY_NOISE_RE.sub(" ", t)
    t = _MONTH_YEAR_GLUED_RE.sub(" ", t)
    t = _MONTH_ALONE_RE.sub(" ", t)
    t = _NUMERIC_MONTH_YEAR_RE.sub(" ", t)
    t = re.sub(r"\b(19|20)\d{2}\b", " ", t)
    t = _OPEN_COVER_LABEL_RE.sub(" ", t)
    t = re.sub(r"\([^()]*\)", " ", t)
    return t


def _strip_non_digit_chunks(t: str) -> str:
    chunks = re.split(r"([/+&,-])", t)
    kept = []
    for c in chunks:
        if c in ("/", "+", "&", ",", "-"):
            kept.append(c)
        elif c.strip() and not re.search(r"\d", c):
            kept.append(" ")
        else:
            kept.append(c)
    return "".join(kept)


def _breakdown_slip_codes(text: str):
    top_chunks = [c.strip() for c in re.split(r"[,+/&-]", text) if c.strip()]
    if len(top_chunks) < 2:
        return None

    codes = []
    anchor_prefix = ""
    anchor_digits = None

    for chunk in top_chunks:
        chunk = chunk.strip()
        if not chunk:
            continue
        m_full = re.fullmatch(r"([A-Za-z]{0,3})(\d{7,})", chunk)
        if m_full:
            anchor_prefix, anchor_digits = m_full.group(1), m_full.group(2)
            full_code = f"{anchor_prefix}{anchor_digits}"
            if full_code not in codes:
                codes.append(full_code)
            continue
        if re.fullmatch(r"\d{1,6}", chunk) and anchor_digits:
            width = len(chunk)
            reconstructed = (
                f"{anchor_prefix}{chunk}"
                if width >= len(anchor_digits)
                else f"{anchor_prefix}{anchor_digits[:-width]}{chunk}"
            )
            if reconstructed not in codes:
                codes.append(reconstructed)
            continue
        return None

    return codes if codes else None


def clean_split_code_impl(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return [], False
    original = text.strip()

    normalized_check = _MULTISPACE_RE.sub(" ", original.upper()).strip()
    if normalized_check in _SLIP_KEEP_AS_IS_LITERALS:
        return [original], True

    if _STANDALONE_KEEP_RE.match(original):
        return [original.upper()], True
        
    if _SUSPENSE_KEEP_AS_IS_RE.search(original):
        cleaned_suspense = _CURRENCY_NOISE_RE.sub(" ", original)
        cleaned_suspense = _MULTISPACE_RE.sub(" ", cleaned_suspense).strip()
        return [cleaned_suspense], True
        
    if _SEE_ATTACHMENT_RE.match(original):
        return [original], True

    if (_MONTH_ALONE_RE.search(original) or _OPEN_COVER_LABEL_RE.search(original) or _NUMERIC_MONTH_YEAR_RE.search(original)) and not _has_real_code(original):
        return [original], True

    t = re.sub(r"(?:^|(?<=[/,+&-]))\s*TBA\s*(?=[/,+&-]|$)", "", original, flags=RE_I)
    t = re.sub(r"(?:^|(?<=[/,+&-]))\s*(?:VAR(?:I(?:O(?:U(?:S)?)?)?)?)\s*(?=[/,+&-]|$)", "", t, flags=RE_I)
    t = re.sub(r"[/,+&-]{2,}", "+", t).strip()
    t = re.sub(r"^[/,+&\-\s]+|[/,+&\-\s]+$", "", t).strip()
    if not t:
        return [original], True

    t = _strip_slip_noise(t)
    t = re.sub(r"[/,+&-]{2,}", "+", t).strip()
    t = re.sub(r"^[/,+&\-\s]+|[/,+&\-\s]+$", "", t).strip()
    if not t:
        return [original], True

    t = _strip_non_digit_chunks(t)
    t = re.sub(r"[/,+&-]{2,}", "+", t).strip()
    t = re.sub(r"^[/,+&\-\s]+|[/,+&\-\s]+$", "", t).strip()
    if not t or _looks_irregular(t):
        return [t if t else original], True

    t = _DOT_COMMA_TAIL_RE.sub("", t).strip()
    if not t:
        return [original], True

    sd_tail_match = re.search(r"([A-Za-z]{0,3}\d{7,})\s*S\s*/\s*D\s*([A-Za-z]{0,3}\d{7,})\s*$", t, RE_I)
    if sd_tail_match:
        prefix_part = t[: sd_tail_match.start()].rstrip("/, ").strip()
        start_tok, end_tok = sd_tail_match.group(1), sd_tail_match.group(2)
        leading_codes = [c.strip() for c in re.split(r"[/,+]", prefix_part) if c.strip()] if prefix_part else []
        sd_codes = _expand_or_pair_sd_full_codes(start_tok, end_tok)
        combined = leading_codes + sd_codes
        return (combined, True) if len(combined) <= MAX_BREAKDOWN_CODES else ([original], True)

    codes = _breakdown_slip_codes(t)
    if codes is not None:
        return ([original], True) if len(codes) > MAX_BREAKDOWN_CODES else (codes, False)

    t_final = _DOT_COMMA_TAIL_RE.sub("", t).strip()
    if not t_final:
        return [original], True
    return ([t_final], True) if not _has_real_code(t_final) else ([t_final], False)


def clean_split_code(text: str) -> list:
    results, _ = clean_split_code_impl(text)
    return [r.upper() if isinstance(r, str) else r for r in results]


# ==============================================================================
# 4. CLEANSING POLIS & CERTIFICATE
# ==============================================================================
_SHORT_SUFFIX_RE = re.compile(r"^\d{1,4}$")
_MAIN_TOKEN_RE = re.compile(r"^(?:[A-Za-z]{1,3}\d{7,16}|\d{5,13})$")
_SPACE_SEPARATED_POLIS_RE = re.compile(r"^(?:\d{7,10}\s+){1,}\d{7,10}$")
# rev4: hardcase literal baru -- "DE3200002000100-42000100100-DE32000100100"
# HARUS dibreakdown jadi 3 polis persis sesuai instruksi user, TIDAK
# digeneralisasi ke rule "-" umum (polanya campur antara token penuh
# berprefix "DE" & token pure-digit dgn panjang berbeda-beda, terlalu
# spesifik/ambigu utk dijadikan rule umum baru).
_POLIS_HARDCASE_OVERRIDES = {
    "2000373-372": ["2000373", "2000372"],
    "DE3200002000100-42000100100-DE32000100100": [
        "DE3200002000100", "DE3242000100100", "DE32000100100"
    ],
}


def _cert_pad6(num_str: str) -> str:
    return num_str.zfill(6)


def _try_extend_cert_range_pair(pairs: list, idx: int, piece: str) -> bool:
    prev = pairs[idx][1]
    if not prev:
        return False
    prev_cert_raw = prev.split(" SD ")[-1].split(", ")[-1]
    if not (len(prev_cert_raw) == 6 and prev_cert_raw.isdigit()):
        return False
    prev_n = int(prev_cert_raw)
    padded_piece = _cert_pad6(piece)
    piece_n = int(padded_piece)
    gap = piece_n - prev_n
    if not (0 < gap <= 999):
        return False
    existing = prev.split(", ")
    pairs[idx][1] = ", ".join(existing + [padded_piece])
    return True


def _collapse_consecutive_cert_list(cert_str: str) -> str:
    if not cert_str or " SD " in cert_str or ", " not in cert_str:
        return cert_str
    items = cert_str.split(", ")
    if not all(len(v) == 6 and v.isdigit() for v in items):
        return cert_str
    nums = [int(v) for v in items]
    runs = [[nums[0]]]
    for n in nums[1:]:
        if n == runs[-1][-1] + 1:
            runs[-1].append(n)
        else:
            runs.append([n])
    parts = []
    for run in runs:
        if len(run) > 3:
            parts.append(f"{run[0]:06d} SD {run[-1]:06d}")
        else:
            parts.extend(f"{n:06d}" for n in run)
    return ", ".join(parts)


def _get_anchor_prefix_digits(token: str):
    m = re.match(r"^([A-Za-z]{0,3})(\d+)$", token)
    return (m.group(1), m.group(2)) if m else ("", token)


def _reconstruct_trailing_digits(anchor_prefix: str, anchor_digits: str, piece: str) -> str:
    width = len(piece)
    if width >= len(anchor_digits):
        return f"{anchor_prefix}{piece}"
    return f"{anchor_prefix}{anchor_digits[:-width]}{piece}"


def _expand_or_pair_sd_full_codes(start_token: str, end_token: str) -> list:
    start_prefix, start_digits = _get_anchor_prefix_digits(start_token)
    end_prefix, end_digits = _get_anchor_prefix_digits(end_token)

    if (start_prefix == end_prefix and len(start_digits) == len(end_digits)
            and start_digits.isdigit() and end_digits.isdigit()):
        start_n, end_n = int(start_digits), int(end_digits)
        if 0 <= end_n - start_n <= MAX_BREAKDOWN_CODES:
            width = len(start_digits)
            return [f"{start_prefix}{n:0{width}d}" for n in range(start_n, end_n + 1)]

    return [f"{start_token.upper()} SD {end_token.upper()}"]


def process_policy_certificate(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return [], []
    original = text.strip()

    normalized_check_polis = _MULTISPACE_RE.sub(" ", original.upper()).strip()
    if normalized_check_polis in _POLIS_HARDCASE_OVERRIDES:
        codes = list(_POLIS_HARDCASE_OVERRIDES[normalized_check_polis])
        return [""] * len(codes), codes

    if _SUSPENSE_KEEP_AS_IS_RE.search(original):
        return [""], [original]
    if (_MONTH_ALONE_RE.search(original) or _OPEN_COVER_LABEL_RE.search(original)) and not _has_real_code(original):
        return [""], [original]

    if _SPACE_SEPARATED_POLIS_RE.match(original.strip()):
        tokens = original.strip().split()
        return [""] * len(tokens), tokens

    t = re.sub(r"S\s*/\s*D", "SD", original, flags=RE_I)
    t = re.sub(r"(?<=\d)(SD)", r" \1", t, flags=RE_I)
    t = re.sub(r"(SD)(?=\d)", r"\1 ", t, flags=RE_I)

    if _STANDALONE_KEEP_RE.match(t):
        return [""], [t.upper()]

    t = re.sub(
        r"(?:^|(?<=[+,/:\s-]))\s*(?:VAR(?:I(?:O(?:U(?:S)?)?)?)?|VAIOUS|TBA|TPI|P\d{0,2})\s*(?=[+,/:\s-]|$)",
        "", t, flags=RE_I
    )
    t = re.sub(r"[+,/:]{2,}", "+", t).strip()
    t = re.sub(r"^[+,/:\-\s]+|[+,/:\-\s]+$", "", t).strip()
    if not t:
        return [""], [original]

    if _MONTH_ALONE_RE.search(t) and _has_real_code(t):
        t = _MONTH_YEAR_GLUED_RE.sub(" ", t)
        t = _MONTH_ALONE_RE.sub(" ", t)
        t = _NUMERIC_MONTH_YEAR_RE.sub(" ", t)
        t = re.sub(r"\b(19|20)\d{2}\b", " ", t)
        t = _CURRENCY_NOISE_RE.sub(" ", t)
        t = re.sub(r"[+,/:-]{2,}", "+", t).strip()
        t = re.sub(r"^[+,/:\-\s]+|[+,/:\-\s]+$", "", t).strip()
        if not t:
            return [""], [original]

    t = re.sub(r"[*#]+", " ", t)
    t = re.sub(r"[()]", " ", t)
    t = re.sub(r"[+,/:-]{2,}", "+", t).strip()
    t = re.sub(r"^[+,/:\-\s]+|[+,/:\-\s]+$", "", t).strip()
    if not t:
        return [""], [original]

    t_noise_stripped = t
    t = re.sub(r"[\u2012\u2013\u2014\u2015]", "-", t)

    if " " in t and not re.search(r"[,+/:&-]", t):
        space_tokens = t.split()
        if all(_MAIN_TOKEN_RE.match(tok) for tok in space_tokens) and len(space_tokens) > 1:
            t = ",".join(space_tokens)

    top_chunks = [c.strip() for c in re.split(r"[,+/:&]", t) if c.strip()]
    if not top_chunks:
        return [""], [t_noise_stripped]

    pairs: list[list] = []
    fallback = False
    active_idx = None

    for chunk in top_chunks:
        pieces = [p.strip() for p in re.split(r"(\bSD\b|-)", chunk.strip().strip("/").strip()) if p.strip()]
        if not pieces:
            continue

        i = 0
        local_main_prefix, local_main_digits = "", None
        pending_op = None

        while i < len(pieces):
            piece = pieces[i]
            if piece in ("-", "SD"):
                pending_op = piece
                i += 1
                continue

            if pending_op is None and _MAIN_TOKEN_RE.match(piece):
                local_main_prefix, local_main_digits = _get_anchor_prefix_digits(piece)
                if local_main_prefix == "" and len(local_main_digits) < 8:
                    active_idx = None
                else:
                    existing_idx = next((idx for idx, p in enumerate(pairs) if p[0] == piece), None)
                    if existing_idx is not None:
                        active_idx = existing_idx
                    else:
                        pairs.append([piece, ""])
                        active_idx = len(pairs) - 1
                pending_op = None
                i += 1
                continue

            if piece.isdigit():
                if local_main_digits is not None:
                    anchor_digits, anchor_prefix = local_main_digits, local_main_prefix
                elif active_idx is not None and pairs:
                    anchor_prefix, anchor_digits = _get_anchor_prefix_digits(pairs[active_idx][0])
                else:
                    anchor_digits, anchor_prefix = None, ""

                if pending_op == "SD":
                    if _MAIN_TOKEN_RE.match(piece):
                        start_token = (
                            f"{local_main_prefix}{local_main_digits}" if local_main_digits is not None
                            else pairs[active_idx][0] if active_idx is not None and pairs
                            else piece
                        )
                        expanded = _expand_or_pair_sd_full_codes(start_token, piece)
                        if active_idx is not None and pairs and pairs[active_idx][0] == start_token:
                            pairs[active_idx][0] = expanded[0]
                            for extra in expanded[1:]:
                                pairs.append([extra, ""])
                            active_idx = len(pairs) - 1
                        else:
                            for extra in expanded:
                                pairs.append([extra, ""])
                            active_idx = len(pairs) - 1
                        pending_op = None
                        i += 1
                        continue

                    if anchor_digits:
                        reconstructed = _reconstruct_trailing_digits(anchor_prefix, anchor_digits, piece)
                        existing_idx = next((idx for idx, p in enumerate(pairs) if p[0] == reconstructed), None)
                        if existing_idx is None:
                            pairs.append([reconstructed, ""])
                            active_idx = len(pairs) - 1
                        else:
                            active_idx = existing_idx
                    pending_op = None
                    i += 1
                    continue

                if anchor_digits and (_SHORT_SUFFIX_RE.match(piece) or (re.fullmatch(r"\d{1,11}", piece) and len(piece) < len(anchor_digits))):
                    if pending_op == "-" and active_idx is not None and _try_extend_cert_range_pair(pairs, active_idx, piece):
                        pending_op = None
                        i += 1
                        continue

                    reconstructed = _reconstruct_trailing_digits(anchor_prefix, anchor_digits, piece)
                    existing_idx = next((idx for idx, p in enumerate(pairs) if p[0] == reconstructed), None)
                    if existing_idx is None:
                        pairs.append([reconstructed, ""])
                        active_idx = len(pairs) - 1
                    else:
                        active_idx = existing_idx
                    pending_op = None
                    i += 1
                    continue

                if _MAIN_TOKEN_RE.match(piece):
                    existing_idx = next((idx for idx, p in enumerate(pairs) if p[0] == piece), None)
                    if existing_idx is None:
                        pairs.append([piece, ""])
                        active_idx = len(pairs) - 1
                    else:
                        active_idx = existing_idx
                    pending_op = None
                    i += 1
                    continue

                fallback = True
                break

            fallback = True
            break

        if local_main_digits is not None and active_idx is None:
            full_token = f"{local_main_prefix}{local_main_digits}"
            if not any(p[0] == full_token for p in pairs):
                pairs.append([full_token, ""])

        if fallback:
            break

    polis_codes = [p[0] for p in pairs]
    if fallback or not polis_codes or len(polis_codes) > MAX_BREAKDOWN_CODES:
        return [""], [t_noise_stripped]

    cert_list = [_collapse_consecutive_cert_list(p[1]) for p in pairs]
    return cert_list, polis_codes


# ==============================================================================
# 5. HELPER COMPACT & I/O
# ==============================================================================
def _compact_value(value: str) -> str:
    return _NON_ALNUM_RE.sub("", value) if isinstance(value, str) else value


def _compact_split_code_results(results: list, is_kept_as_is: bool) -> list:
    if is_kept_as_is:
        return results
    return [_compact_value(v) for v in results]


def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)


def filter_cedant(df_raw: pd.DataFrame, keyword: str, cedant_column: str) -> pd.DataFrame:
    if cedant_column not in df_raw.columns:
        raise ValueError(f"Kolom '{cedant_column}' tidak ditemukan. Tersedia: {list(df_raw.columns)}")
    mask = df_raw[cedant_column].astype(str).str.strip() == keyword.strip()
    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant '{keyword}': {len(df_filtered)} baris ditemukan.")
    return df_filtered


# ==============================================================================
# 6. BUILD OUTPUT
# ==============================================================================
def _breakdown_all_rows(df: pd.DataFrame, col_insured: str, col_polis: str, col_slip: str):
    ins_arr = df[col_insured].to_numpy()
    pol_arr = df[col_polis].to_numpy()
    slp_arr = df[col_slip].to_numpy()

    insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val in zip(ins_arr, pol_arr, slp_arr):
        ins_cln = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []

        if pd.notna(pol_val) and str(pol_val).strip():
            cert_cln, pol_cln = process_policy_certificate(str(pol_val))
            pol_cln = [v.upper() if isinstance(v, str) else v for v in pol_cln]
        else:
            pol_cln, cert_cln = [], []

        if pd.notna(slp_val) and str(slp_val).strip():
            slp_results, slp_kept = clean_split_code_impl(str(slp_val))
            slp_cln = _compact_split_code_results(
                [r.upper() if isinstance(r, str) else r for r in slp_results], slp_kept
            )
        else:
            slp_cln = []

        insured_cln_all.append(ins_cln)
        polis_cln_all.append(pol_cln)
        slip_cln_all.append(slp_cln)
        cert_cln_all.append(cert_cln)

        max_ins = max(max_ins, len(ins_cln))
        max_pol = max(max_pol, len(pol_cln))
        max_slp = max(max_slp, len(slp_cln))

    return insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all, max_ins, max_pol, max_slp


def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_insured, col_polis, col_slip = INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN
    for col, label in ((col_insured, "INSURED"), (col_polis, "POLIS"), (col_slip, "SLIP")):
        if col not in df_filtered.columns:
            raise ValueError(f"Kolom {label}='{col}' tidak ditemukan.")

    ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp = _breakdown_all_rows(
        df_filtered, col_insured, col_polis, col_slip
    )

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)

    if BROKER_NAME_COLUMN in df_filtered.columns and CEDANT_COLUMN in df_filtered.columns:
        is_direct = df_filtered[BROKER_NAME_COLUMN].astype(str).str.strip().str.upper() == DIRECT_MARKER
        bp_series = df_filtered[CEDANT_COLUMN].where(is_direct, df_filtered[BROKER_NAME_COLUMN])
        business_partner_values = bp_series.astype(str).str.replace(r"^(\s*PT)\.\s*", r"\1 ", regex=True, flags=RE_I).values
    else:
        business_partner_values = None

    output_cols_data = {}
    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == BROKER_NAME_COLUMN and business_partner_values is not None:
            output_cols_data["BUSINESS_PARTNER"] = business_partner_values

        if col == col_insured:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in ins_all]

        if col == col_polis:
            for i in range(1, max_pol + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in pol_all]
                output_cols_data[f"CERTIFICATE_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_all]

        if col == col_slip:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slp_all]

    df_master = pd.DataFrame(output_cols_data)
    drop_cols = [
        c for c in df_master.columns
        if ("_CLN_" in c or c.startswith("CERTIFICATE_"))
        and c != "CERTIFICATE_1"
        and df_master[c].astype(str).str.strip().eq("").all()
    ]
    if "CERTIFICATE_1" not in df_master.columns:
        df_master["CERTIFICATE_1"] = ""
    return df_master.drop(columns=drop_cols)


# ==============================================================================
# 7. EXPORT TO EXCEL
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook = load_workbook(output_path)
    worksheet = workbook.active
    text_col_indices = [
        idx for idx, col in enumerate(df.columns, start=1)
        if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)
    ]

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)
    workbook.save(output_path)


# ==============================================================================
# 8. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70 + "\n CLEANSING DATA 1 - FACULTATIVE | CEDANT: MSIG INDONESIA\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
        if candidates:
            input_path = candidates[0]
            print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")
        else:
            raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    print(f"[INFO] Total baris source: {len(df_raw)}")

    df_msig = filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)
    if df_msig.empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter cedant. Proses dihentikan.")
        return

    df_hasil = build_output(df_msig)
    save_with_text_format(df_hasil, OUTPUT_FILE)

    print("=" * 70 + f"\n[SUCCESS] Selesai. Total baris output: {len(df_hasil)}\n[SUCCESS] File hasil: '{OUTPUT_FILE}'\n" + "=" * 70)


if __name__ == "__main__":
    main()