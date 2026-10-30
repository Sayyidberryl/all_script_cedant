import os
import re
import sys
import time
import functools

sys.modules['numexpr'] = None
sys.modules['bottleneck'] = None

# pyrefly: ignore [missing-import]
import numpy as np
import pandas as pd

try:
    # pyrefly: ignore [missing-import]
    from dotenv import load_dotenv
    # pyrefly: ignore [missing-import]
    from sqlalchemy import create_engine
except ImportError:
    load_dotenv = None
    create_engine = None


# --- Konfigurasi (MSIG) ---

_BASE_DIR  = os.path.dirname(os.path.abspath(__file__))   # folder root MSIG/
_CLEAN_DIR = os.path.join(_BASE_DIR, "data_clean")        # hasil cleansing
_RAW_DIR   = os.path.join(_BASE_DIR, "data_raw")          # data mentah


def _find_file_in(directory: str, basenames: list) -> str:
    """Cari file pertama yang ada dari daftar kandidat nama dalam direktori tertentu."""
    for name in basenames:
        p = os.path.join(directory, name)
        if os.path.exists(p):
            return p
    return os.path.join(directory, basenames[0])


# Path file untuk cedant MSIG
SUSPEND_FILE = _find_file_in(_CLEAN_DIR, ["msig_output_suspense.xlsx", "msig_output_suspense_V2.xlsx"])
OSBAL_FILE   = _find_file_in(_CLEAN_DIR, ["msig_output_osbal.xlsx", "msig_output_osbal_V2.xlsx"])
FACUL_FILE   = _find_file_in(_CLEAN_DIR, ["msig_output_facul.xlsx", "msig_output_facul_V2.xlsx"])
SLIPDB_FILE  = _find_file_in(_RAW_DIR, [
    "5. Database Rekap RI Slip.xlsx",
    "5.Database Rekap RI Slip.xlsx",
    "Database Rekap RI Slip.xlsx",
])
BORDERO_FILE = _find_file_in(_RAW_DIR, [
    "(Februari 2022 - Juni 2026) Bordore Marine Cargo MSIG.xlsx",
    "Bordore Marine Cargo MSIG.xlsx",
    "Bordero Marine Cargo MSIG.xlsx",
])
OUTPUT_FILE  = os.path.join(_BASE_DIR, "matching_msig_v1.xlsx")

# Kolom di tabel Bordero MSIG (Marine Cargo)
BORDERO_FAC_COL     = "FAC CODE"
BORDERO_POLIS_COL   = "POLIS"
BORDERO_SLIP_COL    = "SLIP"
BORDERO_CERT_COL    = "CERTIFICATE"
BORDERO_INSURED_COL = "INSURED"
BORDERO_CURR_COL    = "CURR"
BORDERO_BULAN_COL   = "BULAN"
BORDERO_TAHUN_COL   = "TAHUN"
BORDERO_NET_COL     = "Net"

# Toleransi matching AMOUNT ORI vs NET bordero
NET_TOLERANCE_PCT = 0.00

# Kontrol FASE 2 (SLIP): default False untuk MSIG karena slip dummy master
ENABLE_SLIP_MATCHING = False

# Mapping nama bulan Indonesia -> angka bulan (1-12)
_BULAN_MAP = {
    "JANUARI": 1, "FEBRUARI": 2, "MARET": 3, "APRIL": 4,
    "MEI": 5, "JUNI": 6, "JULI": 7, "AGUSTUS": 8,
    "SEPTEMBER": 9, "OKTOBER": 10, "NOVEMBER": 11, "DESEMBER": 12,
}

# Kolom FAC CODE di masing-masing tabel referensi
OSBAL_FACODE_COL  = "CCOS_REF_CODE"
FACUL_FACODE_COL  = "FAC_CODE"
SLIPDB_FACODE_COL = "FAC. CODE"

# Kolom polis/slip standar
CLSDT_POLIS_COL = "CLSDT_POLICY_NO"
CLSDT_SLIP_COL  = "CLSDT_SLIP_NO"
FAC_POLIS_COL   = "FAC_POLICY_NO"
FAC_SLIP_COL    = "FAC_SLIP"

# Nama kolom currency/periode
SUSPEND_CURR_COL = "CURR ORI"
OSBAL_CURR_COL   = "CCOS_CURR"
SUSPEND_DATE_COL = "RECEIPT DATE"
OSBAL_DATE_COL   = "FAC_COM_DATE"

# Marker baris administratif yang dikecualikan dari matching
_EXCLUDED_REF_MARKERS = [
    "PENYELESAIAN SUSPEND",
    "HUTANG PIUTANG",
    "HUTANG",
    "UTANG",
    "DATA SUSPENSE",
    "SUSPENSE",
    "LINESLIP",
    "LINE SLIP",
]

# Marker LINESLIP - pakai RECEIPT DATE - 1 bulan sebagai tanggal efektif
_LINESLIP_MARKERS = ("LINESLIP", "LINE SLIP")

_MIN_TOKEN_LEN = 3
_TOKEN_RE      = re.compile(r"[^A-Z0-9]+")
_DELIMITER_RE  = re.compile(r"[,;|+\s]+")

# Token generik yang dikecualikan saat menghitung overlap INSURED (kurangi false positive)
_INSURED_STOPWORDS: frozenset = frozenset({
    "PT", "CV", "TBK", "UD", "PD", "NV", "LLC", "LTD", "INC",
    "GROUP", "GROUPS",
    "INDONESIA", "INDO",
    "INDUSTRY", "INDUSTRIES", "INDUSTRI",
    "INTERNATIONAL", "INTERNASIONAL",
    "COMPANY", "COMPANIES", "CORP", "CORPORATION",
    "ENTERPRISE", "ENTERPRISES",
    "TRADING", "TRADE",
    "UTAMA", "MANDIRI", "MAKMUR", "SEJAHTERA",
    "ASURANSI", "INSURANCE",
    "SERVICE", "SERVICES",
    "OPERATIONS", "OPERATION", "OPERASIONAL",
})

_INSURED_LIKE_THRESHOLD: float = 0.65
_POLIS_LIKE_THRESHOLD: float   = 0.80

# Kolom join saat akumulasi suspend
_SUSPEND_JOIN_COLS = [
    "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "RECEIPT DATE",
    "CEDANT NAME", "CEDANT SHRT NAME",
    "INSURED", "INSURED_CLN_1", "INSURED_CLN_2",
    "CURR ORI", "CURR PAY", "AMOUNT PAY",
    "POLIS", "POLIS_CLN_1",
    "SLIP NO", "SLIP_NO_CLN_1",
    "DESC 1", "DESC 2", "DESC 3", "DESC 4",
    "STATUS", "REC_TYPE",
]
_SUSPEND_SUM_COLS = ["AMOUNT ORI"]

# Kolom output final MSIG
FINAL_COLUMNS = [
    "CCOS_DOC_NO", "CCOS_REF_CODE",
    "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "RECEIPT DATE",
    "CEDANT NAME", "CEDANT SHRT NAME",
    "INSURED_ORI", "INSURED_1", "INSURED_2",
    "CURR ORI", "AMOUNT ORI", "AMOUNT_ORI_MIN1",
    "CURR PAY", "AMOUNT PAY",
    "CCOS_OR_BAL", "CCOS_BAL_DUE", "DIFERENCE",
    "FLAG_PROD",
    "POLIS_ORI",
    "POLICY_CLEAN_1", "SERTIF_CLEAN_1",
    "POLICY_CLEAN_2", "SERTIF_CLEAN_2",
    "POLICY_CLEAN_3", "SERTIF_CLEAN_3",
    "POLICY_CLEAN_4", "POLICY_CLEAN_5",
    "SLIP_NO_ORI", "SLIP_NO_CLN",
    "DESC 1", "DESC 2", "DESC 3", "DESC 4",
    "STATUS", "REC_TYPE",
    "SKENARIO",
]


# --- String Utilities ---

@functools.lru_cache(maxsize=131072)
def _norm_cached(s: str) -> str:
    text = s.strip().upper().replace("S/D", "SD")
    return re.sub(r"\s+", " ", text) if "  " in text else text


def _normalize(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return _norm_cached(str(value))


def _format_sertif(val) -> str:
    if pd.isna(val) or str(val).strip() == "":
        return ""
    if isinstance(val, (int, float)):
        return str(int(val)).zfill(6)
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s.zfill(6) if s.isdigit() else s


def _tokenize(text: str) -> frozenset:
    cleaned = _TOKEN_RE.sub(" ", _normalize(text))
    return frozenset(t for t in cleaned.split() if len(t) >= _MIN_TOKEN_LEN)


def _insured_overlap_coeff(tokens_a: frozenset, tokens_b: frozenset, stopwords: frozenset) -> float:
    sig_a = tokens_a - stopwords
    sig_b = tokens_b - stopwords
    if not sig_a or not sig_b:
        return 0.0
    denom = min(len(sig_a), len(sig_b))
    return len(sig_a & sig_b) / denom if denom > 0 else 0.0


def _standardize_label(label: str) -> str:
    lbl = label.upper().strip()
    if "POLIS" in lbl and "SLIP" in lbl:
        return "Polis + Slip"
    if "POLIS" in lbl:
        return "Polis only"
    if "SLIP" in lbl:
        return "Slip only"
    if "INSURED" in lbl:
        return "Insured only"
    return label


def _clean_scenario_label(scenario: str) -> str:
    s = scenario.strip()
    s = re.sub(r"\s+only\s*\+", " +", s)
    s = re.sub(r"\s+only\s*$", "", s)
    return s.strip()


def _map_scenario(label: str) -> str:
    if "SLIP" in label and "POLIS" in label:
        return "Polis + Slip"
    if "SLIP" in label:
        return "Slip only"
    if "POLIS" in label:
        return "Polis only"
    if "INSURED" in label:
        return "Insured only"
    return label


# --- Row Dict Helpers ---

def _is_excluded_ref_row(row: dict, polis_cols: list, slip_cols: list,
                         insured_cols: list = None) -> bool:
    """Cek apakah baris referensi mengandung marker administratif pada kolom kunci."""
    check_cols = (
        list(polis_cols) + ["polis_ori", "POLIS"]
        + list(slip_cols) + ["slip_ori", "SLIP"]
        + (list(insured_cols) + ["insured_ori", "INSURED", "FAC_INSURED"] if insured_cols else
           ["insured_ori", "INSURED", "FAC_INSURED"])
    )
    for col in check_cols:
        val = _normalize(row.get(col, ""))
        if val and any(m in val for m in _EXCLUDED_REF_MARKERS):
            return True
    return False


def _excluded_facodes(rows: list, facode_col: str, polis_cols: list, slip_cols: list,
                      insured_cols: list = None) -> set:
    """Kumpulkan fac code dari baris referensi dengan marker administratif."""
    result = set()
    if not rows or not facode_col:
        return result
    for r in rows:
        if _is_excluded_ref_row(r, polis_cols, slip_cols, insured_cols):
            fc = _normalize(r.get(facode_col, ""))
            if fc:
                result.add(fc)
    return result


def _is_lineslip_row(sus_row: dict, insured_cols: list) -> bool:
    for c in list(insured_cols) + ["INSURED", "insured_ori"]:
        val = _normalize(sus_row.get(c, ""))
        if val and any(m in val for m in _LINESLIP_MARKERS):
            return True
    return False


def _collect_clean_values(row: dict, cols: list) -> list:
    seen, vals = set(), []
    for c in cols:
        v = _normalize(row.get(c, ""))
        if v and v not in seen:
            seen.add(v)
            vals.append(v)
    return vals


def _collect_all_values(row: dict, clean_cols: list, ori_col: str) -> list:
    seen, vals = set(), []
    for c in clean_cols:
        v = _normalize(row.get(c, ""))
        if v and v not in seen:
            seen.add(v)
            vals.append(v)
    v_ori = _normalize(row.get(ori_col, ""))
    if v_ori and v_ori not in seen:
        vals.append(v_ori)
    return vals


def _get_effective_sus_cols(sus_row: dict, polis_cols: list, slip_cols: list) -> tuple:
    c_polis = _normalize(sus_row.get(CLSDT_POLIS_COL, ""))
    c_slip  = _normalize(sus_row.get(CLSDT_SLIP_COL, ""))

    eff_polis  = ([CLSDT_POLIS_COL] if c_polis else []) + polis_cols
    eff_slip   = ([CLSDT_SLIP_COL]  if c_slip  else []) + slip_cols

    # Ensure unique
    eff_polis = list(dict.fromkeys(eff_polis))
    eff_slip  = list(dict.fromkeys(eff_slip))

    like_polis = ([FAC_POLIS_COL] + eff_polis) if _normalize(sus_row.get(FAC_POLIS_COL, "")) else eff_polis
    like_slip  = ([FAC_SLIP_COL]  + eff_slip)  if _normalize(sus_row.get(FAC_SLIP_COL, ""))  else eff_slip

    like_polis = list(dict.fromkeys(like_polis))
    like_slip  = list(dict.fromkeys(like_slip))

    return eff_polis, eff_slip, like_polis, like_slip


def _get_clean_cols(header: list, prefix: str) -> list:
    prefix_upper = prefix.upper()
    cols = [c for c in header if isinstance(c, str) and c.upper().startswith(prefix_upper)]
    return sorted(cols, key=lambda x: [int(t) if t.isdigit() else t for t in re.split(r'(\d+)', str(x))])


def _slip_in_range(slip_val: str, osbal_row: dict) -> bool:
    if not slip_val:
        return False
    digits = re.findall(r"\d+", slip_val)
    if not digits:
        return False
    try:
        sus_num = int(digits[-1])
    except ValueError:
        return False

    ref_slip = _normalize(osbal_row.get("FAC_SLIP", ""))
    range_match = re.search(r"(\d+)\s*(?:-|SD|S/D|TO)\s*(\d+)", ref_slip)
    if range_match:
        try:
            start_num = int(range_match.group(1))
            end_num   = int(range_match.group(2))
            if start_num <= sus_num <= end_num:
                return True
        except ValueError:
            pass
    return False


def _ref_has_value(row: dict, cols: list, ori_col: str, candidate_vals: list) -> bool:
    cand_set = set(candidate_vals)
    for c in cols:
        v = _normalize(row.get(c, ""))
        if v and v in cand_set:
            return True
    v_ori = _normalize(row.get(ori_col, ""))
    if v_ori:
        for cv in candidate_vals:
            if cv and cv in v_ori:
                return True
    return False


# --- Index Builders ---

def _build_exact_index(rows: list, cols: list, excluded: set = None) -> dict:
    index: dict = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        for col in cols:
            val = _normalize(row.get(col, ""))
            if val:
                index.setdefault(val, []).append(i)
    return index


_RE_POLIS_SEP = re.compile(r"[&,/]|\s+(?=\d)")

def _split_polis_compound(val: str) -> list:
    """Split nilai polis compound menjadi token individual."""
    if not val:
        return []
    parts = [p.strip() for p in _RE_POLIS_SEP.split(val)]
    result = []
    for p in parts:
        p = p.strip()
        if p:
            result.append(p)
    # Tambahkan juga nilai asli (compound) agar reverse lookup tetap bisa match
    if len(result) > 1:
        result.append(val)  # val = compound asli
    return result if result else [val]


def _build_exact_index_polis(rows: list, cols: list, excluded: set = None) -> dict:
    """Index polis dengan compound splitting."""
    index: dict = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        for col in cols:
            val = _normalize(row.get(col, ""))
            if not val:
                continue
            for token in _split_polis_compound(val):
                if token:
                    index.setdefault(token, []).append(i)
    return index


def _build_facode_index(rows: list, col: str, excluded: set = None) -> dict:
    index: dict = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        val = _normalize(row.get(col, ""))
        if val:
            index.setdefault(val, []).append(i)
    return index


def _build_token_index(rows: list, clean_cols: list, ori_col: str, excluded: set = None) -> tuple:
    token_index: dict = {}
    token_cache: dict = {}
    for i, row in enumerate(rows):
        if excluded and i in excluded:
            continue
        tokens: frozenset = frozenset()
        for c in clean_cols:
            t = _normalize(row.get(c, ""))
            if t:
                tokens = tokens | _tokenize(t)
        token_cache[i] = tokens
        for tok in tokens:
            token_index.setdefault(tok, set()).add(i)
    return token_index, token_cache


def _build_lookup(rows: list, polis_cols: list, slip_cols: list,
                  insured_cols: list, facode_col: str = None, label: str = "") -> tuple:
    """Bangun 10-tuple inverted index untuk tabel referensi."""
    if not rows:
        if label:
            print(f"  {label:<6} done (0.0s)", flush=True)
        return ({}, {}, {}, {}, {}, {}, ({}, {}), ({}, {}), ({}, {}), {})

    t0 = time.perf_counter()
    excluded = {i for i, r in enumerate(rows) if _is_excluded_ref_row(r, polis_cols, slip_cols, insured_cols)}

    facode_index = _build_facode_index(rows, facode_col, excluded=excluded) if facode_col else {}
    lookup = (
        _build_exact_index(rows, slip_cols,    excluded=excluded),
        _build_exact_index_polis(rows, polis_cols, excluded=excluded),
        _build_exact_index(rows, insured_cols, excluded=excluded),
        {}, {}, {},
        _build_token_index(rows, slip_cols,    "slip_ori",    excluded=excluded),
        _build_token_index(rows, polis_cols,   "polis_ori",   excluded=excluded),
        _build_token_index(rows, insured_cols, "insured_ori", excluded=excluded),
        facode_index,
    )

    elapsed = time.perf_counter() - t0
    if label:
        exc_info = f" - {len(excluded):,} baris dikecualikan" if excluded else ""
        print(f"  {label:<6} done ({elapsed:.1f}s){exc_info}", flush=True)
    return lookup


# --- Match Engine ---

def _exact_match(query_values: list, index: dict) -> set:
    return {i for v in query_values if v and v in index for i in index[v]}


def _like_match(
    query_values: list,
    token_index: dict,
    rows: list,
    ref_cols: list,
    ori_col: str = None,
    stopwords: frozenset = frozenset(),
    overlap_threshold: float = _INSURED_LIKE_THRESHOLD,
    match_log: list = None,
) -> list:
    valid_queries = [q for q in query_values if q]
    if not valid_queries:
        return []

    all_query_tokens: frozenset = frozenset()
    for q in valid_queries:
        all_query_tokens = all_query_tokens | _tokenize(q)

    sig_query_tokens = all_query_tokens - stopwords
    if not sig_query_tokens:
        return []

    candidate_indices: set = set()
    for tok in sig_query_tokens:
        if tok in token_index:
            candidate_indices.update(token_index[tok])

    strict_matched: dict = {}
    relaxed_matched: dict = {}
    valid_queries_ns = [q.replace(" ", "") for q in valid_queries if len(q.replace(" ", "")) >= 5]
    is_insured_match = bool(stopwords)

    for i in candidate_indices:
        r = rows[i]
        ref_values = [_normalize(r.get(c, "")) for c in ref_cols]
        if ori_col and ori_col in r:
            ref_values.append(_normalize(r[ori_col]))

        # Kriteria 1: SUBSTR (dua arah: q in rv atau rv in q)
        substr_relaxed = False
        substr_strict = False
        has_fix_a_mismatch = False

        for rv in ref_values:
            if not rv: continue
            for q in valid_queries:
                if q in rv:
                    substr_relaxed = True
                    substr_strict = True
                    break
                if len(rv) >= 6 and rv in q:
                    substr_relaxed = True
                    if is_insured_match:
                        # Fix A: cek sisa string setelah rv dibuang
                        remainder = q.replace(rv, '', 1).strip()
                        rem_tokens = [t for t in _TOKEN_RE.split(remainder) if len(t) >= _MIN_TOKEN_LEN and t not in stopwords]
                        short_suffix = len(remainder) <= 2 and remainder.isalpha()
                        if not rem_tokens or short_suffix:
                            substr_strict = True
                        else:
                            has_fix_a_mismatch = True
                    else:
                        substr_strict = True
                    break
            if substr_strict: break

            if valid_queries_ns:
                rv_ns = rv.replace(" ", "")
                for q_ns in valid_queries_ns:
                    if q_ns in rv_ns:
                        substr_relaxed = True
                        substr_strict = True
                        break
                    if len(rv_ns) >= 6 and rv_ns in q_ns:
                        substr_relaxed = True
                        if is_insured_match:
                            remainder = q_ns.replace(rv_ns, '', 1).strip()
                            rem_tokens = [t for t in _TOKEN_RE.split(remainder) if len(t) >= _MIN_TOKEN_LEN and t not in stopwords]
                            short_suffix = len(remainder) <= 2 and remainder.isalpha()
                            if not rem_tokens or short_suffix:
                                substr_strict = True
                            else:
                                has_fix_a_mismatch = True
                        else:
                            substr_strict = True
                        break
            if substr_strict: break

        if substr_strict:
            strict_matched[i] = "SUBSTR"
        if substr_relaxed:
            relaxed_matched[i] = "SUBSTR"

        # Kriteria 2: Evaluasi per-entitas token overlap
        if not has_fix_a_mismatch:
            token_relaxed = False
            token_strict = False
            for rv in ref_values:
                if not rv: continue
                rv_tokens = _tokenize(rv)
                sig_rv_tokens = rv_tokens - stopwords
                if not sig_rv_tokens:
                    continue

                # SUBSET 1: query adalah subset dari entitas referensi
                if sig_query_tokens.issubset(sig_rv_tokens):
                    if len(sig_query_tokens) >= 2 or (len(sig_query_tokens) == 1 and len(next(iter(sig_query_tokens))) >= 5):
                        token_relaxed = True
                        token_strict = True
                        break

                # SUBSET 2: entitas referensi adalah subset dari query
                if sig_rv_tokens.issubset(sig_query_tokens):
                    if len(sig_rv_tokens) >= 2 or (len(sig_rv_tokens) == 1 and len(next(iter(sig_rv_tokens))) >= 5):
                        token_relaxed = True
                        # Fix B: ref 1 token vs multi query ditolak
                        if is_insured_match and (len(sig_rv_tokens) == 1 and len(sig_query_tokens) > 1):
                            pass
                        else:
                            token_strict = True
                            break

                # OVERLAP per-entitas
                inter = sig_query_tokens & sig_rv_tokens
                denom = min(len(sig_query_tokens), len(sig_rv_tokens))
                if denom > 0 and (len(inter) / denom) >= overlap_threshold:
                    token_relaxed = True
                    # Fix B: ref 1 token vs multi query ditolak
                    if is_insured_match and (len(sig_rv_tokens) == 1 and len(sig_query_tokens) > 1):
                        pass
                    else:
                        token_strict = True
                        break

            if token_strict and i not in strict_matched:
                strict_matched[i] = "TOKEN"
            if token_relaxed and i not in relaxed_matched:
                relaxed_matched[i] = "TOKEN"

            # Kriteria 3: Fallback terhadap gabungan seluruh ref_tokens
            if i not in strict_matched:
                ref_tokens: frozenset = frozenset()
                for rv in ref_values:
                    ref_tokens = ref_tokens | _tokenize(rv)

                if ref_tokens:
                    sig_ref_tokens = ref_tokens - stopwords
                    if sig_query_tokens and sig_query_tokens.issubset(ref_tokens):
                        if i not in relaxed_matched: relaxed_matched[i] = "SUBSET"
                        strict_matched[i] = "SUBSET"
                    elif sig_query_tokens and sig_ref_tokens:
                        intersection = len(sig_query_tokens & sig_ref_tokens)
                        denom = min(len(sig_query_tokens), len(sig_ref_tokens))
                        if denom > 0 and (intersection / denom) >= overlap_threshold:
                            if i not in relaxed_matched: relaxed_matched[i] = "OVERLAP"
                            if is_insured_match and (len(sig_ref_tokens) == 1 and len(sig_query_tokens) > 1):
                                pass
                            else:
                                strict_matched[i] = "OVERLAP"
        else:
            # Tetap masukkan ke relaxed_matched untuk safety net
            if i not in relaxed_matched:
                relaxed_matched[i] = "SUBSTR"

    # Prioritaskan strict match; fallback ke relaxed jika strict kosong
    chosen = strict_matched if strict_matched else relaxed_matched
    if not chosen:
        return []

    if match_log is not None:
        for idx, method in chosen.items():
            match_log.append((idx, method))

    return list(chosen.keys())


def _rematch_strict(
    suspend_row: dict, osbal_rows: list, lookup_osbal: tuple,
    polis_sus_cols: list, slip_sus_cols: list, insured_sus_cols: list,
    polis_ref_cols: list, slip_ref_cols: list, insured_ref_cols: list,
    facode_osbal_idx: dict,
    slipdb_rows: list = None, lookup_slipdb: tuple = None,
    facul_rows: list = None, lookup_facul: tuple = None,
    polis_sus_like_cols: list = None,
    effective_sus_date=None, is_lineslip: bool = False,
    allow_insured_fallback: bool = False,
) -> tuple:
    sus_curr = _normalize(suspend_row.get(SUSPEND_CURR_COL, ""))

    def _strict_filter(idxs) -> list:
        res = []
        for i in idxs:
            if sus_curr and _normalize(osbal_rows[i].get(OSBAL_CURR_COL, "")) != sus_curr:
                continue
            if pd.notna(effective_sus_date):
                com_date = osbal_rows[i].get("_com_date_parsed")
                if pd.notna(com_date):
                    if is_lineslip:
                        if com_date.year != effective_sus_date.year or com_date.month != effective_sus_date.month:
                            continue
                    else:
                        if effective_sus_date < com_date:
                            continue
            res.append(i)
        return res

    polis_vals = _collect_clean_values(suspend_row, polis_sus_cols)

    # R1: OSBAL polis exact
    curr_hits = _strict_filter(_exact_match(polis_vals, lookup_osbal[1]))
    if curr_hits:
        slip_vals = _collect_clean_values(suspend_row, slip_sus_cols)
        confirmed = [i for i in curr_hits if any(
            _normalize(osbal_rows[i].get(c, "")) in slip_vals
            for c in slip_ref_cols if _normalize(osbal_rows[i].get(c, ""))
        )]
        chosen = confirmed if confirmed else curr_hits
        return "OSBAL", ("Polis + Slip" if confirmed else "Polis only"), chosen, True

    # R2: OSBAL polis LIKE
    like_cols = polis_sus_like_cols if polis_sus_like_cols is not None else polis_sus_cols
    if lookup_osbal[7] and lookup_osbal[7][0] and like_cols:
        like_vals = _collect_clean_values(suspend_row, like_cols)
        if like_vals:
            curr_hits = _strict_filter(_like_match(
                like_vals, lookup_osbal[7][0], osbal_rows, polis_ref_cols,
                ori_col="polis_ori",
                stopwords=frozenset(),
                overlap_threshold=_POLIS_LIKE_THRESHOLD,
            ))
            if curr_hits:
                return "OSBAL", "Polis only", curr_hits, True

    # R3: SLIPDB polis exact
    if slipdb_rows and lookup_slipdb:
        hits = _exact_match(polis_vals, lookup_slipdb[1])
        if hits:
            res_idx, ok = _resolve_facode([slipdb_rows[i] for i in hits], SLIPDB_FACODE_COL, facode_osbal_idx, osbal_rows)
            if ok:
                curr_res = _strict_filter(res_idx)
                if curr_res:
                    return "SLIPDB", "Polis only", curr_res, True

    # R4: FACUL polis exact
    if facul_rows and lookup_facul:
        hits = _exact_match(polis_vals, lookup_facul[1])
        if hits:
            res_idx, ok = _resolve_facode([facul_rows[i] for i in hits], FACUL_FACODE_COL, facode_osbal_idx, osbal_rows)
            if ok:
                curr_res = _strict_filter(res_idx)
                if curr_res:
                    return "FACUL", "Polis only", curr_res, True

    # R5: OSBAL insured exact & LIKE (hanya jika allow_insured_fallback == True)
    if allow_insured_fallback:
        insured_vals_r5 = _collect_clean_values(suspend_row, insured_sus_cols)
        curr_hits = _strict_filter(_exact_match(insured_vals_r5, lookup_osbal[2]))
        if curr_hits:
            return "OSBAL", "Insured only", curr_hits, True

        if lookup_osbal[8] and lookup_osbal[8][0] and insured_vals_r5:
            curr_hits = _strict_filter(_like_match(
                insured_vals_r5, lookup_osbal[8][0], osbal_rows, insured_ref_cols,
                ori_col="insured_ori",
                stopwords=_INSURED_STOPWORDS,
                overlap_threshold=_INSURED_LIKE_THRESHOLD,
            ))
            if curr_hits:
                return "OSBAL", "Insured only", curr_hits, True

    return None


# --- Narrowing ---

def _narrow_by_currency(suspend_row: dict, matched: list, rows: list) -> list:
    if len(matched) <= 1:
        return matched
    sus_curr = _normalize(suspend_row.get(SUSPEND_CURR_COL, ""))
    if not sus_curr:
        return matched
    filtered = [i for i in matched if _normalize(rows[i].get(OSBAL_CURR_COL, "")) == sus_curr]
    return filtered if filtered else matched


def _narrow_by_periode(
    suspend_row: dict, matched: list, rows: list,
    effective_sus_date=None, is_lineslip: bool = False,
) -> list:
    if len(matched) <= 1:
        return matched
    sus_date = effective_sus_date
    if pd.isna(sus_date):
        return matched

    filtered = []
    for i in matched:
        com_date = rows[i].get("_com_date_parsed")
        if pd.isna(com_date):
            filtered.append(i)
        else:
            if is_lineslip:
                if com_date.year == sus_date.year and com_date.month == sus_date.month:
                    filtered.append(i)
            else:
                if sus_date >= com_date:
                    filtered.append(i)
    return filtered if filtered else matched


# --- Bordero Matching (Marine Cargo MSIG) ---

def _build_bordero_index_msig(bordero_rows: list) -> tuple:
    """Bangun index bordero Marine Cargo MSIG (fac_code dan polis)."""
    bordero_idx: dict = {}        # fac -> [entry, ...]
    bordero_polis_idx: dict = {}  # polis_normalized -> [(fac, entry), ...]

    for row in bordero_rows:
        fac = _normalize(row.get(BORDERO_FAC_COL, ""))
        if not fac:
            continue

        bulan_str = str(row.get(BORDERO_BULAN_COL, "")).strip().upper()
        tahun_str = str(row.get(BORDERO_TAHUN_COL, "")).strip()
        # Hapus trailing .0 dari tahun numeric
        tahun_str = re.sub(r'\.0+$', '', tahun_str)
        month = _BULAN_MAP.get(bulan_str, 0)
        ym = (int(tahun_str), month) if month and tahun_str.isdigit() else None

        polis_val = _normalize(row.get(BORDERO_POLIS_COL, ""))
        insured_val = _normalize(row.get(BORDERO_INSURED_COL, ""))
        curr_val = _normalize(row.get(BORDERO_CURR_COL, ""))

        net_raw = row.get(BORDERO_NET_COL)
        try:
            net_f = float(net_raw) if net_raw is not None else None
        except (ValueError, TypeError):
            net_f = None

        entry = {
            "polis":   polis_val,
            "insured": insured_val,
            "curr":    curr_val,
            "period":  ym,
            "net":     net_f,
        }
        bordero_idx.setdefault(fac, []).append(entry)

        if polis_val:
            bordero_polis_idx.setdefault(polis_val, []).append((fac, entry))

    return bordero_idx, bordero_polis_idx


def _find_fac_codes_in_bordero_msig(
    suspend_row: dict,
    bordero_idx: dict,
    bordero_polis_idx: dict,
    polis_sus_cols: list,
) -> tuple:
    """Cari FAC CODE di Bordero MSIG berdasarkan POLIS suspend."""
    if not bordero_polis_idx:
        return set(), False

    # Kumpulkan nilai polis dari baris suspend (semua clean cols + ori)
    sus_polis_set: set = set()
    for c in polis_sus_cols:
        v = _normalize(suspend_row.get(c, ""))
        if v:
            sus_polis_set.add(v)
    # Tambahkan dari kolom prioritas
    for c in [CLSDT_POLIS_COL, FAC_POLIS_COL, "POLIS"]:
        v = _normalize(suspend_row.get(c, ""))
        if v:
            sus_polis_set.add(v)

    if not sus_polis_set:
        return set(), False

    sus_curr_val = _normalize(suspend_row.get(SUSPEND_CURR_COL, ""))
    sus_date = suspend_row.get("_sus_date_parsed")
    sus_ym = (sus_date.year, sus_date.month) if sus_date is not None else None

    sus_amount_raw = suspend_row.get("AMOUNT ORI")
    sus_amount_f = None
    if sus_amount_raw is not None:
        try:
            s_amt = str(sus_amount_raw).strip().replace(' ', '')
            if ',' in s_amt and '.' in s_amt:
                s_amt = (s_amt.replace('.', '').replace(',', '.')
                         if s_amt.rfind(',') > s_amt.rfind('.')
                         else s_amt.replace(',', ''))
            elif ',' in s_amt:
                s_amt = s_amt.replace(',', '.')
            sus_amount_f = float(s_amt)
        except (ValueError, TypeError):
            sus_amount_f = None

    polis_f = {p for p in sus_polis_set if len(p) >= 5}
    if not polis_f:
        return set(), False

    # Stage 1: Exact match
    candidates: dict = {}
    for sp in polis_f:
        if sp in bordero_polis_idx:
            for fac, entry in bordero_polis_idx[sp]:
                candidates.setdefault(fac, []).append(entry)

    # Stage 2: Substring match (fallback jika exact kosong)
    if not candidates:
        for bp, entries_list in bordero_polis_idx.items():
            for sp in polis_f:
                if sp in bp or (len(bp) >= 5 and bp in sp):
                    for fac, entry in entries_list:
                        candidates.setdefault(fac, []).append(entry)
                    break

    if not candidates:
        return set(), False

    # Filter currency
    if sus_curr_val:
        curr_filtered: dict = {}
        for fac, entries in candidates.items():
            valid = [e for e in entries
                     if not e.get("curr") or e.get("curr") == sus_curr_val]
            if valid:
                curr_filtered[fac] = valid
        if curr_filtered:
            candidates = curr_filtered

    # Filter periode: periode bordero <= periode suspend
    if sus_ym:
        per_filtered: dict = {}
        for fac, entries in candidates.items():
            valid = [e for e in entries
                     if not e.get("period") or e["period"] <= sus_ym]
            if valid:
                per_filtered[fac] = valid
        if per_filtered:
            candidates = per_filtered

    # Filter net amount jika masih >1 FAC CODE
    if len(candidates) > 1 and sus_amount_f is not None:
        amt_abs = abs(sus_amount_f)
        net_filtered: dict = {}
        for fac, entries in candidates.items():
            valid = []
            for e in entries:
                try:
                    b_net = abs(float(e.get("net") or 0))
                    if abs(b_net - amt_abs) <= (amt_abs * NET_TOLERANCE_PCT + 0.05):
                        valid.append(e)
                except (ValueError, TypeError):
                    continue
            if valid:
                net_filtered[fac] = valid
        if net_filtered:
            candidates = net_filtered

    return set(candidates.keys()), True


# --- FAC Code Resolution ---

def _resolve_facode(ref_rows_matched: list, facode_col: str,
                    facode_osbal_idx: dict, osbal_rows: list) -> tuple:
    fac_codes = set()
    for r in ref_rows_matched:
        val = r.get(facode_col)
        if not val:
            for alt in ["FAC. CODE", "FAC_CODE", "FAC CODE", "CCOS_REF_CODE"]:
                if alt in r:
                    val = r[alt]
                    break
        norm_val = _normalize(val)
        if norm_val:
            fac_codes.add(norm_val)

    fac_codes.discard("")
    if not fac_codes:
        return [], False

    osbal_indices: set = set()
    for fac in fac_codes:
        if fac in facode_osbal_idx:
            osbal_indices.update(facode_osbal_idx[fac])
    result = list(osbal_indices)
    return result, bool(result)


# --- Output Builders ---

def _aggregate_ccos(osbal_rows: list, all_fac_osbal_rows: list = None) -> dict:
    fac_codes = {str(r.get("CCOS_REF_CODE", "")).strip() for r in osbal_rows
                 if str(r.get("CCOS_REF_CODE", "")).strip()}
    multi_fac = len(fac_codes) > 1
    bal_due_rows = all_fac_osbal_rows if (all_fac_osbal_rows is not None and not multi_fac) else osbal_rows

    def join_unique(key):
        seen, vals = set(), []
        for r in osbal_rows:
            v = str(r.get(key, "")).strip()
            if v and v not in seen:
                seen.add(v)
                vals.append(v)
        return ", ".join(vals)

    def sum_col(key, rows):
        if multi_fac:
            return np.nan
        total = 0.0
        for r in rows:
            v = r.get(key, 0)
            if pd.isna(v) or v == "":
                continue
            if isinstance(v, (int, float)):
                total += float(v)
            else:
                s = str(v).strip()
                if not s:
                    continue
                if ',' in s and '.' in s:
                    s = s.replace('.', '').replace(',', '.') if s.rfind(',') > s.rfind('.') else s.replace(',', '')
                elif ',' in s:
                    s = s.replace(',', '.')
                try:
                    total += float(s)
                except ValueError:
                    pass
        return round(total, 2)

    return {
        "CCOS_DOC_NO":   join_unique("CCOS_DOC_NO"),
        "CCOS_REF_CODE": join_unique("CCOS_REF_CODE"),
        "CCOS_OR_BAL":   np.nan if multi_fac else join_unique("CCOS_OR_BAL"),
        "CCOS_BAL_DUE":  sum_col("CCOS_BAL_DUE", bal_due_rows),
    }


def _facode_label(osbal_rows: list, source: str) -> str:
    if not osbal_rows or not osbal_rows[0]:
        return ""
    col = OSBAL_FACODE_COL if source == "OSBAL" else (
          SLIPDB_FACODE_COL if source == "SLIPDB" else FACUL_FACODE_COL)
    codes = {str(r.get(col, "")).strip() for r in osbal_rows if str(r.get(col, "")).strip()}
    return "facode lebih dari 1" if len(codes) > 1 else next(iter(codes), "")

def _clean_base_match(scenario: str) -> str:
    s = str(scenario or "").strip()
    for to_remove in [
        " -> Bordero OC Override", " (Bordero)", " (Beda Currency)",
        " (Beda Periode)", " (> 1 Fac)", " (> 1 Fac Bordero)",
        " (DESC 1)", " + Amount", " (Bordero MC)",
    ]:
        s = s.replace(to_remove, "")
    s = s.strip()
    if s.startswith("Slip + Polis"):
        s = s.replace("Slip + Polis", "Polis + Slip")
    if " only + " in s:
        s = s.replace(" only + ", " + ")
    if s == "Polis + Sertif":
        s = "Polis + Cert"
    if s in ("Insured", "Polis", "Slip"):
        s = s + " only"
    if s in ("Insured + Cert", "Insured + Cert only"):
        s = "Insured only"
    if s in ("Slip + Cert", "Slip + Cert only"):
        s = "Slip only"
    return s


def _format_final_scenario(
    base_scenario: str,
    source: str,
    is_bordero: bool = False,
    is_beda_curr: bool = False,
    is_beda_per: bool = False,
    num_fac: int = 1,
) -> str:
    if not source or not base_scenario or str(base_scenario).strip().upper() in ("UNMATCHING", "UNMATCHED"):
        return "Unmatching"

    has_override = "Bordero OC Override" in str(base_scenario)
    base = _clean_base_match(base_scenario)
    if not base or base.upper() in ("UNMATCHING", "UNMATCHED"):
        return "Unmatching"

    if has_override:
        base = "Polis only"
        if is_beda_curr:
            return f"{base} -> Bordero OC Override (Beda Currency)"
        elif is_beda_per:
            return f"{base} -> Bordero OC Override (Beda Periode)"
        elif num_fac > 1:
            return f"{base} -> Bordero OC Override (> 1 Fac)"
        else:
            return f"{base} -> Bordero OC Override"

    if is_beda_curr:
        return f"{base} (Beda Currency)"
    elif is_beda_per:
        return f"{base} (Beda Periode)"
    elif is_bordero and num_fac > 1:
        return f"{base} (> 1 Fac Bordero)"
    elif is_bordero and num_fac <= 1:
        return f"{base} (Bordero)"
    elif not is_bordero and num_fac > 1:
        return f"{base} (> 1 Fac)"
    else:
        return base


def _build_output_row(
    suspend_row: dict, source: str, scenario: str, osbal_rows: list,
    osbal_count: int = 0, resolved: bool = True,
    suspend_count: int = 1, all_fac_osbal_rows: list = None,
) -> dict:
    """Bangun satu baris output dictionary."""
    has_match = bool(source and osbal_rows and osbal_rows[0])
    has_osbal = has_match and (source in ("OSBAL", "SLIPDB", "FACUL") or resolved)

    ccos = (
        _aggregate_ccos(osbal_rows, all_fac_osbal_rows=all_fac_osbal_rows)
        if has_osbal
        else {"CCOS_DOC_NO": "", "CCOS_REF_CODE": "", "CCOS_OR_BAL": np.nan, "CCOS_BAL_DUE": np.nan}
    )
    if "Beda Currency" in str(scenario) or "Beda Periode" in str(scenario):
        ccos["CCOS_OR_BAL"] = np.nan
        ccos["CCOS_BAL_DUE"] = np.nan

    final_scenario = "Unmatching" if not source or scenario in ("Unmatching", "UNMATCHED") else scenario
    if final_scenario == "INSURED_LIKE":
        final_scenario = "Insured only"
    elif final_scenario == "POLIS_LIKE":
        final_scenario = "Polis only"
    elif final_scenario == "SLIP_LIKE":
        final_scenario = "Slip only"

    out_dict = {
        "CCOS_DOC_NO":       ccos["CCOS_DOC_NO"] if has_osbal else "",
        "CCOS_REF_CODE":     ccos["CCOS_REF_CODE"],
        "RECEIPT NO":        suspend_row.get("RECEIPT NO",        ""),
        "CREDIT NOTES":      suspend_row.get("CREDIT NOTES",      ""),
        "DETAIL RINCIAN NO": suspend_row.get("DETAIL RINCIAN NO", ""),
        "RECEIPT DATE":      suspend_row.get("RECEIPT DATE",      ""),
        "CEDANT NAME":       suspend_row.get("CEDANT NAME",       ""),
        "CEDANT SHRT NAME":  suspend_row.get("CEDANT SHRT NAME",  ""),

        # Kolom insured
        "INSURED_ORI": suspend_row.get("INSURED",       ""),
        "INSURED_1":   suspend_row.get("INSURED_CLN_1", ""),
        "INSURED_2":   suspend_row.get("INSURED_CLN_2", ""),

        "CURR ORI":   suspend_row.get("CURR ORI",   ""),
        "AMOUNT ORI": suspend_row.get("AMOUNT ORI", ""),
        "CURR PAY":   suspend_row.get("CURR PAY",   ""),
        "AMOUNT PAY": suspend_row.get("AMOUNT PAY", ""),

        "CCOS_OR_BAL":  ccos["CCOS_OR_BAL"],
        "CCOS_BAL_DUE": ccos["CCOS_BAL_DUE"],

        # Kolom polis dan sertifikat
        "POLIS_ORI": suspend_row.get("POLIS", ""),
        "POLICY_CLEAN_1": str(suspend_row["POLIS_CLN_1"]).strip() if suspend_row.get("POLIS_CLN_1") and str(suspend_row.get("POLIS_CLN_1")).strip() not in ("", "nan", "None") else (str(suspend_row["POLIS"]).strip() if suspend_row.get("POLIS") and str(suspend_row.get("POLIS")).strip() not in ("", "nan", "None") else ""),
        "POLICY_CLEAN_2": suspend_row.get("POLIS_CLN_2", ""),
        "POLICY_CLEAN_3": suspend_row.get("POLIS_CLN_3", ""),
        "POLICY_CLEAN_4": suspend_row.get("POLIS_CLN_4", ""),
        "POLICY_CLEAN_5": suspend_row.get("POLIS_CLN_5", ""),
        
        "SERTIF_CLEAN_1": str(suspend_row.get("CERTIFICATE", "") or suspend_row.get("clean sertif 1", "")).strip(),
        "SERTIF_CLEAN_2": str(suspend_row.get("clean sertif 2", "")).strip(),
        "SERTIF_CLEAN_3": str(suspend_row.get("clean sertif 3", "")).strip(),
        
        "SLIP_NO_ORI": suspend_row.get("SLIP NO", ""),

        "DESC 1": suspend_row.get("DESC 1", ""),
        "DESC 2": suspend_row.get("DESC 2", ""),
        "DESC 3": suspend_row.get("DESC 3", ""),
        "DESC 4": suspend_row.get("DESC 4", ""),

        "STATUS":   suspend_row.get("STATUS",   ""),
        "REC_TYPE": suspend_row.get("REC_TYPE", ""),

        "SKENARIO":           final_scenario,
        "_OSBAL_ROW_COUNT":   osbal_count,
        "_SUSPEND_ROW_COUNT": suspend_count,
    }

    # Tambahkan kolom clean slip dari suspend_row
    for k, v in suspend_row.items():
        k_up = str(k).upper()
        if k_up.startswith("SLIP_NO_CLN") or k_up.startswith("CLEAN SLIP"):
            out_dict[k] = v

    return out_dict


def _merge_suspend_rows(rows: list) -> dict:
    if len(rows) == 1:
        return rows[0]
    merged: dict = {}
    for col in _SUSPEND_SUM_COLS:
        total = 0.0
        for r in rows:
            v = r.get(col, 0)
            if pd.isna(v) or v == "":
                continue
            if isinstance(v, (int, float)):
                total += float(v)
            else:
                s = str(v).strip()
                if not s:
                    continue
                if ',' in s and '.' in s:
                    s = s.replace('.', '').replace(',', '.') if s.rfind(',') > s.rfind('.') else s.replace(',', '')
                elif ',' in s:
                    s = s.replace(',', '.')
                try:
                    total += float(s)
                except ValueError:
                    pass
        merged[col] = total
    for col in _SUSPEND_JOIN_COLS:
        seen, vals = set(), []
        for r in rows:
            v = str(r.get(col, "")).strip()
            if v and v not in seen:
                seen.add(v)
                vals.append(v)
        merged[col] = ", ".join(vals)
    return merged


def _compute_derived_cols(df: pd.DataFrame) -> pd.DataFrame:
    """Hitung AMOUNT_ORI_MIN1, DIFERENCE, dan FLAG_PROD via matriks np.select."""
    amount_ori_base = df.get("AMOUNT ORI", pd.Series([0]*len(df), index=df.index))
    if "_MERGED_AMOUNT_ORI" in df.columns:
        amount_ori_raw = df["_MERGED_AMOUNT_ORI"].combine_first(amount_ori_base)
    else:
        amount_ori_raw = amount_ori_base

    amount_ori_clean = amount_ori_raw.replace(r'^\s*$', np.nan, regex=True)
    if amount_ori_clean.dtype == object:
        amount_ori_clean = amount_ori_clean.astype(str).str.strip().str.replace(r'\s+', '', regex=True)
        amount_ori_clean = amount_ori_clean.replace("nan", np.nan)
        has_comma = amount_ori_clean.str.contains(',', na=False)
        amount_ori_eu = (amount_ori_clean
                         .str.replace('.', '', regex=False)
                         .str.replace(',', '.', regex=False))
        amount_ori_clean = amount_ori_eu.where(has_comma, amount_ori_clean)
    amount_ori_numeric = pd.to_numeric(amount_ori_clean, errors="coerce")

    df["AMOUNT_ORI_MIN1"] = np.where(amount_ori_numeric.notna(), (amount_ori_numeric * -1).round(2), np.nan)
    amount_ori  = amount_ori_numeric.fillna(0).round(2)
    amount_neg  = df["AMOUNT_ORI_MIN1"].fillna(0).round(2)

    balance_due     = pd.to_numeric(df.get("CCOS_BAL_DUE",      0), errors="coerce").fillna(0).round(2)
    osbal_count_col = pd.to_numeric(df.get("_OSBAL_ROW_COUNT",  0), errors="coerce").fillna(0)
    sus_count_col   = pd.to_numeric(df.get("_SUSPEND_ROW_COUNT",1), errors="coerce").fillna(1)
    ccos_ref        = df.get("CCOS_REF_CODE", pd.Series("", index=df.index)).fillna("").astype(str)
    skenario_col    = df.get("SKENARIO",      pd.Series("", index=df.index)).fillna("")

    df["DIFERENCE"] = (amount_neg - balance_due).round(2)
    accumulated = (osbal_count_col > 1) | (sus_count_col > 1)
    is_equal    = amount_neg == balance_due

    is_beda_periode  = skenario_col.str.contains("Beda Periode", na=False, regex=False)
    if "_IS_BEDA_CURR" in df.columns:
        is_beda_currency = df["_IS_BEDA_CURR"].fillna(False).astype(bool)
    else:
        is_beda_currency = skenario_col.str.contains("Beda Currency", na=False, regex=False)

    is_gt1_ocmc          = (
        (skenario_col.str.contains("Bordero OC Override", na=False, regex=False) & ccos_ref.str.contains(",", na=False))
        | skenario_col.str.contains("> 1 Fac Bordero", na=False, regex=False)
        | (skenario_col.str.contains("Bordero", na=False, regex=False) & ccos_ref.str.contains(",", na=False))
    )
    is_gt1_fac           = (skenario_col.str.contains("> 1 Fac", na=False, regex=False) | ccos_ref.str.contains(",", na=False)) & ~is_gt1_ocmc
    is_new_entry_total   = (amount_ori != 0) & (balance_due == 0)
    is_adj_total         = is_equal
    is_adj_sebagian      = ~is_equal & (amount_neg < balance_due)
    is_new_entry_sebagian= ~is_equal & (amount_neg > balance_due)

    df["FLAG_PROD"] = np.select(
        [is_beda_currency, is_beda_periode, is_gt1_ocmc, is_gt1_fac,
         is_new_entry_total,
         is_adj_total & ~accumulated, is_adj_total & accumulated,
         is_adj_sebagian & ~accumulated, is_adj_sebagian & accumulated,
         is_new_entry_sebagian],
        ["Beda Currency", "Beda Periode",
         "Matching >1 Fac code OC MC", "Matching >1 fac code",
         "New Entry total",
         "Adjustment total tanpa akumulasi", "Adjustment total dengan akumulasi",
         "Adjustment sebagian tanpa akumulasi", "Adjustment sebagian dengan akumulasi",
         "New Entry sebagian"],
        default="Unmatching",
    )

    for col in ["_OSBAL_ROW_COUNT", "_SUSPEND_ROW_COUNT", "_MERGED_AMOUNT_ORI", "_IS_BEDA_CURR"]:
        if col in df.columns:
            df.drop(columns=[col], inplace=True)
    return df


# --- Pipeline ---

def _load_excel(path: str, label: str, optional: bool = False) -> tuple:
    """Load Excel ke (list[dict], list[str]) dengan auto-detect header dan pickle cache."""
    if not os.path.exists(path):
        print(f"  {label}: {path} ({'tidak ada, dilewati' if optional else 'TIDAK DITEMUKAN'})", flush=True)
        return [], []

    cache = path + ".cache.pkl"
    if os.path.exists(cache) and os.path.getmtime(cache) >= os.path.getmtime(path):
        t = time.perf_counter()
        print(f"  {label}: loading dari cache ...", flush=True)
        try:
            data, header = pd.read_pickle(cache)
            print(f"  -> {len(data):,} rows, {len(header)} cols ({time.perf_counter()-t:.2f}s)", flush=True)
            return data, header
        except Exception:
            print(f"  {label}: cache rusak, baca ulang dari Excel ...", flush=True)

    t = time.perf_counter()
    print(f"  {label}: {path} ...", flush=True)
    try:
        df = pd.read_excel(path)
    except Exception as e:
        if isinstance(e, PermissionError) and os.path.exists(cache):
            print(f"  [FALLBACK] File sedang dibuka user, loading dari cache yang ada ...", flush=True)
            data, header = pd.read_pickle(cache)
            return data, header
        print(f"  [RETRY] engine default gagal ({e}), coba openpyxl ...", flush=True)
        try:
            df = pd.read_excel(path, engine="openpyxl")
        except PermissionError:
            if os.path.exists(cache):
                print(f"  [FALLBACK] File sedang dibuka user, loading dari cache yang ada ...", flush=True)
                data, header = pd.read_pickle(cache)
                return data, header
            raise

    # Auto-detect baris header jika row 0 bukan header sebenarnya
    header_cols = {str(c).strip().upper() for c in df.columns}
    if not any(kw in header_cols for kw in ["FAC_CODE", "FAC. CODE", "FAC CODE", "CCOS_REF_CODE", "POLIS_CLN_1", "RECEIPT NO"]):
        for r_idx in range(min(5, len(df))):
            row_vals = {str(v).strip().upper() for v in df.iloc[r_idx] if pd.notna(v)}
            if any(kw in row_vals for kw in ["FAC. CODE", "FAC CODE", "POLICY NO", "POLICY_NO", "SLIP"]):
                try:
                    df = pd.read_excel(path, header=r_idx + 1)
                except Exception:
                    df = pd.read_excel(path, header=r_idx + 1, engine="openpyxl")
                break

    header = list(df.columns)
    data   = df.to_dict("records")
    print(f"  -> {len(data):,} rows, {len(header)} cols ({time.perf_counter()-t:.1f}s)", flush=True)

    try:
        pd.to_pickle((data, header), cache)
        print(f"  -> cache disimpan: {cache}", flush=True)
    except Exception:
        pass
    return data, header


def run() -> None:
    """Pipeline utama MSIG: matching 6 fase (termasuk Bordero Marine Cargo)."""
    t_start = time.perf_counter()
    print("\n" + "=" * 70, flush=True)
    print("  PRODUCTION SCRIPT -- SUSPEND MATCHING  (MSIG INDONESIA)", flush=True)
    print("=" * 70, flush=True)

    # [1] Load data
    print("\n[1/6] Loading data ...", flush=True)
    suspend_rows, suspend_cols = _load_excel(SUSPEND_FILE,  "SUSPEND")
    osbal_rows,   osbal_cols   = _load_excel(OSBAL_FILE,    "OSBAL")
    facul_rows,   facul_cols   = _load_excel(FACUL_FILE,    "FACUL")
    slipdb_rows,  slipdb_cols  = _load_excel(SLIPDB_FILE,   "SLIPDB",  optional=True)
    bordero_rows, bordero_cols = _load_excel(BORDERO_FILE,  "BORDERO", optional=True)

    # Diagnostik kolom kunci
    print("\n  --- Cek kolom SUSPEND ---", flush=True)
    for col in ["INSURED", "INSURED_CLN_1", "INSURED_CLN_2",
                "POLIS", "POLIS_CLN_1", "SLIP NO", "SLIP_NO_CLN_1",
                SUSPEND_CURR_COL, SUSPEND_DATE_COL]:
        status = "OK" if col in suspend_cols else "[MISSING]"
        print(f"    {col:<22}: {status}", flush=True)

    print("\n  --- Cek kolom OSBAL ---", flush=True)
    for col in [OSBAL_FACODE_COL, OSBAL_CURR_COL, OSBAL_DATE_COL,
                "CCOS_DOC_NO", "CCOS_OR_BAL", "CCOS_BAL_DUE"]:
        status = "OK" if col in osbal_cols else "[MISSING]"
        print(f"    {col:<22}: {status}", flush=True)

    # Pre-parse tanggal OSBAL FAC_COM_DATE
    print("  Pre-parsing OSBAL FAC_COM_DATE ...", flush=True)
    for r in osbal_rows:
        raw = r.get(OSBAL_DATE_COL, "")
        r["_com_date_parsed"] = pd.to_datetime(raw, errors="coerce") \
            if raw and not (isinstance(raw, float) and pd.isna(raw)) else None

    # Pre-parse tanggal Suspend RECEIPT DATE
    print("  Pre-parsing Suspend RECEIPT DATE ...", flush=True)
    for r in suspend_rows:
        raw = r.get(SUSPEND_DATE_COL, "")
        r["_sus_date_parsed"] = pd.to_datetime(raw, errors="coerce") \
            if raw and not (isinstance(raw, float) and pd.isna(raw)) else None

    # [2] Deteksi kolom bersih
    print("\n[2/6] Detecting clean columns ...", flush=True)

    # Kolom polis/slip/insured suspend
    polis_sus   = _get_clean_cols(suspend_cols, "POLIS_CLN") or _get_clean_cols(suspend_cols, "clean polis") or []
    slip_sus    = _get_clean_cols(suspend_cols, "SLIP_NO_CLN") or _get_clean_cols(suspend_cols, "clean slip")  or []
    insured_sus = _get_clean_cols(suspend_cols, "INSURED_CLN") or _get_clean_cols(suspend_cols, "clean insured") or []

    # Pre-parse LINESLIP flags
    print("  Pre-parsing LINESLIP flags ...", flush=True)
    lineslip_count = 0
    for r in suspend_rows:
        r["_is_lineslip"] = _is_lineslip_row(r, insured_sus)
        if r["_is_lineslip"]:
            lineslip_count += 1
            parsed = r.get("_sus_date_parsed")
            try:
                r["_sus_date_lineslip"] = (parsed - pd.DateOffset(months=1)) if parsed is not None else None
            except Exception:
                r["_sus_date_lineslip"] = None
        else:
            r["_sus_date_lineslip"] = None
    print(f"  -> {lineslip_count:,} baris LINESLIP terdeteksi", flush=True)

    polis_osbal   = _get_clean_cols(osbal_cols, "POLIS_CLN") or _get_clean_cols(osbal_cols, "clean polis") or []
    slip_osbal    = _get_clean_cols(osbal_cols, "SLIP_NO_CLN") or _get_clean_cols(osbal_cols, "clean slip") or []
    insured_osbal = _get_clean_cols(osbal_cols, "INSURED_CLN") or _get_clean_cols(osbal_cols, "clean insured") or []
    sertif_osbal  = _get_clean_cols(osbal_cols, "clean sertif") or []

    polis_facul   = _get_clean_cols(facul_cols, "POLIS_CLN") or _get_clean_cols(facul_cols, "clean polis") or []
    slip_facul    = _get_clean_cols(facul_cols, "SLIP_NO_CLN") or _get_clean_cols(facul_cols, "clean slip") or []
    insured_facul = _get_clean_cols(facul_cols, "INSURED_CLN") or _get_clean_cols(facul_cols, "clean insured") or []

    polis_slipdb   = _get_clean_cols(slipdb_cols, "POLIS_CLN") or _get_clean_cols(slipdb_cols, "clean polis") or []
    slip_slipdb    = _get_clean_cols(slipdb_cols, "SLIP_NO_CLN") or _get_clean_cols(slipdb_cols, "clean slip") or []
    insured_slipdb = _get_clean_cols(slipdb_cols, "INSURED_CLN") or _get_clean_cols(slipdb_cols, "clean insured") or []

    if not polis_slipdb:
        polis_slipdb = [c for c in ["CLSDT_POLICY_NO", "FAC_POLICY_NO", "POLICY NO", "POLICY_NO", "POLIS"] if c in slipdb_cols]
    if not slip_slipdb:
        slip_slipdb  = [c for c in ["CLSDT_SLIP_NO", "FAC_SLIP", "SLIP", "SLIP_NO", "SLIP NO"] if c in slipdb_cols]
    if not insured_slipdb:
        insured_slipdb = [c for c in ["FAC_INSURED", "INSURED", "CEDANT_NAME", "CEDANT NAME"] if c in slipdb_cols]

    global SLIPDB_FACODE_COL
    if   "FAC. CODE"     in slipdb_cols: SLIPDB_FACODE_COL = "FAC. CODE"
    elif "FAC_CODE"      in slipdb_cols: SLIPDB_FACODE_COL = "FAC_CODE"
    elif "FAC CODE"      in slipdb_cols: SLIPDB_FACODE_COL = "FAC CODE"
    elif "CCOS_REF_CODE" in slipdb_cols: SLIPDB_FACODE_COL = "CCOS_REF_CODE"

    def _prepend(col, cols_list, header):
        return ([col] + cols_list) if col in header and col not in cols_list else cols_list

    polis_osbal  = _prepend("CLSDT_POLICY_NO", polis_osbal,  osbal_cols)
    slip_osbal   = _prepend("CLSDT_SLIP_NO",   slip_osbal,   osbal_cols)
    polis_facul  = _prepend("CLSDT_POLICY_NO", polis_facul,  facul_cols)
    slip_facul   = _prepend("CLSDT_SLIP_NO",   slip_facul,   facul_cols)
    polis_slipdb = _prepend("CLSDT_POLICY_NO", polis_slipdb, slipdb_cols)
    slip_slipdb  = _prepend("CLSDT_SLIP_NO",   slip_slipdb,  slipdb_cols)

    # [3] Bangun lookup index
    print("\n[3/6] Building lookup indexes ...", flush=True)
    lookup_osbal  = _build_lookup(osbal_rows,  polis_osbal,  slip_osbal,  insured_osbal,  OSBAL_FACODE_COL,  label="OSBAL")
    lookup_facul  = _build_lookup(facul_rows,  polis_facul,  slip_facul,  insured_facul,  FACUL_FACODE_COL,  label="FACUL")
    lookup_slipdb = _build_lookup(slipdb_rows, polis_slipdb, slip_slipdb, insured_slipdb, SLIPDB_FACODE_COL, label="SLIPDB")

    facode_osbal_idx = lookup_osbal[9]

    # Bangun index Bordero MSIG (Marine Cargo -- polis-based)
    bordero_idx, bordero_polis_idx = (
        _build_bordero_index_msig(bordero_rows) if bordero_rows else ({}, {})
    )
    if bordero_idx:
        print(f"  Bordero MSIG: {len(bordero_idx):,} FAC CODE unik, "
              f"{len(bordero_polis_idx):,} polis unik", flush=True)
    else:
        print("  Bordero MSIG: tidak tersedia, FASE 6 dilewati", flush=True)

    # Fac code milik baris marker (hutang/suspense) -> tidak boleh dipakai di Slip only
    bad_fc_osbal  = _excluded_facodes(osbal_rows,  OSBAL_FACODE_COL,  polis_osbal,  slip_osbal,  insured_osbal)
    bad_fc_slipdb = _excluded_facodes(slipdb_rows, SLIPDB_FACODE_COL, polis_slipdb, slip_slipdb, insured_slipdb)
    bad_fc_facul  = _excluded_facodes(facul_rows,  FACUL_FACODE_COL,  polis_facul,  slip_facul,  insured_facul)
    bad_fc_all    = bad_fc_osbal | bad_fc_slipdb | bad_fc_facul

    def _drop_bad_facode_osbal(idx_list):
        return [i for i in idx_list
                if _normalize(osbal_rows[i].get(OSBAL_FACODE_COL, "")) not in bad_fc_all]

    # [4] Matching per baris Suspend (Pipeline MSIG)
    total = len(suspend_rows)
    print(f"\n[4/6] Matching {total:,} suspend rows ...", flush=True)
    t4 = time.perf_counter()
    raw_results = []
    stat_like_substr  = 0
    stat_like_subset  = 0
    stat_like_overlap = 0

    for n, sus in enumerate(suspend_rows, 1):
        if n % 500 == 0:
            elapsed_so_far = time.perf_counter() - t4
            rate = n / elapsed_so_far if elapsed_so_far > 0 else 0
            eta  = (total - n) / rate if rate > 0 else 0
            print(f"    Memproses baris {n:,} / {total:,}  ({rate:.0f} rows/s, ETA {eta:.0f}s)", flush=True)

        if str(sus.get("STATUS", "")).strip().upper() == "MATCHING":
            continue

        eff_polis, eff_slip, like_polis, like_slip = _get_effective_sus_cols(sus, polis_sus, slip_sus)
        _is_ls = sus.get("_is_lineslip", False)
        effective_date = sus.get("_sus_date_lineslip") if _is_ls else sus.get("_sus_date_parsed")

        matched_osbal_idx = set()
        source = None
        scenario = ""
        resolved = False
        _like_log: list = []

        # FASE 1: Pencarian berdasarkan POLIS
        found_by_polis = False
        polis_vals = _collect_clean_values(sus, eff_polis)
        if polis_vals:
            # Step 1a: Cek OSBAL terlebih dahulu
            osb_hits = _exact_match(polis_vals, lookup_osbal[1])
            if osb_hits:
                matched_osbal_idx.update(osb_hits)
                found_by_polis = True

            # Step 1b: Hanya jika OSBAL miss -> cek SLIPDB
            if not found_by_polis and slipdb_rows:
                sld_hits = _exact_match(polis_vals, lookup_slipdb[1])
                if sld_hits:
                    res_idx, _ = _resolve_facode(
                        [slipdb_rows[idx] for idx in sld_hits],
                        SLIPDB_FACODE_COL, facode_osbal_idx, osbal_rows
                    )
                    res_idx = _drop_bad_facode_osbal(res_idx)
                    if res_idx:
                        matched_osbal_idx.update(res_idx)
                        found_by_polis = True

            # Step 1c: Hanya jika OSBAL & SLIPDB miss -> cek FACUL
            if not found_by_polis and facul_rows:
                fac_hits = _exact_match(polis_vals, lookup_facul[1])
                if fac_hits:
                    res_idx, _ = _resolve_facode(
                        [facul_rows[idx] for idx in fac_hits],
                        FACUL_FACODE_COL, facode_osbal_idx, osbal_rows
                    )
                    res_idx = _drop_bad_facode_osbal(res_idx)
                    if res_idx:
                        matched_osbal_idx.update(res_idx)
                        found_by_polis = True

        if found_by_polis:
            source = "OSBAL"
            scenario = "Polis only"

        # FASE 2: Pencarian berdasarkan SLIP (fallback jika FASE 1 miss)
        found_by_slip = False
        if ENABLE_SLIP_MATCHING and not found_by_polis:
            slip_vals = _collect_clean_values(sus, eff_slip)
            if slip_vals:
                # 1. Coba OSBAL
                osb_hits = _exact_match(slip_vals, lookup_osbal[0])
                osb_hits = [i for i in osb_hits if _normalize(osbal_rows[i].get(OSBAL_FACODE_COL, "")) not in bad_fc_all]
                if osb_hits:
                    matched_osbal_idx.update(osb_hits)
                    source = "OSBAL"
                    scenario = "Slip only"
                    found_by_slip = True
                else:
                    # 2. Coba SLIPDB
                    if slipdb_rows:
                        sld_hits = _exact_match(slip_vals, lookup_slipdb[0])
                        sld_hits = [i for i in sld_hits if _normalize(slipdb_rows[i].get(SLIPDB_FACODE_COL, "")) not in bad_fc_all]
                        if sld_hits:
                            res_idx, _ = _resolve_facode(
                                [slipdb_rows[idx] for idx in sld_hits],
                                SLIPDB_FACODE_COL, facode_osbal_idx, osbal_rows
                            )
                            res_idx = _drop_bad_facode_osbal(res_idx)
                            if res_idx:
                                matched_osbal_idx.update(res_idx)
                                source = "SLIPDB"
                                scenario = "Slip only"
                                resolved = True
                                found_by_slip = True
                    # 3. Coba FACUL
                    if not matched_osbal_idx and facul_rows:
                        fac_hits = _exact_match(slip_vals, lookup_facul[0])
                        fac_hits = [i for i in fac_hits if _normalize(facul_rows[i].get(FACUL_FACODE_COL, "")) not in bad_fc_all]
                        if fac_hits:
                            res_idx, _ = _resolve_facode(
                                [facul_rows[idx] for idx in fac_hits],
                                FACUL_FACODE_COL, facode_osbal_idx, osbal_rows
                            )
                            res_idx = _drop_bad_facode_osbal(res_idx)
                            if res_idx:
                                matched_osbal_idx.update(res_idx)
                                source = "FACUL"
                                scenario = "Slip only"
                                resolved = True
                                found_by_slip = True

        # FASE 3: Pencarian berdasarkan INSURED (Fallback)
        if not matched_osbal_idx:
            ins_vals = _collect_clean_values(sus, insured_sus)
            if ins_vals:
                # 1. Coba OSBAL Exact & Like
                osb_hits = _exact_match(ins_vals, lookup_osbal[2])
                # Filter bad fac code (admin row yang lolos exclude)
                osb_hits = [i for i in osb_hits
                            if _normalize(osbal_rows[i].get(OSBAL_FACODE_COL, "")) not in bad_fc_all]
                if not osb_hits and lookup_osbal[8] and lookup_osbal[8][0]:
                    osb_hits = _like_match(
                        ins_vals, lookup_osbal[8][0], osbal_rows, insured_osbal,
                        ori_col="insured_ori",
                        stopwords=_INSURED_STOPWORDS,
                        overlap_threshold=_INSURED_LIKE_THRESHOLD,
                        match_log=_like_log,
                    )
                    osb_hits = [i for i in osb_hits
                                if _normalize(osbal_rows[i].get(OSBAL_FACODE_COL, "")) not in bad_fc_all]
                if osb_hits:
                    matched_osbal_idx.update(osb_hits)
                    source = "OSBAL"
                    scenario = "Insured only"

                # 2. Coba SLIPDB Exact & Like
                elif slipdb_rows and not matched_osbal_idx:
                    sld_hits = _exact_match(ins_vals, lookup_slipdb[2])
                    sld_hits = [i for i in sld_hits
                                if _normalize(slipdb_rows[i].get(SLIPDB_FACODE_COL, "")) not in bad_fc_all]
                    if not sld_hits and lookup_slipdb[8] and lookup_slipdb[8][0]:
                        sld_hits = _like_match(
                            ins_vals, lookup_slipdb[8][0], slipdb_rows, insured_slipdb,
                            ori_col="insured_ori",
                            stopwords=_INSURED_STOPWORDS,
                            overlap_threshold=_INSURED_LIKE_THRESHOLD,
                            match_log=_like_log,
                        )
                        sld_hits = [i for i in sld_hits
                                    if _normalize(slipdb_rows[i].get(SLIPDB_FACODE_COL, "")) not in bad_fc_all]
                    if sld_hits:
                        res_idx, _ = _resolve_facode(
                            [slipdb_rows[idx] for idx in sld_hits],
                            SLIPDB_FACODE_COL, facode_osbal_idx, osbal_rows
                        )
                        res_idx = _drop_bad_facode_osbal(res_idx)
                        if res_idx:
                            matched_osbal_idx.update(res_idx)
                            source = "SLIPDB"
                            scenario = "Insured only"
                            resolved = True

                # 3. Coba FACUL Exact & Like
                elif facul_rows and not matched_osbal_idx:
                    fac_hits = _exact_match(ins_vals, lookup_facul[2])
                    fac_hits = [i for i in fac_hits
                                if _normalize(facul_rows[i].get(FACUL_FACODE_COL, "")) not in bad_fc_all]
                    if not fac_hits and lookup_facul[8] and lookup_facul[8][0]:
                        fac_hits = _like_match(
                            ins_vals, lookup_facul[8][0], facul_rows, insured_facul,
                            ori_col="insured_ori",
                            stopwords=_INSURED_STOPWORDS,
                            overlap_threshold=_INSURED_LIKE_THRESHOLD,
                            match_log=_like_log,
                        )
                        fac_hits = [i for i in fac_hits
                                    if _normalize(facul_rows[i].get(FACUL_FACODE_COL, "")) not in bad_fc_all]
                    if fac_hits:
                        res_idx, _ = _resolve_facode(
                            [facul_rows[idx] for idx in fac_hits],
                            FACUL_FACODE_COL, facode_osbal_idx, osbal_rows
                        )
                        res_idx = _drop_bad_facode_osbal(res_idx)
                        if res_idx:
                            matched_osbal_idx.update(res_idx)
                            source = "FACUL"
                            scenario = "Insured only"
                            resolved = True

        # FASE 4: Narrowing Wajib (Currency & Periode) + Strict Rematch
        is_beda_curr = False
        is_beda_per  = False
        if matched_osbal_idx:
            matched_list = list(matched_osbal_idx)

            sus_curr = _normalize(sus.get(SUSPEND_CURR_COL, ""))
            if sus_curr:
                filtered_curr = [i for i in matched_list if _normalize(osbal_rows[i].get(OSBAL_CURR_COL, "")) == sus_curr]
                if not filtered_curr:
                    is_beda_curr = True
                elif len(filtered_curr) < len(matched_list):
                    matched_list = filtered_curr

            if not is_beda_curr:
                sus_date = effective_date
                if pd.notna(sus_date):
                    filtered_per = []
                    for i in matched_list:
                        com_date = osbal_rows[i].get("_com_date_parsed")
                        if pd.isna(com_date):
                            filtered_per.append(i)
                        else:
                            if _is_ls:
                                if com_date.year == sus_date.year and com_date.month == sus_date.month:
                                    filtered_per.append(i)
                            else:
                                if sus_date >= com_date:
                                    filtered_per.append(i)
                    if not filtered_per:
                        is_beda_per = True
                    elif len(filtered_per) < len(matched_list):
                        matched_list = filtered_per

            if is_beda_curr or is_beda_per:
                # Jangan downgrade ke Insured jika awalnya sudah ketemu via Polis/Slip
                allow_ins = not (found_by_polis or found_by_slip)
                rm = _rematch_strict(
                    sus, osbal_rows, lookup_osbal,
                    eff_polis, eff_slip, insured_sus,
                    polis_osbal, slip_osbal, insured_osbal,
                    facode_osbal_idx,
                    slipdb_rows, lookup_slipdb,
                    facul_rows, lookup_facul,
                    like_polis,
                    effective_sus_date=effective_date, is_lineslip=_is_ls,
                    allow_insured_fallback=allow_ins,
                )
                if rm:
                    src, scen, hits, _ = rm
                    source = src
                    scenario = scen
                    matched_list = hits
                    is_beda_curr = False
                    is_beda_per  = False

            matched_osbal_idx = set(matched_list)

        # FASE 5: Narrowing Non-Destruktif (Insured)
        if len(matched_osbal_idx) > 1:
            matched_list = list(matched_osbal_idx)

            # Step 1: Insured Matching (Slip tidak dipakai di MSIG)
            ins_vals = _collect_all_values(sus, insured_sus, "insured_ori")
            if ins_vals:
                confirmed = [idx for idx in matched_list if _ref_has_value(osbal_rows[idx], insured_osbal, "insured_ori", ins_vals)]
                if confirmed:
                    matched_list = confirmed
                    if "Insured" not in scenario:
                        insured_vals_before_narrowing: set = set()
                        for _idx in list(matched_osbal_idx):
                            for _c in insured_osbal:
                                _v = _normalize(osbal_rows[_idx].get(_c, ""))
                                if _v:
                                    insured_vals_before_narrowing.add(_v)
                        if len(insured_vals_before_narrowing) > 1:
                            base = scenario.replace(" only", "") if scenario.endswith(")") or not scenario.endswith("only") else scenario
                            scenario = base + " + Insured"

            matched_osbal_idx = set(matched_list)

        # FASE 6: Bordero Marine Cargo MSIG (polis-based override dan narrowing)
        is_bordero = False
        if bordero_polis_idx:
            bordero_facs, ok = _find_fac_codes_in_bordero_msig(
                sus, bordero_idx, bordero_polis_idx, eff_polis
            )
            if ok and bordero_facs:
                # Resolve FAC CODE bordero -> baris OSBAL
                res_idx: list = []
                for fac in bordero_facs:
                    if fac in facode_osbal_idx:
                        res_idx.extend(facode_osbal_idx[fac])
                res_idx = list(dict.fromkeys(res_idx))  # deduplicate, preserve order

                if res_idx:
                    # Skenario a (Unmatching) atau skenario b (narrowing >1 fac)
                    apply_bordero = False

                    if not matched_osbal_idx:
                        # Skenario a): suspend sebelumnya tidak ketemu match
                        apply_bordero = True
                    else:
                        # Skenario b): narrowing: cek apakah FAC CODE bordero ada di match saat ini
                        current_facs = {
                            _normalize(osbal_rows[i].get(OSBAL_FACODE_COL, ""))
                            for i in matched_osbal_idx
                        }
                        current_facs.discard("")
                        overlap = bordero_facs & current_facs
                        if overlap and len(overlap) < len(current_facs):
                            # Bordero mempersempit kandidat FAC CODE
                            res_idx = [
                                i for i in res_idx
                                if _normalize(osbal_rows[i].get(OSBAL_FACODE_COL, "")) in overlap
                            ]
                            apply_bordero = True
                        elif not matched_osbal_idx or (len(current_facs) > 1 and bordero_facs):
                            # Tidak ada overlap tapi bordero punya kandidat -- override
                            apply_bordero = True

                    if apply_bordero:
                        # Validasi Currency & Periode untuk kandidat Bordero
                        b_curr_diff = False
                        b_per_diff  = False

                        sus_curr_b = _normalize(sus.get(SUSPEND_CURR_COL, ""))
                        if sus_curr_b:
                            filtered_curr_b = [
                                i for i in res_idx
                                if _normalize(osbal_rows[i].get(OSBAL_CURR_COL, "")) == sus_curr_b
                            ]
                            if not filtered_curr_b:
                                b_curr_diff = True
                            elif len(filtered_curr_b) < len(res_idx):
                                res_idx = filtered_curr_b

                        if not b_curr_diff:
                            sus_date_b = effective_date
                            if pd.notna(sus_date_b):
                                filtered_per_b = []
                                for i in res_idx:
                                    com_date_b = osbal_rows[i].get("_com_date_parsed")
                                    if pd.isna(com_date_b):
                                        filtered_per_b.append(i)
                                    else:
                                        if _is_ls:
                                            if (com_date_b.year == sus_date_b.year
                                                    and com_date_b.month == sus_date_b.month):
                                                filtered_per_b.append(i)
                                        else:
                                            if sus_date_b >= com_date_b:
                                                filtered_per_b.append(i)
                                if not filtered_per_b:
                                    b_per_diff = True
                                elif len(filtered_per_b) < len(res_idx):
                                    res_idx = filtered_per_b

                        if not b_curr_diff and not b_per_diff and res_idx:
                            is_override = bool(matched_osbal_idx and set(res_idx) != matched_osbal_idx)
                            matched_osbal_idx = set(res_idx)
                            source   = "OSBAL"
                            if is_override:
                                scenario = "Polis only -> Bordero OC Override"
                            else:
                                scenario = "Polis only"
                            is_bordero = True
                            resolved   = True
                            is_beda_curr = False
                            is_beda_per  = False
                        elif b_curr_diff and len(bordero_facs) == 1 and res_idx:
                            is_override = bool(matched_osbal_idx and set(res_idx) != matched_osbal_idx)
                            if is_override or not matched_osbal_idx:
                                matched_osbal_idx = set(res_idx)
                                source   = "OSBAL"
                                scenario = "Polis only -> Bordero OC Override"
                                is_bordero = True
                                resolved   = True
                                is_beda_curr = True
                                is_beda_per  = False

        # Post-matching label normalisation
        if is_beda_curr or is_beda_per:
            osbal_ref = [osbal_rows[i] for i in list(matched_osbal_idx)] if matched_osbal_idx else [{}]
            osbal_count = 0
            resolved = False
        elif not matched_osbal_idx:
            source = None
            scenario = "Unmatching"
            osbal_ref = [{}]
            osbal_count = 0
            resolved = False
        else:
            osbal_ref = [osbal_rows[i] for i in matched_osbal_idx]
            osbal_count = len(osbal_ref)

        if scenario and " only + " in scenario:
            scenario = scenario.replace(" only + ", " + ")

        fac_codes = {_normalize(r.get(OSBAL_FACODE_COL, "")) for r in osbal_ref if r}
        fac_codes.discard("")

        raw_results.append({
            "suspend": sus, "source": source, "scenario": scenario,
            "base_scenario": scenario,
            "osbal_idx": list(matched_osbal_idx),
            "osbal_count": osbal_count, "resolved": resolved, "fac_codes": fac_codes,
            "_like_log": [m for _, m in _like_log],
            "is_beda_curr": is_beda_curr,
            "is_beda_per": is_beda_per,
            "is_bordero": is_bordero,
            "_orig_idx": n,
        })

    print(f"  Matching selesai dalam {time.perf_counter()-t4:.1f}s", flush=True)
    total_like = 0
    for _row in raw_results:
        for _method in _row.get("_like_log", []):
            if _method == "SUBSTR":  stat_like_substr  += 1
            elif _method == "SUBSET":  stat_like_subset  += 1
            elif _method == "OVERLAP": stat_like_overlap += 1
        total_like += len(_row.get("_like_log", []))

    # [DEBUG] Analisis penyebab Matching >1 fac code
    gt1_fac_rows = [r for r in raw_results if len(r.get("fac_codes", set())) > 1]
    print(f"\n  [DEBUG] Matching >1 fac code (sebelum fac-claim): {len(gt1_fac_rows):,} baris")
    if gt1_fac_rows:
        # Breakdown per skenario
        from collections import Counter
        scen_counts = Counter(r["scenario"] for r in gt1_fac_rows)
        print(f"  Distribusi skenario penyebab >1 fac code:")
        for scen, cnt in scen_counts.most_common():
            print(f"    {scen:<40}: {cnt:,}")
        # Berapa FAC code rata-rata
        avg_fac = sum(len(r["fac_codes"]) for r in gt1_fac_rows) / len(gt1_fac_rows)
        max_fac = max(len(r["fac_codes"]) for r in gt1_fac_rows)
        print(f"  Rata-rata FAC code per baris >1: {avg_fac:.1f} | Maks: {max_fac}")
        # Penyebab: apakah banyak dari OSBAL saja, atau cross-tabel?
        osbal_only = sum(1 for r in gt1_fac_rows if r.get("source") == "OSBAL" and not r.get("resolved", False))
        print(f"  Baris >1 fac dari OSBAL saja (polis/insured ambiguous): {osbal_only:,}")
    print(f"  [DEBUG] _like_match total hits    : {total_like}")
    print(f"    SUBSTR  (substring lama)         : {stat_like_substr}")
    print(f"    SUBSET  (token subset)           : {stat_like_subset}")
    print(f"    OVERLAP (overlap coeff >={_INSURED_LIKE_THRESHOLD:.0%}): {stat_like_overlap}")

    # [4c] Fac code claim & osbal accumulation
    print("\n[4c/6] Fac code claim & osbal accumulation (row asli dipertahankan) ...", flush=True)

    claimed: set = set()
    _EXCL = {"Unmatching", "Beda Periode", "UNMATCHED", "Beda Currency"}
    for r in raw_results:
        if len(r["fac_codes"]) == 1 and not any(exc in r["scenario"] for exc in _EXCL) and not r.get("is_beda_curr", False):
            claimed |= r["fac_codes"]

    for r in raw_results:
        if len(r["fac_codes"]) > 1:
            remaining = r["fac_codes"] - claimed
            if len(remaining) == 1 and not any(exc in r["scenario"] for exc in _EXCL) and not r.get("is_beda_curr", False):
                r["_originally_multi_fac"] = True
                claimed |= remaining

    groups: dict = {}
    final:  list = []

    for r in raw_results:
        r.setdefault("suspend_count", 1)
        _is_beda_curr_row = r.get("is_beda_curr", False) or "Beda Currency" in str(r.get("scenario", ""))
        if (len(r["fac_codes"]) == 1 and r["source"] is not None
                and not any(exc in r["scenario"] for exc in _EXCL)
                and not _is_beda_curr_row):
            groups.setdefault(next(iter(r["fac_codes"])), []).append(r)
        else:
            final.append(r)

    for fac, group in groups.items():
        if len(group) == 1:
            group[0]["suspend_count"] = 1
            final.append(group[0])
            continue
        merged_osbal_idx = set()
        for g in group:
            merged_osbal_idx.update(g["osbal_idx"])
        merged_sus = _merge_suspend_rows([g["suspend"] for g in group])
        for g in group:
            g["osbal_idx"]          = list(merged_osbal_idx)
            g["osbal_count"]        = len(merged_osbal_idx)
            g["suspend_count"]      = len(group)
            g["_merged_amount_ori"] = merged_sus.get("AMOUNT ORI", 0)
            final.append(g)

    final.sort(key=lambda r: r.get("_orig_idx", 0))
    print(f"  {len(raw_results):,} suspend rows -> {len(final):,} output rows (row asli dipertahankan)", flush=True)

    # [5] Build output DataFrame
    print(f"\n[5/6] Building output ({len(final):,} rows) ...", flush=True)
    output_rows = []
    for r in final:
        ref_osbal = [osbal_rows[i] for i in r["osbal_idx"]] if r["source"] else [{}]
        all_fac_osbal_rows = None
        osbal_cnt = r["osbal_count"]
        if r["source"] and len(r["fac_codes"]) == 1:
            fac = next(iter(r["fac_codes"]))
            all_fac_idx = facode_osbal_idx.get(fac, [])
            if len(all_fac_idx) > 0:
                all_fac_osbal_rows = [osbal_rows[i] for i in all_fac_idx]
                osbal_cnt = len(all_fac_idx)

        final_scenario = _format_final_scenario(
            base_scenario=r.get("base_scenario", r.get("scenario", "")),
            source=r["source"],
            is_bordero=r.get("is_bordero", False),
            is_beda_curr=r.get("is_beda_curr", False),
            is_beda_per=r.get("is_beda_per", False),
            num_fac=len(r["fac_codes"]),
        )

        out = _build_output_row(
            r["suspend"], source=r["source"], scenario=final_scenario,
            osbal_rows=ref_osbal, osbal_count=osbal_cnt,
            resolved=r["resolved"], suspend_count=r.get("suspend_count", 1),
            all_fac_osbal_rows=all_fac_osbal_rows,
        )
        if "_merged_amount_ori" in r:
            out["_MERGED_AMOUNT_ORI"] = r["_merged_amount_ori"]
        out["_IS_BEDA_CURR"] = r.get("is_beda_curr", False)
        output_rows.append(out)

    df = pd.DataFrame(output_rows)

    # Fix format angka koma-desimal
    for col in ["AMOUNT ORI", "CCOS_OR_BAL", "CCOS_BAL_DUE"]:
        if col in df.columns:
            s = df[col]
            if s.dtype == object:
                s = s.astype(str).str.strip().str.replace(r'\s+', '', regex=True)
                has_comma = s.str.contains(',', na=False)
                s_eu = s.str.replace('.', '', regex=False).str.replace(',', '.', regex=False)
                s = s_eu.where(has_comma, s)
            df[col] = pd.to_numeric(s, errors="coerce")

    df = _compute_derived_cols(df)
    df.loc[df["SKENARIO"].isin(["UNMATCHED", "Unmatching"]), "FLAG_PROD"] = "Unmatching"
    df.loc[df["SKENARIO"] == "UNMATCHED", "SKENARIO"] = "Unmatching"
    df.loc[df["SKENARIO"] == "INSURED_LIKE", "SKENARIO"] = "Insured only"
    df.loc[df["SKENARIO"] == "POLIS_LIKE", "SKENARIO"] = "Polis only"
    df.loc[df["SKENARIO"] == "SLIP_LIKE", "SKENARIO"] = "Slip only"

    # Kumpulkan semua kolom clean slip yang terdeteksi
    slip_cols_in_df = [c for c in df.columns if str(c).upper().startswith("SLIP_NO_CLN") or str(c).upper().startswith("CLEAN SLIP")]
    global FINAL_COLUMNS
    if "SLIP_NO_ORI" in FINAL_COLUMNS:
        idx = FINAL_COLUMNS.index("SLIP_NO_ORI") + 1
        if "SLIP_NO_CLN" in FINAL_COLUMNS:
            FINAL_COLUMNS.remove("SLIP_NO_CLN")
        for sc in reversed(slip_cols_in_df):
            if sc not in FINAL_COLUMNS:
                FINAL_COLUMNS.insert(idx, sc)

    for col in FINAL_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    df = df[FINAL_COLUMNS]

    # [6] Export
    print(f"\n[6/6] Saving to: {OUTPUT_FILE} ...", flush=True)
    for c in df.select_dtypes(include=['object', 'string']).columns:
        df[c] = df[c].astype(str).str.replace(r'[\x00-\x08\x0B\x0C\x0E-\x1F]', '', regex=True)
    target_out = OUTPUT_FILE
    try:
        df.to_excel(target_out, index=False)
    except PermissionError:
        print(f"\n[PERINGATAN] File '{OUTPUT_FILE}' terkunci karena sedang dibuka di Microsoft Excel.", flush=True)
        target_out = os.path.join(_BASE_DIR, "matching_msig_v1_new.xlsx")
        try:
            df.to_excel(target_out, index=False)
            print(f"Hasil berhasil disimpan ke file alternatif: {target_out}", flush=True)
        except PermissionError:
            target_out = os.path.join(_BASE_DIR, "matching_msig_v1_updated.xlsx")
            df.to_excel(target_out, index=False)
            print(f"Hasil berhasil disimpan ke file alternatif: {target_out}", flush=True)
        print("Silakan tutup Excel jika ingin menimpa file utama matching_msig_v1.xlsx.\n", flush=True)

    try:
        from openpyxl import load_workbook
        workbook = load_workbook(target_out)
        worksheet = workbook.active
        text_cols = [
            "POLIS_ORI", "POLICY_CLEAN_1", "SERTIF_CLEAN_1",
            "POLICY_CLEAN_2", "SERTIF_CLEAN_2",
            "POLICY_CLEAN_3", "SERTIF_CLEAN_3",
            "POLICY_CLEAN_4", "POLICY_CLEAN_5",
            "SLIP_NO_ORI", 
            "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO"
        ] + slip_cols_in_df
        text_col_indices = [idx for idx, col in enumerate(df.columns, start=1) if col in text_cols]
        for col_idx in text_col_indices:
            for row in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=col_idx)
                cell.number_format = "@"
                if cell.value is not None:
                    cell.value = str(cell.value)
        workbook.save(target_out)
    except Exception as e:
        print(f"  [WARN] Format text gagal: {e}")

    try:
        # pyrefly: ignore [missing-import]
        from excel_styler import apply_purple_column_style
        apply_purple_column_style(target_out, "matching")
    except Exception:
        pass

    # Export ke PostgreSQL (opsional)
    try:
        load_dotenv()
        conn_str = "postgresql+psycopg2://{}:{}@{}:{}/{}".format(
            os.environ.get("DB_USER", "postgres"), os.environ.get("DB_PASS", "postgres"),
            os.environ.get("DB_HOST", "localhost"), os.environ.get("DB_PORT", "5432"),
            os.environ.get("DB_NAME", "postgres"),
        )
        engine = create_engine(conn_str)
        df.to_sql("SUSPENSE_DATA_SUSPENSE_MSIG", con=engine, if_exists="append", index=False)
        print(f"      -> {len(df):,} rows berhasil di-export ke 'SUSPENSE_DATA_SUSPENSE_MSIG'")
    except Exception:
        pass

    elapsed = time.perf_counter() - t_start
    print(f"\n{'=' * 70}")
    print(f"  PRODUCTION SCRIPT -- SUSPEND MATCHING  (MSIG INDONESIA) -- SELESAI")
    print(f"{'=' * 70}")
    print(f"  Selesai dalam {elapsed:.1f}s ({elapsed/60:.1f} menit)")
    print(f"  Rows  : {len(df):,}  |  Cols: {len(df.columns)}  |  File: {OUTPUT_FILE}")
    print(f"\n  FLAG_PROD:")
    for flag, count in df["FLAG_PROD"].value_counts().items():
        print(f"    {flag:<50}: {count:,}")
    print(f"\n  SKENARIO:")
    for sce, count in df["SKENARIO"].value_counts().items():
        print(f"    {sce:<50}: {count:,}")

    # Ringkasan info >1 fac code
    gt1 = df[df["FLAG_PROD"] == "Matching >1 fac code"]
    if not gt1.empty:
        print(f"\n  [INFO] 'Matching >1 fac code' = {len(gt1):,} baris ({len(gt1)/len(df)*100:.1f}%)")
        print(f"  Kemungkinan penyebab:")
        print(f"    1. Insured name terlalu generik -> match ke banyak OSBAL (FAC berbeda)")
        print(f"    2. Polis di OSBAL ada di beberapa FAC_CODE sekaligus (reuse polis)")
        print(f"    3. Cek kolom CCOS_REF_CODE yang berisi koma untuk validasi manual")
    print("=" * 70)


if __name__ == "__main__":
    run()