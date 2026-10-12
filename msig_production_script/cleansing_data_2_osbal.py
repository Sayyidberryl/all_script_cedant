"""
================================================================================
 SCRIPT CLEANSING DATA 2 - OSBAL (CEDANT: MSIG )
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI UTAMA
# ==============================================================================
INPUT_FILE = os.path.join("data_raw", "2b. transaksi Osbal.xlsx")
SHEET_NAME = "Query result"
HEADER_ROW = 0
OUTPUT_FILE = os.path.join("data_clean", "msig_output_osballl.xlsx")
CEDANT_FILTER = "MSIG"

CEDANT_COLUMN = "CCOS_COMP_NAME"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"

CLSDT_POLIS_COLUMN = "CLSDT_POLICY_NO"
CLSDT_SLIP_COLUMN = "CLSDT_SLIP_NO"
CLSDT_SERTF_COLUMN = "CLSDT_SERTF_NO"
CERTIFICATE_OUTPUT_COL = "CERTIFICATE"

MAX_BREAKDOWN_CODES = 5
TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CLSDT", "CERTIFICATE",
)

CERT_DIGIT_LEN = 6
CERT_RANGE_JOIN_THRESHOLD = 3
MSIG_INDEPENDENT_MIN_LEN = 8
MSIG_SUFFIX_BASE_MIN_LEN = 6


# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
_BYPASS_ALL_RE = re.compile(r"\bPENYELESAIAN\s+HUTANG\s+PIUTANG\b", flags=re.IGNORECASE)
_TRANSACTION_HEADER_RE = re.compile(
    r"^\s*(?:Pembayaran|Penagihan)?\s*(?:dan\s+Penagihan)?\s*Premi\s*(?:Fakultatif|SOA|Pensesian)?\s*(?:PAR\s*&?\s*EQ|BIT|GIT|CECR|QS)?\s*",
    flags=re.IGNORECASE,
)
_TRUNCATE_TRIGGER_RE = re.compile(
    r"\bsubsidiar|\bassociat|\baffiliat|\bfiliated\b|related\s+compan|respective\s+right|\bincluding\s+an[yd]|\bindu?ding\s+any|as\s+the\s+owner|as\s+owners?\b|as\s+main\s+contractor|as\s+contractor\b|\(as\s+mortgagee\s+and\s+loss\s+payee\)",
    flags=re.IGNORECASE,
)
_REMOVE_ONLY_RE = re.compile(r"\bas\s+principal\b|\bsebagai\s+principal\b|\bsebagai\s+kontraktor\b", flags=re.IGNORECASE)
_ONLY_WORD_RE = re.compile(r"\bONLY\b", flags=re.IGNORECASE)
_GENERIC_BOILERPLATE_TAIL_RE = re.compile(
    r"\bsub\s*-?\s*kontraktor\b.*$|\banak\s+perusahaan\b.*$|\bafiliasi\s+perusahaan\b.*$|\bdan\s+semua\b\s*[–-]?\s*$",
    flags=re.IGNORECASE,
)
_TITLE_RE = re.compile(
    r"\b(?:S\.?H\.?|S\.?I\.?Kom\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b",
    flags=re.IGNORECASE,
)
_LEGAL_ENTITY_RE = re.compile(
    r"\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|CO\.?|PERSERO|Perseroan\s+Terbatas)\b",
    flags=re.IGNORECASE,
)
_ENTITY_SPLIT_RE = re.compile(
    r"\b(?:QQ|Q\.Q\.|CQ|C\.Q\.?)\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|\bin\s+this\s+case\b|\bdalam\s+hal\s+ini\b|,|/",
    flags=re.IGNORECASE,
)
_LEGAL_ENTITY_DASH_SIGNAL_RE = re.compile(
    r",\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO)\s*[\r\n]*\s*-\s*",
    flags=re.IGNORECASE,
)
_GENERIC_DASH_SEPARATOR_RE = re.compile(r"\s*[\r\n]+\s*-\s*|\s+-\s+")
_TRAILING_PAREN_MERGE_RE = re.compile(r",?\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?)?\s*\(([^)]+)\)\s*$", flags=re.IGNORECASE)
_INLINE_PAREN_MERGE_RE = re.compile(r",?\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?)?\s*\(([^)]+)\)", flags=re.IGNORECASE)
_QQ_PAREN_UNWRAP_RE = re.compile(r"\bQQ\s*\(\s*([^)]*)\)", flags=re.IGNORECASE)

_MONTH_WORDS = (
    "JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI", "AGUSTUS",
    "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER", "DES", "JANUARY", "FEBRUARY",
    "MARCH", "MAY", "JUNE", "JULY", "AUGUST", "OCTOBER", "DECEMBER", "JAN", "FEB",
    "MAR", "APR", "JUN", "JUL", "AUG", "SEP", "SEPT", "OCT", "NOV", "DEC",
)
_DATE_BATCH_TOKEN_RE = re.compile(
    r"\b(?:" + "|".join(_MONTH_WORDS) + r")\b|\b(?:19|20)\d{2}\b|\bBATCH\s*\d*\b",
    flags=re.IGNORECASE,
)

_PAITON_DASH_PROTECT_RE = re.compile(r"\bPAITON\s+(\d+)-(\d+)\b", flags=re.IGNORECASE)
_SPC_PAREN_PROTECT_RE = re.compile(r"\(\s*SPC\s*\)", flags=re.IGNORECASE)
_INSURED_DESCRIPTIVE_RE = re.compile(r"\bVARIOUS\s+INSUREDS?\b|\bACCEPTED\s+BY\b", flags=re.IGNORECASE)
_SUSPENSE_WORD_RE = re.compile(r"\bSUSPENSE\b", flags=re.IGNORECASE)

_INSURED_DASH_SPLIT_EXCEPTIONS = (
    "EBDI VAUNDRI - OSCAR OMEGA",
    "ADHI - HUTAMA - NINDYA - ABIPRAYA KSO",
    "PEMERINTAH PROVINSI SULAWESI SELATAN - DINAS BINA MARGA DAN BINA KONSTRUKSI",
)

_INSURED_DASH_MERGE_SUBSTRINGS = (
    "PP - WASKITA - WIJAYA KARYA",
)

_ROMAN_TO_ARABIC = {
    "I": 1, "II": 2, "III": 3, "IV": 4, "V": 5,
    "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10,
}
_ASTRA_DAIHATSU_RE = re.compile(
    r"^\s*ASTRA\s+DAIHATSU\s+MOTOR\s*,?\s*PT\s*\(([^)]+)\)\s*,?\s*(.+)$",
    flags=re.IGNORECASE,
)
_MSIG_GROUP_SUSPENSE_RE = re.compile(r"^\s*MSIG\s+GROUP\s+2016\b", flags=re.IGNORECASE)
_MSIG_CURRENCY_STRIP_RE = re.compile(r"\b(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\b", flags=re.IGNORECASE)

_TOYOTA_XOL_RE = re.compile(
    r"^\s*([A-Z0-9 .&\-]+?)\s*,\s*([A-Z0-9 .&\-]+?)\s*\(\s*([A-Z0-9]{2,10})\s*\)\s*$",
    flags=re.IGNORECASE,
)
_LEGAL_ENTITY_ONLY_RE = re.compile(
    r"^(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|CO\.?|PERSERO)$",
    flags=re.IGNORECASE,
)
_ABBREVIATION_MERGE_RE = re.compile(
    r"^\s*([A-Z0-9 .&\-]+?)\s*,\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|CO\.?|PERSERO)\s*\(\s*([A-Z0-9]{2,10})\s*\)\s*$",
    flags=re.IGNORECASE,
)
_LETTER_SUFFIX_SLASH_RE = re.compile(
    r"^\s*([A-Z][A-Z ]*?)\s+([A-Z](?:\s*(?:&|DAN)\s*[A-Z])+)\s*/\s*(.+)$",
    flags=re.IGNORECASE,
)
_NUMBER_SUFFIX_DAN_RE = re.compile(
    r"^\s*([A-Z][A-Z ]*?)\s+(\d+(?:\s*(?:&|DAN)\s*\d+)+)\s*$",
    flags=re.IGNORECASE,
)
_PAITON_UNIT_RANGE_RE = re.compile(
    r"^\s*(.*?)\s+UNIT\s+(\d+)\s*-\s*(\d+)\s*,?\s*PT\.?\s*$",
    flags=re.IGNORECASE,
)
_PAITON_NAME_RANGE_RE = re.compile(
    r"^\s*([A-Z0-9 .&\-]+?)\s*,\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?)\s+([A-Z0-9 .&\-]*?[A-Za-z])\s+(\d+)\s*-\s*(\d+)\s*$",
    flags=re.IGNORECASE,
)

_KNOWN_ENTITIES = []
_KNOWN_ENTITIES_SORTED = sorted(_KNOWN_ENTITIES, key=len, reverse=True)
_KNOWN_ENTITY_RE = re.compile("|".join(re.escape(name) for name in _KNOWN_ENTITIES_SORTED)) if _KNOWN_ENTITIES_SORTED else re.compile(r"(?!)")
_ENTITY_SLASH_NORMALIZE_PATTERNS = [
    (name, re.compile(r"\b" + r"(?:\s*/\s*|\s+)".join(re.escape(w) for w in name.split()) + r"\b", flags=re.IGNORECASE))
    for name in _KNOWN_ENTITIES_SORTED if len(name.split()) > 1
]


def _unwrap_qq_paren(text: str) -> str:
    def _replace(match):
        inner = re.sub(r"\s+", " ", re.sub(r"[:;]", " ", match.group(1).strip())).strip()
        return f"QQ , {inner}" if inner else "QQ"
    return _QQ_PAREN_UNWRAP_RE.sub(_replace, text)


def _merge_trailing_paren(text: str) -> str:
    match = _TRAILING_PAREN_MERGE_RE.search(text)
    if not match:
        return text
    inner = match.group(1).strip()
    if re.search(r"\d", inner) or inner.strip().upper() == "PERSERODA":
        return text[: match.start()] + text[match.end():]
    if inner.strip().upper() == "SPC":
        return text
    if not inner or len(inner.split()) > 4 or _TRUNCATE_TRIGGER_RE.search(inner):
        return text[: match.start()] + text[match.end():]
    return text[: match.start()] + " " + inner


def _merge_inline_paren(text: str) -> str:
    def _replace(match):
        inner = match.group(1).strip()
        if re.search(r"\d", inner) or inner.strip().upper() == "PERSERODA":
            return ""
        if inner.strip().upper() == "SPC":
            return match.group(0)
        return "" if not inner or len(inner.split()) > 4 or _TRUNCATE_TRIGGER_RE.search(inner) else f" {inner}"
    return _INLINE_PAREN_MERGE_RE.sub(_replace, text)


def _strip_date_batch_info(text: str) -> str:
    return _DATE_BATCH_TOKEN_RE.sub(" ", text)


def _truncate_from_first_trigger(text: str) -> str:
    match = _TRUNCATE_TRIGGER_RE.search(text)
    return text[: match.start()] if match else text


def _normalize_dash_entity_list(text: str) -> str:
    return _GENERIC_DASH_SEPARATOR_RE.sub(", ", text) if _LEGAL_ENTITY_DASH_SIGNAL_RE.search(text) else text


def _strip_transaction_header(text: str) -> str:
    return _TRANSACTION_HEADER_RE.sub("", text)


def _strip_titles(text: str) -> str:
    return _TITLE_RE.sub(" ", text)


def _strip_legal_entity(text: str) -> str:
    return _LEGAL_ENTITY_RE.sub(" ", text)


def _truncate_boilerplate_tail(text: str) -> str:
    text = _truncate_from_first_trigger(text)
    match_generic = _GENERIC_BOILERPLATE_TAIL_RE.search(text)
    if match_generic:
        text = text[: match_generic.start()]
    return _REMOVE_ONLY_RE.sub(" ", text)


def _final_polish(text: str) -> str:
    protected = _PAITON_DASH_PROTECT_RE.sub(lambda m: f"PAITON {m.group(1)}\x00{m.group(2)}", text)
    protected = _SPC_PAREN_PROTECT_RE.sub("\x01SPC\x02", protected)
    t = re.sub(r"\s+", " ", re.sub(r"[-/]", " ", re.sub(r"[().;:\"']", " ", protected))).strip()
    t = t.replace("\x00", "-").replace("\x01", "(").replace("\x02", ")")

    dangling_pattern = r"^\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*|\s*\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$"
    previous = None
    while previous != t:
        previous, t = t, re.sub(dangling_pattern, "", t, flags=re.IGNORECASE).strip()
    return t.upper()


def _normalize_entity_slash_variants(text: str) -> str:
    for canonical, pattern in _ENTITY_SLASH_NORMALIZE_PATTERNS:
        text = pattern.sub(canonical, text)
    return text


def _split_by_known_entities(text: str):
    pos, n, found = 0, len(text), []
    while pos < n:
        while pos < n and text[pos] == " ":
            pos += 1
        if pos >= n:
            break
        match = _KNOWN_ENTITY_RE.match(text, pos)
        if not match:
            return None
        found.append(match.group(0))
        pos = match.end()
    return found if len(found) > 1 else None


def _apply_dash_merge_substrings(text: str) -> str:
    for original_sub in _INSURED_DASH_MERGE_SUBSTRINGS:
        if original_sub.upper() in text.upper():
            merged = re.sub(r"\s*-\s*", " ", original_sub)
            idx = text.upper().find(original_sub.upper())
            if idx != -1:
                text = text[:idx] + merged + text[idx + len(original_sub):]
    return text


def _try_split_insured_by_dash(text: str):
    stripped = text.strip()
    if _SUSPENSE_WORD_RE.search(stripped):
        return None
    for exception in _INSURED_DASH_SPLIT_EXCEPTIONS:
        if stripped.upper() == exception.upper():
            return [p.strip() for p in re.split(r"\s*-\s*", stripped) if p.strip()]
    return None


def _try_handle_astra_daihatsu(text: str):
    match = _ASTRA_DAIHATSU_RE.match(text)
    if not match:
        return None
    ho_label, roman_part = match.group(1).strip(), match.group(2).strip()
    romans = re.split(r"\s*(?:&|,|DAN)\s*", roman_part, flags=re.IGNORECASE)
    results = []
    for r in romans:
        r_clean = r.strip().upper()
        if r_clean in _ROMAN_TO_ARABIC:
            val = f"ASTRA DAIHATSU MOTOR ({ho_label}) {_ROMAN_TO_ARABIC[r_clean]}"
            if val not in results:
                results.append(val)
    return results if results else None


def _try_handle_msig_group_suspense(text: str):
    if not _MSIG_GROUP_SUSPENSE_RE.match(text) or re.search(r"PENYELESAIAN\s+SUSPENSE", text, flags=re.IGNORECASE):
        return None
    cleaned = _MSIG_CURRENCY_STRIP_RE.sub("", text)
    cleaned = re.sub(r"\s{2,}", " ", cleaned).strip()
    return [f"{cleaned} PENYELESAIAN SUSPENSE"]


def _try_handle_toyota_xol_pattern(text: str):
    match = _TOYOTA_XOL_RE.match(text)
    if not match:
        return None
    name1, name2, code = match.group(1).strip(), match.group(2).strip(), match.group(3).strip()
    if not re.search(r"[A-Za-z]", code) or _LEGAL_ENTITY_ONLY_RE.match(name2):
        return None
    results = []
    for part in (name1, name2, code):
        cleaned = _final_polish(part)
        if cleaned and cleaned not in results:
            results.append(cleaned)
    return results if len(results) > 1 else None


def _try_handle_abbreviation_merge(text: str):
    match = _ABBREVIATION_MERGE_RE.match(text)
    if not match:
        return None
    name, abbrev = match.group(1).strip(), match.group(2).strip()
    if not re.search(r"[A-Za-z]", abbrev):
        return None
    merged = _final_polish(f"{name} {abbrev}")
    return [merged] if merged else None


def _try_handle_letter_suffix_slash(text: str):
    match = _LETTER_SUFFIX_SLASH_RE.match(text)
    if not match:
        return None
    base_name, letters_part, other_entity = match.group(1).strip(), match.group(2).strip(), match.group(3).strip()
    letters = [l.strip().upper() for l in re.split(r"\s*(?:&|DAN)\s*", letters_part, flags=re.IGNORECASE) if l.strip()]
    if not letters or any(len(l) != 1 or not l.isalpha() for l in letters):
        return None
    results = []
    for letter in letters:
        val = _final_polish(f"{base_name} {letter}")
        if val not in results:
            results.append(val)
    other_cleaned = _final_polish(other_entity)
    if other_cleaned and other_cleaned not in results:
        results.append(other_cleaned)
    return results if len(results) > 1 else None


def _try_handle_number_suffix_dan(text: str):
    match = _NUMBER_SUFFIX_DAN_RE.match(text)
    if not match:
        return None
    base_name, numbers_part = match.group(1).strip(), match.group(2).strip()
    numbers = [n.strip() for n in re.split(r"\s*(?:&|DAN)\s*", numbers_part, flags=re.IGNORECASE) if n.strip()]
    if not numbers or any(not n.isdigit() for n in numbers):
        return None
    results = []
    for n in numbers:
        val = _final_polish(f"{base_name} {n}")
        if val not in results:
            results.append(val)
    return results if len(results) > 1 else None


def _try_handle_paiton_unit_range(text: str):
    match = _PAITON_UNIT_RANGE_RE.match(text)
    if not match:
        return None
    prefix, n1, n2 = match.group(1).strip(), match.group(2), match.group(3)
    val1 = _final_polish(f"{prefix} UNIT {n1}")
    val2 = _final_polish(f"{prefix} {n2}")
    results = []
    for v in (val1, val2):
        if v not in results:
            results.append(v)
    return results if len(results) > 1 else None


def _try_handle_paiton_name_range(text: str):
    match = _PAITON_NAME_RANGE_RE.match(text)
    if not match:
        return None
    name1, name2, n1, n2 = match.group(1).strip(), match.group(2).strip(), match.group(3), match.group(4)
    val1 = _final_polish(f"{name1} {name2} {n1}")
    val2 = _final_polish(f"{name1} {name2} {n2}")
    results = []
    for v in (val1, val2):
        if v not in results:
            results.append(v)
    return results if len(results) > 1 else None


def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []
    original = text.strip()

    # Case khusus
    original = re.sub(r"FUJISEAT", "FUJI SEAT", original, flags=re.IGNORECASE)
    original = re.sub(r"MULIAGLASS", "MULIA GLASS", original, flags=re.IGNORECASE)

    astra_res = _try_handle_astra_daihatsu(original)
    if astra_res is not None:
        return astra_res

    msig_res = _try_handle_msig_group_suspense(original)
    if msig_res is not None:
        return msig_res

    toyota_res = _try_handle_toyota_xol_pattern(original)
    if toyota_res is not None:
        return toyota_res

    abbreviation_res = _try_handle_abbreviation_merge(original)
    if abbreviation_res is not None:
        return abbreviation_res

    letter_res = _try_handle_letter_suffix_slash(original)
    if letter_res is not None:
        return letter_res

    number_res = _try_handle_number_suffix_dan(original)
    if number_res is not None:
        return number_res

    paiton_unit_res = _try_handle_paiton_unit_range(original)
    if paiton_unit_res is not None:
        return paiton_unit_res

    paiton_name_res = _try_handle_paiton_name_range(original)
    if paiton_name_res is not None:
        return paiton_name_res

    if _BYPASS_ALL_RE.search(original) or _INSURED_DESCRIPTIVE_RE.search(original) or _SUSPENSE_WORD_RE.search(original):
        return [original]

    dash_split_res = _try_split_insured_by_dash(original)
    if dash_split_res is not None:
        results = []
        for part in dash_split_res:
            cleaned = _final_polish(_strip_date_batch_info(_strip_legal_entity(_truncate_boilerplate_tail(part))))
            if cleaned and cleaned not in results:
                results.append(cleaned)
        return [original] if len(results) > MAX_BREAKDOWN_CODES or not results else results

    original_for_processing = _apply_dash_merge_substrings(original)
    t = _truncate_boilerplate_tail(
        _normalize_entity_slash_variants(
            _normalize_dash_entity_list(
                _strip_titles(
                    _strip_transaction_header(
                        _merge_inline_paren(
                            _merge_trailing_paren(
                                _unwrap_qq_paren(_ONLY_WORD_RE.sub(" ", original_for_processing))
                            )
                        )
                    )
                )
            )
        )
    )

    raw_parts = _ENTITY_SPLIT_RE.split(t)
    results = []
    for part in raw_parts:
        cleaned = _final_polish(_strip_date_batch_info(_strip_legal_entity(_truncate_boilerplate_tail(part))))
        if not cleaned or len(cleaned) < 2 or not re.search(r"[A-Z]", cleaned):
            continue

        known_split = _split_by_known_entities(cleaned)
        if known_split is not None:
            for name in known_split:
                if name not in results:
                    results.append(name)
            continue
        if cleaned not in results:
            results.append(cleaned)

    return [original] if len(results) > MAX_BREAKDOWN_CODES else results


# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO
# ==============================================================================
_MSIG_NOISE_WORDS = {"VAR", "VARIOUS", "TPI", "P1", "P2", "P2B", "P3", "P4", "P5", "MD", "TBA"}
_MSIG_SEPARATOR_SPLIT_RE = re.compile(r"[&/+,\-]|\s+")
_MSIG_LETTER_DIGIT_RE = re.compile(r"^[A-Za-z]+\d+$")
_MSIG_LEADING_LETTERS_RE = re.compile(r"^[A-Za-z]+")
_POLIS_THOUSAND_DOT_RE = re.compile(r"^\s*\d+(?:\.\d+){2,}\s*$")

_MSIG_MONTH_WORDS_ID = (
    "JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI",
    "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER",
)
_MSIG_CURRENCY_WORDS = ("IDR", "USD", "SGD", "EUR", "JPY", "GBP", "AUD", "MYR")

_MSIG_SLIP_LABEL_RE = re.compile(
    r"\bACCOUNT\s*NUMBER\s*:|\bACCOUNT\s*NUMNER\s*:|\bACCOUNT\s*NO\s*:|\bACC\s*NO\s*:|\bTREATY\s*NUMBER\s*:",
    flags=re.IGNORECASE,
)
_MSIG_MONTH_YEAR_RE = re.compile(r"\s*/?\s*(?:" + "|".join(_MSIG_MONTH_WORDS_ID) + r")\s+\d{4}\s*", flags=re.IGNORECASE)
_MSIG_CURRENCY_RE = re.compile(r"\b(?:" + "|".join(_MSIG_CURRENCY_WORDS) + r")\b", flags=re.IGNORECASE)
_MSIG_BORDERO_WORD_RE = re.compile(r"\bBORDERO\b", flags=re.IGNORECASE)
_MSIG_DASH_VARIANTS_RE = re.compile(r"[\u2012\u2013\u2014\u2015]")
_MSIG_SD_WORD_RE = re.compile(r"\bS\s*/\s*D\b|\bSD\b|(?<=\d)SD(?=\d)", flags=re.IGNORECASE)

_MSIG_SLIP_BORD_TOKEN_RE = re.compile(r"\bBORD\.?\s*\d{4}\b", flags=re.IGNORECASE)
_MSIG_SLIP_MONTH_ABBR_YEAR_RE = re.compile(
    r"\b(?:JAN|FEB|MAR|APR|MEI|MAY|JUN|JUL|AGU|AUG|SEP|OKT|OCT|NOV|DES|DEC)\.?\s*\d{2,4}\b",
    flags=re.IGNORECASE,
)
_MSIG_SLIP_TRAILING_CODE_RE = re.compile(r"\b(?:XGEN|XOL|BIT|GIT|CECR|QS)\b", flags=re.IGNORECASE)
_MSIG_SLIP_MONTHNUM_YEAR_RE = re.compile(r"\b\d{1,2}\s+(?:19|20)\d{2}\b")
_MSIG_SLIP_YEAR_MONTHNUM_RE = re.compile(r"\b(?:19|20)\d{2}\s+\d{1,2}\b")
_MSIG_SLIP_DASH_PAIR_RE = re.compile(r"^(\d{6,})-(\d{6,})$")
_MSIG_DOT_SUFFIX_RE = re.compile(r"^(\d+)\.(\d+)$")

# Aturan suffix 5-digit: ganti 5 digit terakhir bagian tengah base polis
_MSIG_POLIS_SHORT_SUFFIX_LEN = 5
_MSIG_POLIS_TAIL_FIXED_LEN = 6

# Hardcase spesifik base + suffix polis
_MSIG_POLIS_SHORT_SUFFIX_HARDCASE = {
    ("41701296000100", "70029"): "41700296000100",
}


def _try_msig_polis_short_suffix_replace(base_digit_part: str, suffix: str):
    if len(suffix) != _MSIG_POLIS_SHORT_SUFFIX_LEN:
        return None

    hardcase_result = _MSIG_POLIS_SHORT_SUFFIX_HARDCASE.get((base_digit_part, suffix))
    if hardcase_result is not None:
        return hardcase_result

    min_len_needed = _MSIG_POLIS_TAIL_FIXED_LEN + _MSIG_POLIS_SHORT_SUFFIX_LEN
    if len(base_digit_part) < min_len_needed:
        return None
    tail = base_digit_part[-_MSIG_POLIS_TAIL_FIXED_LEN:]
    tengah = base_digit_part[:-_MSIG_POLIS_TAIL_FIXED_LEN]
    new_tengah = tengah[:-_MSIG_POLIS_SHORT_SUFFIX_LEN] + suffix
    return new_tengah + tail



def _strip_polis_thousand_dots(text: str) -> str:
    return text.replace(".", "") if _POLIS_THOUSAND_DOT_RE.match(text) else text


def _strip_msig_slip_noise(text: str) -> str:
    t = _MSIG_SLIP_BORD_TOKEN_RE.sub(" ", text)
    t = _MSIG_SLIP_MONTH_ABBR_YEAR_RE.sub(" ", t)
    t = _MSIG_SLIP_TRAILING_CODE_RE.sub(" ", t)
    t = _MSIG_SLIP_MONTHNUM_YEAR_RE.sub(" ", t)
    t = _MSIG_SLIP_YEAR_MONTHNUM_RE.sub(" ", t)
    return re.sub(r"\s{2,}", " ", t).strip()


def _strip_msig_noise(text: str) -> str:
    t = _MSIG_DASH_VARIANTS_RE.sub("-", text)
    t = _MSIG_SLIP_LABEL_RE.sub("", t)
    t = _MSIG_MONTH_YEAR_RE.sub(" ", t)
    t = _MSIG_BORDERO_WORD_RE.sub(" ", t)
    t = _MSIG_CURRENCY_RE.sub(" ", t)
    t = _strip_msig_slip_noise(t)
    t = re.sub(r"\s{2,}", " ", t).strip()
    return re.sub(r"^[\s/,\-]+|[\s/,\-]+$", "", t)


def _try_split_slip_dash_pair(token: str):
    match = _MSIG_SLIP_DASH_PAIR_RE.match(token)
    if not match:
        return None
    awal, akhir = match.group(1), match.group(2)
    return [awal, akhir] if len(awal) == len(akhir) else None


def _try_split_dot_suffix_replace(token: str):
    match = _MSIG_DOT_SUFFIX_RE.match(token)
    if not match:
        return None
    base, suffix = match.group(1), match.group(2)
    if len(suffix) >= len(base) or len(base) < MSIG_SUFFIX_BASE_MIN_LEN:
        return None
    return [base, base[:-len(suffix)] + suffix]


def _msig_tokenize_plain(text: str):
    stripped_token = text.strip()
    dash_pair = _try_split_slip_dash_pair(stripped_token)
    if dash_pair is not None:
        return dash_pair

    dot_suffix = _try_split_dot_suffix_replace(stripped_token)
    if dot_suffix is not None:
        return dot_suffix

    parts = [p.strip() for p in _MSIG_SEPARATOR_SPLIT_RE.split(text) if p.strip()]
    if not parts:
        return None

    codes = []
    current_base = None

    for part in parts:
        if part.upper() in _MSIG_NOISE_WORDS:
            continue

        if not part.isdigit():
            if _MSIG_LETTER_DIGIT_RE.match(part):
                current_base = part.upper()
                if current_base not in codes:
                    codes.append(current_base)
                continue
            return None

        if current_base is None:
            current_base = part
            if current_base not in codes:
                codes.append(current_base)
            continue

        base_digit_part = _MSIG_LEADING_LETTERS_RE.sub("", current_base)
        # Suffix digit 5 karakter diolah sebagai suffix-replace bagian tengah
        short_suffix_rebuilt = _try_msig_polis_short_suffix_replace(base_digit_part, part) if len(part) == _MSIG_POLIS_SHORT_SUFFIX_LEN else None
        if short_suffix_rebuilt is not None:
            prefix_match = _MSIG_LEADING_LETTERS_RE.match(current_base)
            prefix_str = prefix_match.group(0) if prefix_match else ""
            rebuilt = prefix_str + short_suffix_rebuilt
            if rebuilt not in codes:
                codes.append(rebuilt)
            continue
        if len(part) < len(base_digit_part) and len(base_digit_part) >= MSIG_SUFFIX_BASE_MIN_LEN:
            prefix_match = _MSIG_LEADING_LETTERS_RE.match(current_base)
            prefix_str = prefix_match.group(0) if prefix_match else ""
            rebuilt = prefix_str + base_digit_part[:-len(part)] + part
            if rebuilt not in codes:
                codes.append(rebuilt)
        elif len(part) >= MSIG_INDEPENDENT_MIN_LEN:
            current_base = part
            if current_base not in codes:
                codes.append(current_base)
        else:
            return None

    return codes if codes else None


def clean_split_code_impl(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return [], False
    raw_original = text.strip()

    if _BYPASS_ALL_RE.search(raw_original):
        return [raw_original], True

    thousand_dot_result = _strip_polis_thousand_dots(raw_original)
    if thousand_dot_result != raw_original:
        return [thousand_dot_result], False

    stripped = _strip_msig_noise(raw_original)
    if not stripped:
        return [raw_original], True

    sd_match = _MSIG_SD_WORD_RE.search(stripped)
    if sd_match:
        left, right = stripped[:sd_match.start()], stripped[sd_match.end():]
        left_before_codes = _msig_tokenize_plain(left) if left.strip() else []
        left_nums = re.findall(r"\d+", left)
        right_nums = re.findall(r"\d+", right)
        if not left_nums or not right_nums:
            return [raw_original], True
        start_s, end_s = left_nums[-1], right_nums[0]
        if len(start_s) != len(end_s):
            return [raw_original], True
        try:
            start_n, end_n = int(start_s), int(end_s)
        except ValueError:
            return [raw_original], True

        prior_codes = left_before_codes[:-1] if left_before_codes else []

        if start_n > end_n:
            sd_expanded = [f"{start_s} SD {end_s}"]
        else:
            count = end_n - start_n + 1
            if count <= MAX_BREAKDOWN_CODES:
                sd_expanded = [str(n).zfill(len(start_s)) for n in range(start_n, end_n + 1)]
            else:
                sd_expanded = [f"{start_s} SD {end_s}"]

        result = prior_codes + sd_expanded
        deduped = []
        for c in result:
            if c not in deduped:
                deduped.append(c)
        if len(deduped) > MAX_BREAKDOWN_CODES:
            return [raw_original], True
        return deduped, False

    codes = _msig_tokenize_plain(stripped)
    if codes is None or len(codes) > MAX_BREAKDOWN_CODES:
        return [raw_original], True
    return codes, False


def clean_split_code(text: str) -> list:
    results, _ = clean_split_code_impl(text)
    return results


# ==============================================================================
# 2B. CLEANSING CERTIFICATE
# ==============================================================================
_CERT_SD_NORMALIZE_RE = re.compile(r"\bS\s*/\s*D\b", flags=re.IGNORECASE)
_CERT_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9])(0\d{5})(?![A-Za-z0-9])")
_CERT_SD_PAIR_RE = re.compile(r"(0\d{5})\s+SD\s+(0\d{5})", flags=re.IGNORECASE)


def _pad_cert_digits(digits: str) -> str:
    return digits.zfill(CERT_DIGIT_LEN)


def _extract_cert_numbers_from_raw(text: str):
    if not text:
        return []

    found = []
    consumed_spans = []

    for m in _CERT_SD_PAIR_RE.finditer(text):
        start_s, end_s = m.group(1), m.group(2)
        before = text[: m.start(1)]
        if not re.search(r"[A-Za-z0-9]", before):
            continue
        try:
            start_n, end_n = int(start_s), int(end_s)
        except ValueError:
            continue
        if start_n > end_n or (end_n - start_n + 1) > MAX_BREAKDOWN_CODES:
            continue
        for n in range(start_n, end_n + 1):
            found.append(_pad_cert_digits(str(n)))
        consumed_spans.append((m.start(), m.end()))

    if consumed_spans:
        masked_chars = list(text)
        for s, e in consumed_spans:
            for i in range(s, e):
                masked_chars[i] = "#"
        scan_text = "".join(masked_chars)
    else:
        scan_text = text

    for m in _CERT_TOKEN_RE.finditer(scan_text):
        before = scan_text[: m.start()]
        if re.search(r"[A-Za-z0-9]", before):
            found.append(_pad_cert_digits(m.group(1)))

    return found


def _single_certificate_from_clsdt(clsdt_sertf_value) -> str:
    if clsdt_sertf_value is None or (isinstance(clsdt_sertf_value, float) and pd.isna(clsdt_sertf_value)):
        return ""

    text = str(clsdt_sertf_value).strip()
    if not text or text.lower() in ("nan", "none", "-"):
        return ""

    text = _CERT_SD_NORMALIZE_RE.sub("SD", text)
    _, is_kept_as_is = clean_split_code_impl(text)
    if not is_kept_as_is:
        return ""

    if re.fullmatch(r"0\d{5}", text):
        return _pad_cert_digits(text)

    cert_candidates = _extract_cert_numbers_from_raw(text)
    return cert_candidates[0] if cert_candidates else ""


def _join_certificates(cert_list: list) -> str:
    if not cert_list:
        return ""

    seen = []
    for c in cert_list:
        if c not in seen:
            seen.append(c)

    if len(seen) == 1:
        return seen[0]

    try:
        nums = [int(c) for c in seen]
    except ValueError:
        return ", ".join(seen)

    segments = []
    current_seg = [seen[0]]
    for i in range(1, len(seen)):
        if nums[i] - nums[i - 1] == 1:
            current_seg.append(seen[i])
        else:
            segments.append(current_seg)
            current_seg = [seen[i]]
    segments.append(current_seg)

    parts = []
    for seg in segments:
        if len(seg) > CERT_RANGE_JOIN_THRESHOLD:
            parts.append(f"{seg[0]} SD {seg[-1]}")
        else:
            parts.extend(seg)

    return ", ".join(parts)


def build_certificates_for_row(clsdt_sertf_value, fac_policy_value, num_polis: int):
    # Certificate diekstrak bila memenuhi format 6-digit angka
    slots = [""] * max(num_polis, 1)

    clsdt_cert = _single_certificate_from_clsdt(clsdt_sertf_value)
    if clsdt_cert:
        slots[0] = clsdt_cert
        return slots

    if _is_blank_value(fac_policy_value):
        return slots
    fac_text = str(fac_policy_value).strip()
    fac_text = _strip_msig_noise(fac_text)
    fac_text = _CERT_SD_NORMALIZE_RE.sub("SD", fac_text)

    single_candidates = _extract_cert_numbers_from_raw(fac_text)
    if single_candidates:
        slots[0] = _join_certificates(single_candidates)

    return slots


# ==============================================================================
# 3. RESOLVER CLSDT vs FAC (POLIS & SLIP)
# ==============================================================================
_PENYELESAIAN_RE = re.compile(r"\bPENYELESAIAN\s+(?:HUTANG|UTANG)\s+PIUTANG\b", flags=re.IGNORECASE)
_TRAILING_CURRENCY_WORD_RE = re.compile(r"\s+(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\s*$", flags=re.IGNORECASE)


def _is_blank_value(value) -> bool:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return True
    s = str(value).strip()
    return s == "" or s.lower() in ("nan", "none")


def _is_valid_clean_result(codes: list) -> bool:
    if not codes:
        return False
    for c in codes:
        if re.fullmatch(r"\d+", c):
            if len(c) >= MSIG_SUFFIX_BASE_MIN_LEN:
                continue
            return False
        if re.fullmatch(r"[A-Za-z]+\d+", c) or re.fullmatch(r"\d+\s+SD\s+\d+", c, flags=re.IGNORECASE):
            continue
        return False
    return True


def _pick_best_source_codes(clsdt_value, fac_value):
    clsdt_codes = clean_split_code(str(clsdt_value)) if not _is_blank_value(clsdt_value) else []
    if _is_valid_clean_result(clsdt_codes):
        return clsdt_codes, "CLSDT"

    fac_codes = clean_split_code(str(fac_value)) if not _is_blank_value(fac_value) else []
    if fac_codes:
        return fac_codes, "FAC"

    return clsdt_codes, "CLSDT"


def _resolve_polis_or_slip(clsdt_value, fac_value):
    fac_str = "" if _is_blank_value(fac_value) else str(fac_value).strip()

    if _PENYELESAIAN_RE.search(fac_str):
        clsdt_codes = clean_split_code(str(clsdt_value)) if not _is_blank_value(clsdt_value) else []
        if _is_valid_clean_result(clsdt_codes):
            keterangan = _TRAILING_CURRENCY_WORD_RE.sub("", fac_str).strip()
            keterangan = re.sub(r"\s{2,}", " ", keterangan)
            return [f"{code} {keterangan}" for code in clsdt_codes], False
        fac_codes = clean_split_code(fac_str)
        return fac_codes, _is_valid_clean_result(fac_codes)

    codes, _ = _pick_best_source_codes(clsdt_value, fac_value)
    return codes, _is_valid_clean_result(codes)


# ==============================================================================
# 4. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)


def filter_cedant(df_raw: pd.DataFrame, keyword: str, cedant_column: str = "") -> pd.DataFrame:
    if cedant_column:
        if cedant_column not in df_raw.columns:
            raise ValueError(f"Kolom {cedant_column} tidak ada.")
        cedant_cols = [cedant_column]
    else:
        cedant_cols = [c for c in df_raw.columns if "CEDANT" in str(c).upper()]
        if not cedant_cols:
            raise ValueError("Kolom CEDANT tidak ditemukan lewat auto-detect.")

    mask = pd.Series(False, index=df_raw.index)
    for col in cedant_cols:
        mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)

    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant '{keyword}': {len(df_filtered)} dari total {len(df_raw)} baris.")
    return df_filtered


# ==============================================================================
# 5. BUILD OUTPUT
# ==============================================================================
def _detect_key_columns(df: pd.DataFrame, insured_col: str, polis_col: str, slip_col: str):
    def _res(col: str, kw_or_fn, lbl: str):
        if col:
            if col not in df.columns:
                raise ValueError(f"{lbl}='{col}' tidak ada.")
            return col
        return next((c for c in df.columns if (kw_or_fn(str(c).upper()) if callable(kw_or_fn) else kw_or_fn in str(c).upper())), None)

    col_ins = _res(insured_col, "INSURED", "INSURED_COLUMN")
    col_pol = _res(polis_col, lambda c: c.strip() == "POLIS", "POLIS_COLUMN")
    col_slp = _res(slip_col, "SLIP", "SLIP_COLUMN")

    if None in (col_ins, col_pol, col_slp):
        raise ValueError("Kolom kunci tidak lengkap.")

    print(f"[INFO] Kolom terdeteksi -> INSURED='{col_ins}', POLIS='{col_pol}', SLIP='{col_slp}'")
    return col_ins, col_pol, col_slp


_NON_ALNUM_RE = re.compile(r"[^A-Za-z0-9]")
_SD_NOTATION_RE = re.compile(r"\bSD\b", flags=re.IGNORECASE)


def _compact_value(value: str) -> str:
    if not isinstance(value, str):
        return value
    if _SD_NOTATION_RE.search(value):
        return value
    return _NON_ALNUM_RE.sub("", value)


def _compact_list(values: list) -> list:
    return [_compact_value(v) for v in values]


def _breakdown_all_rows(
    df: pd.DataFrame,
    col_ins: str,
    col_pol: str,
    col_slp: str,
    col_clsdt_pol: str = None,
    col_clsdt_slp: str = None,
    col_clsdt_sertf: str = None,
):
    ins_arr = df[col_ins].to_numpy()
    pol_arr = df[col_pol].to_numpy()
    slp_arr = df[col_slp].to_numpy()

    clsdt_pol_arr = df[col_clsdt_pol].to_numpy() if col_clsdt_pol and col_clsdt_pol in df.columns else [None] * len(df)
    clsdt_slp_arr = df[col_clsdt_slp].to_numpy() if col_clsdt_slp and col_clsdt_slp in df.columns else [None] * len(df)
    clsdt_sertf_arr = df[col_clsdt_sertf].to_numpy() if col_clsdt_sertf and col_clsdt_sertf in df.columns else [None] * len(df)

    ins_all, pol_all, slp_all, cert_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val, clsdt_pol_val, clsdt_slp_val, clsdt_sertf_val in zip(
        ins_arr, pol_arr, slp_arr, clsdt_pol_arr, clsdt_slp_arr, clsdt_sertf_arr
    ):
        ins = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []
        raw_pol = pol_val if pd.notna(pol_val) else ""
        raw_slp = slp_val if pd.notna(slp_val) else ""

        if col_clsdt_pol and pd.notna(clsdt_pol_val):
            pol, pol_should_compact = _resolve_polis_or_slip(clsdt_pol_val, raw_pol)
        else:
            pol = clean_split_code(str(raw_pol).strip()) if not _is_blank_value(raw_pol) else []
            pol_should_compact = _is_valid_clean_result(pol)

        if col_clsdt_slp and pd.notna(clsdt_slp_val):
            slp, slp_should_compact = _resolve_polis_or_slip(clsdt_slp_val, raw_slp)
        else:
            slp = clean_split_code(str(raw_slp).strip()) if not _is_blank_value(raw_slp) else []
            slp_should_compact = _is_valid_clean_result(slp)

        if pol_should_compact:
            pol = _compact_list(pol)
        if slp_should_compact:
            slp = _compact_list(slp)

        cert = build_certificates_for_row(clsdt_sertf_val, raw_pol, num_polis=len(pol) if pol else 1)
        if len(cert) < len(pol):
            cert = cert + [""] * (len(pol) - len(cert))
        elif len(cert) > len(pol):
            cert = cert[: len(pol)] if pol else cert

        ins_all.append(ins)
        pol_all.append(pol)
        slp_all.append(slp)
        cert_all.append(cert)

        max_ins = max(max_ins, len(ins))
        max_pol = max(max_pol, len(pol))
        max_slp = max(max_slp, len(slp))

    return ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp


def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_ins, col_pol, col_slp = _detect_key_columns(df_filtered, INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN)

    col_clsdt_pol = CLSDT_POLIS_COLUMN if CLSDT_POLIS_COLUMN in df_filtered.columns else None
    col_clsdt_slp = CLSDT_SLIP_COLUMN if CLSDT_SLIP_COLUMN in df_filtered.columns else None
    col_clsdt_sertf = CLSDT_SERTF_COLUMN if CLSDT_SERTF_COLUMN in df_filtered.columns else None

    if col_clsdt_pol or col_clsdt_slp:
        print(f"[INFO] Kolom CLSDT terdeteksi -> POLIS='{col_clsdt_pol}', SLIP='{col_clsdt_slp}' (prioritas utama)")
    else:
        print("[INFO] Kolom CLSDT tidak ditemukan di file ini -> pakai FAC_POLICY_NO/FAC_SLIP saja.")

    if col_clsdt_sertf:
        print(f"[INFO] Kolom CERTIFICATE source terdeteksi -> '{col_clsdt_sertf}' (prioritas), fallback ke FAC_POLICY_NO.")
    else:
        print(f"[WARNING] Kolom '{CLSDT_SERTF_COLUMN}' tidak ditemukan -> CERTIFICATE hanya dari fallback FAC_POLICY_NO.")

    ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp = _breakdown_all_rows(
        df_filtered, col_ins, col_pol, col_slp, col_clsdt_pol, col_clsdt_slp, col_clsdt_sertf
    )

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)
    output_cols_data = {}

    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == col_ins:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in ins_all]
        if col == col_pol:
            # Pastikan minimal CERTIFICATE_CLN_1 terbuat
            for i in range(1, max(max_pol, 1) + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in pol_all]
                output_cols_data[f"{CERTIFICATE_OUTPUT_COL}_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_all]
        if col == col_slp:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slp_all]

    df_master = pd.DataFrame(output_cols_data)
    # Drop kolom _CLN_ kosong, kecuali CERTIFICATE_CLN_1 yang wajib tetap ada
    drop_cols = [
        c for c in df_master.columns
        if ("_CLN_" in c)
        and c != f"{CERTIFICATE_OUTPUT_COL}_CLN_1"
        and df_master[c].astype(str).str.strip().eq("").all()
    ]
    return df_master.drop(columns=drop_cols)


# ==============================================================================
# 6. PENYIMPANAN DATA
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook = load_workbook(output_path)
    text_col_indices = [
        idx for idx, col in enumerate(df.columns, start=1)
        if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)
    ]
    worksheet = workbook.active

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)
    workbook.save(output_path)


# ==============================================================================
# 7. RUNNER EKSEKUSI UTAMA
# ==============================================================================
def main() -> None:
    print("=" * 70 + "\n CLEANSING DATA 2 - OSBAL | CEDANT: MSIG - MITSUI SUMITOMO INSURANCE\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        raw_dir = "data_raw"
        if os.path.exists(raw_dir):
            candidates = [f for f in os.listdir(raw_dir) if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
            if not candidates:
                raise FileNotFoundError(f"File Excel sumber tidak ditemukan di folder '{raw_dir}'.")
            input_path = os.path.join(raw_dir, candidates[0])
            print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")
        else:
            raise FileNotFoundError(f"Folder '{raw_dir}' atau file sumber tidak ditemukan.")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    df_msig = filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)

    if not df_msig.empty:
        df_hasil = build_output(df_msig)
        os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
        save_with_text_format(df_hasil, OUTPUT_FILE)
        print("=" * 70 + f"\n[SUCCESS] Selesai. Output: '{OUTPUT_FILE}' ({len(df_hasil)} baris)\n" + "=" * 70)
    else:
        print("[WARNING] Tidak ada baris yang cocok dengan filter.")


if __name__ == "__main__":
    main()