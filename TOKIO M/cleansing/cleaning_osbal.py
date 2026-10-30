import datetime
import os
import re

import numpy as np
import pandas as pd


INPUT_FILE = os.path.join("data_raw", "2b. transaksi Osbal 01.01.23 - 17.08.26.xlsx")
OUTPUT_FILE = os.path.join("data_cleaned", "tokio_output_osbal.xlsx")

CEDANT_COL   = "CCOS_COMP_NAME"
CEDANT_VALUE = "PT.ASURANSI TOKIO MARINE INDONESIA"

POLIS_COL   = "FAC_POLICY_NO"
SLIP_COL    = "FAC_SLIP"
INSURED_COL = "FAC_INSURED"

# CLSDT (data closing) = sumber UTAMA nomor polis / slip / sertifikat.
# FAC_* hanya cadangan kalau CLSDT kosong / tidak valid.
CLSDT_POLICY_COL = "CLSDT_POLICY_NO"
CLSDT_SLIP_COL   = "CLSDT_SLIP_NO"

CLSDT_INVALID_VALUES = frozenset({
    "", "NAN", "NONE", "NULL", "NA", "N/A", "-", "TBA", "VAR", "VARIOUS",
})

MAX_SPLIT_COLS = 5

# Header output osbal (urutan & nama kolom final)
OUTPUT_HEADERS = [
    "CCOS_DOC_NO", "CCOS_DATE", "CCOS_REF_CODE", "CCOS_COMP", "CCOS_COMP_NAME",
    "CCOS_REF_COMP", "CCOS_REF_COMP_NAME",
    "FAC_INSURED",
    "FAC_INSURED_CLN_1", "FAC_INSURED_CLN_2", "FAC_INSURED_CLN_3",
    "FAC_INSURED_CLN_4", "FAC_INSURED_CLN_5",
    "CCOS_CURR", "CCOS_OR_BAL", "CCOS_BAL_DUE",
    "CCOS_OR_BAL_IN_IDR", "CCOS_BAL_DUE_IN_IDR",
    "FAC_COM_DATE", "FAC_EXP_DATE", "FAC_DUE_DATES", "FAC_SUB_CLASS",
    "FAC_POLICY_NO", "FAC_POLICY_NO_CLEAN_1", "CERTIFICATE_1",
    "FAC_POLICY_NO_CLEAN_2", "FAC_POLICY_NO_CLEAN_3",
    "FAC_POLICY_NO_CLEAN_4", "FAC_POLICY_NO_CLEAN_5",
    "FAC_SLIP",
    "FAC_SLIP_CLEAN_1", "FAC_SLIP_CLEAN_2", "FAC_SLIP_CLEAN_3",
    "FAC_SLIP_CLEAN_4", "FAC_SLIP_CLEAN_5",
    "CLASS_CODE", "CLASS_NAME",
    "CLSDT_POLICY_NO", "CLSDT_SLIP_NO", "CLSDT_SERTF_NO",
]

# Nama internal -> nama header final
_HEADER_RENAME = {
    "polis_ori": "FAC_POLICY_NO",
    "slip_ori": "FAC_SLIP",
    "insured_ori": "FAC_INSURED",
    "CERTIFICATE": "CERTIFICATE_1",
    **{f"clean polis {i}": f"FAC_POLICY_NO_CLEAN_{i}" for i in range(1, 6)},
    **{f"clean slip {i}": f"FAC_SLIP_CLEAN_{i}" for i in range(1, 6)},
    **{f"clean insured {i}": f"FAC_INSURED_CLN_{i}" for i in range(1, 6)},
}

# Polis: kata/prefix yang menyebabkan nilai dibiarkan apa adanya
POLIS_EXCEPTION_RE = re.compile(
    r"""
    MOP\s*MARINE
  | (?:LINE\s*SLIP|LINESLIP)
  | \b(?:P1|P2|P3|P73)\s*CANCEL
  | \b(?:P1|P2|P3|P73)\b
  | \bCANCEL\b
  | PENYELESAIAN(?:\s+SUSPENSE)?
  | HUTANG
  | UTANG
    """,
    re.IGNORECASE | re.VERBOSE,
)

# Slip: kata yang menyebabkan nilai dibiarkan apa adanya
SLIP_EXCEPTION_RE = re.compile(
    r"""
    \bSUMMARY\b
  | \bBORDER[OA]\b
  | \bBORDRO\b
  | \bSINGGLESHIPMENT\b
  | \bP1\b
  | \bP2\b
  | \bP3\b
  | PENYELESAIAN(?:\s+SUSPENSE)?
  | HUTANG
  | UTANG
    """,
    re.IGNORECASE | re.VERBOSE,
)

INSURED_JUNK_WORDS = frozenset({
    "PT", "CV", "TBK", "PERSERO", "LTD", "INC", "LLC",
    "AND", "OR", "THE", "OF", "AS",
    "NON FOOD", "DIV",
    "MR", "MR.", "MRS", "MRS.", "MS", "MS.",
})

INSURED_SUFFIX_RE = re.compile(
    r"""
    ,?\s*\bTBK\b\s*(?:,?\s*PT\.?)?
  | ,?\s*\bPT\.?\s*$
  | ,?\s*\bCV\.?\s*$
    """,
    re.IGNORECASE | re.VERBOSE,
)

INSURED_PT_QQ_RE = re.compile(
    r"""
    \bPT\.?\b
  | \bQQ\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

_SLIP_NOISE_RE = re.compile(
    r"""
    \b(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS
        |SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\b
  | \b(?:JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST
        |SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)\b
  | \b(?:IDR|USD|ENG)\b
  | \b(?:P1|P2|P3|P73)\b
  | \bNEW\b
  | \bVARIOUS\b
  | \b20[0-9]{2}\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

_SLIP_TOKEN_RE = re.compile(r"[A-Z0-9][A-Z0-9\-]{6,}", re.IGNORECASE)

# "... + P2" (penanda periode) dibuang dari breakdown polis; "P1 CANCEL" tetap exception
_PLUS_P_RE = re.compile(r"\s*\+\s*P\d+\b(?!\s*CANCEL)", re.IGNORECASE)

# satu bagian breakdown polis: tanpa spasi, mengandung angka
_PLUS_PART_RE = re.compile(r"[A-Z0-9./\-]*\d[A-Z0-9./\-]*")


# ─────────────────────────────────────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _normalize_spaces(text: str) -> str:
    return re.sub(r"\s{2,}", " ", text).strip()


def _clean_insured_name(name: str) -> str:
    name = _normalize_spaces(name)

    if re.fullmatch(r"AEON\s+MALL\s*\(TENANTS\)", name, re.IGNORECASE):
        return name.upper()

    for _ in range(3):
        cleaned = _normalize_spaces(INSURED_SUFFIX_RE.sub("", name).strip().strip(","))
        if cleaned == name:
            break
        name = cleaned

    name = INSURED_PT_QQ_RE.sub(" ", name)

    if re.fullmatch(r"\s*AEON\s+MALL\s+\(TENANTS\)\s*", name, flags=re.IGNORECASE):
        return "AEON MALL (TENANTS)"

    name = re.sub(r"\(\s*PERSERO\s*\)", "", name, flags=re.IGNORECASE)
    name = re.sub(r"[()]", " ", name)
    name = name.replace("(", " ").replace(")", " ")
    name = _normalize_spaces(name)
    name = re.sub(r"^[\s.,\-]+|[\s.,\-]+$", "", name)

    return name


def _cap_or_join(items: list) -> list:
    if len(items) > MAX_SPLIT_COLS:
        return [",".join(items)]
    return items


def _unique(items):
    out = []
    for x in items:
        if x and x not in out:
            out.append(x)
    return out



# ─────────────────────────────────────────────────────────────────────────────
# MAPPING PREFIX COB BERDASARKAN CLASS_NAME (Dictionary Tokio M)
# ─────────────────────────────────────────────────────────────────────────────
CLASS_TO_COB_PREFIX = {
    "MARINE CARGO":            "MINL",
    "PAR":                     "FPAR",
    "PROPERTY ALL RISK":       "FPAR",
    "IAR":                     "FIAR",
    "INDUSTRIAL ALL RISK":     "FIAR",
    "EAR":                     "ECAR",
    "COMPREHENSIVE L":         "LGEN",
    "PUBLIC LIABILITY":        "LGEN",
    "MB":                      "EMCB",
    "MACHINERY BREAKDOWN":     "EMCB",
    "CAR":                     "EEAR",
    "PROF INDEMNITY":          "LPIC",
    "PROFESSIONAL INDEMNITY":  "LPIC",
    "MARINE HULL":             "HMAR",
}

_RE_PREFIX_YEAR_CORE = re.compile(r"^([A-Z]+)(\d{2})([A-Z]\d{7,8})$")
_RE_POLIS_YEAR_CORE   = re.compile(r"^(\d{2})([A-Z]\d{7,8})$")
_RE_POLIS_CORE_ONLY   = re.compile(r"^([A-Z]\d{7,8})$")


def _extract_year2(val) -> str:
    """Ekstrak 2 digit tahun dari tanggal (Timestamp, datetime, atau string format tanggal)."""
    if val is None or pd.isna(val):
        return ""
    if isinstance(val, (pd.Timestamp, datetime.date, datetime.datetime)):
        return f"{val.year % 100:02d}"
    s = str(val).strip()
    if not s or s.upper() in ("NAN", "NONE", "NAT"):
        return ""
    # Format YYYY-... di awal, contoh '2026-09-09', '2023-01-01'
    m_year_start = re.match(r"^(?:19|20)(\d{2})", s)
    if m_year_start:
        return m_year_start.group(1)
    # Format DD/MM/YYYY atau DD-MM-YYYY
    m_year_end = re.search(r"[/.-](?:19|20)(\d{2})\b", s)
    if m_year_end:
        return m_year_end.group(1)
    return ""


def augment_cleaned_polis(c_polis: list, class_name, fac_com_date) -> list:
    """
    Melengkapi format polis OSBAL yang tidak lengkap berdasarkan CLASS_NAME & FAC_COM_DATE,
    serta membetulkan typo 2 digit tahun pada nomor polis jika berbeda dengan tahun FAC_COM_DATE:
    1. Koreksi typo tahun: jika polis sudah punya prefix (misal TMDFIAR23F0041968) namun FAC_COM_DATE
       berada di tahun 2022, tahun pada polis dibetulkan menjadi TMDFIAR22F0041968.
    2. Format 10 digit (2 digit tahun + 1 huruf + 7-8 angka, misal 26F0003746):
       -> tambahkan prefix format COB dari CLASS_NAME (dan sesuaikan tahun jika berbeda dari FAC_COM_DATE).
          Contoh: CLASS_NAME = PAR -> FPAR26F0003746
    3. Format 8 digit (1 huruf + 7-8 angka, misal F0003746 / F5000953):
       -> tambahkan prefix format COB dari CLASS_NAME + 2 digit tahun dari FAC_COM_DATE.
          Contoh: CLASS_NAME = PAR, FAC_COM_DATE = 2026-09-09 -> FPAR26F0003746
          Contoh: CLASS_NAME = PAR, FAC_COM_DATE tahun 2023   -> FPAR23F5000953
    """
    if not c_polis:
        return c_polis

    cls_str = "" if pd.isna(class_name) else str(class_name).strip().upper()
    prefix = CLASS_TO_COB_PREFIX.get(cls_str, "")
    year2 = _extract_year2(fac_com_date)

    augmented = []
    for item in c_polis:
        item_str = str(item).strip().upper()

        # Case 0: Polis sudah punya prefix huruf + 2 digit tahun + core
        # Contoh: TMDFIAR23F0041968 dengan FAC_COM_DATE tahun 2022 -> TMDFIAR22F0041968
        m_full = _RE_PREFIX_YEAR_CORE.match(item_str)
        if m_full:
            pol_pref, pol_year, pol_core = m_full.group(1), m_full.group(2), m_full.group(3)
            if year2 and pol_year != year2:
                augmented.append(f"{pol_pref}{year2}{pol_core}")
            else:
                augmented.append(item_str)
            continue

        # Case 1: 10-11 char (2 digit tahun + huruf + angka, misal 26F0003746)
        m10 = _RE_POLIS_YEAR_CORE.match(item_str)
        if m10:
            pol_year, pol_core = m10.group(1), m10.group(2)
            eff_year = year2 if year2 else pol_year
            if prefix:
                augmented.append(f"{prefix}{eff_year}{pol_core}")
            else:
                augmented.append(f"{eff_year}{pol_core}")
            continue

        # Case 2: 8-9 char (huruf + angka, misal F0003746 / F5000953)
        m8 = _RE_POLIS_CORE_ONLY.match(item_str)
        if m8 and prefix and year2:
            augmented.append(f"{prefix}{year2}{item_str}")
            continue

        augmented.append(item)

    return augmented


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN POLIS TOKIO
# ─────────────────────────────────────────────────────────────────────────────

def clean_polis_raw(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).strip().upper()

    if not val:
        return []

    val = re.sub(r"\bVARIOUS\b", "", val)
    val = re.sub(r"\s*/\s*VAR$", "", val)

    if val != "TBA":
        val = re.sub(r"\bTBA\b", "", val)

    if val == "TBA":
        return ["TBA"]

    val = val.replace("&", "+")
    val = _normalize_spaces(val)
    val = val.replace(".", "")
    val = re.sub(r"-\d{3}-\d{2}$", "", val)
    val = _normalize_spaces(_PLUS_P_RE.sub("", val))

    if POLIS_EXCEPTION_RE.search(val):
        return [_normalize_spaces(val)]

    if re.search(
        r"(JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\s+\d{4}",
        val,
        flags=re.IGNORECASE,
    ):
        return [_normalize_spaces(val)]

    m = re.fullmatch(
        r"([A-Z]{3}/[A-Z]{4}/\d{2}-[A-Z]\d{6,8})(?:-000-\d+)?",
        val,
        flags=re.IGNORECASE,
    )
    if m:
        return [m.group(1).upper()]

    m = re.fullmatch(
        r"(\d{2}-?[A-Z]\d{7,8})(?:-000-\d+)?",
        val,
        flags=re.IGNORECASE,
    )
    if m:
        return [m.group(1).upper().replace("-", "")]

    m = re.fullmatch(
        r"([A-Z]\d{7,8})(?:-000-\d+)?",
        val,
        flags=re.IGNORECASE,
    )
    if m:
        return [m.group(1).upper()]

    if "S/D" in val:
        hasil = re.findall(r"[A-Z]\d{7,8}", val)
        if hasil:
            return _cap_or_join(_unique(hasil))

    val = re.sub(
        r"([A-Z]\d{7,8})-\d{3}(?=\s|\+|$)",
        r"\1",
        val,
        flags=re.IGNORECASE,
    )

    val = re.sub(
        r"([A-Z]\d{7,8})\s+(TMD/|TMI/)",
        r"\1+\2",
        val,
        flags=re.IGNORECASE,
    )

    if "+" in val:
        parts = [p.strip() for p in val.split("+") if p.strip()]
        results = []
        tokio_prefix = None
        tokio_company = None
        derived = set()   # index hasil yang berasal dari perulangan angka pendek

        for p in parts:
            m_company = re.match(r"^(TMD|TMI)/", p, flags=re.IGNORECASE)
            if m_company:
                tokio_company = m_company.group(1).upper() + "/"

            m = re.fullmatch(r"([A-Z]{4}/\d{2}[A-Z]\d{7,8})", p, flags=re.IGNORECASE)
            if m and tokio_company:
                results.append(tokio_company + m.group(1).upper())
                continue

            m = re.fullmatch(
                r"^([A-Z]{3}/[A-Z]{4}/\d{2}-)([A-Z]\d{7,8})((?:-[A-Z]\d{7,8})+)$",
                p,
                flags=re.IGNORECASE,
            )
            if m:
                tokio_prefix = m.group(1).upper()
                results.append(tokio_prefix + m.group(2).upper())
                suffixes = re.findall(r"[A-Z]\d{7,8}", m.group(3), flags=re.IGNORECASE)
                for suffix in suffixes:
                    results.append(tokio_prefix + suffix.upper())
                continue

            m = re.fullmatch(r"([A-Z]{4}/\d{2}[A-Z]\d{7,8})", p, flags=re.IGNORECASE)
            if m and tokio_prefix:
                results.append("TMD/" + m.group(1).upper())
                continue
            
            m = re.match(
                r"^([A-Z]{3}/[A-Z]{4}/\d{2}-)([A-Z]\d{7,8})$",
                p,
                flags=re.IGNORECASE,
            )
            if m:
                tokio_prefix = m.group(1).upper()
                results.append(tokio_prefix + m.group(2).upper())
                continue

            m = re.match(
                r"^([A-Z]{7,8}\d{2})([A-Z]\d{7,8})$",
                p,
                flags=re.IGNORECASE,
            )
            if m:
                tokio_prefix = m.group(1).upper()
                results.append(tokio_prefix + m.group(2).upper())
                continue
                
            m = re.fullmatch(r"(\d{2}-?[A-Z]\d{7,8})", p, flags=re.IGNORECASE)
            if m:
                code = m.group(1).upper().replace("-", "")
                results.append(tokio_prefix + code if tokio_prefix else code)
                continue

            m = re.fullmatch(r"[A-Z]\d{7,8}", p, flags=re.IGNORECASE)
            if m:
                code = m.group(0).upper()
                results.append(tokio_prefix + code if tokio_prefix else code)
                continue

            # Perulangan polis: angka pendek menimpa digit paling belakang
            # dari polis lengkap terakhir (bukan hasil perulangan sebelumnya).
            if re.fullmatch(r"\d{1,7}", p):
                base = next(
                    (results[i] for i in range(len(results) - 1, -1, -1)
                     if i not in derived
                     and re.search(r"[A-Z]\d{7,8}$", results[i])),
                    None,
                )
                if base:
                    tail_len = len(re.search(r"\d+$", base).group(0))
                    if len(p) <= tail_len:
                        results.append(base[:-len(p)] + p)
                        derived.add(len(results) - 1)
                        continue

            if p.upper() not in ["VAR", "VARIOUS", "TBA"]:
                results.append(p.upper())

        return _cap_or_join(_unique(results))

    m = re.fullmatch(
        r"([A-Z]{3}/[A-Z]{4}/\d{2}-)([A-Z]\d{7,8})((?:-[A-Z]\d{7,8})+)",
        val,
        flags=re.IGNORECASE,
    )
    if m:
        tokio_prefix = m.group(1).upper()
        first_code = m.group(2).upper()
        results = [tokio_prefix + first_code]
        suffixes = re.findall(r"[A-Z]\d{7,8}", m.group(3), flags=re.IGNORECASE)
        for suffix in suffixes:
            results.append(tokio_prefix + suffix.upper())
        return _cap_or_join(_unique(results))

    m = re.fullmatch(
        r"((?:TMD|TMI)[A-Z]{4}\d{2})([A-Z]\d{7,8})-(\d{4})-((?:TMD|TMI)[A-Z]{4}\d{2}[A-Z]\d{7,8})",
        val,
        flags=re.IGNORECASE,
    )
    if m:
        tokio_prefix = m.group(1).upper()
        first_code = m.group(2).upper()
        second_suffix = m.group(3)
        third_policy = m.group(4).upper()
        code_prefix = first_code[:-4]

        return _cap_or_join(_unique([
            tokio_prefix + first_code,
            tokio_prefix + code_prefix + second_suffix,
            third_policy,
        ]))

    m = re.fullmatch(
        r"((?:TMD|TMI)[A-Z]{4}\d{2})([A-Z]\d{7,8})-([A-Z]\d{7,8})",
        val,
        flags=re.IGNORECASE
    )
    if m:
        tokio_prefix = m.group(1).upper()
        return [
            tokio_prefix + m.group(2).upper(),
            tokio_prefix + m.group(3).upper()
        ]

    m = re.fullmatch(
        r"((?:TMD|TMI)[A-Z]{4}\d{2})([A-Z]\d{7,8})((?:-\d{3,7})+)",
        val,
        flags=re.IGNORECASE
    )
    if m:
        tokio_prefix = m.group(1).upper()
        first_code = m.group(2).upper()
        results = [tokio_prefix + first_code]
        suffixes = re.findall(r"\d{3,7}", m.group(3))

        for suffix in suffixes:
            # suffix menimpa digit paling belakang polis pertama
            code_prefix = first_code[:-len(suffix)]
            results.append(tokio_prefix + code_prefix + suffix)

        return _cap_or_join(_unique(results))

    if re.fullmatch(
        r"(?:TMD|TMI)[A-Z]{4}\d{2}[A-Z]\d{7,8}(?:-\d+)+",
        val,
        flags=re.IGNORECASE
    ):
        return [val]

    hasil = re.findall(
        r"(?:TMD|TMI)?[A-Z]{3,5}\d{2}[A-Z]\d{7,8}",
        val,
        flags=re.IGNORECASE
    )
    if len(hasil) > 1:
        return _cap_or_join(_unique([x.upper() for x in hasil]))

    hasil_10 = re.findall(
        r"\b\d{2}[A-Z]\d{7,8}\b",
        val,
        flags=re.IGNORECASE
    )
    if len(hasil_10) > 1:
        return _cap_or_join(_unique([x.upper() for x in hasil_10]))
    if len(hasil_10) == 1 and not re.search(r"(?:TMD|TMI)[A-Z]", val, flags=re.IGNORECASE):
        return [hasil_10[0].upper()]

    hasil = re.findall(
        r"[A-Z]\d{7,8}",
        val,
        flags=re.IGNORECASE
    )
    if len(hasil) > 1:
        return _cap_or_join(_unique([x.upper() for x in hasil]))

    if re.match(r"^(TMD|TMI)[A-Z0-9]+$", val):
        return [val.upper()]

    m = re.search(r"\b\d{2}[A-Z]\d{7,8}\b", val)
    if m:
        return [m.group(0).upper()]

    m = re.search(r"[A-Z]\d{7,8}", val)
    if m:
        return [m.group(0).upper()]

    return []


def _polis_asis_if_many(val):
    """
    Breakdown '+' dengan lebih dari MAX_SPLIT_COLS (5) polis -> dibiarkan apa adanya
    (tidak di-clean, tanda '+' tidak dibuang). Return None kalau tidak berlaku.
    """
    if pd.isna(val):
        return None

    s = _normalize_spaces(str(val).strip().upper())
    s2 = _PLUS_P_RE.sub("", s.replace("&", "+"))
    if "+" not in s2 or POLIS_EXCEPTION_RE.search(s2):
        return None

    parts = [p.strip() for p in s2.split("+") if p.strip()]
    if len(parts) > MAX_SPLIT_COLS and all(_PLUS_PART_RE.fullmatch(p) for p in parts):
        return s
    return None


def _polis_asis_if_many_dash(val):
    """
    Breakdown '-' (TMDFIAR21F0039255-688-312-...) dengan lebih dari 5 polis
    -> dibiarkan apa adanya (tanda '-' tidak dibuang). Return None kalau tidak berlaku.
    """
    if pd.isna(val):
        return None

    s = _normalize_spaces(str(val).strip().upper())
    m = re.fullmatch(
        r"(?:TMD|TMI)[A-Z]{4}\d{2}[A-Z]\d{7,8}((?:-\d+)+)",
        s.replace(".", "").replace("/", ""),
    )
    if m and 1 + len(re.findall(r"-\d+", m.group(1))) > MAX_SPLIT_COLS:
        return s
    return None


def clean_polis(val) -> list:
    if pd.notna(val):
        val_str = str(val).strip().upper()
        if re.fullmatch(r"TMD\.FIAR\.22\.F\d{7,8}(?:\+\d{3})+", val_str):
            return [val_str.replace(".", "")]

    asis = _polis_asis_if_many(val)
    if asis is None:
        asis = _polis_asis_if_many_dash(val)
    if asis is not None:
        return [asis]

    hasil = clean_polis_raw(val)
    hasil_clean = []

    for x in hasil:
        x = re.sub(r"[^A-Z0-9]", "", str(x).upper())
        if x:
            hasil_clean.append(x)

    return hasil_clean

# ─────────────────────────────────────────────────────────────────────────────
# CLEAN SLIP
# ─────────────────────────────────────────────────────────────────────────────

def _clean_slip_core(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).strip()
    if not val:
        return []

    val = _normalize_spaces(val)

    val = re.sub(
        r"([A-Z]{3}/[A-Z]{4}/\d{2}-[A-Z]\d{7,8})-\d{3}-\d{2}",
        r"\1",
        val,
        flags=re.IGNORECASE,
    )

    val = re.sub(
        r"([A-Z]\d{7,8})-\d{3}-\d{2}",
        r"\1",
        val,
        flags=re.IGNORECASE,
    )

    m = re.match(r"^(\d+)\s+S/D\s+\d+", val, flags=re.IGNORECASE)
    if m:
        return [m.group(1)]

    m = re.search(r"(TM[DI]/[A-Z]{4}/\d{2}-[A-Z]\d{7,8})", val, flags=re.IGNORECASE)
    if m:
        return [m.group(1)]

    if "FACILITY" in val:
        return [val]

    if POLIS_EXCEPTION_RE.search(val):
        return [_normalize_spaces(val)]

    if re.fullmatch(r"P\d+(?:\s*\+\s*P\d+)*(\s+CANCEL)?", val, flags=re.IGNORECASE):
        return [_normalize_spaces(val)]

    val = re.sub(r"\s*/\s*END\b", "", val, flags=re.IGNORECASE).strip()

    val = re.sub(
        r"\s*/\s*(VAR\s+ACCOUNT|FACILITY\s+AUTOMOBILE|FACILITY)\b",
        "",
        val,
        flags=re.IGNORECASE,
    )

    val = re.sub(r"\bVARIOUS\b", "", val, flags=re.IGNORECASE)
    val = re.sub(r"\bVAR\b", "", val, flags=re.IGNORECASE)
    val = _normalize_spaces(val)

    val = re.sub(r"/\s*VAR\s+ACCOUNT\b", "", val, flags=re.IGNORECASE)
    val = re.sub(r"/\s*FACILITY\s+AUTOMOBILE\b", "", val, flags=re.IGNORECASE)
    val = re.sub(r"/\s*FACILITY\b", "", val, flags=re.IGNORECASE)
    val = _normalize_spaces(val)

    parts = re.split(r"\s{2,}", val)
    parts = [p.strip() for p in parts if p.strip()]

    if len(parts) >= 2:
        results = [_normalize_spaces(p) for p in parts]
        return _cap_or_join(results)

    if SLIP_EXCEPTION_RE.search(val):
        return [_normalize_spaces(val)]

    if re.fullmatch(r"\d+(?:\.\d+)+(?:-\d+){3,}", val):
        return [_normalize_spaces(val)]

    if "+" in val:
        parts = [p.strip() for p in re.split(r"\s*\+\s*", val) if p.strip()]
        parts = [p for p in parts if not re.fullmatch(r"P\d+", p, flags=re.IGNORECASE)]

        if len(parts) > 1:
            first = parts[0]
            rest = parts[1:]

            if all(re.fullmatch(r"\d+", p) for p in parts):
                return _cap_or_join(parts)

            rest_is_digit = all(re.fullmatch(r"\d+", r) for r in rest)
            suffix_lengths = {len(r) for r in rest} if rest_is_digit else set()

            if rest_is_digit and len(suffix_lengths) == 1:
                suffix_len = suffix_lengths.pop()
                m = re.match(rf"^(.*?)(\d{{{suffix_len}}})$", first)

                if m:
                    prefix = m.group(1)
                    results = [first]
                    for s in rest:
                        s = s.zfill(suffix_len)
                        results.append(prefix + s)
                    return _cap_or_join(results)

            if all(re.match(r"^\d+(?:\.\d+)+$", p) for p in parts):
                return _cap_or_join(parts)

            return [_normalize_spaces(val)]

    m = re.match(r"^(.*?)(\d+)-(\d+)$", val)
    if m:
        prefix, first, second = m.group(1), m.group(2), m.group(3)
        results = [prefix + first]
        if len(second) < len(first):
            second = second.zfill(len(first))
            results.append(prefix + second)
        else:
            results.append(second)
        return _cap_or_join(results)

    slips = re.findall(r"\d+(?:\.\d+)+", val)
    if len(slips) >= 2:
        return _cap_or_join(slips)

    if re.search(
        r"(JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\s+\d{4}",
        val,
        flags=re.IGNORECASE,
    ):
        return [_normalize_spaces(val)]

    m = re.search(r"\d+(?:\.\d+)+", val)
    if m:
        return [m.group(0)]

    return [_normalize_spaces(val)]     


def clean_slip_raw(val) -> list:
    results = _clean_slip_core(val)
    cleaned = []

    for r in results:
        if pd.isna(r):
            continue
        r = str(r).upper()
        r = _normalize_spaces(r)
        r = r.replace(".", "")
        r = re.sub(r'(?<=\d)\s*TBA\b', '', r).strip()

        if r and r not in cleaned:
            cleaned.append(r)

    val_str = "" if pd.isna(val) else str(val).strip()
    if re.fullmatch(r"\d{4}", val_str):
        return ["0000" + val_str]

    return cleaned


def clean_slip(val) -> list:
    hasil = clean_slip_raw(val)
    hasil_clean = []

    for x in hasil:
        x = re.sub(r"[^A-Z0-9]", "", str(x).upper())
        if x:
            hasil_clean.append(x)

    return hasil_clean

# ─────────────────────────────────────────────────────────────────────────────
# CLEAN INSURED
# ─────────────────────────────────────────────────────────────────────────────

def clean_insured(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).strip()
    if not val:
        return []

    bracket_parts = re.findall(r"\(([^)]*)\)", val)
    val = re.sub(r"\([^)]*\)", "", val)

    parts = re.split(r"/|,|QQ", val, flags=re.IGNORECASE)
    parts.extend(bracket_parts)

    cleaned = []

    for p in parts:
        p = _normalize_spaces(p.strip())
        p = re.sub(r"^(MR|MRS|MS|DR|IR)\.?\s+", "", p, flags=re.IGNORECASE)

        if len(p) <= 2 or p.upper() in INSURED_JUNK_WORDS:
            continue
        if re.match(r"^[^A-Za-z0-9]+$", p):
            continue

        p_clean = _clean_insured_name(p)
        if not p_clean or p_clean.upper() in INSURED_JUNK_WORDS:
            continue

        if p_clean not in cleaned:
            cleaned.append(p_clean.upper())

    if not cleaned:
        fallback = _clean_insured_name(_normalize_spaces(val))
        return [fallback.upper()] if fallback else []

    return _cap_or_join(cleaned)


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN CERTIFICATE
# ─────────────────────────────────────────────────────────────────────────────

def _is_valid_clsdt(val) -> bool:
    """CLSDT dianggap valid kalau tidak NaN/blank/TBA/VAR/VARIOUS."""
    if val is None or pd.isna(val):
        return False
    return str(val).strip().upper() not in CLSDT_INVALID_VALUES


def clean_polis_clsdt(clsdt_val, fac_val, class_name=None, fac_com_date=None) -> tuple:
    """
    Prioritas CLSDT_POLICY_NO; kalau tidak valid (atau hasil cleaning kosong)
    fallback ke FAC_POLICY_NO. Return (list clean polis, sumber).
    Nomor polis pendek yang tidak lengkap diaugmentasi berdasarkan CLASS_NAME & FAC_COM_DATE.
    """
    if _is_valid_clsdt(clsdt_val):
        hasil = clean_polis(clsdt_val)
        if hasil:
            return augment_cleaned_polis(hasil, class_name, fac_com_date), "CLSDT"
    hasil = clean_polis(fac_val)
    return augment_cleaned_polis(hasil, class_name, fac_com_date), "FAC"


def clean_slip_clsdt(clsdt_val, fac_val) -> tuple:
    """
    Prioritas CLSDT_SLIP_NO; kalau tidak valid (atau hasil cleaning kosong)
    fallback ke FAC_SLIP. Return (list clean slip, sumber).
    """
    if _is_valid_clsdt(clsdt_val):
        hasil = clean_slip(clsdt_val)
        if hasil:
            return hasil, "CLSDT"
    return clean_slip(fac_val), "FAC"


_CERT_RE = re.compile(r"-(\d{3})-\d+$")


def clean_certificate(clsdt_polis) -> str:
    """
    Certificate diambil dari segmen 3 digit di ujung CLSDT_POLICY_NO:
      TMD/EEAR/16-E0006258-009-03 -> 000009
      TMD/FIAR/24-F5040056-005-01 -> 000005
      TMI/FIAR/23-F5041278-000-01 -> "" (000 = tidak ada certificate)
    """
    if not _is_valid_clsdt(clsdt_polis):
        return ""
    m = _CERT_RE.search(str(clsdt_polis).strip().upper())
    if not m or int(m.group(1)) == 0:
        return ""
    return str(int(m.group(1))).zfill(6)


# ─────────────────────────────────────────────────────────────────────────────
# PROCESS DATA (OPTIMIZED EXCEL READ/WRITE)
# ─────────────────────────────────────────────────────────────────────────────

def _insert_clean_columns(
    df: pd.DataFrame,
    all_lists: list,
    prefix: str,
    max_cols: int,
) -> list:
    added = []
    for i in range(1, max_cols + 1):
        col_name = f"clean {prefix} {i}"
        df[col_name] = [lst[i - 1] if i - 1 < len(lst) else None for lst in all_lists]
        added.append(col_name)
    return added


def _fast_read_excel(path: str, sheet_name=None, header: int = 0) -> pd.DataFrame:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    for _ in range(header):
        next(rows_iter)
    cols = next(rows_iter)
    data = list(rows_iter)
    wb.close()
    return pd.DataFrame(data, columns=cols)


def _fast_write_excel(df: pd.DataFrame, path: str) -> str:
    import openpyxl
    os.makedirs(os.path.dirname(path), exist_ok=True)

    df2 = df.astype(object).where(pd.notnull(df), None)

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    ws.append(list(df2.columns))
    for row in df2.itertuples(index=False, name=None):
        ws.append(row)

    target_path = path
    try:
        wb.save(target_path)
    except PermissionError:
        base, ext = os.path.splitext(path)
        target_path = f"{base}_new{ext}"
        print(f"\n[WARNING] File '{path}' sedang dibuka di aplikasi lain (Excel).")
        print(f"          Menyimpan hasil ke file alternatif: '{target_path}'")
        wb.save(target_path)

    return target_path


def process_data(input_file: str, output_file: str) -> None:
    print(f"[1/5] Membaca data cepat dari: {input_file} ...")
    df = _fast_read_excel(input_file, header=0)
    print(f"      Total baris keseluruhan: {len(df):,}")

    if CEDANT_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{CEDANT_COL}' tidak ditemukan!")
        print(f"        Kolom tersedia: {list(df.columns)}")
        return

    for col in [POLIS_COL, SLIP_COL, INSURED_COL,
                CLSDT_POLICY_COL, CLSDT_SLIP_COL]:
        if col not in df.columns:
            print(f"\n[ERROR] Kolom '{col}' tidak ditemukan di file!")
            print(f"        Kolom tersedia: {list(df.columns)}")
            return

    df[CEDANT_COL] = df[CEDANT_COL].fillna("").astype(str).str.strip()
    is_tokio = df[CEDANT_COL] == CEDANT_VALUE
    print(f"[2/5] Filter cedant '{CEDANT_VALUE}': {is_tokio.sum():,} baris TOKIO dari total {len(df):,} baris.")

    # Buang semua baris selain cedant Tokio -> output hanya Tokio
    total_awal = len(df)
    df = df[is_tokio].reset_index(drop=True)
    is_tokio = pd.Series(True, index=df.index)
    print(f"      Baris non-Tokio dibuang: {total_awal - len(df):,} | sisa: {len(df):,}")

    df.rename(columns={POLIS_COL: "polis_ori", SLIP_COL: "slip_ori", INSURED_COL: "insured_ori"},
              inplace=True)

    print("[3/5] Menjalankan proses cleaning cepat hanya untuk baris TOKIO ...")

    n = len(df)
    all_clean_polis = [[] for _ in range(n)]
    all_clean_slip  = [[] for _ in range(n)]
    all_clean_ins   = [[] for _ in range(n)]
    all_certificates = ["" for _ in range(n)]

    max_polis = max_slip = max_ins = 1

    tokio_idx = np.flatnonzero(is_tokio.to_numpy())
    polis_vals   = df["polis_ori"].to_numpy()
    slip_vals    = df["slip_ori"].to_numpy()
    insured_vals = df["insured_ori"].to_numpy()
    clsdt_polis_vals = df[CLSDT_POLICY_COL].to_numpy()
    clsdt_slip_vals  = df[CLSDT_SLIP_COL].to_numpy()
    class_name_vals  = df["CLASS_NAME"].to_numpy() if "CLASS_NAME" in df.columns else np.array([None] * n)
    fac_date_vals    = df["FAC_COM_DATE"].to_numpy() if "FAC_COM_DATE" in df.columns else np.array([None] * n)

    src_polis = {"CLSDT": 0, "FAC": 0}
    src_slip  = {"CLSDT": 0, "FAC": 0}

    total_t = len(tokio_idx)
    for n_done, pos in enumerate(tokio_idx, 1):
        if n_done % 20_000 == 0:
            print(f"      Progress: {n_done:,} / {total_t:,} baris TOKIO diproses...")

        p_ori = polis_vals[pos]
        c_polis, s_polis = clean_polis_clsdt(
            clsdt_polis_vals[pos],
            p_ori,
            class_name=class_name_vals[pos],
            fac_com_date=fac_date_vals[pos],
        )
        c_slip,  s_slip  = clean_slip_clsdt(clsdt_slip_vals[pos], slip_vals[pos])
        c_ins   = clean_insured(insured_vals[pos])
        c_cert  = clean_certificate(clsdt_polis_vals[pos])
        src_polis[s_polis] += 1
        src_slip[s_slip]   += 1

        max_polis = max(max_polis, len(c_polis))
        max_slip  = max(max_slip,  len(c_slip))
        max_ins   = max(max_ins,   len(c_ins))

        all_clean_polis[pos] = c_polis
        all_clean_slip[pos]  = c_slip
        all_clean_ins[pos]   = c_ins
        all_certificates[pos] = c_cert

    print(f"      Selesai diproses!")
    print(f"      -> Jumlah kolom clean polis  : {max_polis}")
    print(f"      -> Jumlah kolom clean slip   : {max_slip}")
    print(f"      -> Jumlah kolom clean insured: {max_ins}")
    print(f"      -> Sumber polis : CLSDT {src_polis['CLSDT']:,} | fallback FAC {src_polis['FAC']:,}")
    print(f"      -> Sumber slip  : CLSDT {src_slip['CLSDT']:,} | fallback FAC {src_slip['FAC']:,}")

    print("[4/5] Menyusun kolom output ...")

    df["CERTIFICATE"] = all_certificates

    new_columns = []
    for col in df.columns:
        if col == "CERTIFICATE":
            continue

        new_columns.append(col)
        if col == "polis_ori":
            clean_polis_cols = _insert_clean_columns(df, all_clean_polis, "polis", MAX_SPLIT_COLS)
            if clean_polis_cols:
                new_columns.append(clean_polis_cols[0])
                new_columns.append("CERTIFICATE")
                new_columns += clean_polis_cols[1:]
            else:
                new_columns.append("CERTIFICATE")
        elif col == "slip_ori":
            new_columns += _insert_clean_columns(df, all_clean_slip,  "slip",    MAX_SPLIT_COLS)
        elif col == "insured_ori":
            new_columns += _insert_clean_columns(df, all_clean_ins,   "insured", MAX_SPLIT_COLS)

    df = df[new_columns]

    # Ganti nama kolom sesuai header final & urutkan
    df = df.rename(columns=_HEADER_RENAME)

    missing = [c for c in OUTPUT_HEADERS if c not in df.columns]
    if missing:
        raise ValueError(f"Kolom berikut tidak ada di data: {missing}")

    extra = [c for c in df.columns if c not in OUTPUT_HEADERS]
    if extra:
        print(f"      [INFO] Kolom di luar header yang ditentukan (ditaruh di paling kanan): {extra}")

    df = df[OUTPUT_HEADERS + extra]

    print(f"[5/5] Menyimpan hasil cepat ke: {output_file} ...")
    saved_path = _fast_write_excel(df, output_file)

    print(f"\n{'=' * 55}")
    print(f"  [OK] Selesai!")
    print(f"  Total baris output  : {len(df):,}")
    print(f"  Total kolom output  : {len(df.columns)}")
    print(f"  File disimpan di    : {saved_path}")
    print(f"{'=' * 55}")


# ─────────────────────────────────────────────────────────────────────────────
# ENTRY POINT
# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    process_data(INPUT_FILE, OUTPUT_FILE)