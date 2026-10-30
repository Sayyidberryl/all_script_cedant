"""Audit Insured TF-IDF mandiri untuk TRIPA (PT Asuransi Tripakarta).
Tidak mengubah data input atau keputusan matching.
Menjalankan pencarian kandidat nama Insured menggunakan Char N-Gram TF-IDF & Cosine Similarity
terhadap 3 sumber data referensi Tripa: OSBAL, FACUL, dan RI SLIP (tanpa BORDERO).

Jalankan dari folder project:
    python TFIDF-TRIPA.py
Dependensi: pandas, openpyxl, numpy, scikit-learn.
"""
from pathlib import Path
from collections import defaultdict
import re
import unicodedata
import time
import numpy as np
import pandas as pd
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.neighbors import NearestNeighbors

ROOT = Path(__file__).resolve().parent
CLEAN_DIR = ROOT / 'data_cleaned'
RAW_DIR = ROOT / 'data_raw'
OUT = ROOT / 'output'
TOP_N = 10
MIN_SCORE = 0.0

# Jalur pencarian file Matching hasil eksekusi Tripa
MATCH_PATHS = [
    ROOT / 'matching_tripa.xlsx',
    ROOT / 'matching_tripa_new.xlsx',
    OUT / 'matching_tripa.xlsx',
    CLEAN_DIR / 'tripa_output_suspend.xlsx',
]

OSBAL_PATHS = [
    CLEAN_DIR / 'tripa_output_osbal.xlsx',
    OUT / 'tripa_output_osbal.xlsx',
    ROOT / 'tripa_output_osbal.xlsx',
]

FACUL_PATHS = [
    CLEAN_DIR / 'tripa_output_facul.xlsx',
    OUT / 'tripa_output_facul.xlsx',
    ROOT / 'tripa_output_facul.xlsx',
]

RI_SLIP_PATHS = [
    RAW_DIR / '5. Database Rekap RI Slip.xlsx',
    CLEAN_DIR / '5. Database Rekap RI Slip.xlsx',
    OUT / '5. Database Rekap RI Slip.xlsx',
    ROOT / '5. Database Rekap RI Slip.xlsx',
    OUT / 'tripa_output_ri_slip.xlsx',
]

DEST = ROOT / 'tripa_insured_tfidf_audit_3_sumber.xlsx'

# Stopwords badan hukum & entitas generik
COMMON_ENTITIES = {
    'PT', 'CV', 'UD', 'TBK', 'PERSERO', 'LTD', 'INC', 'LLC', 'CO', 'CORP',
    'CORPORATION', 'COMPANY', 'NV', 'PD'
}

# Stopwords industri / kata bisnis generik untuk evaluasi WEAK_SHARED_TOKENS
COMMON_BUSINESS = {
    'GROUP', 'GROUPS', 'GRUP', 'INDONESIA', 'INDO', 'INDUSTRY', 'INDUSTRIES', 'INDUSTRI',
    'INTERNATIONAL', 'INTERNASIONAL', 'ENTERPRISE', 'ENTERPRISES',
    'TRADING', 'TRADE', 'UTAMA', 'MANDIRI', 'MAKMUR', 'SEJAHTERA',
    'LINE', 'LINES', 'SHIPPING', 'SHIP', 'SHIPS', 'LOGISTICS', 'LOGISTIC',
    'TRANSPORT', 'TRANSPORTATION', 'PELAYARAN', 'PERKAPALAN', 'MARITIM', 'MARITIME',
    'CARRIER', 'CARRIERS', 'SERVICES', 'SERVICE', 'HOLDING', 'HOLDINGS'
}

COMMON_ALL = COMMON_ENTITIES | COMMON_BUSINESS


def txt(value):
    if value is None or (not isinstance(value, (list, tuple)) and pd.isna(value)):
        return ''
    value = str(value).strip()
    return '' if value.upper() in {'NAN', 'NONE', 'NAT'} else value


def norm(value):
    """Normalisasi string: NFKC, uppercase, hilangkan non-alfanumerik dan badan usaha."""
    value = unicodedata.normalize('NFKC', txt(value)).upper()
    value = re.sub(r'[^A-Z0-9]+', ' ', value)
    return ' '.join(token for token in value.split() if token not in COMMON_ENTITIES)


def col(df, *names):
    """Mencari nama kolom secara case-insensitive & simbol-insensitif."""
    lookup = {re.sub(r'[^A-Z0-9]+', '', str(key).upper()): key for key in df.columns}
    for name in names:
        result = lookup.get(re.sub(r'[^A-Z0-9]+', '', name.upper()))
        if result is not None:
            return result
    return None


def insured_clean_columns(df):
    """Mendeteksi kolom insured clean: INSURED_1..5, INSURED_CLEAN_1..5, FAC_INSURED_CLN_1..5, clean insured 1..5."""
    result = []
    for column in df.columns:
        name = re.sub(r'[^A-Z0-9]+', '', str(column).upper())
        if re.fullmatch(r'(?:INSUREDCLEAN|INSUREDCLN|CLEANINSURED|FACINSUREDCLN|FACINSURED|INSURED)[1-5]', name):
            result.append(column)
    return sorted(result, key=lambda c: int(re.search(r'[1-5]$', str(c)).group()))


def name_columns(df, source):
    """Ambil kolom-kolom penampung nama Tertanggung sesuai sumber."""
    clean = insured_clean_columns(df)
    if clean:
        return clean
    raw = col(df, 'FAC_INSURED', 'INSURED_ORI', 'INSURED', 'INSURED NAME', 'INSURED_NAME')
    if raw is None:
        raise ValueError(f'{source}: tidak ditemukan kolom insured (clean maupun raw): {list(df.columns)}')
    return [raw]


def variants(row, columns, source):
    """Ekstraksi variasi nama insured per baris.
    Mendukung pemecahan co-insured/bank QQ yang dipisahkan garis miring (/) atau QQ.
    """
    for column in columns:
        raw = txt(row[column])
        if not raw:
            continue
        # Pecah bila ada pemisah nama gabungan (/ atau QQ)
        parts = re.split(r'\s*(?:/|QQ|C/O)\s*', raw, flags=re.IGNORECASE)
        # Yield versi utuh jika ternormalisasi
        raw_norm = norm(raw)
        if raw_norm:
            yield str(column), raw, raw_norm
        # Jika terpecah lebih dari 1 entitas, yield masing-masing entitas individu
        if len(parts) > 1:
            for part in parts:
                part = part.strip()
                part_norm = norm(part)
                if part_norm and part_norm != raw_norm:
                    yield f"{column}_PART", part, part_norm


def locate_file(candidates, label):
    for path in candidates:
        if path.exists():
            return path
    raise FileNotFoundError(f'{label} tidak ditemukan di lokasi kandidat: {", ".join(str(p) for p in candidates)}')


def read_required(path, label):
    if not path.exists():
        raise FileNotFoundError(f'{label} tidak ditemukan: {path}')
    print(f"[{label}] Membaca {path.name} ...")
    return pd.read_excel(path, dtype=object)


def read_ri_slip(path):
    """Deteksi baris header RI Slip pada baris 1-4 Excel."""
    print(f"[RI SLIP] Membaca {path.name} ...")
    for header in (1, 0, 2, 3):
        try:
            frame = pd.read_excel(path, header=header, dtype=object)
            if insured_clean_columns(frame) or col(frame, 'FAC_INSURED', 'INSURED_ORI',
                                                    'INSURED', 'INSURED NAME', 'INSURED_NAME') is not None:
                if header != 0:
                    print(f"  -> RI SLIP terdeteksi pada header baris ke-{header + 1}")
                return frame
        except Exception:
            continue
    raise ValueError(f'RI SLIP: tidak ditemukan kolom insured di header baris 1-4: {path}')


def references(source_df, source):
    """Kompilasi bank nama referensi unik dari tiap tabel sumber Tripa."""
    namecols = name_columns(source_df, source)
    fc = col(source_df, 'CCOS_REF_CODE', 'FAC_CODE', 'FAC CODE', 'FACODE', 'FAC. CODE')
    currency = col(source_df, 'CCOS_CURR', 'FAC_CURRENCY', 'CURR', 'CURR ORI', 'CURRENCY')
    records = {}
    for _, row in source_df.iterrows():
        for column, raw, key in variants(row, namecols, source):
            record = records.setdefault(key, {
                'source': source,
                'name': raw,
                'key': key,
                'codes': set(),
                'currencies': set(),
                'count': 0,
                'columns': set()
            })
            record['count'] += 1
            record['columns'].add(column)
            if fc and txt(row[fc]):
                record['codes'].add(txt(row[fc]))
            if currency and txt(row[currency]):
                record['currencies'].add(txt(row[currency]).upper())
    return list(records.values())


def chosen_rows(result):
    """Pilih baris-baris hasil matching yang membutuhkan audit (Unmatching, Insured only, >1 fac code, beda)."""
    scenario_col = col(result, 'SKENARIO')
    flag_col = col(result, 'FLAG_PROD')

    if scenario_col is None and flag_col is None:
        print("  [INFO] File input adalah data suspend murni (tidak ada SKENARIO/FLAG_PROD). Seluruh baris akan diaudit.")
        return result.copy(), None, None

    scenarios = result[scenario_col].fillna('').astype(str) if scenario_col else pd.Series('', index=result.index)
    flags = result[flag_col].fillna('').astype(str) if flag_col else pd.Series('', index=result.index)

    # Filter: Unmatching, Insured only, >1 fac code, Beda Periode, Beda Currency
    mask = (
        scenarios.str.contains('insured|unmatch|>1 fac', case=False, regex=True) |
        flags.str.contains('unmatch|>1 fac|beda', case=False, regex=True)
    )
    selected = result.loc[mask].copy()
    print(f"  -> Terpilih {len(selected):,} baris dari {len(result):,} baris untuk audit (Unmatching / Insured only / >1 fac / Beda).")
    return selected, scenario_col, flag_col


def main():
    t_start = time.perf_counter()
    print("=" * 65)
    print("  AUDIT INSURED TF-IDF MANDIRI -- TRIPA MARINE (3 SUMBER)")
    print("=" * 65)

    match_file = locate_file(MATCH_PATHS, 'File Matching / Suspend Tripa')
    osbal_file = locate_file(OSBAL_PATHS, 'File OSBAL Tripa')
    facul_file = locate_file(FACUL_PATHS, 'File FACUL Tripa')
    ri_file = locate_file(RI_SLIP_PATHS, 'File Database RI Slip')

    result = read_required(match_file, 'Matching Tripa')
    data = [
        ('OSBAL', read_required(osbal_file, 'OSBAL Tripa')),
        ('FACUL', read_required(facul_file, 'FACUL Tripa')),
        ('RI SLIP', read_ri_slip(ri_file)),
    ]

    selected, scenario_col, flag_col = chosen_rows(result)
    qcols = name_columns(result, 'SUSPEND')
    ori_col = col(result, 'INSURED_ORI', 'FAC_INSURED', 'INSURED')

    refs = []
    stats = []
    for source, frame in data:
        group = references(frame, source)
        stats.append({'METRIK': f'Referensi unik {source}', 'NILAI': len(group)})
        refs.extend(group)

    if not refs:
        raise ValueError('Ketiga sumber referensi Tripa tidak memiliki nama insured valid')

    print(f"\n[TF-IDF] Membangun indeks n-gram TF-IDF untuk {len(refs):,} nama referensi unik ...")
    vectorizer = TfidfVectorizer(analyzer='char_wb', ngram_range=(2, 5), min_df=1, sublinear_tf=True)
    matrix = vectorizer.fit_transform([r['key'] for r in refs])
    neighbors = NearestNeighbors(metric='cosine', algorithm='brute', n_neighbors=min(TOP_N, len(refs)))
    neighbors.fit(matrix)

    queries = {}
    details = []
    fac_result = col(result, 'CCOS_REF_CODE', 'FAC_CODE', 'FAC CODE')
    curr_result = col(result, 'CURR ORI', 'CURR')
    receipt = col(result, 'RECEIPT NO', 'RECEIPT_NO')
    amount_col = col(result, 'AMOUNT ORI', 'AMOUNT_ORI')
    polis_col = col(result, 'POLIS_ORI', 'POLIS', 'POLICY_CLEAN_1')

    for idx, row in selected.iterrows():
        row_variants = list(variants(row, qcols, 'SUSPEND'))
        if not row_variants and ori_col:
            row_variants = list(variants(row, [ori_col], 'SUSPEND'))

        for column, raw, key in row_variants:
            q = queries.setdefault(key, {
                'name': raw,
                'key': key,
                'receipts': set(),
                'scenarios': set(),
                'flags': set(),
                'columns': set(),
                'currencies': set(),
            })
            q['receipts'].add(idx)
            q['columns'].add(column)
            if scenario_col and txt(row[scenario_col]):
                q['scenarios'].add(txt(row[scenario_col]))
            if flag_col and txt(row[flag_col]):
                q['flags'].add(txt(row[flag_col]))
            if curr_result and txt(row[curr_result]):
                q['currencies'].add(txt(row[curr_result]).upper())

            details.append({
                'EXCEL_ROW': idx + 2,
                'RECEIPT_NO': txt(row[receipt]) if receipt else '',
                'AMOUNT_ORI': txt(row[amount_col]) if amount_col else '',
                'POLIS_ORI': txt(row[polis_col]) if polis_col else '',
                'INSURED_VARIANT': raw,
                'INSURED_COLUMN': column,
                'INSURED_KEY': key,
                'INSURED_ORI': txt(row[ori_col]) if ori_col else '',
                'SKENARIO': txt(row[scenario_col]) if scenario_col else '',
                'FLAG_PROD': txt(row[flag_col]) if flag_col else '',
                'FAC_CODE_HASIL': txt(row[fac_result]) if fac_result else '',
                'CURR_ORI': txt(row[curr_result]) if curr_result else '',
            })

    print(f"[TF-IDF] Mencari kandidat untuk {len(queries):,} nama Insured Suspend unik ...")
    candidates = []
    if queries:
        all_queries = list(queries.values())
        query_vectors = vectorizer.transform([q['key'] for q in all_queries])
        distances, indices = neighbors.kneighbors(query_vectors)

        for q, ds, ids in zip(all_queries, distances, indices):
            qt = set(q['key'].split())
            q_curr = q['currencies']

            for rank, (distance, ri) in enumerate(zip(ds, ids), 1):
                ref = refs[int(ri)]
                common = qt & set(ref['key'].split())
                score = round(max(0.0, 1.0 - float(distance)) * 100, 2)
                if score < MIN_SCORE:
                    continue

                rt = set(ref['key'].split())
                risks = []

                if (qt ^ rt) & {'SYARIAH', 'SHARIA'}:
                    risks.append('SYARIAH_DIFFERENCE')
                if (qt ^ rt) & {'GROUP', 'GRUP'}:
                    risks.append('GROUP_DIFFERENCE')
                if (qt ^ rt) & {'PALM', 'SAWIT'}:
                    risks.append('BUSINESS_WORD_DIFFERENCE')
                if len(common - COMMON_ALL) <= 1:
                    risks.append('WEAK_SHARED_TOKENS')
                if len(ref['codes']) > 1:
                    risks.append('REFERENCE_MULTI_FAC_CODE')
                if q_curr and ref['currencies'] and not (q_curr & ref['currencies']):
                    risks.append('CURRENCY_MISMATCH')

                candidates.append({
                    'INSURED_SUSPEND': q['name'],
                    'NORMALIZED_SUSPEND': q['key'],
                    'CLEAN_COLUMNS': ', '.join(sorted(q['columns'])),
                    'JUMLAH_RECEIPT': len(q['receipts']),
                    'SKENARIO_ASAL': ' | '.join(sorted(x for x in q['scenarios'] if x)),
                    'FLAG_ASAL': ' | '.join(sorted(x for x in q['flags'] if x)),
                    'RANK': rank,
                    'TFIDF_SCORE': score,
                    'KANDIDAT_INSURED': ref['name'],
                    'KANDIDAT_NORMALIZED': ref['key'],
                    'SUMBER': ref['source'],
                    'SHARED_TOKENS': ', '.join(sorted(common)),
                    'RISK_FLAGS': ' | '.join(risks),
                    'FAC_CODE_REFERENSI': ', '.join(sorted(ref['codes'])),
                    'JUMLAH_FAC_CODE_REFERENSI': len(ref['codes']),
                    'CURRENCY_REFERENSI': ', '.join(sorted(ref['currencies'])),
                    'CURRENCY_SUSPEND': ', '.join(sorted(q_curr)),
                    'JUMLAH_REFERENSI': ref['count'],
                    'KOLOM_REFERENSI': ', '.join(sorted(ref['columns'])),
                    'STATUS_VALIDASI': 'PENDING',
                    'CATATAN_REVIEW': '',
                    'KEPUTUSAN_FAC_CODE': '',
                })

    summary = [
        {'METRIK': 'Cedant / Perusahaan', 'NILAI': 'TRIPA (PT Asuransi Tripakarta)'},
        {'METRIK': 'File Matching Asal', 'NILAI': match_file.name},
        {'METRIK': 'Total baris data matching', 'NILAI': len(result)},
        {'METRIK': 'Baris terpilih untuk audit', 'NILAI': len(selected)},
        {'METRIK': 'Nama Insured clean unik yang diaudit', 'NILAI': len(queries)},
        *stats,
        {'METRIK': 'Total nama referensi unik seluruh sumber (OSBAL, FACUL, RI SLIP)', 'NILAI': len(refs)},
        {'METRIK': 'Jumlah rekomendasi kandidat TF-IDF', 'NILAI': len(candidates)},
        {'METRIK': 'Top N kandidat per nama Insured', 'NILAI': TOP_N},
    ]

    guide = [
        {'PANDUAN': '1. Sumber referensi Tripa: OSBAL, FACUL, dan RI SLIP (BORDERO tidak digunakan pada Tripa).'},
        {'PANDUAN': '2. Skrip ini adalah Decision-Support Tool (AUDIT) dan tidak mengubah keputusan matching secara otomatis.'},
        {'PANDUAN': '3. TFIDF_SCORE mengukur kemiripan sub-kata (char n-gram 2-5). Nilai mendekati 100 menunjukkan kemiripan sangat tinggi.'},
        {'PANDUAN': '4. Perhatikan kolom RISK_FLAGS (mis. SYARIAH_DIFFERENCE, REFERENCE_MULTI_FAC_CODE, WEAK_SHARED_TOKENS, CURRENCY_MISMATCH).'},
        {'PANDUAN': '5. Jika kandidat disetujui, isikan kolom STATUS_VALIDASI ("APPROVED") dan masukkan FAC Code yang dipilih pada KEPUTUSAN_FAC_CODE.'},
        {'PANDUAN': '6. Sheet RECEIPT_DETAILS memetakan kembali nomor baris Excel asal dan RECEIPT NO untuk memudahkan penyesuaian.'},
    ]

    print(f"\n[Excel] Menulis laporan audit ke {DEST.name} ...")
    with pd.ExcelWriter(DEST, engine='openpyxl') as writer:
        for sheet, rows in [
            ('AUDIT_SUMMARY', summary),
            ('TFIDF_CANDIDATES', candidates),
            ('RECEIPT_DETAILS', details),
            ('CARA_PAKAI', guide),
        ]:
            frame = pd.DataFrame(rows)
            if frame.empty:
                frame = pd.DataFrame({'INFO': ['Tidak ada data yang memenuhi filter']})
            frame.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions

            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )

            # Styling Header: Dark Blue (#1F4E78)
            for cell in ws[1]:
                cell.font = Font(name='Segoe UI', size=11, color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='1F4E78')
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

            # Auto-fit width kolom
            for column in ws.columns:
                max_len = 0
                for cell in list(column)[:200]:
                    val = str(cell.value or '')
                    max_len = max(max_len, len(val))
                col_letter = column[0].column_letter
                ws.column_dimensions[col_letter].width = min(54, max(14, max_len + 3))

    elapsed = time.perf_counter() - t_start
    print(f"\n{'=' * 65}")
    print(f"  SELESAI DALAM {elapsed:.1f} DETIK")
    print(f"  Output file : {DEST}")
    print(f"  Baris diaudit: {len(selected):,}  |  Insured unik: {len(queries):,}  |  Kandidat: {len(candidates):,}")
    print("=" * 65)


if __name__ == '__main__':
    main()
